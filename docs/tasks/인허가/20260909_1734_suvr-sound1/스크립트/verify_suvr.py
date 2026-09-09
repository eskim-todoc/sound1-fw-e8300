# -*- coding: utf-8 -*-
"""SUVC / SUVR 자체 검증.

위조 방지 관점의 검증이 핵심이다.
  - 코드 분석 유닛의 측정 내용이 실제 findings 와 같은가
  - 실기 유닛의 결과 열이 비어 있는가 (지어낸 통과가 없는가)
  - 작성자와 일자가 비어 있는가
"""
import io, os, sys, json

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import load_workbook
from suvc_rows import build
import suvr_findings as SF

ROOT = u'E:/workspace/projects/sound1-fw-e8300-backup'
OUT = ROOT + u'/docs/참고/인허가/사운드1'
P_C = OUT + u'/[04] SUVC-FW-SND1 소프트웨어 유닛 검증 기준서.xlsx'
P_R = OUT + u'/[05] SUVR-FW-SND1 소프트웨어 유닛 검증 보고서.xlsx'


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    rows = build()
    by_su = {r['su_id']: r for r in rows}

    for tag, path in ((u'SUVC', P_C), (u'SUVR', P_R)):
        if not os.path.exists(path):
            print(u'%s 없음: %s' % (tag, path))
            return
        wb = load_workbook(path, data_only=True)
        print(u'%s 시트 %s' % (tag, wb.sheetnames))

    wb = load_workbook(P_R, data_only=True)
    ws = wb[u'보고서_유닛 시험']

    # 본표 머리 행 찾기
    hdr = None
    for r in range(1, 60):
        if ws.cell(r, 2).value == u'CORE':
            hdr = r
            break
    print(u'\n[검증_1] 본표 머리 행: %s' % hdr)
    assert hdr, u'머리 행을 찾지 못했다'
    cols = [ws.cell(hdr, c).value for c in range(2, 14)]
    print(u'  열: %s' % cols)
    ok = cols[10] and u'측정' in cols[10] and cols[11] == u'결과'
    print(u'  결과: %s' % (u'통과' if ok else u'실패'))

    # 행 수집
    data = []
    r = hdr + 1
    while ws.cell(r, 4).value:
        data.append({
            'test_id': ws.cell(r, 4).value,
            'method': ws.cell(r, 6).value,
            'su_id': ws.cell(r, 7).value,
            'proc': ws.cell(r, 10).value or u'',
            'crit': ws.cell(r, 11).value or u'',
            'meas': ws.cell(r, 12).value or u'',
            'result': ws.cell(r, 13).value or u'',
        })
        r += 1
    print(u'\n[검증_2] 행 수 = 유닛 수')
    print(u'  본표 %d행 / 유닛 %d건 : %s'
          % (len(data), len(rows), u'통과' if len(data) == len(rows) else u'실패'))

    print(u'\n[검증_3] 실기(D) 유닛의 결과 열이 비어 있는가 (지어낸 통과 없음)')
    bad = [d['su_id'] for d in data if d['method'] == 'D' and d['result'].strip()]
    print(u'  D 유닛 %d건 / 결과가 채워진 것: %s'
          % (len([d for d in data if d['method'] == 'D']), bad or u'없음'))
    print(u'  결과: %s' % (u'통과' if not bad else u'실패'))

    print(u'\n[검증_4] 코드 분석(A) 유닛의 측정 내용이 실제 확인 결과와 같은가')
    mism = []
    empty = []
    for d in data:
        if d['method'] != 'A':
            continue
        row = by_su.get(d['su_id'])
        if not row:
            mism.append((d['su_id'], u'유닛 없음'))
            continue
        f = SF.F.get(u'%s-%03d' % (row['core'], row['orig_id']))
        if not f:
            mism.append((d['su_id'], u'findings 없음'))
            continue
        if d['meas'].strip() != u'\n'.join(f['meas']).strip():
            mism.append((d['su_id'], u'내용 불일치'))
        if not d['meas'].strip():
            empty.append(d['su_id'])
    print(u'  A 유닛 %d건 / 불일치 %s / 빈 측정 %s'
          % (len([d for d in data if d['method'] == 'A']),
             mism[:5] or u'없음', empty[:5] or u'없음'))
    print(u'  결과: %s' % (u'통과' if not mism and not empty else u'실패'))

    print(u'\n[검증_5] 시험 절차가 방법과 맞는가')
    bad = []
    for d in data:
        if d['method'] == 'A' and d['proc'].strip() != u'코드 분석':
            bad.append((d['su_id'], d['proc'][:20]))
        if d['method'] == 'D' and d['proc'].strip() == u'코드 분석':
            bad.append((d['su_id'], u'D 인데 코드 분석'))
    print(u'  어긋남: %s' % (bad[:5] or u'없음'))

    print(u'\n[검증_6] 작성자 · 검토자 · 승인자 · 시험 일자가 비어 있는가')
    cov = wb[u'표지']
    filled = []
    for r in (12, 14, 16):
        v = cov.cell(r, 3).value
        if v and v not in (u'토닥/인공와우사업부',):
            filled.append((r, v))
        if cov.cell(r, 4).value or cov.cell(r, 5).value:
            filled.append((r, u'성명/서명/일자 채워짐'))
    for r in range(1, 40):
        if ws.cell(r, 3).value == u'시험 일자' and ws.cell(r, 4).value:
            filled.append((r, u'시험 일자 채워짐'))
        if ws.cell(r, 3).value == u'검사자' and ws.cell(r, 4).value:
            filled.append((r, u'검사자 채워짐'))
    print(u'  채워진 항목: %s' % (filled or u'없음 (의도대로)'))

    print(u'\n[검증_7] 폐기 부품 언급')
    print(u'  Sound1 이 버린 부품이 «설명·통과 기준» 에 남으면 문서 잘못이다.')
    print(u'  «측정/검토 내용» 에 나오는 것은 코드에 남은 잔재를 검증이 기록한 것이므로 정당하다.')
    PARTS = (u'nRF', u'NRF', u'가속도', u'이어피스', u'텔레코일')
    bad = 0
    rec = 0
    for d in data:
        for b in PARTS:
            if b in (d['crit'] or u''):
                bad += 1
                print(u'  X %s 통과 기준에 «%s»' % (d['su_id'], b))
            if b in (d['meas'] or u''):
                rec += 1
    for d in data:
        for b in PARTS:
            if b in (d['test_id'] or u''):
                bad += 1
                print(u'  X %s 설명에 «%s»' % (d['su_id'], b))
    print(u'  기준·설명에 남은 것: %d건 -> %s' % (bad, u'통과' if bad == 0 else u'실패'))
    print(u'  측정 내용의 코드 잔재 기록: %d건 (의도된 기록)' % rec)

    print(u'\n[검증_8] SUVC 와 SUVR 의 기준서 부분이 같은가')
    wc = load_workbook(P_C, data_only=True)
    wsc = wc[u'기준서_유닛 검증']
    hc = None
    for r in range(1, 60):
        if wsc.cell(r, 2).value == u'CORE':
            hc = r
            break
    diff = 0
    for i in range(len(data)):
        for c in range(2, 12):          # B~K
            a = wsc.cell(hc + 1 + i, c).value
            b = ws.cell(hdr + 1 + i, c).value
            if (a or u'') != (b or u''):
                diff += 1
                if diff <= 3:
                    print(u'  X 행%d 열%d' % (i + 1, c))
    print(u'  차이 %d건 : %s' % (diff, u'통과' if diff == 0 else u'실패'))


if __name__ == '__main__':
    main()
