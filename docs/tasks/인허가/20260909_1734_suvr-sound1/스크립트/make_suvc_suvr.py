# -*- coding: utf-8 -*-
"""SUVC(유닛 검증 기준서)와 SUVR(유닛 검증 보고서) xlsx 생성.

1.5세대 SUVC-FW-TD20B / SUVR-FW-TD20B 의 구조를 그대로 따른다.
  SUVC = 표지 + 기준서 시트 (B~K 열)
  SUVR = 표지 + 보고서 시트 (B~K 열 + L 측정/검토 내용 + M 결과)

측정/검토 내용은 suvr_findings 가 갖고 있다. 실제 코드를 읽고 확인한 것만 들어 있다.
실기(D) 유닛은 결과 열을 비운다. 은수님이 실기로 채우신다.
"""
import io, os, sys, json

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from suvc_rows import build, short_name
import suvr_findings as SF

ROOT = u'E:/workspace/projects/sound1-fw-e8300-backup'
OUT = ROOT + u'/docs/참고/인허가/사운드1'
HERE = os.path.dirname(os.path.abspath(__file__))

DOC_SUVC = u'SUVC-FW-SND1'
DOC_SUVR = u'SUVR-FW-SND1'
SYSTEM = u'SND1-A-EXT'
SW_NAME = u'OTE_SND1'
SW_VER = u'1.0.0'
HW = u'인공달팽이관 외부장치 Sound1 시제품'

THIN = Side(style='thin', color='999999')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill('solid', fgColor='D9D9D9')
TOP_LEFT = Alignment(horizontal='left', vertical='top', wrap_text=True)
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

# 열 정의: (제목, 폭)
COLS_C = [(u'CORE', 7), (u'Module', 10), (u'TEST ID', 26), (u'설명', 22),
          (u'방법', 5), (u'상세 설계 ID', 16), (u'시험 준비 및 초기 조건', 42),
          (u'가정 및 제한 조건', 38), (u'시험 절차', 44), (u'통과 기준', 62)]
COLS_R = COLS_C + [(u'측정/ 검토 내용', 46), (u'결과', 7)]

# 시험 준비 / 가정 제한 / 절차 / 통과 기준은 expected_values 가 유닛별로 갖는다.


def load_expected():
    p = os.path.join(HERE, 'expected_values.json')
    if not os.path.exists(p):
        return {}
    return {v['su']: v for v in json.load(io.open(p, encoding='utf-8'))}


EXP = load_expected()


def su_key(r):
    """expected_values 는 재번호된 SU-ID 뒤 세 자리를 쓴다."""
    return u'%s-%03d' % (r['core'], r['id'])


def fkey(r):
    return u'%s-%03d' % (r['core'], r['orig_id'])


def criteria(r):
    f = SF.F.get(fkey(r))
    if f:
        return u'\n'.join(f['crit'])
    return u'- %s 구현' % short_name(r['name'])


def procedure(r):
    """시험 절차. 실기 유닛은 유닛별 절차를 그대로 싣는다."""
    if r['method'] == 'A':
        return u'코드 분석'
    e = EXP.get(su_key(r))
    if e and e.get('proc'):
        return e['proc']
    return (u'- 해당 유닛을 실행시키고 디버깅용 중단점을 적용하여 출력 데이터를 확인한다.\n'
            u'- 라이브러리 출력에 의존하므로 실측값을 기록한다.')


def preparation(r):
    """시험 준비 및 초기 조건."""
    if r['method'] == 'A':
        return u'NA'
    e = EXP.get(su_key(r))
    return e['prep'] if e and e.get('prep') else u'NA'


def limitation(r):
    """가정 및 제한 조건."""
    if r['method'] == 'A':
        return u'NA'
    e = EXP.get(su_key(r))
    return e['limit'] if e and e.get('limit') else u'NA'


def pass_criteria(r):
    """통과 기준.

    실기 유닛은 «입력 -> 예상 출력» 을 하나도 빼지 않고 싣는다.
    이 문서만 보고 시험할 수 있어야 하므로 다른 문서를 참조하게 하지 않는다.
    """
    base = criteria(r)
    if r['method'] != 'D':
        return base
    e = EXP.get(su_key(r))
    if not e:
        return base
    lines = [base, u'', u'[산출 근거] ' + e['note']]
    if e['kind'] == u'실측':
        lines.append(u'')
        lines.append(u'[판정] 기댓값을 미리 낼 수 없다. 아래 항목의 실측값을 기록한다.')
        lines.append(u'       고정소수점 연산 차이를 고려하여 판정한다.')
    else:
        lines.append(u'')
        lines.append(u'[판정] 아래 입력에 대해 예상 출력과 일치해야 한다. (고정소수점 연산 고려)')
    lines.append(u'')
    for i, c in enumerate(e['cases'], 1):
        lines.append(u'%d) 입력 : %s' % (i, c['in']))
        lines.append(u'   예상 출력 : %s' % c['out'])
    return u'\n'.join(lines)


def measurement(r):
    f = SF.F.get(fkey(r))
    if r['method'] == 'A':
        return u'\n'.join(f['meas']) if f else u''
    e = EXP.get(su_key(r))
    if e and e['kind'] in (u'재현', u'상수'):
        return u'(실기 시험 후 기록)'
    return u'(실기 시험 후 기록 - 라이브러리 출력 의존)'


def result(r):
    return u'통과' if r['method'] == 'A' else u''


def cover(ws, doc_no, title):
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 22
    for c in 'CDE':
        ws.column_dimensions[c].width = 18
    ws['B2'] = u'문서번호'
    ws['C2'] = doc_no
    ws['B3'] = u'버전'
    ws['C3'] = u'Rev. 1'
    ws['B5'] = u'%s\n(%s)' % (title, SYSTEM)
    ws['B5'].font = Font(size=16, bold=True)
    ws['B5'].alignment = CENTER
    ws.merge_cells('B5:E8')
    ws.row_dimensions[5].height = 60

    ws['B11'] = u'소속/부서'
    ws['C11'] = u'성명'
    ws['D11'] = u'서명'
    ws['E11'] = u'일자'
    for i, role in enumerate((u'작성', u'검토', u'승인')):
        r = 12 + i * 2
        ws.cell(r, 2, role)
        ws.cell(r, 3, u'토닥/인공와우사업부')
    for r in (11, 12, 14, 16):
        for c in range(2, 6):
            ws.cell(r, c).border = BOX
    for c in range(2, 6):
        ws.cell(11, c).fill = HDR_FILL
        ws.cell(11, c).alignment = CENTER

    ws['B20'] = u'개정이력'
    ws['B20'].font = Font(bold=True)
    ws['B21'] = u'개정번호'
    ws['C21'] = u'개정날짜'
    ws['D21'] = u'개정내용'
    ws['B22'] = 0
    ws['D22'] = u'최초 작성'
    for c in range(2, 6):
        ws.cell(21, c).fill = HDR_FILL
        ws.cell(21, c).alignment = CENTER
    for r in (21, 22):
        for c in range(2, 6):
            ws.cell(r, c).border = BOX


def header_block(ws, title, with_env):
    """1 시험 대상 소프트웨어 ~ 4 시험 환경. 다음에 쓸 행 번호를 반환한다."""
    ws['B2'] = u'%s\n인공달팽이관 외부장치 펌웨어' % title
    ws['B2'].font = Font(size=13, bold=True)
    ws['B2'].alignment = CENTER
    ws.merge_cells('B2:F3')
    ws.row_dimensions[2].height = 34

    ws['B5'] = u'1. 시험 대상 소프트웨어'
    ws['B5'].font = Font(bold=True)
    ws['C7'] = u'소프트웨어 식별 이름 :'
    ws['D7'] = SW_NAME
    ws['C8'] = u'소프트웨어 버전 :'
    ws['D8'] = SW_VER

    ws['B10'] = u'2. 동작 하드웨어'
    ws['B10'].font = Font(bold=True)
    ws['C11'] = HW

    ws['B13'] = u'3. 필요 장비 및 지그'
    ws['B13'].font = Font(bold=True)
    ws['C15'] = u'장비 및 지그, 소프트웨어 툴'
    ws['D15'] = u'사용목적'
    for c in (3, 4):
        ws.cell(15, c).fill = HDR_FILL
        ws.cell(15, c).alignment = CENTER
        ws.cell(15, c).border = BOX
    tools = [
        (u'EDK Ezairo', u'Ezairo system 코딩, 컴파일, 링크, 다운로드 및 디버깅 프로그램 툴'),
        (u'I2C 디버거', u'Ezairo system 디버깅 장치'),
        (u'J-Link 디버거', u'Cortex-M3 디버깅 장치'),
        (u'스피커', u'소리(음파)를 생성한다.'),
        (u'음악 재생기(PC)', u'음악 파일을 재생한다.'),
        (u'음악 파일', u'시험에 필요한 특정 주파수 및 음악 신호 파일'),
        (u'유닛 시험 하네스', u'시험 벡터를 주입하고 결과를 메모리에 적재한다. '
                        u'평시 빌드에서 제외되며 기존 소스를 수정하지 않는다.'),
    ]
    r = 16
    for name, use in tools:
        ws.cell(r, 3, name).border = BOX
        ws.cell(r, 4, use).border = BOX
        ws.cell(r, 4).alignment = TOP_LEFT
        r += 1

    r += 1
    if with_env:
        ws.cell(r, 2, u'4. 시험 환경').font = Font(bold=True)
        r += 1
        for label in (u'시험 장소', u'시험 일자', u'검사자'):
            ws.cell(r, 3, label)
            ws.cell(r, 3).border = BOX
            ws.cell(r, 4).border = BOX     # 비워 둔다
            r += 1
        r += 1
        ws.cell(r, 2, u'5. 시험 내용').font = Font(bold=True)
    else:
        ws.cell(r, 2, u'4. 시험 내용').font = Font(bold=True)
    return r + 1


def sheet_body(ws, rows, cols, with_result, start_row):
    for i, (title, w) in enumerate(cols):
        col = get_column_letter(2 + i)
        ws.column_dimensions[col].width = w
        c = ws.cell(start_row, 2 + i, title)
        c.fill = HDR_FILL
        c.alignment = CENTER
        c.border = BOX
        c.font = Font(bold=True, size=9)

    r = start_row + 1
    prev_core = prev_mod = None
    for row in rows:
        vals = [
            row['core'] if row['core'] != prev_core else u'',
            row['mod_label'] if (row['core'], row['mod']) != prev_mod else u'',
            row['test_id'],
            short_name(row['name']),
            row['method'],
            row['su_id'],
            preparation(row),
            limitation(row),
            procedure(row),
            pass_criteria(row),
        ]
        if with_result:
            vals.append(measurement(row))
            vals.append(result(row))
        for i, v in enumerate(vals):
            c = ws.cell(r, 2 + i, v)
            c.border = BOX
            c.font = Font(size=9)
            c.alignment = CENTER if i in (0, 1, 4, 11) else TOP_LEFT
        # 줄 수에 맞춰 행 높이를 준다. 자동 맞춤이 없으므로 직접 계산한다.
        nl = max(str(v).count('\n') + 1 for v in vals)
        ws.row_dimensions[r].height = min(12.5 * nl, 640)
        prev_core = row['core']
        prev_mod = (row['core'], row['mod'])
        r += 1
    return r


def save(wb, path):
    tmp = path + u'.tmp'
    wb.save(tmp)
    try:
        if os.path.exists(path):
            os.remove(path)
        os.rename(tmp, path)
        return path
    except OSError:
        print(u'[경고] 원본이 열려 있어 교체하지 못했다. 임시 파일로 두었다:')
        print(u'       %s' % tmp)
        return tmp


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    rows = build()

    # --- SUVC
    wb = Workbook()
    cover(wb.active, DOC_SUVC, u'소프트웨어 유닛 검증 기준서')
    wb.active.title = u'표지'
    ws = wb.create_sheet(u'기준서_유닛 검증')
    start = header_block(ws, u'유닛 시험 기준서', with_env=False)
    end = sheet_body(ws, rows, COLS_C, False, start)
    p1 = save(wb, OUT + u'/[04] %s 소프트웨어 유닛 검증 기준서.xlsx' % DOC_SUVC)

    # --- SUVR
    wb = Workbook()
    cover(wb.active, DOC_SUVR, u'소프트웨어 유닛 검증 보고서')
    wb.active.title = u'표지'
    ws = wb.create_sheet(u'보고서_유닛 시험')
    start = header_block(ws, u'유닛 시험 보고서', with_env=True)
    end2 = sheet_body(ws, rows, COLS_R, True, start)
    p2 = save(wb, OUT + u'/[05] %s 소프트웨어 유닛 검증 보고서.xlsx' % DOC_SUVR)

    na = len([r for r in rows if r['method'] == 'A'])
    nd = len([r for r in rows if r['method'] == 'D'])
    print(u'SUVC: %s' % p1)
    print(u'SUVR: %s' % p2)
    print(u'유닛 %d건 (A %d / D %d) · 본표 %d행' % (len(rows), na, nd, end - start))
    print(u'결과 채움: 통과 %d / 비움(실기 대상) %d' % (na, nd))


if __name__ == '__main__':
    main()
