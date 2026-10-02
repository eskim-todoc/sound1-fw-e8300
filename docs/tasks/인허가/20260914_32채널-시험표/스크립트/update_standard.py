"""Preserve the existing Word layout, replace ID-6/7 numeric models and ADC setup."""
from pathlib import Path
from copy import deepcopy
import json,shutil
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import simulate as sim

root=Path.cwd(); tmp=root/'tmp/sound1-32ch'
dest=root/'docs/참고/인허가/사운드1/[Sound1] EPTC-SND1 필수성능시험 기준서_rev0.docx'
backup=tmp/'standard-before.docx'
if not backup.exists():
    candidate=Document(dest)
    if len(candidate.paragraphs)!=312 or len(candidate.tables)!=28:
        raise RuntimeError('Original 312-paragraph standard required; do not patch an already edited document.')
    shutil.copy2(dest,backup)
d=Document(backup);p=d.paragraphs;t=d.tables
m=json.loads((tmp/'model.json').read_text(encoding='utf-8'))
def setp(par,text):
    style=deepcopy(par.runs[-1]._r.rPr) if par.runs and par.runs[-1]._r.rPr is not None else None
    par.clear();run=par.add_run(text)
    if style is not None:run._r.insert(0,style)
def para(n,text):setp(p[n-1],text)
def cell(c,text):
    setp(c.paragraphs[0],str(text))
    for extra in c.paragraphs[1:]:extra._element.getparent().remove(extra._element)
    for run in c.paragraphs[0].runs:run.font.size=Pt(8)
def fill_table(tab,rows):
    sample=deepcopy(tab.rows[-1]._tr)
    while len(tab.rows)<len(rows):tab._tbl.append(deepcopy(sample))
    while len(tab.rows)>len(rows):tab._tbl.remove(tab.rows[-1]._tr)
    for row,values in zip(tab.rows,rows):
        for c,v in zip(row.cells,values):cell(c,v)
        trPr=row._tr.get_or_add_trPr()
        for el in list(trPr):
            if el.tag==qn('w:trHeight'):trPr.remove(el)
        if trPr.find(qn('w:cantSplit')) is None:trPr.append(OxmlElement('w:cantSplit'))
    hp=tab.rows[0]._tr.get_or_add_trPr()
    if hp.find(qn('w:tblHeader')) is None:hp.append(OxmlElement('w:tblHeader'))
def bounds(v):return f'{v[0]:,}~{v[2]:,}'
def case(f,lo,vol):return next(c for c in m['cases'] if c['f']==f and c['lo']==lo and c['vol']==vol)

para(206,'• 프로그램 1(32-band CIS), 자극 볼륨 4, 펄스폭 50 us를 사용한다. 기본 맵의 T/C는 500/1,500 uA이며, 16번 전극만 T=C=1,500 uA다. 표의 볼륨 1/10은 오디오 볼륨이다.')
para(209,'5.6.3. 시험 진행 및 예상 출력')
rows=[['#','조건','볼륨','입력 음압\n[dBSPL]','Sound1 ADC\n중간 / 범위','전류 중간\n[uA]','PCM\n범위','계산 전류 범위\n[uA]']]
labels={34:'기존 저음압',40:'추가 저음압',69:'일상',74:'75 dB',98:'추가 고음압',119:'사양 밖 비교'}
for lo in [34,40,69,74,98,119]:
    for vol in [1,10]:
        c=case(3815,lo,vol);v=c['currents'][25];a=c['amps'][25]
        rows.append([len(rows),labels[lo],vol,f"{c['lo']}~{c['hi']}",f"{c['adc'][1]:,} / {bounds(c['adc'])}",v[1],bounds(a),bounds(v)])
# Preserve the maximum-spec 109~110 dBSPL condition, separately from the new broad range.
for vol in [1,10]:
    import numpy as np
    allrep=np.concatenate([sim.simulate(3815,x,vol)[0] for x in [109,109.5,110]])
    center=sim.simulate(3815,110,vol)[0]
    rv=np.array([allrep[:,25].min(),int(np.median(center[:,25])),allrep[:,25].max()])
    av,iv=sim.current(rv,vol,26)
    rows.append([len(rows),'사양 상한',vol,'109~110',f'31,907 / 28,437~31,907',int(iv[1]),bounds(av),bounds(iv)])
fill_table(t[17],rows)
para(211,'1. 3,815 Hz 정현파를 재생하고 표의 오디오 볼륨과 음압을 설정한다. 예: 69~71 dBSPL은 Sound1 ADC 284~358, 40~42 dBSPL은 ADC 10~12에 대응한다.')
para(212,'2. 교정된 음압계로 입력을 맞추고, 2초 이상 안정화 후 1초 이상 음압·ADC·전류를 함께 기록한다. ADC 표는 진폭 환산값이며 16샘플 순간 최대값의 모든 무선 응답이 범위 안에 있어야 한다는 뜻은 아니다.')
para(214,'4. 26번 전극의 전류를 계산 중간값·범위와 비교하고 편차를 기록한다. 나머지 31개 전극도 함께 확인하려면 「Sound1 ID-6 ID-7 32채널 예상 전류표」에서 같은 주파수·음압·볼륨을 선택한다.')
para(215,'결과 비교 및 판정 적용')
para(216,'• 표의 전류는 Sound1 코드 기반 예상값이다. 입력 범위와 시간에 따른 계산 변동만 포함하며 실기·계측 오차 허용차는 아니다. 전류 합격 판정은 별도로 승인된 허용차를 적용하며, 승인 전에는 전류 적합/부적합을 확정하지 않는다.')
para(217,'• 최대 입력 사양은 110 dBSPL이다. 119~120은 비교용 외삽으로 합격 판정에서 제외한다. 98~110은 추가 구간이며 기존 상한 109~110 조건도 유지한다. 상한 행의 중간 전류는 110 dBSPL 값이다.')
cell(t[18].cell(0,0),'산출 근거  입력 ADC가 달라지면 AGC와 비선형 매핑 입력도 달라지므로 1.5세대 전류 범위를 그대로 승계하지 않는다. Sound1 완제품 ADC → 16샘플 AGC → 512점 Hann FFT/32-band → x^0.15 매핑 → PCM Amplitude로 계산한다. 전류 표시는 코드 환산식 497+floor(19333×Amplitude/4096) [uA]이며, 물리 DAC 실측치가 아니다. HEAR FFT와 DSP log/power 내부는 이상 함수로 근사했다.')
para(223,'• 프로그램 1(32-band CIS), 오디오 볼륨 1, 자극 볼륨 4, 펄스폭 50 us를 사용한다. 기본 맵 T/C는 500/1,500 uA이고 16번 전극만 T=C=1,500 uA다. 볼륨 10 비교는 같은 입력 조건으로 별도 기록한다.')
para(225,'5.7.3. 시험 진행 및 예상 출력')
id7=[]
rows=[['#','입력 주파수\n[Hz] (±0.5)','입력 음압\n[dBSPL]','Sound1 ADC\n중간 / 범위','주·인접 전극 예상 [uA]\n중간 / 계산 범위','기타 전극 예상 [uA]']]
for f in [100,114,128,2869,3038,3217]:
    c=case(f,74 if f==3217 else 69,1);id7.append(c)
    active=[(i+1,v) for i,v in enumerate(c['currents']) if i!=15 and v[2]>497]
    text='\n'.join(f'ch{e}: {v[1]:,} / {bounds(v)}' for e,v in active)
    rows.append([len(rows),f,f"{c['lo']}~{c['hi']}",f"{c['adc'][1]:,} / {bounds(c['adc'])}",text,'ch16: 1,492 고정\n그 외: 497'])
fill_table(t[19],rows)
para(227,'1. 아래 여섯 조건을 순서대로 재생한다. 각 조건은 음압계로 설정하고 Sound1 ADC 환산값을 함께 기록한다.')
for c in id7:
    new=p[227].insert_paragraph_before(f"• {c['f']:,} Hz 정현파: {c['lo']}~{c['hi']} dBSPL, Sound1 ADC {c['adc'][0]:,}~{c['adc'][2]:,} (중간 음압 {c['mid']:g} dBSPL에서 {c['adc'][1]:,}).")
    new.style=p[226].style
para(228,'2. 2초 이상 안정화 후 1초 이상 기록한다. 저역은 전극 1~3, 고역은 전극 23~25를 우선 측정하고 나머지 전극까지 확인한다.')
para(229,'3. 부하저항 양단의 양·음 위상 중심 전압을 측정하고 실측 부하저항으로 나누어 전류를 산출한다. 기본 맵의 16번 전극은 고정 전류 예외로 분리한다.')
para(230,'4. 표의 전극별 예상값과 비교한다. 볼륨 10 또는 추가 음압 조건은 「Sound1 ID-6 ID-7 32채널 예상 전류표」의 동일 선택 조건을 사용한다.')
para(231,'결과 비교 및 판정 적용')
para(232,'• 주파수·음압 조건을 확인하고 32전극별 전류 편차를 기록한다. 전류 계산 범위는 합격 허용차가 아니며, 승인된 실기·계측 허용차가 확정되기 전에는 전류 적합/부적합을 확정하지 않는다.')
para(233,'• 16번 전극의 높은 고정 전류를 오검출하지 않는다. 기존 1.5세대의 비대상 전극 450~550 uA 판정은 이 기본 맵의 16번 전극에 적용하지 않는다.')
cell(t[20].cell(0,0),'계산 조건  1.5세대의 입력 주파수 100/114/128/2,869/3,038/3,217 Hz를 유지하고, Sound1 passbin/Hann 창에서 각 대역에 분산되는 에너지를 계산했다. 100 Hz의 최근접 bin 3은 제외되지만 bin 4 누설 성분은 전극 1에 반영된다. 중간값은 중간 음압의 시간 중앙값이며 범위는 1 dB 격자와 중간 음압에서 관측한 계산 최소/최대다. 위상 0, 잡음·음향 주파수 응답 미포함이므로 연속 조건의 보증 한계가 아니다.')
# Keep the appendix's generation comparison consistent with the requested frequencies.
cell(t[27].rows[6].cells[3],'100 / 114 / 128 Hz 유지')
cell(t[27].rows[6].cells[4],'음압 기준 유지, ADC·전류를 Sound1으로 재산출')
cell(t[27].rows[7].cells[3],'2,869 / 3,038 / 3,217 Hz 유지')
cell(t[27].rows[7].cells[4],'Sound1 대역 간 에너지 분산 반영')
cell(t[27].rows[10].cells[3],'입력·ADC·전류 예상값 내장')
cell(t[27].rows[10].cells[4],'ID-6·7 전류 허용차는 승인 후 판정에 적용')
d.save(dest)
assert len(d.inline_shapes)==len(Document(backup).inline_shapes)
print(json.dumps({'output':str(dest),'paragraphs':len(d.paragraphs),'tables':len(d.tables),'images':len(d.inline_shapes),'id6Rows':len(t[17].rows)},ensure_ascii=False))
