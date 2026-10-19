import json, hashlib, re
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
TASK = HERE.parents[2]
PLAN = json.loads((TASK / 'builders/qcc_tests/payload.json').read_text(encoding='utf-8'))
BOOKS = {b['key']: b for b in PLAN['books']}

def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def values(w):
    return {s.title: {c.coordinate: str(c.value) if c.value is not None else ''
        for row in s.iter_rows() for c in row if c.value is not None} for s in w}

def load(key):
    b=BOOKS[key]
    return openpyxl.load_workbook(b['target'],rich_text=True,data_only=False)

def main():
    HERE.mkdir(parents=True, exist_ok=True)
    output={'books': {}, 'pairs': []}
    opened={}
    patterns=re.compile(r'원문|제공.{0,8}자료|병합|SW-FW-TEST-U-|SW-(?:FW|BT)-(?:DS|RI)-ID|실제.*경로.*미제공|예상값과 일치함|\[QCC 단독\]|\[E8300/QCC 각각 확인\]')
    for k,b in BOOKS.items():
        w=load(k); opened[k]=w
        original=openpyxl.load_workbook(b['original'],rich_text=True,data_only=False)
        data=values(w); source=values(original)
        allcells=[(s,a,v) for s,cells in data.items() for a,v in cells.items()]
        flags=[{'sheet':s,'cell':a,'value':v} for s,a,v in allcells if patterns.search(v)]
        s=w[b['sheet']]
        output['books'][k]={'path':b['target'],'hash':sha(b['target']), 'originalHash':sha(b['original']),
            'all_cells':data,'original_cells':source,'flags':flags,
            'sheets': [{'name':t.title,'rows':t.max_row,'cols':t.max_column,'merges':[str(m) for m in t.merged_cells.ranges]} for t in w],
            'height_exceeds': {r:d.height for r,d in s.row_dimensions.items() if d.height and d.height>409.5},
            'groups':b['groups']}
        original.close()
    for ck,rk in [('SUVC','SUVR'),('SITC','SITR'),('SSTC','SSTR')]:
        c=opened[ck][BOOKS[ck]['sheet']]; r=opened[rk][BOOKS[rk]['sheet']]
        start=BOOKS[ck]['commonStart'];end=max(BOOKS[ck]['newRows']); last=openpyxl.utils.column_index_from_string(BOOKS[ck]['commonLast'])
        diff=[{'cell':c.cell(rr,cc).coordinate,'criteria':str(c.cell(rr,cc).value or ''),'report':str(r.cell(rr,cc).value or '')}
            for rr in range(start,end+1) for cc in range(2,last+1) if str(c.cell(rr,cc).value or '') != str(r.cell(rr,cc).value or '')]
        cm={str(m) for m in c.merged_cells.ranges if m.max_col<=last and m.min_col>=2 and m.min_row>=start and m.max_row<=end}
        rm={str(m) for m in r.merged_cells.ranges if m.max_col<=last and m.min_col>=2 and m.min_row>=start and m.max_row<=end}
        output['pairs'].append({'pair':[ck,rk],'range':f'B{start}:{BOOKS[ck]["commonLast"]}{end}', 'differences':diff, 'merge_difference':sorted(cm^rm)})
    (HERE/'audit-current.json').write_text(json.dumps(output,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'pairs':output['pairs'],'books':{k:{'flags':len(v['flags']),'cells':sum(map(len,v['all_cells'].values())),'height_exceeds':v['height_exceeds'],'hash':v['hash']} for k,v in output['books'].items()}},ensure_ascii=True))
    for w in opened.values(): w.close()

if __name__=='__main__':main()
