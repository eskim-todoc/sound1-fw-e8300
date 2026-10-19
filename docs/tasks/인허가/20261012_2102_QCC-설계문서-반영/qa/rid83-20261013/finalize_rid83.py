"""Reconcile the current manifest and task only after verified G publication."""
from pathlib import Path
from zipfile import ZipFile
from collections import Counter
from lxml import etree
import hashlib, json, datetime, re, sys
import pypdfium2 as pdfium
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
QA=Path(__file__).resolve().parent
TASK=QA.parent.parent
ROOT=TASK.parents[3]
def read(path): return json.loads(Path(path).read_text(encoding='utf-8-sig'))
def save(path,data): Path(path).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest().upper()
stamp=datetime.datetime.now().astimezone().isoformat()
published=read(QA/'published.json')
assert len(published['files'])==12 and sum(x['changed'] for x in published['files'])==5
assert len(published['protected_originals'])==23 and all(x['unchanged'] for x in published['protected_originals'])
for item in published['files']: assert digest(item['path'])==item['sha256']
paths={x['code']:Path(x['path']) for x in published['files']}
post={'checked_at':stamp,'files_hash_match':12,'protected_originals_unchanged':23,'checks':[]}
srs=openpyxl.load_workbook(paths['SRS'],read_only=True,data_only=False)
sheet=srs['소프트웨어 입력 요구사항 사양서']
assert sheet['H93'].value in (None,'')
assert sheet['G93'].data_type=='f'
post['checks'].append('SRS H93 blank; G93 requirement formula retained')
srs.close()
books={c:openpyxl.load_workbook(paths[c],read_only=True,data_only=False) for c in ('SSTC','SSTR')}
sc=books['SSTC']['기준서_시스템 시험']; sr=books['SSTR']['보고서_시스템 시험']
for col in range(2,11):
    assert sc.cell(285,col).value==sr.cell(285,col).value, f'Common test83 cell mismatch {col}'
assert sc['F285'].value in (None,'') and sr['F285'].value in (None,'')
assert '83' in str(sr['E285'].value) and str(sr['D285'].value).strip()=='I'
assert str(sr['L285'].value).strip()=='판정 대기'
assert all(t in sr['K285'].value for t in ['형성평가용 초안','경고 및 안전사항','안전하게 사용하는 작동법','안전 예방조치','문제 해결 방법'])
post['checks'].append('TEST83 upper blank; internal RID83 and inspection I retained; criteria/report B:J match; actual draft observations and pending final verdict retained')
for code,book in books.items():
    for tab in book:
        if tab.title.endswith('_비교'): continue
        for row in tab:
            for cell in row:
                text=str(cell.value or '')
                assert 'SYS-RID-ACDOC' not in text
                assert not re.search(r'\[[-+]\]',text)
    book.close()
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
for code in ('SR','SAD'):
    with ZipFile(paths[code]) as z:
        tree=etree.fromstring(z.read('word/document.xml'))
        active=''.join(t.text or '' for t in tree.iter('{'+W+'}t') if not any(a.tag=='{'+W+'}del' for a in t.iterancestors()))
        assert 'SYS-RID-ACDOC' not in active and 'SW-FW-RID-83' in active
        assert tree.findall('.//{'+W+'}ins') or tree.findall('.//{'+W+'}del')
post['checks'].append('Both saved Word current bodies omit ACDOC links and retain RID83 with tracked history')
save(QA/'post-publish-checks.json',post)

positions=read(QA/'current-positions.json')
pos={x['code']:x for x in positions}
assert set(pos)==set(paths)
excel_native={x['code']:x for x in read(QA/'native-checks.json')}
word_native={x['code']:x for x in read(QA/'word/native-checks.json')}
package={x['code']:x for x in read(QA/'package-checks.json')}
manifest=read(QA/'source-manifest.json')
manifest['status']='RID83 direct-parent correction and draft manual inspection recorded; document checks complete; final distribution-manual verdict remains pending'
manifest['checked_at']=stamp
manifest['scope']='Current 12 E8300-QCC DHF documents; this follow-up changes SRS/SR/SAD/SSTC/SSTR and preserves the other 7 combined documents and all 23 protected originals'
manifest['current_location_count']=sum(len(x['entries']) for x in positions)
manifest['product_execution']='Actual inspection of the identified formative-evaluation manual draft, also confirmed by the user; no new hardware/software execution, final-distribution-manual pass, or approval/signature'
for entry in manifest['files']:
    current=next(x for x in published['files'] if x['code']==entry['code'])
    code=entry['code']
    entry['changed_from_current_baseline']=current['changed']
    if not current['changed']: continue
    entry['previous_followup_sha256']=entry['sha256']
    entry['sha256']=current['sha256']
    entry['markers']=dict(Counter(x['prefix'] for x in pos[code]['entries']))
    native=excel_native.get(code) or word_native[code]
    assert native['sha256'].upper()==entry['sha256']
    previews=native.get('previews') or [{'pdf':native['pdf'],'sheet':None,'range':None,'scope':'Actual Word read-only in-memory accepted full-document preview; delivered revision history retained'}]
    entry['native_previews']=[]
    for preview in previews:
        pdf=Path(preview['pdf'])
        doc=pdfium.PdfDocument(pdf)
        pages=len(doc)
        doc.close()
        entry['native_previews'].append({**preview,'pdf_sha256':digest(pdf),'pages':pages,
                                        'source_sha256':entry['sha256'],'checked_at':stamp,
                                        'candidate_matches_published_bytes':True})
    if code in package:
        entry['changes_from_previous_followup']={'scope':'RID83 parent correction, manual inspection steps/results, comparison-only appended history',
            'changed_parts':package[code]['changedParts'],'unchanged_cells_and_parts_preserved':True,
            'previous_comparison_rich_text_preserved':True,'native_candidate_matches_published_bytes':True}
    else:
        entry['changes_from_previous_followup']={'scope':'RID83 parent-reference wording only',
            'tracked_revision_author':'김은수','previous_history_preserved':True,
            'native_accept_and_QA_projection_reject_checked':True,
            'native_direct_selective_reject_limitation':'Word Revisions collection grouped earlier same-author nested changes; new-change rejection tested on an explicitly identified QA projection, not claimed directly on the full-history candidate'}
old_evidence=manifest['evidence']
manifest['evidence']={
    'previous_followup_manifest':str(QA/'source-manifest.json'),
    'current_requirements_audit':str(QA/'requirements-audit.json'),
    'manual_document_inspection':str(QA/'qcc-design-manual-audit.json'),
    'latest_user_confirmation':str(QA/'user-confirmation.json'),
    'current_package_checks':str(QA/'package-checks.json'),
    'independent_test_review':str(QA/'tests-final-review.json'),
    'native_excel_checks':str(QA/'native-checks.json'),
    'native_word_acceptance':str(QA/'word/native-checks.json'),
    'new_word_revision_QA_projection':str(QA/'word/native-projection-checks.json'),
    'visual_review':str(QA/'visual-review.json'),
    'policy_review':str(QA/'policy-review.json'),
    'current_positions':str(QA/'current-positions.json'),
    'published_hashes_and_protected_files':str(QA/'published.json'),
    'saved_G_checks':str(QA/'post-publish-checks.json')}
manifest['RID83_followup']={
    'cause':'Actual matrix ACDOC17/18/20 accompanied-document requirements were incorrectly promoted to direct SW parents because their wording resembled the manual requirement; QCC source parent trace is N/A',
    'current_parent':'blank','own_requirement':'SW-FW-RID-83','system_test':83,'method':'I',
    'observed_manual':'260824_Sound1 사용 설명서 초안_사용적합성 형성평가용; IFU-TD24A-01 Rev.0 (2026.08.12)',
    'user_confirmed_all_three_contents':True,'actual_distribution_manual_is_separate':True,
    'result_scope':'Draft inspection only','final_verdict':'판정 대기',
    'affected_codes':['SRS','SR','SAD','SSTC','SSTR'],
    'historical_ACDOC_mapping_claims_superseded':True}
save(TASK/'evidence/final-output-manifest.json',manifest)

task_file=TASK/'작업.md'
text=task_file.read_text(encoding='utf-8-sig')
start='<!-- TDC-TASK-STATUS:BEGIN v1 -->'; end='<!-- TDC-TASK-STATUS:END -->'
block='''<!-- TDC-TASK-STATUS:BEGIN v1 -->
| 항목 | 내용 |
|---|---|
| 상태 | 완료(RID83 상위 추적 정정 및 초안 검사 결과 반영) |
| 현재 결과 | SRS/SR/SAD/SSTC/SSTR 5종 저장·교차 검증 완료. RID83 상위는 빈칸, 자체 요구와 시스템83 유지. 초안 세 항목 확인 결과 기록 |
| 남은 일 | 이번 문서 정정 없음. 제품 검증의 별도 배포본 확인과 시스템83 최종 판정은 대기 |
| 다음 행동 | 별도 사용자 배포본이 제공되면 현재 기준서의 세 항목을 다시 검사하고 해당 버전·절과 최종 판정을 기록 |
| 다음 담당 | 별도 배포본 확인: 담당자 / 이후 문서 반영: Codex |
| 확인 시각 | '''+stamp+''' |
<!-- TDC-TASK-STATUS:END -->'''
text=text[:text.index(start)]+block+text[text.index(end)+len(end):]
heading='## 2026-10-13 후속 정정: RID83 자체 요구와 사용설명서 초안 검사'
assert heading not in text
text+='''

'''+heading+'''

앞선 기록의 “RID83을 실제 상위 ACDOC17/18/20에 연결”했다는 판단을 정정한다. 세 ID는 실제 매트릭스에 있으나 동봉문서(사용설명서)의 라벨·심볼·기재사항 분류에 속한다. QCC 원문의 사용자 설명서 요구 발생원은 N/A이며 직접 SW 상위 관계를 확인하지 못했다. 문구 유사성을 근거로 직접 상위 추적을 만든 것이 오류였다. 이전 수행 기록과 비교/변경 이력은 과거 이력으로 보존하며 현재 상태의 근거로 재사용하지 않는다.

SRS H93, 시스템 기준서/보고서 F285를 비웠다. 자체 SW-FW-RID-83과 시스템 시험83, 검사(I)를 유지하고 SR/SAD 및 사이버보안 연계의 해당 ACDOC 참조도 제거했다. 기준서와 보고서의 같은 행에는 경고·안전사항, 안전한 작동법, 안전 예방조치·사용 방법·문제 해결을 확인하고 문서 버전과 해당 절을 기록하는 동일 절차를 적용했다. 본문에는 적용된 내용만 두고 원래 비교 이력 뒤의 `_비교` 시트에 이번 정정을 추가했다.

IFU-TD24A-01 Rev.0(2026.08.12)의 260824 형성평가용 초안에서 세 내용을 직접 검사했다. 은수님도 모두 기재되어 있음을 직접 확인했다고 답했으며, 실제 배포본은 별도이므로 현재는 초안 검사 결과만 기록하라고 확인했다. 보고서 K285에는 실제 초안의 관찰 내용과 표시 쪽 번호를 기록하고 “미실행”은 제거했다. 최종 배포본의 내용 확인은 남아 있어 L285의 “판정 대기”를 유지했다. 초안 검사 결과를 별도 배포본의 통과로 확대하지 않았다.

Excel 정적 134개 독립 검사와 시험 변경 화면 8개, SRS 본문/비교 2개 화면을 확인했다. 전체 주요 시트를 실제 Excel로 읽기 전용 출력하고 본문에 편집 비교 기호나 잘못된 ACDOC 참조가 없는지 검사했다. Word는 실제 수락 결과와 SAD의 249개 셀을 대조했다. 전체 68쪽을 렌더링해 이전 검증본과 같은 64쪽은 픽셀 일치를 확인하고 달라진 SR13쪽/SAD51~53쪽은 직접 검토했다. 기존 이력을 가진 원본 후보에서 Word COM의 Revisions가 같은 작성자의 중첩 변경을 과거 변경과 묶어 노출하여 새 변경만 고르는 검사는 통과하지 못했다. 따라서 실제 후보의 새 변경만 XML로 거부하면 이전 문단이 정확히 복원되는 검사와, 과거 이력만 수락한 명시적 QA 사본에서 새 변경을 실제 Word로 수락/거부하는 검사를 분리했다. 직접 후보의 선택 거부가 통과했다고 주장하지 않는다. 납품 Word의 기존 이력과 작성자 김은수는 그대로 보존했고 전역 Word 사용자 설정을 변경하지 않았다.

검증된 후보와 같은 SHA256으로 G 현재본 5종을 저장했다. 나머지 현재본 7종과 보호 원본 23개는 불변이다. 저장본의 현재 본문으로 중앙 QCC 위치 목록을 재수집하고 현재 manifest를 갱신했다. SV 작성 지침 2번에는 직접 상위가 없는 자체 요구의 발생원 칸을 비워 두고 비슷한 문구만으로 다른 분류의 ID를 끌어오지 않는 검사 기준을 추가했다.

- [요구사항 조사](qa/rid83-20261013/requirements-audit.json), [사용자 최종 확인](qa/rid83-20261013/user-confirmation.json), [초안 본문 검사](qa/rid83-20261013/qcc-design-manual-audit.json)
- [독립 시험 문서 검토](qa/rid83-20261013/tests-final-review.json), [Word 새 변경 수락/거부 QA](qa/rid83-20261013/word/native-projection-checks.json), [표시 검사](qa/rid83-20261013/visual-review.json)
- [29개 원칙 검사](qa/rid83-20261013/policy-review.json), [저장 해시/원본 보존](qa/rid83-20261013/published.json), [최종 G 내용 대조](qa/rid83-20261013/post-publish-checks.json), [현재 산출물 목록](evidence/final-output-manifest.json)
'''
task_file.write_text(text,encoding='utf-8')
print(json.dumps({'manifest_updated':True,'task_updated':True,'current_locations':manifest['current_location_count'],'post_publish_checks':len(post['checks'])},ensure_ascii=False))
