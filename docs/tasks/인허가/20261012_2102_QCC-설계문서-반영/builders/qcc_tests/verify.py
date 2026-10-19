import json,re,hashlib,zipfile
from pathlib import Path
from lxml import etree as ET
import openpyxl
HERE=Path(__file__).resolve().parent;TASK=HERE.parent.parent
PLAN=json.loads((HERE/'payload.json').read_text(encoding='utf-8'))
EVIDENCE=TASK/'evidence/qcc_tests_authored.json'
EV=json.loads(EVIDENCE.read_text(encoding='utf-8'))
N={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def value(v):return str(v)if v is not None else''
checks=[]
books={b['key']:b for b in PLAN['books']}
for criterion,report in [('SUVC','SUVR'),('SITC','SITR'),('SSTC','SSTR')]:
 bc=books[criterion];br=books[report]
 wc=openpyxl.load_workbook(bc['target'],data_only=False,rich_text=True)
 wr=openpyxl.load_workbook(br['target'],data_only=False,rich_text=True)
 sc=wc[bc['sheet']];sr=wr[br['sheet']]
 last=max(bc['newRows']);first=bc['commonStart'];cols=openpyxl.utils.column_index_from_string(bc['commonLast'])
 differences=[]
 for row in range(first,last+1):
  for c in range(2,cols+1):
   vc=value(sc.cell(row,c).value);vr=value(sr.cell(row,c).value)
   if vc!=vr:differences.append({'cell':sc.cell(row,c).coordinate,'criterion':vc,'report':vr})
 cm={str(r)for r in sc.merged_cells.ranges if r.min_row>=first and r.max_row<=last and r.min_col>=2 and r.max_col<=cols}
 rm={str(r)for r in sr.merged_cells.ranges if r.min_row>=first and r.max_row<=last and r.min_col>=2 and r.max_col<=cols}
 ids=[]
 idcol='D'if bc['kind']=='unit'else'C'if bc['kind']=='integration'else'B'
 for row in range(first,last+1):
  t=value(sc[f'{idcol}{row}'].value).strip()
  if re.match(r'SW-FW-(?:UnitTest|IntegrationTest|SYS-Test)',t):ids.append(t)
 oldcount=115 if bc['kind']=='unit'else 17 if bc['kind']=='integration'else 81
 newcount=15 if bc['kind']=='unit'else 7 if bc['kind']=='integration'else 2
 verdicts={str(r):value(sr[f'{br["verdictCol"]}{r}'].value).strip()for r in br['newRows']}
 results={str(r):value(sr[f'{br["resultCol"]}{r}'].value).strip()for r in br['newRows']}
 heights={str(r):[sc.row_dimensions[r].height,sr.row_dimensions[r].height]for r in bc['newRows']}
 nativeFeatures=[]
 for b in[bc,br]:
  with zipfile.ZipFile(b['target'])as z,zipfile.ZipFile(b['original'])as zo:
   feat=[n for n in zo.namelist()if n.startswith('xl/media/')or n.startswith('xl/drawings/')]
   assert all(z.read(n)==zo.read(n)for n in feat)
   nativeFeatures.append({'key':b['key'],'feature_members_preserved':len(feat),'original_hash_unchanged':sha(b['original'])==b['originalHash']})
 auxdif=[]
 if bc['kind']=='system':
  ac=wc.worksheets[-1];ar=wr.worksheets[-1]
  for row in range(1,40):
   for c in range(1,11):
    if value(ac.cell(row,c).value)!=value(ar.cell(row,c).value):auxdif.append(ac.cell(row,c).coordinate)
 assert not differences,(criterion,'common value differences',differences[:3])
 assert cm==rm,(criterion,'common merge differences',list(cm^rm)[:8])
 assert not auxdif,(criterion,'aux common differences',auxdif)
 assert len(ids)==oldcount+newcount and len(ids)==len(set(ids)),(criterion,len(ids))
 assert all('미실행'in v and'판정 대기'in v for v in verdicts.values())
 assert all('미실행'in v and'실제 결과 미제공'in v for v in results.values())
 assert all(a==b and 0<a<=409.5 for a,b in heights.values())
 if br['history']:
  hist=wr['시험 보조자료']
  for addr,expected in br['history'].items():assert value(hist[addr].value)==expected,('history current value',addr)
 checks.append({'pair':[criterion,report],'common_columns':'B:'+bc['commonLast'],'comparison_rows':[first,last],'common_value_mismatches':differences,'common_merge_equal':True,'common_merge_count':len(cm),'cybersecurity_common_mismatches':auxdif,'preserved_old_test_count':oldcount,'new_test_count':newcount,'all_test_ids':ids,'new_condition_count':len(bc['newRows']),'new_verdicts':verdicts,'row_height_criteria_report':heights,'native_features':nativeFeatures})
 wc.close();wr.close()
# Actual notices are rich runs with color, underline and +2pt from the saved cell font.
noticecheck=[]
for b in PLAN['books']:
 with zipfile.ZipFile(b['target'])as z:
  st=ET.fromstring(z.read('xl/styles.xml'));fonts=st.find('x:fonts',N);xfs=st.find('x:cellXfs',N)
  for sheetpart in ['xl/worksheets/sheet2.xml']+([f'xl/worksheets/sheet{len(openpyxl.load_workbook(b["target"],read_only=True).sheetnames)}.xml']if b['aux']else[])+(['xl/worksheets/sheet3.xml']if b['history']else[]):
   s=ET.fromstring(z.read(sheetpart))
   for c in s.findall('.//x:sheetData/x:row/x:c',N):
    isn=c.find('x:is',N)
    if isn is None:continue
    for r in isn.findall('x:r',N):
     t=r.findtext('x:t',namespaces=N)or''
     p=r.find('x:rPr',N)
     if t.startswith(('[QCC 자료 입력]','[QCC/E8300 자료 입력]','[E8300 보완]')):
      base=float(fonts[int(xfs[int(c.get('s','0'))].get('fontId','0'))].find('x:sz',N).get('val','11'))
      size=float(p.find('x:sz',N).get('val'));color=p.find('x:color',N).get('rgb')
      expected='FF7030A0'if t.startswith('[QCC/E8300')else'FF0070C0'if t.startswith('[E8300')else'FFC00000'
      assert size==base+2 and color==expected and p.find('x:u',N)is not None,(b['key'],c.get('r'),base,size,color)
      noticecheck.append({'key':b['key'],'part':sheetpart,'cell':c.get('r'),'base':base,'notice_size':size,'color':color,'underline':True})
for f in EV['files']:f['sha256_after']=sha(f['path'])
readsource=json.loads((TASK/'evidence/qcc_tests.json').read_text(encoding='utf-8'))['source_files'][0]
assert sha(readsource['path'])==readsource['sha256_start']
EV['verification']={'status':'PASS','pairs':checks,'new_rich_notice_runs_verified':len(noticecheck),'notice_styles':noticecheck,'all_original_hashes_unchanged':all(sha(b['original'])==b['originalHash']for b in PLAN['books']),'qcc_source_docx_sha256_end':sha(readsource['path']),'native_render':'Parent root owns final native Excel rendering/visual QA; Artifact staging previews viewed.','reference_dates':'Existing general time input2026-05-02/result2026-05-29 preserved; no new approval/signature copied.'}
EVIDENCE.write_text(json.dumps(EV,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'status':'PASS','pairs':[{'pair':c['pair'],'ids':len(c['all_test_ids']),'new_conditions':c['new_condition_count'],'common_merges':c['common_merge_count']}for c in checks],'notice_runs':len(noticecheck)},ensure_ascii=True))
