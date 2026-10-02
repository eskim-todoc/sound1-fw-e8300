import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def text_of_cell(cell) -> str:
    return "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())


def inspect(path: Path) -> dict:
    doc = Document(str(path))
    sections = []
    for idx, sec in enumerate(doc.sections, 1):
        sections.append(
            {
                "index": idx,
                "width_twip": sec.page_width.twips,
                "height_twip": sec.page_height.twips,
                "left_twip": sec.left_margin.twips,
                "right_twip": sec.right_margin.twips,
                "top_twip": sec.top_margin.twips,
                "bottom_twip": sec.bottom_margin.twips,
                "header_twip": sec.header_distance.twips,
                "footer_twip": sec.footer_distance.twips,
                "start_type": str(sec.start_type),
            }
        )

    paragraphs = []
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip():
            paragraphs.append(
                {
                    "index": idx,
                    "style": p.style.name if p.style else None,
                    "text": p.text.strip(),
                }
            )

    tables = []
    for tidx, table in enumerate(doc.tables, 1):
        rows = []
        for ridx, row in enumerate(table.rows, 1):
            rows.append(
                {
                    "row": ridx,
                    "cells": [text_of_cell(cell) for cell in row.cells],
                }
            )
        grid = []
        tbl_grid = table._tbl.tblGrid
        if tbl_grid is not None:
            grid = [int(col.get(qn("w:w"))) for col in tbl_grid]
        tables.append({"index": tidx, "rows": rows, "grid_twip": grid})

    styles = []
    for style in doc.styles:
        if style.type != 1:
            continue
        font = style.font
        pf = style.paragraph_format
        styles.append(
            {
                "name": style.name,
                "font": font.name,
                "size_pt": font.size.pt if font.size else None,
                "bold": font.bold,
                "italic": font.italic,
                "space_before_pt": pf.space_before.pt if pf.space_before else None,
                "space_after_pt": pf.space_after.pt if pf.space_after else None,
                "line_spacing": pf.line_spacing,
                "alignment": str(pf.alignment),
            }
        )

    headers = []
    footers = []
    for sidx, sec in enumerate(doc.sections, 1):
        headers.append(
            {
                "section": sidx,
                "text": "\n".join(p.text for p in sec.header.paragraphs if p.text.strip()),
                "tables": [
                    [[text_of_cell(c) for c in row.cells] for row in t.rows]
                    for t in sec.header.tables
                ],
            }
        )
        footers.append(
            {
                "section": sidx,
                "text": "\n".join(p.text for p in sec.footer.paragraphs if p.text.strip()),
                "tables": [
                    [[text_of_cell(c) for c in row.cells] for row in t.rows]
                    for t in sec.footer.tables
                ],
            }
        )

    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        package = {
            name: {
                "size": zf.getinfo(name).file_size,
                "sha256": hashlib.sha256(zf.read(name)).hexdigest().upper(),
            }
            for name in names
        }
        all_xml_text = {}
        for name in names:
            if name.startswith("word/") and name.endswith(".xml"):
                raw = zf.read(name)
                try:
                    from lxml import etree

                    root = etree.fromstring(raw)
                    texts = root.xpath("//*[local-name()='t']/text()")
                    if texts:
                        all_xml_text[name] = texts
                except Exception:
                    pass

    return {
        "path": str(path.resolve()),
        "size": path.stat().st_size,
        "sha256": sha256(path),
        "sections": sections,
        "paragraphs": paragraphs,
        "tables": tables,
        "styles": styles,
        "headers": headers,
        "footers": footers,
        "inline_shapes": len(doc.inline_shapes),
        "package": package,
        "xml_text": all_xml_text,
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("output")
    args = ap.parse_args()
    data = inspect(Path(args.input))
    Path(args.output).write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "paragraphs": len(data["paragraphs"]),
                "tables": len(data["tables"]),
                "sections": len(data["sections"]),
                "inline_shapes": data["inline_shapes"],
                "parts": len(data["package"]),
                "sha256": data["sha256"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
