#!/usr/bin/env python3
"""Sound1 필수성능시험 기준서의 내용·구조·산식 정합성을 검사한다."""

from __future__ import annotations

import math
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")


ROOT = Path(__file__).resolve().parents[5]
DOCX = ROOT / "docs" / "참고" / "인허가" / "사운드1" / "[Sound1] EPTC-SND1 필수성능시험 기준서_rev0.docx"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def expected_adc(spl: int, attenuation_db: float) -> int:
    return math.floor(262143 * (10 ** ((spl - 120 + attenuation_db) / 20)))


def cell_int(text: str) -> int:
    return int(text.replace(",", "").strip())


def main() -> None:
    require(DOCX.exists(), f"문서가 없습니다: {DOCX}")
    doc = Document(DOCX)
    body_text = "\n".join(p.text for p in doc.paragraphs)
    margin_text = "\n".join(
        p.text
        for section in doc.sections
        for story in (section.header, section.footer)
        for p in story.paragraphs
    )
    all_text = body_text + "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    ) + "\n" + margin_text

    # 핵심 문서 식별자와 시험 항목
    for token in ("EPTC-SND1", "SND1-A-EXT", "2026.09.10"):
        require(token in all_text, f"문서 식별 정보 누락: {token}")
    for test_id in range(1, 10):
        heading = f"5.{test_id}. 시험 ID : SYS-EP-TEST-ID-{test_id}"
        require(sum(p.text.strip() == heading for p in doc.paragraphs) == 1, f"시험 ID 제목 오류: {heading}")

    # Sound1 변경사항과 독립 수행에 필요한 핵심 내용
    required = (
        "7,968.75 Hz",
        "-8.293 dB",
        "-3.644 dB",
        "31,907",
        "17,942~31,907",
        "28,437~31,907",
        "최대 입력 사양은 110 dB SPL",
        "고정 입력",
        "명목 ADC",
        "허용오차",
        "합격 범위",
        "N-of-M",
        "상위 16",
        "512점 FFT",
        "31.25 Hz",
        "2-DMIC",
        "매핑 Live",
        "부록 1. Sound1 입력 음압 대비 ADC 출력 테이블",
        "부록 2. 32-band 주파수 배정 및 시험 톤",
        "부록 3. N-of-M 자극 방식 시험 근거",
        "부록 4. 1.5세대 대비 Sound1 변경사항",
    )
    for token in required:
        require(token in all_text, f"필수 내용 누락: {token}")

    forbidden = (
        "작성 필요", "확인 필요", "TD23D-A-EXT", "EPTC-TD23D", "7990", "`",
        "115~120", "119~120", "89,926~100,899", "29~120 dB SPL",
        "120 dB SPL의 판정 기준값", "120 dB SPL 근처로 조정",
    )
    for token in forbidden:
        require(token not in all_text, f"금지된 잔여 문자열 발견: {token}")

    # 부록 1의 29~110 dB SPL 전 범위와 세 열의 계산값 검증
    adc_tables = [t for t in doc.tables if len(t.columns) == 8 and t.cell(0, 0).text.strip() == "dB SPL" and t.cell(0, 4).text.strip() == "dB SPL"]
    require(len(adc_tables) == 1, "ADC 부록 표를 유일하게 찾지 못했습니다.")
    adc_table = adc_tables[0]
    require(len(adc_table.rows) == 42 and len(adc_table.columns) == 8, "ADC 부록 표 구조가 42x8이 아닙니다.")
    seen: dict[int, tuple[int, int, int]] = {}
    for row in adc_table.rows[1:]:
        for offset in (0, 4):
            spl = cell_int(row.cells[offset].text)
            values = tuple(cell_int(row.cells[offset + i].text) for i in (1, 2, 3))
            seen[spl] = values
    require(sorted(seen) == list(range(29, 111)), "ADC 부록의 음압 범위가 29~110 dB SPL 연속값이 아닙니다.")
    for spl, values in seen.items():
        expected = (
            expected_adc(spl, 0.0),
            expected_adc(spl, -3.644),
            expected_adc(spl, -8.293),
        )
        require(values == expected, f"ADC 산식 불일치: {spl} dB SPL, 문서={values}, 계산={expected}")

    # ID-8 프로그램 2의 24채널 명목 전류와 ±10% 범위 검증
    map2_tables = [t for t in doc.tables if len(t.columns) == 6 and t.cell(0, 0).text.strip() == "채널" and "명목 전류" in t.cell(0, 1).text]
    require(len(map2_tables) == 1, "ID-8 채널별 전류표를 유일하게 찾지 못했습니다.")
    expected_map2 = [500, 543, 586, 630, 673, 717, 760, 804, 847, 891, 934, 978, 1021, 1065, 1108, 1500, 1195, 1239, 1282, 1326, 1369, 1413, 1456, 1500]
    actual_map2: dict[int, tuple[int, str]] = {}
    for row in map2_tables[0].rows[1:]:
        for offset in (0, 3):
            channel = cell_int(row.cells[offset].text)
            nominal = cell_int(row.cells[offset + 1].text)
            accepted = row.cells[offset + 2].text.strip()
            actual_map2[channel] = (nominal, accepted)
    require(sorted(actual_map2) == list(range(1, 25)), "ID-8 채널 번호가 1~24 연속값이 아닙니다.")
    for channel, nominal in enumerate(expected_map2, 1):
        expected_range = f"{math.ceil(nominal * 0.9):,}~{math.floor(nominal * 1.1):,}"
        require(actual_map2[channel] == (nominal, expected_range), f"ID-8 전류표 불일치: 채널 {channel}")

    # 핵심 고정 시험벡터가 수치로 남아 있는지 확인
    vector_tokens = (
        "110 / -1/+0", "31,907 / 28,437~31,907",
        "93.75", "ch1: 872 (-17/+16) / 855~888",
        "ch25: 934 (-17/+16) / 917~950",
        "band 0~15: 72.92", "band 16~31: 72.92",
        "band 16만 52.96", "band 20만 52.96",
        "불일치 허용 횟수는 0회",
    )
    for token in vector_tokens:
        require(token in all_text, f"고정 시험벡터 누락: {token}")

    # A4 단일 절, 표 폭, 바닥글 필드
    require(len(doc.sections) == 1, "문서 절(section)은 하나여야 합니다.")
    section = doc.sections[0]
    require(abs(section.page_width.cm - 21.0) < 0.05, "용지 폭이 A4가 아닙니다.")
    require(abs(section.page_height.cm - 29.7) < 0.05, "용지 높이가 A4가 아닙니다.")
    content_width = section.page_width - section.left_margin - section.right_margin
    for idx, table in enumerate(doc.tables):
        grid = table._tbl.tblGrid
        if grid is None:
            continue
        widths = [int(col.get(qn("w:w"))) for col in grid.gridCol_lst if col.get(qn("w:w"))]
        if widths:
            require(sum(widths) <= int(content_width / 635) + 20, f"표 {idx}의 폭이 본문 영역을 초과합니다.")

    with zipfile.ZipFile(DOCX) as zf:
        names = set(zf.namelist())
        rel_names = [n for n in names if n.endswith(".rels")]
        rel_xml = "\n".join(zf.read(n).decode("utf-8", errors="replace") for n in rel_names)
        require('TargetMode="External"' not in rel_xml, "외부 링크 관계가 남아 있습니다.")
        document_xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
        require("INCLUDEPICTURE" not in document_xml and "LINK " not in document_xml, "외부 개체 필드가 남아 있습니다.")
        footer_xml = "\n".join(
            zf.read(n).decode("utf-8", errors="replace") for n in names if n.startswith("word/footer") and n.endswith(".xml")
        )
        for token in ("PAGE", "NUMPAGES", "F7030-02", "EPTC-SND1"):
            require(token in footer_xml, f"바닥글 요소 누락: {token}")

    print(f"검증 완료: {DOCX}")
    print(f"- 시험 ID: 1~9")
    print(f"- ADC 산식: 29~110 dB SPL, 82행")
    print("- 고정 시험벡터: ID-1~9")
    print("- ID-8 맵 전류: 24채널")
    print(f"- 표: {len(doc.tables)}개, 그림: {len(doc.inline_shapes)}개, 절: {len(doc.sections)}개")
    print("- 외부 링크: 없음")


if __name__ == "__main__":
    try:
        main()
    except AssertionError as exc:
        print(f"검증 실패: {exc}", file=sys.stderr)
        raise SystemExit(1)
