"""Guarded comparison and publication for the three document-only edits."""
from pathlib import Path
import json
import re
import shutil
import zipfile
import hashlib
import openpyxl
import pypdfium2 as pdfium
from lxml import etree as E
from patch_audit_query_guard import QA, KINDS, FILES, sha, TASK, FIELDS, PURPOSE
from inspect_sad_unit_trace import N, q, text
from preserve_audit_workbooks import unpack
from publish_audit_documents import page_caches

def verify():
    base=json.loads((QA/'baseline.json').read_text(encoding='utf8'))
    plan=json.loads((QA/'sheet-edits.json').read_text(encoding='utf8'))
    saved={}
    report={'new_unit_ids':0,'device_execution':False,'gate_results':'pending code change and verification'}
    for kind in ('suvc','suvr'):
        old=openpyxl.load_workbook(QA/(kind+'-before.xlsx'))
        new=openpyxl.load_workbook(QA/(kind+'.xlsx'))
        assert old.sheetnames==new.sheetnames
        diff=[]
        for before,after in zip(old,new):
            assert before.max_row==after.max_row and before.max_column==after.max_column
            assert set(map(str,before.merged_cells.ranges))==set(map(str,after.merged_cells.ranges))
            for row in before:
                for a in row:
                    b=after[a.coordinate]
                    assert a._style==b._style,(kind,a.coordinate,'style')
                    if a.value!=b.value:diff.append(a.coordinate)
        edits={**plan['common'],**(plan['report'] if kind=='suvr' else {})}
        assert set(diff)==set(edits),(kind,diff)
        s=new.worksheets[1]
        for a,want in edits.items():assert s[a].value==want,(kind,a)
        ids=[str(s.cell(r,4).value).strip() for r in range(1,s.max_row+1) if str(s.cell(r,4).value).strip().startswith('SW-FW-UnitTest-')]
        oldids=[str(old.worksheets[1].cell(r,4).value).strip() for r in range(1,s.max_row+1) if str(old.worksheets[1].cell(r,4).value).strip().startswith('SW-FW-UnitTest-')]
        assert len(ids)==115 and ids==oldids
        oldp,newp=unpack(QA/(kind+'-before.xlsx')),unpack(QA/(kind+'.xlsx'))
        assert set(oldp)==set(newp)
        assert [n for n in oldp if oldp[n]!=newp[n]]==['xl/worksheets/sheet2.xml']
        ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        a,b=[E.fromstring(p['xl/worksheets/sheet2.xml']) for p in (oldp,newp)]
        for addr in edits:
            node=b.find(f'.//s:c[@r="{addr}"]',ns)
            node.getparent().replace(node,a.find(f'.//s:c[@r="{addr}"]',ns))
        for row in plan['heights']:
            oa=a.find(f'.//s:row[@r="{row}"]',ns);nb=b.find(f'.//s:row[@r="{row}"]',ns)
            nb.attrib.clear();nb.attrib.update(oa.attrib)
        # Cell nodes moved above must also be restored to the original tree for
        # a normalized equality comparison.
        a=E.fromstring(oldp['xl/worksheets/sheet2.xml'])
        assert E.tostring(a)==E.tostring(b),'Unexpected worksheet XML change'
        checks=json.loads((QA/(kind+'-native-check.json')).read_text(encoding='utf-8-sig'))
        assert all(x['needed']<=x['available']+1 for x in checks)
        assert max(d.height or 0 for d in s.row_dimensions.values())<=409.5
        saved[kind]=new
        report[kind]={'changed_cells':diff,'test_ids':115,'rows':s.max_row,'native_fit_checks':len(checks)}
    a,b=saved['suvc'].worksheets[1],saved['suvr'].worksheets[1]
    assert all(a.cell(r,c).value==b.cell(r,c).value for r in range(25,a.max_row+1) for c in range(1,12))
    report['common_A_K_equal']=True
    before=openpyxl.load_workbook(QA/'suvr-before.xlsx').worksheets[1]
    now=saved['suvr'].worksheets[1]
    assert before['L811'].value==now['L810'].value
    assert before['L809'].value.strip() in now['L809'].value and before['L810'].value.strip() in now['L809'].value
    assert all(f'검증 [{i}]' in now['L811'].value for i in (4,5,6,7))
    assert '통과' not in now['L811'].value and '코드 반영 후 검증' in now['M809'].value
    assert all(f'조건 [{i}]' in now['K811'].value for i in (4,5,6,7))
    assert now['K811'].value.count('[LOG] IDENT=')==1
    assert now['L792'].value==before['L792'].value and now['L1213'].value==before['L1213'].value
    oldp,newp=unpack(QA/'sddd-before.docx'),unpack(QA/'sddd.docx')
    assert [n for n in oldp if oldp[n]!=newp[n]]==['word/document.xml']
    a,b=[E.fromstring(p['word/document.xml']) for p in (oldp,newp)]
    ab,bb=[e.find(q('body')) for e in (a,b)]
    changes=json.loads((QA/'design-edits.json').read_text(encoding='utf8'))['body_indices']
    assert len(ab)==len(bb)
    assert [i for i in range(len(ab)) if E.tostring(ab[i])!=E.tostring(bb[i])]==changes
    oldids=re.findall(r'SU-ID-(?:CFX|CM3|HEAR|BOOT)-\d+',text(a))
    newids=re.findall(r'SU-ID-(?:CFX|CM3|HEAR|BOOT)-\d+',text(b))
    assert set(oldids)==set(newids)
    assert newids.count('SU-ID-CM3-030')==oldids.count('SU-ID-CM3-030')+2
    assert newids.count('SU-ID-CM3-031')==oldids.count('SU-ID-CM3-031')+2
    assert text(b).count('[QCC 자료 입력]')==4
    assert E.fromstring(newp['word/settings.xml']).find(q('trackRevisions')) is not None
    rejected=unpack(QA/'sddd-rejected.docx')
    previous=unpack(TASK/'qa/audit-units-four-docs-20261012/sddd-rejected.docx')
    assert text(E.fromstring(rejected['word/document.xml']))==text(E.fromstring(previous['word/document.xml']))
    assert all(sha(Path(p))==h for p,h in base['sources'].items())
    assert all(sha(Path(p))==h for p,h in base['protected'].items())
    report['protected_files_unchanged']=len(base['protected'])
    report['source_files_unchanged']=len(base['sources'])
    report['sddd']={'changed_body_blocks':changes,'rejection_baseline_equal':True,'QCC_markers':4}
    (QA/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

def visual():
    old=pdfium.PdfDocument(str(TASK/'qa/audit-units-four-docs-20261012/sddd.pdf'))
    new=pdfium.PdfDocument(str(QA/'sddd.pdf'))
    def body_hash(p):
        im=p.render(scale=1.5).to_pil().convert('RGB')
        return hashlib.sha256(im.crop((0,195,im.width,im.height-100)).tobytes()).hexdigest()
    old_hashes={body_hash(p) for p in old}
    changed=[]
    for i,p in enumerate(new):
        if body_hash(p) not in old_hashes:
            changed.append(i+1)
            p.render(scale=1.5).to_pil().save(QA/f'sddd-page-{i+1}.png')
    for src in QA.glob('*-native-*.pdf'):
        pdf=pdfium.PdfDocument(str(src))
        for i,p in enumerate(pdf):p.render(scale=2).to_pil().save(QA/(src.stem+f'-{i+1}.png'))
    report={'pages':len(new),'new_body_pages':changed,'pixel_identical_body_pages':len(new)-len(changed)}
    (QA/'page-comparison.json').write_text(json.dumps(report,indent=2),encoding='utf8')
    print(report)

def publish():
    base=json.loads((QA/'baseline.json').read_text(encoding='utf8'))
    check=json.loads((QA/'verification.json').read_text(encoding='utf8'))
    visual_check=json.loads((QA/'visual-review.json').read_text(encoding='utf8'))
    pages=json.loads((QA/'page-comparison.json').read_text(encoding='utf8'))
    assert visual_check['issues']==[] and visual_check['reviewed_sddd_pages']==pages['new_body_pages']
    assert check['common_A_K_equal']
    assert all(sha(Path(p))==h for p,h in {**base['protected'],**base['sources']}.items())
    assert all(sha(FILES[k])==base[k]['sha256'] for k in KINDS),'External target edit'
    package=unpack(QA/'sddd.docx')
    caches=page_caches(package,unpack(QA/'sddd-accepted.docx'),pages['pages'])
    with zipfile.ZipFile(QA/'sddd-publish.docx','w',zipfile.ZIP_DEFLATED) as z:
        for name,data in package.items():z.writestr(name,data)
    staged={}
    for k in KINDS:
        src=QA/('sddd-publish.docx' if k=='sddd' else k+'.xlsx')
        tmp=FILES[k].with_name(FILES[k].name+'.pending')
        assert not tmp.exists()
        staged[k]=(src,tmp)
    for k,(src,tmp) in staged.items():shutil.copyfile(src,tmp)
    for k,(src,tmp) in staged.items():
        assert sha(FILES[k])==base[k]['sha256']
        tmp.replace(FILES[k])
        assert sha(FILES[k])==sha(src)
    assert all(sha(Path(p))==h for p,h in {**base['protected'],**base['sources']}.items())
    result={'outputs':{k:{'path':str(FILES[k]),'sha256':sha(FILES[k])} for k in KINDS},'cache_updates':caches,'code_unchanged':True,'IDs_unchanged':True,'new_tests_executed':False}
    (QA/'published-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=='__main__':
    import sys
    if '--visual' in sys.argv:visual()
    elif '--publish' in sys.argv:publish()
    else:verify()
