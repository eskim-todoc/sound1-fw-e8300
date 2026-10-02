"""Prepare/review SITC values; no openpyxl authoring. Excel/Artifact write data."""
from pathlib import Path
import json, hashlib, shutil, re, zipfile, difflib, sys, copy
from lxml import etree as E
import openpyxl

T=Path(__file__).resolve().parent
ROOT=T.parents[3]
F=ROOT/'docs/참고/인허가/사운드1/DHF/03_설계 출력/SW/펌웨어'
QA=T/'qa/sitc-current-design'
TARGET=F/'SITC-FW-TD24A 소프트웨어 통합시험 기준서.xlsx'
MAIN='기준서_통합 시험'; DIFF=MAIN+'_변경 비교'
DETAIL=F/'[작성자 확인] SDDD-FW-TD24A 소프트웨어 상세 설계서.docx'
SIP=F/'[작성자 확인] SIP-FW-TD24A 소프트웨어 통합 계획서.docx'
PROTECT=[DETAIL,SIP,F/'[작성자 확인] SAD-FW-TD24A 소프트웨어 아키텍쳐 기술서.docx',F/'SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(n,v):(QA/n).write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8')
def lines(*v):return '\n'.join('- '+s for s in v)
def accepted(p):
    with zipfile.ZipFile(p) as z:x=E.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    return [''.join(n.xpath('.//w:t[not(ancestor::w:del)]/text()',namespaces=ns)) for n in x.xpath('.//w:body//w:p',namespaces=ns)]

# (core, name, units by core, hardware, method, preparation, limits, procedure, expected)
CASES=[
('BOOT/CFX/CM3','코어 기동 및 초기화 연결',{'BOOT':[1,5,6],'CFX':list(range(1,7)),'CM3':[1,2,3,4]},'Sound1, 디버거','D',
 lines('정상 시험 이미지를 적재한다. CFX 실행 허가=0에서 준비 상태와 초기화 호출 순서를 기록한다.'),
 lines('초기 기동 로그는 저장한 뒤 읽는다. 중단점 때문에 생긴 지연을 정상 기동 시간으로 판정하지 않는다.'),
 lines('CFX 기동/대기→CM3 공유 자료 및 주변장치 준비→CFX 허가=1→CFX 정상 초기화를 차례로 확인한다.','마이크, HEAR 체인, PCM/I2S 설정과 첫 반복 처리의 선후 관계를 확인한다.'),
 lines('허가=0에는 정상 오디오/자극 처리가 없다. CM3 준비 뒤 허가=1이 되고 첫 처리 전에 CFX 초기화가 끝난다.','HEAR CMD 0/1, 마이크/PCM/I2S 버퍼가 각 경로에 연결된다. 적재 전에 응용 코어를 실행하지 않는다.')),
('CFX/CM3','저전력 요청, 정지 및 재시작',{'CFX':[9],'CM3':[7,8,9,12,15]},'Sound1, 터치 센서','D',
 lines('정상 오디오/PCM 출력 중인 시험 장치를 준비한다. 저전력 요청과 완료 상태, 반복 처리 횟수를 기록한다.'),
 lines('절전 진입 중 긴 RTT 출력을 하지 않는다. 복귀 경로는 정상 터치 입력으로 시험한다.'),
 lines('CM3 절전 요청→CFX 주변장치 정리/요청 해제→CM3 절전 진입을 관찰한다.','정상 터치 조건으로 재시작하고 코어 준비, 실행 허가 및 오디오 처리를 확인한다.'),
 lines('CFX의 PCM/오디오/I2S/DMIC/FIFO/HEAR 정리가 CM3 절전 진입보다 먼저다. 처리된 요청은 0이다.','대기 중 정상 자극 처리 횟수는 증가하지 않는다. 재시작은 초기화/허가를 거친 뒤 정상 입력을 처리한다.')),
('CFX','오디오 수집과 입력 경로 전환',{'CFX':[1,3,7,8,14,15,19,21]},'마이크, I2S 입력 장치','D',
 lines('마이크 블록 A[i]=i/B[i]=100+i(i=0~15), 다른 마이크는 각 값+1000을 준비한다. I2S는 ON→OFF로 시험한다.'),
 lines('버퍼 값은 시험 입력으로 고정한다. 현재 I2S 복사 코드에는 완료=1 설정문이 없다. 완료 플래그를 대신 1로 넣어 통과시키지 않는다.'),
 lines('FIFO 인터럽트와 두 번의 블록 복사 후 현재/직전 자료를 읽는다.','I2S 시작, 안정 수신, 중단에서 마이크 수/입력 경로와 완만한 전환을 확인한다.'),
 lines('두 번째 복사 후 현재=B, 직전=A이며 다른 마이크와 섞이지 않는다. 각 블록은 16표본이다.','I2S 수신 시 SINGLE+I2S, 중단 후 정상 두 마이크 경로로 돌아간다. 시작/중단 전환 뒤 이전 입력을 계속 쓰지 않는다.','상세 설계의 완료 조회는 복사 후 1/해제 후 0이다. 실제 복사/조회 경로로 확인한다.')),
('CFX/CM3','PCM 모드 및 지정 명령 전달',{'CFX':[1,38,39,40,41],'CM3':[17]},'Sound1','D',
 lines('지정 명령=[0x11,0x22,0x33], 길이=3을 준비한다. 대기, 프리앰블, 지정 명령, 실시간 및 0 채움 모드를 각각 요청한다.'),
 lines('PCM 전송 전 버퍼를 저장하여 판정한다. 접근할 수 없는 PCM 핀을 직접 측정하는 절차는 요구하지 않는다.'),
 lines('CM3 모드 요청을 CFX에서 읽고 해당 처리 함수와 생성 프레임을 확인한다.','지정 명령 송신 중/완료 뒤의 버퍼 사용 상태와 프리앰블 종료 후 다음 모드를 읽는다.'),
 lines('각 모드의 처리 경로만 실행된다. 지정 명령은 0x11→0x22→0x33 순서이고 사용 중 덮어쓰지 않는다.','프리앰블 완료 뒤 요청한 다음 모드로 전환한다. 대기는 NOP, 0 채움은 0이며 이전 자극 데이터를 재사용하지 않는다.')),
('CFX/HEAR/CM3','신호처리부터 CIS 패킷까지 연결',{'CFX':[21,22,23,24,25,26,27,28,29,30,31,32,33,36,38,39],'HEAR':[1,2],'CM3':[19,20,29]},'Sound1, 오디오 입력 장치','D',
 lines('혼합 전 16표본=[0,-1024,544,-3200,3168,나머지0]을 단일 입력에 넣는다. 별도 FFT 시험은 512개 0표본 및 1000Hz 순음을 준비한다.','패킷 확인은 CIS/32대역/전극[i]=i+1, 크기=100, 2프레임/자극, 12채널/ms로 고정한다.'),
 lines('FFT 최대 위치와 패킷은 별도 고정 입력으로 확인한다. 크기=100은 매핑 출력 시험값이며 실제 음압의 전류값으로 간주하지 않는다.'),
 lines('혼합→HEAR CMD 0/Event 0→AGC/FFT 입력 갱신→CMD 1/Event 1→대역/매핑→CIS PCM을 기록한다.','512표본 x[n]=round(65536*sin(2*pi*32*n/512))과 0 입력을 각각 적용한다. 패킷 조건을 적용해 연속 세 구간을 읽는다.'),
 lines('혼합 결과=[0,-32,17,-100,99,나머지0], Event 0 뒤 maxAudioInput=100이다. 0블록 다음 값은 0으로 갱신된다.','FFT 0 입력의 결과는 0, 순음의 최대 bin은32(1000Hz)다. 결과는 해당 완료 이벤트 뒤 사용한다.','CIS 전극 인덱스/크기는0~11/100,12~23/100,24~31 다음0~3/100이며 구간마다12개다.')),
('CFX/CM3','맵, 방향 및 볼륨 변경 반영',{'CFX':[1,3,11,12,13,21,29,30,31],'CM3':[6,20,22,23]},'Sound1','D',
 lines('시험 맵1/2에 구분할 수 있는 전극/대역 조건을 넣는다. 디지털 T=50/C=150, 자극 볼륨1→4, 오디오 볼륨1→10을 준비한다.'),
 lines('맵 변경과 볼륨 변경은 각각 확인한 뒤 동시에 요청한다. 완료 전 자료를 새 설정으로 판정하지 않는다.'),
 lines('CM3 맵 요청과 CFX 복사/완료 응답을 읽는다. 좌/우 방향과 마이크 수를 바꾼다.','볼륨 변경 시 NOP 전환→C 상한→매핑 계수 갱신을 확인하고 같은 맵을 다시 선택한다.'),
 lines('맵 요청 시 재적재 완료=0, CFX 반영 후 완료=1/요청=0이다. 좌는DMIC2, 우는DMIC1이 지연 입력이다.','자극 볼륨1/4의 C 상한은119/150, T는50이다. 오디오 볼륨 계수는 선택 값으로 갱신되며 변경 전 계수로 자극하지 않는다.','같은 맵 유지에는 불필요한 새 변경 요청을 만들지 않는다.')),
('CM3/CFX','CM3 초기화와 공유 상태 연결',{'CM3':[1,2,3,4,11,16,26],'CFX':[1]},'Sound1, 디버거','D',
 lines('기동 준비 전 공유 실행 허가=0과 오류 초기 상태를 기록한다. QCC 제어 출력과 SPI 응답 버퍼를 관찰한다.'),
 lines('GPIO는 출력 상태 스냅샷으로 확인할 수 있다. 실제 QCC 무선 동작은 이 기동 시험의 판정 대상이 아니다.'),
 lines('CM3 주변장치/터치/QCC 핀 및 SPI 버퍼 초기화 후 공유 주소와 CFX 준비 확인 순서를 기록한다.','공유 자료 설정 뒤 허가 상태를 CFX에서 조회한다.'),
 lines('QCC_CTRL/ISD_CHECK 초기 출력은Low/Low다. 초기화 중 터치를 정상 입력으로 처리하지 않는다.','공유 주소/자료 준비 후 CFX 허가가0→1로 바뀌고 양측 값이 같다. SPI 송수신 버퍼를 혼동하지 않는다.')),
('CM3/CFX','내부기 실행, 연결 검사 및 해제',{'CM3':[17,18,25,30,31,32,34,36,37,38,39,40,41],'CFX':[10,38,40,41,42]},'FPGA, 송신 PMIC, 시험 내부기','D',
 lines('격리된 시험 내부기 또는 정상/응답 없음의 시험 응답을 준비한다. 초기화, 역방향 검사 및 PCM 완료 상태를 기록한다.'),
 lines('환자에게 연결하지 않는다. 내부기 설정 완료와 RF 링크 검사 결과는 따로 읽는다.'),
 lines('송신 전원/FPGA 초기화→명령 PCM 전송→완료 통보→역방향 응답 처리→자극 설정 순서를 확인한다.','정상 응답과 응답 없음/오류를 각각 적용하고 해제/재연결을 진행한다.'),
 lines('준비가 끝나기 전에 설정 완료를 반환하지 않는다. CFX 송신 완료와 CM3 수신/설정 완료는 단계별로 구분된다.','응답이 링크 상태에 반영되고 처리한 검사 이벤트는 정리된다. 해제 후 이전 연결/완료값으로 Live를 계속하지 않는다.','송신 전원 설정은 해당 PMIC/FPGA로 전달되며 T/C 설정과 혼동하지 않는다.')),
('CM3/CFX','매핑 명령과 맵/자극 설정 연결',{'CM3':[17,20,27,28,29,33,41,48,49],'CFX':[11,12,29,36,39,40,41]},'SPI 시험 송신기, 시험 내부기','D',
 lines('매핑 연결 상태에서 지정 자극/Live 맵/종료 요청을 준비한다. Live 시험은32대역 NofM, 펄스폭14us, 전극1~32로 설정한다.'),
 lines('지정 자극과 Live 명령을 별도로 시작한다. 저장은 시험 저장소만 사용한다.'),
 lines('명령 수신→CM3 맵 요청→CFX 반영 확인→CM3 파라미터 계산/공유→CFX 적용→내부기 준비→Live 전환을 기록한다.','지정 자극 요청을 전달하고 매핑 종료를 적용한다.'),
 lines('NofM 계산 결과는자극당3프레임/1ms당8채널이다. 공유한32개T/C와 CFX 수신 배열이 같다.','Live는 맵/계산/내부기 준비 후다. 지정 자극은 지정 명령 경로만 사용한다. 종료 요청 후 매핑 상태와 출력 모드가 종료 경로로 전환된다.')),
('CM3','블루투스 명령 수신, 처리 및 응답',{'CM3':[22,23,26,45,46,47,48,49]},'QCC 또는 SPI 시험 송신기','D',
 lines('일반 명령 허용=false/true를 독립 준비한다. 허용 상태에서 맵2, 오디오 볼륨10, 자극 볼륨4 설정/조회 요청을 준비한다.'),
 lines('E8300의 SPI 수신 이후를 판정한다. 비활성 패스키 비교를 인증 기능으로 쓰지 않으며 QCC 내부 무선 시험은 별도다.'),
 lines('수신 분류→허용 판정→요청 처리→조회 응답 버퍼→명령 정리 순서를 읽는다.','허용=false의 일반 명령과 별도 연결 확인 명령을 각각 확인한다.'),
 lines('허용=true의 설정 후 조회는맵2/오디오10/자극4다. 응답은 대응 명령의 결과이고 처리 후 명령은IDLE다.','허용=false에서 일반 설정은 실행되지 않는다. 연결 확인은 별도 경로를 사용한다. SPI 사용 중 응답 버퍼를 덮어쓰지 않는다.')),
('CM3/CFX','정상 반복 제어 및 자동 종료 연결',{'CM3':[5,6,7,8,9,11,12,13,15,20,30,45],'CFX':[1,9,10,11]},'Sound1, 터치 센서','D',
 lines('다른 종료 조건 없이 미연결 카운터 호출 전179999/180000과 연결=true를 독립 준비한다. 맵/볼륨 변경 중 종료 요청도 준비한다.'),
 lines('1ms tick을 사용한다. 연결 상태는 내부기 자극 설정의 완료 상태이며 Bluetooth 연결 여부로 대신하지 않는다.'),
 lines('입력 수집→시스템 제어→맵/내부기→통신 처리의 호출 순서를 읽는다. 각 카운터 조건으로1tick 진행한다.','변경 처리와 종료 요청을 함께 넣어 CFX 정지 확인 후 CM3 대기를 확인한다.'),
 lines('미연결179999는180000이 되어3분 종료 요청이 없다. 미연결180000은180001 검사로종료 요청/enable_ISD=false가 되고 카운터가0이다.','연결=true는카운터0이다. 종료가 결정되면 새Live 처리를 계속하지 않고 CFX 정리 후 CM3가 대기한다.')),
('E8300/QCC','QCC 시스템 정보의 상태 및 응답',{'CM3':[11,13,26,45,50],'CFX':[21]},'QCC 또는 SPI 시험 송신기','T',
 lines('0x34 데이터=[80,1,1]/[40,0,2], 0x36 데이터=[1,1]/[1,2]/[0,1]을 준비한다. 0x35는 OTA_EZAIRO 상태와 일반 표시 상태를 준비한다.'),
 lines('E8300 수신 버퍼/응답과 시스템 상태를 RTT로 읽는다. QCC 내부 구현은 자료 확보 후 별도 추가한다.'),
 lines('각 0x34/0x35/0x36 요청을 순서대로 전달하고 상태, 공유 값, SPI 응답을 저장한다.','상태 변경 후 리모콘 배터리 조회와 입력 종류 판정이 같은 자료를 쓰는지 확인한다.'),
 lines('0x34 후 잔량은 80→40, 충전은 연결→해제, 뚜껑은 열림→닫힘이다. 응답 데이터는 [0x34,1]이다.','0x36의 크래들 플래그는 1/0/0, 응답은 [0x36,1]이다. 0x35 OTA 상태는 DFU 연결, 일반 상태는 해제이며 응답은 [0x35,1]이다.')),
('E8300/QCC','I2S 혼합 게인 및 스트리밍 전환',{'CFX':[3,7,8,15,16,18,19,21,22],'HEAR':[1],'CM3':[45,47,48]},'마이크, I2S 입력 장치','T',
 lines('I2S/마이크 각16표본=320, Gain A=65536(1.0)/B=131072(2.0)을 고정한다. 안정 수신 상태에서 일반/크래들/매핑 연결을 각각 적용한다.'),
 lines('수신 시작/중단 전환이 끝난 블록으로 수치를 비교한다. 시험 게인 직접 고정과 명령 설정 결과를 구분한다.'),
 lines('0x36으로 입력 종류를 바꾸고 선택한 A/B 게인과 혼합 버퍼/최대값을 읽는다.','매핑 연결 시 두 게인 1.0, 해제 후 저장 게인 복귀를 확인한다. I2S 시작/중단도 확인한다.'),
 lines('크래들은 마이크 A/I2S B로 출력 30, 일반은 마이크 A/I2S 1.0으로 20, 매핑은 각 1.0으로 20이다. maxAudioInput도 각각 30/20/20이다.','시작은 크기가 증가, 중단은 감소한 뒤 마이크 경로로 돌아간다. 이전 크래들/매핑 게인을 다음 경로에 잘못 유지하지 않는다.')),
('E8300/QCC','QCC 제어 및 내부기 상태 통보',{'CM3':[3,7,8,30,41]},'Sound1','T',
 lines('정상/종료와 내부기 설정 완료=false/true를 독립 적용한다. 두 출력의 GPIO 상태를 RTT로 읽을 준비를 한다.'),
 lines('핀 접근이 가능하면 오실로스코프로 추가 확인한다. 핀 출력 요청만으로 QCC 내부 동작을 검증했다고 판정하지 않는다.'),
 lines('초기화 후 정상/종료 요청과 내부기 완료/미완료를 각각 적용하고 QCC_CTRL/ISD_CHECK 출력을 읽는다.','초기화 중과 설정 완료 뒤의 통보 시점도 비교한다.'),
 lines('초기화는 두 출력 Low다. QCC_CTRL은 정상 High/종료 Low, ISD_CHECK는 완료 High/미완료 Low다.','내부기 설정 미완료를 완료로 통보하지 않는다. GPIO 출력 상태와 적용한 조건이 일치한다.')),
('CFX/HEAR/CM3','빔포밍, NofM 선택 및 분할 패킷',{'CFX':[11,13,17,21,22,28,32,34,35,36,37,39],'HEAR':[1,2],'CM3':[20,27,29]},'마이크, Sound1, 시험 내부기','T',
 lines('NofM/32대역/전극[i]=i+1, 3프레임/자극, 8채널/ms, phase=0, 묵음 꺼짐을 준비한다.','빔포밍은 기준 0/지연 현재[i]=64*i/지연 직전[0]=2048이다. 선택은 매핑 출력 32개 모두 100 또는 값[i]=i로 고정한다.'),
 lines('선택 시험값은 FFT/매핑 이후 경계에서 고정한다. 디지털 패킷 판정과 실제 전극 전류/음압 시험을 구분한다.'),
 lines('좌/우 지연 입력 연결과 혼합/HEAR 최대값을 읽는다. 32개 크기를 고정해 선택/교차 배치를 읽는다.','직전 마지막 인덱스 15/0을 독립 준비하고 연속 두 1ms 출력의 전극/크기/NOP를 해석한다.'),
 lines('빔포밍 출력[0~14]=1~15/[15]=32, 최대값 32다. 동률 100은 인덱스 0~15, 증가 크기는 16~31 선택이다.','동률 교차 순서=[0,8,1,9,2,10,3,11,4,12,5,13,6,14,7,15]. 직전 15는 이 순서, 직전 0은 마지막 15를 맨 앞으로 옮긴다.','각 1ms는 8자극/24프레임, 자극마다 2프레임 NOP다. 미선택 채널은 자극 패킷에 없어야 하며 내부 배열의 크기 0을 요구하지 않는다.')),
('BOOT/CFX/CM3','이미지 검사, 적재 및 코어 기동',{'BOOT':[1,3,5,6,7,8,9,10,11,12,13,14,15,16],'CFX':[1],'CM3':[4]},'격리된 시험 저장소, 디버거','D',
 lines('정상 이미지, 헤더 불일치, 구획 읽기 부족, CRC 손상 시험 사본을 준비한다. 암호화 사본은 정상 AES/복호화 후 CRC 오류를 각각 준비한다.'),
 lines('시험 저장소만 사용한다. 키/제조 고유값은 로그에 남기지 않는다. 초기 분기/잠금은 코드 분석을 함께 기록한다.'),
 lines('저장장치 준비→목록/헤더→구획 검사/복호화→코어별 전송/잠금→기동 순서와 반환값을 읽는다.','각 실패 사본은 독립 리셋해 실행한다. 서비스 진입 조건도 별도 적용한다.'),
 lines('정상 검사/적재 성공 뒤 해당 코어가 기동한다. 헤더/짧은 읽기/최종 CRC 실패는 오류이며 해당 이미지를 정상 기동하지 않는다.','CRC는 헤더 포함 잔여값 0으로 판정한다. 암호화는 복호화 후 잔여값 0/암호화 표시 일치일 때만 통과한다.','대상 주소/길이와 잠금 시점이 적재 조건과 같다. 서비스 조건은 서비스 경로만 선택한다.')),
('BOOT/CM3','갱신 파일 기록과 대체 슬롯 기동',{'BOOT':[6,10,13,19,20],'CM3':[26,45,50,51]},'SPI 시험 송신기, 시험 저장소','T',
 lines('기본 슬롯 1/대체 2, 시도 상한 3, 정상/불완전 갱신 사본을 준비한다. TRY 호출 전 횟수 2/3과 TRY_DONE 성공/실패는 독립 상태로 준비한다.'),
 lines('운영 이미지에 덮어쓰지 않는다. OTA 명령 지원 슬롯 1~2와 부트 상태 슬롯 1~4를 구분한다. 실패 1회에 즉시 기본 복귀한다고 가정하지 않는다.'),
 lines('갱신 시작→시험 파일 수신/기록→종료→기동 위치 요청을 확인한다. 저장 실패/지원 밖 슬롯도 시험한다.','재기동 후 선택 슬롯/횟수/결과와 이미지 검사를 읽는다. 각 TRY/TRY_DONE 조건을 적용한다.'),
 lines('정상 기록은 수신 파일 종류/슬롯/내용과 같고 시작/종료 후 DFU 상태는 연결/해제다. 저장 실패/지원 밖 요청은 오류 응답이다.','TRY 2는 횟수 3/슬롯 2, TRY 3은 횟수 4/슬롯 1/BOOT/TRY_DONE/FAIL이다. TRY_DONE 성공은 2, 실패는 1/UNKNOWN이다.','불완전 이미지는 정상 기동으로 판정하지 않는다. 다음 기동 선택은 저장 상태/시도 상한에 따른다.'),
)
]

def prepare():
    QA.mkdir(parents=True,exist_ok=True)
    assert len(CASES)==17
    base={'workbook_sha256':sha(TARGET),'protected':{str(p):sha(p) for p in PROTECT},'source':{}}
    # Record read-only source and test-file hashes as an explicit no-code-write guard.
    for p in (ROOT/'src').rglob('*'):
        if p.suffix.lower() in {'.c','.h','.hct'}:base['source'][str(p.relative_to(ROOT))]=sha(p)
    if (QA/'baseline.json').exists():
        saved=json.loads((QA/'baseline.json').read_text('utf-8'))
        assert saved==base,'Baseline changed; do not overwrite user changes'
    else:shutil.copy2(TARGET,QA/'before.xlsx');dump('baseline.json',base)
    model=json.loads((T/'qa/unit-reconciliation/model.json').read_text('utf-8'));by={u['su']:u for u in model}
    texts=accepted(DETAIL);units={}
    for t in texts:
        m=re.fullmatch(r'(SU-ID-(?:CFX|HEAR|CM3|BOOT)-\d{3})(?:,\s*(.+))?',t)
        if m:units[m[1]]=m[2] or by[m[1]]['name']
    assert len(units)==115,len(units)
    assert set(units)==set(by)
    dump('inputs.json',{'detail_units':units,'integration_plan':accepted(SIP)})
    w=openpyxl.load_workbook(TARGET);old=[list(row) for row in w[MAIN].iter_rows(min_row=25,max_row=41,max_col=13,values_only=True)]
    rows=[];refs=[]
    for i,case in enumerate(CASES,1):
        core,name,group,hw,method,prep,limit,procedure,expected=case
        us=[f'SU-ID-{c}-{n:03d}' for c,nums in group.items() for n in nums]
        assert len(set(us))==len(us) and all(u in by for u in us)
        sis=list(dict.fromkeys(by[u]['si'] for u in us))
        mods=list(dict.fromkeys(by[u]['module'] for u in us))
        r=[None,core,old[i-1][2],'\n'.join(sis),name,'\n'.join(us),hw,'\n'.join(mods),method,prep,limit,procedure,expected]
        assert r[2]==f'SW-FW-IntegrationTest-ID-{i:03d}'
        assert not any('·' in str(v) for v in r)
        assert all(r[c] for c in range(1,13))
        rows.append(r);refs.append({'test':r[2],'si':sis,'su':us,'sources':list({d['path'] for u in us for d in by[u]['refs']})})
    # Existing equipment block is extended without adding rows or changing headings.
    header=[
        ['장비 및 지그, 소프트웨어 툴','사용목적'],
        ['I2C 디버거(AMIS CCA)','E8300 디버깅/중단점/메모리 확인'],
        ['EDK Ezairo / SEGGER RTT','빌드/다운로드 및 CM3 상태 확인. CFX/HEAR 값은 처리 후 별도 저장하여 CM3 RTT로 읽는다. 고정 입력/응답 주입은 시험 코드에 준비한다.'],
        ['QCC 또는 SPI 시험 송신기','E8300 명령 수신/응답 연결. QCC 내부 무선 기능과 구분'],
        ['QCC 펌웨어 자료','[자료 추가 위치] QCC 내부 통합시험은 자료 확보 후 기존 마지막 시험 다음에 같은 열 구성으로 추가한다.'],
        ['오실로스코프 / 시험 내부기','UART RX/TX, Calibration 핀 및 필요 시 PMIC 출력 확인. 실제 자극은 시험 내부기만 사용'],
        ['I2S 오디오 입력 장치 / OD 이어폰','I2S 수신/혼합 확인 및 관찰 출력 청취. 청취만으로 수치 판정하지 않음'],
    ]
    diff=[]
    for a,b in zip(old,rows):
        aa=list(a);bb=list(b);aa[0]='[-]';bb[0]='[+]';diff.extend([aa,bb])
    dump('plan.json',{'rows':rows,'old':old,'diffrows':diff,'header':header,'refs':refs,'sheet_names':w.sheetnames})
    print('Prepared17 tests,115-unit design checked. No workbook/source writes.')

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda s:'{'+NS+'}'+s
def replace_zip(path,parts):
    tmp=path.with_suffix('.tmp.xlsx')
    with zipfile.ZipFile(path) as z,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as o:
        for info in z.infolist():o.writestr(info,parts.get(info.filename,z.read(info.filename)))
    tmp.replace(path)
def richdiff():
    p=json.loads((QA/'plan.json').read_text('utf-8'));path=QA/'candidate.xlsx'
    with zipfile.ZipFile(path) as z:x=E.fromstring(z.read('xl/worksheets/sheet3.xml'))
    cells={c.get('r'):c for c in x.iter(q('c'))};count=0;changes=0
    for i,(a,b) in enumerate(zip(p['old'],p['rows'])):
        for c in range(13):
            va=str(a[c] or '');vb=str(b[c] or '')
            op=difflib.SequenceMatcher(None,va,vb,autojunk=False).get_opcodes()
            for side,value,row,color in [(0,va,25+2*i,'FF9C0006'),(1,vb,26+2*i,'FF006100')]:
                address=f'{chr(65+c)}{row}';cell=cells[address]
                if c==0:value='[-]' if side==0 else '[+]';runs=[(value,color)]
                else:
                    runs=[]
                    for tag,i1,i2,j1,j2 in op:
                        sub=value[i1:i2] if side==0 else value[j1:j2]
                        if sub:runs.append((sub,'FF000000' if tag=='equal' else color))
                    changes+=sum(len(s) for s,col in runs if col!='FF000000')
                for child in list(cell):
                    if child.tag!=q('extLst'):cell.remove(child)
                cell.set('t','inlineStr');inline=E.SubElement(cell,q('is'))
                for text,col in runs or [('', 'FF000000')]:
                    r=E.SubElement(inline,q('r'));pr=E.SubElement(r,q('rPr'))
                    E.SubElement(pr,q('rFont'),val='맑은 고딕');E.SubElement(pr,q('sz'),val='11');E.SubElement(pr,q('color'),rgb=col)
                    if col!='FF000000':E.SubElement(pr,q('b'))
                    t=E.SubElement(r,q('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=text
                count+=1
    parts={'xl/worksheets/sheet3.xml':E.tostring(x,xml_declaration=True,encoding='UTF-8',standalone=True)}
    # Preserve native Excel's ID compaction. Unchanged cover content/style is
    # checked semantically below rather than restoring stale style/string IDs.
    replace_zip(path,parts)
    dump('richdiff.json',{'cells':count,'changed_characters':changes});print('Rich diff',count,changes)

def verify():
    p=json.loads((QA/'plan.json').read_text('utf-8'));b=json.loads((QA/'baseline.json').read_text('utf-8'))
    w=openpyxl.load_workbook(QA/'candidate.xlsx',data_only=False,rich_text=True);old=openpyxl.load_workbook(QA/'before.xlsx',data_only=False)
    assert w.sheetnames==old.sheetnames==p['sheet_names']
    for sheet in (MAIN,DIFF):
        for i,values in enumerate(p['header'],15):
            assert [w[sheet].cell(i,c).value for c in (3,4)]==values,(sheet,i,'equipment')
        for r in range(1,15):
            for c in range(1,14):
                assert str(w[sheet].cell(r,c).value or '')==str(old[sheet].cell(r,c).value or ''),(sheet,r,c,'unrelated header')
        for c in range(1,14):
            letter=chr(64+c)
            assert w[sheet].column_dimensions[letter].width==old[sheet].column_dimensions[letter].width,(sheet,letter,'width')
    for row in old['표지']:
        for c in row:
            after=w['표지'][c.coordinate]
            assert str(c.value or '')==str(after.value or ''),(c.coordinate,'cover value')
            for attr in ('font','fill','border','alignment','number_format','protection'):
                assert copy.copy(getattr(c,attr))==copy.copy(getattr(after,attr)),(c.coordinate,attr)
    for i,r in enumerate(p['rows'],25):
        assert [w[MAIN].cell(i,c).value for c in range(1,14)]==r,(i,'values')
        assert all(w[MAIN].cell(i,c).value for c in range(2,14))
    for i,r in enumerate(p['diffrows'],25):
        assert [str(w[DIFF].cell(i,c).value or '') for c in range(1,14)]==[str(v or '') for v in r],(i,'diff')
    with zipfile.ZipFile(QA/'candidate.xlsx') as z:
        x=E.fromstring(z.read('xl/worksheets/sheet3.xml'))
    cells={c.get('r'):c for c in x.iter(q('c'))}
    for i,(a,brow) in enumerate(zip(p['old'],p['rows'])):
        for col in range(1,13):
            va=str(a[col] or '');vb=str(brow[col] or '')
            op=difflib.SequenceMatcher(None,va,vb,autojunk=False).get_opcodes()
            for side,text,row,color in [(0,va,25+2*i,'FF9C0006'),(1,vb,26+2*i,'FF006100')]:
                expected=[]
                for tag,i1,i2,j1,j2 in op:
                    sub=text[i1:i2] if side==0 else text[j1:j2]
                    if sub:expected.append((sub,'FF000000' if tag=='equal' else color))
                actual=[(r.find(q('t')).text or '',r.find(q('rPr')).find(q('color')).get('rgb')) for r in cells[f'{chr(65+col)}{row}'].find(q('is')).findall(q('r'))]
                assert actual==(expected or [('', 'FF000000')]),(row,col,'changed characters')
    # Cover values/styles/topology and original object bytes must be preserved.
    with zipfile.ZipFile(QA/'candidate.xlsx') as z,zipfile.ZipFile(QA/'before.xlsx') as oz:
        keep=[n for n in oz.namelist() if n.startswith(('xl/media/','xl/drawings/'))]
        assert all(z.read(n)==oz.read(n) for n in keep),'Cover/objects changed'
    assert str(w['표지'].merged_cells)==str(old['표지'].merged_cells)
    assert w['표지'].page_setup==old['표지'].page_setup
    assert w['표지'].page_margins==old['표지'].page_margins
    for path,h in b['protected'].items():assert sha(Path(path))==h,path
    for path,h in b['source'].items():assert sha(ROOT/path)==h,path
    assert sha(TARGET)==b['workbook_sha256'],'User edited target; stop publication'
    # Verify illustrative arithmetic independently. Fixed inputs are not target test results.
    raw=[0,-1024,544,-3200,3168]+[0]*11
    assert [v>>5 for v in raw]==[0,-32,17,-100,99]+[0]*11
    assert max(map(abs,[v>>5 for v in raw]))==100
    bf=[((0>>5)+((64*(i+1))>>5))>>1 for i in range(15)]+[((0>>5)+(2048>>5))>>1]
    assert bf==list(range(1,16))+[32]
    assert (320>>5)+(320>>5)==20
    assert (320>>5)+((320>>5)*2)==30
    assert sum(v!='NOP' for v in [x for k in range(8) for x in (k,'NOP','NOP')])==8
    dump('checks.json',{'tests':17,'paired_diff_rows':34,'valid_unit_references':sum(len(r['su']) for r in p['refs']),'distinct_units':len(set(u for r in p['refs'] for u in r['su'])),'participating_items':len(set(si for r in p['refs'] for si in r['si'])),'cover_semantics_preserved':True,'preserved_object_parts':len(keep),'protected_documents':len(b['protected']),'source_files_unchanged':len(b['source']),'candidate_sha256':sha(QA/'candidate.xlsx'),'illustrative_arithmetic_checks':6,'target_tests_executed':False})
    print('Verified17 tests/34 diff rows; cover, inputs and all source files preserved.')

def publish():
    verify();v=json.loads((QA/'visual-check.json').read_text('utf-8'))
    assert v['candidate_sha256']==sha(QA/'candidate.xlsx') and v['reviewed_all_main_pages']
    shutil.copy2(QA/'candidate.xlsx',TARGET.with_suffix('.tmp.xlsx'))
    TARGET.with_suffix('.tmp.xlsx').replace(TARGET)
    dump('publication.json',{'path':str(TARGET),'sha256':sha(TARGET),'actual_tests':False});print('Published same filename',sha(TARGET))

if __name__=='__main__':{'prepare':prepare,'richdiff':richdiff,'verify':verify,'publish':publish}[sys.argv[1]]()
