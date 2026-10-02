"""Revise the live style pack and unit presentation without new design claims."""
from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import zipfile, json, hashlib, re, sys
from lxml import etree as E

T=Path(__file__).resolve().parent
PROJECT=T.parents[3]
R=Path('E:/workspace/rules-for-codex')
RT=R/'작업/문서/20261001_0918_Word-공통-스타일'
Q=T/'qa/design-readability/purpose-and-compact-style'
RQ=RT/'qa/compact-unit-styles'
LIVE=PROJECT/'docs/참고/인허가/사운드1/DHF/03_설계 출력/SW/펌웨어/SDDD-FW-TD24A 소프트웨어 상세 설계서.docx'
TPL=R/'지침/양식/Word/TDC 공통 스타일.dotx'
OLDQ=T/'qa/design-readability/structured-units'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
q=lambda t:'{'+W+'}'+t
NS={'w':W,'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
ANCHORS=[0,.35,.7,1.05,1.4,1.75]
SIZES=[14,13,12,11,10,10]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):
    with zipfile.ZipFile(p) as z:return {n:z.read(n) for n in z.namelist()}
def write(p,parts):
    with zipfile.ZipFile(p,'w',zipfile.ZIP_DEFLATED) as z:
        for n,v in parts.items():z.writestr(n,v)
def el(p,t,**a):
    x=E.SubElement(p,q(t))
    for k,v in a.items():x.set(q(k),str(v))
    return x
def child(p,t,**a):
    x=p.find(q(t))
    if x is None:x=el(p,t)
    for k,v in a.items():x.set(q(k),str(v))
    if p.tag==q('tblPr'):
        order=['tblStyle','tblpPr','tblOverlap','bidiVisual','tblStyleRowBandSize','tblStyleColBandSize','tblW','jc','tblCellSpacing','tblInd','tblBorders','shd','tblLayout','tblCellMar','tblLook','tblCaption','tblDescription','tblPrChange']
        for item in sorted(list(p),key=lambda e:order.index(E.QName(e).localname) if E.QName(e).localname in order else 99):p.remove(item);p.append(item)
    return x
def tw(cm):return round(cm*1440/2.54)
def txt(p):return ''.join(p.xpath('.//w:t[not(ancestor::w:del) and not(ancestor::w:tr[w:trPr/w:del])]/text()',namespaces=NS))
def indent(p,n):
    x=child(p,'ind');x.attrib.clear();x.set(q('left'),str(n));x.set(q('right'),'0');x.set(q('hanging'),'0')
def jsonout(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf8')
def xmlbytes(r):return E.tostring(r,xml_declaration=True,encoding='UTF-8',standalone=True)

def template():
    from docx import Document
    from docx.shared import Cm
    Q.mkdir(parents=True,exist_ok=True);RQ.mkdir(parents=True,exist_ok=True)
    base=RQ/'before.dotx'
    if not base.exists():base.write_bytes(TPL.read_bytes())
    assert sha(TPL)==sha(base),'External template edit'
    parts=read(base);st=E.fromstring(parts['word/styles.xml']);num=E.fromstring(parts['word/numbering.xml'])
    names={s.find(q('name')).get(q('val')):s for s in st.findall(q('style'))}
    for name,s in names.items():
        if not name.startswith('TDC '):continue
        pp=child(s,'pPr');rp=child(s,'rPr')
        if name!='TDC 문서 제목':child(rp,'sz',val=20);child(rp,'szCs',val=20)
        m=re.fullmatch(r'TDC 제목 ([1-6])',name)
        if m:
            level=int(m[1]);child(rp,'sz',val=SIZES[level-1]*2);child(rp,'szCs',val=SIZES[level-1]*2)
            indent(pp,tw(ANCHORS[level-1]));child(pp,'spacing',before=160,after=80,line=276,lineRule='auto')
            for x in pp.findall(q('tabs')):pp.remove(x)
        elif re.search(r' [1-6]$',name):
            level=int(name[-1]);indent(pp,0 if s.get(q('type'))=='table' else tw(ANCHORS[level-1]))
            if s.get(q('type'))=='table':child(child(s,'tblPr'),'tblInd',w=tw(ANCHORS[level-1])+100,type='dxa')
        if name in ['TDC 표 머리행','TDC 표 항목']:child(pp,'jc',val='center')
        if name in ['TDC 표 본문','TDC 표 설명']:child(pp,'jc',val='left')
        if s.get(q('type'))=='table':
            child(child(s,'trPr'),'jc',val='left')
            for x in s.findall(q('tblStylePr')):s.remove(x)
            for role in ['firstRow','firstCol']:
                cond=el(s,'tblStylePr',type=role)
                child(child(cond,'pPr'),'jc',val='center')
                cr=child(cond,'rPr');child(cr,'b',val=1);child(cr,'sz',val=20)
                tc=child(cond,'tcPr');child(tc,'shd',val='clear',color='auto',fill='F5F5F5');child(tc,'vAlign',val='center')
    for lv in num.findall('.//'+q('lvl')):
        ps=lv.find(q('pStyle'))
        if ps is None:continue
        name=next((n for n,s in names.items() if s.get(q('styleId'))==ps.get(q('val'))),'')
        m=re.fullmatch(r'TDC 제목 ([1-6])',name)
        if m:
            i=int(m[1]);child(lv,'suff',val='space');pp=child(lv,'pPr');indent(pp,tw(ANCHORS[i-1]))
            for x in pp.findall(q('tabs')):pp.remove(x)
            rp=child(lv,'rPr');child(rp,'sz',val=SIZES[i-1]*2);child(rp,'szCs',val=SIZES[i-1]*2)
    parts['word/styles.xml']=xmlbytes(st);parts['word/numbering.xml']=xmlbytes(num)
    write(RQ/'candidate.dotx',parts)
    parts['[Content_Types].xml']=parts['[Content_Types].xml'].replace(b'wordprocessingml.template.main+xml',b'wordprocessingml.document.main+xml')
    write(RQ/'base.docx',parts)
    d=Document(RQ/'base.docx')
    for p in list(d.paragraphs):p._element.getparent().remove(p._element)
    d.add_paragraph('공통 스타일 간격 확인','TDC 문서 제목')
    for i in range(1,7):
        d.add_paragraph('번호 뒤 공백과 제목 '+str(i),'TDC 제목 '+str(i))
        purpose=d.add_paragraph('이 유닛은 입력을 받아 처리 결과를 전달한다. 목적은 표 밖 일반 문장으로 읽으며 본문, 표와 캡션의 기준선은 제목 번호 시작에 맞춘다.','TDC 본문 '+str(i))
        purpose.paragraph_format.keep_with_next=True
        d.add_paragraph('표 '+str(i)+'. 입력/출력 및 처리','TDC 표 캡션 '+str(i))
        tb=d.add_table(rows=4,cols=2);tb.style='TDC 표 '+str(i);tb.autofit=False
        widths=[2.3,16.6-ANCHORS[i-1]-2.3]
        for col,w in zip(tb.columns,widths):col.width=Cm(w)
        for ri,vals in enumerate([['항목','설계 내용'],['입력','입력 조건과 데이터'],['출력','처리한 결과'],['처리','데이터를 처리하고 결과를 전달한다.']]):
            for ci,(c,v) in enumerate(zip(tb.rows[ri].cells,vals)):
                c.width=Cm(widths[ci]);c.text=v;c.paragraphs[0].style='TDC 표 머리행' if ri==0 else 'TDC 표 항목' if ci==0 else 'TDC 표 설명'
                tc=c._tc.get_or_add_tcPr();child(tc,'vAlign',val='center')
                if ri==0 or ci==0:child(tc,'shd',val='clear',color='auto',fill='F5F5F5')
        child(tb._tbl.tblPr,'tblInd',w=tw(ANCHORS[i-1])+100,type='dxa')
        child(tb._tbl.tblPr,'jc',val='left')
        child(tb.rows[0]._tr.get_or_add_trPr(),'tblHeader')
        d.add_paragraph('','TDC 본문 '+str(i))
        d.add_paragraph('참고: 표 뒤 빈 문단도 정상 본문 크기를 유지한다.','TDC 주석 '+str(i))
    d.save(RQ/'specimen.docx')
    jsonout(RQ/'style-spec.json',{'style_count':59,'anchors_cm':ANCHORS,'headings_pt':SIZES,'body_pt':10,'table_pt':10,'suffix':'space','classification_fill':'F5F5F5','baseline_sha256':sha(base)})

def prepare():
    template()
    base=Q/'before.docx'
    if not base.exists():base.write_bytes(LIVE.read_bytes())
    assert sha(LIVE)==sha(base),'External SDDD edit'
    parts=read(base);tp=read(RQ/'candidate.dotx');doc=E.fromstring(parts['word/document.xml']);st=E.fromstring(parts['word/styles.xml']);ts=E.fromstring(tp['word/styles.xml']);num=E.fromstring(parts['word/numbering.xml'])
    existing={s.find(q('name')).get(q('val')):s for s in st.findall(q('style'))}
    customs={s.find(q('name')).get(q('val')):s for s in ts.findall(q('style')) if s.find(q('name')).get(q('val')).startswith('TDC ')}
    ids={n:existing[n].get(q('styleId')) for n in customs};remap={s.get(q('styleId')):ids[n] for n,s in customs.items()}
    # Keep the expanded inheritance from the previous application; copy only changed
    # formatting and conditional table properties from the actual template.
    for n,s in customs.items():
        dest=existing[n]
        for tag in ['pPr','rPr','tblPr','tblStylePr']:
            sp=s.findall(q(tag))
            if tag=='tblStylePr':
                for x in dest.findall(q(tag)):dest.remove(x)
                for x in sp:dest.append(deepcopy(x))
            else:
                for prop in sp:
                    dp=child(dest,tag)
                    for x in prop:
                        if x.tag==q('numPr'):continue
                        old=dp.find(x.tag)
                        if old is not None:dp.remove(old)
                        dp.append(deepcopy(x))
                if tag=='pPr':
                    for x in child(dest,'pPr').findall(q('tabs')):child(dest,'pPr').remove(x)
        if n.startswith('TDC 제목 '):
            pp=child(dest,'pPr');np=child(pp,'numPr');child(np,'numId',val=7);child(np,'ilvl',val=int(n[-1])-1)
    aid=num.xpath('./w:num[@w:numId="7"]/w:abstractNumId/@w:val',namespaces=NS)[0]
    abstract=num.xpath('./w:abstractNum[@w:abstractNumId="'+aid+'"]',namespaces=NS)[0]
    for lv in abstract.findall(q('lvl')):
        i=int(lv.get(q('ilvl')))+1
        if i>6:continue
        child(lv,'suff',val='space');pp=child(lv,'pPr');indent(pp,tw(ANCHORS[i-1]))
        for x in pp.findall(q('tabs')):pp.remove(x)
        rp=child(lv,'rPr');child(rp,'sz',val=SIZES[i-1]*2);child(rp,'szCs',val=SIZES[i-1]*2)
    units=json.loads((OLDQ/'fragments.json').read_text(encoding='utf8'))
    originals=json.loads((OLDQ/'units-before.json').read_text(encoding='utf8'))
    unitnames={u['id']:u['name'] for u in originals}
    body=doc.find(q('body'));bookid=max([int(v) for v in doc.xpath('.//w:bookmarkStart/@w:id',namespaces=NS) if v.isdigit()]+[0])+1
    manifest=[]
    for u in units:
        purpose=u['fields']['목적']
        matches=[t for t in body.findall(q('tbl')) if any(txt(c)==purpose for c in t.xpath('./w:tr/w:tc',namespaces=NS))]
        if matches:
            assert len(matches)==1,u['unit']
            start=matches[0].getprevious();assert start.tag==q('p') and '유닛 설계 내용' in txt(start),u['unit']
            last=matches[0]
            nxt=last.getnext()
            if nxt is not None and nxt.tag==q('p') and '구현 조건' in txt(nxt):
                last=nxt.getnext();assert last.tag==q('tbl'),u['unit']
            nxt=last.getnext()
            if nxt is not None and txt(nxt).startswith('입출력 예 :'):last=nxt
        else:
            ps=[p for p in body.findall(q('p')) if txt(p)=='목적 : '+purpose]
            assert len(ps)==1,(u['unit'],len(ps));start=last=ps[0]
            while last.getnext() is not None and re.match('입력 :|출력 :|처리 :|관련 함수 :|동작 조건 :|필수 데이터/상수 :|입출력 예 :',txt(last.getnext())):last=last.getnext()
        end=last.getnext();assert end is not None and end.tag==q('p'),u['unit']
        name='CompactUnit_'+u['unit'].replace('-','_')
        start.insert(1,E.Element(q('bookmarkStart'),{q('id'):str(bookid),q('name'):name}))
        pos=1 if end.find(q('pPr')) is not None else 0
        end.insert(pos,E.Element(q('bookmarkEnd'),{q('id'):str(bookid)}));bookid+=1
        manifest.append({**u,'bookmark':name,'name':unitnames[u['unit']]})
    for n,r in [('word/document.xml',doc),('word/styles.xml',st),('word/numbering.xml',num)]:parts[n]=xmlbytes(r)
    write(Q/'styled.docx',parts)
    folder=Q/'fragments';folder.mkdir(exist_ok=True)
    def para(parent,text,style):
        p=el(parent,'p');pp=el(p,'pPr');el(pp,'pStyle',val=ids[style]);el(pp,'outlineLvl',val=9)
        if text:
            r=el(p,'r');t=el(r,'t');t.text=text;t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
        return p
    def caption(parent,title):
        p=para(parent,'표 ','TDC 표 캡션 4')
        for typ in ['begin','separate','end']:
            r=el(p,'r');el(r,'fldChar',fldCharType=typ)
            if typ=='begin':el(el(p,'r'),'instrText').text=' SEQ 표 \\* ARABIC '
            if typ=='separate':el(el(p,'r'),'t').text='1'
        el(el(p,'r'),'t').text='. '+title
    def table(parent,rows,head):
        total=tw(16.6-ANCHORS[3]);widths=[tw(2.3),total-tw(2.3)]
        tb=el(parent,'tbl');pr=el(tb,'tblPr');el(pr,'tblStyle',val=ids['TDC 표 4']);el(pr,'tblW',w=total,type='dxa');el(pr,'jc',val='left');el(pr,'tblInd',w=tw(ANCHORS[3])+100,type='dxa');el(pr,'tblLayout',type='fixed')
        grid=el(tb,'tblGrid')
        for w in widths:el(grid,'gridCol',w=w)
        for ri,row in enumerate([head]+rows):
            tr=el(tb,'tr');rp=el(tr,'trPr');el(rp,'cantSplit')
            if ri==0:el(rp,'tblHeader')
            for ci,(v,w) in enumerate(zip(row,widths)):
                tc=el(tr,'tc');cp=el(tc,'tcPr');el(cp,'tcW',w=w,type='dxa');el(cp,'vAlign',val='center')
                if ri==0 or ci==0:el(cp,'shd',val='clear',color='auto',fill='F5F5F5')
                para(tc,v,'TDC 표 머리행' if ri==0 else 'TDC 표 항목' if ci==0 else 'TDC 표 설명')
        para(parent,'','TDC 본문 4')
    PKG='http://schemas.microsoft.com/office/2006/xmlPackage';REL='http://schemas.openxmlformats.org/package/2006/relationships'
    for u in manifest:
        fields=u['fields'];root=E.Element(q('document'),nsmap={'w':W});b=el(root,'body')
        purpose=para(b,fields['목적'],'TDC 본문 4');child(child(purpose,'pPr'),'keepNext',val=1)
        caption(b,u['name']+' / 유닛 설계 내용');table(b,[(k,fields[k]) for k in ['입력','출력','처리']],['항목','설계 내용'])
        conditions=[(k,fields[k]) for k in ['관련 함수','동작 조건','필수 데이터/상수'] if k in fields]
        if conditions:caption(b,u['name']+' / 구현 조건');table(b,conditions,['구분','구현 조건'])
        if '입출력 예' in fields:para(b,'입출력 예 : '+fields['입출력 예'],'TDC 주석 4')
        pkg=E.Element('{'+PKG+'}package',nsmap={'pkg':PKG})
        def part(n,ct,node):
            p=E.SubElement(pkg,'{'+PKG+'}part',{'{'+PKG+'}name':n,'{'+PKG+'}contentType':ct});E.SubElement(p,'{'+PKG+'}xmlData').append(node)
        rr=E.Element('{'+REL+'}Relationships',nsmap={None:REL});E.SubElement(rr,'{'+REL+'}Relationship',Id='rId1',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument',Target='word/document.xml');part('/_rels/.rels','application/vnd.openxmlformats-package.relationships+xml',rr)
        rr=E.Element('{'+REL+'}Relationships',nsmap={None:REL});E.SubElement(rr,'{'+REL+'}Relationship',Id='rIdStyles',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles',Target='styles.xml');part('/word/_rels/document.xml.rels','application/vnd.openxmlformats-package.relationships+xml',rr)
        part('/word/document.xml','application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',root)
        part('/word/styles.xml','application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml',deepcopy(st))
        path=folder/(u['unit']+'.xml');path.write_bytes(xmlbytes(pkg));u['fragment']=str(path);u['tables']=1+bool(conditions)
    jsonout(Q/'fragments.json',manifest)
    jsonout(Q/'prepare.json',{'baseline_sha256':sha(base),'template_baseline_sha256':sha(RQ/'before.dotx'),'styles':ids,'units':115,'unit_tables':sum(u['tables'] for u in manifest)})
    print(json.dumps({'units':115,'unit_tables':sum(u['tables'] for u in manifest),'template':str(RQ/'candidate.dotx')},ensure_ascii=False))

if __name__=='__main__':prepare()
