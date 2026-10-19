"""Inspect the current report and design, preserving user-owned baselines."""
from pathlib import Path
import hashlib
import json
import shutil
import zipfile
import openpyxl
from lxml import etree as E

ROOT = Path('E:/workspace/projects/sound1-fw-e8300-backup')
TASK = ROOT / 'docs/tasks/인허가/20260929_1755_DHF-펌웨어-설계출력'
DHF = ROOT / 'docs/참고/인허가/사운드1/DHF'
QA = TASK / 'qa/audit-expected-report-20261012'
PREFIX = '[김은수 작성] [Rev.0] '
REPORT = DHF / ('04_설계 검증/SW/펌웨어/' + PREFIX + 'SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서.xlsx')
DESIGN = DHF / ('03_설계 출력/SW/펌웨어/' + PREFIX + 'SDDD-FW-TD24A 소프트웨어 상세 설계서.docx')
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def current_text(elem):
    pieces = []
    for node in elem.iter():
        if any(a.tag == '{' + NS['w'] + '}del' for a in node.iterancestors()):
            continue
        if node.tag == '{' + NS['w'] + '}t': pieces.append(node.text or '')
        elif node.tag == '{' + NS['w'] + '}br': pieces.append('\n')
        elif node.tag == '{' + NS['w'] + '}tab': pieces.append('\t')
    return ''.join(pieces)

def inspect():
    QA.mkdir(parents=True, exist_ok=True)
    before = QA / 'report-before.xlsx'
    if not before.exists(): shutil.copyfile(REPORT, before)
    else: assert sha(before) == sha(REPORT), 'Report changed after baseline'
    protected = {str(p): sha(p) for p in DHF.rglob('*')
                 if p.is_file() and p.suffix in ('.xlsx', '.docx')
                 and p != REPORT and not p.name.startswith('~$')}
    sources = {str(p): sha(p) for folder in ('src', 'tests') for p in (ROOT / folder).rglob('*')
               if p.is_file() and p.suffix.lower() in ('.c', '.h', '.cpp', '.py', '.mjs', '.ps1', '.cproject', '.project')}
    (QA / 'baseline.json').write_text(json.dumps({'report': str(REPORT), 'sha256': sha(REPORT),
        'protected': protected, 'sources': sources}, ensure_ascii=False, indent=2), encoding='utf8')
    wb = openpyxl.load_workbook(REPORT, data_only=False, rich_text=True)
    ws = wb.worksheets[1]
    units = []
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 4).value or '').strip() not in [f'SW-FW-UnitTest-CM3-ID-{n}' for n in (1, 5, 46)]: continue
        end = r + 1
        while end <= ws.max_row and not str(ws.cell(end, 4).value or '').strip().startswith('SW-FW-UnitTest-'): end += 1
        units.append({'start': r, 'end': end - 1, 'rows': [{'row': x, 'height': ws.row_dimensions[x].height,
            'cells': {ws.cell(x,c).column_letter: str(ws.cell(x,c).value) for c in range(1,14) if ws.cell(x,c).value is not None}}
            for x in range(r,end)], 'merges': [str(m) for m in ws.merged_cells.ranges if m.min_row < end and m.max_row >= r]})
    with zipfile.ZipFile(DESIGN) as z: doc = E.fromstring(z.read('word/document.xml'))
    body = doc.find('w:body', NS)
    blocks = []
    for n in (1, 5, 46):
        start = next(i for i,p in enumerate(body) if current_text(p).strip().startswith(f'SU-ID-CM3-{n:03},'))
        end = next(i for i in range(start+1,len(body)) if current_text(body[i]).strip().startswith('SU-ID-CM3-'))
        blocks.append({'unit': n, 'blocks': [current_text(body[i]) for i in range(start,end)]})
    output = {'sheets': wb.sheetnames, 'max_row': ws.max_row, 'units': units, 'design': blocks,
              'missing': [{ 'cell': c.coordinate, 'text': str(c.value)} for row in ws for c in row
                          if c.value and any(x in str(c.value) for x in ('[E8300', '結果未', '결과 미기록', '코드 반영 후'))]}
    (QA / 'inspection.json').write_text(json.dumps(output,ensure_ascii=False,indent=2), encoding='utf8')
    print(json.dumps({'sha256': sha(REPORT), 'sheets': wb.sheetnames, 'rows': ws.max_row,
        'units': [(u['start'],u['end']) for u in units], 'missing_cells': [x['cell'] for x in output['missing']],
        'protected_files': len(protected), 'source_files':len(sources)}, ensure_ascii=False))

def make_plan():
    d = json.loads((QA / 'inspection.json').read_text(encoding='utf8'))
    old = {f'{col}{r["row"]}': v for u in d['units'] for r in u['rows'] for col,v in r['cells'].items()}
    pad = lambda s: '\n' + s.strip() + '\n'
    edits = {}
    def put(cell, value): edits[cell] = pad(value)
    header = '[LOG] IDENT=EVENTLOG.TXT, WRITE INDEX=1, ENTITY COUNT=1'
    record = '      <  0> : TIME {26-10-02-13-14-15} ADDR {12:34:56:78:9A:BC} EVENT {00000040}'
    fixture = header + '\n' + record

    # Keep the old observations; move the existing third condition into the spare row above.
    k3 = old['K792'].split('\n\n조건 [4]')[0].strip()
    l3 = old['L792'].split('\n\n검증 [4]')[0].strip()
    put('K791', old['K791'].strip() + '\n\n' + k3)
    put('L791', old['L791'].strip() + '\n\n' + l3)
    put('K792', '''조건 [4] 감사기록 파일 준비
각 입력은 독립된 4,096바이트 파일로 시험한다.
입력 ① 파일 없음, ② 식별자=INVALID0.TXT,
③ 다음 위치=255, ④ 기록 수=256.
예상: 파일 크기=4,096, 식별자=EVENTLOG.TXT,
기록 수=0, 다음 위치=0, 반환=0. 기록 영역은 모두 0.
입력 ⑤ 정상 기록 1개, 다음 위치=1, 기록 수=1.
예상: 크기=4,096, 기록 수=1, 다음 위치=1, 반환=0,
파일 내용 동일, 파일 쓰기=0회.

조건 [5] 변경 내용 저장 실패
입력: 새 파일 초기화 중 f_sync 응답=1(저장 실패).
예상: 초기화 함수 반환=-1.

조건 [6] 파일 열기 실패
입력: f_open 응답=1(열기 실패).
예상: 초기화 함수 반환=-1.

조건 [7] 초기화 정보 쓰기 실패
입력: f_write 응답=1 또는 기록 크기=4,095바이트.
각 경우 초기화 함수 반환=-1.

조건 [8] 파일 닫기 실패
입력: 새 파일 준비 후 f_close 응답=1(닫기 실패).
예상: 초기화 함수 반환=-1.''')
    l4 = old['L792'].split('\n\n검증 [4]')[1].split('\n\n검증 [5]')[0]
    l5 = old['L792'].split('\n\n검증 [5]')[1].split('\n[E8300 보완]')[0]
    put('L792', '검증 [4]' + l4 + '\n\n검증 [5]' + l5 + '''

[E8300 보완] 코드 반영 후 확인할 예상값
조건 [5]: 저장 실패 입력 → 초기화 반환=-1.
조건 [6]: 열기 실패 입력 → 초기화 반환=-1.
조건 [7]: 쓰기 실패 입력 → 초기화 반환=-1.
조건 [7]: 4,095바이트 쓰기 → 초기화 반환=-1.
조건 [8]: 닫기 실패 입력 → 초기화 반환=-1.
위 값은 후속 시험의 예상값이다.''')
    put('H789', old['H789'].strip() + '\n- 오류 응답 1은 저장 함수의 실패 값이다. 열기/쓰기/저장/닫기는 한 번에 한 곳만 실패시킨다. 조건 [4]의 정상 기록은 시간=26-10-02-13-14-15, 주소=12:34:56:78:9A:BC, 이벤트=0x40이다.')
    put('J789', old['J789'].strip() + '\n- 조건 [4]는 파일 크기/관리 정보/기록 내용을 RTT로 출력하여 비교한다. 조건 [5]~[8]은 초기화 함수의 반환값을 출력해 -1인지 확인한다.')
    put('M789', '''기존 [1]~[3] 통과
기존 보충 [4] 통과

[E8300 보완]
[5] 기존 반환=0
코드 반영 후 재시험
[6]~[8] 예상값 작성
코드 반영 후 시험''')

    # Complete the access-gate cases without changing their existing condition IDs.
    put('I809', '''조건 [3]은 감사기록 읽기 함수의 독립 시험이다.
조건 [4]~[8]은 조회 명령의 접근 제한 시험이다.
[E8300 보완] 조회 제한 코드를 반영한 뒤 조건 [4]~[8]을 확인한다.''')
    put('K811', '''조건 [4] 정보 일치 / 정상 연결
등록 ID=실제 ID=0x12345678, 정보 확인 완료,
초기화 완료, 정상 연결. --dump_log 1회 입력.
예상: 조회 허용, 읽기=1회, 읽기 반환=0, 기록 출력=1건.
''' + fixture + '''

조건 [5] 정보 일치 / 미연결
등록 ID=실제 ID=0x12345678, 미연결. 조회 1회.
예상: 차단, 읽기=0회, 저장 기록 출력=0건.

조건 [6] 정보 불일치
등록 ID=0x12345678, 실제 ID=0x12345679.
정상 연결/미연결에서 각각 조회 1회.
예상: 각각 차단, 읽기=0회, 저장 기록 출력=0건.

조건 [7] 허용된 조회 후 연결 해제
입력: 조건 [4] 실행 후 연결 해제, 조회 1회 추가.
예상: 차단, 추가 읽기=0회, 추가 기록 출력=0건.

조건 [8] 정보 확인 전 또는 초기화 중
정상 연결 상태를 넣더라도 정보 미확인/초기화 중을
각각 설정하고 조회 1회. 둘 다 차단,
읽기=0회, 저장 기록 출력=0건.
조건 [4]~[8] 모두 파일 쓰기=0회, 파일 내용 동일.''')
    put('L811', '''[E8300 보완] 코드 반영 후 확인할 예상 출력

조건 [4]의 예상값
조회=허용, 읽기=1회, 반환=0, 기록 출력=1건.
''' + fixture + '''

조건 [5]의 예상값
조회=차단, 읽기=0회, 기록 출력=0건.

조건 [6]의 예상값
정상 연결: 차단, 읽기=0회, 기록 출력=0건.
미연결: 차단, 읽기=0회, 기록 출력=0건.

조건 [7]의 예상값
연결 해제 후: 차단, 추가 읽기=0회, 추가 출력=0건.

조건 [8]의 예상값
정보 미확인: 차단, 읽기=0회, 기록 출력=0건.
초기화 중: 차단, 읽기=0회, 기록 출력=0건.

조건 [4]~[8]: 파일 쓰기=0회, 파일 내용 동일.
차단 시 읽기 함수는 실행하지 않으므로 반환값 없음.
위 값은 후속 시험의 예상값이다.''')
    put('H809', old['H809'].strip() + '\n- 정보 확인 완료와 초기화 완료 상태도 독립 입력으로 설정한다. 조건 [4]~[8]은 같은 기록 1개를 사용하고 요청마다 연결 조건을 다시 입력한다.')
    put('J809', old['J809'].strip().replace('[4]~[7]', '[4]~[8]') + '\n- 정보 미확인/초기화 중에도 조회를 요청한다. RTT로 읽기 호출 수/기록 출력 수/파일 쓰기 수와 파일 비교 결과를 확인한다. 차단 안내 문구는 기록 출력 건수에서 제외한다.')
    put('M809', '''기존 [1]~[3]
통과

[E8300 보완]
추가 [4]~[8]
예상값 작성
코드 반영 후 시험''')

    # Separate precise design expectations from earlier supplemental executions.
    normal_k = '''조건 [5] 연결 확인 / 감사기록 저장
일반 명령 허용 분기와 별도로 연결 확인을 처리한다.
입력: 연결 확인 명령, 시간={26,10,2,13,14,15},
주소={12,34,56,78,9A,BC}(16진수), 빈 정상 기록 파일.
예상: 시간 갱신 → 주소 설정 → 연결 이벤트 저장.
시간=26-10-02-13-14-15, 주소=12:34:56:78:9A:BC,
이벤트=0x40, 기록 수=1, 다음 위치=1, 반환=0.
갱신 적용 후 999ms: 초=15, 1,000ms: 초=16.
다음 위치=254/기록 수=255에서 1회 저장하면
다음 위치=0, 기록 수=255. 다음 저장 위치는 0이다.'''
    put('K1212', old['K1212'].strip() + '\n\n' + normal_k)
    put('K1213', '''조건 [6] 파일 닫기 실패
입력: 기록을 쓴 뒤 f_close 응답=1.
예상: 기록 함수 반환=-1.

조건 [7] 파일 열기 실패
입력: f_open 응답=1. 예상: 기록 함수 반환=-1.

조건 [8] 기록 파일 읽기/관리 정보 오류
입력: f_read 응답=1, 읽은 크기=4,095바이트,
식별자=INVALID0.TXT, 다음 위치=255, 기록 수=256.
각 입력은 독립 시험. 모두 기록 함수 반환=-1.

조건 [9] 기록 쓰기 실패
f_write 응답=1 / 기록 크기=4,095바이트를 각각 입력.
예상: 기록 함수 반환=-1.

조건 [10] 시간 경계
각 시간을 새 기준으로 설정하고 갱신 적용 후 1,000ms 입력.
그 뒤 이벤트 0x40을 저장한다. 예상 기록 시간:
26-10-02-13-14-59 → 26-10-02-13-15-00
26-10-02-23-59-59 → 26-10-03-00-00-00
26-04-30-23-59-59 → 26-05-01-00-00-00
28-02-28-23-59-59 → 28-02-29-00-00-00
26-02-28-23-59-59 → 26-03-01-00-00-00
26-12-31-23-59-59 → 27-01-01-00-00-00''')
    normal_l = old['L1213'].split('\n\n검증 [6]')[0].strip()
    put('L1212', old['L1212'].strip() + '\n\n' + normal_l + '''

[E8300 보완] 조건 [5]의 추가 예상값
999ms: 시간=26-10-02-13-14-15.
1,000ms: 시간=26-10-02-13-14-16.
위 값은 후속 시험의 예상값이다.''')
    error_l = '검증 [6]' + old['L1213'].split('\n\n검증 [6]')[1].split('\n[E8300 보완]')[0]
    put('L1213', error_l + '''

[E8300 보완] 코드 반영 후 확인할 예상값
조건 [6] 닫기 실패: 기록 함수 반환=-1.
조건 [7] 열기 실패: 기록 함수 반환=-1.
조건 [8] 읽기 실패/4,095바이트/식별자 오류/
위치255/기록 수256: 각각 기록 함수 반환=-1.
조건 [9] 쓰기 실패/4,095바이트: 각각 기록 함수 반환=-1.

조건 [10]의 예상 기록 시간
분 경계: 26-10-02-13-15-00
하루 경계: 26-10-03-00-00-00
월 경계: 26-05-01-00-00-00
윤년 2월: 28-02-29-00-00-00
평년 2월: 26-03-01-00-00-00
연도 경계: 27-01-01-00-00-00
위 값은 후속 시험의 예상값이다.''')
    put('H1208', old['H1208'].strip() + '\n- 감사기록 오류 입력은 각 조건마다 빈 정상 파일에서 시작한다. 오류 응답 1은 저장 함수의 실패 값이고 읽기/쓰기 크기는 4,096 또는 4,095바이트로 설정한다.')
    put('J1208', old['J1208'].strip() + '\n- 시간/주소/이벤트/기록 수/다음 위치와 저장 함수 반환값을 RTT로 확인한다. 시간 갱신이 적용된 뒤 999ms와 1,000ms를 각각 입력한다. 저장 오류는 한 조건씩 넣어 반환=-1을 비교한다.')
    put('M1208', '''기존 [1]~[5] 통과
기존 감사기록 정상값 통과

[E8300 보완]
[6] 기존 반환=0
코드 반영 후 재시험
[7]~[10] 예상값 작성
코드 반영 후 시험''')
    plan = {'edits': edits, 'heights': {'791':175.5,'792':409.5,'811':409.5,'1212':330,'1213':409.5},
            'units':[1,5,46], 'new_condition_ids':{'1':[5,6,7,8], '5':[4,5,6,7,8], '46':[6,7,8,9,10]},
            'full_blue':['L811'], 'expected_prefix':'[E8300 보완]'}
    assert not any('·' in x or '。' in x for x in edits.values())
    (QA / 'sheet-edits.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({'changed_cells':list(edits),'condition_counts':{1:8,5:8,46:10}},ensure_ascii=False))

if __name__ == '__main__':
    import sys
    if '--plan' in sys.argv: make_plan()
    else: inspect()
