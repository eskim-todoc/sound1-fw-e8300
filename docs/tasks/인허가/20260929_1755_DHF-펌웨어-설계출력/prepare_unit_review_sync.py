"""검토용 노란 배경 복구와 보고서 A~K의 기준서 동기화.

Artifact Tool의 작성값/배경을 기존 OOXML 보존 도구로 결합한다.
이전 사본에서 배경만 복구하며 이전 문구/결과는 가져오지 않는다.
"""
from pathlib import Path
import copy
import hashlib
import json
import shutil
import sys
import zipfile
from openpyxl import load_workbook
import preserve_aligned_report_packages as native
import sync_unit_criterion_first as sync

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
DHF = ROOT / 'docs/참고/인허가/사운드1/DHF'
QA = TASK / 'qa/unit-review-color-criteria-sync-20261012'
REPORT = DHF / '04_설계 검증/SW/펌웨어/[김은수 작성] [Rev.0] SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서.xlsx'
CRITERION = DHF / '03_설계 출력/SW/펌웨어/[김은수 작성] [Rev.0] SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx'
REFERENCE = TASK / 'qa/audit-report-dates-format-20261012/report-before.xlsx'
RS, CS = sync.RS, sync.CS
N = native.N


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def put(name, value):
    (QA / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def get(name):
    return json.loads((QA / name).read_text(encoding='utf-8'))


def configure():
    sync.QA = QA / 'criterion'
    sync.REPORT = QA / 'report.xlsx'
    sync.TARGET = CRITERION
    sync.USE_SOURCE_THEME = True


def prepare():
    QA.mkdir(parents=True, exist_ok=False)
    shutil.copy2(REPORT, QA / 'report-before.xlsx')
    shutil.copy2(CRITERION, QA / 'criterion-before.xlsx')
    old, current = native.parts(REFERENCE), native.parts(REPORT)
    reference_cells = {c.get('r'):c for c in native.parse(old[native.sheet_path(old, RS)]).findall('.//s:c', N)}
    cells = native.parse(current[native.sheet_path(current, RS)]).findall('.//s:c', N)
    styles, ref_styles = native.parse(current['xl/styles.xml']), native.parse(old['xl/styles.xml'])
    fills, xfs = styles.find('s:fills', N), styles.find('s:cellXfs', N)
    ref_fills, ref_xfs = ref_styles.find('s:fills', N), ref_styles.find('s:cellXfs', N)
    restore = []
    for c in cells:
        fill = fills[int(xfs[int(c.get('s','0'))].get('fillId','0'))]
        fg = fill.find('s:patternFill/s:fgColor', N)
        if fg is None or fg.get('rgb') != 'FFFFF2CC':
            continue
        address = c.get('r')
        rc = reference_cells[address]
        original = ref_fills[int(ref_xfs[int(rc.get('s','0'))].get('fillId','0'))]
        key = native.serialize(original)
        fill_id = next(i for i,f in enumerate(fills) if native.serialize(f) == key)
        original_fg = original.find('s:patternFill/s:fgColor', N)
        color = '#' + original_fg.get('rgb')[-6:] if original_fg is not None and original_fg.get('rgb') else '#FFFFFF'
        restore.append(dict(cell=address, fill_id=fill_id, color=color))
    assert restore, 'No review highlight found'
    # 다른 채팅의 시스템 문서 작업과 충돌하지 않도록 보호 범위는 사용자
    # 전용 사본/과거 사본과 제품/시험 소스로 잡고, 다른 문서는 쓰지 않는다.
    protected = {str(p):sha(p) for p in DHF.rglob('*') if p.is_file() and p.suffix in ('.docx','.xlsx')
                 and any(s in p.name for s in ('전용','수정전','수정중')) and not p.name.startswith('~$')}
    source = {str(p):sha(p) for directory in ('src','tests') for p in (ROOT/directory).rglob('*')
              if p.is_file() and p.suffix in ('.c','.h','.hct','.py','.ps1','.mjs')}
    report, criterion = load_workbook(REPORT, rich_text=True)[RS], load_workbook(CRITERION, rich_text=True)[CS]
    differences = [c.coordinate for row in report.iter_rows(max_col=11) for c in row
                   if str(c.value) != str(criterion[c.coordinate].value) and c.coordinate != 'B2']
    assert report.max_row == criterion.max_row == 1734
    assert {str(m) for m in report.merged_cells.ranges if m.max_col <= 11} == set(map(str, criterion.merged_cells.ranges))
    put('baseline.json', dict(reportHash=sha(REPORT), criterionHash=sha(CRITERION),
        protected=protected, source=source, commonDifferences=differences, referenceHash=sha(REFERENCE)))
    put('sheet-edits.json', dict(edits={}, heights={}, restore_fills=restore))
    print(json.dumps(dict(restoredNativeCells=len(restore), criterionValueUpdates=len(differences),
                         protectedDocuments=len(protected), protectedSourceFiles=len(source)), ensure_ascii=False))


def build_report():
    package = native.parts(QA / 'report-before.xlsx')
    artifact = load_workbook(QA / 'report-artifact.xlsx')[RS]
    path = native.sheet_path(package, RS)
    root = native.parse(package[path])
    cells = {c.get('r'):c for c in root.findall('.//s:c',N)}
    styles = native.parse(package['xl/styles.xml'])
    xfs = styles.find('s:cellXfs',N)
    restored = get('sheet-edits.json')['restore_fills']
    for item in restored:
        c = cells[item['cell']]
        ac = artifact[item['cell']]
        if ac.__class__.__name__ != 'MergedCell':
            assert ac.fill.fgColor.rgb[-6:] == item['color'][1:], item['cell']
        # 채우기 없음과 흰색 채우기의 차이를 보존한다. 글꼴/테두리/문구는 유지한다.
        xf = copy.deepcopy(xfs[int(c.get('s','0'))])
        xf.set('fillId', str(item['fill_id']))
        key = native.serialize(xf)
        match = next((i for i,v in enumerate(xfs) if native.serialize(v) == key), None)
        if match is None:
            match = len(xfs)
            xfs.append(xf)
        c.set('s', str(match))
    xfs.set('count',str(len(xfs)))
    package[path] = native.serialize(root)
    package['xl/styles.xml'] = native.serialize(styles)
    with zipfile.ZipFile(QA / 'report.xlsx','w',zipfile.ZIP_DEFLATED) as z:
        for name,raw in package.items():
            z.writestr(name,raw)
    before,after = load_workbook(QA/'report-before.xlsx', rich_text=True),load_workbook(QA/'report.xlsx', rich_text=True)
    changed = {p['cell'] for p in restored}
    reference = load_workbook(REFERENCE, rich_text=True)[RS]
    for os,ns in zip(before,after):
        assert os.title == ns.title
        assert (os.max_row,os.max_column) == (ns.max_row,ns.max_column)
        assert set(map(str,os.merged_cells.ranges)) == set(map(str,ns.merged_cells.ranges))
        for row in os:
            for c in row:
                nc = ns[c.coordinate]
                assert sync.rich_key(c.value) == sync.rich_key(nc.value), c.coordinate
                for attr in ('font','alignment','border','number_format','protection'):
                    assert copy.copy(getattr(c,attr)) == copy.copy(getattr(nc,attr)), (c.coordinate,attr)
                wanted = reference[c.coordinate].fill if os.title == RS and c.coordinate in changed else c.fill
                assert copy.copy(nc.fill) == copy.copy(wanted), (c.coordinate,'fill')
    assert not any(c.fill.fgColor.rgb == 'FFFFF2CC' for row in after[RS] for c in row)
    old_package = native.parts(QA/'report-before.xlsx')
    assert {n for n in package if package[n] != old_package[n]} == {path,'xl/styles.xml'}
    put('report-checks.json', dict(changedCells=len(changed), textResultsVerdictsPreserved=True,
        onlyFillChanged=True, originalFillsRestored=True, hash=sha(QA/'report.xlsx')))
    print('Report fills restored; every value/result/verdict preserved')


def prepare_criterion():
    configure()
    sync.prepare()
    # 기존 도구의 광범위 보호 목록 대신 이번 작업의 보호 범위를 사용한다.
    base=sync.read('baseline.json')
    base['protected']=get('baseline.json')['protected']
    sync.save('baseline.json',base)


def build_criterion():
    configure()
    sync.build()
    sync.verify()


def reset_height_fixes():
    configure()
    plan=sync.read('plan.json')
    plan['heights']=[]
    sync.save('plan.json',plan)
    path=sync.QA/'native-heights.json'
    if path.exists():
        path.rename(sync.QA/'native-heights-first.json')


def fit_criterion():
    configure()
    path=QA/'native-review-checks.json'
    measurements=json.loads(path.read_text(encoding='utf-8-sig'))
    shutil.copy2(path,QA/'native-review-checks-first.json')
    fixes=[]
    for item in measurements:
        if item['book']=='criterion' and item['needed'] > item['available']+1:
            row=int(''.join(c for c in item['cell'] if c.isdigit()))
            height=min(409.5,item['needed']+8.5)
            fixes.append(dict(row=row,height=height))
    assert fixes
    sync.save('native-heights.json',fixes)
    plan=sync.read('plan.json')
    plan['heights']=[dict(key='criterion',sheet=RS,**item) for item in fixes]
    sync.save('plan.json',plan)
    print(json.dumps(fixes,ensure_ascii=False))


def checks():
    baseline = get('baseline.json')
    for field in ('protected','source'):
        for path,h in baseline[field].items():
            assert sha(path)==h, ('Protected file changed',path)
    assert sha(REPORT)==baseline['reportHash'], 'Report modified externally'
    assert sha(CRITERION)==baseline['criterionHash'], 'Criterion modified externally'
    report,criterion=load_workbook(QA/'report.xlsx',rich_text=True)[RS],load_workbook(QA/'criterion/criterion-candidate.xlsx',rich_text=True)[CS]
    import re
    for start,end,count in ((789,792,8),(809,811,8),(1208,1213,10)):
        text='\n'.join(str(criterion[f'K{r}'].value or '') for r in range(start,end+1))
        found={int(n) for n in re.findall(r'조건\s*\[(\d+)\]',text)}
        assert found == set(range(1,count+1)), (start,found)
        assert str(report[f'M{start}'].value).strip()=='통과'
    assert '26-05-02' in str(criterion['K810'].value)
    assert '26-05-29' in str(report['L810'].value)
    assert criterion['C19'].value == report['C19'].value
    assert not any(c.fill.fgColor.rgb=='FFFFF2CC' for row in criterion for c in row)
    put('checks.json',dict(commonColumns='A:K',commonValuesAndFormattingVerified=True,
        testIDs=115,rows=1734,criterionColumns=11,reviewYellowRemoved=True,
        conditionCounts=[8,8,10],reportResultsUnchanged=True,QCCMarkerUnchanged=True,
        protectedDocuments=len(baseline['protected']),protectedSourceFiles=len(baseline['source']),
        reportHash=sha(QA/'report.xlsx'),criterionHash=sha(QA/'criterion/criterion-candidate.xlsx')))
    print(json.dumps(get('checks.json'),ensure_ascii=False,indent=2))


def publish():
    checks()
    assert (QA/'native-review-checks.json').exists(), 'Native Excel QA not yet performed'
    native_checks=json.loads((QA/'native-review-checks.json').read_text(encoding='utf-8-sig'))
    assert all(p['needed'] <= p['available']+1 for p in native_checks), 'Native text clipping remains'
    paths=[(REPORT,QA/'report.xlsx'),(CRITERION,QA/'criterion/criterion-candidate.xlsx')]
    handles=[]
    writing=False
    try:
        for target,_ in paths:
            handles.append(target.open('r+b'))
        baseline=get('baseline.json')
        assert sha(REPORT)==baseline['reportHash'] and sha(CRITERION)==baseline['criterionHash']
        writing=True
        for handle,(_,candidate) in zip(handles,paths):
            handle.seek(0)
            handle.write(candidate.read_bytes())
            handle.truncate()
            handle.flush()
    except Exception:
        if writing:
            for handle,backup in zip(handles,(QA/'report-before.xlsx',QA/'criterion-before.xlsx')):
                handle.seek(0)
                handle.write(backup.read_bytes())
                handle.truncate()
        raise
    finally:
        for handle in handles:
            handle.close()
    for target,candidate in paths:
        assert sha(target)==sha(candidate)
    put('published.json',dict(reportHash=sha(REPORT),criterionHash=sha(CRITERION)))
    print('Published reviewed report and synchronized criterion under existing names')


if __name__=='__main__':
    {'prepare':prepare,'report':build_report,'prepare-criterion':prepare_criterion,
     'criterion':build_criterion,'fit-criterion':fit_criterion,'reset-heights':reset_height_fixes,
     'checks':checks,'publish':publish}[sys.argv[1]]()
