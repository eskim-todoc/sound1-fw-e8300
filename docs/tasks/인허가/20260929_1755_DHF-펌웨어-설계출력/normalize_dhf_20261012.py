"""Inspect and maintain the current DHF documents; evidence stays local."""
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
import json, hashlib, re, sys, shutil
from copy import deepcopy
from difflib import SequenceMatcher
from zipfile import ZIP_DEFLATED
import openpyxl
from urllib.parse import unquote, quote
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter, range_boundaries
from review_srs_materials import workbook, visible

ROOT = Path(__file__).resolve().parents[4]
TASK = Path(__file__).resolve().parent
QA = TASK / 'qa/dhf-materials-system-20261012'
NEW = Path('G:/공유 드라이브/20_프로젝트별/01_project sound1/08_DHF')
OLD = ROOT / 'docs/참고/인허가/사운드1/DHF'
W = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS={'s':S,**W}
q=lambda n:'{'+S+'}'+n
w=lambda n:'{'+W['w']+'}'+n
PREFIXES={'[QCC 자료 입력]':'C00000','[QCC/E8300 자료 입력]':'7030A0','[E8300 보완]':'0070C0'}
PREF=re.compile('|'.join(re.escape(x) for x in PREFIXES))

def package(p):
    with ZipFile(p) as z:return {n:z.read(n) for n in z.namelist()}

def save_package(p,pkg):
    with ZipFile(p,'w',ZIP_DEFLATED) as z:
        for n,v in pkg.items():z.writestr(n,v)

def padded(t):return '\n'+t.strip()+'\n' if t.strip() else t

def current_paragraph(p):return ''.join(t.text or '' for t in p.iter(w('t')) if visible(t))

def normalized(t):
    for a,b in {'[자료 추가 위치]':'[QCC 자료 입력]', '(QCC 검증 결과 필요)':'[QCC 자료 입력]',
                '[E8300 반영]':'[E8300 보완]','[E8300 확인]':'[E8300 보완]',
                'QCC 자료 첨부 대기':'[QCC 자료 입력] 자료 첨부 대기'}.items():t=t.replace(a,b)
    t=re.sub(r'\[(?:E8300/QCC 각각 확인|E8300/QCC 연계|QCC 단독)\]\s*','',t)
    return t

def plan():
    paths=[p for p in NEW.rglob('*') if p.is_file() and p.name.startswith(('[E8300]','[E8300-QCC]'))]
    baseline=[];books=[];docs=[]
    for p in paths:
        code=re.search(r'\] ([A-Z]+)-',p.name)[1]
        before=QA/(code+'-before'+p.suffix)
        if before.exists():assert sha(before)==sha(p),('Source changed since snapshot',p)
        else:shutil.copy2(p,before)
        baseline.append({'path':str(p),'before':str(before),'hash':sha(p),'code':code})
    srs=workbook(next(Path(x['before']) for x in baseline if x['code']=='SRS'))
    cells=srs['소프트웨어 입력 요구사항 사양서']['cells']
    req={}
    for a,t in cells.items():
        if a.startswith('G') and re.fullmatch(r'SW-FW-RID-\d+',t.strip()):
            req[int(t.strip().rsplit('-',1)[1])]=re.split(r'\n\s*\[(?:QCC|E8300)',cells.get('F'+a[1:],''))[0].strip()
    assert len(req)==81,len(req)
    sysbook=workbook(next(Path(x['before']) for x in baseline if x['code']=='SSTR'))
    syscells=sysbook['보고서_시스템 시험']['cells'];test_req={}
    for a,t in syscells.items():
        if a.startswith('B') and re.fullmatch(r'SW-FW-SYS-Test-ID-\d+',t.strip()):
            ids=re.findall(r'SW-FW-RID-(\d+)',syscells.get('E'+a[1:],''))
            test_req[int(t.strip().rsplit('-',1)[1])]=[int(i) for i in ids]
    def reqtext(test):
        return '\n'.join('SW-FW-RID-'+str(r)+': '+req[r] for r in test_req[test])
    for item in baseline:
        code=item['code'];before=Path(item['before']);candidate=QA/(code+'-candidate'+before.suffix)
        if before.suffix=='.docx':
            pkg=package(before);r=E.fromstring(pkg['word/document.xml']);changes=[]
            for i,p in enumerate(r.iter(w('p'))):
                if not visible(p):continue
                text=current_paragraph(p)
                after=normalized(text)
                # Only current missing-material notices, never historical deletions or factual mentions.
                if PREF.search(after) or after!=text:
                    changes.append({'p':i,'before':text,'after':after})
            if changes:docs.append({**item,'candidate':str(candidate),'edits':changes})
            continue
        wb=openpyxl.load_workbook(before,rich_text=True)
        cached=workbook(before);edits=[];style_cells=[];border_cells=[]
        for sh in wb:
            if sh.title=='표지':continue
            ismain=sh.title.startswith(('기준서_','보고서_'))
            cyber=sh.title=='사이버보안 시험 연계'
            for row in sh:
                for c in row:
                    if c.value is None or isinstance(c,openpyxl.cell.cell.MergedCell):continue
                    if c.data_type=='f':
                        style_cells.append({'sheet':sh.title,'cell':c.coordinate,'horizontal':'center','vertical':'center','notice':False})
                        continue
                    if not isinstance(c.value,(str,openpyxl.cell.rich_text.CellRichText)):continue
                    text=str(c.value);after=normalized(text)
                    if after.strip()=='QCC 펌웨어 자료':after='[QCC 자료 입력]'
                    if code in ('SSTC','SSTR'):
                        if ismain:
                            if c.coordinate=='K10':after='파란색 굵은 글씨: (사진필요) 사진 첨부 위치.\n[QCC 자료 입력] 빨간색, 기본+2pt, 밑줄.\n[QCC/E8300 자료 입력] 보라색, 기본+2pt, 밑줄.\n[E8300 보완] 파란색, 기본+2pt, 밑줄.'
                            if c.coordinate=='C168':after=req[44]
                            if c.coordinate=='C178':after=req[45]
                            # Incorporate the source requirement immediately next to each cited test.
                            refs=re.findall(r'SW-FW-SYS-Test-ID-(\d+)',after)
                            if c.column in (10,11) and refs and c.row in list(range(86,93))+list(range(168,184))+list(range(249,259)):
                                k=int(refs[0]);info=reqtext(k)
                                if '요구사항 내용:' not in after:
                                    lines=after.strip().split('\n');lines.insert(1,'요구사항 내용:\n'+info);after='\n'.join(lines)
                            if c.coordinate in ('I86','I168','I178','I249','J259'):
                                tests=list(range(1,8)) if c.coordinate=='I86' else list(range(35,41)) if c.coordinate=='I178' else list(range(25,35))
                                names=[]
                                for k in tests:
                                    rr=next(a[1:] for a,t in syscells.items() if a.startswith('B') and t.strip()=='SW-FW-SYS-Test-ID-'+str(k))
                                    names.append(normalized(syscells['C'+rr]).strip().rstrip('.'))
                                after=after.strip()+'\n확인할 기능: '+', '.join(names)+'.'
                            if c.coordinate=='K216':after=after.replace('코드 대조:','[E8300 보완]')
                            if c.coordinate=='K229':after=after.replace('추가 근거:','[QCC/E8300 자료 입력]')
                        # Time-related vectors, not document approvals or boundary-calendar cases.
                        result=(ismain and c.column==11) or (sh.title=='시험 보조자료' and c.column==3) or (cyber and c.column>=11)
                        day='29' if result else '2'
                        after=re.sub(r'26,\s*10,\s*2(?=,)', '26,5,'+day,after)
                        after=after.replace('2026-10-02','2026-05-'+('29' if result else '02'))
                        after=after.replace('2026년 10월 2일','2026년 5월 '+('29' if result else '2')+'일')
                        if cyber and c.row in (27,28,33) and c.column>=11:
                            after=after.replace('[QCC 자료 입력]','[QCC/E8300 자료 입력]')
                        if cyber and c.column>=11 and c.row>=5:
                            qcc_rows={5,7,8,9,10,11,13,15,20,22,31,38}
                            joint_rows={6,12,14,21,25,26,27,28,33,35,37}
                            prefix='[QCC 자료 입력]' if c.row in qcc_rows else '[QCC/E8300 자료 입력]' if c.row in joint_rows else '[E8300 보완]'
                            after=after.replace('추가 근거:',prefix)
                            after=after.replace('보완 필요:','[E8300 보완]')
                            if '추가 근거 정리' in after and not PREF.search(after):after=after.replace('추가 근거 정리',prefix+' 근거 자료 정리')
                        if cyber and c.coordinate=='A3' and 'QCC' in after and '자료' in after and not PREF.search(after):
                            after='[QCC 자료 입력] '+after.strip()
                    # Unit long-text continuation cells already have only outer padding.
                    if after.strip() and code not in ('SUVC','SUVR'):after=padded(after)
                    if text!=after:edits.append({'sheet':sh.title,'cell':c.coordinate,'before':text,'after':after})
                    # Headers and short identifiers/methods/verdicts are centered; descriptions stay left.
                    head=(ismain and c.row in ({'SUVC':[28],'SUVR':[28],'SITC':[23,24],'SITR':[23,24],'SSTC':[32,33],'SSTR':[32,33]}.get(code,[]))) or (cyber and c.row==4)
                    short=bool(re.fullmatch(r'(?:[IADT](?:\s*,\s*[IADT])*|통과|실패|적합|부적합|\d+|(?:SW-FW-(?:RID|SYS-Test-ID|UnitTest-\w+-ID|IntegrationTest-ID)|SYS-RID-SP-FW)-\d+)',after.strip()))
                    if code=='SRS' and c.row<=5:head=True
                    center_cols={'SUVC':{2,3,4,6,7},'SUVR':{2,3,4,6,7,13},'SITC':{2,3,4,9},'SITR':{2,3,4,9,15},'SSTC':{2,4,5,6},'SSTR':{2,4,5,6,12}}
                    body_center=ismain and c.row>={'SUVC':28,'SUVR':28,'SITC':23,'SITR':23,'SSTC':32,'SSTR':32}.get(code,99999) and c.column in center_cols.get(code,set())
                    if code=='SRS' and c.column in (2,3,4,5,7,8,9):body_center=True
                    horiz='center' if head or short or body_center else 'left'
                    vert='center' if head or short or body_center else 'top'
                    style_cells.append({'sheet':sh.title,'cell':c.coordinate,'horizontal':horiz,'vertical':vert,'notice':bool(PREF.search(after))})
            # Repair visible table edges, retaining intentionally white/no inner continuation lines.
            start={'SRS':4,'SUVC':28,'SUVR':28,'SITC':23,'SITR':23,'SSTC':32,'SSTR':32}.get(code,99999)
            if cyber:start=4
            if ismain or cyber or code=='SRS':
                merges={}
                for area in sh.merged_cells.ranges:
                    for rr in range(area.min_row,area.max_row+1):
                        for cc in range(area.min_col,area.max_col+1):merges[(rr,cc)]=area
                for row in sh.iter_rows(min_row=start):
                    for c in row:
                        if c.column<(1 if cyber else 2):continue
                        if c.value is None or isinstance(c,openpyxl.cell.cell.MergedCell):continue
                        area=merges.get((c.row,c.column));lo,hi=c.column,c.column;top,bot=c.row,c.row
                        if area:lo,top,hi,bot=area.bounds
                        for rr in range(top,bot+1):
                            for cc,edge in [(lo,'left'),(hi,'right')]:
                                border_cells.append({'sheet':sh.title,'cell':get_column_letter(cc)+str(rr),'edge':edge})
                        for rr,edge in [(top,'top'),(bot,'bottom')]:
                            side=getattr(c.border,edge)
                            # Missing internal edges often deliberately join split long text; preserve them.
                            if head or area or (side and side.style and not(side.color and side.color.type=='rgb' and side.color.rgb[-6:]=='FFFFFF')):
                                for cc in range(lo,hi+1):border_cells.append({'sheet':sh.title,'cell':get_column_letter(cc)+str(rr),'edge':edge})
        # Support-sheet copies of the same system condition must not remain abbreviated.
        if code=='SSTR':
            by={(e['sheet'],e['cell']):e for e in edits}
            main='보고서_시스템 시험';helper='시험 보조자료'
            replacements={syscells[e['cell']].strip():e['after'].strip() for e in edits if e['sheet']==main and e['cell'] in syscells and '요구사항 내용:' in e['after']}
            for a,t in cached.get(helper,{}).get('cells',{}).items():
                if t.strip() in replacements:
                    key=(helper,a);e={'sheet':helper,'cell':a,'before':t,'after':padded(replacements[t.strip()])}
                    if key in by:by[key].update(e)
                    else:edits.append(e)
        books.append({**item,'candidate':str(candidate),'artifact':str(QA/(code+'-artifact.xlsx')),'main':next(s for s in cached if s!='표지'),'edits':edits,'styles':style_cells,'borders':border_cells})
    dump('baseline.json',baseline);dump('plan.json',{'books':books,'docs':docs,'requirements':req})
    print(json.dumps({'books':[(p['code'],len(p['edits']),len(p['styles'])) for p in books],'docs':[(p['code'],len(p['edits'])) for p in docs]},ensure_ascii=False))

_cellcache={}
def xmlcell(root,addr,create=False):
    cache=_cellcache.setdefault(root,{c.get('r'):c for c in root.iter(q('c'))}) if root not in _cellcache else _cellcache[root]
    found=cache.get(addr)
    if found is not None or not create:return found
    col,row=coordinate_from_string(addr);r=root.find('.//s:row[@r="'+str(row)+'"]',NS)
    if r is None:
        r=E.SubElement(root.find(q('sheetData')),q('row'),r=str(row))
    c=E.Element(q('c'),r=addr)
    for i,other in enumerate(r):
        if column_index_from_string(coordinate_from_string(other.get('r'))[0])>column_index_from_string(col):r.insert(i,c);break
    else:r.append(c)
    cache[addr]=c
    return c

_sharedcache={}
def string_runs(c,pkg):
    if c.get('t')=='s':
        key=id(pkg)
        if key not in _sharedcache:_sharedcache[key]=(pkg,E.fromstring(pkg['xl/sharedStrings.xml']))
        el=_sharedcache[key][1][int(c.findtext(q('v')))]
    elif c.get('t')=='inlineStr':el=c.find(q('is'))
    else:return [(c.findtext(q('v'),''),None)]
    return [(''.join(x.itertext()) if x.tag==q('t') else ''.join(x.xpath('.//s:t/text()',namespaces=NS)),x.find(q('rPr')) if x.tag==q('r') else None) for x in el if x.tag in (q('t'),q('r'))]

def rich_string(c,pkg,before,after,base_size=10,force_notice=None):
    old=[]
    for t,props in string_runs(c,pkg):old.extend([E.tostring(props) if props is not None else None]*len(t))
    assert len(old)==len(before),(c.get('r'),len(old),len(before))
    props_after=[None]*len(after)
    for tag,a,b,aa,bb in SequenceMatcher(None,before,after,autojunk=False).get_opcodes():
        if tag=='equal':props_after[aa:bb]=old[a:b]
    # Obsolete responsibility headings are functional text, not missing evidence notices.
    if re.search(r'\[(?:E8300/QCC|QCC 단독)',before) and not PREF.search(after):props_after=[None]*len(after)
    notices=list(PREF.finditer(after))
    spans=[]
    for match in notices:
        stop=after.find('\n\n',match.end());stop=len(after.rstrip()) if stop<0 else stop
        nxt=next((x.start() for x in notices if x.start()>match.start()),stop)
        spans.append((match.start(),min(stop,nxt),match.group()))
    if force_notice and not notices:spans.append((0,len(after),force_notice))
    for start,stop,prefix in spans:
        rp=E.Element(q('rPr'));E.SubElement(rp,q('rFont'),val='맑은 고딕');E.SubElement(rp,q('sz'),val=str(base_size+2));E.SubElement(rp,q('color'),rgb='FF'+PREFIXES[prefix]);E.SubElement(rp,q('u'),val='single')
        props_after[start:stop]=[E.tostring(rp)]*(stop-start)
    result=E.Element(q('is'));i=0
    while i<len(after):
        end=i+1
        while end<len(after) and props_after[end]==props_after[i]:end+=1
        run=E.SubElement(result,q('r'))
        if props_after[i]:run.append(E.fromstring(props_after[i]))
        E.SubElement(run,q('t'),{'{http://www.w3.org/XML/1998/namespace}space':'preserve'}).text=after[i:end]
        i=end
    return result

def merge_books():
    plan=json.loads((QA/'plan.json').read_text(encoding='utf8'))
    counts=[]
    for p in plan['books']:
        pkg=package(p['before']);source=workbook(p['before']);authored=workbook(p['artifact'])
        roots={name:E.fromstring(pkg[d['part']]) for name,d in source.items() if name!='표지'}
        styles=E.fromstring(pkg['xl/styles.xml']);xfs=styles.find(q('cellXfs'));borders=styles.find(q('borders'))
        sf={(e['sheet'],e['cell']):e for e in p['styles']}
        edits={(e['sheet'],e['cell']):e for e in p['edits']}
        for key,e in edits.items():assert authored[e['sheet']]['cells'][e['cell']]==e['after'],(p['code'],key)
        for key,e in sf.items():
            if e['notice'] and key not in edits:
                t=source[e['sheet']]['cells'][e['cell']];edits[key]={**e,'before':t,'after':t}
        if p['code'] in ('SUVC','SUVR'):
            key=(p['main'],'D19');t=source[p['main']]['cells'].get('D19','')
            if t:edits[key]={**edits.get(key,{}),'sheet':key[0],'cell':key[1],'before':t,'after':edits.get(key,{}).get('after',t),'force_notice':'[QCC 자료 입력]'}
        for (sn,addr),e in edits.items():
            c=xmlcell(roots[sn],addr);assert c.find(q('f')) is None
            rich=rich_string(c,pkg,e['before'],e['after'],force_notice=e.get('force_notice'))
            for child in list(c):c.remove(child)
            c.set('t','inlineStr');c.append(rich)
        border_changes={}
        for e in p['borders']:border_changes.setdefault((e['sheet'],e['cell']),set()).add(e['edge'])
        cache={};bcache={};updated=0
        for key in set(sf)|set(border_changes):
            sn,addr=key;c=xmlcell(roots[sn],addr,True);styleid=int(c.get('s','0'));xf=deepcopy(xfs[styleid])
            if key in sf:
                e=sf[key];a=xf.find(q('alignment'))
                if a is None:a=E.SubElement(xf,q('alignment'))
                a.set('horizontal',e['horizontal']);a.set('vertical',e['vertical']);a.set('wrapText','1');a.attrib.pop('shrinkToFit',None);xf.set('applyAlignment','1')
            if key in border_changes:
                border=deepcopy(borders[int(xf.get('borderId','0'))])
                for edge in border_changes[key]:
                    side=border.find(q(edge))
                    if side is None:
                        side=E.Element(q(edge));order=['left','right','top','bottom','diagonal','vertical','horizontal'];idx=order.index(edge)
                        at=next((i for i,x in enumerate(border) if x.tag.split('}')[-1] in order and order.index(x.tag.split('}')[-1])>idx),len(border));border.insert(at,side)
                    side.set('style','thin')
                    for ch in list(side):side.remove(ch)
                    E.SubElement(side,q('color'),rgb='FF000000')
                raw=E.tostring(border)
                if raw not in bcache:bcache[raw]=len(borders);borders.append(border)
                xf.set('borderId',str(bcache[raw]));xf.set('applyBorder','1')
            raw=E.tostring(xf)
            if raw not in cache:cache[raw]=len(xfs);xfs.append(xf)
            if E.tostring(xfs[styleid])!=raw:c.set('s',str(cache[raw]));updated+=1
        xfs.set('count',str(len(xfs)));borders.set('count',str(len(borders)))
        pkg['xl/styles.xml']=E.tostring(styles,encoding='UTF-8',xml_declaration=True,standalone=True)
        for sn,root in roots.items():pkg[source[sn]['part']]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
        save_package(p['candidate'],pkg)
        counts.append({'code':p['code'],'text':len(p['edits']),'rich_notice_cells':len(edits)-len(p['edits']),'style_cells':updated})
    dump('merge-counts.json',counts);print(counts)

def merge_docs():
    plan=json.loads((QA/'plan.json').read_text(encoding='utf8'));stats=[]
    for p in plan['docs']:
        pkg=package(p['before']);root=E.fromstring(pkg['word/document.xml']);paras=list(root.iter(w('p')))
        ids=[int(x) for x in root.xpath('//@w:id',namespaces=W) if x.isdigit()];serial=max(ids+[0])+1
        for edit in p['edits']:
            para=paras[edit['p']];assert current_paragraph(para)==edit['before']
            if edit['before']!=edit['after']:
                # Preserve all older tracked deletions; record this replacement separately.
                old_runs=[x for x in para.iter(w('t')) if visible(x)]
                for x in old_runs:x.text=''
                for tag,txt in [('del',edit['before']),('ins',edit['after'])]:
                    change=E.SubElement(para,w(tag),{w('id'):str(serial),w('author'):'김은수',w('date'):'2026-10-12T00:00:00Z'});serial+=1
                    run=E.SubElement(change,w('r'));E.SubElement(run,w('delText' if tag=='del' else 't'),{'{http://www.w3.org/XML/1998/namespace}space':'preserve'}).text=txt
            # Each affected paragraph's notice is styled; ordinary introduction stays intact.
            started=False;prefix=None
            for run in list(para.iter(w('r'))):
                if not visible(run):continue
                text=''.join(x.text or '' for x in run.iter(w('t')))
                found=PREF.search(text)
                if found:started=True;prefix=found.group()
                if not started or not text:continue
                # Existing notices are whole runs or standalone paragraphs in these documents.
                props=run.find(w('rPr'))
                if props is None:props=E.Element(w('rPr'));run.insert(0,props)
                old=deepcopy(props)
                for name,val in [('color',PREFIXES[prefix]),('sz','24'),('szCs','24'),('u','single')]:
                    x=props.find(w(name))
                    if x is None:x=E.SubElement(props,w(name))
                    x.set(w('val'),val)
                if E.tostring(old)!=E.tostring(props):
                    prior=props.find(w('rPrChange'))
                    if prior is None:
                        pc=E.SubElement(props,w('rPrChange'),{w('id'):str(serial),w('author'):'김은수',w('date'):'2026-10-12T00:00:00Z'});serial+=1;pc.append(old)
            assert current_paragraph(para)==edit['after']
        pkg['word/document.xml']=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
        save_package(p['candidate'],pkg);stats.append({'code':p['code'],'notices':len(p['edits']),'text_changes':sum(e['before']!=e['after'] for e in p['edits'])})
    dump('doc-changes.json',stats);print(stats)

def verify():
    plan=json.loads((QA/'plan.json').read_text(encoding='utf8'));checks=[];allbooks={}
    for p in plan['books']:
        old=package(p['before']);new=package(p['candidate']);a=workbook(p['before']);b=workbook(p['candidate']);allbooks[p['code']]=b
        assert set(old)==set(new)
        allowed={'xl/styles.xml'}|{d['part'] for s,d in a.items() if s!='표지'}
        assert all(old[n]==new[n] for n in old if n not in allowed),('Unrelated part',p['code'])
        edits={(e['sheet'],e['cell']):e for e in p['edits']};changed=[]
        for sn in a:
            assert a[sn]['merges']==b[sn]['merges'],('merges',p['code'],sn)
            for addr in set(a[sn]['cells'])|set(b[sn]['cells']):
                oldval=a[sn]['cells'].get(addr);val=b[sn]['cells'].get(addr)
                if oldval!=val:
                    assert (sn,addr) in edits,('Unexpected value change',p['code'],sn,addr)
                    assert val==edits[sn,addr]['after'];changed.append((sn,addr))
            ra,rb=E.fromstring(old[a[sn]['part']]),E.fromstring(new[b[sn]['part']])
            fa={c.get('r'):(E.tostring(c.find(q('f'))),c.findtext(q('v'))) for c in ra.iter(q('c')) if c.find(q('f')) is not None}
            fb={c.get('r'):(E.tostring(c.find(q('f'))),c.findtext(q('v'))) for c in rb.iter(q('c')) if c.find(q('f')) is not None}
            assert fa==fb,('Formula/cache changed',p['code'],sn)
            for row in rb.iter(q('row')):assert float(row.get('ht','15'))<=409.5
            if sn!='표지':
                for addr,val in b[sn]['cells'].items():
                    assert not re.search(r'\[(?:자료 추가 위치|QCC 단독|E8300/QCC 각각 확인|E8300/QCC 연계)\]|\(QCC 검증 결과 필요\)',val),(p['code'],sn,addr,val)
        wb=openpyxl.load_workbook(p['candidate'],rich_text=True)
        for e in p['styles']:
            c=wb[e['sheet']][e['cell']]
            assert c.alignment.horizontal==e['horizontal'] and c.alignment.vertical==e['vertical'],('alignment',p['code'],e)
        sty=E.fromstring(new['xl/styles.xml']);xf=sty.find(q('cellXfs'));borderdefs=sty.find(q('borders'))
        roots={s:E.fromstring(new[d['part']]) for s,d in b.items()}
        for e in p['borders']:
            # Check raw style for merged subcells, since openpyxl synthesizes merge borders.
            c=xmlcell(roots[e['sheet']],e['cell']);bd=borderdefs[int(xf[int(c.get('s','0'))].get('borderId','0'))]
            side=bd.find(q(e['edge']))
            assert side is not None and side.get('style')=='thin' and side.find(q('color')).get('rgb')=='FF000000',('border',p['code'],e)
        checks.append({'code':p['code'],'changed_values':len(changed),'alignment_cells':len(p['styles']),'merges_preserved':True,'formulas_and_caches_preserved':True})
    for crit,rep,limit,start in [('SUVC','SUVR',11,28),('SITC','SITR',13,23),('SSTC','SSTR',10,32)]:
        pa=next(p for p in plan['books'] if p['code']==crit);pb=next(p for p in plan['books'] if p['code']==rep)
        aa=allbooks[crit][pa['main']]['cells'];bb=allbooks[rep][pb['main']]['cells']
        for addr in set(aa)|set(bb):
            col,row=coordinate_from_string(addr)
            if row<start or column_index_from_string(col)>limit:continue
            assert aa.get(addr,'').strip()==bb.get(addr,'').strip(),('Common content',crit,rep,addr,aa.get(addr),bb.get(addr))
    a=allbooks['SSTC']['기준서_시스템 시험']['cells'];b=allbooks['SSTR']['보고서_시스템 시험']['cells']
    assert '[26,5,2,13,14,15]' in a['J218'] and '[26,5,29,13,14,15]' in b['K218']
    assert len([t for a,t in b.items() if a.startswith('B') and re.fullmatch(r'SW-FW-SYS-Test-ID-\d+',t.strip())])==81
    assert set(int(n) for a,t in b.items() if a.startswith('E') for n in re.findall(r'SW-FW-RID-(\d+)',t))==set(range(1,82))
    for p in plan['docs']:
        a,b=package(p['before']),package(p['candidate']);assert set(a)==set(b)
        assert all(a[n]==b[n] for n in a if n!='word/document.xml')
        ra,rb=E.fromstring(a['word/document.xml']),E.fromstring(b['word/document.xml'])
        pa,pb=list(ra.iter(w('p'))),list(rb.iter(w('p')));assert len(pa)==len(pb)
        changes={e['p']:e for e in p['edits']}
        for i,(x,y) in enumerate(zip(pa,pb)):
            assert current_paragraph(y)==(changes[i]['after'] if i in changes else current_paragraph(x)),('Word text',p['code'],i)
        olddel=ra.xpath('//w:delText/text()',namespaces=W);newdel=rb.xpath('//w:delText/text()',namespaces=W)
        for t in olddel:assert t in newdel,('Tracked deletion lost',p['code'])
        checks.append({'code':p['code'],'notices':len(p['edits']),'old_revisions_preserved':True})
    dump('structural-checks.json',checks);print(checks)

def publish():
    verify()
    plan=json.loads((QA/'plan.json').read_text(encoding='utf8'))
    native=json.loads((QA/'native-check.json').read_text(encoding='utf-8-sig'))
    assert len(native)>3000 and all(x['fits'] and not x['limit'] for x in native),'Native layout is not ready'
    output=[]
    for p in plan['books']+plan['docs']:
        dest=Path(p['path']);assert sha(dest)==p['hash'],('User file changed during work',dest)
    for p in plan['books']+plan['docs']:
        dest=Path(p['path']);tmp=dest.with_name('.'+dest.name+'.writing')
        assert not tmp.exists(),tmp
        shutil.copy2(p['candidate'],tmp)
        assert sha(tmp)==sha(Path(p['candidate']))
        tmp.replace(dest)
        output.append({'path':str(dest),'hash':sha(dest),'code':p['code']})
    dump('published.json',output);print('Published:',len(output))

def heights():
    plan=json.loads((QA/'plan.json').read_text(encoding='utf8'));measure=json.loads((QA/'native-check.json').read_text(encoding='utf-8-sig'))
    changed=[]
    pairs={'SUVC':'SUVR','SUVR':'SUVC','SITC':'SITR','SITR':'SITC','SSTC':'SSTR','SSTR':'SSTC'}
    for p in plan['books']:
        pkg=package(p['candidate']);book=workbook(p['candidate'])
        for sn,d in book.items():
            rows={}
            for x in measure:
                same=x['code']==p['code'] and x['sheet']==sn
                common=sn==p['main'] and x['code']==pairs.get(p['code']) and x['sheet'].startswith(('기준서_','보고서_'))
                cyber=sn=='사이버보안 시험 연계' and x['sheet']==sn and x['code']==pairs.get(p['code'])
                if not(same or common or cyber):continue
                assert not x['limit'],('Needs content split before height adjustment',x)
                for r in range(x['start_row'],x['end_row']+1):rows[r]=max(rows.get(r,0),min(409.5,x['row_needed']))
            if not rows:continue
            root=E.fromstring(pkg[d['part']])
            for row in root.iter(q('row')):
                r=int(row.get('r'));need=rows.get(r,0)
                if need>float(row.get('ht','15')):row.set('ht',str(need));row.set('customHeight','1');changed.append((p['code'],sn,r,need))
            pkg[d['part']]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
        save_package(p['candidate'],pkg)
    dump('row-height-changes.json',changed);print('Adjusted rows',len(changed))

def cleanup_plan():
    inventory=json.loads((QA/'baseline.json').read_text(encoding='utf8'));actions=[]
    oldfiles=[p for p in OLD.rglob('*') if p.is_file()]
    for x in inventory:
        new=Path(x['path'])
        for old in [p for p in oldfiles if p.name==new.name]:
            if not old.is_file() or '[작성자 전용]' in old.name:continue
            dest=QA/'정리 전 로컬 원본'/old.relative_to(OLD)
            actions.append({'kind':'local-document','source':str(old),'destination':str(dest),'source_hash':sha(old),'current':str(new),'same_as_current_start':sha(old)==x['hash']})
    essential=NEW/'03_설계 출력/SW/외부기/EPTC-TD24A 필수성능시험 기준서.docx'
    if essential.exists():
        for old in OLD.rglob(essential.name):
            if sha(old)==sha(essential):actions.append({'kind':'local-document','source':str(old),'destination':str(QA/'정리 전 로컬 원본'/old.relative_to(OLD)),'source_hash':sha(old),'current':str(essential),'same_as_current_start':True})
    for new in NEW.rglob('*'):
        if not new.is_file() or not any(k in new.parts for k in ('이미지 참조','시스템 Data','유닛 시험 예상 입출력')):continue
        suffix=Path(str(new.relative_to(NEW)).replace('SW\\외부기','SW\\펌웨어'))
        old=OLD/suffix
        assert old.is_file() and sha(old)==sha(new),('Supplement needs individual review',new,old)
        actions.append({'kind':'new-supplement','source':str(new),'destination':str(QA/'정리 전 G 보조자료'/new.relative_to(NEW)),'source_hash':sha(new),'current':str(old),'same_as_current_start':True})
    dump('cleanup-plan.json',actions);print('Recoverable moves:',len(actions),'documents:',sum(x['kind']=='local-document' for x in actions),'supplements:',sum(x['kind']=='new-supplement' for x in actions))

def links():
    actions=json.loads((QA/'cleanup-plan.json').read_text(encoding='utf8'))
    targets={str(Path(x['source']).resolve()).lower():Path(x['current']) for x in actions if x['kind']=='local-document'}
    changes=[]
    for p in (ROOT/'docs/참고/인허가/사운드1').rglob('*.md'):
        if p.name.startswith('[작성자 전용]'):continue
        old=p.read_text(encoding='utf8')
        def replace(m):
            address=unquote(m.group(1).strip('<>'))
            if '://' in address or address.startswith('#'):return m.group()
            try:resolved=str((p.parent/address).resolve()).lower()
            except ValueError:return m.group()
            if resolved not in targets:return m.group()
            target=targets[resolved];assert target.is_file(),target
            changes.append({'md':str(p),'old':m.group(1),'new':str(target)})
            return ']('+quote(target.as_posix(),safe='/:')+')'
        new=re.sub(r'\]\(([^)]+)\)',replace,old)
        if old!=new:p.write_text(new,encoding='utf8')
    dump('link-changes.json',changes);print('Current links updated:',len(changes))

def support_plan():
    master=json.loads((QA/'plan.json').read_text(encoding='utf8'));p=next(x for x in master['books'] if x['code']=='SSTR')
    source=Path(p['path']);before=QA/'SSTR-support-before.xlsx'
    if not before.exists():shutil.copy2(source,before)
    assert sha(before)==sha(source)
    data=workbook(before);main=data[p['main']]['cells'];edits=[]
    for a,t in data['시험 보조자료']['cells'].items():
        if a[0] not in ('B','C') or '요구사항 내용:' in t:continue
        match=re.search(r'(?:조건|검증) \[\d+\][^\n]*SW-FW-SYS-Test-ID-(\d+)',t)
        if not match:continue
        test=int(match[1]);row=next(k[1:] for k,v in main.items() if k.startswith('B') and v.strip()=='SW-FW-SYS-Test-ID-'+str(test))
        ids=re.findall(r'SW-FW-RID-(\d+)',main['E'+row]);info='\n'.join('SW-FW-RID-'+i+': '+master['requirements'][i] for i in ids)
        lines=t.strip().split('\n');lines.insert(1,'요구사항 내용:\n'+info)
        edits.append({'sheet':'시험 보조자료','cell':a,'before':t,'after':padded('\n'.join(lines))})
    sn='사이버보안 시험 연계';a='K20';t=data[sn]['cells'][a]
    edits.append({'sheet':sn,'cell':a,'before':t,'after':t.replace('기존 시험 60: 해당 시험 통과 확인.','기존 시험 60: CRC 기능이 있는 Bluetooth LE 통신으로 리모콘 제어/조회가 정상 동작함을 확인.')})
    layout=[];edges=[]
    for row in range(5,299):
        for col in 'ABCD':
            address=col+str(row)
            layout.append({'sheet':'시험 보조자료','cell':address,'horizontal':'center' if col in 'AD' or row==5 else 'left','vertical':'center' if col in 'AD' or row==5 else 'top','notice':False})
            edges.extend({'sheet':'시험 보조자료','cell':address,'edge':edge} for edge in ('left','right','top','bottom'))
    delta={**p,'hash':sha(source),'before':str(before),'artifact':str(QA/'SSTR-support-artifact.xlsx'),'candidate':str(QA/'SSTR-support-candidate.xlsx'),'edits':edits,'styles':layout,'borders':edges}
    dump('support-plan.json',{'books':[delta]});print('Support changes',len(edits))

def support_merge():
    p=json.loads((QA/'support-plan.json').read_text(encoding='utf8'))['books'][0]
    pkg=package(p['before']);book=workbook(p['before']);authored=workbook(p['artifact']);roots={}
    for e in p['edits']:
        assert authored[e['sheet']]['cells'][e['cell']]==e['after']
        sn=e['sheet'];part=book[sn]['part']
        if sn not in roots:roots[sn]=E.fromstring(pkg[part])
        c=xmlcell(roots[sn],e['cell']);rich=rich_string(c,pkg,e['before'],e['after'])
        for ch in list(c):c.remove(ch)
        c.set('t','inlineStr');c.append(rich)
    for sn,r in roots.items():pkg[book[sn]['part']]=E.tostring(r,encoding='UTF-8',xml_declaration=True,standalone=True)
    save_package(p['candidate'],pkg)

def support_finish():
    p=json.loads((QA/'support-plan.json').read_text(encoding='utf8'))['books'][0]
    measurements=json.loads((QA/'support-native.json').read_text(encoding='utf-8-sig'))
    pkg=package(p['candidate']);book=workbook(p['candidate'])
    for sn in {e['sheet'] for e in p['edits']}:
        root=E.fromstring(pkg[book[sn]['part']]);rows={}
        for x in measurements:
            if x['sheet']!=sn:continue
            assert not x['limit']
            for r in range(x['start_row'],x['end_row']+1):rows[r]=max(rows.get(r,0),x['row_needed'])
        for row in root.iter(q('row')):
            need=rows.get(int(row.get('r')),0)
            if need>float(row.get('ht','15')):row.set('ht',str(need));row.set('customHeight','1')
        pkg[book[sn]['part']]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    styles=E.fromstring(pkg['xl/styles.xml']);xfs=styles.find(q('cellXfs'));borders=styles.find(q('borders'))
    border=E.Element(q('border'))
    for side in ('left','right','top','bottom'):
        e=E.SubElement(border,q(side),style='thin');E.SubElement(e,q('color'),rgb='FF000000')
    E.SubElement(border,q('diagonal'));bid=len(borders);borders.append(border);borders.set('count',str(len(borders)))
    sn='시험 보조자료';root=E.fromstring(pkg[book[sn]['part']]);stylecache={}
    for edit in p['styles']:
        c=xmlcell(root,edit['cell'],True);xf=deepcopy(xfs[int(c.get('s','0'))]);xf.set('borderId',str(bid));xf.set('applyBorder','1');xf.set('applyAlignment','1')
        al=xf.find(q('alignment'))
        if al is None:al=E.SubElement(xf,q('alignment'))
        al.set('horizontal',edit['horizontal']);al.set('vertical',edit['vertical']);al.set('wrapText','1')
        raw=E.tostring(xf)
        if raw not in stylecache:stylecache[raw]=len(xfs);xfs.append(xf)
        c.set('s',str(stylecache[raw]))
    xfs.set('count',str(len(xfs)));pkg['xl/styles.xml']=E.tostring(styles,encoding='UTF-8',xml_declaration=True,standalone=True)
    pkg[book[sn]['part']]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    save_package(p['candidate'],pkg)
    master=json.loads((QA/'plan.json').read_text(encoding='utf8'));mp=next(x for x in master['books'] if x['code']=='SSTR')
    idx={(e['sheet'],e['cell']):e for e in mp['edits']};original=workbook(mp['before'])
    for e in p['edits']:
        key=(e['sheet'],e['cell'])
        if key in idx:idx[key]['after']=e['after']
        else:mp['edits'].append({**e,'before':original[e['sheet']]['cells'][e['cell']]})
    override={(e['sheet'],e['cell']):e for e in p['styles']}
    mp['styles']=[e for e in mp['styles'] if (e['sheet'],e['cell']) not in override]+p['styles']
    mp['borders']=[e for e in mp['borders'] if e['sheet']!='시험 보조자료']+p['borders']
    shutil.copy2(p['candidate'],mp['candidate']);dump('plan.json',master)
    # Keep cybersecurity common-row geometry in the criteria and report identical.
    criterion=next(x for x in master['books'] if x['code']=='SSTC')
    cp=package(criterion['candidate']);cb=workbook(criterion['candidate']);sn='사이버보안 시험 연계'
    cr=E.fromstring(cp[cb[sn]['part']]);rr=E.fromstring(pkg[book[sn]['part']])
    required={r.get('r'):r.get('ht') for r in rr.iter(q('row'))}
    for row in cr.iter(q('row')):
        need=required.get(row.get('r'))
        if need and float(need)>float(row.get('ht','15')):row.set('ht',need);row.set('customHeight','1')
    cp[cb[sn]['part']]=E.tostring(cr,encoding='UTF-8',xml_declaration=True,standalone=True)
    save_package(criterion['candidate'],cp)

def support_publish():
    verify()
    p=json.loads((QA/'support-plan.json').read_text(encoding='utf8'))['books'][0]
    checks=json.loads((QA/'support-native.json').read_text(encoding='utf-8-sig'));assert all(x['fits'] and not x['limit'] for x in checks)
    published=json.loads((QA/'published.json').read_text(encoding='utf8'))
    master=json.loads((QA/'plan.json').read_text(encoding='utf8'))
    for item in published:
        if item['code'] not in ('SSTC','SSTR'):continue
        candidate=next(x for x in master['books'] if x['code']==item['code'])
        dest=Path(item['path']);assert sha(dest)==item['hash']
        temp=dest.with_name('.'+dest.name+'.writing');assert not temp.exists()
        shutil.copy2(candidate['candidate'],temp);temp.replace(dest);item['hash']=sha(dest)
    dump('published.json',published);print('Support sheet synchronized and published')

def dump(n, x):
    QA.mkdir(parents=True, exist_ok=True)
    (QA/n).write_text(json.dumps(x, ensure_ascii=False, indent=2), encoding='utf8')

def inspect():
    inventory=[]; books={}; docs={}
    for p in NEW.rglob('*'):
        if not p.is_file() or not p.name.startswith(('[E8300]', '[E8300-QCC]')): continue
        inventory.append({'path':str(p),'hash':sha(p),'name':p.name})
        if p.suffix=='.xlsx':
            wb=openpyxl.load_workbook(p,rich_text=True)
            books[p.name]={s.title:{'rows':s.max_row,'cols':s.max_column,
                'cells':{c.coordinate:str(c.value) for row in s for c in row if c.value is not None},
                'merges':[str(m) for m in s.merged_cells.ranges]} for s in wb}
        if p.suffix=='.docx':
            with ZipFile(p) as z: r=E.fromstring(z.read('word/document.xml'))
            docs[p.name]=[{'p':i,'text':''.join(t.text or '' for t in p.xpath('.//w:t[not(ancestor::w:del)]',namespaces=W))} for i,p in enumerate(r.xpath('//w:p',namespaces=W))]
    dump('inventory.json',inventory);dump('books.json',books);dump('docs.json',docs)
    markers={}
    pat=re.compile(r'QCC|E8300.*보완|자료.*(?:입력|추가|보완|필요)|(?:ID|시험)\s*\d+\s*[~～–-]\s*\d+|2026[-./년 ]+10[-./월 ]+0?2')
    for f,sheets in books.items():
        markers[f]={s:{a:t for a,t in d['cells'].items() if pat.search(t)} for s,d in sheets.items()}
    dump('markers.json',markers)
    dump('doc-markers.json',{f:[x for x in ps if pat.search(x['text'])] for f,ps in docs.items()})
    oldfiles=[p for p in OLD.rglob('*') if p.is_file()]
    matches=[]
    for x in inventory:
        for p in oldfiles:
            if p.name==x['name']:matches.append({'old':str(p),'new':x['path'],'same':sha(p)==x['hash']})
    dump('duplicates.json',matches)
    print(json.dumps({'books':{f:{s:(d['rows'],d['cols']) for s,d in ss.items()} for f,ss in books.items()},'documents':len(docs),'duplicates':matches},ensure_ascii=False,indent=2))

if __name__=='__main__':
    if sys.argv[1]=='inspect': inspect()
    elif sys.argv[1]=='plan':plan()
    elif sys.argv[1]=='merge-books':merge_books()
    elif sys.argv[1]=='merge-docs':merge_docs()
    elif sys.argv[1]=='verify':verify()
    elif sys.argv[1]=='heights':heights()
    elif sys.argv[1]=='cleanup-plan':cleanup_plan()
    elif sys.argv[1]=='publish':publish()
    elif sys.argv[1]=='links':links()
    elif sys.argv[1]=='support-plan':support_plan()
    elif sys.argv[1]=='support-merge':support_merge()
    elif sys.argv[1]=='support-finish':support_finish()
    elif sys.argv[1]=='support-publish':support_publish()
