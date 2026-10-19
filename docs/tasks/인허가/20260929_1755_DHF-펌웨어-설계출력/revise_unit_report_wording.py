"""사용자가 편집한 유닛 보고서를 기준으로 문장만 교정한다.

워크북 값 작성은 Artifact Tool을 사용한다. 원본 XML의 글자별 서식,
병합, 행 높이, 그림을 보존하기 위해 검증된 값만 원본 패키지에 결합한다.
실측값, 원시 로그와 판정은 생성하거나 바꾸지 않는다.
"""
from __future__ import annotations
import collections
import copy
import difflib
import hashlib
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path
from lxml import etree as E
from openpyxl import load_workbook

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
QA = TASK / 'qa/unit-reader-wording'
REPORT = ROOT / 'docs/참고/인허가/사운드1/DHF/04_설계 검증/SW/펌웨어/SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서_수정중.xlsx'
DESIGN = ROOT / 'docs/참고/인허가/사운드1/DHF/03_설계 출력/SW/펌웨어/[작성자 확인] SDDD-FW-TD24A 소프트웨어 상세 설계서.docx'
SHEET = '보고서_유닛 시험'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
POLITE = re.compile(r'입니다|합니다|됩니다|습니다|십시오|하세요|[가-힣]+ㅂ니다')
LOG = re.compile(r'^\s*\[UT-[^\r\n]+', re.M)

def unit_names():
    with zipfile.ZipFile(DESIGN) as z:
        root = E.fromstring(z.read('word/document.xml'))
    names = {}
    for row in root.findall('.//w:tr', NS):
        fields = [''.join(t.text or '' for t in c.findall('.//w:t', NS)
                  if not any(a.tag == '{'+NS['w']+'}del' for a in t.iterancestors()))
                  for c in row.findall('w:tc', NS)]
        for i, value in enumerate(fields[:-1]):
            if re.fullmatch(r'SU-ID-[A-Z0-9]+-\d+', value.strip()):
                ident, name = value.strip(), fields[i+1].strip()
                if name:
                    assert ident not in names or names[ident] == name, (ident, name, names.get(ident))
                    names[ident] = name
    return names

def endings(value):
    for old, new in [('해야 합니다','해야 한다'), ('두었습니다','두었다'),
                     ('했습니다','했다'), ('않았습니다','않았다'),
                     ('없습니다','없다'), ('있습니다','있다'), ('아닙니다','아니다'),
                     ('않습니다','않는다'), ('읽습니다','읽는다'), ('넣습니다','넣는다'),
                     ('덮습니다','덮는다'), ('넘습니다','넘는다'), ('남습니다','남는다'),
                     ('찾습니다','찾는다'), ('바꿉니다','바꾼다'), ('둡니다','둔다'),
                     ('남깁니다','남긴다'), ('같습니다','같다'), ('합니다','한다'), ('됩니다','된다'),
                     ('입니다','이다')]:
        value = value.replace(old, new)
    return value

def case_number(number, unit):
    # HEAR C80/C81는 검증1의 원시값 기록이며 검증80/81가 아니다.
    return '1' if re.search(r'(?:HEAR-ID-2|HEAR-002)$',unit) and number in ('80','81') else number

def revised(value, sheet, coord, unit):
    reasons = []
    def record(old, new, why):
        nonlocal value
        updated = value.replace(old, new)
        if updated != value:
            value = updated
            reasons.append(why)
    # 줄 전체가 로그이면 문자열과 숫자를 그대로 보존한다.
    pieces = value.splitlines(keepends=True)
    for i, line in enumerate(pieces):
        if line.lstrip().startswith('[UT-'): continue
        converted = endings(line)
        if line != converted: reasons.append('문장 종결')
        pieces[i] = converted
    value = ''.join(pieces)
    for old, new in [('예상 RTT 전문','예상 RTT 출력 결과'),
                     ('예상 출력 전문','예상 출력 결과'), ('RTT 전문','RTT 출력 결과'),
                     ('기대값','예상값'), ('기댓값','예상값'),
                     ('기대 출력','예상 출력'), ('기대이다','예상값이다'),
                     ('전문','출력 내용')]:
        record(old, new, '예상 출력 용어')
    def heading(m):
        reasons.append('검증 번호와 RTT 연결')
        n = m.group(1)
        return f'검증 [{case_number(n, unit)}]의 RTT 출력 (C{n})'
    value = re.sub(r'(?m)^RTT C(\d+)(?=\s*$)', heading, value)
    # 원소별 입력/결과에도 같은 C번호를 붙여 RTT 제목과 함께 찾을 수 있게 한다.
    def input_heading(m):
        reasons.append('검증 번호와 입력 연결')
        n = m.group(2)
        return f'검증 [{case_number(n, unit)}]의 입력/원소별 {m.group(1)} (C{n})'
    value = re.sub(r'입력/원소별 (기준|결과): C(\d+)(?=\D|$)', input_heading, value)
    for old in ['그룹 전체 원본 .log를 저장한다.', '전체 .log를 보관한다.']:
        record(old, '', '로그 저장 지시 삭제')
    record('배열과 판정 줄을 원본 .log로 보관한다.', '배열과 판정 줄을 비교한다.', '로그 저장 지시 삭제')
    record('배열/불일치/요약 및 END 원문을 저장하고 예상값과 비교한다.',
           '배열/불일치/요약 및 END 출력을 예상값과 비교한다.', '로그 저장 지시 삭제')
    record('구형 Board1_2 ISL91128와 현재 소자를 혼동하지 않는다.', '', '레거시 주의문 삭제')
    record('준비하고 구형 PC 보충과 현재 보드를 구분한다.', '준비한다.', '레거시 주의문 삭제')
    record('구형 Board_1_3 보충 시험 사진은 현재 Board_1_5와 별도 표시하며 전압/전류 수치를 임의 기재하지 않음.',
           '', '레거시 주의문 삭제')
    # 실제 시험 범위/보류/실패를 설명하는 문장은 이 목록에 포함하지 않는다.
    for sentence in ['실제 핀 전압/QCC 동작이 필요한 별도 시험 자료와 구분한다.',
                     'QCC 측 수신/처리 및 실제 핀 전압은 별도 검증이며 사진 없이 텍스트로 기록.',
                     'QCC 수신/중계는 별도 자료다.',
                     '예상값/제품 오류는 유지하고 실제 SPI/QCC/파일 완료는 별도 검증한다.']:
        record(sentence, '', '불필요한 QCC 별도 자료 문장 삭제')
    record('PC 파서 접수와 대응 기능 실행/완료/QCC 중계를 별도 경계로 둔다.',
           '파서 접수와 대응 기능의 실행 완료를 구분한다.', 'QCC 별도 자료 문장 정리')
    record('실제 파일/물리 SPI/QCC/재시작은 별도 자료로 확인하며 위험한 tail>16/잘못된 종류 WRITE는 실행하지 않는다.',
           'tail>16/잘못된 종류 WRITE는 실행하지 않는다.', '중복 자료 요구 삭제/안전 조건 유지')
    record('실제 핀을 확인한 경우 측정 지점/조건/화면을 별도 붙인다.', '', '중복 선택 사진 안내 삭제')
    if sheet == SHEET and coord in ('J800','J835'):
        record('사진 없이 텍스트로 기록.', '', '중복 사진 안내 삭제')
    # 실제 파형/핀 사진은 기준에서 요구되는 경우에만 유지한다.
    optional_photos = {
        (SHEET,'L160'): '(사진필요) 저역통과 필터 후 OUT+/OUT- 차동 파형(보조자료).',
        (SHEET,'L1594'): '(사진필요) 대상 메모리별 실제 전송 내용/주소/길이와 24비트 패킹 비교 화면',
        ('시험 보조자료','C89'): '파형을 추가 확인하는 경우에만 (사진필요) 1000 Hz 사인파/무음 입력에서 저역통과 필터 후 OUT+/OUT- 차동 파형.',
        ('시험 보조자료','C207'): '(사진필요) 실제 16표본 결과와 최대값, Event0/1 및 오류 시나리오 대조 화면',
        ('시험 보조자료','C218'): '(사진필요) 실제 창/FFT/크기 결과, scale 지수 E, Event1 및 허용오차 비교 화면',
        ('시험 보조자료','C567'): '(사진필요) 대상 메모리별 실제 전송 내용/주소/길이와 24비트 패킹 비교 화면',
        ('예상 그림','C6'): '파형을 추가 확인할 때만 (사진필요) 필터 후 차동 파형을 첨부한다.',
        ('예상 그림','C21'): '(사진필요) 같은 입력 조건의 실제 사진/파형을 첨부한다.',
        ('예상 그림','C27'): '(사진필요) 같은 입력 조건의 실제 사진/파형을 첨부한다.',
        ('예상 그림','C33'): '(사진필요) 같은 입력 조건의 실제 사진/파형을 첨부한다.',
    }
    if (sheet,coord) in optional_photos:
        record(optional_photos[(sheet,coord)], '', '중복/선택 사진 요청 삭제')
    for old,new in [
        ('(사진필요) 실제 기동 LED 표시/정지 사진/영상은 별도 관찰 원본. PC tick 4000/14440 결과는 사진 없이 텍스트로 실행 원문 보존',
         '(사진필요) 실제 기동 LED의 표시/정지 사진 또는 영상.'),
        ('(사진필요) 실제 서비스 LED 표시 사진/시간표시 영상. PC reset 호출 1회는 사진 없이 텍스트로 별도 기록',
         '(사진필요) 실제 서비스 LED의 표시 사진 또는 시간 표시 영상.'),
        ('(사진필요) 실제 기동 LED 표시/정지 모습과 tick 4000/14440의 결과 화면',
         '(사진필요) 실제 기동 LED의 표시/정지 사진 또는 영상.'),
        ('(사진필요) 서비스 LED 표시와 실제 reset 요청 훅 1회 결과 화면',
         '(사진필요) 실제 서비스 LED의 표시 사진 또는 시간 표시 영상.')]:
        record(old,new,'필수 사진 안내 간소화')
    if reasons:
        if not value.strip(): value = ''
        else:
            value = re.sub(r'(?m)^[ \t]+$', '', value)
            value = re.sub(r'^\n +', '\n', value)
            value = re.sub(r'\n{3,}$', '\n', value)
    if sheet == 'RTT 결과 대조' and re.fullmatch(r'\s*C\d+\s*', value):
        n = value.strip()[1:]
        value = '\n'+f'검증 [{case_number(n, unit)}]의 RTT 출력 (C{n})'+'\n'
        reasons.append('검증 번호와 RTT 연결')
    if sheet == '예상 그림' and coord == 'A3':
        record('본문의 시험 ID로 찾고, 같은 입력의 실제 사진을 오른쪽에 첨부한다.',
               '본문의 시험 ID로 찾는다. 실제 사진이 필요한 위치는 본문의 (사진필요) 표시로 확인한다.',
               '중복/선택 사진 요청 삭제')
    if sheet == '예상 그림' and coord == 'A27':
        record('실제 시간 수치는 스코프 사진에 기록한다.', '', '중복/선택 사진 요청 삭제')
    return value, sorted(set(reasons))

def prepare():
    QA.mkdir(parents=True, exist_ok=True)
    baseline_path = QA/'unit-before.xlsx'
    if baseline_path.exists():
        assert digest(baseline_path) == digest(REPORT), '작성자 원본이 달라졌다.'
    else: shutil.copy2(REPORT, baseline_path)
    book = load_workbook(baseline_path, rich_text=True)
    names = unit_names()
    changes, name_map = [], []
    for sheet in book:
        unit = ''
        for row in sheet:
            for cell in row:
                value = plain(cell.value)
                if not value or cell.data_type == 'f': continue
                if sheet.title == SHEET and cell.column == 4 and 'SW-FW-UnitTest' in value:
                    unit = value.strip()
                if sheet.title == 'RTT 결과 대조' and cell.column == 1 and re.search(r'(?:CFX|HEAR|CM3|BOOT)-(?:\d{3}|ID-\d+)',value):
                    unit = value.strip()
                revised_value, reasons = revised(value, sheet.title, cell.coordinate, unit)
                if sheet.title == SHEET and cell.column == 7 and 'SU-ID-' in value:
                    ident = value.strip()
                    assert ident in names, (cell.coordinate, ident)
                    revised_value = '\n'+ident+'\n('+names[ident]+')\n'
                    reasons.append('상세 설계 유닛 이름 추가')
                    name_map.append(dict(cell=cell.coordinate, unit=unit, ident=ident, name=names[ident]))
                assert LOG.findall(value) == LOG.findall(revised_value), (sheet.title,cell.coordinate,'로그 변화')
                if revised_value != value:
                    changes.append(dict(key='unit',sheet=sheet.title,cell=cell.coordinate,
                                        old=value,value=revised_value,reason=' / '.join(sorted(set(reasons))),unit=unit))
    assert len(name_map) == 115, len(name_map)
    remaining = [(p['sheet'],p['cell'],POLITE.findall(p['value'])) for p in changes if POLITE.search(p['value'])]
    assert not remaining, remaining[:20]
    protected = {str(p):digest(p) for p in REPORT.parent.glob('*.xlsx')
                 if p != REPORT and not p.name.startswith('~$')}
    protected[str(DESIGN)] = digest(DESIGN)
    for p in DESIGN.parent.glob('*SUVC*.xlsx'):
        if not p.name.startswith('~$'):protected[str(p)]=digest(p)
    for p in (ROOT/'src/test').rglob('*'):
        if p.is_file(): protected[str(p)] = digest(p)
    views = ['G30:G37','K31:L32','K683:L683','K160:L160','G1594:L1594']
    save('baseline.json',dict(books=[dict(key='unit',path=str(REPORT),backup=str(baseline_path),sheet=SHEET,views=views)],
                              sourceHash=digest(REPORT),protected=protected))
    save('plan.json',dict(changes=changes,shades=[],nameMap=name_map))
    save('summary.json',dict(changedCells=len(changes),bySheet=dict(collections.Counter(p['sheet'] for p in changes)),
                            byReason=dict(collections.Counter(r for p in changes for r in p['reason'].split(' / ')))))
    print((QA/'summary.json').read_text(encoding='utf-8'))

def strings_in(package, cell, shared):
    if cell.get('t') == 's':
        return shared[int(cell.find('s:v',NS).text)]
    if cell.get('t') == 'inlineStr': return cell.find('s:is',NS)
    return None

def patch_text(package, cell, value, shared):
    original = strings_in(package, cell, shared)
    text, props = '', []
    if original is not None:
        for node in original:
            if node.tag == '{'+NS['s']+'}r':
                t = node.find('s:t',NS); rp = node.find('s:rPr',NS)
                chunk = '' if t is None else t.text or ''
            elif node.tag == '{'+NS['s']+'}t':
                chunk, rp = node.text or '', None
            else: continue
            text += chunk; props.extend([rp]*len(chunk))
    uniform = {None if p is None else E.tostring(p) for p in props}
    out = [None]*len(value)
    if len(uniform) == 1:
        out = [props[0]]*len(value)
    else:
        for tag,i,j,k,l in difflib.SequenceMatcher(None,text,value,autojunk=False).get_opcodes():
            if tag == 'equal': out[k:l] = props[i:j]
            elif tag in ('insert','replace') and props:
                source = props[min(i,len(props)-1)]
                out[k:l] = [source]*(l-k)
    # 새 유닛 이름에는 셀 기본 서식을 사용하며 사진/QCC 강조를 물려주지 않는다.
    if value.strip().startswith('SU-ID-') and '\n(' in value:
        out[value.index('\n('):] = [None]*(len(value)-value.index('\n('))
    for node in list(cell): cell.remove(node)
    cell.set('t','inlineStr')
    ins = E.SubElement(cell,'{'+NS['s']+'}is')
    groups = []
    for ch, prop in zip(value,out):
        key = None if prop is None else E.tostring(prop)
        if groups and groups[-1][0] == key: groups[-1][1] += ch
        else: groups.append([key,ch,prop])
    for _,chunk,prop in groups:
        run = E.SubElement(ins,'{'+NS['s']+'}r')
        if prop is not None: run.append(copy.deepcopy(prop))
        t = E.SubElement(run,'{'+NS['s']+'}t')
        t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=chunk

def build():
    baseline, plan = read('baseline.json'), read('plan.json')
    authored = load_workbook(QA/'unit-authored.xlsx',rich_text=True)
    with zipfile.ZipFile(QA/'unit-before.xlsx') as z:
        package = {n:z.read(n) for n in z.namelist()}
    shared = E.fromstring(package['xl/sharedStrings.xml']).findall('s:si',NS) if 'xl/sharedStrings.xml' in package else []
    wbxml=E.fromstring(package['xl/workbook.xml'])
    rels=E.fromstring(package['xl/_rels/workbook.xml.rels'])
    targets={r.get('Id'):r.get('Target') for r in rels}
    paths={s.get('name'):targets[s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
           for s in wbxml.findall('s:sheets/s:sheet',NS)}
    parsed={}
    for p in plan['changes']:
        assert plain(authored[p['sheet']][p['cell']].value) == p['value'], (p['sheet'],p['cell'])
        path=paths[p['sheet']]
        path=path.lstrip('/') if path.startswith('/') else 'xl/'+path
        if path not in parsed: parsed[path]=E.fromstring(package[path])
        cell=parsed[path].find(f'.//s:c[@r="{p["cell"]}"]',NS)
        assert cell is not None, p['cell']
        patch_text(package,cell,p['value'],shared)
    main_path=paths[SHEET]
    main_path=main_path.lstrip('/') if main_path.startswith('/') else 'xl/'+main_path
    for item in plan.get('heights',[]):
        row=parsed[main_path].find(f'.//s:row[@r="{item["row"]}"]',NS)
        assert row is not None,item['row']
        row.set('ht',str(item['height']));row.set('customHeight','1')
    for path,root in parsed.items(): package[path]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    with zipfile.ZipFile(QA/'unit-candidate.xlsx','w',zipfile.ZIP_DEFLATED) as z:
        for path,data in package.items():z.writestr(path,data)
    print('native candidate:',len(plan['changes']),'cells')

def verify():
    before=load_workbook(QA/'unit-before.xlsx',rich_text=True)
    after=load_workbook(QA/'unit-candidate.xlsx',rich_text=True)
    plan=read('plan.json'); expected={(p['sheet'],p['cell']):p for p in plan['changes']}
    counts=collections.Counter(); logcount=0
    assert before.sheetnames == after.sheetnames
    def result_numbers(value):
        value='\n'.join(line for line in value.splitlines() if '(사진필요)' not in line)
        value=re.sub(r'검증 \[\d+\]|C\d+','',value)
        return re.findall(r'0[xX][0-9A-Fa-f]+|-?\d+(?:\.\d+)?',value)
    for old in before:
        new=after[old.title]
        assert (old.max_row,old.max_column)==(new.max_row,new.max_column)
        assert set(map(str,old.merged_cells.ranges))==set(map(str,new.merged_cells.ranges))
        assert len(old._images)==len(new._images)
        assert {k:dict(v) for k,v in old.column_dimensions.items()}=={k:dict(v) for k,v in new.column_dimensions.items()}
        height_rows={p['row']:p['height'] for p in plan.get('heights',[]) if p['sheet']==old.title}
        assert {k:dict(v) for k,v in old.row_dimensions.items() if k not in height_rows}=={k:dict(v) for k,v in new.row_dimensions.items() if k not in height_rows}
        for row,height in height_rows.items():
            assert new.row_dimensions[row].height==height,(row,height)
        for row in old:
            for c in row:
                nc=new[c.coordinate];a,b=plain(c.value),plain(nc.value)
                key=(old.title,c.coordinate)
                assert c._style==nc._style, key
                if key in expected:
                    assert b==expected[key]['value'],key
                    assert LOG.findall(a)==LOG.findall(b),key
                    counts['changed']+=1
                else: assert a==b,(key,'unplanned cell change')
                logcount+=len(LOG.findall(a))
                assert not POLITE.search(b), (key,POLITE.findall(b))
                if old.title==SHEET and c.column in (4,13):assert a==b,key
                if old.title==SHEET and c.column==12:
                    assert result_numbers(a)==result_numbers(b),(key,'결과 수치 변화')
    with zipfile.ZipFile(QA/'unit-before.xlsx') as a,zipfile.ZipFile(QA/'unit-candidate.xlsx') as b:
        changedparts=[n for n in a.namelist() if a.read(n)!=b.read(n)]
        assert all(re.fullmatch(r'xl/worksheets/sheet\d+\.xml',n) for n in changedparts),changedparts
    for path,h in read('baseline.json')['protected'].items():assert digest(path)==h,path
    for p in plan['nameMap']:
        assert plain(after[SHEET][p['cell']].value).strip()==p['ident']+'\n('+p['name']+')',p['cell']
    for addr,code in [('B145','80'),('B149','81')]:
        assert plain(after['RTT 결과 대조'][addr].value).strip()==f'검증 [1]의 RTT 출력 (C{code})',addr
    result=dict(changedCells=counts['changed'],rawLogLinesPreserved=logcount,
                unitNamesAdded=len(plan['nameMap']),sheetStructuresPreserved=True,
                testIDsAndVerdictsPreserved=True,resultNumbersPreserved=True,
                rowHeightsAdjusted=len(plan.get('heights',[])),changedParts=changedparts)
    save('verification.json',result);print(json.dumps(result,ensure_ascii=False,indent=2))

def publish():
    base=read('baseline.json')
    native=json.loads((QA/'native-checks.json').read_text(encoding='utf-8-sig'))
    assert not native['insufficient'],'표시 높이 부족을 먼저 해결한다.'
    assert read('verification.json')['resultNumbersPreserved']
    assert digest(REPORT)==base['sourceHash'], '작성자 추가 수정이 감지되어 덮어쓰기를 중단한다.'
    # 같은 폴더에서 교체하며 열린 파일은 OS의 잠금에 따라 실패한다.
    staging=REPORT.with_name(REPORT.name+'.tmp')
    shutil.copy2(QA/'unit-candidate.xlsx',staging)
    staging.replace(REPORT)
    assert digest(REPORT)==digest(QA/'unit-candidate.xlsx')
    save('publication.json',dict(path=str(REPORT),hash=digest(REPORT)))
    print('published',str(REPORT))

def fit_heights():
    plan=read('plan.json')
    checks=json.loads((QA/'native-checks.json').read_text(encoding='utf-8-sig'))
    book=load_workbook(QA/'unit-candidate.xlsx',rich_text=True)
    sheet=book[SHEET]; heights={}
    for item in checks['insufficient']:
        assert item['sheet']==SHEET,item
        addr=item['cell'];cell=sheet[addr]
        merged=next((m for m in sheet.merged_cells.ranges if addr in m),None)
        start=cell.row if merged is None else merged.min_row
        end=cell.row if merged is None else merged.max_row
        sizes={r:heights.get(r,sheet.row_dimensions[r].height or sheet.sheet_format.defaultRowHeight)
               for r in range(start,end+1)}
        shortage=item['needed']+2-sum(sizes.values())
        if shortage>0:
            for r in range(start,end+1):
                add=min(shortage,409.5-sizes[r]);sizes[r]+=add;shortage-=add
                heights[r]=sizes[r]
                if shortage<=0:break
        assert shortage<=0,(addr,shortage)
    plan['heights']=[dict(key='unit',sheet=SHEET,row=r,height=h) for r,h in sorted(heights.items())]
    save('plan.json',plan)
    save('height-adjustments.json',dict(rows=len(heights),maximum=max(heights.values()),source=checks['insufficient']))
    print('row heights:',len(heights),'maximum:',max(heights.values()))

def correct_helper_cases():
    plan=read('plan.json'); book=load_workbook(QA/'unit-before.xlsx',rich_text=True)
    sheet=book['RTT 결과 대조']; units={}; unit=''
    for row in range(1,sheet.max_row+1):
        value=plain(sheet.cell(row,1).value).strip()
        if re.fullmatch(r'(?:CFX|HEAR|CM3|BOOT)-\d{3}',value):unit=value
        units[row]=unit
    corrections=[]
    for p in plan['changes']:
        if p['sheet']!='RTT 결과 대조':continue
        row=int(re.search(r'\d+',p['cell']).group());p['unit']=units[row]
        if p['cell'].startswith('B'):
            case=re.search(r'\(C(\d+)\)',p['value'])
            if case and case_number(case.group(1),p['unit'])!=case.group(1):
                corrected=re.sub(r'검증 \[\d+\]',f'검증 [{case_number(case.group(1),p["unit"])}]',p['value'])
                if corrected!=p['value']:
                    p['value']=corrected;corrections.append(p['cell'])
    assert corrections==['B145','B149'],corrections
    save('plan.json',plan);save('helper-case-corrections.json',corrections);print(corrections)

def patch_helper_cases():
    plan=read('plan.json'); cells=read('helper-case-corrections.json')
    authored=load_workbook(QA/'unit-authored.xlsx',rich_text=True)
    with zipfile.ZipFile(QA/'unit-candidate.xlsx') as z:package={n:z.read(n) for n in z.namelist()}
    path='xl/worksheets/sheet4.xml';root=E.fromstring(package[path])
    for p in plan['changes']:
        if p['sheet']=='RTT 결과 대조' and p['cell'] in cells:
            assert plain(authored[p['sheet']][p['cell']].value)==p['value']
            cell=root.find(f'.//s:c[@r="{p["cell"]}"]',NS)
            patch_text(package,cell,p['value'],[])
    package[path]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    temp=QA/'unit-candidate-new.xlsx'
    with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
        for path,data in package.items():z.writestr(path,data)
    temp.replace(QA/'unit-candidate.xlsx');print('helper case links corrected')

def notes():
    plan,summary,checks=read('plan.json'),read('summary.json'),read('verification.json')
    lines=[
        '# 유닛 검증 보고서 문구 수정 내역', '',
        '은수님이 편집한 `SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서_수정중.xlsx`를 수정했다. '
        '수정전 파일과 작성자 전용 파일은 변경하지 않았다.', '',
        '시험값/기존 RTT 문자열/판정과 행/병합/그림은 유지했다. '
        '수정 범위는 문장, 검증 번호 표시, 유닛 이름, 불필요한 자료 안내이다.', '',
        '## 1. 요청별 변경 내용', '',
        '| 요청 | 변경 전 | 변경 후 / 적용 위치 |', '|---|---|---|',
        '| 짧은 문장 종결 | `호출합니다`, `같습니다`, `기대입니다` | `호출한다`, `같다`, `예상값이다`. 표지를 제외한 본문/보조 시트의 해당 문구 전체 |',
        '| 검증 번호 연결 | `RTT C1` | `검증 [1]의 RTT 출력 (C1)`. 본문 K열과 RTT 결과 대조 B열 |',
        '| 입력과 RTT 연결 | `입력/원소별 기준: C1` | `검증 [1]의 입력/원소별 기준 (C1)`. 해당 기준/결과 제목 |',
        '| 로그 저장 지시 삭제 | `그룹 전체 원본 .log를 저장한다.` | 해당 문장 삭제. 실행/배열/판정 비교 절차는 유지 |',
        '| 쉬운 예상 출력 표현 | `예상 RTT 전문`, `기대값` | `예상 RTT 출력 결과`, `예상값`. 본문과 보조 시트 모두 적용 |',
        '| 유닛 이름 추가 | `SU-ID-CFX-001` | 줄 바꿈 후 `(인터럽트와 반복 실행 제어)`. 본문 G열 115개 |',
        '| 레거시 주의문 삭제 | `구형 Board1_2 ISL91128와 현재 소자를 혼동하지 않는다.` | 주의문 삭제. 실제 과거 시험의 입력/결과/적용 범위는 유지 |',
        '| QCC 별도 자료 문장 정리 | `QCC 수신/중계는 별도 자료다.` | E8300 자체 시험에 중복되는 자료 요구 삭제. QCC 내부 자료 추가 자리와 실제 시험 범위는 유지 |',
        '| 불필요한 사진 요청 삭제 | OD 참고 파형, 메모리 전송 비교 화면 | 본문 2곳 삭제. 보조 시트의 중복/선택 안내도 함께 정리 |', '',
        'HEAR-ID-2의 C80/C81은 검증 [1]의 원시값 기록이다. '
        '`검증 [1]의 RTT 출력 (C80)` / `(C81)`로 표시하며 로그 안의 C번호는 바꾸지 않았다. '
        'C101처럼 기존 시험에 사용된 번호도 원래 C번호와 함께 표시했다.', '',
        '## 2. 사진 표시 정리', '',
        '| 시험 | 본문 위치 | 처리 / 이유 |', '|---|---|---|',
        '| CFX-ID-20 | L160 | 참고용 저역통과 필터 후 OD 파형 사진 요청 삭제. RAM 배열 비교가 본 시험 기준이며 청취 조건은 유지 |',
        '| BOOT-ID-15 | L1594 | 메모리 전송 비교 화면 사진 요청 삭제. 전송 값/순서/길이를 텍스트로 비교하는 기준은 유지 |',
        '| CM3-ID-14 | L898 | 유지. 실제 LED 색/점멸/소등의 확인 자료 |',
        '| CM3-ID-39 | L1158 | 유지. 실제 VP 활성/차단 출력 유무의 확인 자료 |',
        '| BOOT-ID-4 | L1312 | 유지. 실제 기동 LED의 표시/정지 확인 자료. 저장 지시와 중복 설명만 삭제 |',
        '| BOOT-ID-18 | L1648 | 유지. 실제 서비스 LED 표시의 확인 자료. reset 호출 횟수와 사진 안내는 분리 |', '',
        '보조 시트에서는 HEAR 배열/FFT 비교 화면, PMIC 코드 비교 그림, 선택 Calibration 그림 등에 '
        '남아 있던 중복/선택 사진 요청도 삭제했다. 예상 그림 11개 자체는 그대로 유지했다. '
        '사진이 필요한 위치는 본문의 파란색 `(사진필요)` 표시로 찾는다.', '',
        '## 3. 주요 변경 위치', '',
        '공통 문장/예상값 용어는 해당 문구가 있는 셀 전체에 적용했다. '
        '아래는 자료 요구와 시험 설명에 영향을 주는 문장 변경 위치이다.', '',
        '| 변경 종류 | 시트 | 셀 |', '|---|---|---|']
    cats=('로그 저장 지시 삭제','레거시 주의문 삭제','불필요한 QCC 별도 자료 문장 삭제',
          'QCC 별도 자료 문장 정리','중복 자료 요구 삭제/안전 조건 유지',
          '중복/선택 사진 요청 삭제','중복 선택 사진 안내 삭제','중복 사진 안내 삭제','필수 사진 안내 간소화')
    for cat in cats:
        bysheet=collections.defaultdict(list)
        for p in plan['changes']:
            if cat in p['reason'].split(' / '):bysheet[p['sheet']].append(p['cell'])
        for sheet,cells in bysheet.items():
            lines.append(f'| {cat} | {sheet} | '+', '.join(cells)+' |')
    lines += ['', '## 4. 유닛 이름을 추가한 위치', '',
              '최신 소프트웨어 상세 설계서의 유닛 목록과 대조했다. '
              '식별 번호는 그대로 두고 다음 이름만 줄 바꿈/괄호로 덧붙였다.', '',
              '| 상세 설계 ID | 유닛 이름 | 보고서 셀 |','|---|---|---|']
    for p in plan['nameMap']:lines.append(f'| {p["ident"]} | {p["name"]} | G열 {p["cell"]} |')
    lines += ['', '## 5. 저장본 확인', '',
              f'- 수정한 셀은 {summary["changedCells"]:,}개이다. 같은 셀에 여러 변경이 함께 적용될 수 있다.',
              f'- 유닛 이름 {checks["unitNamesAdded"]}개와 기존 RTT 줄 {checks["rawLogLinesPreserved"]:,}개를 대조했다.',
              '- 기존 115개 시험 ID/판정, 7개 시트의 행/열/병합, 그림 11개와 셀 기본 서식을 유지했다.',
              f'- 문구 표시가 부족한 본문 {checks["rowHeightsAdjusted"]}개 행은 글씨 크기를 바꾸지 않고 높이만 늘렸다. 행 추가/병합 변경은 없으며 409.50pt 이하이다.',
              '- 수정전/작성자 전용 사본, 유닛 기준서, 상세 설계서와 공통 시험 소스의 파일 해시가 동일하다.',
              '- 검증 번호/유닛 이름의 표시 높이와 대표 화면을 별도로 확인했다.',
              '- 기존 실행 범위, 미완료/불일치 기록과 QCC 내부 자료 추가 안내는 문구 정리로 통과 처리하지 않았다.', '']
    content='\n'.join(lines)
    destination=REPORT.parent/'유닛 검증 보고서 문구 수정 내역.md'
    patch='*** Begin Patch\n*** Add File: '+destination.as_posix()+'\n'+''.join('+'+x+'\n' for x in content.splitlines())+'*** End Patch'
    print(json.dumps(dict(path=str(destination),patch=patch),ensure_ascii=False))

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def save(name, value):
    (QA / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')

def read(name):
    return json.loads((QA / name).read_text(encoding='utf-8'))

def plain(v):
    return '' if v is None else str(v)

def inventory():
    QA.mkdir(parents=True, exist_ok=True)
    book = load_workbook(REPORT, rich_text=True)
    allcells = []
    for s in book:
        unit = ''
        for row in s:
            for c in row:
                val = plain(c.value)
                if not val: continue
                if s.title == SHEET and c.column == 4 and 'SW-FW-UnitTest' in val:
                    unit = val.strip()
                allcells.append(dict(sheet=s.title, cell=c.coordinate, text=val, unit=unit))
    with zipfile.ZipFile(DESIGN) as archive:
        root = E.fromstring(archive.read('word/document.xml'))
    design = []
    for p in root.findall('.//w:p', NS):
        text = ''.join(t.text or '' for t in p.findall('.//w:t', NS)
                       if not any(x.tag == '{'+NS['w']+'}del' for x in t.iterancestors()))
        if 'SU-ID-' in text: design.append(text)
    inventory_data = dict(path=str(REPORT), hash=digest(REPORT),
        sheets=[dict(name=s.title, rows=s.max_row, columns=s.max_column,
                     merges=len(s.merged_cells.ranges), images=len(s._images)) for s in book],
        cells=allcells, design=design)
    save('inventory.json', inventory_data)
    print(json.dumps({k:inventory_data[k] for k in ('hash','sheets')}, ensure_ascii=False, indent=2))
    patterns = {'존댓말':r'입니다|합니다|됩니다|습니다|하십시오|하세요',
                'RTT 제목':r'(?:^|\n)RTT C\d+',
                '로그 저장':r'\.log|로그.{0,15}(?:저장|보관)|원본.{0,20}저장',
                '딱딱한 표현':r'전문|기대입니다|기대값|기대 출력|기댓값',
                '레거시':r'구형|Board1_2|ISL91128|레거시',
                'QCC 부가':r'QCC.{0,70}(?:별도|자료)',
                '자료 첨부':r'사진필요|사진 필요|캡처|첨부|촬영'}
    for name, pattern in patterns.items():
        matches = [x for x in allcells if re.search(pattern, x['text'])]
        print(name, len(matches))
        save('scan-'+name+'.json', matches)
    print('design SU paragraphs:', len(design))

if __name__ == '__main__':
    {'inventory':inventory,'prepare':prepare,'build':build,'verify':verify,'publish':publish,'notes':notes,'fit-heights':fit_heights,
     'correct-helper-cases':correct_helper_cases,'patch-helper-cases':patch_helper_cases}[sys.argv[1]]()
