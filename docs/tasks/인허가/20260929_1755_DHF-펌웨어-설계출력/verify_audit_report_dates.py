"""Transplant Artifact-authored text/fills without flattening native report features."""
from copy import deepcopy
from datetime import datetime, timedelta
from pathlib import Path
import json
import re
import shutil
import sys
import zipfile
import openpyxl
from lxml import etree as E
from prepare_audit_report_dates import QA, REPORT, sha, FILL, MARKER
from verify_audit_expected_report import S, NS, q, unpack, cell_value, font_properties
if '--confirmed-results' in sys.argv:
    from prepare_audit_confirmed_report import QA


def addresses(plan):
    out = set()
    for area in plan['fills']:
        low_c, low_r, high_c, high_r = openpyxl.utils.range_boundaries(area)
        for row in range(low_r, high_r + 1):
            for col in range(low_c, high_c + 1):
                out.add(f'{openpyxl.utils.get_column_letter(col)}{row}')
    return out


def preserve():
    plan = json.loads((QA / 'sheet-edits.json').read_text(encoding='utf8'))
    parts = unpack(QA / 'report-before.xlsx')
    authored = unpack(QA / 'report-artifact.xlsx')
    root = E.fromstring(parts['xl/worksheets/sheet2.xml'])
    src = E.fromstring(authored['xl/worksheets/sheet2.xml'])
    strings = E.fromstring(authored['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in authored else []
    styles = E.fromstring(parts['xl/styles.xml'])
    authored_styles = E.fromstring(authored['xl/styles.xml'])
    for addr, wanted in plan['edits'].items():
        target = root.find(f'.//s:c[@r="{addr}"]', NS)
        source = src.find(f'.//s:c[@r="{addr}"]', NS)
        assert target is not None and source is not None, addr
        value = cell_value(source, strings)
        assert value == wanted, addr
        pr = font_properties(styles, target)
        if addr in plan['clear_marker_fonts']:
            for tag in ('b', 'color', 'u'):
                for child in pr.findall(q(tag)):
                    pr.remove(child)
            E.SubElement(pr, q('color'), {'rgb': 'FF000000'})
        for child in list(target):
            target.remove(child)
        target.set('t', 'inlineStr')
        inline = E.SubElement(target, q('is'))
        run = E.SubElement(inline, q('r'))
        run.append(pr)
        E.SubElement(run, q('t'), {'{http://www.w3.org/XML/1998/namespace}space': 'preserve'}).text = value

    fills = styles.find(q('fills'))
    fid = next((i for i, f in enumerate(fills)
                if f.find(f'.//s:fgColor[@rgb="{FILL}"]', NS) is not None), None)
    if fid is None:
        fid = len(fills)
        fill = E.SubElement(fills, q('fill'))
        pattern = E.SubElement(fill, q('patternFill'), {'patternType': 'solid'})
        E.SubElement(pattern, q('fgColor'), {'rgb': FILL})
        E.SubElement(pattern, q('bgColor'), {'indexed': '64'})
        fills.set('count', str(len(fills)))
    xfs = styles.find(q('cellXfs'))
    map_styles = {}
    for addr in sorted(addresses(plan)):
        target = root.find(f'.//s:c[@r="{addr}"]', NS)
        source = src.find(f'.//s:c[@r="{addr}"]', NS)
        assert target is not None and source is not None, addr
        sxf = authored_styles.find(q('cellXfs'))[int(source.get('s', '0'))]
        sf = authored_styles.find(q('fills'))[int(sxf.get('fillId', '0'))]
        assert sf.find(f'.//s:fgColor[@rgb="{FILL}"]', NS) is not None, addr
        old_id = int(target.get('s', '0'))
        if old_id not in map_styles:
            new_xf = deepcopy(xfs[old_id])
            new_xf.set('fillId', str(fid))
            new_xf.set('applyFill', '1')
            map_styles[old_id] = len(xfs)
            xfs.append(new_xf)
        target.set('s', str(map_styles[old_id]))
    xfs.set('count', str(len(xfs)))
    for row, height in plan['heights'].items():
        node = root.find(f'.//s:row[@r="{row}"]', NS)
        node.set('ht', str(height))
        node.set('customHeight', '1')
    parts['xl/worksheets/sheet2.xml'] = E.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
    if plan['fills']:
        parts['xl/styles.xml'] = E.tostring(styles, encoding='UTF-8', xml_declaration=True, standalone=True)
    with zipfile.ZipFile(QA / 'report.xlsx', 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)


def check(path=None):
    path = path or QA / 'report.xlsx'
    plan = json.loads((QA / 'sheet-edits.json').read_text(encoding='utf8'))
    baseline = json.loads((QA / 'baseline.json').read_text(encoding='utf8'))
    old = openpyxl.load_workbook(QA / 'report-before.xlsx', rich_text=True)
    new = openpyxl.load_workbook(path, rich_text=True)
    assert new.sheetnames == old.sheetnames
    assert new.worksheets[1].max_row == old.worksheets[1].max_row == 1734
    differences = []
    for a, b in zip(old.worksheets, new.worksheets):
        assert [str(m) for m in a.merged_cells.ranges] == [str(m) for m in b.merged_cells.ranges]
        for row in a:
            for cell in row:
                other = b[cell.coordinate]
                if str(cell.value) != str(other.value):
                    assert b == new.worksheets[1] and cell.coordinate in plan['edits'], cell.coordinate
                    differences.append(cell.coordinate)
                if other.value:
                    assert MARKER not in str(other.value), other.coordinate
    ws = new.worksheets[1]
    for addr, wanted in plan['edits'].items():
        assert str(ws[addr].value) == wanted, addr
        assert wanted.startswith('\n') and wanted.endswith('\n'), addr
        assert not re.search('[①-⑳\u3040-\u30ff\u4e00-\u9fff]', wanted), addr
        assert '출력: 출력:' not in wanted
    for area in plan['fills']:
        anchor = area.split(':')[0]
        assert ws[anchor].fill.fgColor.rgb == FILL, anchor
    for addr in plan['clear_marker_fonts']:
        assert all(not (getattr(run, 'font', None) and run.font.color
                        and run.font.color.type == 'rgb' and run.font.color.rgb == 'FF0070C0')
                   for run in ws[addr].value), addr
    for unit, first, last, count in [(1, 789, 792, 8), (5, 809, 811, 8), (46, 1208, 1213, 10)]:
        text = '\n'.join(str(ws[f'K{r}'].value or '') for r in range(first, last+1))
        assert [int(x) for x in re.findall(r'^조건 \[(\d+)\](?:\s|$)', text, re.M)] == list(range(1, count+1)), unit
    confirmed = plan.get('mode') == 'tester-confirmed-results'
    if confirmed:
        assert plan['test_confirmation']
        for first, last, count in ((789, 792, 8), (809, 811, 8), (1208, 1213, 10)):
            results = '\n'.join(str(ws[f'L{r}'].value or '') for r in range(first, last+1))
            labels = [int(x) for x in re.findall(r'^검증 \[(\d+)\](?:\s|$)', results, re.M)]
            assert labels == list(range(1, count+1)), (first, labels)
            assert not any(x in results for x in ('예상', '후속 시험', '코드 반영', '재시험', '불일치', '기존 보충')), first
            assert str(ws[f'M{first}'].value) == '\n통과\n', first
        assert '저장 실패 입력 / 초기화 함수 반환=-1.' in str(ws['L792'].value)
        assert '닫기 실패 입력 / 기록 함수 반환=-1.' in str(ws['L1213'].value)
    else:
        assert '저장 실패 입력 시 반환=0이다. 기준 -1과 불일치한다.' in str(ws['L792'].value)
        assert '닫기 실패 입력 시 반환=0이다. 기준 -1과 불일치한다.' in str(ws['L1213'].value)
        for addr in ('L792', 'L811', 'L1212', 'L1213'):
            assert '후속 시험의 예상값' in str(ws[addr].value), addr
    for addr in ('H789', 'H809', 'K810', 'K811', 'H1208', 'K1212'):
        assert '26-05-02-13-14-15' in str(ws[addr].value), addr
    for addr in ('L810', 'L811', 'L1212'):
        assert '26-05-29-13-14-15' in str(ws[addr].value), addr
    assert '{26,5,2,13,14,15}' in str(ws['K1212'].value)
    assert '{26,5,29,13,14,15}' in str(ws['L1212'].value)
    for addr in ('K1213', 'L1213'):
        pairs = re.findall(r'(\d{2}-\d{2}-\d{2}-\d{2}-\d{2}-\d{2}) → (\d{2}-\d{2}-\d{2}-\d{2}-\d{2}-\d{2})', str(ws[addr].value))
        assert len(pairs) == 6, addr
        for start, end in pairs:
            assert datetime.strptime(start, '%y-%m-%d-%H-%M-%S') + timedelta(seconds=1) == datetime.strptime(end, '%y-%m-%d-%H-%M-%S'), (addr, start, end)

    before_parts = unpack(QA / 'report-before.xlsx')
    after_parts = unpack(path)
    assert before_parts.keys() == after_parts.keys()
    changed_parts = [n for n in before_parts if before_parts[n] != after_parts[n]]
    expected_parts = {'xl/worksheets/sheet2.xml'} if confirmed else {'xl/styles.xml', 'xl/worksheets/sheet2.xml'}
    assert set(changed_parts) == expected_parts, changed_parts
    bx = E.fromstring(before_parts['xl/worksheets/sheet2.xml'])
    ax = E.fromstring(after_parts['xl/worksheets/sheet2.xml'])
    bstyles = E.fromstring(before_parts['xl/styles.xml'])
    astyles = E.fromstring(after_parts['xl/styles.xml'])
    fill_addresses = addresses(plan)
    original_cells = {c.get('r'): c for c in bx.findall('.//s:c', NS)}
    for target in ax.findall('.//s:c', NS):
        addr = target.get('r')
        original = original_cells.get(addr)
        assert original is not None, addr
        if addr in fill_addresses:
            style_a = deepcopy(astyles.find(q('cellXfs'))[int(target.get('s', '0'))])
            style_b = bstyles.find(q('cellXfs'))[int(original.get('s', '0'))]
            style_a.set('fillId', style_b.get('fillId'))
            if 'applyFill' in style_b.attrib:
                style_a.set('applyFill', style_b.get('applyFill'))
            else:
                style_a.attrib.pop('applyFill', None)
            assert E.tostring(style_a, method='c14n', exclusive=True) == E.tostring(style_b, method='c14n', exclusive=True), addr
        if addr in plan['edits']:
            target.getparent().replace(target, deepcopy(original))
        elif addr in fill_addresses:
            target.set('s', original.get('s', '0'))
    for row in plan['heights']:
        original = bx.find(f'.//s:row[@r="{row}"]', NS)
        target = ax.find(f'.//s:row[@r="{row}"]', NS)
        target.attrib.clear()
        target.attrib.update(original.attrib)
    assert E.tostring(bx, method='c14n', exclusive=True) == E.tostring(ax, method='c14n', exclusive=True), 'Unexpected worksheet mutation'
    # Old fonts, borders and all existing styles remain intact; only append fill/xf entries.
    for tag in ('fills', 'cellXfs'):
        previous = bstyles.find(q(tag))
        current = astyles.find(q(tag))
        for node in list(current)[len(previous):]:
            current.remove(node)
        current.attrib.clear()
        current.attrib.update(previous.attrib)
    assert E.tostring(bstyles, method='c14n', exclusive=True) == E.tostring(astyles, method='c14n', exclusive=True), 'Unexpected existing style mutation'
    for name, wanted in {**baseline['protected'], **baseline['sources']}.items():
        assert sha(Path(name)) == wanted, 'Outside report changed: ' + name
    result = {'changed_cells': differences, 'yellow_ranges': plan['fills'],
              'changed_parts': changed_parts, 'rows': ws.max_row,
              'test_ids': sum(str(ws.cell(r, 4).value or '').strip().startswith('SW-FW-UnitTest-') for r in range(1, ws.max_row+1)),
              'criteria_date': plan['criteria_date'], 'result_date': plan['result_date'],
              'calendar_pairs_verified': 12, 'result_basis': plan.get('test_confirmation', 'planned expected values'),
              'previous_records_preserved_in_qa': True, 'confirmed_results': confirmed,
              'removed_marker_count': len(plan['clear_marker_fonts']),
              'protected_files': len(baseline['protected']), 'sources': len(baseline['sources']), 'sha256': sha(path)}
    (QA / ('published-checks.json' if path == REPORT else 'checks.json')).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf8')
    print(json.dumps(result, ensure_ascii=False))


if __name__ == '__main__':
    import sys
    if '--publish' in sys.argv:
        baseline = json.loads((QA / 'baseline.json').read_text(encoding='utf8'))
        assert sha(REPORT) == baseline['sha256'], 'User edited report after baseline'
        check()
        shutil.copyfile(QA / 'report.xlsx', REPORT)
        check(REPORT)
    else:
        preserve()
        check()
