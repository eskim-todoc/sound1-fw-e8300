"""통합/시스템 rev0 자료 추출, 원본 보존 및 독립 검사.

XLSX 값/서식 작성은 같은 이름의 Artifact Tool 스크립트에서 수행한다.
이 파일의 openpyxl 사용은 읽기/검사 전용이다.
"""
from pathlib import Path
import hashlib
import json
import re
import shutil
import sys
import math
import unicodedata
import copy
import posixpath
import zipfile
from lxml import etree as E
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
import preserve_aligned_report_packages as native

N=native.N
TAG='{'+native.NS+'}'

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
QA = TASK / 'qa/integration-system-rev0'
DHF = ROOT / 'docs/참고/인허가/사운드1/DHF'
RDIR = DHF / '04_설계 검증/SW/펌웨어'
CDIR = DHF / '03_설계 출력/SW/펌웨어'
PATHS = {
    'unit': RDIR/'SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서_rev0(김은수 작성).xlsx',
    'integration': RDIR/'SITR-FW-TD24A 소프트웨어 통합시험 보고서_rev0(김은수 작성).xlsx',
    'system': RDIR/'SSTR-FW-TD24A 소프트웨어 시스템 시험 보고서_rev0(김은수 작성).xlsx',
    'integrationCriterion': CDIR/'SITC-FW-TD24A 소프트웨어 통합시험 기준서.xlsx',
    'systemCriterion': CDIR/'[작성자 확인] SSTC-FW-TD24A 소프트웨어 시스템 시험 기준서.xlsx',
}

def digest(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(name,obj):
    (QA/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
def read(name): return json.loads((QA/name).read_text(encoding='utf-8'))
def txt(v): return str(v) if v is not None else ''
def pad(v): return '\n'+str(v).strip()+'\n' if str(v).strip() else ''

def text_lines(text,width):
    # 11pt Calibri/한글 대체 글꼴을 여유 있게 계산한다. 실제 Excel 재검사로 보완한다.
    capacity=max(4,width*0.85)
    return sum(max(1,math.ceil(sum(2 if unicodedata.east_asian_width(c) in 'WF' else 1 for c in line)/capacity)) for line in text.split('\n'))

def layout_integration():
    import integration_rev0_content as content
    base=read('baseline.json'); tests=content.make_tests(base)
    widths={'B':12,'C':27,'D':23,'E':34,'F':23,'G':23,'H':27,'I':6,'J':42,'K':32,'L':48,'M':64.625,'N':64.625,'O':13.875}
    rows=[];merges=[];heights=[];test_ranges=[];row=25
    for t in tests:
        start=row
        for case in t['cases']:
            # 논리 조건 하나는 한 행으로 표현한다. 한 셀이 Excel 높이 한계를 넘는 경우만 후속 분할한다.
            values=['']*14
            values[11]=pad(case['expected']);values[12]=pad(case['result'])
            if row==start:
                for col,v in t['meta'].items():values[ord(col)-66]=pad(v)
                values[13]=pad(t['status'])
            rows.append(values)
            h=max(text_lines(values[11],widths['M']),text_lines(values[12],widths['N']))*16.5+4
            assert h<=409.5,(t['id'],case['number'],h)
            heights.append(dict(row=row,height=max(66,h)))
            case['row']=row;row+=1
        end=row-1
        for col in 'BCDEFGHIJKLO':merges.append(f'{col}{start}:{col}{end}')
        # 병합된 식별/준비/절차 열도 전체 높이에 충분히 들어가게 한다.
        total=sum(x['height'] for x in heights if start<=x['row']<=end)
        needed=max(text_lines(pad(v),widths[col])*16.5+4 for col,v in t['meta'].items())
        if total<needed:
            add=(needed-total)/(end-start+1)
            for h in heights:
                if start<=h['row']<=end:h['height']+=add
        test_ranges.append(dict(id=t['id'],start=start,end=end,status=t['status'],cases=t['cases']))
    front={
      'N5':'기록: 시험자가 확인한 통과 결과를 입력 조건별 수치와 동작으로 정리했다.',
      'N10':'조건 [n]과 같은 행의 검증 [n]을 비교한다.\n짝수 시험 ID: RGB(245,245,245).',
      'D18':'E8300에 전달한 명령과 응답을 확인한다.',
      'C20':'격리된 시험 내부기',
      'D20':'내부기 실행/해제/재연결 및 자극 설정 완료 경로 확인',
      'C21':'I2S 오디오 입력 장치',
      'D21':'I2S 수신/혼합과 마이크 경로 전환 확인',
    }
    plan=dict(key='integration',sheet='보고서_통합시험',first=25,last=row-1,cols=15,rows=rows,
        merges=merges,heights=heights,widths=widths,tests=test_ranges,front=front,
        dropSheets=['시험 보조자료','예상 그림'],views=[dict(sheet='보고서_통합시험',range='M25:O28',name='start')])
    for tid,label in [(3,'copy'),(12,'state'),(15,'nofm'),(17,'ota')]:
        t=next(v for v in test_ranges if v['id']==tid)
        plan['views'].append(dict(sheet=plan['sheet'],range=f'M{t["start"]}:O{min(t["start"]+2,t["end"])}',name=label))
    save('integration-plan.json',plan);save('integration-content.json',tests)
    print('integration',len(tests),'tests',sum(len(t['cases']) for t in tests),'conditions',row-1,'rows')

def split_text(value,width):
    lines=value.strip().splitlines();chunks=[];current=[];used=0
    for line in lines:
        cost=text_lines(line,width)
        if used+cost>20 and current:
            chunks.append('\n'.join(current));current=[];used=0
        current.append(line);used+=cost
    if current:chunks.append('\n'.join(current))
    if not chunks:return ['']
    # 논리 셀의 처음/끝에만 빈 줄을 넣고 중간에 여백이나 '(계속)'을 넣지 않는다.
    chunks[0]='\n'+chunks[0];chunks[-1]+='\n'
    return chunks

def layout_system():
    base=read('baseline.json');source=base['books']['system']['sheets'];main='보고서_시스템 시험'
    content=json.loads((TASK/'system_rev0_content_plan.json').read_text(encoding='utf-8'))
    working=copy.deepcopy(source)
    for change in content['changes']:
        old=working[change['sheet']]['cells'].get(change['cell'],'')
        assert old.strip()==change['before'].strip(),('agent source mismatch',change['sheet'],change['cell'])
        working[change['sheet']]['cells'][change['cell']]=change['after']
    # 최종 대조: ID63의 통과 확인 기록에 상태 요약만 남지 않도록
    # 원래 기준/보고서에 있는 두 RTT 본문 문자열을 그대로 보존한다.
    log_lines='[BT] CMD 0x34, CHARGER STATE: 1, BATT LEVEL 37 PERCENT, LID STATE 1\n[BT] CMD 0x34, CHARGER STATE: 0, BATT LEVEL 72 PERCENT, LID STATE 2'
    for sheet,address in [(main,'K109'),('시험 보조자료','C285')]:
        value=working[sheet]['cells'][address]
        working[sheet]['cells'][address]=value.replace('충전/잔량/뚜껑 상태의 두 입력에 대응한 표시 확인.', 'RTT 출력 결과:\n'+log_lines+'\n충전/잔량/뚜껑 상태가 각각 (1,37,1), (0,72,2)로 확인.')
    for sheet,address in [(main,'J109'),('시험 보조자료','B285')]:
        working[sheet]['cells'][address]=working[sheet]['cells'][address].replace('ci_printv 본문 예상 전문(수집 로그 아님):','예상 RTT 출력 결과:')
    cells=working[main]['cells']; widths={'B':27,'C':54.625,'D':6,'E':24,'F':24,'G':42,'H':32,'I':48,'J':64.625,'K':64.625,'L':13.875}
    rows=[];merges=[];heights=[];tests=[];row=34;maps={}
    def cases(text,label):
        pattern=r'(?:^|\n)'+label+r' \[(\d+)\](?!의)'
        matches=list(re.finditer(pattern,text));prefix=text[:matches[0].start()].strip() if matches else ''
        found=[(int(m[1]),text[m.start():matches[i+1].start() if i+1<len(matches) else len(text)].strip()) for i,m in enumerate(matches)]
        assert found
        return prefix,found
    for oldrow in range(34,115):
        tid=int(cells['B'+str(oldrow)].strip().rsplit('-',1)[1]);start=row
        prefix,leftcases=cases(cells['J'+str(oldrow)],'조건')
        _,rightcases=cases(cells['K'+str(oldrow)],'(?:검증|측정)')
        assert [n for n,v in leftcases]==[n for n,v in rightcases],tid
        condition_rows=[]
        for ci,((number,expected),(_,actual)) in enumerate(zip(leftcases,rightcases)):
            if ci==0 and prefix:expected=expected+'\n\n'+prefix
            left=split_text(expected,widths['J']);right=split_text(actual,widths['K'])
            count=max(len(left),len(right));case_start=row
            for i in range(count):
                values=['']*11
                if row==start:
                    for c in 'BCDEFGHI':values[ord(c)-66]=pad(cells.get(c+str(oldrow),''))
                    values[10]=pad(cells['L'+str(oldrow)])
                if i<len(left):values[8]=left[i]
                if i<len(right):values[9]=right[i]
                rows.append(values)
                h=max(text_lines(values[8],widths['J']),text_lines(values[9],widths['K']))*16.5+18
                assert h<=409.5,('system long single line',tid,number,i,h)
                heights.append(dict(row=row,height=max(66,h)));row+=1
            if count>1:
                for col,chunks in [('J',left),('K',right)]:
                    if len(chunks)<count:merges.append(f'{col}{case_start+len(chunks)-1}:{col}{row-1}')
            condition_rows.append(dict(number=number,start=case_start,end=row-1))
        end=row-1
        if end>start:
            for col in 'BCDEFGHIL':merges.append(f'{col}{start}:{col}{end}')
        maps[oldrow]=(start,end)
        tests.append(dict(id=tid,start=start,end=end,sourceRow=oldrow,cases=condition_rows,status=cells['L'+str(oldrow)].strip()))
    footer_delta=row-115
    for oldrow in range(115,137):
        maps[oldrow]=(oldrow+footer_delta,oldrow+footer_delta)
        rows.append([cells.get(get_column_letter(c)+str(oldrow),'') for c in range(2,13)])
        heights.append(dict(row=oldrow+footer_delta,height=source[main]['heights'].get(str(oldrow)) or 21))
    front={a:v for a,v in cells.items() if int(re.search(r'\d+',a)[0])<34 and v!=source[main]['cells'].get(a,'')}
    for m in source[main]['merges']:
        x1,y1,x2,y2=range_boundaries(m)
        if y1>=115:merges.append(f'{get_column_letter(x1)}{y1+footer_delta}:{get_column_letter(x2)}{y2+footer_delta}')
    views=[dict(sheet=main,range=f'J{t["start"]}:L{t["end"]}',name='id'+str(t['id'])) for t in tests if t['id'] in [1,2,15,47,63,69,72,79,81]]
    views.extend([dict(sheet='시험 보조자료',range='A5:D8',name='photo-guide'),dict(sheet='예상 그림',range='A2:C7',name='expected-figure')])
    plan=dict(key='system',sheet=main,first=34,last=136+footer_delta,bodyLast=row-1,cols=12,rows=rows,
        merges=merges,heights=heights,widths=widths,tests=tests,front=front,dropSheets=[],views=views,
        helperChanges=[dict(sheet=name,cell=a,after=v) for name,sh in working.items() if name!=main for a,v in sh['cells'].items() if v!=source[name]['cells'].get(a,'')],rowMap=maps)
    save('system-plan.json',plan);print('system',len(tests),'tests',row-34,'body rows','max',max(x['height'] for x in heights))

def save_package(package,path):
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,data in package.items():z.writestr(name,data)

def workbook_names(package,sheet,area,titles):
    root=native.parse(package['xl/workbook.xml'])
    index=next(i for i,n in enumerate(root.find('s:sheets',N)) if n.get('name')==sheet)
    names=root.find('s:definedNames',N)
    if names is None:
        names=E.Element(TAG+'definedNames'); pos=list(root).index(root.find('s:sheets',N))+1;root.insert(pos,names)
    for name,text in [('_xlnm.Print_Area',f"'{sheet}'!{area}"),('_xlnm.Print_Titles',f"'{sheet}'!{titles}")]:
        node=next((n for n in names if n.get('name')==name and n.get('localSheetId')==str(index)),None)
        if node is None:node=E.SubElement(names,TAG+'definedName',name=name,localSheetId=str(index))
        node.text=text
    package['xl/workbook.xml']=native.serialize(root)

def remove_sheets(package,names):
    if not names:return
    wb=native.parse(package['xl/workbook.xml']); sheets=wb.find('s:sheets',N)
    rels=native.parse(package['xl/_rels/workbook.xml.rels'])
    removed={i:s for i,s in enumerate(sheets) if s.get('name') in names}
    rid={s.get('{'+native.RNS+'}id') for s in removed.values()}
    for s in removed.values():sheets.remove(s)
    for rel in list(rels):
        if rel.get('Id') in rid:rels.remove(rel)
    defs=wb.find('s:definedNames',N)
    if defs is not None:
        for node in list(defs):
            idx=int(node.get('localSheetId','-1'))
            if idx in removed:defs.remove(node)
            elif idx>=0:node.set('localSheetId',str(idx-sum(i<idx for i in removed)))
    for view in wb.findall('s:bookViews/s:workbookView',N):
        if int(view.get('activeTab','0'))>=len(sheets):view.set('activeTab','1')
    package['xl/workbook.xml']=native.serialize(wb)
    package['xl/_rels/workbook.xml.rels']=native.serialize(rels)
    reachable={'[Content_Types].xml','_rels/.rels'}
    pending=[r.get('Target').lstrip('/') for r in native.parse(package['_rels/.rels']) if r.get('TargetMode')!='External']
    while pending:
        part=pending.pop()
        if part in reachable:continue
        assert part in package,part
        reachable.add(part);rp=native.relpath(part)
        if rp in package:
            reachable.add(rp)
            pending.extend(native.resolve(part,r.get('Target')) for r in native.parse(package[rp]) if r.get('TargetMode')!='External')
    ct=native.parse(package['[Content_Types].xml'])
    for node in list(ct):
        if E.QName(node).localname=='Override' and node.get('PartName','').lstrip('/') not in reachable:ct.remove(node)
    package['[Content_Types].xml']=native.serialize(ct)
    # 앱 속성의 시트 목록도 실제 보존 시트와 맞춘다.
    if 'docProps/app.xml' in package:
        app=native.parse(package['docProps/app.xml']);ns={'e':'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties','v':'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'}
        titles=app.find('e:TitlesOfParts/v:vector',ns)
        if titles is not None:
            for child in list(titles):
                if child.text in names:titles.remove(child)
            titles.set('size',str(len(titles)))
        pairs=app.findall('e:HeadingPairs/v:vector/v:variant',ns)
        for i in range(0,len(pairs),2):
            label=pairs[i].find('v:lpstr',ns)
            if label is not None and label.text in ('워크시트','Worksheets'):pairs[i+1].find('v:i4',ns).text=str(len(sheets))
        package['docProps/app.xml']=native.serialize(app)
    for key in list(package):
        if key not in reachable:del package[key]

def replace_cell(root,source,address):
    new=source.find(f'.//s:c[@r="{address}"]',N)
    assert new is not None,address
    rownum=re.search(r'\d+',address)[0];data=root.find('s:sheetData',N)
    row=data.find(f's:row[@r="{rownum}"]',N)
    if row is None:row=E.SubElement(data,TAG+'row',r=rownum)
    old=row.find(f's:c[@r="{address}"]',N)
    if old is not None:row.replace(old,copy.deepcopy(new))
    else:
        pos=next((i for i,c in enumerate(row) if native.col(c.get('r'))>native.col(address)),len(row));row.insert(pos,copy.deepcopy(new))

def preserve_report(key):
    p=read(key+'-plan.json'); original=native.parts(QA/f'{key}-before.xlsx');authored=native.parts(QA/f'{key}-authored.xlsx')
    imp=native.Importer(original,authored,f'rev0_{key}')
    path=native.sheet_path(original,p['sheet']); ap=native.sheet_path(authored,p['sheet'])
    root=native.parse(original[path]); converted=imp.convert_sheet(native.parse(authored[ap]))
    if key=='integration':
        data=root.find('s:sheetData',N)
        for row in list(data):
            if int(row.get('r'))>=p['first']:data.remove(row)
        for row in converted.findall('s:sheetData/s:row',N):
            if p['first']<=int(row.get('r'))<=p['last']:data.append(copy.deepcopy(row))
        merges=root.find('s:mergeCells',N)
        for m in list(merges):
            if range_boundaries(m.get('ref'))[1]>=p['first']:merges.remove(m)
        for m in converted.findall('s:mergeCells/s:mergeCell',N):
            if range_boundaries(m.get('ref'))[1]>=p['first']:merges.append(copy.deepcopy(m))
        merges.set('count',str(len(merges)))
        root.find('s:dimension',N).set('ref',f'A1:O{p["last"]}')
        for a in list(p['front'])+['C19','D19']:replace_cell(root,converted,a)
        for tag in ['cols','sheetViews']:
            previous=root.find('s:'+tag,N); new=converted.find('s:'+tag,N)
            if previous is not None and new is not None:root.replace(previous,copy.deepcopy(new))
    else:
        old_data=copy.deepcopy(root.find('s:sheetData',N))
        data=root.find('s:sheetData',N)
        for row in list(data):
            if int(row.get('r'))>=p['first']:data.remove(row)
        for row in converted.findall('s:sheetData/s:row',N):
            if p['first']<=int(row.get('r'))<=p['last']:
                newrow=copy.deepcopy(row)
                if int(row.get('r'))>p['bodyLast']:
                    oldnumber=int(row.get('r'))-(p['last']-136)
                    for c in newrow:
                        oldaddr=re.sub(r'\d+$',str(oldnumber),c.get('r'))
                        oldcell=old_data.find(f'.//s:c[@r="{oldaddr}"]',N)
                        if oldcell is not None:c.set('s',oldcell.get('s','0'))
                data.append(newrow)
        merges=root.find('s:mergeCells',N)
        for m in list(merges):
            if range_boundaries(m.get('ref'))[1]>=p['first']:merges.remove(m)
        for m in converted.findall('s:mergeCells/s:mergeCell',N):
            if range_boundaries(m.get('ref'))[1]>=p['first']:merges.append(copy.deepcopy(m))
        merges.set('count',str(len(merges)))
        root.find('s:dimension',N).set('ref',f'A1:L{p["last"]}')
        for a in p['front']:replace_cell(root,converted,a)
        for tag in ['cols','sheetViews']:
            previous=root.find('s:'+tag,N);new=converted.find('s:'+tag,N)
            if previous is not None and new is not None:root.replace(previous,copy.deepcopy(new))
        for name in ['시험 보조자료','예상 그림']:
            hp=native.sheet_path(original,name);ahp=native.sheet_path(authored,name)
            hr=native.parse(original[hp]);acr=imp.convert_sheet(native.parse(authored[ahp]))
            for change in p['helperChanges']:
                if change['sheet']==name:replace_cell(hr,acr,change['cell'])
            if name=='시험 보조자료':
                for h in p.get('helperHeights',[]):
                    row=hr.find('s:sheetData/s:row[@r="'+str(h['row'])+'"]',N)
                    source_row=acr.find('s:sheetData/s:row[@r="'+str(h['row'])+'"]',N)
                    assert row is not None and source_row is not None
                    row.set('ht',source_row.get('ht'));row.set('customHeight','1')
            # 사진/도식과 관계는 원본 그대로, 변경 셀의 글자만 새 결과로 이식한다.
            mark_evidence_runs(hr)
            original[hp]=native.serialize(hr)
        mark_evidence_runs(root)
    original[path]=native.serialize(root);imp.flush()
    workbook_names(original,p['sheet'],f'$B$2:${get_column_letter(p["cols"])}${p["last"]}',f'${p["first"]-2}:${p["first"]-1}')
    remove_sheets(original,p.get('dropSheets',[]))
    save_package(original,QA/f'{key}-candidate.xlsx')
    print(key,'native cover/header/package preserved')

def mark_evidence_runs(root):
    """Artifact에서 작성한 텍스트를 보존하고 사진/QCC 표시의 부분 글꼴만 복원한다."""
    for cell in root.findall('.//s:c',N):
        ins=cell.find('s:is',N)
        if ins is not None:value=''.join(ins.itertext())
        elif cell.get('t')=='str':
            value_node=cell.find('s:v',N)
            value=value_node.text or '' if value_node is not None else ''
        else:continue
        if '(사진필요)' not in value and '(QCC 검증 결과 필요)' not in value:continue
        if ins is None:
            cell.remove(value_node);cell.set('t','inlineStr');ins=E.SubElement(cell,TAG+'is')
        for child in list(ins):ins.remove(child)
        for line in value.splitlines(keepends=True):
            run=E.SubElement(ins,TAG+'r');pr=E.SubElement(run,TAG+'rPr')
            E.SubElement(pr,TAG+'rFont',val='맑은 고딕');E.SubElement(pr,TAG+'sz',val='10')
            color='FFC00000' if '(QCC 검증 결과 필요)' in line else ('FF0000FF' if '(사진필요)' in line else 'FF000000')
            E.SubElement(pr,TAG+'color',rgb=color)
            if color!='FF000000':E.SubElement(pr,TAG+'b')
            text=E.SubElement(run,TAG+'t');text.set('{http://www.w3.org/XML/1998/namespace}space','preserve');text.text=line

def fit_native(key):
    p=read(key+'-plan.json'); all_measurements=read(key+'-native-heights.json')['measurements']
    body_last=p.get('bodyLast',p['last'])
    m=[c for c in all_measurements if c['end']<=body_last]
    single={r:48.0 for r in range(p['first'],body_last+1)}
    for c in m:
        if c['start']==c['end']:single[c['start']]=max(single[c['start']],c['needed']+18)
    for c in m:
        if c['start']!=c['end']:
            total=sum(single[r] for r in range(c['start'],c['end']+1))
            if total<c['needed']+18:
                delta=(c['needed']+18-total)/(c['end']-c['start']+1)
                for r in range(c['start'],c['end']+1):single[r]+=delta
    footer=[h for h in p['heights'] if h['row']>body_last]
    for h in footer:
        needed=max([c['needed']+6 for c in all_measurements if c['start']==c['end']==h['row']]+[h['height']])
        h['height']=math.ceil(needed/1.5)*1.5
    p['heights']=[dict(row=r,height=math.ceil(h/1.5)*1.5) for r,h in single.items()]+footer
    if key=='system':
        helper={}
        for c in read(key+'-native-heights.json').get('helperMeasurements',[]):
            assert c['start']==c['end']
            helper[c['start']]=max(helper.get(c['start'],0),c['current'],c['needed']+18)
        p['helperHeights']=[dict(row=r,height=math.ceil(h/1.5)*1.5) for r,h in helper.items()]
        assert max([h['height'] for h in p['helperHeights']]+[0])<=409.5
    assert max(x['height'] for x in p['heights'])<=409.5
    save(key+'-plan.json',p);print('native heights',key,len(single),max(single.values()))

def plan_criterion(key):
    rp=read(key+'-plan.json'); source=load_workbook(QA/f'{key}-candidate.xlsx',rich_text=True)[rp['sheet']]
    ck=key+'Criterion'; old=read('baseline.json')['books'][ck]
    name=next(n for n in old['sheets'] if n.startswith('기준서_'))
    limit=rp['cols']-2
    assert not any(m.min_col<=limit<m.max_col for m in source.merged_cells.ranges)
    views=[dict(sheet=name,range=v['range'].replace('O','M').replace('N','M') if key=='integration' else v['range'].replace('L','J').replace('K','J'),name=v['name']) for v in rp['views'] if v['sheet']==rp['sheet']]
    p=dict(key=ck,sheet=name,reportSheet=rp['sheet'],first=rp['first'],last=rp['last'],cols=limit,
      unmerges=[str(m) for m in source.merged_cells.ranges if m.min_col>limit],title=str(source['B2'].value).replace('보고서','기준서'),views=views,
      tests=rp.get('tests',[]))
    save(ck+'-plan.json',p);print(ck,'remove',get_column_letter(limit+1),get_column_letter(limit+2))

def preserve_criterion(key):
    ck=key+'Criterion';p=read(ck+'-plan.json')
    source=native.parts(QA/f'{key}-candidate.xlsx');target=native.parts(QA/f'{ck}-before.xlsx')
    sp=native.sheet_path(source,p['reportSheet']);tp=native.sheet_path(target,p['sheet'])
    imp=native.Importer(target,source,ck+'_report')
    root=imp.convert_sheet(native.parse(source[sp])); limit=p['cols']
    assert not root.findall('.//s:f',N) and not root.findall('.//s:drawing',N)
    for row in root.findall('s:sheetData/s:row',N):
        for cell in list(row):
            if native.col(cell.get('r'))>limit:row.remove(cell)
        row.attrib.pop('spans',None)
    merges=root.find('s:mergeCells',N)
    if merges is not None:
        for node in list(merges):
            if range_boundaries(node.get('ref'))[0]>limit:merges.remove(node)
        merges.set('count',str(len(merges)))
    cols=root.find('s:cols',N)
    for c in list(cols):
        if int(c.get('min'))>limit:cols.remove(c)
        elif int(c.get('max'))>limit:c.set('max',str(limit))
    root.find('s:dimension',N).set('ref',f'A1:{get_column_letter(limit)}{p["last"]}')
    authored=load_workbook(QA/f'{ck}-authored.xlsx',rich_text=True)[p['reportSheet']]
    assert txt(authored['B2'].value)==p['title']
    assert not any(c.value is not None for row in authored.iter_rows(min_col=limit+1,max_col=limit+2) for c in row)
    title=root.find('.//s:c[@r="B2"]',N)
    for node in list(title):title.remove(node)
    title.set('t','inlineStr');text=E.SubElement(E.SubElement(title,TAG+'is'),TAG+'t')
    text.set('{http://www.w3.org/XML/1998/namespace}space','preserve');text.text=p['title']
    # main에는 그림이 없으며 프린터 관계만 존재한다. 원본 보고서의 인쇄 설정을 복사한다.
    if native.relpath(sp) in source:
        rels=native.parse(source[native.relpath(sp)])
        for rel in rels:
            assert rel.get('Type').endswith('/printerSettings'),rel.get('Type')
            old=native.resolve(sp,rel.get('Target'));mapped=imp.copy_part(old)
            rel.set('Target',posixpath.relpath(mapped,posixpath.dirname(tp)))
        target[native.relpath(tp)]=native.serialize(rels)
    else:target.pop(native.relpath(tp),None)
    phonetic=root.find('s:phoneticPr',N)
    if phonetic is not None and 'fontId' in phonetic.attrib:phonetic.set('fontId',str(imp.maps['fonts'][int(phonetic.get('fontId'))]))
    target[tp]=native.serialize(root);imp.flush()
    workbook_names(target,p['sheet'],f'$B$2:${get_column_letter(limit)}${p["last"]}',f'${p["first"]-2}:${p["first"]-1}')
    save_package(target,QA/f'{ck}-candidate.xlsx');print(ck,'report common columns copied, original cover/history preserved')

def verify(key):
    p=read(key+'-plan.json');base=read('baseline.json');before=load_workbook(QA/f'{key}-before.xlsx',rich_text=True)
    after=load_workbook(QA/f'{key}-candidate.xlsx',rich_text=True);s=after[p['sheet']]
    old=native.parts(QA/f'{key}-before.xlsx');new=native.parts(QA/f'{key}-candidate.xlsx')
    assert old[native.sheet_path(old,'표지')]==new[native.sheet_path(new,'표지')],'표지 XML 변경'
    assert old['xl/theme/theme1.xml']==new['xl/theme/theme1.xml'],'테마 변경'
    for row in before['표지']:
        for c in row:
            d=after['표지'][c.coordinate]
            assert txt(c.value)==txt(d.value),('cover',c.coordinate)
            for attr in ['font','fill','border','alignment','number_format']:assert copy.copy(getattr(c,attr))==copy.copy(getattr(d,attr)),('cover',c.coordinate,attr)
    ck=key+'Criterion'; cp=read(ck+'-plan.json');crit=load_workbook(QA/f'{ck}-candidate.xlsx',rich_text=True)[cp['sheet']]
    explicit={c.get('r') for c in native.parse(new[native.sheet_path(new,p['sheet'])]).findall('.//s:c',N)}
    checked=0
    for row in s.iter_rows(max_col=cp['cols']):
        for c in row:
            d=crit[c.coordinate]
            expected=cp['title'] if c.coordinate=='B2' else txt(c.value)
            assert txt(d.value)==expected,('criterion value',c.coordinate)
            for attr in ['font','fill','border','alignment','number_format']:
                if c.coordinate in explicit and c.__class__.__name__!='MergedCell':assert copy.copy(getattr(c,attr))==copy.copy(getattr(d,attr)),('criterion style',c.coordinate,attr)
            checked+=1
    assert crit.max_column==cp['cols']
    assert set(map(str,crit.merged_cells.ranges))=={str(m) for m in s.merged_cells.ranges if m.max_col<=cp['cols']}
    for r in range(1,p['last']+1):assert s.row_dimensions[r].height==crit.row_dimensions[r].height,('criterion height',r)
    for col in range(1,cp['cols']+1):
        letter=get_column_letter(col);assert s.column_dimensions[letter].width==crit.column_dimensions[letter].width,('criterion width',letter)
    criterion_old=native.parts(QA/f'{ck}-before.xlsx');criterion_new=native.parts(QA/f'{ck}-candidate.xlsx')
    for name in base['books'][ck]['sheets']:
        if name!=cp['sheet']:assert criterion_old[native.sheet_path(criterion_old,name)]==criterion_new[native.sheet_path(criterion_new,name)],('criterion retained sheet',name)
    idcol='C' if key=='integration' else 'B';expectedcol='M' if key=='integration' else 'J';actualcol='N' if key=='integration' else 'K';verdictcol='O' if key=='integration' else 'L'
    old_ids=[txt(c.value).strip() for row in before[p['sheet']].iter_rows(min_col=native.col(idcol+'1'),max_col=native.col(idcol+'1')) for c in row if txt(c.value).strip().startswith('SW-FW-')]
    new_ids=[txt(c.value).strip() for row in s.iter_rows(min_col=native.col(idcol+'1'),max_col=native.col(idcol+'1')) for c in row if txt(c.value).strip().startswith('SW-FW-')]
    assert old_ids==new_ids and len(set(new_ids))==len(new_ids)
    bad=[];logs=0;case_count=0
    banned=r'결과 대조 필요|결과 기입 필요|기록 확인이 필요|기존 기록 불일치|출력\s*:\s*출력\s*:|^RTT\s+C\d|해야 한다|기록한다|저장한다|확인한다|기대입니다|원본에는|시험 입력에 대응한 값'
    for t in p['tests']:
        assert txt(s[f'{verdictcol}{t["start"]}'].value).strip()=='통과'
        if key=='system':
            for case in t['cases']:
                expected='\n'.join(txt(s[f'J{r}'].value) for r in range(case['start'],case['end']+1)).strip()
                actual='\n'.join(txt(s[f'K{r}'].value) for r in range(case['start'],case['end']+1)).strip()
                number=case['number']
                assert expected.startswith(f'조건 [{number}]'),('system condition',t['id'],number)
                assert re.match(r'(검증|측정) \['+str(number)+r'\]',actual),('system result',t['id'],number)
                assert actual!=expected,('system criteria pasted',t['id'],number)
                case_count+=1
                raw=lambda v:re.findall(r'^\[(?:UT-[^\]]+|BT)\][^\r\n]+',v,re.M)
                if raw(expected):assert raw(expected)==raw(actual),('system RTT mismatch',t['id'],number)
                logs+=len(raw(actual))
        for r in range(t['start'],t['end']+1):
            actual=txt(s[f'{actualcol}{r}'].value);expected=txt(s[f'{expectedcol}{r}'].value)
            if actual.strip():
                if re.search(banned,actual,re.M):bad.append((r,actual))
                assert len(re.sub(r'(검증|측정)\s*\[[^]]+\]|검증 결과|통과|[\s\W]','',actual))>=5,('empty result',r)
                assert actual.strip()!=expected.strip(),('criteria pasted',r)
            if key=='integration':
                n=r-t['start']+1
                assert expected.strip().startswith(f'조건 [{n}]'),('condition',r)
                assert actual.strip().startswith(f'검증 [{n}]'),('verification',r)
                case_count+=1
                raw=lambda value:re.findall(r'^\[UT-[^\r\n]+',value,re.M)
                if raw(expected):assert raw(expected)==raw(actual),('RTT mismatch',r)
                logs+=len(raw(actual))
            fill='FFF5F5F5' if t['id']%2==0 else 'FFFFFFFF'
            assert s[f'{expectedcol}{r}'].fill.fgColor.rgb==fill,('stripe',r)
    assert not bad,bad
    for row in s:
        for c in row:
            if c.data_type=='e':raise AssertionError(('formula error',c.coordinate,c.value))
    assert max((d.height or 0) for d in s.row_dimensions.values())<=409.5
    if key=='integration':
        assert after.sheetnames==['표지','보고서_통합시험']
        assert not any('사진' in txt(c.value) or '시험 보조자료' in txt(c.value) or '예상 그림' in txt(c.value) for row in s for c in row)
        assert sum(len(sh._images) for sh in after)==0
    renamed=[]
    for path,h in base['protected'].items():
        if Path(path).exists():assert digest(path)==h,('protected changed',path)
        else:
            matches=[str(f) for f in Path(path).parent.iterdir() if f.is_file() and f.suffix==Path(path).suffix and not f.name.startswith('~$') and digest(f)==h]
            assert len(matches)==1,('protected missing or changed',path,matches)
            renamed.append(dict(before=path,current=matches[0],hash=h))
    if renamed:save('external-renames.json',renamed)
    for k in ['integration','system','integrationCriterion','systemCriterion']:assert digest(PATHS[k])==base['books'][k]['hash'],('externally edited target',k)
    out=dict(testCount=len(new_ids),conditions=case_count,rawLogLines=logs,commonCells=checked,
        rows=s.max_row,merges=len(s.merged_cells.ranges),criterionCommonColumnsIdentical=True,
        coverThemePreserved=True,protectedFiles=len(base['protected']),allVerdictsPassedByTesterConfirmation=True,
        hash=digest(QA/f'{key}-candidate.xlsx'))
    save(key+'-verification.json',out);print(json.dumps(out,ensure_ascii=False))

def prepare():
    QA.mkdir(parents=True,exist_ok=False)
    books={}
    for key,p in PATHS.items():
        backup=QA/f'{key}-before.xlsx'
        shutil.copy2(p,backup)
        w=load_workbook(p,rich_text=True)
        sheets={}
        for s in w:
            cells={c.coordinate:txt(c.value) for row in s for c in row if c.value is not None}
            sheets[s.title]=dict(rows=s.max_row,columns=s.max_column,merges=list(map(str,s.merged_cells.ranges)),
                images=len(s._images),cells=cells,widths={k:v.width for k,v in s.column_dimensions.items()},
                heights={str(k):v.height for k,v in s.row_dimensions.items()},
                printArea=str(s.print_area),printTitles=s.print_title_rows)
        books[key]=dict(path=str(p),backup=str(backup),hash=digest(p),sheets=sheets)
    protect={str(p):digest(p) for p in DHF.rglob('*') if p.is_file() and p.suffix in ('.xlsx','.docx') and p not in list(PATHS.values())[1:] and not p.name.startswith('~$')}
    save('baseline.json',dict(books=books,protected=protect))
    for key,b in books.items():
        print(key,[(n,s['rows'],s['columns'],s['images']) for n,s in b['sheets'].items()])

def dump(key,sheet,start,end):
    b=read('baseline.json')['books'][key]
    if sheet=='main': sheet=next(n for n in b['sheets'] if n.startswith(('보고서_','기준서_')))
    s=b['sheets'][sheet]
    for r in range(int(start),int(end)+1):
        print(f'--- {sheet} row {r} ---')
        for a,v in s['cells'].items():
            if int(re.search(r'\d+',a)[0])==r: print(a, v)

if __name__=='__main__':
    if sys.argv[1]=='prepare': prepare()
    elif sys.argv[1]=='plan-integration': layout_integration()
    elif sys.argv[1]=='plan-system': layout_system()
    elif sys.argv[1]=='preserve': preserve_report(sys.argv[2])
    elif sys.argv[1]=='fit': fit_native(sys.argv[2])
    elif sys.argv[1]=='plan-criterion': plan_criterion(sys.argv[2])
    elif sys.argv[1]=='preserve-criterion': preserve_criterion(sys.argv[2])
    elif sys.argv[1]=='verify': verify(sys.argv[2])
    elif sys.argv[1]=='dump': dump(*sys.argv[2:])
