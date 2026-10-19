import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[4]
TASK = Path(__file__).parent
REF = Path(r'G:\공유 드라이브\20_프로젝트별\01_project sound1\04_docs\99_참고용DHF\04_설계 검증\SW\펌웨어\SUVR-FW-TD20B 소프트웨어 유닛 검증 보고서.xlsx')
CRITERIA = ROOT / 'docs/참고/인허가/사운드1/DHF/03_설계 출력/SW/펌웨어/SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx'
QA = TASK / 'qa/suvr'
QA.mkdir(parents=True, exist_ok=True)
for label, path in [('reference', REF), ('criteria', CRITERIA)]:
    w = openpyxl.load_workbook(path, data_only=False)
    result = {}
    for s in w:
        rows = [[c.value for c in r] for r in s]
        result[s.title] = {'rows': rows, 'merges': [str(m) for m in s.merged_cells.ranges],
            'widths': {k: v.width for k,v in s.column_dimensions.items()},
            'heights': {k: v.height for k,v in s.row_dimensions.items() if v.height},
            'print_area': str(s.print_area)}
        print(label, s.title, s.max_row, s.max_column, 'merges',len(s.merged_cells.ranges))
        for i,r in enumerate(rows[:15],1):
            if any(v is not None for v in r): print(i, str(r)[:2000])
    (QA / f'{label}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
