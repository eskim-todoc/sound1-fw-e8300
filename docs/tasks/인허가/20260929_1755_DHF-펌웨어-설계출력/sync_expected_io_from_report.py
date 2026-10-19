"""Export the reviewed report to Markdown comparison documents, without editing XLSX.

This is document synchronization, not a test run or an independent oracle.
Source cell markers permit a separate XML-based verifier to check every copied cell.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
from collections import Counter
from pathlib import Path
from urllib.parse import quote
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
DHF = ROOT / 'docs/참고/인허가/사운드1/DHF'
FIRMWARE = DHF / '04_설계 검증/SW/펌웨어'
REPORT = FIRMWARE / 'SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서_1차본.xlsx'
DEST = FIRMWARE / '유닛 시험 예상 입출력'
QA = TASK / 'qa/expected-io-from-report'
SHEET = '보고서_유닛 시험'
ID_RE = re.compile(r'SW-FW-UnitTest-(CFX|HEAR|CM3|BOOT)-ID-(\d+)')
LOG_RE = re.compile(r'^\[UT-(?:CFX|HEAR|CM3|BOOT)-\d+\] C\d+ (?:A\[.*|I\d+ .*|(?:PASS|FAIL|OBSERVED) .*)$')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(value):
    return '' if value is None else str(value).replace('\r\n', '\n').replace('\r', '\n').strip('\n')


def link(label, target):
    return f'[{label}]({quote(str(target).replace(chr(92), "/"), safe="/.-_")})'


def inline(line):
    # Workbook text is literal, not Markdown or HTML. Preserve what a reader sees.
    return re.sub(r'([\\`*_{}\[\]#|])', r'\\\1', html.escape(line, quote=False))


def cell_block(address, value):
    lines = value.split('\n')
    rendered = []
    in_log = False
    for line in lines:
        is_log = bool(LOG_RE.fullmatch(line.strip()))
        if is_log and not in_log:
            rendered.extend(['', '```text'])
        elif in_log and not is_log:
            rendered.extend(['```', ''])
        rendered.append(line if is_log else inline(line) + ('<br>' if line else ''))
        in_log = is_log
    if in_log:
        rendered.append('```')
    return f'<!-- source-cell:{address} -->\n' + '\n'.join(rendered) + f'\n<!-- /source-cell:{address} -->\n'


def read_source():
    wb = load_workbook(REPORT, rich_text=True, data_only=False)
    ws = wb[SHEET]
    cells = {c.coordinate: text(c.value) for row in ws for c in row if c.value is not None}
    d_merges = {m.min_row: m.max_row for m in ws.merged_cells if m.min_col == 4 and m.max_col == 4}
    units = []
    for row in range(1, ws.max_row + 1):
        m = ID_RE.fullmatch(cells.get(f'D{row}', '').strip())
        if not m:
            continue
        core, num = m[1], int(m[2])
        end = d_merges.get(row, row)
        design = cells.get(f'G{row}', '')
        name = re.search(r'\((.*?)\)', design, re.S)
        # Joint K/L boundaries keep the parts of a merged criterion/result together.
        cuts = set(range(row, end + 2))
        for merged in ws.merged_cells:
            if merged.min_col <= 12 and merged.max_col >= 11 and row <= merged.min_row <= end:
                cuts.difference_update(range(merged.min_row + 1, min(merged.max_row, end) + 1))
        ordered = sorted(cuts)
        groups = []
        for a, b in zip(ordered, ordered[1:]):
            pairs = {col: [f'{col}{r}' for r in range(a, b) if cells.get(f'{col}{r}', '')]
                     for col in ('K', 'L')}
            if pairs['K'] or pairs['L']:
                groups.append({'first': a, 'last': b - 1, **pairs})
        units.append({'unit': f'{core}-{num:03d}', 'core': core, 'number': num,
                      'id': m[0], 'name': name[1] if name else design,
                      'first': row, 'last': end, 'groups': groups})
    wb.close()
    assert len(units) == 115 and len({u['id'] for u in units}) == 115
    return cells, units


def protect_hashes():
    paths = [p for p in DHF.rglob('*') if p.is_file() and p.suffix.lower() in ('.xlsx', '.docx')
             and not p.name.startswith('~$')]
    for folder in ('src', 'tests'):
        paths.extend(p for p in (ROOT / folder).rglob('*')
                     if p.is_file() and p.suffix.lower() in ('.c', '.h', '.py'))
    return {p.relative_to(ROOT).as_posix(): sha(p) for p in sorted(set(paths))}


def unit_doc(unit, cells):
    r = unit['first']
    header = [f"# {unit['unit']} {unit['name']}", '',
              f"시험 ID: `{unit['id']}`  ",
              f"원본: {link('유닛 검증 보고서 1차본', '../../' + REPORT.name)}, `{SHEET}` {r}~{unit['last']}행.  ",
              '아래 예상 결과는 보고서에 기록된 결과를 그대로 옮긴 비교값이다.', '']
    for col, title in [('E', '유닛 설명'), ('F', '시험 방법'), ('G', '상세 설계 유닛'),
                       ('H', '시험 준비와 입력'), ('I', '가정 및 제한 조건'), ('J', '시험 절차')]:
        addresses = [f'{col}{i}' for i in range(r, unit['last'] + 1) if cells.get(f'{col}{i}')]
        if addresses:
            header.extend([f'## {title}', ''])
            for addr in addresses:
                header.extend([cell_block(addr, cells[addr]), ''])
    header.extend(['## 입력 조건별 기준과 예상 결과', ''])
    for n, group in enumerate(unit['groups'], 1):
        header.extend([f'### 비교 {n}', ''])
        for col, title in [('K', '입력 조건과 통과 기준'), ('L', '예상 결과 (보고서 기록값)')]:
            if group[col]:
                header.extend([f'#### {title}', ''])
                for addr in group[col]:
                    header.extend([cell_block(addr, cells[addr]), ''])
    header.extend(['## 보고서 판정', '', cell_block(f'M{r}', cells[f'M{r}']), ''])
    return '\n'.join(header).rstrip() + '\n'


def log_doc(unit, cells):
    parts = [f"# {unit['unit']} 예상 RTT 출력", '',
             link('입력 조건과 전체 예상 결과', f"../{unit['core']}/{unit['unit']}.md"), '',
             '보고서 1차본의 결과 칸에서 RTT 관련 기록을 옮겼다. 조건 번호, 배열값과 줄 나눔은 원문과 같다.', '']
    addresses = [f'L{r}' for r in range(unit['first'], unit['last'] + 1)
                 if '[UT-' in cells.get(f'L{r}', '')]
    for i, addr in enumerate(addresses, 1):
        parts.extend([f'## RTT 비교 {i}', '', cell_block(addr, cells[addr]), ''])
    if not addresses:
        parts.extend(['현재 보고서에는 이 유닛의 전체 RTT 문자열이 없다. 위 유닛 문서의 입력 조건과 결과값으로 비교한다.', ''])
    return '\n'.join(parts).rstrip() + '\n'


def make_outputs(cells, units):
    out = {}
    log_units = set(p.stem for p in (DEST / '예상 로그').glob('*.md') if re.fullmatch(r'(CFX|HEAR|CM3|BOOT)-\d+', p.stem))
    log_units.update(u['unit'] for u in units if any('[UT-' in cells.get(f'L{r}', '') for r in range(u['first'], u['last'] + 1)))
    for unit in units:
        out[f"{unit['core']}/{unit['unit']}.md"] = unit_doc(unit, cells)
        if unit['unit'] in log_units:
            out[f"예상 로그/{unit['unit']}.md"] = log_doc(unit, cells)
    counts = Counter(u['core'] for u in units)
    intro = ['# 예상 테스트 입력과 출력', '',
             f"2026-10-08 기준 {link('유닛 검증 보고서 1차본', '../' + REPORT.name)}의 `{SHEET}`를 원본으로 정리했다.", '',
             '115개 시험의 입력 조건, 통과 기준과 결과를 같은 시험 ID로 연결했다. 예상 결과는 보고서의 결과 칸을 그대로 옮긴 값이며, 이번 동기화에서 새로 측정하거나 독립 계산한 값은 아니다.', '',
             '보고서에 있는 수치, 단위, 조건 번호와 RTT 문자열을 유지한다. 보고서에 없는 이전 출력 형식과 보충 시험은 현재 비교 목록에 포함하지 않는다. 통과 기준의 허용 범위와 결과 칸의 단일 값은 원본처럼 구분한다.', '',
             '## 찾는 방법', '',
             f"1. 아래 목록에서 시험 ID를 선택한다. CFX {counts['CFX']}개, HEAR {counts['HEAR']}개, CM3 {counts['CM3']}개, BOOT {counts['BOOT']}개다.",
             '2. 시험 준비와 입력, 시험 절차를 확인한다.',
             '3. 같은 비교 항목의 입력 조건과 예상 결과를 나란히 확인한다.',
             '4. RTT만 확인할 때는 로그 보기를 연다. C99의 로그 출력 없음도 보고서와 같은 정상 결과로 읽는다.', '',
             link('실행과 로그 읽기', '공통/실행과 로그 읽기.md') + ' / ' + link('유닛 시험 일괄 실행 매뉴얼', '../유닛 시험 일괄 실행 매뉴얼.md'), '',
             '## 유닛 목록', '', '| 시험 | 유닛 | 방법 | RTT |', '|---|---|---|---|']
    for u in units:
        log = link('로그 보기', f"예상 로그/{u['unit']}.md") if u['unit'] in log_units else '본문 결과'
        intro.append(f"| {link(u['unit'], u['core'] + '/' + u['unit'] + '.md')} | {inline(u['name'])} | {cells.get('F'+str(u['first']), '').strip()} | {log} |")
    intro.extend(['', '## QCC 자료 추가 위치', '',
                  '현재 115개는 보고서의 E8300 유닛 시험 범위다. 향후 QCC 내부 유닛이 추가되면 기준서와 보고서에 먼저 추가하고 이 목록에 반영한다.', '',
                  link('QCC 자료 추가 위치 관리', '../../../../../QCC 자료 추가 위치 관리.md') + '의 6절과 7절에서 추가 위치를 관리한다.', ''])
    out['README.md'] = '\n'.join(intro)
    log_index = ['# 보고서와 비교하는 예상 RTT 출력', '',
                 '유닛별 문서에는 보고서 결과 칸의 RTT 기록만 담았다. 입력 조건과 수치 설명은 같은 유닛의 전체 문서에서 확인한다.', '',
                 link('전체 입력과 예상 결과', '../README.md'), '', '| 유닛 | 예상 RTT 출력 |', '|---|---|']
    for u in units:
        if u['unit'] in log_units:
            log_index.append(f"| {u['unit']} {inline(u['name'])} | {link('로그 보기', u['unit'] + '.md')} |")
    out['예상 로그/README.md'] = '\n'.join(log_index) + '\n'
    out['공통/실행과 로그 읽기.md'] = '''# 실행과 로그 읽기

각 유닛 문서의 입력 조건, 통과 기준과 예상 결과는 보고서 1차본의 같은 시험에서 가져왔다. 예상 결과의 원본은 보고서의 시험/측정 내용 및 결과 칸이다.

## 실행 순서

1. [유닛 시험 일괄 실행 매뉴얼](../../유닛%20시험%20일괄%20실행%20매뉴얼.md)에서 실행할 그룹을 확인한다.
2. [전체 목록](../README.md)에서 같은 유닛을 열고 시험 준비와 입력을 확인한다.
3. 실행한 조건에 해당하는 결과값과 [예상 RTT 출력](../예상%20로그/README.md)을 비교한다.

## RTT 표기

| 표기 | 의미 |
|---|---|
| UT-CM3-027 | CM3 유닛 시험 27번 |
| C1 | 해당 시험 코드의 입력 조건 번호 |
| A[0:2] | 배열의 0번부터 2번까지 출력한 값 |
| N | 비교값 개수 |
| F | 불일치 개수 |
| PASS / FAIL | 해당 조건의 비교 통과 / 실패 |
| OBSERVED | 관찰값 기록 |

CM3-027의 현재 비교값은 [해당 유닛의 RTT 출력](../예상%20로그/CM3-027.md)에서 확인한다. 이전의 두 값 형식은 현재 비교 목록에서 제외했다.

HEAR-001/002의 C99는 시간 초과 오류 기록이다. 이번 보고서에 대응하는 정상 결과는 C99 로그 출력 없음이다. 이 상태를 PASS 문자열로 새로 만들지 않는다.

## 시험 방법

| 표시 | 약어 | 의미 |
|---|---|---|
| I | Inspection | 제어 또는 시각적 확인 |
| A | Analysis | 분석적인 증거에 기반한 확인 |
| D | Demonstration | 정량적인 측정 없이 동작 특성 확인 |
| T | Test | 정량적인 측정을 통한 정량적인 특성 확인 |
'''
    # Keep existing paths usable, but do not mix test-tool examples with product results.
    for filename, title in [('시험 도구 자체의 입력과 예상 결과.md', '시험 도구 자체의 입력과 예상 결과'),
                            ('감사기록.md', '감사기록'), ('전체 검사 대응표.md', '전체 검사 대응표')]:
        out['공통/' + filename] = f'''# {title}

현재 비교값은 [보고서 기준 전체 유닛 목록](../README.md)과 [유닛별 RTT 출력](../예상%20로그/README.md)에서 확인한다.

이 문서에 있던 시험 도구 자체의 보충 시험은 유닛 보고서 1차본의 115개 시험과 별도 범위이므로 현재 결과 목록에서 분리했다. 이전 자료는 작업 폴더의 복구본에 보존했다. 실제 실행 순서는 [일괄 실행 매뉴얼](../../유닛%20시험%20일괄%20실행%20매뉴얼.md)을 따른다.
'''
    out['예상 로그/공통-시험도구.md'] = '''# 공통 시험 도구 로그

제품 유닛의 비교 로그는 [현재 RTT 목록](README.md)에서 확인한다. 보고서에 없는 시험 도구용 가상 번호와 오류 주입 예시는 현재 목록에서 분리했다. 이전 내용은 작업 폴더의 복구본에 보존했다.
'''
    return out


def build():
    QA.mkdir(parents=True, exist_ok=True)
    baseline = QA / 'baseline.json'
    if baseline.exists():
        raise RuntimeError('Baseline already exists; reuse the saved build or choose an explicit new baseline.')
    original = {p.relative_to(DEST).as_posix(): sha(p) for p in DEST.rglob('*') if p.is_file()}
    assert all(p.endswith('.md') for p in original)
    with ZipFile(QA / 'expected-io-before.zip', 'x', ZIP_DEFLATED) as z:
        for name in original:
            z.write(DEST / name, name)
    info = {'report': str(REPORT), 'report_sha256': sha(REPORT), 'original_files': original,
            'protected': protect_hashes()}
    baseline.write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding='utf-8')
    cells, units = read_source()
    outputs = make_outputs(cells, units)
    assert set(original).issubset(outputs), set(original) - set(outputs)
    for name, content in outputs.items():
        p = QA / 'candidate' / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(content, encoding='utf-8', newline='\n')
    data = {'units': units, 'files': {name: hashlib.sha256(content.encode()).hexdigest() for name, content in outputs.items()}}
    (QA / 'manifest.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'units': len(units), 'files': len(outputs), 'report_sha256': info['report_sha256'],
                      'backup_files': len(original), 'protected_files': len(info['protected'])}, ensure_ascii=False))


def publish():
    baseline = json.loads((QA / 'baseline.json').read_text(encoding='utf-8'))
    manifest = json.loads((QA / 'manifest.json').read_text(encoding='utf-8'))
    assert sha(REPORT) == baseline['report_sha256'], 'Report changed after extraction'
    current = {p.relative_to(DEST).as_posix(): sha(p) for p in DEST.rglob('*') if p.is_file()}
    assert current == baseline['original_files'], 'Destination changed after backup'
    assert protect_hashes() == baseline['protected'], 'Protected files changed during preparation'
    verification = json.loads((QA / 'verification.json').read_text(encoding='utf-8'))
    assert verification['ok'] and verification['report_sha256'] == baseline['report_sha256']
    for name, expected_hash in manifest['files'].items():
        assert sha(QA / 'candidate' / name) == expected_hash, name
    for name in manifest['files']:
        target = (DEST / name).resolve()
        assert target.is_relative_to(DEST.resolve()) and target.suffix == '.md'
        # This is a bulk mechanical export from the reviewed workbook.
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes((QA / 'candidate' / name).read_bytes())
    assert protect_hashes() == baseline['protected']
    print(json.dumps({'published': len(manifest['files']), 'report_unchanged': True}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['build', 'publish'])
    args = parser.parse_args()
    (build if args.action == 'build' else publish)()
