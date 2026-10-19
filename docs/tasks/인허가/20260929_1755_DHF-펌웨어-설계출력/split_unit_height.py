"""409.50pt를 넘는 유닛 비교 칸만 나누고 원본 패키지를 보존한다.

openpyxl은 읽기/검사 전용이다. Artifact Tool 작성값과 Excel 글 높이를
확인한 뒤 승인된 셀/병합/테두리/행 높이만 원본 패키지에 결합한다.
"""
from __future__ import annotations
import collections, copy, json, re, shutil, sys, zipfile
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
import update_report_results as prev
import simplify_evidence_wording as keep
from merge_unit_conditions import norm

TASK=Path(__file__).resolve().parent
QA=TASK/'qa/unit-height-limit-split'
def read(name): return keep.read(QA/name)

def shifted(row,points):return row+sum(n for at,n in points if at<row)
def shifted_range(ref,points):
    def one(m):return m.group(1)+str(shifted(int(m.group(2)),points))
    return re.sub(r'(\$?[A-Z]+\$?)(\d+)',one,ref)

def stage_rows():
    """도구가 지원하지 않는 일반 행 삽입만 native 패키지에서 보존 처리한다."""
    base=read('baseline.json');measure=read('measurements.json');additions={}
    for m in measure:
        extra=len(m['parts'])-(m['end']-m['start']+1)
        if extra>0:additions[m['key'],m['sheet'],m['end']]=extra
    main={b['key']:b['sheet'] for b in base['books']}
    for (key,sn,at),n in list(additions.items()):
        if sn==main[key]:
            for other,sheet in main.items():additions[other,sheet,at]=max(additions.get((other,sheet,at),0),n)
    serial=[dict(key=k,sheet=sn,after=r,count=n) for (k,sn,r),n in additions.items()]
    for b in base['books']:
        package=prev.parts(b['backup']);wb=prev.parse(package['xl/workbook.xml'])
        w=load_workbook(b['backup'],rich_text=True)
        modified=[]
        for sn in {sn for key,sn,_ in additions if key==b['key']}:
            points=sorted((r,n) for (key,sheet,r),n in additions.items() if key==b['key'] and sheet==sn)
            path=prev.sheet_path(package,sn);root=prev.parse(package[path]);data=root.find('s:sheetData',prev.N)
            original={int(row.get('r')):copy.deepcopy(row) for row in data}
            for row in data:
                old=int(row.get('r'));new=shifted(old,points);row.set('r',str(new))
                for c in row:c.set('r',re.sub(r'\d+$',str(new),c.get('r')))
            for at,n in points:
                for offset in range(1,n+1):
                    r=shifted(at,points)+offset;row=copy.deepcopy(original[at]);row.set('r',str(r))
                    for c in row:
                        c.set('r',re.sub(r'\d+$',str(r),c.get('r')));c.attrib.pop('t',None)
                        for child in list(c):c.remove(child)
                    data.append(row)
            ordered=sorted(list(data),key=lambda row:int(row.get('r')))
            for row in list(data):data.remove(row)
            for row in ordered:data.append(row)
            root.find('s:dimension',prev.N).set('ref',shifted_range(root.find('s:dimension',prev.N).get('ref'),points))
            mc=root.find('s:mergeCells',prev.N)
            existing={str(m) for m in w[sn].merged_cells}
            for m in mc:m.set('ref',shifted_range(m.get('ref'),points))
            # 추가한 조건 행의 짧은 반대쪽과 단독 공통 칸을 병합한다.
            for at,n in points:
                start=shifted(at,points);end=start+n
                for col in range(1,w[sn].max_column+1):
                    if col==11:continue
                    addr=f'{get_column_letter(col)}{at}'
                    if any(addr in m for m in w[sn].merged_cells):continue
                    if not norm(w[sn][addr].value):continue
                    prev.E.SubElement(mc,'{'+prev.NS+'}mergeCell',ref=f'{get_column_letter(col)}{start}:{get_column_letter(col)}{end}')
            mc.set('count',str(len(mc)))
            package[path]=prev.serialize(root);modified.append(sn)
            idx=w.sheetnames.index(sn)
            for name in wb.findall('s:definedNames/s:definedName',prev.N):
                if name.get('localSheetId')==str(idx):name.text=shifted_range(name.text,points)
        package['xl/workbook.xml']=prev.serialize(wb)
        b['authorInput']=str(QA/(b['key']+'-layout.xlsx'))
        with zipfile.ZipFile(b['authorInput'],'w',zipfile.ZIP_DEFLATED) as z:
            for name,value in package.items():z.writestr(name,value)
        # 행 추가로 원래 값/스타일이 바뀌지 않았는지 이동한 모든 셀을 대조한다.
        staged=load_workbook(b['authorInput'],rich_text=True)
        for s in w:
            points=sorted((r,n) for (key,sn,r),n in additions.items() if key==b['key'] and sn==s.title)
            for c in s._cells.values():
                new=staged[s.title].cell(shifted(c.row,points),c.column)
                assert norm(c.value)==norm(new.value),(b['key'],s.title,c.coordinate,'row shift text')
                if c.value is not None:assert copy.copy(c.font)==copy.copy(new.font) and copy.copy(c.fill)==copy.copy(new.fill),(b['key'],s.title,c.coordinate,'row shift font/fill')
        w.close();staged.close()
    prev.save_json(QA/'baseline.json',base);prev.save_json(QA/'row-insertions.json',serial)
    print(json.dumps(dict(insertions=serial,original_values_preserved=True),ensure_ascii=True))

def prepare():
    assert not (QA/'baseline.json').exists(), 'existing baseline must not be replaced'
    QA.mkdir(parents=True,exist_ok=True)
    old=keep.read(TASK/'qa/evidence-positive-wording/baseline.json')
    base=dict(books=[],protected={p:prev.digest(p) for p in old['protected'] if Path(p).is_file()})
    jobs=[]
    for oldb in old['books']:
        b=dict(oldb);b['backup']=str(QA/(b['key']+'-before.xlsx'))
        b['sha256']=prev.digest(b['path']);shutil.copy2(b['path'],b['backup'])
        b['views']=['K191:L192' if b['key']=='unit' else 'K191:K192']
        base['books'].append(b)
        w=load_workbook(b['backup'],rich_text=True)
        targets=[(b['sheet'],[11,12] if b['key']=='unit' else [11])]
        if b['key']=='criterion':targets.append(('기준서_유닛 검증_변경 비교',[11]))
        for sn,cols in targets:
            s=w[sn]
            spans={(m.min_col,m.min_row):m.max_row for m in s.merged_cells if m.min_col==m.max_col}
            for (r,c),cell in sorted(s._cells.items()):
                if c not in cols or r<30 or not norm(cell.value):continue
                if '변경 비교' in sn and keep.current_sign(s,r).strip()!='[+]':continue
                end=spans.get((c,r),r)
                jobs.append(dict(key=b['key'],sheet=sn,cell=cell.coordinate,start=r,end=end,col=c,
                    text=norm(cell.value),oldHeight=sum(s.row_dimensions[i].height or 15 for i in range(r,end+1))))
        w.close()
    prev.save_json(QA/'baseline.json',base);prev.save_json(QA/'height-jobs.json',jobs)
    print(json.dumps(dict(jobs=len(jobs),books=len(base['books'])),ensure_ascii=True))

def divide(parts,count):
    """측정한 안전한 조각을 기존 물리 행 수에 맞춰 더 나눈다."""
    text='\n'.join(p['text'] for p in parts)
    # 순수 빈 줄을 별도 조각으로 만들지 않는다. 줄 수와 줄 바꿈 길이를
    # 함께 사용해 기존 행에 고르게 배치한 뒤 Excel에서 실제 글꼴로 재측정한다.
    cumulative=[];weight=0;start=0
    for line in text.split('\n'):
        units=sum(2 if ord(ch)>127 else 1 for ch in line)
        weight+=max(1,(units+74)//75)
        at=start+len(line)
        if at>0 and at<len(text)-1 and text[at-1]!='\n' and text[at+1]!='\n':cumulative.append((at,weight))
        start=at+1
    out=[];offset=0;previous=0
    for remaining in range(count,1,-1):
        choices=[p for p in cumulative if p[0]>=offset]
        assert len(choices)>=remaining-1,('not enough text boundaries',count)
        choices=choices[:len(choices)-(remaining-2)]
        target=previous+(weight-previous)/remaining
        at,w=min(choices,key=lambda p:abs(p[1]-target))
        out.append(dict(text=text[offset:at],offset=offset,height=400.5))
        offset=at+1;previous=w
    out.append(dict(text=text[offset:],offset=offset,height=400.5))
    return out

def plan():
    base=read('baseline.json');jobs=read('height-jobs.json');measure=read('measurements.json')
    indexed={(m['key'],m['sheet'],m['cell']):m for m in measure}
    result=dict(changes=[],unmerges=[],merges=[],heights=[],wraps=[],borders=[],shades=[],groups=[])
    # 같은 통과 기준은 두 본문에서 같은 칸으로 나눈다. 더 큰 글 높이를 기준으로 한다.
    main={b['key']:b for b in base['books']}
    common={}
    for m in measure:
        if m['sheet']!=main[m['key']]['sheet'] or m['col']!=11:continue
        old=common.get(m['cell'])
        if old is None or len(m['parts'])>len(old['parts']) or (len(m['parts'])==len(old['parts']) and m['wholeHeight']>old['wholeHeight']):common[m['cell']]=m
    additions=read('row-insertions.json')
    books={b['key']:load_workbook(b['authorInput'],rich_text=True) for b in base['books']}
    for j in jobs:
        m=indexed[j['key'],j['sheet'],j['cell']]
        if j['sheet']==main[j['key']]['sheet'] and j['col']==11:m=common[j['cell']]
        if len(m['parts'])<=1:continue
        points=sorted((p['after'],p['count']) for p in additions if p['key']==j['key'] and p['sheet']==j['sheet'])
        start=shifted(j['start'],points);end=shifted(j['end'],points)+sum(n for at,n in points if at==j['end'])
        count=end-start+1
        assert len(m['parts'])<=count,('additional rows needed',j['key'],j['sheet'],j['cell'],count,len(m['parts']))
        chunks=divide(m['parts'],count)
        assert '\n'.join(p['text'] for p in chunks)==j['text'],('reconstruction',j['cell'])
        col=get_column_letter(j['col']);span=f"{col}{start}:{col}{end}"
        s=books[j['key']][j['sheet']]
        if span in {str(m) for m in s.merged_cells}:result['unmerges'].append(dict(key=j['key'],sheet=j['sheet'],range=span))
        sourcecell=f'{col}{start}'
        result['groups'].append(dict(key=j['key'],sheet=j['sheet'],cell=sourcecell,start=start,end=end,col=j['col'],sourceText=j['text'],chunks=chunks))
        s=books[j['key']][j['sheet']]
        for n,p in enumerate(chunks):
            addr=f"{col}{start+n}"
            result['changes'].append(dict(key=j['key'],sheet=j['sheet'],cell=addr,before=norm(s[addr].value),
                value=p['text'],sourceCell=sourcecell,sliceStart=p['offset'],sliceEnd=p['offset']+len(p['text'])))
            result['wraps'].append(dict(key=j['key'],sheet=j['sheet'],range=addr))
            result['borders'].append(dict(key=j['key'],sheet=j['sheet'],range=addr,sourceCell=sourcecell,
                edges={edge:dict(style='thin',color='#'+('000000' if edge in ('left','right') or (edge=='top' and n==0) or (edge=='bottom' and n==count-1) else 'FFFFFF')) for edge in ('top','bottom','left','right')}))
    # 짧은 쪽은 기존 세로 병합을 유지한다. 본문 K/L의 조건 구간은 이미 같다.
    for b in base['books']:
        groups=[g for g in result['groups'] if g['key']==b['key'] and g['sheet']==b['sheet']]
        if groups:
            examples=[]
            for g in groups:
                r=g['start'];end=g['end']
                view=f"K{r}:L{end}" if b['key']=='unit' else f"K{r}:K{end}"
                if view not in examples:examples.append(view)
                if len(examples)>=3:break
            b['views']=examples
    prev.save_json(QA/'baseline.json',base);prev.save_json(QA/'plan.json',result)
    for w in books.values():w.close()
    print(json.dumps(dict(groups=len(result['groups']),cells=len(result['changes']),by_book=dict(collections.Counter(g['key'] for g in result['groups']))),ensure_ascii=True))

def fragment_cell(package,cell,p,style,shared,source):
    text,props,default=keep.native_text_props(package,source,style,shared)
    lo,hi=p['sliceStart'],p['sliceEnd']
    assert text[lo:hi]==p['value'],(p['cell'],'exact source slice')
    for ch in list(cell):cell.remove(ch)
    cell.set('t','inlineStr');ins=prev.E.SubElement(cell,'{'+prev.NS+'}is')
    groups=[]
    for ch,rp in zip(p['value'],props[lo:hi]):
        shape=prev.serialize(rp)
        if groups and groups[-1][2]==shape:groups[-1][0]+=ch
        else:groups.append([ch,rp,shape])
    for text,rp,_ in groups:
        run=prev.E.SubElement(ins,'{'+prev.NS+'}r');run.append(copy.deepcopy(rp))
        t=prev.E.SubElement(run,'{'+prev.NS+'}t');t.set(prev.XMLSPACE,'preserve');t.text=text

def build():
    base=read('baseline.json');pl=read('plan.json')
    for b in base['books']:
        assert prev.digest(b['path'])==b['sha256'],('external edit',b['key'])
        package=prev.parts(b['authorInput']);authored=prev.parts(QA/(b['key']+'-authored.xlsx'))
        strings=[''.join(si.itertext()) for si in prev.parse(authored['xl/sharedStrings.xml'])] if 'xl/sharedStrings.xml' in authored else []
        style=prev.parse(package['xl/styles.xml']);shared=prev.parse(package['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in package else None
        xfs=style.find('s:cellXfs',prev.N);borders=style.find('s:borders',prev.N)
        roots={};sources={};cells={};acells={}
        for p in [p for p in pl['changes'] if p['key']==b['key']]:
            sn=p['sheet']
            if sn not in roots:
                roots[sn]=prev.parse(package[prev.sheet_path(package,sn)])
                sources[sn]={c.get('r'):c for c in prev.parse(package[prev.sheet_path(package,sn)]).findall('.//s:c',prev.N)}
                cells[sn]={c.get('r'):c for c in roots[sn].findall('.//s:c',prev.N)}
                acells[sn]={c.get('r'):c for c in prev.parse(authored[prev.sheet_path(authored,sn)]).findall('.//s:c',prev.N)}
            ac=acells[sn].get(p['cell'])
            value=strings[int(ac.find('s:v',prev.N).text)] if ac is not None and ac.get('t')=='s' else prev.xmltext(authored,ac)
            assert norm(value)==p['value'],(b['key'],sn,p['cell'],'artifact value')
            fragment_cell(package,cells[sn][p['cell']],p,style,shared,sources[sn][p['sourceCell']])
        for m in [m for m in pl['unmerges'] if m['key']==b['key']]:
            mc=roots[m['sheet']].find('s:mergeCells',prev.N)
            matches=[c for c in mc if c.get('ref')==m['range']]
            assert len(matches)==1,('merge not found',m)
            mc.remove(matches[0]);mc.set('count',str(len(mc)))
        stylecache={}
        for p in [p for p in pl['borders'] if p['key']==b['key']]:
            source=sources[p['sheet']][p['sourceCell']];oldid=int(source.get('s','0'))
            signature=(oldid,json.dumps(p['edges'],sort_keys=True))
            if signature not in stylecache:
                xf=copy.deepcopy(xfs[oldid]);border=copy.deepcopy(borders[int(xf.get('borderId','0'))])
                for edge,e in p['edges'].items():
                    node=border.find('s:'+edge,prev.N)
                    if node is None:node=prev.E.SubElement(border,'{'+prev.NS+'}'+edge)
                    node.set('style',e['style'])
                    for child in list(node):node.remove(child)
                    prev.E.SubElement(node,'{'+prev.NS+'}color',rgb='FF'+e['color'].lstrip('#'))
                bid=len(borders);borders.append(border);xf.set('borderId',str(bid));xf.set('applyBorder','1')
                align=xf.find('s:alignment',prev.N)
                if align is None:align=prev.E.SubElement(xf,'{'+prev.NS+'}alignment')
                align.set('wrapText','1');align.set('vertical','top');xf.set('applyAlignment','1')
                stylecache[signature]=len(xfs);xfs.append(xf)
            cells[p['sheet']][p['range']].set('s',str(stylecache[signature]))
        xfs.set('count',str(len(xfs)));borders.set('count',str(len(borders)))
        package['xl/styles.xml']=prev.serialize(style)
        for sn,root in roots.items():
            rows={int(row.get('r')):row for row in root.findall('s:sheetData/s:row',prev.N)}
            for h in [h for h in pl['heights'] if h['key']==b['key'] and h['sheet']==sn]:
                rows[h['row']].set('ht',str(h['height']));rows[h['row']].set('customHeight','1')
        for sn,root in roots.items():package[prev.sheet_path(package,sn)]=prev.serialize(root)
        with zipfile.ZipFile(QA/(b['key']+'-candidate.xlsx'),'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in package.items():z.writestr(name,data)
        print('candidate',b['key'])

def apply_heights():
    pl=read('plan.json');base=read('baseline.json');measures=read('fragment-heights.json')
    result={}
    for m in measures:
        assert m['height']<=409.5,('fragment too high',m)
        key=(m['key'],m['sheet'],m['row'])
        result[key]=max(result.get(key,0),m['height'])
    main={b['key']:b['sheet'] for b in base['books']}
    rows={r for key,sn,r in result if sn==main[key]}
    for r in rows:
        h=max(result.get((key,sn,r),0) for key,sn in main.items())
        for key,sn in main.items():result[key,sn,r]=h
    pl['heights']=[dict(key=key,sheet=sn,row=r,height=h) for (key,sn,r),h in result.items()]
    prev.save_json(QA/'plan.json',pl)
    print('fitted rows',len(pl['heights']))

def prepare_fonts():
    """Excel 측정용 글자별 글꼴을 native 원본에서 가져온다."""
    base=read('baseline.json');pl=read('plan.json')
    for b in base['books']:
        package=prev.parts(b['authorInput']);style=prev.parse(package['xl/styles.xml'])
        shared=prev.parse(package['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in package else None
        maps={};cache={}
        for p in [p for p in pl['changes'] if p['key']==b['key']]:
            sn=p['sheet']
            if sn not in maps:maps[sn]={c.get('r'):c for c in prev.parse(package[prev.sheet_path(package,sn)]).findall('.//s:c',prev.N)}
            key=(sn,p['sourceCell'])
            if key not in cache:cache[key]=keep.native_text_props(package,maps[sn][p['sourceCell']],style,shared)
            _,props,default=cache[key]
            runs=[]
            for at,rp in enumerate(props[p['sliceStart']:p['sliceEnd']]):
                def property(tag):
                    ch=rp.find('s:'+tag,prev.N)
                    return ch if ch is not None else default.find('s:'+tag,prev.N)
                name=property('rFont');sz=property('sz');bold=property('b');italic=property('i')
                font=dict(name=name.get('val') if name is not None else 'Arial',size=float(sz.get('val')) if sz is not None else 11,
                    bold=bold is not None and bold.get('val','1') not in ['0','false'],italic=italic is not None and italic.get('val','1') not in ['0','false'])
                if runs and runs[-1]['font']==font:runs[-1]['length']+=1
                else:runs.append(dict(start=at+1,length=1,font=font))
            assert all(ord(ch)<=65535 for ch in p['value']),('UTF16 offset',p['cell'])
            p['fontRuns']=runs
    prev.save_json(QA/'plan.json',pl)
    print('font runs prepared')

def review_views():
    base=read('baseline.json')
    for b in base['books']:
        b['views']=['K81:L83','K109:L112','K1363:L1364','K1063:L1072'] if b['key']=='unit' else ['K81:K83','K109:K112','K1363:K1364']
    prev.save_json(QA/'baseline.json',base)

def verify():
    base=read('baseline.json');pl=read('plan.json');books={};checks=[]
    for b in base['books']:
        old=load_workbook(b['authorInput'],rich_text=True);new=load_workbook(QA/(b['key']+'-candidate.xlsx'),rich_text=True)
        approved={(p['sheet'],p['cell']):p for p in pl['changes'] if p['key']==b['key']}
        assert old.sheetnames==new.sheetnames
        assert keep.logs(old)==keep.logs(new),('RTT full lines',b['key'])
        for s in old:
            t=new[s.title]
            assert (s.max_row,s.max_column,len(s._images))==(t.max_row,t.max_column,len(t._images))
            unmerged={p['range'] for p in pl['unmerges'] if p['key']==b['key'] and p['sheet']==s.title}
            assert {str(m) for m in t.merged_cells}=={str(m) for m in s.merged_cells}-unmerged
            for pos in set(s._cells)|set(t._cells):
                c=s._cells.get(pos);d=t._cells.get(pos);addr=(c or d).coordinate
                p=approved.get((s.title,addr))
                assert norm(d.value if d else None)==(p['value'] if p else norm(c.value if c else None)),(b['key'],s.title,addr,'text')
        for g in [g for g in pl['groups'] if g['key']==b['key']]:
            s=new[g['sheet']];values=[norm(s.cell(r,g['col']).value) for r in range(g['start'],g['end']+1)]
            assert '\n'.join(values)==g['sourceText'],(g['cell'],'joined values')
            assert values[0].startswith('\n') and values[-1].endswith('\n')
            assert all(not v.startswith('\n') for v in values[1:]),(g['cell'],'internal leading newline')
            assert all(not v.endswith('\n') for v in values[:-1]),(g['cell'],'internal trailing newline')
            for r in range(g['start'],g['end']+1):
                c=s.cell(r,g['col']);assert not any(c.coordinate in m for m in s.merged_cells)
                for edge in ['left','right','top','bottom']:
                    side=getattr(c.border,edge)
                    color='FF000000' if edge in ('left','right') or (edge=='top' and r==g['start']) or (edge=='bottom' and r==g['end']) else 'FFFFFFFF'
                    assert side.style=='thin' and side.color.rgb==color,(c.coordinate,edge)
        op,np=prev.parts(b['authorInput']),prev.parts(QA/(b['key']+'-candidate.xlsx'))
        assert set(op)==set(np)
        allowed={'xl/styles.xml'}|{prev.sheet_path(op,sn) for sn,_ in approved}
        for name in op:
            if name not in allowed:assert op[name]==np[name],(b['key'],name)
        os,ns=prev.parse(op['xl/styles.xml']),prev.parse(np['xl/styles.xml'])
        for tag in ['cellXfs','borders']:
            a=os.find('s:'+tag,prev.N);z=ns.find('s:'+tag,prev.N)
            assert all(keep.xmlshape(x)==keep.xmlshape(y) for x,y in zip(a,z)),(tag,'old styles')
            os.remove(a);ns.remove(z)
        assert keep.xmlshape(os)==keep.xmlshape(ns),('fonts/fills changed',b['key'])
        shared=prev.parse(op['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in op else None
        oldstyle=prev.parse(op['xl/styles.xml']);newstyle=prev.parse(np['xl/styles.xml'])
        for sn in {sn for sn,_ in approved}:
            ox,nx=prev.parse(op[prev.sheet_path(op,sn)]),prev.parse(np[prev.sheet_path(np,sn)])
            oc={c.get('r'):c for c in ox.findall('.//s:c',prev.N)};nc={c.get('r'):c for c in nx.findall('.//s:c',prev.N)}
            font_sources={addr:copy.deepcopy(c) for addr,c in oc.items()}
            for (sheet,addr),p in approved.items():
                if sheet!=sn:continue
                original,fonts,_=keep.native_text_props(op,font_sources[p['sourceCell']],oldstyle,shared)
                actual,nfonts,_=keep.native_text_props(np,nc[addr],newstyle,None)
                assert actual==p['value']
                assert [keep.xmlshape(rp) for rp in nfonts]==[keep.xmlshape(rp) for rp in fonts[p['sliceStart']:p['sliceEnd']]],(sn,addr,'full rich fonts')
                # 내용/스타일 변경 승인 칸을 제거하고 나머지 XML을 대조한다.
                for c in [oc[addr],nc[addr]]:
                    c.attrib.pop('t',None);c.attrib.pop('s',None)
                    for child in list(c):c.remove(child)
            for root in [ox,nx]:
                mc=root.find('s:mergeCells',prev.N)
                for child in list(mc):
                    if child.get('ref') in {m['range'] for m in pl['unmerges'] if m['key']==b['key'] and m['sheet']==sn}:mc.remove(child)
                mc.set('count',str(len(mc)))
                for row in root.findall('s:sheetData/s:row',prev.N):
                    if int(row.get('r')) in {h['row'] for h in pl['heights'] if h['key']==b['key'] and h['sheet']==sn}:
                        row.attrib.pop('ht',None);row.attrib.pop('customHeight',None)
            assert keep.xmlshape(ox)==keep.xmlshape(nx),(b['key'],sn,'unapproved package change')
        s=new[b['sheet']]
        tests=sum(norm(c.value).strip().startswith('SW-FW-UnitTest-') for c in s._cells.values() if c.column==4)
        assert tests==115
        checks.append(dict(key=b['key'],tests=tests,rows=s.max_row,splitBoxes=sum(g['key']==b['key'] for g in pl['groups']),images=sum(len(s._images) for s in new)))
        old.close();books[b['key']]=new
    cs=books['criterion'][base['books'][0]['sheet']];rs=books['unit'][base['books'][1]['sheet']]
    for r in range(30,cs.max_row+1):
        for c in range(1,12):assert norm(cs.cell(r,c).value)==norm(rs.cell(r,c).value),('common content',r,c)
    assert {str(m) for m in cs.merged_cells}=={str(m) for m in rs.merged_cells if m.max_col<=11}
    for p,h in base['protected'].items():assert prev.digest(p)==h,('protected',p)
    prev.save_json(QA/'verification.json',dict(status='PASS',books=checks,exact_text_and_full_fonts_preserved=True,common_columns='A:K identical',black_outer_white_inner_borders=True,protected_files=len(base['protected'])))
    for w in books.values():w.close()
    print(json.dumps(checks,ensure_ascii=True))

def publish():
    assert read('verification.json')['status']=='PASS'
    assert read('native-checks.json')['status']=='PASS'
    base=read('baseline.json');handles=[]
    try:
        for b in base['books']:
            assert prev.digest(b['path'])==b['sha256'],('external edit',b['key'])
            handles.append((b,open(b['path'],'r+b')))
        for b,h in handles:
            data=(QA/(b['key']+'-candidate.xlsx')).read_bytes();h.seek(0);h.write(data);h.truncate();h.flush()
    finally:
        for _,h in handles:h.close()
    for b in base['books']:assert prev.digest(b['path'])==prev.digest(QA/(b['key']+'-candidate.xlsx'))
    prev.save_json(QA/'publication.json',dict(status='PASS',books=[dict(key=b['key'],path=b['path'],sha256=prev.digest(b['path'])) for b in base['books']]))
    print('published two original filenames')

def saved_checks():
    base=read('baseline.json');checked=[]
    for b in base['books']:
        assert prev.digest(b['path'])==prev.digest(QA/(b['key']+'-candidate.xlsx'))
        w=load_workbook(b['path'],rich_text=True);s=w[b['sheet']]
        assert 'QCC' in norm(s['C19'].value)+norm(s['D19'].value)
        assert norm(s['D1706'].value).strip().endswith('BOOT-ID-20')
        assert next(m.max_row for m in s.merged_cells if m.min_col==m.max_col==4 and m.min_row==1706)==1746
        if b['key']=='unit':
            assert 'L1252:L1254' in {str(m) for m in s.merged_cells}
            assert 'L1271:L1272' in {str(m) for m in s.merged_cells}
        checked.append(dict(key=b['key'],rows=s.max_row,qcc='C19:D19 / last BOOT-ID-20 D1706, end1746',sha256=prev.digest(b['path'])))
        w.close()
    for p,h in base['protected'].items():assert prev.digest(p)==h,('protected',p)
    prev.save_json(QA/'saved-checks.json',dict(status='PASS',books=checked,protected_files=len(base['protected'])))
    prev.save_json(QA/'visual-review.json',dict(status='PASS',reviewed=['unit-review-1.png','unit-review-2.png','unit-review-3.png','unit-review-4.png'],
        coverage='Long criterion/short merged result, both long, three inserted rows, ten-row vector. Full text/fonts checked separately.',
        note='Artifact renderer can show some b val=0 rich text as bold. Native full rPr and Excel checks preserve the original fonts. This is not a full printed-page review.'))
    print('saved checks PASS')

if __name__=='__main__':
    globals()[sys.argv[1]]()
