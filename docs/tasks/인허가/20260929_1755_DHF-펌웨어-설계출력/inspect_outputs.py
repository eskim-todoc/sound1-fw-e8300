from pathlib import Path
import json,hashlib,zipfile
from lxml import etree as E
import openpyxl
T=Path(__file__).parent
S=Path(r'G:\공유 드라이브\20_프로젝트별\01_project sound1\04_docs\99_참고용DHF\03_설계 출력\SW\펌웨어')
N={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
inv=[]
for f in sorted(S.iterdir()):
 if f.name.startswith('~') or f.suffix not in ('.docx','.xlsx'):continue
 kind=f.name.split('-')[0];data={'name':f.name,'source':str(f),'sha256':hashlib.sha256(f.read_bytes()).hexdigest()}
 if f.suffix=='.docx':
  with zipfile.ZipFile(f) as z:
   root=E.fromstring(z.read('word/document.xml'))
   data['paragraphs']=[{'index':i,'text':''.join(p.xpath('.//w:t/text()',namespaces=N)),'style':p.xpath('./w:pPr/w:pStyle/@w:val',namespaces=N)} for i,p in enumerate(root.xpath('//w:body//w:p',namespaces=N))]
   data['parts']=[{'name':n,'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()]
  (T/f'{kind}-source.txt').write_text('\n'.join(f"{p['index']} {p['style']} {p['text']}" for p in data['paragraphs']),encoding='utf8')
 else:
  w=openpyxl.load_workbook(f,data_only=True)
  data['sheets']=[{'name':s.title,'rows':s.max_row,'cols':s.max_column,'values':[[c.value for c in row] for row in s]} for s in w]
  (T/f'{kind}-source.txt').write_text('\n'.join(f"{s['name']} row {i+1}: "+' | '.join(f'{openpyxl.utils.get_column_letter(c+1)}={v}' for c,v in enumerate(row) if v is not None) for s in data['sheets'] for i,row in enumerate(s['values'])),encoding='utf8')
 inv.append(data)
(T/'inventory.json').write_text(json.dumps(inv,ensure_ascii=False,indent=2,default=str),encoding='utf8')
print([(x['name'],len(x.get('paragraphs',[]))) for x in inv])
