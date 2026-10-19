"""CM3-027의 현재 RTT 두 조건을 남기고 중복된 과거 출력 두 행을 제거한다.

문구는 공통 Artifact Tool 생성기로 먼저 작성/검사한다. 행 삭제 API가
제공되지 않아 기존 패키지의 행 주소/병합/인쇄 영역만 제한적으로 조정한다.
"""
from pathlib import Path
import collections
import copy
import json
import re
import sys
import zipfile
from lxml import etree as E
from openpyxl import load_workbook
import revise_unit_report_wording as native

QA = Path(__file__).resolve().parent / 'qa/unit-cm3-27-rtt'
SHEET = '보고서_유닛 시험'
NS = native.NS
MAIN = 'xl/worksheets/sheet2.xml'
START, END, COUNT = 1010, 1011, 2


def remap(address):
    def shift(match):
        col, raw = match.groups()
        row = int(raw)
        # 삭제 구간을 가로지르는 병합의 끝을 직전 살아 있는 행으로 줄인다.
        row = row - COUNT if row > END else START - 1 if row >= START else row
        return col + str(row)
    return re.sub(r'(\$?[A-Z]{1,3}\$?)(\d+)', shift, address)


def save(name, obj):
    (QA / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')


def build():
    native.QA = QA
    native.build()
    with zipfile.ZipFile(QA / 'unit-candidate.xlsx') as z:
        package = {n: z.read(n) for n in z.namelist()}
    root = E.fromstring(package[MAIN])
    assert not root.findall('.//s:f', NS), '수식이 있으면 참조 이동을 별도로 검토한다.'
    assert not root.findall('.//s:drawing', NS)
    assert not root.findall('.//s:conditionalFormatting', NS)
    data = root.find('s:sheetData', NS)
    rows = {int(row.get('r')): row for row in data}
    # 공통 열의 병합 하단 테두리를 새 마지막 행으로 옮긴다.
    for col in 'BCDEFGHIJM':
        dest = rows[1009].find(f's:c[@r="{col}1009"]', NS)
        src = rows[1011].find(f's:c[@r="{col}1011"]', NS)
        assert dest is not None and src is not None
        dest.set('s', src.get('s'))
    for row in list(data):
        number = int(row.get('r'))
        if START <= number <= END:
            # 값은 Artifact Tool 단계에서 승인된 네 셀만 이미 비웠다.
            for cell in row:
                assert not ''.join(cell.itertext()).strip(), cell.get('r')
            data.remove(row)
        elif number > END:
            row.set('r', str(number - COUNT))
            for cell in row:
                cell.set('r', remap(cell.get('r')))
    for node in root.iter():
        for attr in ('ref', 'sqref', 'activeCell', 'topLeftCell'):
            if attr in node.attrib:
                node.set(attr, remap(node.get(attr)))
    package[MAIN] = E.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)
    workbook = E.fromstring(package['xl/workbook.xml'])
    for name in workbook.findall('s:definedNames/s:definedName', NS):
        if name.get('name') == '_xlnm.Print_Area' and name.get('localSheetId') == '1':
            name.text = remap(name.text)
    package['xl/workbook.xml'] = E.tostring(workbook, encoding='UTF-8', xml_declaration=True, standalone=True)
    with zipfile.ZipFile(QA / 'unit-candidate.xlsx', 'w', zipfile.ZIP_DEFLATED) as z:
        for path, data in package.items():
            z.writestr(path, data)
    print('현재 RTT 두 조건 유지 / 과거 출력 2행 삭제')


def verify():
    before = load_workbook(QA / 'unit-before.xlsx', rich_text=True)
    after = load_workbook(QA / 'unit-candidate.xlsx', rich_text=True)
    plan = json.loads((QA / 'plan.json').read_text(encoding='utf-8'))
    edits = {p['cell']: p['value'] for p in plan['changes']}
    assert before.sheetnames == after.sheetnames
    cells = 0
    for old in before:
        new = after[old.title]
        main = old.title == SHEET
        assert new.max_row == old.max_row - (COUNT if main else 0)
        assert new.max_column == old.max_column
        assert set(map(str, new.merged_cells.ranges)) == {
            remap(str(m)) if main else str(m) for m in old.merged_cells.ranges}
        assert {k:dict(v) for k,v in old.column_dimensions.items()} == {
            k:dict(v) for k,v in new.column_dimensions.items()}
        for row in old:
            for c in row:
                if main and START <= c.row <= END:
                    continue
                address = remap(c.coordinate) if main else c.coordinate
                nc = new[address]
                expected = edits.get(c.coordinate, native.plain(c.value)) if main else native.plain(c.value)
                assert native.plain(nc.value) == expected, (old.title, c.coordinate, address)
                # 수정 대상 밖의 글자별 강조도 그대로 보존한다.
                if not (main and c.coordinate in edits):
                    assert c.value == nc.value, (old.title, c.coordinate, 'rich text')
                bottom = main and (c.row == 1009 and c.column in list(range(2, 11)) + [13])
                expected_style = old.cell(1011, c.column)._style if bottom else c._style
                assert nc._style == expected_style, (old.title, c.coordinate, 'style')
                cells += 1
        for r, dim in old.row_dimensions.items():
            if main and START <= r <= END:
                continue
            dest = r - COUNT if main and r > END else r
            assert {k:v for k,v in dict(dim).items() if k != 'r'} == {
                k:v for k,v in dict(new.row_dimensions[dest]).items() if k != 'r'}, (old.title, r)
    s = after[SHEET]
    ids = [(str(s.cell(r, 4).value).strip(), str(s.cell(r, 13).value).strip())
           for r in range(1, s.max_row + 1) if 'SW-FW-UnitTest-' in str(s.cell(r, 4).value)]
    assert len(ids) == 115 and len({x[0] for x in ids}) == 115
    assert all(v == '통과' for _, v in ids)
    assert not any(re.search(r'^RTT C\d+', str(c.value).strip()) for row in s for c in row)
    block = '\n'.join(native.plain(s.cell(r, c).value) for r in range(1003, 1010) for c in range(2, 14))
    assert 'legacy' not in block and 'N=2' not in block
    for i, vector in [(1,'40,1,24'),(2,'42,2,12')]:
        for col in ('K','L'):
            value = native.plain(s[f'{col}{1007+i}'].value)
            assert f'조건 [{i}]' in value
            assert f'[UT-CM3-027] C{i} A[0:2]={vector}' in value
            assert f'[UT-CM3-027] C{i} PASS N=3 F=0' in value
    assert 'CM3-ID-28' in native.plain(s['D1010'].value)
    with zipfile.ZipFile(QA / 'unit-before.xlsx') as a, zipfile.ZipFile(QA / 'unit-candidate.xlsx') as b:
        assert a.namelist() == b.namelist()
        changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
        assert set(changed) == {MAIN, 'xl/workbook.xml'}, changed
        aa, bb = E.fromstring(a.read('xl/workbook.xml')), E.fromstring(b.read('xl/workbook.xml'))
        for x in aa.findall('s:definedNames/s:definedName', NS):
            if x.get('name') == '_xlnm.Print_Area' and x.get('localSheetId') == '1':
                x.text = remap(x.text)
        assert E.tostring(aa, method='c14n') == E.tostring(bb, method='c14n')
    result = dict(checkedCells=cells, changedTextCells=6, deletedDuplicateRows=2,
                  rows=s.max_row, mergedRanges=len(s.merged_cells.ranges),
                  unchangedTestIDs=115, unchangedPassVerdicts=115,
                  remainingBareRTTHeadings=0, unrelatedContentAndStylesPreserved=True,
                  changedPackageParts=changed,
                  finalHash=native.digest(QA / 'unit-candidate.xlsx'))
    save('verification.json', result)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    {'build': build, 'verify': verify}[sys.argv[1]]()
