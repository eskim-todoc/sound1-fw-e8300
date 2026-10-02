"""Check the saved design-only revision without changing any workbook or source."""
from pathlib import Path
from collections import Counter
from urllib.parse import unquote
import hashlib,json,re,sys,zipfile
from lxml import etree as E
from docx import Document
from pypdf import PdfReader
from PIL import Image
from design_examples import EXAMPLES
from unit_flowcharts import FLOWS

T=Path(__file__).resolve().parent;ROOT=T.parents[3]
QA=T/'qa/unit-reconciliation';REVIEW=T/'qa/design-readability'
OUT=ROOT/'docs/참고/인허가/사운드1/DHF/03_설계 출력/SW/펌웨어'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
N={'w':W,'m':'http://schemas.openxmlformats.org/officeDocument/2006/math',
   'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
   'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
q=lambda s:'{'+W+'}'+s
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
compact=lambda s:re.sub(r'\s+','',s)
getjson=lambda p:json.loads(p.read_text(encoding='utf8'))
live='--live' in sys.argv
units=getjson(QA/'model.json')
before=Document(REVIEW/'before-accepted.docx')
candidate=Document(QA/'SDDD-candidate.docx')
path=QA/('SDDD-final-accepted.docx' if live else 'SDDD-candidate.docx')
doc=Document(path)
body='\n'.join(p.text for p in doc.paragraphs)
all_text='\n'.join(doc.element.xpath('.//w:t/text()'))
unit_headings=[p.text.split(',')[0] for p in doc.paragraphs if p.text.startswith('SU-ID-')]
assert unit_headings==[u['su'] for u in units] and len(unit_headings)==115
assert len([p for p in doc.paragraphs if p.text.startswith('SI-ID-')])==10
assert not re.search(r'\b(?:SAD|SDDD|SUVC)(?!-FW-)',all_text)
assert 'SW-FW-UnitTest-' not in all_text and '·' not in all_text
assert '소프트웨어 아키텍쳐 기술서에서 정의한 기능 모듈' in body
assert '시험 ID' not in body and '유닛 검증 기준서' not in body
assert sum(p.text.startswith('입출력 예 :') for p in doc.paragraphs)==38
assert sum(p.text.startswith('동작 조건 :') for p in doc.paragraphs)==77
for u in units:
    for key in ['purpose','inp','out']:
        assert compact(u[key]) in compact(body),(u['su'],key)
    for symbol in u['symbols']:
        assert symbol in body,(u['su'],symbol)
for key,value in EXAMPLES.items():assert compact(value) in compact(body),key

# Equations, coefficient tables, figures and the original reading order remain.
def tables(d):return [[[compact(c.text) for c in r.cells] for r in t.rows] for t in d.tables]
assert tables(doc)==tables(before),'coefficient and unit-list tables'
def mathtext(d):return [compact(''.join(n.text or '' for n in x.findall('.//{'+N['m']+'}t'))) for x in d.element.xpath('.//m:oMath')]
assert mathtext(doc)==mathtext(before),'native equation values'
def images(p):
    with zipfile.ZipFile(p) as z:
        root=E.fromstring(z.read('word/document.xml'))
        rel=E.fromstring(z.read('word/_rels/document.xml.rels'))
        targets={x.get('Id'):x.get('Target') for x in rel}
        return [hashlib.sha256(z.read('word/'+targets[x])).hexdigest()
                for x in root.xpath('//a:blip/@r:embed',namespaces=N)]
old_images=images(REVIEW/'before-accepted.docx');new_images=images(path)
assert [v for v in new_images if v in old_images]==old_images,'retained figure bytes/order'
flow_manifest=getjson(REVIEW/'flowcharts.json')
flow_images=[f['sha256'] for f in flow_manifest['flows']]
assert [v for v in new_images if v not in old_images]==flow_images,'flow image bytes/order'
assert len(new_images)==7+len(FLOWS)
captions=[p.text for p in doc.paragraphs if re.match(r'<?그림\s+\d',p.text)]
assert [int(re.search(r'그림\s+(\d+)',s).group(1)) for s in captions]==list(range(1,len(new_images)+1))
for key in FLOWS:
    assert sum('SU-ID-'+key+' ' in s for s in captions)==1,key
for f in flow_manifest['flows']:
    assert f['font_pt']>=10 and f['label_pt']>=9
    assert sha(Path(f['path']))==f['sha256']
for p,h in flow_manifest['source_hashes'].items():assert sha(ROOT/p)==h,p

# A Word keep-with-next flag alone does not prove the rendered caption stayed
# with its diagram. Match the decoded PDF image to its source pixels per page.
pdf=PdfReader(QA/('SDDD-final.pdf' if live else 'SDDD-candidate.pdf'))
flow_pixels={f['key']:hashlib.sha256(Image.open(f['path']).convert('RGB').tobytes()).hexdigest()
             for f in flow_manifest['flows']}
flow_pages={}
for i,page in enumerate(pdf.pages):
    keys=re.findall(r'그림\s+\d+[.]\s+SU-ID-([A-Z0-9]+-\d+)',page.extract_text())
    if not keys:continue
    hashes={hashlib.sha256(im.image.convert('RGB').tobytes()).hexdigest()
            for im in page.images if im.image.width>=1000}
    for key in keys:
        assert key not in flow_pages,('duplicate flow caption',key)
        assert flow_pixels[key] in hashes,('caption separated from diagram',key,i+1)
        flow_pages[key]=i+1
assert set(flow_pages)==set(FLOWS),'missing rendered flowchart'
for prefix in ['목적 :','입력 :','출력 :','처리 :','인터페이스 :','구현 위치 :']:
    assert any(p.text.startswith(prefix) for p in doc.paragraphs),prefix
def headings(d):
    return [p.text for p in d.paragraphs if p.style.name.lower().startswith('heading') or
            p.text=='표준 및 규제']
assert headings(doc)==headings(before)

# Check actual effective heading/body properties after Word normalization.
lefts=[400,650,900,1245,1330,1415]
for i in range(1,7):
    style=next(s for s in doc.styles if s.name.lower()=='heading '+str(i))
    assert style.font.bold and style.font.name=='맑은 고딕',i
    assert style.paragraph_format.left_indent.twips==lefts[i-1],i
assert max(lefts)<1418
def effective(p,name):
    value=getattr(p.paragraph_format,name)
    if value is not None:return value
    return getattr(p.style.paragraph_format,name)
start=False;level=1;heading_count=0;body_count=0
for p in doc.paragraphs:
    if p.style.name.lower()=='heading 1' and p.text=='서론':start=True
    if not start:continue
    m=re.fullmatch(r'heading ([1-6])',p.style.name.lower())
    if m:level=int(m[1]);heading_count+=1;continue
    if not p.text:continue
    if p.text.startswith(('목적 :','입력 :','출력 :','처리 :','동작 조건 :','인터페이스','구현 위치 :','입출력 예 :')):
        assert effective(p,'left_indent').twips==lefts[level-1],p.text[:90]
        assert abs(effective(p,'line_spacing')-1.25)<.01,p.text[:90]
        body_count+=1
        for run in p.runs:
            if run.text:assert (run.font.size or p.style.font.size).pt==11,p.text[:90]

for part in ['header','footer']:
    def headertexts(file):
        with zipfile.ZipFile(file) as z:
            return [''.join(E.fromstring(z.read(n)).xpath('//w:t/text()',namespaces=N))
                    for n in sorted(z.namelist()) if re.fullmatch('word/'+part+r'\d+\.xml',n)]
    assert headertexts(path)==headertexts(REVIEW/'before-accepted.docx'),part

modelcheck=getjson(QA/'model-checks.json')
assert not [p for p,h in modelcheck['source_hashes'].items() if sha(ROOT/p)!=h]
base=getjson(QA/'baseline.json')
for name in ['SDDD-source','SUVC-source','SAD-source','SAD-before']:
    assert sha(Path(base[name]['path']))==base[name]['sha256'],name
unchanged={
 'SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx':'498f7a4614b75652428d5e27b64f259c975f9856a147776842c199401c5e785b',
 'SITC-FW-TD24A 소프트웨어 통합시험 기준서.xlsx':'e39d54d1f89b3307cb64e67f9f190cfd6ca9d65db7fea5bd104b30a5aabc05a7',
 'SSTC-FW-TD24A 소프트웨어 시스템 시험 기준서.xlsx':'99a92516b1d3d48025a9b038e9b777fca2f3194a2fbc43812bd54df988754c85'}
for name,h in unchanged.items():assert sha(OUT/name)==h,name

registry=ROOT/'docs/참고/인허가/사운드1/QCC 자료 추가 위치 관리.md'
md=registry.read_text(encoding='utf8');links=0
assert '8.2 `QCC 내부 설계`' in md and '9 `소프트웨어 상세 설계 검증`' in md
assert 'QCC 내부 설계' in body and 'E8300 연계' in body
for source in [registry,T/'작업.md',OUT/'[정리] Sound1 설계 출력 문서 변경사항.md']:
    for target in re.findall(r'\]\(([^)]+)\)',source.read_text(encoding='utf8')):
        if re.match(r'^(https?://|#)',target):continue
        assert (source.parent/unquote(target.split('#',1)[0].strip('<>'))).exists(),target
        links+=1

result={'scope':'design-document wording and layout; no firmware or executed-test result changes',
        'unit_count':115,'item_count':10,'design_examples':38,'operating_conditions':77,
        'headings':heading_count,'unit_body_paragraphs':body_count,'heading_indents_twips':lefts,
        'body_font_pt':11,'line_multiple':1.25,'original_figure_count':7,'flowchart_count':len(FLOWS),
        'flowchart_caption_same_page':True,'flowchart_pages':flow_pages,
        'equations_and_tables_preserved':True,'firmware_files_unchanged':len(modelcheck['source_hashes']),
        'workbooks_unchanged':True,'local_links':links,'qcc_positions_preserved':True}
if live:
    assert [compact(p.text) for p in candidate.paragraphs if p.text]==[compact(p.text) for p in doc.paragraphs if p.text]
    published=OUT/'SDDD-FW-TD24A 소프트웨어 상세 설계서.docx'
    with zipfile.ZipFile(published) as z:
        root=E.fromstring(z.read('word/document.xml'));settings=E.fromstring(z.read('word/settings.xml'))
        assert settings.find(q('trackRevisions')) is not None
        ins=len(root.findall('.//'+q('ins')));deleted=len(root.findall('.//'+q('del')))
        changes=len(root.findall('.//'+q('rPrChange')))+len(root.findall('.//'+q('pPrChange')))
        assert ins and deleted and changes
        assert z.testzip() is None
    source=Document(QA/'SDDD-source.docx');rejected=Document(QA/'SDDD-rejected-qa.docx')
    assert [compact(p.text) for p in source.paragraphs if p.text]==[compact(p.text) for p in rejected.paragraphs if p.text]
    assert tables(source)==tables(rejected),'reject original tables'
    pages=len(PdfReader(QA/'SDDD-final.pdf').pages)
    assert pages==len(PdfReader(QA/'SDDD-candidate.pdf').pages),'accepted layout page count'
    result.update({'word_pages':pages,'insertions':ins,'deletions':deleted,'format_changes':changes,
                   'reject_restores_original':True,'sha256':{'SDDD':sha(published)}})
else:result['candidate_sha256']=sha(path)
(REVIEW/('final-checks.json' if live else 'candidate-checks.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(result,ensure_ascii=False,indent=2))
