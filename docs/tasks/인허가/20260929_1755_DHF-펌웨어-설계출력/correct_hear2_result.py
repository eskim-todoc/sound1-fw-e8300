"""HEAR-2 결과 요약만 수정하고 원본 패키지/개별 기록을 보존한다.

값과 병합은 기존 Artifact Tool 작성기로 작성한다. 이 도구는 작성 결과를
대조한 뒤 해당 두 셀과 병합 하나만 원래 OOXML 패키지에 반영한다.
새 보드 측정값이나 FFT raw 값은 생성하지 않는다.
"""
from __future__ import annotations

import copy
import json
import shutil
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

import update_report_results as prev
from simplify_evidence_wording import native_text_props, xmlshape

TASK = Path(__file__).resolve().parent
QA = TASK / 'qa/hear2-result-summary'
ROOT = TASK.parents[3]
REPORT = ROOT / 'docs/참고/인허가/사운드1/DHF/04_설계 검증/SW/펌웨어/SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서.xlsx'
CRITERION = ROOT / 'docs/참고/인허가/사운드1/DHF/03_설계 출력/SW/펌웨어/SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx'
SHEET = '보고서_유닛 시험'
TARGETS = ('L668', 'L669')
SUMMARY = (
    '\n검증 결과: 5개 입력 벡터와 수치 비교\n'
    '벡터 1: 영 입력의 FFT 512개와 vMag 256개가 모두 0이며, 최대 bin=0이다.\n'
    '벡터 2: S[n%16] 입력의 최대 bin=32이다.\n'
    '벡터 3: S[(2*n)%16] 입력의 최대 bin=64이다.\n'
    '벡터 4: 32768 고정 입력의 최대 bin=0이다.\n'
    '벡터 5: 32768/-32768 교대 입력의 최대 bin=0이다.\n'
    '벡터 2~5 각각의 블록 지수 비교값=1, FFT 성분 18개의 비교값=모두 1, '
    'vMag 256개의 비교값=모두 1로 확인되었다.\n'
    '벡터 1~5의 완료 상태는 모두 chain0=0, chain1=1이다.\n'
    '입력별 출력과 수치 비교 결과가 통과 기준을 만족하여 통과하였다.\n'
)


def read(name):
    return json.loads((QA / name).read_text(encoding='utf-8'))


def prepare():
    assert not (QA / 'baseline.json').exists(), '기존 기준본을 덮어쓰지 않는다.'
    QA.mkdir(parents=True, exist_ok=True)
    book = load_workbook(REPORT, rich_text=True)
    sheet = book[SHEET]
    assert str(sheet['D658'].value).strip() == 'SW-FW-UnitTest-HEAR-ID-2'
    assert '5개 입력 벡터와 수치 비교식' in str(sheet['K668'].value)
    assert '다음 정의를 적용한다' in str(sheet['L668'].value)
    assert '이하여야 한다' in str(sheet['L669'].value)
    evidence = {cell: str(sheet[cell].value) for cell in
                ('L659', 'L670', 'L671', 'L673', 'L675', 'L677')}
    assert '모두 정확히 0' in evidence['L659']
    assert '모든 i=0..1: 0' in evidence['L670']
    for cell, peak in [('L671', 32), ('L673', 64), ('L675', 0), ('L677', 0)]:
        value = evidence[cell]
        assert f'최대 vMag의 인덱스 = {peak}' in value
        assert '모든 i=0..17: 1' in value and '모든 i=0..255: 1' in value
        assert 'chain0=0, chain1=1' in value
    assert not any(m.min_col <= 12 <= m.max_col and
                   m.min_row <= 669 and m.max_row >= 668 for m in sheet.merged_cells)
    backup = QA / 'unit-before.xlsx'
    shutil.copy2(REPORT, backup)
    base = {'books': [{'key': 'unit', 'path': str(REPORT), 'backup': str(backup),
                      'sheet': SHEET, 'views': ['K668:L669']}],
            'sourceHash': prev.digest(REPORT), 'criterionHash': prev.digest(CRITERION),
            'evidence': evidence}
    prev.save_json(QA / 'baseline.json', base)
    prev.save_json(QA / 'plan.json', {
        'changes': [{'key': 'unit', 'sheet': SHEET, 'cell': 'L668', 'value': SUMMARY},
                    {'key': 'unit', 'sheet': SHEET, 'cell': 'L669', 'value': None}],
        'merges': [{'key': 'unit', 'sheet': SHEET, 'range': 'L668:L669'}],
        'shades': []})
    print('prepare: existing vector records verified; two result cells only')


def build():
    base = read('baseline.json')
    assert prev.digest(REPORT) == base['sourceHash'], '작업 중 원본이 변경되었다.'
    backup = Path(base['books'][0]['backup'])
    authored = load_workbook(QA / 'unit-authored.xlsx', rich_text=True)[SHEET]
    assert str(authored['L668'].value) == SUMMARY and authored['L669'].value is None
    assert 'L668:L669' in {str(m) for m in authored.merged_cells}
    package = prev.parts(backup)
    native_sheet = prev.sheet_path(package, SHEET)
    root = prev.parse(package[native_sheet])
    cells = {c.get('r'): c for c in root.findall('s:sheetData/s:row/s:c', prev.N)}
    style = prev.parse(package['xl/styles.xml'])
    shared = prev.parse(package['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in package else None
    old_text, fonts, default = native_text_props(package, cells['L668'], style, shared)
    assert old_text == str(load_workbook(backup, rich_text=True)[SHEET]['L668'].value)
    assert len({repr(xmlshape(p)) for p in fonts}) == 1, '기존 강조 서식을 먼저 검토한다.'
    font = copy.deepcopy(fonts[0] if fonts else default)
    for address in TARGETS:
        node = cells[address]
        for child in list(node):
            if child.tag in {f'{{{prev.NS}}}v', f'{{{prev.NS}}}is', f'{{{prev.NS}}}f'}:
                node.remove(child)
        if address == 'L668':
            node.set('t', 'inlineStr')
            inline = prev.E.SubElement(node, f'{{{prev.NS}}}is')
            run = prev.E.SubElement(inline, f'{{{prev.NS}}}r')
            run.append(font)
            text = prev.E.SubElement(run, f'{{{prev.NS}}}t')
            text.set(prev.XMLSPACE, 'preserve')
            text.text = SUMMARY
        else:
            node.attrib.pop('t', None)
    merged = root.find('s:mergeCells', prev.N)
    prev.E.SubElement(merged, f'{{{prev.NS}}}mergeCell', ref='L668:L669')
    merged.set('count', str(len(merged)))
    package[native_sheet] = prev.serialize(root)
    with zipfile.ZipFile(QA / 'unit-candidate.xlsx', 'w', zipfile.ZIP_DEFLATED) as archive:
        for path, data in package.items():
            archive.writestr(path, data)
    verify()


def verify():
    base = read('baseline.json')
    backup = Path(base['books'][0]['backup'])
    old_package, new_package = prev.parts(backup), prev.parts(QA / 'unit-candidate.xlsx')
    assert old_package.keys() == new_package.keys()
    native_sheet = prev.sheet_path(old_package, SHEET)
    changed = [p for p in old_package if old_package[p] != new_package[p]]
    assert changed == [native_sheet], changed
    a, b = prev.parse(old_package[native_sheet]), prev.parse(new_package[native_sheet])
    for root in (a, b):
        for node in root.findall('s:sheetData/s:row/s:c', prev.N):
            if node.get('r') in TARGETS:
                for child in list(node):
                    if child.tag in {f'{{{prev.NS}}}v', f'{{{prev.NS}}}is'}:
                        node.remove(child)
                node.attrib.pop('t', None)
        merges = root.find('s:mergeCells', prev.N)
        for item in list(merges):
            if item.get('ref') == 'L668:L669':
                merges.remove(item)
        merges.attrib.pop('count', None)
    assert xmlshape(a) == xmlshape(b), '두 결과 칸 이외의 시트 변경'
    old, new = load_workbook(backup, rich_text=True), load_workbook(QA / 'unit-candidate.xlsx', rich_text=True)
    assert old.sheetnames == new.sheetnames
    s = new[SHEET]
    assert str(s['L668'].value) == SUMMARY and s['L669'].value is None
    assert s.max_row == old[SHEET].max_row
    assert {str(m) for m in s.merged_cells} == {str(m) for m in old[SHEET].merged_cells} | {'L668:L669'}
    for cell, value in base['evidence'].items():
        assert str(s[cell].value) == value
    for forbidden in ('기대 최대', '적용한다', '이하여야', '비교한다', '확정하지', 'OBSERVED'):
        assert forbidden not in SUMMARY
    assert prev.digest(CRITERION) == base['criterionHash']
    prev.save_json(QA / 'verification.json', {
        'status': 'PASS', 'changedResultCells': list(TARGETS),
        'addedResultMerge': 'L668:L669', 'allOtherPackagePartsUnchanged': True,
        'allOtherCellsAndNativeStylesUnchanged': True, 'individualResultsUnchanged': True,
        'criterionUnchanged': True, 'newMeasurementsGenerated': False,
        'candidateHash': prev.digest(QA / 'unit-candidate.xlsx')})
    print('verify: PASS; two result cells and one merge only')


def publish():
    base = read('baseline.json')
    assert read('verification.json')['status'] == 'PASS'
    assert prev.digest(REPORT) == base['sourceHash'], '원본 변경을 덮어쓰지 않는다.'
    assert prev.digest(CRITERION) == base['criterionHash']
    candidate = (QA / 'unit-candidate.xlsx').read_bytes()
    # 열려 있는 파일이면 정상적인 쓰기 실패를 반환하고 검증 후보를 보존한다.
    with REPORT.open('r+b') as destination:
        destination.seek(0)
        destination.write(candidate)
        destination.truncate()
    assert prev.digest(REPORT) == prev.digest(QA / 'unit-candidate.xlsx')
    saved = load_workbook(REPORT, rich_text=True)[SHEET]
    assert str(saved['L668'].value) == SUMMARY
    assert 'L668:L669' in {str(m) for m in saved.merged_cells}
    assert str(saved['D1706'].value).strip() == 'SW-FW-UnitTest-BOOT-ID-20'
    assert 'L1252:L1254' in {str(m) for m in saved.merged_cells}
    assert 'L1271:L1272' in {str(m) for m in saved.merged_cells}
    prev.save_json(QA / 'saved-checks.json', {
        'status': 'PASS', 'savedHash': prev.digest(REPORT),
        'QccLocatorsUnchanged': True, 'sourceEvidence': list(base['evidence'])})
    print('publish: saved same report filename; QCC locators unchanged')


if __name__ == '__main__':
    {'prepare': prepare, 'build': build, 'verify': verify, 'publish': publish}[sys.argv[1]]()
