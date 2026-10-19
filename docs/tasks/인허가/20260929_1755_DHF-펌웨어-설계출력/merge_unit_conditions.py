"""분할한 시험 조건을 병합하고 기대값/결과/서식의 보존을 검사한다.

openpyxl은 읽기와 비교에만 사용한다. Artifact Tool이 작성한 값/병합을
기존 패키지에 반영하여 그림, 스타일, 판정과 작성자 편집을 보존한다.
"""
from __future__ import annotations
import collections
import difflib
import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.utils.cell import range_boundaries
import update_report_results as prev

TASK = Path(__file__).resolve().parent
QA = TASK/'qa/unit-merged-conditions'
PART = re.compile(r'^(.*) \((\d+)/(\d+)\)$')
FILLER = '앞 행의 같은 조건에 대한 결과를 이어 기록한다.'
BLACK = '000000'

def read(p):
    return json.loads(Path(p).read_text(encoding='utf-8-sig'))

def plain(v):
    return '' if v is None else str(v)

def norm(v):
    return plain(v).replace('\r\n', '\n')

def split_groups(s, diff=False):
    groups=[]; active=[]; uid=''; key=None
    for r in range(30, s.max_row+1):
        if s.cell(r,4).value:
            uid=plain(s.cell(r,4).value)
        v=norm(s.cell(r,11).value)
        m=PART.fullmatch(v.split('\n')[0]) if v else None
        current=(uid,m[1],int(m[3])) if m and (not diff or plain(s.cell(r,1).value)=='[+]') else None
        if current is None or current!=key or (active and r!=active[-1]+1):
            if active:
                groups.append(dict(uid=key[0],title=key[1],rows=active))
            active=[]
        if current:
            active.append(r)
        key=current
    if active:
        groups.append(dict(uid=key[0],title=key[1],rows=active))
    for g in groups:
        parts=[PART.fullmatch(norm(s.cell(r,11).value).split('\n')[0]) for r in g['rows']]
        assert [int(m[2]) for m in parts]==list(range(1,len(parts)+1)),g
        assert all(int(m[3])==len(parts) for m in parts),g
        assert len(parts)>1,g
    return groups

def criterion_value(s, g):
    bodies=[]
    for r in g['rows']:
        body=norm(s.cell(r,11).value).partition('\n')[2]
        # 과거 높이 확보용 자리 문구만 제거한다. 수치와 조건은 모두 보존한다.
        body='\n'.join(x for x in body.split('\n') if x!=FILLER)
        if body:
            bodies.append(body)
    return g['title']+'\n'+'\n'.join(bodies)

def clean_output(v):
    value=re.sub(r'출력:\s*출력:', '출력:', norm(v))
    return re.sub(r'^(검증[^\n]*|측정[^\n]*) \(\d+/\d+\)$',r'\1',value,flags=re.M)

def result_value(s, g):
    lines=[]; seen_fragments=set(); titles=set(); repeated_output=0
    for r in g['rows']:
        value=clean_output(s.cell(r,12).value)
        if value in seen_fragments:
            repeated_output+=1
            continue
        seen_fragments.add(value)
        for i,line in enumerate(value.split('\n')):
            if i==0 and (line.startswith('검증 ') or line.startswith('측정 ') or line=='검증 결과'):
                if line in titles:
                    continue
                titles.add(line)
            if line.startswith('입력:'):
                # 분할 뒤 칸에 반복된 짧은 입력 설명은 앞쪽 전체 입력에 이미 포함돼 있다.
                if any(x.startswith(line.rstrip('.')) for x in lines if x.startswith('입력:')):
                    continue
            # 같은 조건의 전체 배열을 각 조각에 반복 기록한 경우 한 번만 남긴다.
            if (line.startswith('출력: A[') or line.startswith('출력 A[') or line=='출력:') and line in lines:
                repeated_output+=1
                continue
            if line==FILLER:
                continue
            lines.append(line)
    return '\n'.join(lines),repeated_output

def character_styles(cells, default=BLACK):
    text=''; props=[]
    for n,c in enumerate(cells):
        if n:
            text+='\n';props.append((default,False))
        if isinstance(c.value,CellRichText):
            for run in c.value:
                if isinstance(run,TextBlock):
                    value=run.text
                    color=run.font.color.rgb[-6:] if run.font.color and run.font.color.type=='rgb' else default
                    bold=bool(run.font.b)
                else:
                    value=str(run);color=default;bold=False
                text+=value;props.extend([(color,bold)]*len(value))
        else:
            value=plain(c.value);text+=value;props.extend([(default,False)]*len(value))
    return text,props

def segments(cells, value, default=BLACK):
    old,styles=character_styles(cells,default)
    out=[(default,False)]*len(value)
    for tag,i,j,k,l in difflib.SequenceMatcher(None,old,value,autojunk=False).get_opcodes():
        if tag=='equal':out[k:l]=styles[i:j]
    runs=[]
    for ch,style in zip(value,out):
        if runs and tuple(runs[-1][1:])==style:runs[-1][0]+=ch
        else:runs.append([ch,*style])
    return runs

def prepare():
    QA.mkdir(parents=True,exist_ok=True)
    assert not (QA/'baseline.json').exists(),'기존 기준을 덮어쓰지 않는다.'
    old=read(TASK/'qa/report-reader-revision/baseline.json')
    books=[]
    for key in ['criterion','unit']:
        b=dict(next(b for b in old['books'] if b['key']==key))
        b['sha256']=prev.digest(b['path']);b['backup']=str(QA/(key+'-before.xlsx'))
        b['views']=['K191:L192','K193:L194'] if key=='unit' else ['K191:K192','K193:K194']
        shutil.copyfile(b['path'],b['backup']);books.append(b)
    targets={b['path'] for b in books}
    protected={p:prev.digest(p) for p in old['protected'] if p not in targets and Path(p).is_file()}
    ws={b['key']:load_workbook(b['backup'],rich_text=True) for b in books}
    cs=ws['criterion'][books[0]['sheet']];rs=ws['unit'][books[1]['sheet']]
    assert cs.max_row==rs.max_row
    for r in range(30,cs.max_row+1):
        for col in range(1,12):
            assert plain(cs.cell(r,col).value)==plain(rs.cell(r,col).value),(r,col,'common before')
    gs=split_groups(cs)
    ds=ws['criterion']['기준서_유닛 검증_변경 비교']
    dg=split_groups(ds,True)
    lookup={(g['uid'],g['title']):g for g in dg}
    assert len(lookup)==len(dg)==len(gs)
    changes={};merges=[];repeats=0;outputs=0
    def put(key,s,col,rows,value,reason):
        addr=f'{col}{rows[0]}'
        changes[(key,s.title,addr)]=dict(key=key,sheet=s.title,cell=addr,value=value,
            reason=reason,segments=segments([s[f'{col}{r}'] for r in rows],value))
        for r in rows[1:]:
            tail=f'{col}{r}'
            changes[(key,s.title,tail)]=dict(key=key,sheet=s.title,cell=tail,value='',segments=[],reason='병합 셀의 아래 값 제거')
        if len(rows)>1:
            rng=f'{col}{rows[0]}:{col}{rows[-1]}'
            bounds=range_boundaries(rng)
            for x in s.merged_cells:
                assert x.max_row<bounds[1] or x.min_row>bounds[3] or x.max_col<bounds[0] or x.min_col>bounds[2],(rng,str(x),'overlap')
            merges.append(dict(key=key,sheet=s.title,range=rng))
    for g in gs:
        val=criterion_value(cs,g)
        put('criterion',cs,'K',g['rows'],val,'한 조건의 입력/기대값 병합')
        put('unit',rs,'K',g['rows'],val,'기준서와 동일한 조건 병합')
        res,n=result_value(rs,g);repeats+=n
        put('unit',rs,'L',g['rows'],res,'같은 조건의 결과 병합/반복 출력 정리')
        d=lookup[(g['uid'],g['title'])]
        assert criterion_value(ds,d)==val,(g['uid'],g['title'],'diff mismatch before')
        put('criterion',ds,'K',d['rows'],val,'현재 [+] 기준의 병합 구조 동기화')
    for s in ws['unit']:
        for c in list(s._cells.values()):
            oldv=norm(c.value)
            n=len(re.findall(r'출력:\s*출력:',oldv))
            if not n:continue
            outputs+=n
            if ('unit',s.title,c.coordinate) not in changes:
                changes[('unit',s.title,c.coordinate)]=dict(key='unit',sheet=s.title,cell=c.coordinate,
                    value=clean_output(oldv),reason='중복 출력 표기 제거',segments=segments([c],clean_output(oldv)))
    plan=dict(changes=list(changes.values()),shades=[],merges=merges,heights=[],groups={'criterion':gs,'unit':gs},
        counts=dict(conditions=len(gs),splitCells=sum(len(g['rows']) for g in gs),
                    repeatedOutputs=repeats,duplicateOutputLabels=outputs),diffGroups=dg)
    prev.save_json(QA/'baseline.json',dict(books=books,protected=protected))
    prev.save_json(QA/'plan.json',plan)
    for w in ws.values():w.close()
    print(json.dumps(dict(plan['counts'],changedCells=len(changes),merges=len(merges)),ensure_ascii=False))

def build():
    prev.QA=QA;prev.build()

def review_views():
    base=read(QA/'baseline.json')
    for b in base['books']:
        col='L' if b['key']=='unit' else 'K'
        b['views']+= [f'K109:{col}110',f'K302:{col}303',f'K668:{col}669',f'K1031:{col}1040']
    prev.save_json(QA/'baseline.json',base)

def clean_result_titles():
    """병합한 결과 안에 남은 분할 제목만 정리한다. 수식의 분수는 유지한다."""
    base=read(QA/'baseline.json');plan=read(QA/'plan.json')
    b=next(x for x in base['books'] if x['key']=='unit')
    w=load_workbook(b['backup'],rich_text=True);s=w[b['sheet']]
    changes={(x['key'],x['sheet'],x['cell']):x for x in plan['changes']}
    audits=[]
    for g in plan['groups']['unit']:
        addr=f"L{g['rows'][0]}";p=changes[('unit',s.title,addr)]
        headings=set();out=[];removed=[]
        for line in p['value'].split('\n'):
            m=PART.fullmatch(line)
            if m and (m[1]==g['title'] or m[1].startswith(('검증 ', '측정 ', '입력/원소별 기준:'))):
                removed.append(line)
                if m[1] in headings:continue
                headings.add(m[1]);line=m[1]
            out.append(line)
        value='\n'.join(out)
        if value!=p['value']:
            p['value']=value
            p['segments']=segments([s[f'L{r}'] for r in g['rows']],value)
            audits.append(dict(uid=g['uid'],cell=addr,removedHeadings=removed))
    plan['changes']=list(changes.values());plan['resultTitleRepairs']=audits
    prev.save_json(QA/'plan.json',plan);w.close()
    print(json.dumps(dict(resultTitleRepairs=audits),ensure_ascii=False))

def complete_fields():
    """잘린 라벨/잘못된 원소 번호를 저장된 전체 RTT 배열에서 복원한다."""
    base=read(QA/'baseline.json');plan=read(QA/'plan.json')
    b=next(x for x in base['books'] if x['key']=='unit')
    w=load_workbook(b['backup'],rich_text=True);s=w[b['sheet']]
    changes={(x['key'],x['sheet'],x['cell']):x for x in plan['changes']}
    value=lambda r,c:changes.get(('unit',s.title,f'{c}{r}'),{}).get('value',norm(s[f'{c}{r}'].value))
    recs=read(prev.UPLAN)['records']
    fields=re.compile(r'payload 인덱스: (\d+)\n(?:출력 )?의미/단위: ([^\n]+)\n사전 기대값: ([^\n]+)')
    audits=[];missing_total=0;corrected_total=0
    for rec in recs:
        uid=rec['id'];arrays={}
        repair_rows=set();wanted_cases=set()
        for item in rec['rows']:
            r=item['row'];defs=fields.findall(value(r,'K'))
            if not defs:continue
            existing=re.findall(r'^출력\[(\d+)\] (.+) = ([^\n]+)',value(r,'L'),re.M)
            if {int(i) for i,m,e in defs}<={int(i) for i,m,v in existing} and len(existing)==len(defs):continue
            repair_rows.add(r)
            wanted_cases.add(int(re.search(r'\bC(\d+)\b',value(r,'K').split('\n')[0])[1]))
        if not repair_rows:continue
        for item in rec['rows']:
            if item['kind']!='rtt':continue
            for line in norm(s[f"L{item['row']}"].value).split('\n'):
                m=re.fullmatch(r'\[UT-'+re.escape(uid)+r'\] C(\d+) A\[(\d+):(\d+)\]=(.+)',line)
                if not m:continue
                case,lo,hi=map(int,m.groups()[:3]);vals=m[4].split(',')
                if case not in wanted_cases:continue
                assert hi-lo+1==len(vals),(uid,line)
                target=arrays.setdefault(case,{})
                for i,v in enumerate(vals,lo):
                    assert i not in target or target[i]==v,(uid,case,i,'conflicting raw record')
                    target[i]=v
        for item in rec['rows']:
            r=item['row'];criterion=value(r,'K');actual=value(r,'L')
            if r not in repair_rows:continue
            defs=fields.findall(criterion)
            if not defs:continue
            cm=re.search(r'\bC(\d+)\b',criterion.split('\n')[0])
            assert cm,(uid,r,'case missing')
            case=int(cm[1])
            existing=re.findall(r'^출력\[(\d+)\] (.+) = ([^\n]+)',actual,re.M)
            expected_indices={int(i) for i,meaning,e in defs}
            existing_indices={int(i) for i,meaning,e in existing}
            missing=sorted(expected_indices-existing_indices)
            source='기존 본문 RTT A 배열'
            if case not in arrays:
                pc_record={}
                for match in re.finditer(r'payload 인덱스: (\d+)\n([\s\S]*?)(?=payload 인덱스:|$)',criterion):
                    pc=re.search(r'기존 PC 관찰 O0/O2: (-?\d+)',match[2])
                    if pc:pc_record[int(match[1])]=pc[1]
                # C91/C92는 합격 기준이 아닌, 이전 실패 경로의 명시된 관찰 기록이다.
                # 사전 기대값 대신 기존 원본에 관찰값으로 기록된 수치만 사용한다.
                observed=re.findall(r'payload 인덱스: (\d+)\n관찰값: (-?\d+)',criterion)
                if expected_indices<=set(pc_record):
                    record=pc_record;source='기존 PC 함수 실행 관찰값 (O0/O2)'
                else:
                    assert all('미정' in e for i,m,e in defs) and observed,(uid,r,case,'existing observation missing')
                    record={int(i):v for i,v in observed};source='기존 관찰 기록 (관찰 및 판정 보류 유지)'
            else:record=arrays[case]
            wrong=[];lines=[]
            for i,meaning,e in defs:
                idx=int(i);v=record[idx]
                # 원래 16진 표시가 있으면 기록된 실제 십진수의 16진수로 함께 표시한다.
                suffix=re.search(r' \(0x[0-9A-Fa-f]+\)',e)
                if suffix:
                    digits=len(suffix[0].split('0x')[1].rstrip(')'))
                    hexval=int(v)&((1<<(4*digits))-1)
                    v+=f' (0x{hexval:0{digits}X})'
                lines.append(f'출력[{idx}] {meaning} = {v}')
                oldvals=[ov.split(' ')[0] for oi,om,ov in existing if int(oi)==idx and om==meaning]
                if oldvals and any(ov!=record[idx] for ov in oldvals):wrong.append(idx)
            if not missing and not wrong and len(existing)==len(defs):continue
            first=next((i for i,line in enumerate(actual.split('\n')) if line.startswith('출력[')),1)
            before=actual.split('\n')[:first]
            after=[line for line in actual.split('\n')[first:] if not line.startswith('출력[')]
            if source.startswith('기존 관찰 기록'):
                after=['기존 실패 경로 관찰값을 정리했다. 합격 판정은 보류한다.']
            new='\n'.join(before+lines+after)
            key=('unit',s.title,f'L{r}')
            group=next((g for g in plan['groups']['unit'] if r==g['rows'][0]),None)
            cells=[s[f'L{x}'] for x in (group['rows'] if group else [r])]
            changes[key]=dict(key='unit',sheet=s.title,cell=f'L{r}',value=new,
                reason='원소별 결과 누락/번호를 기존 RTT 전문과 대조하여 복원',segments=segments(cells,new))
            audits.append(dict(uid=uid,case=case,row=r,fields=len(defs),missingIndices=missing,
                              correctedValueIndices=wrong,source=source,record=record))
            missing_total+=len(missing);corrected_total+=len(wrong)
    for p in changes.values():
        if p['key']=='unit' and p['cell'].startswith('L') and p['sheet']==s.title:
            value=clean_output(p['value'])
            if value!=p['value']:
                r=int(p['cell'][1:]);group=next((g for g in plan['groups']['unit'] if r==g['rows'][0]),None)
                cells=[s[f'L{x}'] for x in (group['rows'] if group else [r])]
                p['value']=value;p['segments']=segments(cells,value)
    plan['changes']=list(changes.values());plan['resultFieldRepairs']=audits
    plan['counts']['missingResultFields']=missing_total;plan['counts']['correctedResultFields']=corrected_total
    prev.save_json(QA/'plan.json',plan)
    print(json.dumps(dict(repairedCases=len(audits),missingFields=missing_total,correctedFields=corrected_total),ensure_ascii=False))

def logs(value):
    return [x for x in norm(value).split('\n') if x.startswith('[UT-')]

def verify():
    base=read(QA/'baseline.json');plan=read(QA/'plan.json');checks=[]
    expected=read(prev.UPLAN)['records']
    originals={};newbooks={}
    for b in base['books']:
        key=b['key'];old=load_workbook(b['backup'],rich_text=True);new=load_workbook(QA/(key+'-candidate.xlsx'),rich_text=True)
        originals[key]=old;newbooks[key]=new
        changes={(p['sheet'],p['cell']):p for p in plan['changes'] if p['key']==key}
        ms=[m for m in plan['merges'] if m['key']==key]
        assert old.sheetnames==new.sheetnames
        tails=set()
        for m in ms:
            c,r,c2,r2=range_boundaries(m['range'])
            tails.update((m['sheet'],f'{chr(64+c)}{x}') for x in range(r+1,r2+1))
        for so,sn in zip(old,new):
            wanted={str(m) for m in so.merged_cells}|{m['range'] for m in ms if m['sheet']==sn.title}
            assert wanted=={str(m) for m in sn.merged_cells},(key,sn.title,'merges')
            assert (so.max_row,so.max_column,len(so._images))==(sn.max_row,sn.max_column,len(sn._images))
            heights={h['row']:h['height'] for h in plan['heights'] if h['key']==key and h['sheet']==sn.title}
            for r,d in so.row_dimensions.items():
                assert sn.row_dimensions[r].height==heights.get(r,d.height),(key,sn.title,r,'height')
            for c in list(so._cells.values()):
                p=changes.get((sn.title,c.coordinate));after=sn[c.coordinate]
                assert plain(after.value)==(p['value'] if p else plain(c.value)),(key,sn.title,c.coordinate,'value')
                if (sn.title,c.coordinate) not in tails:
                    assert list(c._style or [0]*9)==list(after._style or [0]*9),(key,sn.title,c.coordinate,'style')
        before=prev.parts(b['backup']);after=prev.parts(QA/(key+'-candidate.xlsx'))
        allowed={'xl/styles.xml'}|{prev.sheet_path(after,s) for s,a in changes}
        assert before.keys()==after.keys()
        assert all(before[p]==after[p] for p in before if p not in allowed),'unrelated package'
        # 병합 아래 셀의 서식은 openpyxl 정규화 대신 원래 OOXML과 직접 대조한다.
        for sn in {s for s,a in changes}:
            ro=prev.parse(before[prev.sheet_path(before,sn)]);rn=prev.parse(after[prev.sheet_path(after,sn)])
            newcells={c.get('r'):c for c in rn.findall('.//s:c',prev.N)}
            for c in ro.findall('.//s:c',prev.N):
                nc=newcells[c.get('r')]
                assert c.get('s')==nc.get('s'),(key,sn,c.get('r'),'native style')
        checks.append(dict(key=key,changedCells=len(changes),addedMerges=len(ms),sheets=len(new.sheetnames),images=sum(len(s._images) for s in new)))
    cs=newbooks['criterion'][base['books'][0]['sheet']];rs=newbooks['unit'][base['books'][1]['sheet']]
    oldcs=originals['criterion'][base['books'][0]['sheet']];oldrs=originals['unit'][base['books'][1]['sheet']]
    for r in range(30,cs.max_row+1):
        for col in range(1,12):assert plain(cs.cell(r,col).value)==plain(rs.cell(r,col).value),(r,col,'common')
    cm={str(m) for m in cs.merged_cells}
    rm={str(m) for m in rs.merged_cells if m.max_col<=11}
    assert cm==rm,'common merge structure'
    audits=[]
    for g in plan['groups']['unit']:
        r=g['rows'][0];full=norm(cs[f'K{r}'].value)
        assert full==criterion_value(oldcs,g)
        # 원래 벡터의 원소별 기대값/인덱스가 모두 원래 순서로 유지된다.
        tokens=lambda v:re.findall(r'(?:I:|payload 인덱스:|수치 출력 순서:|대체 dB 코드:|사전 기대값:|목표 이득 Q8\.16 기대값:)[^\n]*',v)
        assert tokens(full)==tokens('\n'.join(norm(oldcs[f'K{x}'].value) for x in g['rows']))
        # 원시 로그의 수치/인덱스/통과 결과는 병합 전후 그대로여야 한다.
        beforelogs=[ln for x in g['rows'] for ln in logs(oldrs[f'L{x}'].value)]
        assert beforelogs==logs(rs[f'L{r}'].value),(g,'RTT result lines changed')
        assert [ln for x in g['rows'] for ln in logs(oldcs[f'K{x}'].value)]==logs(full),(g,'RTT expected lines changed')
        for line in norm(rs[f'L{r}'].value).split('\n'):
            m=PART.fullmatch(line)
            assert not (m and (m[1]==g['title'] or m[1].startswith(('검증 ', '측정 ', '입력/원소별 기준:')))),(g,'fragmented result title')
        audits.append(dict(uid=g['uid'],condition=g['title'],rows=g['rows'],expectedFieldLines=len(tokens(full)),rttLines=len(beforelogs)))
    # 이전 입력 벡터의 기대값은 보고서 결과가 아니라 기준 원본을 통해 다시 확인한다.
    source_rows={i['row']:i for rec in expected for i in rec['rows']}
    fields_checked=0
    for g in plan['groups']['unit']:
        source='\n'.join(source_rows[r]['expected'] for r in g['rows'])
        f=re.findall(r'(?:대체 dB 코드:|목표 이득 Q8\.16 기대값:|payload 인덱스:|사전 기대값:)[^\n]*',source)
        # 후속 수치 오차 변경은 이번에 되돌리지 않으므로 값과 인덱스별 고정 벡터만 검사한다.
        final=norm(cs[f"K{g['rows'][0]}"].value)
        fixed=[x for x in f if not ('~' in x or '..' in x)]
        assert all(x in final for x in fixed),(g,'original fixed expected field missing')
        fields_checked+=len(fixed)
    # 병합이 필요 없던 항목도 전체 시험 벡터 원본의 고정 기대값을 검사한다.
    merged_anchor={r:g['rows'][0] for g in plan['groups']['unit'] for r in g['rows']}
    unmerged_checked=0
    for rec in expected:
        for item in rec['rows']:
            r=item['row']
            if r in merged_anchor:continue
            source=norm(item['expected']);final=norm(cs[f'K{r}'].value)
            fixed=re.findall(r'(?:대체 dB 코드:|목표 이득 Q8\.16 기대값:|payload 인덱스:|사전 기대값:) *-?\d+(?:\n|$)',source)
            assert all(x.strip() in final for x in fixed),(rec['id'],r,'unmerged original fixed field missing')
            unmerged_checked+=len(fixed)
    ids=[plain(rs.cell(r,4).value) for r in range(30,rs.max_row+1) if 'SW-FW-UnitTest-' in plain(rs.cell(r,4).value)]
    assert ids==[rec['testId'] for rec in expected] and len(ids)==115
    v=norm(cs['K191'].value)
    assert 'I: 5\n대체 dB 코드: -6209351\n목표 이득 Q8.16 기대값: 51' in v
    for w in newbooks.values():
        for s in w:
            if '변경 비교' in s.title:continue
            for c in s._cells.values():
                assert not re.search(r'출력:\s*출력:',norm(c.value)),(s.title,c.coordinate,'duplicate output')
    for a in plan.get('resultFieldRepairs',[]):
        value=norm(rs[f"L{a['row']}"].value)
        pairs=re.findall(r'^출력\[(\d+)\] (.+) = (-?\d+)',value,re.M)
        assert len(pairs)==a['fields'] and len({int(i) for i,m,v in pairs})==a['fields'],a
        assert all(v==a['record'][i] for i,m,v in pairs),(a['uid'],a['case'],'record values')
    fields_complete=0
    for row in rs:
        defs=re.findall(r'payload 인덱스: (\d+)\n(?:출력 )?의미/단위:',norm(row[10].value))
        if not defs:continue
        indices=re.findall(r'^출력\[(\d+)\]',norm(row[11].value),re.M)
        assert set(defs)<=set(indices),(row[10].row,'missing displayed result index',set(defs)-set(indices))
        assert len(indices)==len(set(indices)),(row[11].row,'duplicate result index')
        fields_complete+=1
    assert not split_groups(cs) and not split_groups(newbooks['criterion']['기준서_유닛 검증_변경 비교'],True)
    for p,sha in base['protected'].items():assert prev.digest(p)==sha,('protected',p)
    prev.save_json(QA/'verification.json',dict(books=checks,conditions=audits,originalExpectedFields=fields_checked,
        unmergedExpectedFields=unmerged_checked,tests=115,
        completePayloadCases=fields_complete,resultFieldRepairs=len(plan.get('resultFieldRepairs',[])),
        protectedUnchanged=len(base['protected']),CFX023_I5='-6209351 -> 51',knownVerdictsUnchanged=True))
    for w in list(originals.values())+list(newbooks.values()):w.close()
    print(json.dumps(dict(books=checks,conditions=len(audits),originalExpectedFields=fields_checked),ensure_ascii=False))

def publish():
    verify();base=read(QA/'baseline.json');native=read(QA/'native-checks.json')
    assert (QA/'visual-review.json').exists(),'visual review required'
    visual=read(QA/'visual-review.json')
    for b in base['books']:
        ck=next(x for x in native if x['key']==b['key'])
        assert not ck['overflow'],(b['key'],ck['overflow'])
        assert ck['candidateSha256']==prev.digest(QA/(b['key']+'-candidate.xlsx'))
        assert visual['candidateSha256'][b['key']]==ck['candidateSha256'],'visual review stale'
        assert prev.digest(b['path'])==b['sha256'],('external edit',b['path'])
        with open(b['path'],'r+b'):pass
    for b in base['books']:
        p=Path(b['path']);stage=p.with_name('.unit-merge-'+p.name)
        assert not stage.exists()
        shutil.copyfile(QA/(b['key']+'-candidate.xlsx'),stage);stage.replace(p)
    prev.save_json(QA/'publication.json',dict(files=[dict(path=b['path'],sha256=prev.digest(b['path'])) for b in base['books']]))
    print('기준서/보고서 기존 이름 저장 완료')

def record_visual_review():
    """실제로 확인한 대표 화면과 별도 Excel 높이 검사의 범위를 기록한다."""
    base=read(QA/'baseline.json');native=read(QA/'native-checks.json')
    assert all(not n['overflow'] for n in native)
    prev.save_json(QA/'visual-review.json',dict(
        candidateSha256={b['key']:prev.digest(QA/(b['key']+'-candidate.xlsx')) for b in base['books']},
        artifactRanges={b['key']:b['views'] for b in base['books']},
        nativeVisuallyReviewed=['unit-final-native-1.png','unit-final-native-5.png'],
        checks=['CFX-023 C1의 I5/-6209351/51이 앞 값과 같은 한 셀에 표시됨',
                'CFX-023 C1의 결과 배열과 출력 표기가 한 번만 나타남',
                'CFX-023 C2, CFX-012, CFX-028, CM3-028의 원소별 값/RTT가 병합 셀 안에 유지됨',
                'HEAR-002의 반복 분할 제목이 사라지고 수식/관찰 내용은 유지됨',
                'Native Excel에서 원래 폰트 크기가 유지되고 전체368개 병합 셀의 높이 부족이 없음'],
        scope='전체 셀 내용/구조 및 병합 높이는 자동 대조했다. 육안 확인은 기록한 대표 화면이며 전체 인쇄 쪽의 육안 전수 검사는 아니다.'))
    print('대표 화면/전체 병합 높이 검사 기록 완료')

def verify_saved():
    base=read(QA/'baseline.json');pub=read(QA/'publication.json');checks=[]
    for b in base['books']:
        sha=prev.digest(b['path'])
        assert sha==prev.digest(QA/(b['key']+'-candidate.xlsx'))
        assert sha==next(p['sha256'] for p in pub['files'] if p['path']==b['path'])
        w=load_workbook(b['path'],rich_text=True);s=w[b['sheet']]
        assert 'I: 5\n대체 dB 코드: -6209351\n목표 이득 Q8.16 기대값: 51' in norm(s['K191'].value)
        if b['key']=='unit':
            for ref in ['L1252:L1254','L1271:L1272']:
                assert ref in {str(m) for m in s.merged_cells},('QCC position',ref)
            assert 'QCC' in norm(s['C19'].value)+norm(s['D19'].value),'QCC equipment/addition guidance'
            assert 'BOOT-ID-20' in norm(s['D1703'].value)
            qcc=(TASK.parents[2]/'참고/인허가/사운드1/QCC 자료 추가 위치 관리.md').read_text(encoding='utf-8')
            assert 'L1252:L1254' in qcc and 'L1271:L1272' in qcc
        checks.append(dict(key=b['key'],sha256=sha,savedMatchesCandidate=True))
        w.close()
    prev.save_json(QA/'saved-checks.json',dict(files=checks,QCCPositionsVerified=True))
    print(json.dumps(dict(files=checks,QCCPositionsVerified=True),ensure_ascii=False))

if __name__=='__main__':
    dict(prepare=prepare,build=build,views=review_views,fields=complete_fields,titles=clean_result_titles,
         verify=verify,visual=record_visual_review,publish=publish,saved=verify_saved)[sys.argv[1]]()
