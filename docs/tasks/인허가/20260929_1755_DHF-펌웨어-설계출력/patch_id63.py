from pathlib import Path
from copy import deepcopy
from difflib import SequenceMatcher
import json,re,zipfile,hashlib
import openpyxl
from lxml import etree as E
from highlight_sstc_diff import P,T,NS,q

qa=T/'qa/sstc-id63'
raw=P.read_bytes(); (qa/'before.xlsx').write_bytes(raw)
values=json.loads((qa/'values.json').read_text(encoding='utf8'))
with zipfile.ZipFile(P) as z: parts={n:z.read(n) for n in z.namelist()}
old=dict(parts)
strings=E.fromstring(parts['xl/sharedStrings.xml'])
def value(c):
    if c.get('t')=='s': return ''.join(strings[int(c.find(q('v')).text)].itertext())
    return ''.join(c.xpath('./s:is//s:t/text()',namespaces=NS)) if c.get('t')=='inlineStr' else c.findtext(q('v')) or ''
for name,rownum in [('xl/worksheets/sheet2.xml',109),('xl/worksheets/sheet3.xml',172)]:
    root=E.fromstring(parts[name]); row=root.find(f"s:sheetData/s:row[@r='{rownum}']",NS)
    assert value(row.find(f"s:c[@r='B{rownum}']",NS))=='SW-FW-SYS-Test-ID-63'
    for col,text in zip('GIJ',values):
        c=row.find(f"s:c[@r='{col}{rownum}']",NS)
        assert c is not None
        for child in list(c):c.remove(child)
        c.set('t','inlineStr'); ins=E.SubElement(c,q('is'))
        if rownum==172:
            r=E.SubElement(ins,q('r'));rp=E.SubElement(r,q('rPr'))
            E.SubElement(rp,q('rFont'),val='맑은 고딕');E.SubElement(rp,q('sz'),val='10')
            E.SubElement(rp,q('b'));E.SubElement(rp,q('color'),rgb='FF006100')
            E.SubElement(r,q('t')).text=text
        else:E.SubElement(ins,q('t')).text=text
    row.set('ht',str(max(float(row.get('ht','0')),100)));row.set('customHeight','1')
    parts[name]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
candidate=qa/'candidate.xlsx'
with zipfile.ZipFile(candidate,'w',zipfile.ZIP_DEFLATED) as z:
    for n,data in parts.items():z.writestr(n,data)
before=openpyxl.load_workbook(qa/'before.xlsx',data_only=False)
after=openpyxl.load_workbook(candidate,data_only=False)
changed=[]
for s in before:
    for row in s:
        for c in row:
            if c.value!=after[s.title][c.coordinate].value:changed.append((s.title,c.coordinate))
assert len(changed)==6,changed
assert all(parts[n]==v for n,v in old.items() if n not in ['xl/worksheets/sheet2.xml','xl/worksheets/sheet3.xml'])
# Check current SRS/SSTC identifiers and merged SYS cells, not comparison/source rows.
inputs=P.parents[3]/'02_설계 입력/SW/펌웨어'
srs=openpyxl.load_workbook(next(inputs.glob('*SRS*.xlsx')),data_only=True)['소프트웨어 입력 요구사항 사양서']
upper={}
for row in srs:
    if re.fullmatch(r'SW-FW-RID-\d+',str(row[6].value)):
        cell=row[7];v=cell.value
        if v is None:
            for m in srs.merged_cells.ranges:
                if cell.coordinate in m:v=srs.cell(m.min_row,m.min_col).value;break
        upper[row[6].value]=set(re.findall(r'SYS-RID-FW-\d+',str(v)))
tests={}; blanks=[]
for row in after['기준서_시스템 시험']:
    if re.fullmatch(r'SW-FW-SYS-Test-ID-\d+',str(row[1].value)):
        ident=row[1].value;sw=row[4].value
        tests.setdefault(ident,{'sw':sw,'sys':set(),'pass':[]})
        tests[ident]['sys'].update(re.findall(r'SYS-RID-FW-\d+',str(row[5].value)))
        if row[9].value:tests[ident]['pass'].append(row[9].value)
sw={v['sw'] for v in tests.values()};sys=set().union(*(v['sys'] for v in tests.values()))
mismatch=[k for k,v in tests.items() if v['sys']!=upper.get(v['sw'])]
report={'tests':len(tests),'sw_missing':sorted(set(upper)-sw),'sw_extra':sorted(sw-set(upper)),
 'sys_missing':sorted({f'SYS-RID-FW-{i}' for i in range(1,51)}-sys),'mapping_mismatch':mismatch,
 'blank_pass_conditions':[k for k,v in tests.items() if not v['pass']], 'changed_cells':changed}
assert len(tests)==81 and len(sw)==81 and not report['sw_missing'] and not report['sys_missing'] and not mismatch
assert P.read_bytes()==raw
P.write_bytes(candidate.read_bytes())
report['sha256']=hashlib.sha256(P.read_bytes()).hexdigest()
(qa/'checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=True))
