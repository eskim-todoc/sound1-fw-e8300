"""시험 문서의 공통 열, 보안 연결, 보호 범위 및 원본 패키지를 검증한다."""
import copy
import hashlib
import json
import re
from pathlib import Path
from openpyxl import load_workbook
import preserve_aligned_report_packages as native

TASK = Path(__file__).resolve().parent
QA = TASK / 'qa/system-cybersecurity-rev0'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def text(value):
    return '' if value is None else str(value)


def verify():
    plan = json.loads((QA / 'author-plan.json').read_text('utf8'))
    baseline = json.loads((QA / 'baseline.json').read_text('utf8'))
    books = [load_workbook(QA / file, rich_text=True) for file in ['system-before.xlsx', 'system-candidate.xlsx', 'criterion-candidate.xlsx']]
    before, report, criterion = books
    old = native.parts(QA / 'system-before.xlsx')
    new = native.parts(QA / 'system-candidate.xlsx')
    cp = native.parts(QA / 'criterion-candidate.xlsx')
    archived = native.parts(QA / 'oldCriterion-before.xlsx')
    changed = {c['cell']: c['after'] for c in plan['changes']}
    a, b, c = before['보고서_시스템 시험'], report['보고서_시스템 시험'], criterion['기준서_시스템 시험']
    explicit = {x.get('r') for x in native.parse(new[native.sheet_path(new,b.title)]).findall('.//s:c',native.N)}
    assert set(map(str, a.merged_cells.ranges)) == set(map(str, b.merged_cells.ranges))
    ids = lambda s: [text(s.cell(r, 2).value).strip() for r in range(1,304) if text(s.cell(r,2).value).strip().startswith('SW-FW-SYS-Test-ID-')]
    assert ids(a) == ids(b) == ids(c) and len(ids(b)) == 81 and len(set(ids(b))) == 81
    for source, prefix, count in [(5,'SW-FW-RID-',81),(6,'SYS-RID-FW-',50)]:
        current=set()
        prior=set()
        for r in range(34,282):
            current.update(int(x) for x in re.findall(re.escape(prefix)+r'(\d+)',text(b.cell(r,source).value)))
            prior.update(int(x) for x in re.findall(re.escape(prefix)+r'(\d+)',text(a.cell(r,source).value)))
        assert current==prior==set(range(1,count+1)), ('requirement coverage',prefix,current)
    for r in range(34,282):
        assert text(b.cell(r,12).value)==text(a.cell(r,12).value), ('historic verdict',r)
    check_count = 0
    for row in a:
        for cell in row:
            out = b[cell.coordinate]
            assert text(out.value) == changed.get(cell.coordinate, text(cell.value)), ('unintended value', cell.coordinate)
            for prop in ['font', 'fill', 'border', 'alignment', 'number_format']:
                assert copy.copy(getattr(cell, prop)) == copy.copy(getattr(out, prop)), ('report style', cell.coordinate, prop)
            check_count += 1
    assert set(map(str, c.merged_cells.ranges)) == {str(m) for m in b.merged_cells.ranges if m.max_col <= 10}
    for row in b.iter_rows(max_col=10):
        for cell in row:
            expected = plan['criterionTitle'] if cell.coordinate == 'B2' else text(cell.value)
            assert text(c[cell.coordinate].value) == expected, ('common value', cell.coordinate)
            for prop in ['font', 'fill', 'border', 'alignment', 'number_format']:
                if cell.coordinate in explicit and cell.__class__.__name__ != 'MergedCell':
                    assert copy.copy(getattr(cell, prop)) == copy.copy(getattr(c[cell.coordinate], prop)), ('common style', cell.coordinate, prop)
    for r in range(1,304):
        assert b.row_dimensions[r].height == c.row_dimensions[r].height, ('row height', r)
    for sheet in before.sheetnames:
        if sheet != a.title:
            part = native.sheet_path(old, sheet)
            assert old[part] == new[part], ('retained sheet XML', sheet)
    for path in old:
        if path.startswith('xl/media/') or path.startswith('xl/drawings/') or path.startswith('xl/theme/'):
            assert old[path] == new[path], ('original media/theme', path)
    assert archived[native.sheet_path(archived,'표지')] == cp[native.sheet_path(cp,'표지')]
    for spec in plan['sheets']:
        rs, cs = report[spec['name']], criterion[spec['name']]
        assert [text(rs.cell(r,2).value) for r in range(5,40)] == [x['original_id'] for x in json.loads((QA/'reference-review.json').read_text('utf8'))['items']]
        for r in range(1,40):
            for col in range(1,11):
                assert text(rs.cell(r,col).value) == text(cs.cell(r,col).value), ('trace common', r,col)
            if r >= 5:
                assert rs.cell(r,10).value and rs.cell(r,11).value and rs.cell(r,12).value, ('blank trace',r)
        assert cs.max_column <= 10
        assert all(rs.row_dimensions[r].height<=409.5 for r in range(1,40))
    for book in [report, criterion]:
        for sheet in book:
            for row in sheet:
                for cell in row:
                    assert cell.data_type != 'f', ('formula', sheet.title,cell.coordinate)
                    assert not re.search(r'#REF!|#DIV/0!|#VALUE!', text(cell.value)), ('formula error',cell.coordinate)
    for entry in baseline['sources'].values():
        assert sha(entry['path']) == entry['sha256'], ('source changed', entry['path'])
    for path, expected in baseline['protected'].items():
        assert sha(path) == expected, ('protected changed',path)
    unit_set = set()
    unit = json.loads((QA/'unit.json').read_text('utf8'))['sheets']['보고서_유닛 시험']['cells']
    for value in unit.values():
        unit_set.update(re.findall(r'SW-FW-UnitTest-(?:CFX|CM3|HEAR|BOOT)-ID-\d+', value))
    for sheet in [report[spec['name']] for spec in plan['sheets']]:
        for row in sheet:
            for cell in row:
                for uid in re.findall(r'SW-FW-UnitTest-(?:CFX|CM3|HEAR|BOOT)-ID-\d+', text(cell.value)):
                    assert uid in unit_set, ('missing unit', uid)
                for sid in re.findall(r'SW-FW-SYS-Test-ID-\d+', text(cell.value)):
                    assert sid in ids(b), ('missing system test', sid)
    result = dict(checkedReportCells=check_count, systemTests=81, cybersecurityRequirements=35,
                  protectedFiles=len(baseline['protected']), commonColumns='A:J', sourceHashesPreserved=True,
                  softwareRequirements=81,systemRequirements=50,historicalVerdictsPreserved=True,
                  originalSheetsAndMediaPreserved=True, formulas=0,
                  newSecurityPassesClaimed=False, reportSha256=sha(QA/'system-candidate.xlsx'),
                  criterionSha256=sha(QA/'criterion-candidate.xlsx'))
    (QA/'verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), 'utf8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    verify()
