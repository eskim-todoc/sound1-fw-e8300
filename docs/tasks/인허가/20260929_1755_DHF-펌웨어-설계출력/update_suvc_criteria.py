"""Build data/verification for the SUVC edit. openpyxl is used read-only.

Workbook writing is performed by Artifact Tool and native Excel; rich character
diffs use OOXML because that feature is not exposed by the authoring surface.
"""
from pathlib import Path
import json, re, hashlib, shutil, zipfile, difflib, sys
from collections import Counter
from lxml import etree as E
import openpyxl

T=Path(__file__).resolve().parent
ROOT=T.parents[3]
F=ROOT/'docs/참고/인허가/사운드1/DHF/03_설계 출력/SW/펌웨어'
QA=T/'qa/suvc-practical-criteria'
TARGET=F/'SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx'
DETAIL=F/'SDDD-FW-TD24A 소프트웨어 상세 설계서.docx'
MAIN='기준서_유닛 검증'; DIFF=MAIN+'_변경 비교'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(name,obj): (QA/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2),encoding='utf-8')

# Explicit replacements are grouped by functional unit, not by incidental source
# file versions. Existing fixed-point vectors are retained unless corrected here.
CASES={
'CFX-001': ('허가=0/1, 오디오/HEAR/PCM 완료 플래그=0/1을 각각 독립 초기 상태로 준비한다.','허가 전에는 정상 처리 호출이 없고, 허가 후 수집된 이벤트만 1회 처리한다. 처리한 플래그는 0, 처리하지 않은 플래그는 유지한다.'),
'CFX-002': ('정지/리셋 상태에서 초기화 순서를 기록한다. CMD 0과 CMD 1을 별도로 요청한다.','정지/리셋과 체인 설정이 실행 허용보다 먼저다. CMD 0 완료는 이벤트 0, CMD 1 완료는 이벤트 1이고 두 이벤트를 혼용하지 않는다.'),
'CFX-003': ('1-DMIC/2-DMIC 설정을 각각 적용하고 입력 라우팅 및 활성 마이크 수를 읽는다.','1-DMIC는 SINGLE, 2-DMIC는 DUAL이다. 각 경로의 입력/핀 설정이 보드 정의와 같고 DUAL에서만 빔포밍을 선택한다.'),
'CFX-004': ('PCM 출력 비활성에서 초기화→DIO 설정→활성화 호출을 기록한다.','PCM용 클럭/프레임/데이터 핀 및 FIFO가 연결되고 활성화는 설정 후다. 마이크/I2S FIFO 주소를 PCM 출력으로 사용하지 않는다.'),
'CFX-005': ('I2S 비활성에서 초기화와 활성화를 수행하고 수신 설정을 읽는다.','I2S 수신 FIFO/인터럽트/버퍼가 연결되고 DMIC 버퍼와 겹치지 않는다. 이 시험은 E8300 수신 경로만 판정한다.'),
'CFX-006': ('PCM 명령 상태에 이전 실행 값을 넣은 상태에서 정상 초기화를 수행한다.','첫 오디오/PCM 처리 전에 명령 사용 상태가 초기값이 되고 포화 산술 설정이 적용된다.'),
'CFX-008': ('I2S 수신 상태를 OFF/ON/OFF로 변경한다. 상태 판정 함수와 조회 함수를 각 단계에서 호출한다.','조회 상태는 입력 OFF/ON/OFF와 같다. 이전 ON이 마지막 OFF에 남지 않는다. 실제 수신 판정 카운터는 시작/중단 경계에서 별도로 확인한다.'),
'CFX-009': ('출력이 활성인 시험 상태에서 저전력 요청=1을 넣고 종료 경로의 호출 순서를 기록한다.','PCM/오디오/I2S/DMIC/FIFO/HEAR 정리가 저전력 완료보다 먼저다. 처리한 저전력 요청은 해제되고 정상 처리 루프가 계속 출력하지 않는다.'),
'CFX-010': ('연결/미연결 상태를 차례로 공유하고 정상 전원 모드 처리 분기를 실행한다.','연결 시 연결 처리, 미연결 시 미연결 처리를 선택한다. 공유 연결 상태와 분기가 일치한다. 단순 플래그 관찰을 RF 연결 성공으로 판정하지 않는다.'),
'CFX-011': ('T=50/C=150, 자극 볼륨=1→4, 오디오 볼륨=1→10, 맵 변경 요청을 각각 별도 적용한다.','자극 볼륨 변경 시 Cnew는 119→150이고 T=50은 유지한다. 필요한 매핑 계수 계산은 Cnew 계산 이후다. 오디오 볼륨과 자극 볼륨의 변경 요청/완료 상태를 혼동하지 않는다.'),
'CFX-012': ('맵의 전극/주파수/자극/날짜 필드에 서로 다른 시험 값을 넣는다. 목적지 날짜는 별도 값으로 고정하고 전체 복사와 날짜 제외 복사를 각각 실행한다.','전체 복사는 입력의 모든 대상 필드와 같고, 날짜 제외 복사는 나머지 대상 필드를 복사하되 목적지 날짜를 그대로 유지한다.'),
'CFX-014': ('시험 FIFO 기록에 [1,2,3,4]를 넣고 단일 FIFO 초기화와 전체 초기화를 각각 실행한다.','지정 FIFO의 리셋/초기화 설정이 적용된다. 단일 초기화는 다른 FIFO를 변경하지 않고 전체 초기화는 설계에서 지정한 모든 FIFO를 초기화한다.'),
'CFX-015': ('I2S 블록[i]=100+i(i=0~15), 복사 허용 인터럽트=1, 수신위치0/복사횟수0/복사완료0에서 복사→완료조회→완료해제→출력버퍼초기화를 실행한다.','수신 위치0의16표본은100~115이고 복사 뒤 입력 위치1/복사횟수1/인터럽트0이다. 상세 설계서의 계약에 따라 복사 완료 조회는1, 해제 후0이어야 한다. clear_data는출력 버퍼16개를0으로 만든다. 완료 플래그를 시험에서 대신1로 넣어 통과시키지 않는다.'),
'CFX-019': ('16표본=65536인 새 블록을 매 단계 준비한다. 완만한 전환 단계 수=4로 고정하여 fade-in/out 첫/마지막 단계와 종료 다음 호출을 시험한다.','출력[15-i]는 해당 Q8.16 계수와 같다. fade-in 첫 단계 [0,1023,2047,...,15359], 마지막 [49151,50175,...,63487,65536]이다. fade-out 첫 단계 [65536,63487,...,50175,49151], 마지막 [15359,14335,...,1023,0]이다. fade-out 종료 다음 호출은 16개 모두 0이다.'),
'CFX-020': ('AGC 출력[i]=i(i=0~15)를 넣어 FIFO 기록을 확인한다. 별도로[-131073,-131072,0,131071,131072,나머지0]을 적용한다. 청취는1000Hz 표본과0표본을 각각OD로 보낸다.','첫 시험의 출력[i]=64*i로[0,64,...,960]이다. 포화 시험 앞5개는[-8388608,-8388608,0,8388544,8388607], 나머지0이다. 원본은불변이다. 순음/무음 청취를 구분하되 전류나 절대 이득을 청취로 판정하지 않는다.'),
'CFX-021': ('I2S/매핑/크래들/활성 마이크/방향 조건을 차례로 고정한다. 각 경로 호출 시 게인과 입력 버퍼 인자를 기록한다.','I2S+매핑은 unity/unity, 일반 I2S는 I2S unity/마이크 Gain A, 크래들은 Gain A/B를 적용한다. 비I2S는 SINGLE 단일 입력 또는 DUAL 방향별 빔포밍이다. 조건별 입력 경로를 중복 실행하지 않는다.'),
'CFX-023': ('볼륨 인덱스0/9, 이전 이득0을 독립 초기화한다. 로그 변환 직후d=[-7700297,-7700296,-7700295,-6209353,-6209352,-6209351](Q8.16)을 각각 고정하고 시간 변화 적용 전 목표g를 기록한다.','입력 순서의 목표g는 볼륨1=[-16,0,0,0,0,51], 볼륨10=[973232,973249,973249,973249,973249,973249]다. 첫 값은잡음, 끝 값은감쇠, 나머지는증폭 영역이다. 목표g와 시간 변화 후 반환 이득을 혼동하지 않는다.'),
'CFX-027': ('인덱스[bin]=bin%32(bin=0~255)로 구성한다. 맵 대역 수와 분석 대역 수가 동일/상이한 조건에서 호출한다.','동일하면 256개 인덱스가 모두 입력과 같다. 상이하면 기존 배열 256개가 유지된다. -1과 0~31 외 인덱스는 정상 입력에 사용하지 않는다.'),
'CFX-031': ('T=50/Cnew=150, xmin^p=65536/xmax^p=131072(Q8.16)로 계수 입력을 고정한다. 별도로 p=1에서 x=1/2에 대한 거듭제곱 경로를 분석한다.','계수 A=25600, B=-12800(Q16.8)이다. xmin/xmax 대입 시 각각 50/150으로 복원된다. p=1의 수학 기대값은 x다. 거듭제곱의 근사 정밀도는 라이브러리 사양으로 검토하고 계수 정수 비교와 구분한다.'),
'CFX-033': ('CIS 전략, 전극 수32, 대역i의 전극=i+1, CIS 대역순서[i]=i+1, 크기100, 묵음꺼짐, 채널당2프레임, 1ms12채널, 시작인덱스0을 고정한다. NofM으로 전환할 때 선택 경로를 기록한다.','전극 맵으로 해석한 CIS 패킷의 전극 인덱스는0~31순이며 각크기는100이다. 전송은0~11,12~23,24~31 다음0~3순으로12개씩 이어간다. CIS/NofM 중 선택한 전략만 실행한다.'),
'CFX-036': ('헤더 필드를 모두 0인 조건과 각 필드를 하나씩 1로 바꾸는 조건을 준비한다.','0 조건은 해당 헤더 필드가 0이다. 단일 필드 변경은 정의된 비트 위치만 변경하며 다른 필드를 덮어쓰지 않는다. 전극/크기 인코딩과 헤더를 혼동하지 않는다.'),
'CFX-038': ('초기 모드와 다음 모드를 PREAMBLE/LIVE/SPECIFIC/NOP/FILL_ZERO 지원 모드로 각각 지정한다.','요청 모드의 함수만 호출한다. 프리앰블은 정의된 프레임/횟수 뒤 다음 모드로 전환하며 데이터 프레임과 섞이지 않는다.'),
'CFX-039': ('16표본 블록, HEAR 완료=0/1, 유효 CIS/NofM 계산 자료를 독립 준비하고 실시간 경로를 실행한다.','완료 전 HEAR 결과를 사용하지 않고 완료 뒤 AGC/FFT/대역값/매핑/선택 전략 순서로 진행한다. 생성한 PCM 버퍼는 해당 전략의 별도 시험 기대값과 같다.'),
'CFX-040': ('명령 프레임 [0x11,0x22,0x33]과 전송 길이=3을 준비한다. 버퍼 사용 중과 완료 뒤 요청 상태를 기록한다.','세 프레임이 입력 순서로 기록된다. 사용 중 새 버퍼를 덮어쓰지 않고 완료 상태 처리 뒤 다음 명령을 허용한다.'),
'CFX-041': ('NOP 대기, 역방향 통신 NOP, FILL_ZERO, PCM 리셋을 각각 실행한다.','대기/역방향 통신은 해당 NOP 패턴을 사용한다. FILL_ZERO는 지정 출력이 0이다. 리셋은 FIFO/모드 초기화 경로를 실행하며 네 동작을 혼동하지 않는다.'),
'CFX-042': ('역방향 통신 요청=1에서 시작/채널 전송/종료 단계와 완료 전후를 기록한다.','설계된 프레임 수와 채널별 전송을 완료한 후 완료 플래그를 설정하고 통신 회로를 정리한다. 완료 플래그는 송신 완료이며 수신 연결 성공과 구분한다.'),
'CM3-001': ('초기 주변장치 상태에서 Initialize와 I2C/SPI/DMA 설정 순서를 분석한다.','보드 정의의 인터페이스/인터럽트/DMA 설정을 적용하고 사용 전에 초기화한다. CFX 실행 허가와 주변장치 준비 순서를 혼동하지 않는다.'),
'CM3-002': ('센서 초기화 시작, 설정 단계, 완료/읽기 실패 조건을 준비하여 단계 상태를 기록한다.','리셋 및 설정 적용 뒤 준비 완료가 된다. 초기화 중 읽기 실패를 정상 터치로 처리하지 않고 중복 초기화를 반복 실행하지 않는다.'),
'CM3-004': ('공유 영역 미준비/준비 완료에서 반복 실행 허가를 각각 금지/허용하고 CFX 측 조회값을 기록한다.','미준비에는 허가=0, 준비 완료 뒤 허가=1이다. 금지 호출 뒤 다시 0이며 주소 설정보다 허가가 앞서지 않는다.'),
'CM3-005': ('tick 1회를 입력하고 각 반복 처리 함수의 호출 수/순서를 기록한다.','해당 tick에서 시스템 제어→맵 반영→내부기 처리→통신 처리가 설계 순서로 호출된다. 반복 금지 조건에는 정상 처리를 시작하지 않는다.'),
'CM3-006': ('공유 맵=1, 요청 맵=1/2를 독립 준비한다. 시스템 상태에 구분 가능한 플래그를 설정한다.','동일 맵은 불필요한 변경 요청을 만들지 않고 변경 맵은 새 번호/요청을 반영한다. 시스템 공유 플래그는 입력값과 같다.'),
'CM3-007': ('정상 사용 조건, 터치/저배터리 별도 종료 조건 없음에서 미연결 카운터 호출 전=179999/180000을 독립 준비한다. 연결=true도 시험한다.','미연결 179999는 호출 후180000이고 3분 종료 요청 없음이다. 미연결180000은180001 검사로 종료 요청 후 카운터0/enable_ISD=false다. 연결=true는 카운터0을 유지한다. 실제 대기 시험은1ms tick 180001회와 같아야 한다.'),
'CM3-009': ('t=0 누름 진입 뒤2399/2400/2401ms를 적용한다. 별도로boot-ignore의연속미누름1/2/3회, 읽기실패1~20/21회, 지속누름30/60/90초를 준비한다. 폴링100ms, 읽기hold2000ms, 부팅해제300ms다.','롱터치는2399ms false/2400ms true/2401ms false로1회다. 해제 후재누름은재판정한다. boot-ignore는미누름3회에서해제, 읽기실패1~20회는HOLD/21회는NOT_TOUCH다. 지속누름30/60/90초는RESEED/RE_ATI/MCLR 요청이다. 가상시간/상태시험이며실제재부팅은차단한다.'),
'CM3-010': ('시험 센서 응답을 ATI 미완료/완료, ULP 해제/설정으로 고정한다. 읽기 실패, 재ATI, reseed 요청도 적용한다.','읽기 결과는 해당 상태 필드로 해석한다. 재보정 요청은 정의된 명령/비트로 전송하고 실패 시 유효한 신규 상태로 표시하지 않는다.'),
'CM3-012': ('PMIC 제어 요청 ON/OFF, 저전력 요청0→1을 독립 실행한다.','해당 공유 요청 필드만 입력대로 바뀌고 다른 플래그는 불변이다. 요청 기록을 CFX 처리 완료와 혼동하지 않는다.'),
'CM3-014': ('전원/오류/내부기 연결/표시 금지/종료 상태를 각각 요청하고 tick을 진행한다.','설계된 상태별 LED 패턴과 종료 시점이 선택된다. 표시 금지의 예외 우선순위는 요청 규칙과 같고 종료/강제 소등 뒤 이전 패턴이 남지 않는다.'),
'CM3-015': ('CFX 정리 완료=0/1을 각각 고정하고 sleep 및 wake 경로의 호출을 기록한다.','완료=0에서는 저클럭 전환/최종 절전 진입을 하지 않는다. 완료=1 뒤 정리/클럭 변경을 수행하고 복귀 시 통신 설정과 정상 루프가 복구된다.'),
'CM3-016': ('시험 오류 비트 1개와 공유 주소 동일/상이, CFX 자료 적재0/1을 각각 적용한다.','오류 조회에 해당 비트가 나타나고 clear 뒤0이다. 주소 불일치와 미적재를 정상 상태로 반환하지 않는다. CM3/CFX 오류 필드는 구분된다.'),
'CM3-020': ('자료 준비=0/1, 새 맵 요청=0/1을 독립 설정하고 계산 자료 위치를 조회한다.','요청/자료 완료/내부기 매핑 연결 필드는 각각 입력대로 바뀌고 계산 자료 포인터는 해당 공유 구조를 가리킨다. 데이터 기록 전에 완료를 알리지 않는다.'),
'CM3-021': ('시험 내부기A/B의 ID/이름/위치/시험 연결 정보에 서로 다른 값을 넣고 선택을 전환한다.','선택한 내부기의 각 필드가 정확히 조회되고 A/B 자료가 섞이지 않는다. 실제 키/개인정보는 기록하지 않는다.'),
'CM3-022': ('A는 맵1/2, B는 맵3의 목록/날짜/시험 표식을 준비하고 선택 맵을 전환한다.','사용 가능 개수/인덱스/날짜는 선택 내부기의 저장 자료와 같다. 선택 맵 조회는 마지막 설정값과 같고 목록의 다른 맵을 덮어쓰지 않는다.'),
'CM3-025': ('가상 I2C 응답 [0x12,0x34], 쓰기 [0x56,0x78], 정상 완료/오류 상태를 각각 준비한다.','전송 진행 상태와 완료/오류를 구분하고 수신 버퍼는[0x12,0x34]다. busy 기간의 새 요청 처리는 정의된 분기와 같고 완료 후 사용 상태를 정리한다.'),
'CM3-026': ('SPI Tx 초기 상태에서 [0x12,0x34]를 기록하고 준비 신호 설정/해제 및 버퍼 초기화를 실행한다.','응답 순서는[0x12,0x34]다. 기록 이후 준비 신호High/empty=false이고 해제 뒤Low/empty=true다. 버퍼 초기화는전체0이다. 버퍼 지우기와 준비 신호 해제는별도 기능이다.'),
'CM3-028': ('DAC 자극 상한 A/B/C/D=400/803/1203/1606uA, 오프셋 상한 A/B=295/592uA, A 환산계수=20871/14435(Q1.15), 펄스폭25, NumCh32, 전 채널T100/C300uA로 고정한다.','Stimulation_DAC_A/Offset_DAC_A, DAC_offset_uA100, 오프셋 레지스터44, 32채널 T디지털0/C디지털127, 반환true다. 최대 지원 범위를 넘는 시험값은 실패로 반환한다. 전극 실제 전류 시험은 아니다.'),
'CM3-029': ('펄스폭25, 전극32, T100/C300uA와 앞 DAC 시험 파라미터를 준비한다. NofM 및 CIS 전략을 각각 적용한다.','NofM32대역은 채널당3프레임/1ms8채널이다. CIS는 펄스폭 계산에 따라2프레임/12채널이다. 공유T/C는32채널 각각0/127이고 계산 반환 오류를 무시하지 않는다.'),
'CM3-030': ('현재 상태A에서 A 유지/B 변경을 각각 요청하고 변경 처리 뒤 플래그를 해제한다.','A 유지와 B 신규 진입은 구분된다. 조회는 마지막 상태이며 새 상태 처리가 끝난 뒤 변경 플래그를 해제한다.'),
'CM3-031': ('시험 전원/FPGA/내부기 응답을 정상/실패로 각각 준비하고 초기화 호출을 기록한다.','전원/리셋 준비 뒤 내부기 정보 읽기와 설정이 이어진다. 미완료 단계에서 설정 완료를 반환하지 않는다. 실제 내부기 연결 성공은 응답 확인과 별도로 판정한다.'),
'CM3-032': ('역방향 통신의 정상 응답과 응답 없음/오류를 독립 준비하고 실시간 연결 확인을 실행한다.','수신 판정에 따라 연결 상태를 갱신하고 완료 플래그를 처리 뒤 정리한다. 송신 완료만으로 정상 수신을 판정하지 않는다.'),
'CM3-033': ('매핑 연결 검사 시작/진행/종료와 정상/오류 응답을 준비한다.','시작/진행에는 진행 상태, 종료 뒤에는 미진행 상태다. 응답 판정이 연결 상태에 반영되며 이전 검사 진행 상태가 남지 않는다.'),
'CM3-034': ('FPGA 제어 레지스터의 다른 비트를 구분 가능한 값으로 고정한다. RF ON/OFF 및 리셋을 각각 요청한다.','전송 주소/데이터가 대상 레지스터 정의와 같다. RF/리셋 비트 외 유지 대상 비트는 불변이다. 가상 쓰기 기록을 실제 RF 출력으로 판정하지 않는다.'),
'CM3-035': ('지원 전력 단계의 최소/최대값을 각각 설정하고 조회/PMIC 쓰기 기록을 수집한다.','조회 단계와 설정 단계가 같고 전압 코드가 정의된 단계 대응표와 같다. T/C 배열은 불변이다. 필요 시 동일 설정의 PMIC 출력 변화를 별도 측정하되 단계 기록과 실제 전압을 구분한다.'),
'CM3-036': ('가상 FPGA 레지스터 응답=0 및 각 오류/empty 비트만1인 값을 각각 반환한다. FIFO 응답=[0x12,0x34]를 준비한다.','0에는 오류 없음이다. 단일 비트에는 대응 상태만 나타나고 버전 필드도 마스크/이동 정의대로 반환한다. FIFO 읽기 길이/데이터는 요청과 같다.'),
'CM3-037': ('펄스폭13/14, 최소 펄스폭 요청, 8bit 역방향 통신 ON/OFF를 각각 적용한다.','기록한 레지스터 주소/데이터가 각 입력과 정의된 인코딩에 맞고 OFF 뒤 해당 활성 비트가0이다. 다른 필드는 유지한다.'),
'CM3-038': ('시험 PMIC 레지스터 주소/쓰기0x12, 읽기 응답0x34와 리셋을 각각 준비한다.','해당 PMIC 주소로 지정 데이터0x12를 기록하고 읽기 반환은0x34다. 리셋 경로/순서는 보드 정의와 같고 다른 소자 주소를 사용하지 않는다.'),
'CM3-039': ('격리된 시험 장치에서 VP 활성/비활성 및 두 리셋 경로를 각각 실행한다. 전원부 안전 상태 확인 뒤 PMIC VP 출력도 측정한다.','명령 기록은 각 요청의 출력 허용/차단 설정과 같다. 활성에서는 설정 출력이 나타나고 비활성/차단 리셋에서는 차단 상태다. 전압 수치는 해당 설정의 소자 규격/회로 허용값으로 별도 기록하며 내부기 자극 전류와 구분한다.'),
'CM3-040': ('보조 PMIC 읽기 응답=0x5A를 고정하고 읽기 및 리셋 경로를 실행한다.','지정 소자 주소/레지스터로 접근하고 반환값은0x5A다. 리셋 요청은 정의된 제어 경로를 사용한다.'),
'CM3-041': ('단극/양극, 새 맵 없음/있음 및 마지막 명령 전/후 상태를 독립 준비한다.','선택한 모드의 설정 명령/데이터를 사용한다. 마지막 설정이 끝나기 전에 완료는false, 끝난 뒤true다. 새 맵 요청의 처리/해제 시점을 구분한다.'),
'CM3-042': ('시험 저장소에서 A/B 각각 다른 맵 번호/볼륨을 가진 자료를 저장하고 구조체를0으로 지운 뒤 다시 읽는다.','읽은 시험 자료의 대상 필드가 저장 전과 같고 A/B 파일 경로를 혼동하지 않는다. 파일 오류/짧은 읽기는 해당 실패 경로로 처리한다.'),
'CM3-043': ('격리된 시험 저장소에 A/B 각각 맵1/2를 준비한다. 맵1삭제, A전체삭제, 전체삭제를 독립 초기 상태에서 수행한다.','맵1삭제는 해당 내부기의 맵1만, A전체는A만, 전체삭제는시험A/B전체가 대상이다. 남겨야 할 파일은 그대로이고 제품 저장소는 시험하지 않는다.'),
'CM3-044': ('알림 채널/크기를 시험값으로 고정하고 트리거1회를 발생시킨다. 반복 tick과 종료 조건을 준비한다.','알림이 정의된 채널/크기/횟수로 생성되고 종료 뒤 알림 활성 상태를 해제한다. 실제 펄스 측정이 아니라 생성 파라미터와 상태를 판정한다.'),
'CM3-045': ('명령0x34/0x35/0x36과 각 유효 필드에 구분 가능한 시험값을 준비하여 E8300 수신 함수를 호출한다.','해당 명령의 상태 필드만 입력값으로 변경하고 정의된 응답 바이트/길이가 준비된다. QCC 내부 무선 구현은 판정하지 않는다.'),
'CM3-046': ('연결 허용 상태=false/true에서 일반 조회/설정 명령을 실행한다. 연결 확인 명령은 별도 실행하고 처리 후 명령 상태를 확인한다.','일반 명령은 허용 상태=true에서 대응 처리를 실행하고 처리 뒤 명령을 정리한다. 연결 확인은 별도 경로다. 비활성인 패스키 값 비교를 실제 인증 시험으로 주장하지 않는다.'),
'CM3-047': ('Gain A/B 각각 유효 조회/설정 인덱스와 범위 밖 인덱스를 요청한다.','유효 인덱스는 해당 A/B 게인 조회/설정과 응답에 반영된다. 범위 밖은 오류 응답이며 정상 저장값을 변경하지 않는다.'),
'CM3-048': ('매핑 연결/해제/종료 요청을 독립 적용하고 요청 저장→처리 완료를 구분한다.','요청 명령이 올바르게 저장되고 실제 상태 갱신 후 조회가 연결/해제와 같다. 종료 요청은 해당 처리에 전달하며 요청만으로 완료로 바꾸지 않는다.'),
'CM3-049': ('지정 자극/실시간 자극/저장/조회 명령을 각각 시험 자료와 함께 수신 함수에 넣는다.','명령은 대응 기능에 연결되고 응답 및 명령 정리는 해당 결과에 맞게 이루어진다. 저장 명령은 시험 저장소만 대상으로 한다.'),
'CM3-050': ('OTA 시작→데이터 수신 상태→종료를 적용하고 상태와 응답을 읽는다.','DFU 연결 조회는 시작 후활성/종료 후비활성이다. 데이터 수신만으로 종료 처리하지 않고 각 응답은 해당 명령에 맞다.'),
'CM3-051': ('각 지원 파일 종류와 지원 슬롯, 지원 범위 밖 선택을 시험 자료로 요청한다. 파일 쓰기 성공/실패 응답도 준비한다.','파일 종류/슬롯에 맞는 시험 경로를 사용하고 범위 밖 선택은 오류 응답이다. 실패를 성공 응답으로 바꾸지 않는다. 슬롯 자료를 운영 이미지에 덮어쓰지 않는다.'),
'BOOT-002': ('요청 지연0/1000us에서 지연 계산과 루프 종료를 분석한다. 선택 실측은 Calibration 시험 핀으로 호출 전/후만 표시한다.','시간 단위가us이고 요청 시간/클럭 조건이 지연 계산에 반영된다. 핀 실측에는 토글 오버헤드를 별도 기록한다. 단순 루프 분석을 정확한 실측 시간으로 쓰지 않는다.'),
'BOOT-003': ('서비스 진입 조건false/true를 각각 고정하여 기동 분기를 분석한다.','false는 정상 부트, true는 service_main만 선택한다. 서비스 경로를 정상 자극 루프로 연결하지 않는다.'),
'BOOT-004': ('기동 표시 초기화→시작→tick 진행→정지를 적용한다. 각 단계의 상태/경과 시간/완료를 기록한다.','경과 시간은 진행한tick과 같고 정지 뒤 표시 갱신이 멈춘다. 응용 LED 요청 상태와 별도로 관리한다.'),
'BOOT-007': ('저장장치 초기화/마운트 성공과 실패를 각각 가상 반환한다.','성공 뒤 이미지 읽기가 진행되고 실패가 정상 적재로 이어지지 않는다.'),
'BOOT-008': ('시험 기동/제조 정보 구조에 필드별 구분 값을 준비하고 읽기 대상/구조 크기를 분석한다.','각 저장 위치/크기는 해당 구조와 같고 필드가 올바른 적재 설정에 전달된다. 실제 키/제조 고유값은 로그에 남기지 않는다.'),
'BOOT-009': ('boot_app_load="test", boot_app_ext="app"로 구성하고 빈 이름/빈 확장자/null 인자 조건을 독립 준비한다.','정상 결과는"test.app", 반환길이는8이다. 빈 이름/빈 확장자/null 인자는-1이다. 이 함수는 코어 번호를 판정하지 않고 전달한 파일 이름/확장자를 조합한다.'),
'BOOT-010': ('시험 파일 열기 실패, 짧은 헤더 읽기, 잘못된 헤더, 정상 헤더를 준비하여 적재 함수를 호출한다.','각 실패는 해당 오류로 반환하고 구획 적재 호출은없다. 정상 헤더만 구획 처리로 전달한다.'),
'BOOT-011': ('지원 구획별 주소/길이/데이터에 구분 값을 넣고 미지원 구획도 준비한다.','지원 구획은 대응 코어/메모리로 전달하고 주소/길이가 입력과 같다. 오류 경로를 정상 적재로 처리하지 않는다.'),
'BOOT-012': ('정상 목록에 시험파일A/B 두 항목을 넣는다. 별도로 목록 읽기 실패/파일 열기 실패를 준비한다.','A/B 각 항목을 목록 순서대로 한 번씩 적재한다. 실패는 오류로 반환하며 정상 완료와 혼동하지 않는다.'),
'BOOT-013': ('헤더 포함 CRC 결과를 첫계산0/1, 복호화후0/1, 복호화후ENCRYPTED 비트0/1로 가정하여 각 분기를 분석한다. null헤더와AES초기화/복호화오류, 데이터길이0/1/3/4/5도 검토한다.','첫CRC0은반환0. 첫CRC1 뒤 복호화후CRC0+암호화비트1은0, CRC1 또는암호화비트0은-1이다. null은-1, AES오류는전달한다. 계산에는4바이트워드 길이/잔여8bit바이트를 반영한다. 헤더를포함한CRC잔여값0으로판정하며 별도헤더CRC값과직접비교하지않는다.'),
'BOOT-015': ('시험 버퍼[0x11,0x22,0x33,0x44]를 준비하고 지원 대상별 메모리 주소/전송 길이를 기록한다.','대상별 전송 값/순서/길이는 입력과 같고 전송 단위 변환이 정확하다. 다른 코어 영역을 대상으로 쓰지 않는다.'),
'BOOT-016': ('지원 PRAM 영역을 적재 전/적재 완료 조건으로 나누어 잠금 호출 위치를 분석한다.','잠금은 필요한 적재 완료 뒤 해당 PRAM에 적용한다. 적재 전에 잠그거나 다른 메모리 영역을 잠그지 않는다.'),
'BOOT-017': ('서비스RTT down채널에"dump\\n", "unknown\\n"을 입력한다. 대상시험파일에[0x12,0x34]를넣고파일열기/읽기실패도독립준비한다.','dump는파일시작/종료표시 사이에"12 34"를출력한다. unknown은"invalid command."를출력한다. 파일실패는오류출력이며일반자극루프는호출하지않는다. 서비스콘솔은RTT이며UART콘솔로가정하지않는다.'),
'BOOT-018': ('서비스 표시tick과 재기동 명령을 각각 준비한다. 재기동 호출은 시험 훅으로 기록하되 실제 자동 재시작은 차단한다.','표시는 서비스 상태에 맞게 갱신되고 재기동 요청은 시스템 재시작 함수를 한 번 호출한다. 응용 LED 요청과 분리한다.'),
'BOOT-019': ('시험 상태 파일 없음/짧은읽기/정상읽기를 준비한다. 정상은 state=BOOT, boot_slot_num=1로 고정한다.','없음/짧은읽기는 기본 초기화/기록 경로다. 정상은 저장 상태를 사용하여 슬롯1을 선택한다. CRC필드 존재만으로 파일CRC 검사를 수행했다고 판정하지 않는다.'),
'BOOT-020': ('boot_slot1, alt_slot2, 시도상한3으로 고정한다. BOOT, ALT_BOOT/BOOT_TRY에서호출전횟수2/3, TRY_DONE의성공/실패, 미지원sub_state를각각준비한다.','BOOT는슬롯1. TRY횟수2는새횟수3/슬롯2, 횟수3은새횟수4/슬롯1/stateBOOT/TRY_DONE/FAIL이다. TRY_DONE성공은슬롯2, 실패/미지원은슬롯1 및UNKNOWN상태다. 대체슬롯 번호 자체의 유효성 검사는별도분석하며없는검사를있다고쓰지않는다.'),
}

# These cases must remain analysis-only: hardware/early boot observation would
# require new tools or destructive executable changes that are outside this edit.
ANALYSIS={
'CFX-001','CFX-002','CFX-003','CFX-004','CFX-005','CFX-006','CFX-009',
'CM3-001','CM3-015','CM3-031',
'BOOT-001','BOOT-002','BOOT-003','BOOT-005','BOOT-006','BOOT-008','BOOT-011','BOOT-013','BOOT-016',
}

def chunks(text,limit):
    text=str(text or '')
    out=[]
    while len(text)>limit:
        cut=text.rfind('\n',0,limit)
        if cut<limit//2: cut=text.rfind('。',0,limit)
        if cut<limit//2: cut=limit
        out.append(text[:cut]);text=text[cut:]
    out.append(text)
    return out
def physicalize(rows):
    result=[];pos=[]
    for row in rows:
        pieces=[chunks(v,[9999,9999,9999,9999,9999,9999,9999,230,190,360,420][c]) for c,v in enumerate(row)]
        count=max(map(len,pieces));start=len(result)
        for k in range(count):
            r=[p[k] if k<len(p) else None for p in pieces]
            if k:
                for c in range(1,8):r[c]=None
                r[0]=row[0];r[4]='계속 '+str(k+1)
            result.append(r)
        pos.append([start,count])
    return result,pos

def prepare():
    QA.mkdir(parents=True,exist_ok=True)
    if (QA/'before.xlsx').exists():
        prior=json.loads((QA/'baseline.json').read_text('utf-8'))
        assert sha(TARGET)==prior['workbook_sha256'],'Source changed after baseline'
    else:
        shutil.copy2(TARGET,QA/'before.xlsx')
        dump('baseline.json',{'workbook_sha256':sha(TARGET),'detail_sha256':sha(DETAIL)})
    model=json.loads((T/'qa/unit-reconciliation/model.json').read_text('utf-8'))
    # Use already rendered native-Word accepted view whose published hash matches
    # current source. Avoid accepting revisions in the source document itself.
    oldqa=T/'qa/design-readability/design-intent-and-placeholders'
    checked=json.loads((oldqa/'checks.json').read_text('utf-8'))
    assert sha(DETAIL)==checked['candidate_sha256'],'Re-read current SDDD before using accepted QA'
    with zipfile.ZipFile(oldqa/'candidate-accepted.docx') as z:x=E.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    headings={}
    for p in x.xpath('./w:body/w:p',namespaces=ns):
        text=''.join(p.xpath('.//w:t/text()',namespaces=ns))
        match=re.fullmatch(r'(SU-ID-(?:CFX|HEAR|CM3|BOOT)-\d{3}),\s*(.+)',text)
        if match:headings[match[1]]=match[2]
    assert len(headings)==115,(len(headings),headings)
    assert set(headings)=={u['su'] for u in model}
    wb=openpyxl.load_workbook(QA/'before.xlsx',data_only=False)
    sheet=wb[MAIN];old=[]
    for row in sheet.iter_rows(min_row=30,values_only=True):
        if row[3] and str(row[3]).startswith('SW-FW-UnitTest-'):old.append(list(row[:11]))
        elif row[4] and str(row[4]).startswith('계속 '):
            for c in range(7,11):
                if row[c]:old[-1][c]=str(old[-1][c] or '')+str(row[c])
    assert len(old)==115
    old_by={r[6]:r for r in old};rows=[];units=[]
    for u in model:
        key=f"{u['core']}-{u['id']:03d}";u=dict(u);u['name']=headings[u['su']]
        case,expected=CASES.get(key,(u['case'],u['expected']))
        case=case.replace('유니티','1.0').replace('·','/');expected=expected.replace('유니티','1.0').replace('·','/')
        method='A' if key in ANALYSIS else 'D'
        if method=='A':
            observe='코드에서 조건별 호출, 인자, 반환값과 상태 갱신을 대조한다. 검토한 분기와 예상 결과를 기록한다.'
            limit='정적 분석 시험이다. 실제 핀 파형/전원/실행 시간/자극 출력의 실측을 대신하지 않는다.'
        elif u['core'] in ('CFX','HEAR'):
            observe='시험 입력을 고정하고 대상 처리를 1회 실행한다. 완료 뒤 입력/출력 배열, 상태와 이벤트를 별도 버퍼에 복사하여 CM3 RTT로 기록한다.'
            limit='시험 입력 주입/스냅샷 경로를 준비한다. 처리 중 버퍼 덮어쓰기를 막고 ISR에서 긴 RTT 로그를 출력하지 않는다.'
        elif u['core']=='BOOT':
            observe='시험용 파일/메모리/호출 기록을 준비하여 대상 처리를 실행한다. 반환값, 데이터와 단계 상태를 보존해 RTT 또는 서비스 UART로 기록한다.'
            limit='기동 초기 RTT 미사용 시 서비스 UART 또는 보존 버퍼로 확인한다. 시험 저장소/메모리만 사용하고 운영 이미지를 변경하지 않는다.'
        else:
            observe='입력과 이전 상태를 독립 초기화하고 대상 함수를 실행한다. 실행 전/후 필드, 반환값, 전송 버퍼와 호출 순서를 CM3 RTT로 기록한다.'
            limit='직접 입력/가상 응답을 사용하는 시험 펌웨어를 준비한다. 실제 연결/전원/전류 판정과 내부 상태 비교를 구분한다.'
        if key=='CFX-015':
            limit='설계/코드 차이: 현재 복사 함수에는 복사 완료 플래그 설정문이 없다. 설계서의 완료=1 계약을 판정하고 불일치를 기록한다. 제품 코드는 이번에 수정하지 않는다.'
        refs=', '.join(u['symbols']) or u['name']+'의 처리 경로'
        proc='- '+observe+'\n- 대상: '+refs+'\n- 조건별로 아래 예상 결과와 비교하고 차이가 있으면 입력/상태/실제값을 기록한다.'
        if key=='CFX-020':proc+='\n- OD 순음/무음 청취는 별도 확인으로 기록한다.'
        if key=='CM3-039':proc+='\n- 안전한 부하/프로브 조건에서 VP 활성/차단 전압을 측정한다.'
        if key=='BOOT-019':proc+='\n- UART 출력이 활성인 시험 구성에서는 TX를 보드 속도로 디코딩하여 기록 바이트와 비교한다. LED와 같은 핀 기능을 동시에 구동하지 않는다.'
        row=[None,u['core'],u['module'],u['test'],u['name'],method,u['su'],'- '+case,limit,proc,'- '+expected]
        assert not any('isd_ver' in str(v) for v in row)
        rows.append(row);units.append({'su':u['su'],'test':u['test'],'name':u['name'],'method':method,'case':case,'expected':expected,'refs':u['refs']})
    diff=[];pairs=[]
    for row in rows:
        prev=old_by[row[6]]
        if prev[1:]==row[1:]:diff.append([None]+row[1:])
        else:
            ni=len(diff);diff.append(['[-]']+prev[1:]);diff.append(['[+]']+row[1:]);pairs.append([ni,ni+1])
    physical,positions=physicalize(rows);diffphysical,diffpositions=physicalize(diff)
    plan={'rows':rows,'physical':physical,'positions':positions,'diff':diff,'diffphysical':diffphysical,'diffpositions':diffpositions,'pairs':pairs,'units':units,'old_last':max(r for r in range(30,sheet.max_row+1) if sheet.cell(r,4).value),'counts':dict(Counter(u['core'] for u in model)),'methods':dict(Counter(u['method'] for u in units)),
      'header':[
       ['장비 및 지그, 소프트웨어 툴','사용목적'],
       ['개발/다운로드 도구','EDK Ezairo 및 시험 펌웨어 다운로드/정적 코드 검토'],
       ['SEGGER RTT','CM3 상태, 전송 버퍼, 실행 전후 값 확인. CFX/HEAR는 스냅샷을 CM3로 전달'],
       ['오실로스코프','UART RX/TX, Calibration 시험 표시. 필요 시 PMIC 출력 전압 측정'],
       ['QCC 펌웨어 자료','여기에 QCC 내부 유닛 시험을 추가한다. 현재 시험은 E8300 연계만 판정'],
       ['시험 입력/응답 준비','배열/공유 메모리 직접 입력, 가상 주변장치 응답/호출 기록. 시험 코드 준비 필요'],
       ['시험 저장소/메모리','저장/삭제/OTA/기동 시험은 운영 파일과 분리. 실제 내부기 자극 차단'],
       ['이어폰/OD DAC','순음/무음과 연속 재생 청취. 수치 이득/전극 전류의 정량 판정은 별도'],
       ['방법 및 판정','D: 실행/관찰, A: 코드/계산 분석. 디지털 값은 정확히 비교하고 로그 누락은 재시험'],
       ['기대값 및 비교 표시','입력별 예상 결과는 통과 기준에 기재. [-] 이전/[+] 수정, 진한 글씨는 실제 차이. 표지 비교 제외'],
      ]}
    assert len(rows)==115 and set(r[6] for r in rows)==set(headings)
    dump('plan.json',plan)
    dump('inspection.json',{'sheets':[(s.title,s.max_row,s.max_column) for s in wb],'counts':plan['counts'],'methods':plan['methods'],'detail_names_changed':[(u['su'],u['name'],headings[u['su']]) for u in model if u['name']!=headings[u['su']]],'detail_sha256':sha(DETAIL),'baseline_sha256':sha(TARGET)})
    print(json.dumps({'units':len(rows),'counts':plan['counts'],'methods':plan['methods'],'physical':len(physical),'diffphysical':len(diffphysical)},ensure_ascii=False))

def highlight():
    plan=json.loads((QA/'plan.json').read_text('utf-8'));path=QA/'candidate.xlsx'
    S='http://schemas.openxmlformats.org/spreadsheetml/2006/main';q=lambda n:'{'+S+'}'+n
    N={'s':S,'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    wb=E.fromstring(parts['xl/workbook.xml']);rels=E.fromstring(parts['xl/_rels/workbook.xml.rels'])
    rid=wb.xpath('//s:sheet[@name="'+DIFF+'"]/@r:id',namespaces=N)[0]
    target=next(x.get('Target') for x in rels if x.get('Id')==rid)
    entry=target.lstrip('/') if target.startswith('/') else 'xl/'+target
    sheet=E.fromstring(parts[entry]);cells={c.get('r'):c for c in sheet.iter(q('c'))};flags={}
    for ni,pi in plan['pairs']:
        for col in range(1,11):
            old=str(plan['diff'][ni][col] or '');new=str(plan['diff'][pi][col] or '')
            neg=[True]*len(old);pos=[True]*len(new)
            for m in difflib.SequenceMatcher(None,old,new,autojunk=False).get_matching_blocks():
                neg[m.a:m.a+m.size]=[False]*m.size;pos[m.b:m.b+m.size]=[False]*m.size
            flags[ni,col]=neg;flags[pi,col]=pos
    stats=Counter()
    for logical,row in enumerate(plan['diff']):
        first,count=plan['diffpositions'][logical]
        for col in range(11):
            whole=str(row[col] or '');mask=flags.get((logical,col),[bool(row[0])]*len(whole));offset=0
            for k in range(count):
                text=str(plan['diffphysical'][first+k][col] or '')
                if not text:continue
                addr=f'{chr(65+col)}{30+first+k}';cell=cells[addr]
                start=whole.find(text,offset)
                cm=mask[start:start+len(text)] if start>=0 else [bool(row[0])]*len(text)
                if start>=0:offset=start+len(text)
                if col==4 and k:cm=[False]*len(text)
                for child in list(cell):cell.remove(child)
                cell.set('t','inlineStr');inline=E.SubElement(cell,q('is'));begin=0
                while begin<len(text):
                    stop=begin+1
                    while stop<len(text) and cm[stop]==cm[begin]:stop+=1
                    run=E.SubElement(inline,q('r'));rp=E.SubElement(run,q('rPr'))
                    E.SubElement(rp,q('rFont'),val='맑은 고딕');E.SubElement(rp,q('sz'),val='11')
                    E.SubElement(rp,q('b'),val='1' if cm[begin] else '0')
                    color=('FF9C2020' if row[0]=='[-]' else 'FF14632A') if cm[begin] else 'FF666666'
                    E.SubElement(rp,q('color'),rgb=color)
                    t=E.SubElement(run,q('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=text[begin:stop]
                    stats['dark_chars' if cm[begin] else 'neutral_chars']+=stop-begin;begin=stop
    parts[entry]=E.tostring(sheet,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,data in parts.items():z.writestr(n,data)
    dump('highlight.json',dict(stats));print(dict(stats))

def verify():
    plan=json.loads((QA/'plan.json').read_text('utf-8'));before=openpyxl.load_workbook(QA/'before.xlsx');after=openpyxl.load_workbook(QA/'candidate.xlsx',rich_text=True)
    assert before.sheetnames==after.sheetnames
    for name in ['표지','원본 대응표','1차본 대응표']:
        a,b=before[name],after[name]
        assert list(a.values)==list(b.values),name+' changed'
    s=after[MAIN]
    assert [str(s.cell(30+i,c+1).value or '') for i,r in enumerate(plan['physical']) for c in range(11)]==[str(v or '') for r in plan['physical'] for v in r]
    ids=[s.cell(r,7).value for r in range(30,s.max_row+1) if s.cell(r,4).value]
    assert ids==[r[6] for r in plan['rows']]
    assert all(r[c] for r in plan['rows'] for c in range(1,11))
    assert len(set(r[3] for r in plan['rows']))==115
    for r in plan['rows']:
        assert r[5] in ['A','D'] and '·' not in ''.join(str(v or '') for v in r)
    assert sha(DETAIL)==json.loads((QA/'baseline.json').read_text('utf-8'))['detail_sha256']
    diff=after[DIFF];rich=0;neutral=0;dark=0
    from openpyxl.cell.rich_text import CellRichText,TextBlock
    for row in diff.iter_rows(min_row=30):
        for c in row:
            if isinstance(c.value,CellRichText):
                rich+=1
                for block in c.value:
                    if isinstance(block,TextBlock):
                        rgb=block.font.color.rgb
                        if rgb=='FF666666':neutral+=len(block.text)
                        elif rgb in ['FF9C2020','FF14632A']:dark+=len(block.text)
    assert rich>0 and neutral>0 and dark>0
    result={'units':115,'counts':plan['counts'],'methods':plan['methods'],'rows':len(plan['physical']),'diff_rows':len(plan['diffphysical']),'rich_cells':rich,'neutral_chars':neutral,'dark_chars':dark,'cover_and_lineages_unchanged':True,'detail_unchanged':True,'main_last_test_row':max(r for r in range(30,s.max_row+1) if s.cell(r,4).value),'candidate_sha256':sha(QA/'candidate.xlsx')}
    dump('checks.json',result);print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':
    {'prepare':prepare,'highlight':highlight,'verify':verify}[sys.argv[1]]()
