"""통합/시스템 기준과 과거 보고서를 읽기 전용으로 추출한다."""
import hashlib
import json
from pathlib import Path

import openpyxl

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
QA = TASK / 'qa/validation-reports'
DHF = ROOT / 'docs/참고/인허가/사운드1/DHF'
REF = Path(r'G:\공유 드라이브\20_프로젝트별\01_project sound1\04_docs\99_참고용DHF\04_설계 검증\SW\펌웨어')


def extract():
    QA.mkdir(parents=True, exist_ok=True)
    paths = {
        'integration-reference': REF / 'SITR-FW-TD20B 소프트웨어 통합시험 보고서.xlsx',
        'system-reference': REF / 'SSTR-FW-TD20B 소프트웨어 시스템 시험 보고서.xlsx',
        'integration-criteria': DHF / '03_설계 출력/SW/펌웨어/SITC-FW-TD24A 소프트웨어 통합시험 기준서.xlsx',
        'system-criteria': DHF / '03_설계 출력/SW/펌웨어/[작성자 확인] SSTC-FW-TD24A 소프트웨어 시스템 시험 기준서.xlsx',
    }
    sources = {}
    for label, path in paths.items():
        book = openpyxl.load_workbook(path, data_only=False)
        sheets = {}
        for sheet in book:
            rows = [[c.value for c in row] for row in sheet]
            sheets[sheet.title] = {
                'rows': rows,
                'merges': [str(m) for m in sheet.merged_cells.ranges],
                'widths': {k: v.width for k, v in sheet.column_dimensions.items()},
                'heights': {k: v.height for k, v in sheet.row_dimensions.items() if v.height},
                'print_area': str(sheet.print_area),
                'images': len(sheet._images),
            }
            print(label, sheet.title, sheet.max_row, sheet.max_column, 'images', len(sheet._images))
        (QA / f'{label}.json').write_text(json.dumps(sheets, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
        sources[label] = {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    (QA / 'sources.json').write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding='utf-8')
    protected = {}
    for base in [ROOT / 'src', DHF]:
        for p in base.rglob('*'):
            if not p.is_file() or '~$' in p.name:
                continue
            if base.name == 'src' and p.suffix.lower() not in {'.c','.h','.s','.asm','.ld','.hct','.cproject','.project'}:
                continue
            if base.name == 'DHF' and p.suffix.lower() not in {'.xlsx','.docx'}:
                continue
            protected[str(p.relative_to(ROOT))] = hashlib.sha256(p.read_bytes()).hexdigest()
    (QA / 'protected-before.json').write_text(json.dumps(protected, ensure_ascii=False, indent=2), encoding='utf-8')
    print('protected', len(protected))


if __name__ == '__main__':
    extract()
