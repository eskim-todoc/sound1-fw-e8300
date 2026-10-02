from pathlib import Path
import json, sys, zipfile
import openpyxl
from lxml import etree as E
T=Path(__file__).resolve().parent
ROOT=T.parents[3]
D=ROOT/'docs/참고/인허가/사운드1/DHF'
paths={'input': next((D/'02_설계 입력/SW/펌웨어').glob('*SSTC*.xlsx')),
       'system': next((D/'03_설계 출력/SW/펌웨어').glob('*SSTC*.xlsx')),
       'unit': D/'03_설계 출력/SW/펌웨어/SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx',
       'integration':D/'03_설계 출력/SW/펌웨어/SITC-FW-TD24A 소프트웨어 통합시험 기준서.xlsx'}
if __name__!='__main__':
 key=None
else:key=sys.argv[1]
if key=='word':
 p=D/'03_설계 출력/10_성능검사기준서/EPTC-TD24A 필수성능시험 기준서.docx'
 with zipfile.ZipFile(p) as z:r=E.fromstring(z.read('word/document.xml'))
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 for i,p in enumerate(r.xpath('//w:p',namespaces=ns)):
  s=''.join(p.xpath('.//w:t[not(ancestor::w:del)]/text()',namespaces=ns))
  if int(sys.argv[2])<=i<int(sys.argv[3]): print(i,s)
elif key:
 w=openpyxl.load_workbook(paths[key],rich_text=True)
 s=w.worksheets[1]
 start=int(sys.argv[2]) if len(sys.argv)>2 else 1
 end=int(sys.argv[3]) if len(sys.argv)>3 else s.max_row+1
 for row in s.iter_rows(min_row=start,max_row=end-1):
  vals=[str(c.value) if c.value is not None else '' for c in row]
  if any(vals):print(row[0].row,json.dumps(vals,ensure_ascii=False))
