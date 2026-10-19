"""기준서/보고서의 현재 구조를 조사하고 동일 양식 후보를 검증한다."""
import hashlib
import json
import sys
import math
import collections
import shutil
import os
import copy
from pathlib import Path
import openpyxl

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
DHF = ROOT / 'docs/참고/인허가/사운드1/DHF'
QA = TASK / 'qa/criterion-report-alignment'
SPECS = [
    ('unit', 'SUVC', 'SUVR', '소프트웨어 유닛 검증', '기준서_유닛 검증', '보고서_유닛 시험', 4, 11),
    ('integration', 'SITC', 'SITR', '소프트웨어 통합시험', '기준서_통합 시험', '보고서_통합시험', 3, 13),
    ('system', 'SSTC', 'SSTR', '소프트웨어 시스템 시험', '기준서_시스템 시험', '보고서_시스템 시험', 2, 10),
]

def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def locate(folder, prefix, name, suffix):
    found = [p for p in folder.glob('*.xlsx') if not p.name.startswith(('~$', '[작성자 전용]'))
             and f'{prefix}-FW-TD24A {name} {suffix}.xlsx' in p.name]
    assert len(found) == 1, found
    return found[0]

def read_book(path):
    wb = openpyxl.load_workbook(path)
    sheets = {}
    for s in wb:
        sheets[s.title] = {
            'rows': [[s.cell(r,c).value for c in range(1,s.max_column+1)] for r in range(1,s.max_row+1)],
            'merges': [str(m) for m in s.merged_cells.ranges],
            'heights': {str(r): d.height for r,d in s.row_dimensions.items()},
            'widths': {c: d.width for c,d in s.column_dimensions.items()},
            'images': len(s._images), 'printArea': str(s.print_area),
        }
    return sheets

def inspect():
    QA.mkdir(parents=True, exist_ok=True)
    payload = {'pairs': []}
    for kind, cp, rp, name, cs, rs, idcol, endcol in SPECS:
        criterion = locate(DHF/'03_설계 출력/SW/펌웨어', cp, name, '기준서')
        report = locate(DHF/'04_설계 검증/SW/펌웨어', rp, name, '보고서')
        a,b = read_book(criterion),read_book(report)
        source_ids = [(i+1, row[idcol-1]) for i,row in enumerate(a[cs]['rows']) if row[idcol-1] and str(row[idcol-1]).startswith('SW-FW-')]
        report_ids = [(i+1, row[idcol-1]) for i,row in enumerate(b[rs]['rows']) if row[idcol-1] and str(row[idcol-1]).startswith('SW-FW-')]
        entry = dict(kind=kind, criterion=str(criterion), report=str(report), criterionHash=digest(criterion), reportHash=digest(report),
                     criterionSheet=cs, reportSheet=rs, idCol=idcol, commonEnd=endcol,
                     criterionSheets=a, reportSheets=b, sourceIds=source_ids, reportIds=report_ids)
        payload['pairs'].append(entry)
        print(kind, criterion.name, report.name, len(source_ids),len(report_ids))
        for tag,data,sname in [('criterion',a,cs),('report',b,rs)]:
            main=data[sname]
            print(tag,'sheets',list(data),'rows',len(main['rows']),'columns',len(main['rows'][0]),'images',main['images'])
            for i,row in enumerate(main['rows'][:38]):
                nonempty={openpyxl.utils.get_column_letter(j+1):str(v)[:160] for j,v in enumerate(row) if v is not None}
                if nonempty: print(i+1,nonempty)
            print('merges first',main['merges'][:12])
        for row,id_ in source_ids[:2]:
            print('TEST',row,id_,a[cs]['rows'][row-1])
    protected={str(p):digest(p) for p in DHF.rglob('*') if p.is_file() and (p.name.startswith('[작성자 전용]') or p.name.startswith('EPTC') or p in [Path(e['criterion']) for e in payload['pairs']])}
    payload['protected']=protected
    (QA/'baseline.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=str),encoding='utf-8')

def prepare():
    baseline=json.loads((QA/'baseline.json').read_text(encoding='utf-8'))
    unit_plan=json.loads((TASK/'qa/suvr/plan.json').read_text(encoding='utf-8'))
    unit_pairs={t['id']:t for suffix in ['cfx-hear','cm3','boot'] for t in json.loads((TASK/f'qa/suvr/pairs-{suffix}.json').read_text(encoding='utf-8'))}
    unit_meta={t['id']:t for t in unit_plan['tests']}
    plans=[]
    for e in baseline['pairs']:
        assert digest(Path(e['criterion']))==e['criterionHash']
        assert digest(Path(e['report']))==e['reportHash']
        kind=e['kind']; main=e['criterionSheets'][e['criterionSheet']]
        header=e['sourceIds'][0][0]-2
        # Latest files, not a reconstructed reference version, own the result content.
        old=e['reportSheets'][e['reportSheet']]['rows']
        starts=e['reportIds']; records=[]
        for idx,(row,test_id) in enumerate(e['sourceIds']):
            assert starts[idx][1]==test_id
            start=starts[idx][0]; stop=starts[idx+1][0]-1 if idx+1<len(starts) else len(old)
            old_plan=json.loads((TASK/('qa/suvr/direct-plan.json' if kind=='unit' else f'qa/validation-reports/{kind}-plan.json')).read_text(encoding='utf-8'))
            old_item=next(u for u in old_plan['units' if kind=='unit' else 'tests'] if (u['id']==test_id if kind!='unit' else test_id.endswith('-'+str(int(u['id'].split('-')[-1]))) and ('-'+u['id'].split('-')[0]+'-') in test_id))
            stop=old_item['end']
            comparison_end=old_item['comparisonEnd']
            comparisons=[]; notes=[]
            for r in range(start,comparison_end+1):
                exp=old[r-1][e['commonEnd']-1]; act=old[r-1][e['commonEnd']]
                if exp or act: comparisons.append(dict(expected=exp or '', actual=act or ''))
            for r in range(comparison_end+1,stop+1):
                label=old[r-1][7 if kind=='unit' else 9 if kind=='integration' else 6]
                act=old[r-1][e['commonEnd']]
                if (label and ('확인' in str(label) or '그림' in str(label))) or (act and '(QCC 검증 결과 필요)' in str(act)):
                    if act: notes.append(str(act))
            uid=old_item['id']
            if kind=='unit':
                extra=unit_meta[uid]['result']
                for k in ['exception','qcc','photo']:
                    if extra.get(k) and not any(extra[k] in n for n in notes): notes.append(extra[k])
            # Put each condition in a paragraph within the same result cell.
            text='\n'.join(c['actual'].replace('\n',' / ') for c in comparisons)
            compact_notes=[]
            for n in notes:
                n=n.replace('\n',' / ')
                n=n.replace('왼쪽은 예상 도식이며 실제 사진이 아니다.','').replace('왼쪽은 예상 도식이며 실제 측정 결과가 아니다.','')
                if n and n not in compact_notes: compact_notes.append(n.strip())
            if compact_notes: text+='\n'+'\n'.join(compact_notes)
            if kind=='unit' and any(l[0]==uid for l in unit_plan['logs']): text+='\nRTT 전문: 「RTT 결과 대조」 시트의 '+uid
            # Detailed/long observations remain in the support sheet, not new main rows.
            if kind=='system' and test_id.endswith('-10'):
                text='대역 1~32: 실제 출력 전극, 출력 유무 및 캡처 결과 기입 필요. 「시험 보조자료」의 같은 시험 ID에 대역별 기록 칸을 유지한다.\n'+comparisons[-1]['actual'].replace('\n',' / ')+'\n'+'\n'.join(compact_notes)
            actual_width=1450
            count=sum(max(1,math.ceil(sum(13.4 if ord(ch)>127 else 7.1 for ch in line)/(actual_width-18))) for line in text.split('\n'))
            height=max(main['heights'].get(str(row)) or 20.1,count*14+18)
            if height>409:
                # Keep the criterion-facing result concise, preserving all detailed values below.
                if kind=='unit':
                    text=unit_meta[uid]['result']['result']+'\n'+ '\n'.join(compact_notes)
                    text+='\n조건별 전체 결과: 「시험 보조자료」 / '+test_id
                    if any(l[0]==uid for l in unit_plan['logs']): text+='\nRTT 전문: 「RTT 결과 대조」 / '+uid
                else:
                    text='\n'.join(c['actual'].split(' / 현재 코드:')[0].replace('\n',' / ') for c in comparisons)
                    text+='\n조건별 전체 결과 및 확인 범위: 「시험 보조자료」 / '+test_id
                    for n in compact_notes:
                        if '(사진필요)' in n or '(QCC 검증 결과 필요)' in n: text+='\n'+n
                count=sum(max(1,math.ceil(sum(13.4 if ord(ch)>127 else 7.1 for ch in line)/(actual_width-18))) for line in text.split('\n'))
                height=max(main['heights'].get(str(row)) or 20.1,count*14+18)
            assert height<=409,(kind,test_id,height)
            judgment=old[start-1][e['commonEnd']+1]
            assert judgment
            records.append(dict(id=test_id,uid=uid,row=row,actual=text,judgment=judgment,height=height,
                                comparisons=comparisons,notes=notes,qcc='(QCC 검증 결과 필요)' in text,
                                photo='(사진필요)' in text))
        images=json.loads((TASK/'qa/suvr/expected-images/manifest.json').read_text(encoding='utf-8')) if kind=='unit' else [{**im,'tests':[im['test']]} for im in old_plan['images']]
        if kind=='unit':
            for im in images:
                im['path']=str(TASK/'qa/suvr/expected-images'/im['file'])
                im['tests']=[next(r['id'] for r in records if r['uid']==im['id'])]
        for record in records:
            matched=[i+1 for i,im in enumerate(images) if record['id'] in im['tests']]
            if matched:
                record['actual']+='\n예상 그림: 「예상 그림」 시트의 그림 '+', '.join(map(str,matched))
                record['height']=min(409,record['height']+28)
        e.update(records=records,header=header,actualWidth=1450,images=images)
        plans.append(e)
        print(kind,'rows',len(records),'maxHeight',max(r['height'] for r in records),'long',[(r['uid'],len(r['actual'])) for r in records if r['height']>370])
    (QA/'plan.json').write_text(json.dumps(plans,ensure_ascii=False,indent=2),encoding='utf-8')

def verify(published=False):
    baseline=json.loads((QA/'baseline.json').read_text(encoding='utf-8'))
    plans=json.loads((QA/'plan.json').read_text(encoding='utf-8'))
    checks=[]
    for p in plans:
        path=Path(p['report']) if published else QA/f"{p['kind']}-candidate.xlsx"
        a=openpyxl.load_workbook(p['criterion']); b=openpyxl.load_workbook(path)
        source=a[p['criterionSheet']]; main=b[p['reportSheet']]
        assert main.max_row==source.max_row,(p['kind'],'main row count',main.max_row,source.max_row)
        assert {str(m) for m in source.merged_cells.ranges}=={str(m) for m in main.merged_cells.ranges if m.max_col<=p['commonEnd']},p['kind']
        common=0
        for row in source:
            for c in row:
                if c.column>p['commonEnd']:continue
                out=main.cell(c.row,c.column)
                expected=str(c.value).replace('기준서','보고서') if c.coordinate=='B2' else c.value
                assert out.value==expected,(p['kind'],c.coordinate,expected,out.value)
                if c.has_style or c.value is not None:
                    attributes=['border'] if isinstance(c,openpyxl.cell.cell.MergedCell) else ['font','fill','border','alignment','protection','number_format']
                    for attr in attributes:
                        left,right=copy.copy(getattr(c,attr)),copy.copy(getattr(out,attr))
                        if attr=='border':
                            def visible_border(b):
                                return [(side,getattr(b,side).style,str(getattr(b,side).color)) for side in ['left','right','top','bottom','diagonal','start','end'] if getattr(b,side) is not None and getattr(b,side).style]
                            assert visible_border(left)==visible_border(right),(p['kind'],c.coordinate,attr)
                        else: assert left==right,(p['kind'],c.coordinate,attr)
                common+=1
        for k,d in source.column_dimensions.items():
            if openpyxl.utils.column_index_from_string(k)<=p['commonEnd']:
                assert d.width==main.column_dimensions[k].width,(p['kind'],k,'width')
        for record in p['records']:
            assert main.cell(record['row'],p['idCol']).value==record['id']
            assert main.cell(record['row'],p['commonEnd']+1).value==record['actual']
            assert main.cell(record['row'],p['commonEnd']+2).value==record['judgment']
        counts=collections.Counter(r['judgment'] for r in p['records'])
        assert len(p['records'])==len(set(r['id'] for r in p['records']))
        # The full, original comparison values are preserved in support rows.
        support=b['시험 보조자료']
        data=[(support.cell(r,1).value,support.cell(r,2).value,support.cell(r,3).value) for r in range(6,support.max_row+1)]
        for r in p['records']:
            for c in r['comparisons']:assert (r['id'],c['expected'],c['actual']) in data,(p['kind'],r['id'],'missing detail')
            for note in r['notes']:assert (r['id'],'확인 범위 및 첨부 자료',note) in data
        original=openpyxl.load_workbook(p['report']) if not published else openpyxl.load_workbook(QA/f"{p['kind']}-before.xlsx")
        for sheet in original:
            if sheet.title==p['reportSheet']:continue
            same=b[sheet.title]
            assert same.max_row==sheet.max_row and same.max_column==sheet.max_column
            assert {str(m) for m in same.merged_cells.ranges}=={str(m) for m in sheet.merged_cells.ranges}
            for row in sheet:
                for c in row:
                    d=same.cell(c.row,c.column)
                    assert c.value==d.value,(p['kind'],sheet.title,c.coordinate)
                    for attr in ['font','fill','border','alignment','number_format']:
                        assert copy.copy(getattr(c,attr))==copy.copy(getattr(d,attr)),(p['kind'],sheet.title,c.coordinate,attr)
        actual_images=b['예상 그림']._images
        expected=[hashlib.sha256(Path(im['path']).read_bytes()).hexdigest() for im in p['images']]
        actual=[hashlib.sha256(im._data()).hexdigest() for im in actual_images]
        assert collections.Counter(expected)==collections.Counter(actual),(p['kind'],'image bytes')
        assert len(actual_images)==len(original[p['reportSheet']]._images),(p['kind'],'picture count')
        assert all(c.data_type!='e' for s in b for row in s for c in row),(p['kind'],'Excel error')
        assert max(main.row_dimensions[r['row']].height for r in p['records'])<=409
        checks.append(dict(kind=p['kind'],tests=len(p['records']),mainRows=main.max_row,commonCells=common,
                           commonMerges=len(source.merged_cells.ranges),counts=dict(counts),images=len(actual_images),
                           criterionHash=digest(Path(p['criterion'])),reportHash=digest(path)))
    for path,before in baseline['protected'].items():assert digest(Path(path))==before,('protected file changed',path)
    result=dict(status='PASS',published=published,pairs=checks,protectedFiles=len(baseline['protected']))
    (QA/('published-checks.json' if published else 'checks.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))

def publish():
    verify()
    plans=json.loads((QA/'plan.json').read_text(encoding='utf-8'))
    for p in plans:
        assert digest(Path(p['report']))==p['reportHash'],('report edited externally',p['report'])
        assert digest(Path(p['criterion']))==p['criterionHash']
    for p in plans:
        before=QA/f"{p['kind']}-before.xlsx"
        if not before.exists():shutil.copy2(p['report'],before)
        temp=Path(p['report']).with_suffix('.alignment.tmp')
        shutil.copy2(QA/f"{p['kind']}-candidate.xlsx",temp)
        os.replace(temp,p['report'])
    verify(published=True)

if __name__=='__main__':
    mode=sys.argv[1] if len(sys.argv)>1 else 'inspect'
    {'prepare':prepare,'verify':verify,'publish':publish,'inspect':inspect}[mode]()
