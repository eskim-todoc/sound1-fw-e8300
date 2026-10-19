"""Snapshot the four authorized artifacts and inspect existing audit-unit rows."""
from pathlib import Path
import hashlib
import json
import shutil
import openpyxl
from inspect_sad_unit_trace import ROOT, TASK, DHF, SAD, SDDD, blocks

QA = TASK / 'qa/audit-units-four-docs-20261012'
FILES = {'sad': SAD, 'sddd': SDDD,
         'suvc': DHF / '03_설계 출력/SW/펌웨어/SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서_rev0(김은수 작성).xlsx',
         'suvr': DHF / '04_설계 검증/SW/펌웨어/SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서_rev0(김은수 작성).xlsx'}

def main():
    QA.mkdir(parents=True, exist_ok=True)
    manifest = {}
    result = {}
    for kind, path in FILES.items():
        target = QA / (kind + '-before' + path.suffix)
        if target.exists():
            assert target.read_bytes() == path.read_bytes(), 'Current file differs from snapshot'
        else:
            shutil.copyfile(path, target)
        manifest[kind] = {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
        if path.suffix == '.xlsx':
            wb = openpyxl.load_workbook(path, data_only=False)
            result[kind] = {'sheets': wb.sheetnames, 'units': []}
            for ws in wb:
                for row in ws:
                    if any(str(c.value).strip() in ['SW-FW-UnitTest-CM3-ID-1', 'SW-FW-UnitTest-CM3-ID-5', 'SW-FW-UnitTest-CM3-ID-46'] for c in row):
                        start = row[0].row
                        end = start + 1
                        while end <= ws.max_row and not str(ws.cell(end, 4).value).strip().startswith('SW-FW-UnitTest-'):
                            end += 1
                        result[kind]['units'].append({'sheet': ws.title, 'start': start, 'end': end - 1,
                            'rows': [{'row': r, 'height': ws.row_dimensions[r].height, 'cells': {ws.cell(r,c).column_letter: ws.cell(r,c).value for c in range(1,ws.max_column+1) if ws.cell(r,c).value is not None}} for r in range(start,end)],
                            'merges': [str(m) for m in ws.merged_cells.ranges if m.min_row <= end-1 and m.max_row >= start]})
        else:
            result[kind] = blocks(path)
    protected = [p for p in DHF.rglob('*') if p.suffix in ('.docx','.xlsx') and (p.name.startswith(('SIP-', 'SITC-', 'SITR-', 'SRS-', '[작성자 전용]')))]
    manifest['protected'] = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in protected}
    (QA / 'baseline.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
    (QA / 'inspection.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k: v for k,v in result.items() if k in ('suvc','suvr')},ensure_ascii=False,indent=2))

if __name__ == '__main__':
    main()
