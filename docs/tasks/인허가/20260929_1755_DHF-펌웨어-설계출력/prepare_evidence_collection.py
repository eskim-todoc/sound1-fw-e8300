"""검토한 문구를 편집 계획으로 연결한다. workbook 읽기/대조만 openpyxl로 수행한다."""
from __future__ import annotations
import json
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from review_evidence_collection import QA

BASE = json.loads((QA / 'baseline.json').read_text(encoding='utf-8'))
BOOKS = {b['key']: b for b in BASE['books']}
PATCHES = {}
ORIGINAL_BOOKS = {}

def put(key, cell, value, sheet=None, **extra):
    sheet = sheet or BOOKS[key]['sheet']
    if isinstance(value,str):
        value=value.replace('·','/').replace('전사','기록').replace('no photo, text record','사진 없이 텍스트로 기록')
        value=value.replace('사진 없이 텍스트로 기록로','사진 없이 텍스트로')
        value=value.replace('FFT raw','FFT 원본값').replace('vMag raw','vMag 원본값').replace('raw OD','OD 핀의 PDM 출력')
        if key.endswith('-report') and (any(mark in value for mark in (
            '(사진필요)', '(QCC 검증 결과 필요)', '(시험 결과 기입 필요)',
            '(자료 필요)', '[계측 준비 필요]')) or (len(value)<120 and ('필요' in value or '대기' in value))):
            extra.setdefault('red', True)
    PATCHES[key, sheet, cell] = {'key': key, 'sheet': sheet, 'cell': cell, 'value': value, **extra}

def original(key, cell, sheet=None):
    if key not in ORIGINAL_BOOKS:
        ORIGINAL_BOOKS[key] = load_workbook(BOOKS[key]['backup'], rich_text=True)
    b = ORIGINAL_BOOKS[key]
    value = b[sheet or BOOKS[key]['sheet']][cell].value
    return '' if value is None else str(value)

def current(key, cell, sheet=None):
    p = PATCHES.get((key, sheet or BOOKS[key]['sheet'], cell))
    return p['value'] if p else original(key, cell, sheet)

def main():
    for path in sorted(QA.glob('audit-*.json')):
        data = json.loads(path.read_text(encoding='utf-8-sig'))
        for p in data.get('patches', []):
            p = dict(p)
            key = p.pop('key'); cell = p.pop('cell'); value = p.pop('value')
            if key.startswith('unit-report-'):
                key = 'unit-report'
            put(key, cell, value, **p)

    # Remove obsolete statements that would contradict the revised procedure/verdict.
    put('unit-report','L37',current('unit-report','L37').replace(
        '기존 판정은 통과로 유지한다.',
        '상태 설정/조회 결과는 보존하되 실제 수신 판정은 확인 필요로 남긴다.'))
    put('unit-report','L115',current('unit-report','L115').replace(
        '기준서의 A/B 맵번호와 볼륨 입력은 구체 숫자가 없으므로 새 값을 만들지 않았다. 숫자로 비교 가능한 저장 자료는 MD의 T/C 및 파일 길이 벡터로 제시함.',
        '아래 A/B 맵번호와 볼륨은 새 자료 수집용 입력 예다. 기존 T/C와 파일 길이 기록을 맵번호/볼륨의 실측값으로 옮기지 않는다.'))
    put('unit-report','L142',current('unit-report','L142').replace(
        '실제 재시작은 시험 훅으로 차단',
        'PC 전용 시험에서 하드웨어 reset을 호출 횟수 기록으로 대체'),red=True)

    # Keep preparation cells readable within Excel's fixed 409-point row limit.
    verbose=('수동 입력 위치와 준비된 입력 전달 수단(시험 전용 함수/디버거 입력/실제 송신기)을 시험 전에 확인하고 이미지/함수/단계를 기록한다. '
             '정상 루프/다른 코어가 쓰는 배열을 실행 중 임의 덮어쓰지 않는다. 안전한 입력 준비 뒤 필요한 실제 모듈 연결을 실행하고 완료한 자료를 보관한다. '
             '준비된 통합 입력 경로가 없으면 시험 준비 미완료로 남기며 유닛 ut auto/observe를 전체 통합 시험으로 대체하지 않는다.')
    brief=('입력 수단과 적용 이미지/단계를 기록한다. 공유 버퍼는 관련 루프/ISR을 정지한 상태에서 준비한다. '
           '실제 모듈 연결을 실행하고 완료 자료를 수집한다. 통합 입력 경로가 준비되지 않았으면 미실행으로 남긴다.')
    for patch in list(PATCHES.values()):
        if patch['key']=='integration-criterion' and patch['sheet']==BOOKS['integration-criterion']['sheet']:
            put(patch['key'],patch['cell'],patch['value'].replace(verbose,brief),patch['sheet'])
    compact_prep={
        'J29':('- 혼합 입력: 단일 경로에 16표본=[0,-1024,544,-3200,3168,나머지0].\n'
                '- HEAR FFT 입력: 512개 0 또는 1000 Hz 순음 x[n]=round(65536*sin(2*pi*32*n/512)).\n'
                '- PCM 입력: 매핑 뒤 크기100, CIS/32대역/전극[i]=i+1, 2프레임/자극, 12채널/ms.\n'
                '세 입력은 별도 디지털 조건이다. 관련 루프/ISR 정지 후 입력을 준비하고 실제 모듈 연계를 실행한다. 입력 경로/이미지와 HEAR 완료 Event/전체 PCM을 저장한다. 유닛 RAM 결과만으로 전체 통합을 통과시키지 않는다.'),
        'J33':('- 매핑 연결에서 지정 자극/Live 맵/종료를 요청한다. Live는 32대역 NofM, 펄스폭14us, 전극1~32로 설정한다.\n'
                '적용 SPI/매핑 프로토콜의 전체 필드/길이를 확인한 명령을 시험 송신기로 보내고 RX/TX 원문을 저장한다. CM3 요청→CFX 완료와 공유32개 T/C/모드를 같은 조건으로 기록한다. '
                '공유 입력은 관련 루프/ISR 정지 후 준비한다. QCC를 대신한 SPI 입력은 무선/앱 검증이 아니다. 입력 경로 미준비 시 미실행으로 남긴다.'),
        'J37':('- I2S/마이크 각각 16표본=320, Gain A=65536(1.0)/B=131072(2.0). 안정 수신에서 일반/크래들/매핑 연결을 각각 적용한다.\n'
                '고정 표본은 혼합 입력 경계에 관련 루프/ISR 정지 후 준비한다. 실제 I2S 공급과 시험 버퍼 공급을 구분한다. '
                '0x36 종류 변경과 게인 직접 고정을 별도 조건으로 실행하고 입력 경로/이미지/수신 상태/혼합 결과를 저장한다. 통합 입력 경로 미준비 시 미실행으로 남긴다.'),
        'J39':('- NofM/32대역/전극[i]=i+1, 3프레임/자극, 8채널/ms, phase=0, 묵음 꺼짐.\n'
                '- 빔포밍 입력: 기준0/지연 현재[i]=64*i/지연 직전[0]=2048. 선택 입력: 매핑 출력32개 모두100 또는 값[i]=i.\n'
                '두 벡터는 혼합 전과 FFT/매핑 뒤의 별도 입력이다. 관련 루프/ISR 정지 후 직전 버퍼/맵/볼륨을 초기화한다. 실제 연계를 실행하고 입력 경로/이미지, 선택/배치/분할 PCM 전체 배열을 저장한다. 전극 전류/음향 실측이 아니다. 경로 미준비 시 미실행으로 남긴다.'),
    }
    for cell,value in compact_prep.items():
        put('integration-criterion',cell,value)
    for cell,note in {
        'N31':'(QCC 검증 결과 필요) QCC 초기화/오디오 입력 준비 상태와 E8300 기동 단계의 실제 대응 자료.',
        'N38':'(QCC 검증 결과 필요) QCC_CTRL/ISD_CHECK 상태를 QCC가 사용한 결과와 실제 연결/해제 단계의 대응 자료.',
    }.items():
        value=current('integration-report',cell)
        if '(QCC 검증 결과 필요)' not in value:
            put('integration-report',cell,value+'\n'+note,red=True)

    # Preserve numerical acceptance limits. Record the input/output disagreement instead
    # of deciding a new release requirement in an evidence-collection wording task.
    for cell in ('J34','J35'):
        PATCHES.pop(('system-criterion',BOOKS['system-criterion']['sheet'],cell),None)
    put('system-criterion','H34',current('system-criterion','H34')+'\n설계 입력에는 20 kHz 이상 외에 24 kHz/±5% 조건도 함께 있다. 수치 기준의 적용 관계는 확인 필요하며 기존 기준을 임의로 바꾸지 않는다.')
    put('system-criterion','H35',current('system-criterion','H35')+'\n설계 입력의 1000 Hz 이상 조건과 출력 기준의 1000 Hz ±5% 조건은 다르므로 적용 관계를 확인해야 한다. 측정값은 두 조건과 각각 비교해 남긴다.')

    # CFX-020: two short RAM vectors do not generate a continuous audio waveform.
    key = 'unit-criterion'
    put(key, 'F49', 'T, D')
    put(key, 'H49', 'C1: 입력 src[i]=i(i=0~15), 출력용 RAM sink[0~15]=0.\nC2: src[0~4]=[-131073,-131072,0,131071,131072], src[5~15]=0. C1과 C2를 각각 초기화하여 실행한다.\n청취 보조 확인: 자동 시험을 종료한 정상 오디오 경로에 연속 1000 Hz 사인파와 무음을 차례로 입력한다.')
    put(key, 'I49', '자동 시험은 lib_loopback_AGC_out의 RAM 계산값을 확인한다. C1은 증가 배열, C2는 포화 경계 배열이며 사인파 표본이 아니다. DAC FIFO/OD 핀으로 자동 재생되지 않는다.\nOD1_P/OD1_N은 PDM 차동 출력이다. 핀에서 직접 사인파를 찾지 않는다. 청취는 보드에 맞는 수신기를 사용한다. 스코프 확인을 추가할 경우 차동 저역통과 필터의 출력을 측정한다.')
    put(key, 'J49', '1. 준비된 CFX 자동 그룹(ut cfx 1 또는 ut auto 1)을 실행한다. RTT에서 UT-CFX-020으로 찾고 C1/C2의 배열과 PASS/FAIL 줄을 비교한다. 전체 .log를 보관하며 정수 시험에 스코프 사진은 필요 없다.\n2. 자동 시험을 종료하고 정상 OD 출력 경로에서 1000 Hz 사인파/무음을 각각 재생한다. 순음이 들리는지와 무음에서 순음이 사라지는지를 글로 기록한다.\n3. 파형을 추가 확인할 때만 필터 후 OUT+/OUT- 차동 파형과 입력 주파수를 저장한다. 스코프 GND는 보드 GND에 연결하며 OD1_N에 연결하지 않는다.')
    put(key, 'K49', 'C1: RAM sink[0~15]=[0,64,128,192,256,320,384,448,512,576,640,704,768,832,896,960], src[0~15]=0~15 유지.\nC2: RAM sink[0~4]=[-8388608,-8388608,0,8388544,8388607], sink[5~15]=0.\nRTT 요약: [UT-CFX-020] C1 PASS N=32 F=0 / [UT-CFX-020] C2 PASS N=16 F=0.\n청취 보조: 1000 Hz 순음/무음이 구분되어야 한다. 사진은 청취 증거가 아니다. 필터 후 파형의 참고 주기는 1 ms이며 청취로 정확한 이득, 전압, 전극 전류를 판정하지 않는다.')
    result = original('unit-report', 'L49').replace('FIFO 출력', 'RAM sink 출력')
    result = result.replace('[3] OD 청취, 1000 Hz 순음과 무음 입력 / 순음 입력 시 청취=순음 구분; 무음 입력 시 청취=무음 구분', '[3] 정상 OD 경로의 보조 청취, 연속 1000 Hz 사인파와 무음 입력\n(시험 결과 기입 필요) 실제 입력/볼륨/연결한 수신기=____; 순음 청취=____; 무음 청취=____')
    result = result.replace('(사진필요) 1 kHz 신호/무음의 실제 출력 관찰 화면 및 포화 경계 정수 결과', '수집: C1/C2는 UT-CFX-020 배열/판정 줄이 포함된 RTT 원본 .log를 보관한다. 정수 결과의 스코프 사진은 불필요하다.\n청취 기록: 정상 OD 경로의 1000 Hz 사인파/무음 입력과 들린 소리를 글로 기록한다. 자동 벡터가 재생한 파형으로 취급하지 않는다.\n파형을 추가 측정하는 경우에만 (사진필요) 저역통과 필터 후 OUT+/OUT-의 차동 파형. 원래 OD 핀의 PDM 파형은 사인파 비교에 사용하지 않는다.')
    put('unit-report', 'L49', result, red=True)
    put('unit-report', 'M49', '수치 통과 / 청취 확인 필요', red=True)
    for row in (85, 86):
        for column in 'BC':
            a=f'{column}{row}';put('unit-report',a,original('unit-report',a,'시험 보조자료').replace('FIFO 출력','RAM sink 출력'), '시험 보조자료')
    put('unit-report','B87','[3] 정상 OD 경로의 보조 청취: 연속 1000 Hz 사인파와 무음\n순음/무음이 구분되어야 한다. C1/C2의 RAM 벡터와 별개다.','시험 보조자료')
    put('unit-report','C87','(시험 결과 기입 필요) 정상 Live 자극 경로에서 사용한 1000 Hz 입력/볼륨/수신기=____; 순음 청취=____; 무음 청취=____. 자동 RAM 시험의 계산 통과를 청취 완료로 취급하지 않는다.','시험 보조자료',red=True)
    put('unit-report','D87','청취 확인 필요','시험 보조자료',red=True)
    put('unit-report','C89','RTT: UT-CFX-020 C1/C2 배열과 판정 줄을 원본 .log로 보관한다. 청취는 입력 조건과 소리를 글로 기록하며 사진을 요구하지 않는다.\n파형을 추가 확인하는 경우에만 (사진필요) 1000 Hz 사인파/무음 입력에서 저역통과 필터 후 OUT+/OUT- 차동 파형.','시험 보조자료',red=True)
    put('unit-report','A6','그림 1. 정상 OD 경로에 1000 Hz 사인파/무음을 입력했을 때 저역통과 필터 후 차동 파형의 참고 모습. 자동 RAM 시험의 출력이 아님.\nSW-FW-UnitTest-CFX-ID-20','예상 그림')
    put('unit-report','C6','청취만 확인할 때는 사진 불필요. 파형을 추가 확인할 때만 (사진필요) 필터 후 차동 파형을 첨부한다. OD 핀의 원래 출력은 PDM이다.','예상 그림',red=True)
    put('unit-report','C651',original('unit-report','C651','입력 및 예상값')+'\nC1은 증가 배열, C2는 포화 경계 배열로 연속 사인파가 아니다. 물리 청취는 자동 시험 종료 후 정상 경로에서 별도로 한다. OD 핀은 PDM이며 필터 후 차동 출력에서만 사인파 모양을 확인한다.','입력 및 예상값')

    # One saved RTT file can cover a whole ready-made group. Do not ask for one scope
    # or screen photo per calculation vector and do not invent commands for other IDs.
    auto={
        'CFX':{7,8,10,11,12,13,*range(15,21),*range(22,38),40},
        'CM3':{4,6,7,9,10,11,12,13,*range(16,25),25,26,27,28,29,30,34,35,36,37,40,44,45,47,50,51},
        'HEAR':{1,2}, 'BOOT':{8,9,10,11,12,13,14,19,20},
    }
    cmd={'CFX':'ut cfx 1','CM3':'ut cm3 1','HEAR':'ut hear 1','BOOT':'ut boot 1'}
    for t in BOOKS['unit-criterion']['tests']:
        m=re.fullmatch(r'SW-FW-UnitTest-([A-Z0-9]+)-ID-(\d+)',t['id']);core,num=m[1],int(m[2])
        if num not in auto.get(core,set()) or (core,num)==('CFX',20):continue
        a=f"J{t['row']}";tag=f'UT-{core}-{num:03d}'
        put('unit-criterion',a,current('unit-criterion',a)+f'\n수집: 준비된 {core} 그룹({cmd[core]}) 실행 후 {tag}의 조건별 배열/상태와 판정 줄을 찾는다. 그룹 전체 원본 .log를 저장한다. RTT 화면은 필요할 때만 캡처하며 스코프 사진을 요구하지 않는다.')
    for row,tag in ((72,'HEAR-001'),(73,'HEAR-002')):
        a=f'L{row}';v=current('unit-report',a)
        v=re.sub(r'\(사진필요\)[^\n]*',f'(자료 필요) {tag}의 실제 RTT 배열/완료 상태와 판정 줄이 포함된 원본 .log. 화면 캡처는 선택이며 스코프 사진은 불필요.',v)
        put('unit-report',a,v,red=True)
    put('unit-report','M139','요청 기록 확인 / 실제 전송 확인 필요',red=True)
    put('unit-report','M142','PC 요청 확인 / LED 자료 필요',red=True)

    # The old generic "photo of this test" wording does not specify what to collect.
    photo_kinds={
        20:'조건별 LED의 실제 색/점등 사진. 점멸 조건은 한 반복 이상 보이는 영상과 관찰 기록',
        21:'충전 연결/해제와 잔량 조건이 구분된 실제 LED 사진. 점멸은 영상',
        22:'잔량 구간별 실제 LED 사진. 점멸 구간은 영상',
        27:'LED 알림 켜짐/꺼짐 설정별 실제 LED 사진과 보낸 명령/설정 조회값',
        33:'리모콘 LED 알림 설정/조회 화면과 해당 조건의 실제 LED 사진',
        62:'충전 연결/해제와 잔량 구간별 실제 LED 사진. 점멸은 영상',
        64:'같은 색의 두 점멸 패턴이 각각 한 반복 이상 보이는 실제 영상과 구별 결과',
        29:'리모콘의 맵 번호 조회 화면과 송신 명령/수신 응답 원문',
        30:'리모콘의 시스템 상태 표시 화면과 같은 시점 장치 상태/응답 원문',
        31:'리모콘 오디오 볼륨 조회 화면과 송신 명령/응답 원문',
        32:'리모콘 자극 볼륨 조회 화면과 송신 명령/응답 원문',
        35:'매핑 프로그램의 맵 표시 화면과 읽어온 맵 데이터',
        38:'매핑 프로그램의 시스템 상태 화면과 같은 시점 장치 상태',
        39:'매핑 프로그램의 선택 전극/설정 표시 화면과 출력 조회값',
        40:'매핑 프로그램의 실시간 자극 상태 화면과 같은 입력의 장치 기록',
        63:'표시된 배터리 잔량 화면과 같은 시점 QCC 잔량 통보/장치 조회값',
        75:'오류를 발생시킨 입력과 실제 앱 오류 화면/응답 원문. 키/비밀값은 포함하지 않음',
    }
    for t in BOOKS['system-report']['tests']:
        num=int(t['id'].rsplit('-',1)[-1]);a=f"K{t['row']}";v=current('system-report',a)
        generic='(사진필요) 이 시험의 실제 화면/파형/표시 사진을 첨부한다.'
        if generic not in v:continue
        what=photo_kinds.get(num)
        if what is None:
            what='시험 맵/음원/전극/부하가 기록된 PBS 또는 32채널 시험 지그의 실제 전극 파형. 해당 입력의 출력 유무/순서/변화를 확인할 수 있는 시간축 사용'
        put('system-report',a,v.replace(generic,'(사진필요) '+what+'.'),red=True)

    # Copy changed common fields to the corresponding report. Never add main test rows.
    for kind in ('unit', 'integration', 'system'):
        ck=f'{kind}-criterion'; rk=f'{kind}-report'
        cr={r['id']:r for r in BOOKS[ck]['tests']}; rr={r['id']:r for r in BOOKS[rk]['tests']}
        for patch in list(PATCHES.values()):
            if patch['key']!=ck or patch['sheet']!=BOOKS[ck]['sheet']:continue
            row=int(re.search(r'\d+',patch['cell'])[0]); column=re.match(r'[A-Z]+',patch['cell'])[0]
            record=next((r for r in cr.values() if r['row']==row),None)
            if record:put(rk,column+str(rr[record['id']]['row']),patch['value'])
            elif row<=BOOKS[rk]['maxRow'] and len(column)==1 and ord(column)-64<=BOOKS[ck]['common'] and patch['cell']!='B2':
                put(rk,patch['cell'],patch['value'])

    # Keep the current [+] rows synchronized and color only genuinely differing text strongly.
    for key in ('unit-criterion','integration-criterion','system-input','system-criterion'):
        book=load_workbook(BOOKS[key]['backup'],rich_text=True)
        diff=next((s for s in book if '변경 비교' in s.title),None)
        if diff is None:continue
        idcol=BOOKS[key]['idCol']; byid={r['id']:r for r in BOOKS[key]['tests']}
        for patch in list(PATCHES.values()):
            if patch['key']!=key or patch['sheet']!=BOOKS[key]['sheet']:continue
            row=int(re.search(r'\d+',patch['cell'])[0]); column=re.match(r'[A-Z]+',patch['cell'])[0]
            record=next((r for r in byid.values() if r['row']==row),None)
            if not record:continue
            matches=[r for r in range(1,diff.max_row+1) if str(diff.cell(r,idcol).value)==record['id'] and str(diff.cell(r,1).value)!='[-]']
            assert len(matches)==1,(key,record['id'],matches)
            dr=matches[0]; minus=dr-1 if str(diff.cell(dr-1,1).value)=='[-]' else None
            old='' if minus is None or diff[f'{column}{minus}'].value is None else str(diff[f'{column}{minus}'].value)
            put(key,f'{column}{dr}',patch['value'],diff.title, diffOld=old, diffMinus=f'{column}{minus}' if minus else None)
        book.close()

    # Re-evaluate heights without shrinking existing rows. Excel has a 409 point row limit.
    for key,b in BOOKS.items():
        book=load_workbook(b['backup'],rich_text=True)
        sheets={p['sheet'] for p in PATCHES.values() if p['key']==key}
        for sn in sheets:
            s=book[sn]; rows={int(re.search(r'\d+',p['cell'])[0]) for p in PATCHES.values() if p['key']==key and p['sheet']==sn}
            for row in rows:
                old=s.row_dimensions[row].height or 15
                needed=old
                for col in range(1,s.max_column+1):
                    a=f'{get_column_letter(col)}{row}'
                    p=PATCHES.get((key,sn,a))
                    if not p:continue
                    width=(s.column_dimensions[get_column_letter(col)].width or 13)*7
                    val=str(p['value']); lines=0
                    for line in val.split('\n'):
                        size=sum(13.4 if ord(ch)>127 else 7.1 for ch in line)
                        lines+=max(1,-(-int(size)//max(1,int(width-18))))
                    needed=max(needed,lines*14+18)
                # Preserve the original row height if cells are already laid out wider than this estimate.
                needed=min(409,needed)
                for p in PATCHES.values():
                    if p['key']==key and p['sheet']==sn and int(re.search(r'\d+',p['cell'])[0])==row:p['height']=needed
        book.close()
    data=list(PATCHES.values())
    (QA/'patch-plan.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
    print('patches',len(data),'books',sorted({p['key'] for p in data}))
    for book in ORIGINAL_BOOKS.values():
        book.close()

if __name__=='__main__':main()
