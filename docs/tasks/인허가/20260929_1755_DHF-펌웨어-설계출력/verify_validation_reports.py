"""저장한 보고서를 다시 읽어 기준 보존, 결과 구분과 그림 배치를 검사한다."""
import hashlib
import json
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import openpyxl

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
QA = TASK / 'qa/validation-reports'
NS = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def read(name):
    return json.loads((QA / name).read_text(encoding='utf-8'))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def verify(kind):
    plan = read(f'{kind}-plan.json')
    criteria = read(f'{kind}-criteria.json')
    path = QA / f'{kind}-candidate.xlsx'
    book = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    name = '보고서_통합시험' if kind == 'integration' else '보고서_시스템 시험'
    c_name = '기준서_통합 시험' if kind == 'integration' else '기준서_시스템 시험'
    sheet = book[name]
    id_col = ord(plan['idCol'])-64
    expected_col = ord(plan['expectedCol'])-64
    actual_col = expected_col+1
    final_col = actual_col+1
    source_tests = {r[id_col-1]:r for r in criteria[c_name]['rows'] if str(r[id_col-1] or '').startswith('SW-FW-')}
    tests = {t['id']:t for t in plan['tests']}
    assert set(source_tests) == set(tests), 'Criterion ID set changed'
    for test_id, test in tests.items():
        raw = source_tests[test_id]
        for col in range(2, expected_col):
            assert sheet.cell(test['start'],col).value == raw[col-1], (test_id,col,'criterion field')
        assert sheet.cell(test['scopeRow'],expected_col).value == raw[expected_col-1], (test_id,'original pass criterion')
        assert sheet.cell(test['start'],final_col).value == test['judgment']
        if test['judgment'] == '통과':
            method = raw[8 if kind == 'integration' else 3]
            assert method == 'A', (test_id,'unconfirmed hardware method marked pass',method)
        for row in range(test['start'],test['comparisonEnd']+1):
            exp,act = [sheet.cell(row,col).value for col in [expected_col,actual_col]]
            assert exp and act, (test_id,row,'empty comparison')
            assert exp.split('\n')[0] == act.split('\n')[0], (test_id,row,'input differs')
        if test['qcc']:
            cell = sheet.cell(test['scopeRow'],actual_col)
            assert '(QCC 검증 결과 필요)' in cell.value
            assert cell.font.color and cell.font.color.type == 'rgb' and cell.font.color.rgb[-6:] == 'C00000'
    for item in plan['meta']:
        row = item['row']
        assert 40 <= sheet.row_dimensions[row].height <= 400, (row,'row height')
        for col in [expected_col,actual_col]:
            c = sheet.cell(row,col)
            assert c.font.sz == 10 and c.alignment.wrap_text, (row,col,'font/wrap')
        if item['pending']:
            c = sheet.cell(row,actual_col)
            assert c.font.color and c.font.color.type == 'rgb' and c.font.color.rgb[-6:] == 'C00000'
    assert len(sheet._images) == len(plan['images'])
    assert len(book['표지']._images) == 0, 'Old signature remains'
    for expected, actual in zip(plan['images'],sheet._images):
        assert actual.anchor._from.row == expected['row']-1
        assert actual.anchor._from.col == plan['expectedIndex']
        assert hashlib.sha256(actual._data()).hexdigest() == sha(expected['path'])
        assert '(사진필요)' in sheet.cell(expected['row'],actual_col).value
    counts = Counter(t['judgment'] for t in plan['tests'])
    for cell, formula in plan['formulas'].items():
        assert sheet[cell].value == formula
        status = sheet.cell(sheet[cell].row,actual_col).value
        assert values[name][cell].value == counts[status], (cell,values[name][cell].value,counts[status])
    formulas=[]
    errors=[]
    for sh in book:
        for row in sh:
            for c in row:
                if c.data_type == 'f': formulas.append((sh.title,c.coordinate,c.value))
                if c.data_type == 'e': errors.append((sh.title,c.coordinate,c.value))
    assert not errors, errors
    assert len(formulas) == len(plan['formulas']), ('stale formula count',formulas[:15])
    for c in ['J12','M12','D14','H14','J14','M14','D16','H16','J16','M16','D24']:
        assert book['표지'][c].value is None, ('old cover value',c)
    assert book['표지']['B24'].value == 0
    assert book['표지']['H12'].value == '김은수'
    with ZipFile(path) as z:
        z.testzip()
        for xml in [n for n in z.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml')]:
            root=ET.fromstring(z.read(xml))
            fs=root.findall('.//s:f',NS)
            assert all(f.text for f in fs), ('empty/shared formula',xml)
    return {'kind':kind,'sha256':sha(path),'tests':len(tests),'comparisons':plan['comparisonCount'],
            'images':len(plan['images']),'judgments':dict(counts),'formulas':len(formulas),'qcc':sum(t['qcc'] for t in plan['tests'])}


if __name__ == '__main__':
    reports=[verify(k) for k in ['integration','system']]
    protected=read('protected-before.json')
    changes=[p for p,h in protected.items() if not (ROOT/p).exists() or sha(ROOT/p)!=h]
    assert not changes, ('Protected files changed',changes)
    for item in read('sources.json').values():
        assert sha(item['path']) == item['sha256'], ('Source changed',item['path'])
    output=ROOT/'docs/참고/인허가/사운드1/DHF/04_설계 검증/SW/펌웨어'
    tracker=(ROOT/'docs/참고/인허가/사운드1/QCC 자료 추가 위치 관리.md').read_text(encoding='utf-8')
    published=[]
    for kind, name in [('integration','SITR-FW-TD24A 소프트웨어 통합시험 보고서.xlsx'),
                       ('system','SSTR-FW-TD24A 소프트웨어 시스템 시험 보고서.xlsx')]:
        p=output/name
        if p.exists():
            assert sha(p)==sha(QA/f'{kind}-candidate.xlsx'), ('Publication differs',name)
            plan=read(f'{kind}-plan.json')
            for t in plan['tests']:
                if t['qcc']:
                    assert f'{plan["actualCol"]}{t["scopeRow"]}' in tracker, (name,t['id'],'tracker location')
            published.append({'path':str(p),'sha256':sha(p)})
    result={'reports':reports,'protectedFiles':len(protected),'protectedChanges':changes,'sourceFiles':4,'published':published}
    (QA/'checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
