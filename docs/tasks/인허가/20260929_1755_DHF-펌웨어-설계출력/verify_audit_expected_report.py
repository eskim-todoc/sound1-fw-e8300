"""Preserve native XLSX parts; check planned values and scope before publication."""
from copy import deepcopy
import json
import re
import zipfile
from lxml import etree as E
import openpyxl
from prepare_audit_expected_report import QA, REPORT, sha

# Another active task updated only the SRS workbook while this report was prepared.
# Keep the initial snapshot intact and record that external change separately.
CONCURRENT_CHANGES = {
    str(REPORT.parents[3] / '02_설계 입력' / 'SW' / '펌웨어' /
        '[김은수 작성] [Rev.0] SRS-FW-TD24A 소프트웨어 요구사항 사양서.xlsx'):
        'e49854acc42635ef00997aa1a047f384a3d0c4e8f1a966179da1775d4353a38c',
}

S = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS = {'s': S}
def q(x): return '{' + S + '}' + x
def unpack(path):
    with zipfile.ZipFile(path) as z: return {n:z.read(n) for n in z.namelist()}
def cell_value(cell, strings):
    if cell.get('t') == 's': return ''.join(strings[int(cell.find(q('v')).text)].itertext())
    if cell.get('t') == 'inlineStr': return ''.join(cell.find(q('is')).itertext())
    v = cell.find(q('v'))
    return v.text if v is not None else ''
def font_properties(styles, cell):
    xf = styles.find(q('cellXfs'))[int(cell.get('s','0'))]
    font = styles.find(q('fonts'))[int(xf.get('fontId','0'))]
    pr = E.Element(q('rPr'))
    for child in font:
        c = deepcopy(child)
        if c.tag == q('name'): c.tag = q('rFont')
        pr.append(c)
    return pr
def rich_text(cell, text, styles, full_blue=False):
    pr = font_properties(styles, cell)
    # Text/attributes other than type remain; the original cell style owns borders/layout.
    for c in list(cell): cell.remove(c)
    cell.set('t','inlineStr')
    target = E.SubElement(cell,q('is'))
    idx = text.find('[E8300 보완]')
    segments = [(text, True)] if full_blue else ([(text[:idx],False),(text[idx:],True)] if idx >= 0 else [(text,False)])
    for segment,blue in segments:
        if not segment: continue
        r = E.SubElement(target,q('r'))
        rp = deepcopy(pr)
        if blue:
            for tag in ('b','color','u'):
                for child in rp.findall(q(tag)): rp.remove(child)
            E.SubElement(rp,q('b'))
            E.SubElement(rp,q('color'),{'rgb':'FF0070C0'})
        r.append(rp)
        E.SubElement(r,q('t'),{'{http://www.w3.org/XML/1998/namespace}space':'preserve'}).text = segment

def preserve():
    plan = json.loads((QA/'sheet-edits.json').read_text(encoding='utf8'))
    original = unpack(QA/'report-before.xlsx')
    authored = unpack(QA/'report-artifact.xlsx')
    root = E.fromstring(original['xl/worksheets/sheet2.xml'])
    src = E.fromstring(authored['xl/worksheets/sheet2.xml'])
    strings = E.fromstring(authored['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in authored else []
    styles = E.fromstring(original['xl/styles.xml'])
    for address,want in plan['edits'].items():
        dst = root.find(f'.//s:c[@r="{address}"]',NS)
        source = src.find(f'.//s:c[@r="{address}"]',NS)
        assert dst is not None and source is not None, address
        actual = cell_value(source,strings)
        assert actual == want, address
        rich_text(dst,actual,styles,address in plan['full_blue'])
    for row,height in plan['heights'].items():
        node = root.find(f'.//s:row[@r="{row}"]',NS)
        node.set('ht',str(height)); node.set('customHeight','1')
    original['xl/worksheets/sheet2.xml'] = E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    with zipfile.ZipFile(QA/'report.xlsx','w',zipfile.ZIP_DEFLATED) as z:
        for name,data in original.items(): z.writestr(name,data)

def check(path=None):
    path = path or QA/'report.xlsx'
    plan = json.loads((QA/'sheet-edits.json').read_text(encoding='utf8'))
    baseline = json.loads((QA/'baseline.json').read_text(encoding='utf8'))
    old = openpyxl.load_workbook(QA/'report-before.xlsx',rich_text=True,data_only=False)
    new = openpyxl.load_workbook(path,rich_text=True,data_only=False)
    assert new.sheetnames == old.sheetnames
    assert new.worksheets[1].max_row == old.worksheets[1].max_row == 1734
    assert [str(x) for x in new.worksheets[1].merged_cells.ranges] == [str(x) for x in old.worksheets[1].merged_cells.ranges]
    differences=[]
    for a,b in zip(old.worksheets,new.worksheets):
        for row in a:
            for c in row:
                other = b[c.coordinate]
                if str(c.value) != str(other.value):
                    differences.append((a.title,c.coordinate))
                    assert b.title == new.worksheets[1].title and c.coordinate in plan['edits'], differences[-1]
                assert c.style_id == other.style_id, c.coordinate
    ws=new.worksheets[1]
    for cell,want in plan['edits'].items():
        assert str(ws[cell].value) == want, cell
        assert want.startswith('\n') and want.endswith('\n'),cell
        assert all(x not in want for x in ('결과 미기록','왼쪽의 출력','출력: 출력:','[E8300 확인]','[E8300 반영]','·','。','接続','記録')),cell
    for unit,first,last,n in [(1,789,792,8),(5,809,811,8),(46,1208,1213,10)]:
        content='\n'.join(str(ws[f'K{r}'].value or '') for r in range(first,last+1))
        labels=[int(x) for x in re.findall(r'^조건 \[(\d+)\](?:\s|$)',content,re.M)]
        assert labels == list(range(1,n+1)),(unit,labels)
        assert '예상값' in '\n'.join(str(ws[f'L{r}'].value or '') for r in range(first,last+1)),unit
    # Never turn known actual error returns into the planned -1 result.
    assert '저장 실패 입력 시 반환=0이다. 기준 -1과 불일치한다.' in str(ws['L792'].value)
    assert '닫기 실패 입력 시 반환=0이다. 기준 -1과 불일치한다.' in str(ws['L1213'].value)
    for cell in ['L792','L811','L1212','L1213']:
        assert '후속 시험의 예상값' in str(ws[cell].value),cell
    for cell in ('L792','L811','L1212','L1213','M789','M809','M1208','I809'):
        value=ws[cell].value
        marked=[x for x in value if hasattr(x,'font') and x.font.b and x.font.color and x.font.color.rgb=='FF0070C0']
        assert marked,cell
        assert all(float(x.font.sz)==float(ws[cell].font.sz) for x in marked),cell
    before_parts=unpack(QA/'report-before.xlsx');after_parts=unpack(path)
    assert before_parts.keys() == after_parts.keys()
    changed=[k for k in before_parts if before_parts[k] != after_parts[k]]
    assert changed == ['xl/worksheets/sheet2.xml'],changed
    before_xml=E.fromstring(before_parts['xl/worksheets/sheet2.xml'])
    after_xml=E.fromstring(after_parts['xl/worksheets/sheet2.xml'])
    for addr in plan['edits']:
        b=before_xml.find(f'.//s:c[@r="{addr}"]',NS);a=after_xml.find(f'.//s:c[@r="{addr}"]',NS)
        a.getparent().replace(a,deepcopy(b))
    for row in plan['heights']:
        b=before_xml.find(f'.//s:row[@r="{row}"]',NS);a=after_xml.find(f'.//s:row[@r="{row}"]',NS)
        a.attrib.clear();a.attrib.update(b.attrib)
    assert E.tostring(before_xml,method='c14n') == E.tostring(after_xml,method='c14n'),'Unexpected XML change'
    assert all(float(r.height or 0)<=409.5 for r in ws.row_dimensions.values())
    concurrent=[]
    for name,want in {**baseline['protected'],**baseline['sources']}.items():
        from pathlib import Path
        actual=sha(Path(name))
        if actual != want and name in CONCURRENT_CHANGES:
            assert actual==CONCURRENT_CHANGES[name],'Concurrent file changed again: '+name
            concurrent.append({'path':name,'before':want,'now':actual,'origin':'separate SRS task'})
        else:
            assert actual==want,'Outside scope change: '+name
    result={'changed_cells':differences,'changed_parts':changed,'rows':ws.max_row,'test_ids':sum(str(ws.cell(r,4).value or '').strip().startswith('SW-FW-UnitTest-') for r in range(1,ws.max_row+1)),
            'condition_counts':{'CM3-001':8,'CM3-005':8,'CM3-046':10},'protected_files':len(baseline['protected']),'sources':len(baseline['sources']),
            'known_error_returns_preserved':True,'new_outputs_labelled_expected':True,'concurrent_external_changes':concurrent,'sha256':sha(path)}
    (QA/('published-checks.json' if path==REPORT else 'checks.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result,ensure_ascii=False))

if __name__ == '__main__':
    import sys
    if '--publish' in sys.argv:
        baseline=json.loads((QA/'baseline.json').read_text(encoding='utf8'))
        assert sha(REPORT)==baseline['sha256'],'User changed report after baseline'
        check()
        import shutil
        shutil.copyfile(QA/'report.xlsx',REPORT)
        check(REPORT)
    else:
        preserve();check()
