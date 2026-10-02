"""Read-only cross-document/character-mask/reference checks of QA candidates."""
from __future__ import annotations
import argparse
from difflib import SequenceMatcher
import hashlib
import json
from pathlib import Path
import re
import sys
import openpyxl
from lxml import etree as E
import apply_validation_review as package

Q=package.Q

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--live',action='store_true',help='also check published bytes and protected files')
    args=parser.parse_args()
    plan=json.loads((Q/'plan.json').read_text('utf-8'))
    books={b['key']:b for b in plan['books']}
    if args.live:
        for b in plan['books']:
            assert package.sha(Path(b['path']))==package.sha(Q/b['key']/'candidate.xlsx'),b['key']
        for p,digest in plan['protected'].items():
            assert package.sha(Path(p))==digest,('Protected document changed',p)
    def load(key,which='candidate'):
        return openpyxl.load_workbook(Q/key/(which+'.xlsx'),data_only=True,rich_text=True)
    srs=load('srs');input_=load('input-system');output=load('output-system');unit=load('unit');integration=load('integration')
    def ids(wb,col,prefix):
        return [str(r[col-1].value) for r in wb.worksheets[1] if str(r[col-1].value).startswith(prefix)]
    def merged_value(sheet,row,col):
        for bounds in sheet.merged_cells.ranges:
            if bounds.min_row<=row<=bounds.max_row and bounds.min_col<=col<=bounds.max_col:
                return sheet.cell(bounds.min_row,bounds.min_col).value
        return sheet.cell(row,col).value
    expected_sw={f'SW-FW-RID-{n}' for n in range(1,82)}
    srs_ids=ids(srs,7,'SW-FW-RID-')
    assert len(srs_ids)==81 and set(srs_ids)==expected_sw
    # Read through merged hierarchy cells; the requirement row owns the SW ID.
    sys_trace={}
    for row in srs.worksheets[1]:
        sw=str(row[6].value or '')
        if sw in expected_sw:
            sys_trace[sw]=set(re.findall(r'SYS-RID-FW-\d+',str(merged_value(srs.worksheets[1],row[0].row,8) or '')))
    assert set().union(*sys_trace.values())=={f'SYS-RID-FW-{n}' for n in range(1,51)}
    assert len(set(ids(unit,4,'SW-FW-UnitTest-')))==115
    assert len(set(ids(integration,3,'SW-FW-IntegrationTest-ID-')))==17
    for key,wb in [('input-system',input_),('output-system',output)]:
        system_ids=ids(wb,2,'SW-FW-SYS-Test-ID-')
        assert len(system_ids)==81 and set(system_ids)=={f'SW-FW-SYS-Test-ID-{n}' for n in range(1,82)}
        tracked=[]
        for row in wb.worksheets[1]:
            identifier=str(row[1].value or '')
            if identifier.startswith('SW-FW-SYS-Test-ID-'):
                sws=set(re.findall(r'SW-FW-RID-\d+',str(row[4].value or '')))
                assert sws and sws<=expected_sw,(key,identifier,'SW trace')
                uppers=set(re.findall(r'SYS-RID-FW-\d+',str(merged_value(wb.worksheets[1],row[0].row,6) or '')))
                assert uppers==set().union(*(sys_trace[sw] for sw in sws)),(key,identifier,'SYS trace')
                tracked.extend(sws)
        assert set(tracked)==expected_sw,(key,'Uncovered software requirement')
        for identifier,target in books[key]['targets'].items():
            mate=books['output-system' if key=='input-system' else 'input-system']['targets'][identifier]
            assert target['changes']==mate['changes']
            for col,text in target['changes'].items():assert str(wb.worksheets[1][col+str(target['row'])].value)==text
    sys.path.insert(0,str(package.ROOT/'tests/unit'))
    import tdc_ut_hear_reference as hear
    _,reference=hear.build();h=unit['HEAR 기대값']
    for n,value in enumerate(hear.source_window()):assert h.cell(n+6,2).value==value
    for i in range(256):
        for v in range(5):assert h.cell(i+6,v+6).value==reference['vectors'][v]['vmag_unscaled'][i]
    for v in range(5):
        for i,k in enumerate(hear.FFT_BINS):
            r=267+9*v+i
            assert [h.cell(r,c).value for c in range(5,9)]==[v+1,k,*reference['vectors'][v]['fft_unscaled_packed_re_im'][i*2:i*2+2]]
    masks=0
    for key,b in books.items():
        raw=package.parts(Q/key/'candidate.xlsx');ss=package.strings(raw)
        tree=E.fromstring(raw[package.sheets(raw)[b['comparison']]])
        groups={}
        for row in tree.find(package.tag('sheetData')):
            cells={re.sub(r'\d','',c.get('r')):c for c in row}
            identifier=package.value(cells.get(openpyxl.utils.get_column_letter(b['idcol'])),ss)
            mark=package.value(cells.get('A'),ss)
            if identifier in b['targets'] and mark in ('[-]','[+]'):groups.setdefault(identifier,{})[mark]=cells
        assert set(groups)==set(b['targets'])
        for identifier,pair in groups.items():
            assert set(pair)=={'[-]','[+]'},(key,identifier)
            for mark,cells in pair.items():
                other=pair['[+]' if mark=='[-]' else '[-]']
                for col,c in cells.items():
                    if c.find(package.tag('f')) is not None:continue
                    text=package.value(c,ss);mate=package.value(other.get(col),ss)
                    expected=[False]*len(text)
                    for block in SequenceMatcher(None,text,mate,autojunk=False).get_opcodes():
                        if block[0]!='equal':expected[block[1]:block[2]]=[True]*(block[2]-block[1])
                    if col=='A':expected=[True]*len(text)
                    offset=0
                    for run in c.xpath('./s:is/s:r',namespaces=package.NS):
                        value=''.join(run.xpath('./s:t/text()',namespaces=package.NS));end=offset+len(value)
                        active=run.find(package.tag('rPr')).find(package.tag('b')).get('val')=='1'
                        color=run.find(package.tag('rPr')).find(package.tag('color')).get('rgb')
                        wanted=('FF9C0006' if mark=='[-]' else 'FF006100') if active else 'FF666666'
                        assert expected[offset:end]==[active]*len(value) and color==wanted,(key,identifier,col)
                        offset=end;masks+=1
                    if text:assert offset==len(text)
        if key=='unit':
            before=package.parts(Q/key/'before.xlsx')
            old=E.fromstring(before['xl/styles.xml']);new=E.fromstring(raw['xl/styles.xml'])
            for section in ('fonts','fills','borders','cellXfs','cellStyleXfs','cellStyles'):
                oldnodes=old.find(package.tag(section));newnodes=new.find(package.tag(section))
                assert all(E.tostring(a)==E.tostring(newnodes[i]) for i,a in enumerate(oldnodes)),section
    result={'current_srs_ids':81,'system_ids_each':81,'unit_ids':115,'integration_ids':17,
            'input_output_common_criteria_equal':True,'sys_requirements_traced':50,
            'sw_requirements_traced_each':81,'hear_window_coefficients':512,
            'hear_vmag_references':1280,'hear_fft_components':90,'rich_text_runs_checked':masks,
            'old_styles_preserved':True,'device_execution':False,'published_and_protected_checked':args.live}
    package.save_json(Q/'consistency-checks.json',result);print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':main()
