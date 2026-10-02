"""검토 의견의 좁은 셀 변경을 준비/검증/게시한다.

본문 값은 Artifact Tool로 작성한다. OOXML은 기존 서식/병합/도형과
셀 부분 비교 서식 보존에만 사용한다. openpyxl은 읽기 검증 전용이다.
"""
from __future__ import annotations
from copy import deepcopy
from difflib import SequenceMatcher
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import sys
import zipfile
import openpyxl
from lxml import etree as E

T = Path(__file__).resolve().parent
ROOT = T.parents[3]
D = ROOT / 'docs/참고/인허가/사운드1/DHF'
Q = T / 'qa/validation-review'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
def tag(t): return '{' + NS['s'] + '}' + t
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def save_json(p, v): p.write_text(json.dumps(v, ensure_ascii=False, indent=2), encoding='utf-8')
def parts(p):
    with zipfile.ZipFile(p) as z: return {n: z.read(n) for n in z.namelist()}
def sheets(raw):
    links = {x.get('Id'): x.get('Target') for x in E.fromstring(raw['xl/_rels/workbook.xml.rels'])}
    return {x.get('name'): (links[x.get('{'+RNS+'}id')].lstrip('/') if links[x.get('{'+RNS+'}id')].startswith('/') else 'xl/'+links[x.get('{'+RNS+'}id')])
            for x in E.fromstring(raw['xl/workbook.xml']).find(tag('sheets'))}
def strings(raw):
    return [''.join(x.itertext()) for x in E.fromstring(raw['xl/sharedStrings.xml'])] if 'xl/sharedStrings.xml' in raw else []
def value(c, ss):
    if c is None: return ''
    if c.get('t') == 's': return ss[int(c.findtext(tag('v')))]
    if c.get('t') == 'inlineStr': return ''.join(c.xpath('./s:is//s:t/text()', namespaces=NS))
    return c.findtext(tag('v')) or ''

SYSTEM = {
    12: {'I': '시험용 타이머로 연속 오디오 블록 처리 시작 시각을 기록한다. 인접한 두 시작 시각의 차이를 처리 주기로 비교하고, 결과는 처리 후 RTT로 읽는다.',
         'J': '처리 주기는 1.00ms, 허용 오차 +/-5%로 0.95~1.05ms이다. 입력부터 출력까지의 지연 시간과 구분한다.'},
    58: {'I': '정상 실시간 자극 상태에서 연속 블록의 처리 시작/완료 시각을 저장한다. 반복 주기와 다음 블록 시작 전 처리 완료 여부를 확인하고, 결과는 처리 후 RTT로 읽는다.',
         'J': '실시간 처리 주기는 1.00ms, 허용 오차 +/-5%로 0.95~1.05ms이며 각 블록은 다음 처리 시작 전에 완료되어야 한다. 개별 연산 시간이 0.95ms보다 짧은 것은 실패가 아니다.'},
    56: {'H': '오디오 볼륨은 고정하고 자극 볼륨을 낮은 단계로 설정한다. 정상 시험 내부기를 사용한다.',
         'I': '자극 볼륨을 높은 단계로 변경하고 실제 자극 크기의 변화와 최종 상태를 저장한 디버그 기록으로 확인한다.',
         'J': '자극 크기는 설정한 최종 단계까지 단계적으로 증가해야 한다. 요청 순간에 최종 크기로 뛰지 않아야 하며 오디오 볼륨은 바뀌지 않아야 한다.'},
    50: {'I': '리모콘 연결/해제 이벤트를 발생시킨 뒤 생성된 감사기록에서 해당 이벤트와 연결 정보를 확인한다.',
         'J': '발생시킨 연결/해제 이벤트와 대상 연결 정보가 감사기록으로 남아야 한다. 현재 연결 상태 조회만으로 기록 생성을 통과시키지 않는다.'},
    75: {'I': '인증/연결/저장 오류를 각각 발생시켜 사용자에게 전달되는 오류 내용을 확인한다.',
         'J': '오류 내용에 인증 키, 암호 또는 보호 대상의 원문 정보가 직접 표시되지 않아야 한다. 오류 메시지가 존재한다는 이유만으로 실패로 판정하지 않는다.'},
    79: {'I': '인가된 연결과 비인가 연결에서 감사기록 읽기를 각각 요청한다. 허용된 전송의 보호 여부와 원격 수정 거부를 확인한다.',
         'J': '인가된 연결에서만 감사기록을 읽을 수 있어야 하고 비인가 요청은 거부되어야 한다. 외부 전송에는 기밀성 보호가 적용되어야 하며 원격 편집 기능이 없어야 한다.'},
}
UNIT = {
    'CM3-015': {'H': '저전력 진입 요청 플래그(enter_ULP_mode_Command_CM3_to_CFX)를 1(요청/정리 대기)과 0(CFX 정리 후 요청 해제)로 나누어 절전 진입/복귀 분기를 검토한다.',
                'K': '요청 플래그가 1인 동안 CM3는 CFX 정리를 기다린다. CFX가 주변장치를 정리한 뒤 0으로 해제하면 저클럭 전환/절전 진입을 진행한다. 복귀는 재기동/초기화를 거쳐 정상 처리를 준비한다.'},
    'HEAR-001': {'J': '중간 vAbsolute 배열은 함수체인 순서와 절대값 수식으로 분석한다. 실제 CMD 0/Event 0 완료 후 배열과 최대값을 별도 저장하여 RTT로 비교한다. 완료 후 배열을 중간 시점의 실측으로 쓰지 않는다.'},
    'CFX-040': {'H': '24프레임 버퍼의 앞 3칸을 [0x11,0x22,0x33], 나머지 21칸을 NOP(0x25555)로 준비한다. 별도로 i=0~23에 0x50080+513*i를 넣어 모든 단어를 구분한다. 초기 읽기 인덱스=17, 읽기 중 플래그=0, 요청 모드=SpecificCommand, 다음 모드=NopStandby다.',
                'K': '두 조건 모두 24프레임이 입력 순서대로 전달된다. 쓰기 순간 읽기 중 플래그=1, 인덱스=[17,0,1,...,22]이며 완료 후 인덱스=23/플래그=0/모드=NopStandby다. 입력 버퍼는 바뀌지 않는다. 유효 데이터 3개와 전송 길이 24를 구분하며 송신 중 버퍼를 호출자가 변경하지 않는다.'},
    'CFX-036': {'H': '4대역 CIS 순서=[1,2,3,4], 배정 전극=[1,2,17,32], 물리 전극 표 map[i]=i, Amplitude=[0,1,127,255], T 미만 묵음 꺼짐으로 설정한다. 첫 펄스 위상=0/1을 각각 실행한다.',
                'K': '위상 0의 헤더=0x40000, 패킷=[0x40000,0x40404,0x441FC,0x47FFC], 위상 1의 헤더=0x48000, 패킷=[0x48000,0x48404,0x4C1FC,0x4FFFC]다. 완성 단어의 비트 19~16=4, 15=위상, 14~10=물리 전극 인덱스, 9~2=Amplitude, 1~0=0을 각각 비교한다. 고정 헤더 비트를 가변 입력으로 가정하지 않는다.'},
    'CFX-037': {'H': '교차 대역 순서=[0,8,1,9,2,10,3,11,4,12,5,13,6,14,7,15], 배정 전극=32-i, 물리 전극 표 map[j]=(7*j)%32, 대역 i의 Amplitude=[0,1,127,255][i%4]로 설정한다. 묵음 꺼짐, 프레임 간격=3, 블록당 자극=8, 전송 인덱스=0. (위상,직전 마지막)=(0,15)/(1,0)을 각각 실행한다.',
                'K': '직전 마지막=15는 원래 순서/새 마지막=15, 직전 마지막=0은 [15,0,8,1,9,2,10,3,11,4,12,5,13,6,14,7]/새 마지막=7이다. 각 대역 i의 패킷 W=0x40000|(위상<<15)|(map[31-i]<<10)|(Amplitude[i]<<2)다. 두 블록은 앞/뒤 8개를 각각 3프레임 간격으로 내보내고 사이 2프레임은 NOP(0x25555)다. 전송 인덱스는 8 다음 0이며 전극/크기/위상/순서를 개별 비교한다.'},
    'HEAR-002': {'H': '512점 Hanning 창, n=0~511. 표본 배열 P=[0,25080,46341,60547,65536,60547,46341,25080,0,-25080,-46341,-60547,-65536,-60547,-46341,-25080]. 입력은 (1)0, (2)P[n%16], (3)P[(2*n)%16], (4)32768, (5)짝수 n은32768/홀수 n은-32768이다. 매 조건 FIFO 전부를 다시 채운다.',
                 'J': '실제 CMD1/Event1 완료 후 FFT 512단어, vMag 256단어, BFP 블록 지수 E를 보존한다. HEAR 기대값 시트의 사전 DFT 값과 비교하며 FFT 실수/허수는 bin [0,1,31,32,33,63,64,65,255], vMag는 256개 전부를 판정한다. 원본 배열도 RTT로 남긴다.',
                 'K': '입력 1은 FFT/vMag 전부 정확히 0이다. 나머지는 FFT 원시값*2^E, vMag 원시값*2^(E+1)을 사전 DFT 기준과 비교한다(E=0~13). FFT 성분 허용 오차=64+ceil(|기준|/1000), vMag=64+ceil(기준/100) 정수 단위다. 첫 FFT 단어/둘째 단어는 DC/Nyquist가 들어가며 vMag[0]은 이 두 값의 제곱합 제곱근이다. 입력 2/3의 최대 bin은 32/64, 4/5는 0이다. 완료 이벤트는 1이다. 이 오차는 사전 설정한 기술 검토 기준이며 SDK 제조사 보장값이 아니다. 실기 확인 전 전체 합격으로 확정하지 않는다.'},
    'CM3-050': {'H': 'OTA 시작 [0xC0,0,7]/종료 [0xC1,9] 명령을 각각 입력하고 기존 DFU 플래그는 각각 0/1로 둔다. 별도로 QCC 시스템 정보 [0x35,4](OTA_EZAIRO)/[0x35,0](NONE)을 입력한다. 입력 상태 4와 연결 플래그 1을 구분한다.',
                'K': '시작 ACK=[0xC0,0,7], 종료 ACK=[0xC1,0,9], 길이=3이며 기존 DFU 플래그 0/1은 각각 유지된다. QCC [0x35,4]/[0x35,0]의 ACK는 [0x35,1], 길이=2이며 연결 조회 플래그는 각각 1/0이다. 시작/종료 응답만으로 연결 상태가 바뀐다고 가정하지 않는다.'},
}
INTEGRATION = {
    4: {'J': '24프레임 버퍼의 앞 3칸에 [0x11,0x22,0x33], 나머지 21칸에 NOP(0x25555)를 넣는다. 대기/프리앰블/지정 명령/실시간/0 채움 모드를 각각 요청한다.',
        'M': '각 모드의 처리 경로만 실행된다. 지정 명령은 고정 24프레임이며 앞 3칸과 나머지 NOP가 입력 순서와 같다. 송신 처리 후 버퍼 요청을 해제한다. 프리앰블 후 다음 모드로 전환하고 대기는 NOP, 0 채움은 0이다.'},
    17: {'L': '갱신 시작/시험 파일 수신 및 기록/종료/기동 위치 요청의 응답을 확인한다. QCC [0x35,4](OTA_EZAIRO)/[0x35,0](NONE)을 별도로 입력하여 DFU 상태를 확인한다. 격리된 시험 저장소의 재기동 선택과 저장 실패/지원 밖 슬롯, TRY/TRY_DONE 조건도 각각 적용한다.',
         'M': '정상 기록의 파일 종류/슬롯/내용과 응답은 입력과 같다. DFU 상태는 시작/종료 ACK가 아니라 QCC 0x35 입력 4/0에 따라 연결 플래그 1/0으로 바뀐다. 저장 실패/지원 밖 요청은 오류다. TRY 2는 3/슬롯2, TRY 3은 4/슬롯1/BOOT/TRY_DONE/FAIL, TRY_DONE 성공은 슬롯2, 실패는 슬롯1/UNKNOWN이다. 불완전 이미지를 정상 기동으로 판정하지 않는다.'},
}

def prepare():
    Q.mkdir(parents=True, exist_ok=True)
    specs = [('srs', '02_설계 입력', 'SRS', 7), ('input-system', '02_설계 입력', 'SSTC', 2),
             ('output-system', '03_설계 출력', 'SSTC', 2), ('unit', '03_설계 출력', 'SUVC', 7),
             ('integration', '03_설계 출력', 'SITC', 3)]
    plan = {'books': [], 'protected': {str(p): sha(p) for p in D.rglob('*') if p.is_file() and not p.name.startswith('~$') and ('[작성자 전용]' in p.name or 'EPTC-' in p.name)}}
    for key, directory, prefix, idcol in specs:
        matches = [p for p in (D/directory/'SW/펌웨어').glob('*'+prefix+'*.xlsx') if '[작성자 전용]' not in p.name and not p.name.startswith('~$')]
        assert len(matches) == 1, matches
        p = matches[0]; folder = Q/key; folder.mkdir(exist_ok=True)
        before = folder/'before.xlsx'
        if before.exists(): assert sha(before) == sha(p), 'Baseline externally changed: '+str(p)
        else: shutil.copy2(p, before)
        wb = openpyxl.load_workbook(p, data_only=True); s = wb.worksheets[1]
        targets = {}
        for row in s:
            identifier = str(row[idcol-1].value or '')
            changes = {}
            if key == 'srs':
                if identifier == 'SW-FW-RID-12': changes = {'F': '사운드 처리 주기는 1.00ms이며 허용 오차는 +/-5%(0.95~1.05ms)이다.'}
                if identifier == 'SW-FW-RID-58': changes = {'F': '실시간 신호처리 주기는 1.00ms, 허용 오차 +/-5%(0.95~1.05ms)이며 각 블록 처리는 다음 처리 시작 전에 완료되어야 한다.'}
            elif key.endswith('system') and identifier.startswith('SW-FW-SYS-Test-ID-'):
                changes = SYSTEM.get(int(identifier.rsplit('-',1)[1]), {})
            elif key == 'unit' and identifier.startswith('SU-ID-'):
                changes = UNIT.get(identifier[6:], {})
            elif key == 'integration' and identifier.startswith('SW-FW-IntegrationTest-ID-'):
                changes = INTEGRATION.get(int(identifier.rsplit('-',1)[1]), {})
            if changes:
                targets[identifier] = {'row': row[0].row, 'changes': changes}
        expected = 2 if key == 'srs' else len(SYSTEM) if key.endswith('system') else len(UNIT) if key == 'unit' else len(INTEGRATION)
        assert len(targets) == expected, (key, targets)
        if key == 'unit':
            # Keep the original midpoint criteria and add independently prepared
            # SDK-path cases. The workbook itself contains every numeric bound.
            sys.path.insert(0, str(ROOT/'tests/unit'))
            import cfx_numeric_reference as nr
            references = nr.references()
            def interval(pair):
                return str(pair[0]) if pair[0] == pair[1] else f'{pair[0]}~{pair[1]}'
            additions = {
                'SU-ID-CFX-023': {
                    'H': '\n- 별도 SDK 경로: 입력=[1,8,128,4096,65536,262143], 볼륨 1/10. 이전 이득=0, 어택=62259/상보=3276, 릴리즈=655/상보=64880(Q8.16)으로 매 조건을 초기화한다. 로그/거듭제곱 결과 고정은 해제한다.',
                    'J': '\n- SDK 경로의 목표 이득/평활 이득/선형 이득과 사전 하한/상한을 RTT로 보존한다. 실제값과 허용 구간을 먼저 읽고, 구간 판정 결과를 별도로 비교한다.',
                    'K': '\n- SDK 경로 허용 구간. 각 줄은 볼륨/입력: 목표(Q8.16), 평활(Q8.16), 선형(Q12.12) 순서다.\n' + '\n'.join(f"{r['volume']+1}/{r['input']}: "+', '.join(interval(v) for v in r['bounds']) for r in references['agc']) + '\n- SDK Log2/Exp2 정밀도, 변환 상수 오차 및 정수 절삭을 포함해 실행 전에 산출한 구간이다. 관측값에 맞추어 구간을 늘리지 않는다.'
                },
                'SU-ID-CFX-031': {
                    'H': '\n- 별도 SDK 경로: 채널 i=0~31에서 T=50+(i%4), C=150+(i%4), xMin=16, xMax=4096, p=1258291(Q1.23), 정규화 계수=5163793(Q8.16). 중간 거듭제곱 입력 고정을 해제한다.',
                    'J': '\n- 실제 SDK 거듭제곱/나눗셈으로 계산한 32채널 A/B와 각 하한/상한을 보존한다.',
                    'K': '\n- SDK 경로의 A 허용 구간은 전 채널 13016~13019(Q16.8). i%4=0/1/2/3일 때 B 허용 구간은 각각 -6935~-6923 / -6679~-6667 / -6423~-6411 / -6167~-6155(Q16.8)다. 각 채널을 개별 비교하며 한 채널의 결과를 복제하지 않는다.'
                },
                'SU-ID-CFX-032': {
                    'H': '\n- 별도 SDK 경로: 채널 i=0~31에 [0,1,8,64,512,4096,16384,262143][i%8]을 넣는다. xMin=16/xMax=4096, p=1258291(Q1.23), 정규화 계수=5163793(Q8.16), T=50/C=150으로 실제 계수를 계산한다. 묵음 모드 1/2를 각각 적용한다.',
                    'J': '\n- 실제 SDK 경로의 PCM 크기와 알림 처리 전 공유값을 32채널 각각 비교한다. 원본 계산값/하한/상한도 별도 RTT로 보존한다.',
                    'K': '\n- 위 SDK 입력 순서에서 모드 1은 [0,0,0,67,102,149~150,150,150], 모드 2는 [50,50,50,67,102,149~150,150,150]다. 8개 입력을 4번 반복하며 PCM/공유값이 각각 이 구간 안에 있어야 한다. 4096의 한 코드 여유는 SDK 근사/정수 절삭을 포함한 사전 계산이다.'
                },
                'SU-ID-CM3-028': {
                    'H': '\n- 전하량 경계는 T=100uA/C=300uA, 펄스폭 249/250/251us로 나눈다. 별도 DAC 범위 초과는 T=100uA/C=2500uA, 펄스폭 25us로 넣는다.',
                    'K': '\n- 전하량 Q[pC]=전류[uA]*펄스폭[us], 한계=75000pC(75nC). 74700pC(249us)와 75000pC(250us)는 성공, 75300pC(251us)는 실패다. T/C 디지털 값은 정상 조건에서 0/127이다. DAC 범위 초과 조건은 전하량 62500pC로 한계 이하이나 전류 상한 2198uA를 넘으므로 실패다. 전하량 초과와 DAC 범위 초과를 구분한다.'
                },
                'SU-ID-CM3-029': {
                    'H': '\n- CIS의 펄스폭을 249/250/251us로 나누어 전하량 경계를 시험한다. 별도 DAC 실패는 T=100uA/C=2500uA, 펄스폭 25us다. 실패 상태 관찰에서는 기존 완료 플래그를 0/1로 각각 둔다.',
                    'J': '\n- 계산 반환값과 공유 갱신을 구분하여 저장한다. DAC 범위 실패의 복사 직전 시험 경계에서는 초기화되지 않은 T/C 지역 배열을 읽지 않고 복사 시도/프레임/완료 상태만 OBSERVED로 남긴다. 이 관찰에서 시험 경계가 보존한 값은 제품의 실패 후 보존 증거가 아니다.',
                    'K': '\n- CIS 249us/250us는 각각 성공, 자극당 13프레임/1ms당 1채널이다. 251us(75300pC)는 실패다. 계산 함수 자체는 완료 플래그를 설정하지 않는다. 실패 후 공유값 유지/초기화 정책은 상세 설계에서 결정한 뒤 별도 판정하며 현재 관찰을 합격으로 처리하지 않는다.'
                },
            }
            for identifier, changes in additions.items():
                match = [r for r in s if r[idcol-1].value == identifier]
                assert len(match) == 1
                row = match[0]
                targets[identifier] = {'row': row[0].row, 'changes': {
                    col: str(s[col+str(row[0].row)].value or '') + text
                    for col, text in changes.items()}}
        plan['books'].append({'key': key, 'path': str(p), 'sha256': sha(p), 'sheet': s.title, 'comparison': wb.worksheets[2].title,
                              'idcol': idcol, 'targets': targets})
    save_json(Q/'plan.json', plan)
    print('Prepared 5 baseline-preserving workbooks')

def add_hear_sheet(raw, authored):
    """Attach the Artifact-authored independent reference sheet, not results.

    It has no drawings, formulas, named styles or external relationships.
    Merge its low-level cell formats without replacing any existing style.
    """
    name='HEAR 기대값'; source=E.fromstring(authored[sheets(authored)[name]])
    ass=strings(authored)
    for c in source.iter(tag('c')):
        if c.get('t')=='s':write_value(c,value(c,ass))
    styles=E.fromstring(raw['xl/styles.xml']); ast=E.fromstring(authored['xl/styles.xml'])
    offsets={}
    for section in ('fonts','fills','borders'):
        dest=styles.find(tag(section)); extra=ast.find(tag(section)); offsets[section]=len(dest)
        dest.extend(deepcopy(list(extra)));dest.set('count',str(len(dest)))
    fmts=styles.find(tag('numFmts')); afmts=ast.find(tag('numFmts')); fmtmap={}
    if afmts is not None and len(afmts):
        if fmts is None:fmts=E.Element(tag('numFmts'));styles.insert(0,fmts)
        nextid=max([163]+[int(n.get('numFmtId')) for n in fmts])+1
        for fmt in afmts:
            new=deepcopy(fmt);fmtmap[int(fmt.get('numFmtId'))]=nextid
            new.set('numFmtId',str(nextid));fmts.append(new);nextid+=1
        fmts.set('count',str(len(fmts)))
    xfs=styles.find(tag('cellXfs')); offset=len(xfs)
    for xf in ast.find(tag('cellXfs')):
        new=deepcopy(xf)
        for field,section in (('fontId','fonts'),('fillId','fills'),('borderId','borders')):
            if new.get(field):new.set(field,str(int(new.get(field))+offsets[section]))
        if new.get('numFmtId') and int(new.get('numFmtId')) in fmtmap:new.set('numFmtId',str(fmtmap[int(new.get('numFmtId'))]))
        new.set('xfId','0');xfs.append(new)
    xfs.set('count',str(len(xfs)))
    for node in source.iter():
        field='s' if node.tag==tag('c') else 'style' if node.tag in (tag('row'),tag('col')) else None
        if field and node.get(field):node.set(field,str(int(node.get(field))+offset))
    wb=E.fromstring(raw['xl/workbook.xml']);sh=wb.find(tag('sheets'))
    rel=E.fromstring(raw['xl/_rels/workbook.xml.rels']);reluri='http://schemas.openxmlformats.org/package/2006/relationships'
    sheetid=max(int(n.get('sheetId')) for n in sh)+1
    rid='rId'+str(max([0]+[int(n.get('Id')[3:]) for n in rel if n.get('Id','').startswith('rId')])+1)
    path=f'xl/worksheets/validation_hear_{sheetid}.xml'
    E.SubElement(sh,tag('sheet'),name=name,sheetId=str(sheetid),attrib={'{'+RNS+'}id':rid})
    E.SubElement(rel,'{'+reluri+'}Relationship',Id=rid,Type=RNS+'/worksheet',Target=path[3:])
    ct=E.fromstring(raw['[Content_Types].xml']);cturi='http://schemas.openxmlformats.org/package/2006/content-types'
    E.SubElement(ct,'{'+cturi+'}Override',PartName='/'+path,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
    for path_,tree in [('xl/styles.xml',styles),('xl/workbook.xml',wb),('xl/_rels/workbook.xml.rels',rel),('[Content_Types].xml',ct),(path,source)]:
        raw[path_]=E.tostring(tree,xml_declaration=True,encoding='utf-8',standalone=True)
    return {'xl/styles.xml','xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml',path}

def write_value(cell, text, runs=None):
    for child in list(cell):
        if child.tag in (tag('f'), tag('v'), tag('is')): cell.remove(child)
    cell.set('t','inlineStr'); inline = E.SubElement(cell,tag('is'))
    if runs is None:
        t=E.SubElement(inline,tag('t')); t.text=text; t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    else:
        for txt, active, mark in runs:
            r=E.SubElement(inline,tag('r')); rp=E.SubElement(r,tag('rPr'))
            E.SubElement(rp,tag('rFont'),val='맑은 고딕'); E.SubElement(rp,tag('sz'),val='11')
            E.SubElement(rp,tag('b'),val='1' if active else '0')
            E.SubElement(rp,tag('color'),rgb=('FF9C0006' if mark=='[-]' else 'FF006100') if active else 'FF666666')
            t=E.SubElement(r,tag('t')); t.text=txt; t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')

def move_row(row, number):
    row.set('r',str(number))
    for c in row: c.set('r',re.sub(r'\d+$',str(number),c.get('r')))
def fit(row, ss, changed):
    # Only increase an affected row. Existing widths/point sizes remain intact.
    lines=max((sum(max(1,math.ceil(len(line)/32)) for line in value(c,ss).split('\n')) for c in row if re.sub(r'\d','',c.get('r')) in changed), default=1)
    row.set('ht',str(min(409,max(float(row.get('ht','18')),lines*16+12))))
    row.set('customHeight','1')

def finish():
    plan=json.loads((Q/'plan.json').read_text('utf-8')); authored=parts(Q/'authored.xlsx'); ass=strings(authored)
    ar=E.fromstring(authored[sheets(authored)['수정 내용']])
    approved={value(row[0],ass):value(row[1],ass) for row in ar.find(tag('sheetData')) if len(row)>=2}
    results=[]
    for b in plan['books']:
        raw=parts(Q/b['key']/'before.xlsx'); original=dict(raw); ss=strings(raw); links=sheets(raw)
        main=E.fromstring(raw[links[b['sheet']]]); data=main.find(tag('sheetData')); rows={int(r.get('r')):r for r in data}
        mainallowed=set()
        for identifier,target in b['targets'].items():
            row=rows[target['row']]; cells={re.sub(r'\d','',c.get('r')):c for c in row}
            for col,txt in target['changes'].items():
                assert approved[b['key']+'|'+identifier+'|'+col]==txt
                if col not in cells:
                    c=E.Element(tag('c'),r=col+str(target['row'])); c.set('s',cells.get('J',next(iter(cells.values()))).get('s','0')); row.append(c); cells[col]=c
                    row[:]=sorted(row,key=lambda x:openpyxl.utils.column_index_from_string(re.sub(r'\d','',x.get('r'))))
                write_value(cells[col],txt); mainallowed.add(cells[col].get('r'))
            fit(row,ss,target['changes'])
        diff=E.fromstring(raw[links[b['comparison']]]); dd=diff.find(tag('sheetData')); new=[]; positions={}; pairs=[]
        idletter=openpyxl.utils.get_column_letter(b['idcol'])
        for oldrow in dd:
            row=deepcopy(oldrow); oldnum=int(row.get('r')); cells={re.sub(r'\d','',c.get('r')):c for c in row}; identifier=value(cells.get(idletter),ss)
            mark=value(cells.get('A'),ss); positions[oldnum]=[len(new)+1,len(new)+1]
            if identifier in b['targets'] and mark not in ('[-]','[+]'):
                # A previously unchanged comparison row becomes one before/after pair.
                before=deepcopy(row)
                bc={re.sub(r'\d','',c.get('r')):c for c in before}
                if 'A' not in bc: bc['A']=E.Element(tag('c'),r='A'+str(oldnum)); before.insert(0,bc['A'])
                write_value(bc['A'],'[-]'); new.append(before)
                if 'A' not in cells: cells['A']=E.Element(tag('c'),r='A'+str(oldnum)); row.insert(0,cells['A'])
                write_value(cells['A'],'[+]'); mark='[+]'; positions[oldnum][1]+=1
            if identifier in b['targets'] and mark=='[+]':
                for col,txt in b['targets'][identifier]['changes'].items():
                    if col not in cells:
                        cells[col]=E.Element(tag('c'),r=col+str(oldnum)); cells[col].set('s',cells.get('J',next(iter(cells.values()))).get('s','0')); row.append(cells[col])
                    write_value(cells[col],txt)
                fit(row,ss,b['targets'][identifier]['changes'])
            new.append(row)
        # Row numbers may contain existing gaps; assign based on original row plus inserted count.
        inserted=0; mapping={}; out=[]
        for oldrow in dd:
            n=int(oldrow.get('r')); a,z=positions[n]; group=new[a-1:z]; mapping[n]=(n+inserted,n+inserted+len(group)-1)
            for i,row in enumerate(group): move_row(row,n+inserted+i); out.append(row)
            inserted+=len(group)-1
        dd[:]=out
        merged=diff.find(tag('mergeCells'))
        def shift_ref(ref):
            out=[]
            for i,point in enumerate(ref.split(':')):
                m=re.fullmatch(r'([A-Z]+)(\d+)',point); n=int(m[2]); bounds=mapping.get(n)
                if bounds: n=bounds[1 if i else 0]
                else: n+=sum(1 for r in mapping if r<n and mapping[r][1]!=mapping[r][0])
                out.append(m[1]+str(n))
            return ':'.join(out)
        if merged is not None:
            for m in merged: m.set('ref',shift_ref(m.get('ref')))
        for node in diff.xpath('./s:dimension|./s:autoFilter',namespaces=NS):
            if node.get('ref'):node.set('ref',shift_ref(node.get('ref')))
        # Preserve historical [-] rows. Only affected pairs get new character masks.
        groups={}
        for row in dd:
            cells={re.sub(r'\d','',c.get('r')):c for c in row}; identifier=value(cells.get(idletter),ss); mark=value(cells.get('A'),ss)
            if identifier in b['targets'] and mark in ('[-]','[+]'): groups.setdefault(identifier,{})[mark]=cells
        for identifier,pair in groups.items():
            for mark,cells in pair.items():
                other=pair.get('[+]' if mark=='[-]' else '[-]',{})
                for col,c in cells.items():
                    if c.find(tag('f')) is not None:continue
                    txt=value(c,ss); mate=value(other.get(col),ss); mask=[False]*len(txt)
                    for op,a,z,_,_ in SequenceMatcher(None,txt,mate,autojunk=False).get_opcodes():
                        if op!='equal':mask[a:z]=[True]*(z-a)
                    if col=='A':mask=[True]*len(txt)
                    runs=[]; i=0
                    while i<len(txt):
                        stop=i+1
                        while stop<len(txt) and mask[stop]==mask[i]:stop+=1
                        runs.append((txt[i:stop],mask[i],mark)); i=stop
                    if txt:write_value(c,txt,runs)
        raw[links[b['sheet']]]=E.tostring(main,xml_declaration=True,encoding='utf-8',standalone=True)
        raw[links[b['comparison']]]=E.tostring(diff,xml_declaration=True,encoding='utf-8',standalone=True)
        extra=add_hear_sheet(raw,authored) if b['key']=='unit' else set()
        candidate=Q/b['key']/'candidate.xlsx'
        with zipfile.ZipFile(candidate,'w',zipfile.ZIP_DEFLATED) as z:
            for name,content in raw.items():z.writestr(name,content)
        assert all(raw[k]==v for k,v in original.items() if k not in {links[b['sheet']],links[b['comparison']]}|extra)
        before=openpyxl.load_workbook(Q/b['key']/'before.xlsx',data_only=False,rich_text=True)
        after=openpyxl.load_workbook(candidate,data_only=False,rich_text=True)
        assert before.sheetnames+(['HEAR 기대값'] if b['key']=='unit' else [])==after.sheetnames
        for name in before.sheetnames:
            if name==b['comparison']:continue
            for r in before[name]:
                for c in r:
                    if name==b['sheet'] and c.coordinate in mainallowed:continue
                    assert str(c.value)==str(after[name][c.coordinate].value),(name,c.coordinate,'Unrequested content changed')
        for identifier,target in b['targets'].items():
            for col,txt in target['changes'].items(): assert str(after[b['sheet']][col+str(target['row'])].value)==txt
        assert not any(c.data_type=='e' for s in after for r in s for c in r)
        results.append({'key':b['key'],'sha256':sha(candidate),'cells':len(mainallowed),'added_comparison_rows':inserted,'unrelated_parts_unchanged':True})
    save_json(Q/'checks.json',results); print(json.dumps(results,ensure_ascii=False))

def publish():
    plan=json.loads((Q/'plan.json').read_text('utf-8')); checks=json.loads((Q/'checks.json').read_text('utf-8'))
    reviewed=json.loads((Q/'visual-reviewed.json').read_text('utf-8'))
    for p,h in plan['protected'].items():assert sha(Path(p))==h,'Protected document changed'
    for b,c in zip(plan['books'],checks):
        assert sha(Path(b['path']))==b['sha256'],'External workbook edit'
        assert sha(Q/b['key']/'candidate.xlsx')==c['sha256']==reviewed[b['key']]
    # Publish design input first, then common output. File locks are not bypassed.
    for b in plan['books']:
        dest=Path(b['path']); tmp=dest.with_name('.validation-review-publication.tmp')
        shutil.copy2(Q/b['key']/'candidate.xlsx',tmp)
        try:
            assert sha(dest)==b['sha256']; os.replace(tmp,dest)
        finally:
            if tmp.exists():tmp.unlink()
    save_json(Q/'publication.json',checks); print('Published same five filenames; protected docs untouched')

if __name__=='__main__': {'prepare':prepare,'finish':finish,'publish':publish}[sys.argv[1]]()
