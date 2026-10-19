"""통합 기준서/보고서의 배열 표시를 좁은 범위에서 수정하고 검증한다.

내용 작성은 Artifact Tool에서 수행한다. 원본 패키지는 수정한 셀의 텍스트와
필요한 행 높이만 이식하여 표지/서식/병합/이력 및 사용자 편집을 보존한다.
"""
import copy
import hashlib
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

from lxml import etree as E
from openpyxl import load_workbook

import preserve_aligned_report_packages as native
from integration_rev0_content import FRAME_ARRAY, readable_frames

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
DHF = ROOT/'docs/참고/인허가/사운드1/DHF'
QA = TASK/'qa/integration-readable-frames'
BOOKS = {
    'report': (DHF/'04_설계 검증/SW/펌웨어/SITR-FW-TD24A 소프트웨어 통합시험 보고서_rev0(김은수 작성).xlsx', '보고서_통합시험'),
    'criterion': (DHF/'03_설계 출력/SW/펌웨어/SITC-FW-TD24A 소프트웨어 통합시험 기준서.xlsx', '기준서_통합 시험'),
}
N = native.N
TAG = '{'+native.NS+'}'

def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def save(name, value):
    (QA/name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
def read(name): return json.loads((QA/name).read_text(encoding='utf-8'))
def write_package(path, package):
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in package.items(): z.writestr(name, data)

def audit(path):
    wb=load_workbook(path, rich_text=True)
    pkg=native.parts(path)
    formulas=[]; arrays=[]; references=[]; errors=[]
    for s in wb:
        for row in s:
            for c in row:
                value=str(c.value or '')
                if c.data_type == 'f': formulas.append({'sheet':s.title, 'cell':c.coordinate,'value':value})
                if re.search(r'\[\s*"E\d+/A\d+"',value): arrays.append({'sheet':s.title,'cell':c.coordinate})
                if re.search(r"(?:'[^']+'|[A-Za-z_][\w ]*)!\$?[A-Z]{1,3}\$?\d+",value):
                    references.append({'sheet':s.title,'cell':c.coordinate,'value':value})
                if c.data_type == 'e': errors.append({'sheet':s.title,'cell':c.coordinate,'value':value})
    # 인쇄 영역/반복 행의 자기 시트 참조는 문서 설정이며 셀 계산 참조와 구분한다.
    external=[n for n in pkg if n.startswith('xl/externalLinks/')]
    return dict(formulas=formulas, arrays=arrays, cellReferences=references, errors=errors,
                externalLinkParts=external, sheets=wb.sheetnames)

def prepare():
    QA.mkdir(parents=True,exist_ok=True)
    assert not (QA/'plan.json').exists(), '원본 보존 자료가 이미 존재한다. 새 원본으로 덮어쓰지 않는다.'
    plan={'books':{}, 'protected':{}}
    targets={p.resolve() for p,_ in BOOKS.values()}
    for p in DHF.rglob('*'):
        if p.is_file() and not p.name.startswith('~$') and p.suffix.lower() in {'.xlsx','.docx'} and p.resolve() not in targets:
            plan['protected'][str(p)] = digest(p)
    for key,(path,sheet) in BOOKS.items():
        backup=QA/f'{key}-before.xlsx'; shutil.copy2(path,backup)
        wb=load_workbook(backup,rich_text=True); s=wb[sheet]; changes=[]
        for sh in wb:
            for row in sh:
                for c in row:
                    if c.value is None: continue
                    before=str(c.value); after=readable_frames(before)
                    if after!=before:
                        assert sh.title==sheet, ('unexpected array in another sheet',sh.title,c.coordinate)
                        assert c.coordinate in {'M107','M108','M109','M110'}
                        frames=json.loads(FRAME_ARRAY.search(before).group(1))
                        # 수치와 순서는 유지하고 배열의 표시 방식만 바꾼다.
                        got=[]
                        for electrode,amplitude in re.findall(r'프레임 \d+~\d+: 전극 (\d+)\(크기 (\d+)\)',after):
                            got.extend([f'E{electrode}/A{amplitude}','NOP','NOP'])
                        assert got==frames
                        rowheight=max(s.row_dimensions[c.row].height or 15, 270)
                        changes.append(dict(cell=c.coordinate,row=c.row,before=before,after=after,frames=frames,height=rowheight))
        assert len(changes)==4, (key,len(changes))
        check=audit(backup)
        assert not check['formulas'] and not check['cellReferences'] and not check['externalLinkParts']
        assert len(check['arrays'])==4
        plan['books'][key]=dict(path=str(path),backup=str(backup),hash=digest(path),sheet=sheet,changes=changes,audit=check)
    save('plan.json',plan)
    print('원본 2파일/각 4셀 수정, 수식/다른 시트 참조 0, 보호 파일',len(plan['protected']))

def preserve():
    plan=read('plan.json')
    for key,b in plan['books'].items():
        target=native.parts(b['backup']); authored=native.parts(QA/f'{key}-authored.xlsx')
        sp=native.sheet_path(target,b['sheet']); ap=native.sheet_path(authored,b['sheet'])
        root=native.parse(target[sp]); ar=native.parse(authored[ap])
        sst=native.parse(authored['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in authored else None
        for change in b['changes']:
            addr=change['cell']; old=root.find(f'.//s:c[@r="{addr}"]',N); new=ar.find(f'.//s:c[@r="{addr}"]',N)
            assert old is not None and new is not None
            old.attrib.pop('t',None)
            for child in list(old): old.remove(child)
            if new.get('t')=='s':
                si=sst[int(new.find('s:v',N).text)]; value=''.join(si.itertext())
            elif new.get('t')=='inlineStr': value=''.join(new.find('s:is',N).itertext())
            else: value=new.find('s:v',N).text
            assert value==change['after'],(key,addr,value)
            old.set('t','inlineStr'); ins=E.SubElement(old,TAG+'is'); text=E.SubElement(ins,TAG+'t')
            text.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); text.text=value
            row=root.find(f's:sheetData/s:row[@r="{change["row"]}"]',N)
            row.set('ht',str(change['height']));row.set('customHeight','1')
        target[sp]=native.serialize(root)
        write_package(QA/f'{key}-candidate.xlsx',target)
    print('원본 패키지의 해당 4셀/행 높이만 반영')

def verify():
    plan=read('plan.json'); summary={}
    for key,b in plan['books'].items():
        before=load_workbook(b['backup'],rich_text=True); after=load_workbook(QA/f'{key}-candidate.xlsx',rich_text=True)
        assert before.sheetnames==after.sheetnames
        addresses={c['cell']:c for c in b['changes']}; altered=[]; counts=0
        for old,new in zip(before,after):
            assert old.max_row==new.max_row and old.max_column==new.max_column
            assert list(map(str,old.merged_cells.ranges))==list(map(str,new.merged_cells.ranges))
            for row in old:
                for c in row:
                    n=new[c.coordinate]; counts+=1
                    if c.value!=n.value:
                        assert old.title==b['sheet'] and c.coordinate in addresses,(old.title,c.coordinate)
                        assert n.value==addresses[c.coordinate]['after']; altered.append(c.coordinate)
                    assert c._style==n._style,(key,old.title,c.coordinate,'style')
            assert {k:dict(v) for k,v in old.column_dimensions.items()}=={k:dict(v) for k,v in new.column_dimensions.items()}
            for r,d in old.row_dimensions.items():
                expected=addresses.get('M'+str(r),{}).get('height',d.height) if old.title==b['sheet'] else d.height
                assert new.row_dimensions[r].height==expected,(key,old.title,r)
        assert set(altered)==set(addresses)
        check=audit(QA/f'{key}-candidate.xlsx')
        assert not any(check[x] for x in ['formulas','arrays','cellReferences','errors','externalLinkParts'])
        bp=native.parts(b['backup']); np=native.parts(QA/f'{key}-candidate.xlsx')
        changed=[p for p in bp if bp[p]!=np[p]]
        assert changed==[native.sheet_path(bp,b['sheet'])],changed
        summary[key]=dict(checkedCells=counts,changedCells=altered,packageChanges=changed,audit=check,sha256=digest(QA/f'{key}-candidate.xlsx'))
    report=load_workbook(QA/'report-candidate.xlsx',rich_text=True)['보고서_통합시험']
    criterion=load_workbook(QA/'criterion-candidate.xlsx',rich_text=True)['기준서_통합 시험']
    for row in report.iter_rows(max_col=13):
        for c in row:
            if c.coordinate=='B2':continue
            # 작성자가 보고서 앞뒤 줄바꿈을 지운 편집은 원복하지 않는다.
            # 의미 있는 본문 내용이 같은지를 대조한다.
            assert str(c.value or '').strip()==str(criterion[c.coordinate].value or '').strip(),c.coordinate
    for row in [107,108,109,110]:assert report.row_dimensions[row].height==criterion.row_dimensions[row].height
    for p,h in plan['protected'].items():assert digest(p)==h,p
    save('verification.json',summary);print(json.dumps(summary,ensure_ascii=False,indent=2))

def publish():
    plan=read('plan.json');verification=read('verification.json')
    for key,b in plan['books'].items():
        assert digest(b['path'])==b['hash'],('외부 편집 발견',b['path'])
        assert digest(QA/f'{key}-candidate.xlsx')==verification[key]['sha256']
    for p,h in plan['protected'].items():assert digest(p)==h,p
    staged=[]
    for key,b in plan['books'].items():
        destination=Path(b['path']); stage=destination.with_name(destination.name+'.tdc-stage')
        assert not stage.exists();shutil.copy2(QA/f'{key}-candidate.xlsx',stage);staged.append((key,stage,destination))
    published=[]
    try:
        for key,stage,destination in staged:stage.replace(destination);published.append(key)
    except OSError:
        for key in published:shutil.copy2(plan['books'][key]['backup'],plan['books'][key]['path'])
        raise
    saved={key:dict(path=b['path'],sha256=digest(b['path'])) for key,b in plan['books'].items()}
    for key,b in saved.items():assert b['sha256']==verification[key]['sha256']
    save('publication.json',saved);print('동일 이름으로 2파일 저장 완료')

if __name__=='__main__':globals()[sys.argv[1]]()
