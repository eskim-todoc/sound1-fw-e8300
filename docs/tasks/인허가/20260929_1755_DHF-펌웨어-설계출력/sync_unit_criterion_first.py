"""보고서 최신 본문의 A~K를 기준서 1차본에 서식과 함께 옮긴다.

Artifact Tool에서 결과 열 내용 제거와 제목 편집을 수행한다. 지원되지 않는
시트 간 원본 서식 이식/물리 열 제거는 기존 OOXML 보존 도구로 보완한다.
표지, 대상 문서명, 기존 기준서 시트 이름은 보존한다.
"""
from pathlib import Path
import copy
import hashlib
import json
import re
import shutil
import sys
import zipfile
from lxml import etree as E
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
import preserve_aligned_report_packages as native

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
QA = TASK / 'qa/unit-criterion-from-report'
DHF = ROOT / 'docs/참고/인허가/사운드1/DHF'
REPORT = DHF / '04_설계 검증/SW/펌웨어/SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서_1차본.xlsx'
TARGET = DHF / '03_설계 출력/SW/펌웨어/SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서_1차본.xlsx'
RS, CS = '보고서_유닛 시험', '기준서_유닛 검증'
USE_SOURCE_THEME = False
N = native.N
TAG = '{' + native.NS + '}'


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save(name, value):
    (QA / name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def read(name):
    return json.loads((QA / name).read_text(encoding='utf-8'))


def prepare():
    QA.mkdir(parents=True, exist_ok=False)
    shutil.copy2(REPORT, QA / 'report-before.xlsx')
    shutil.copy2(TARGET, QA / 'criterion-before.xlsx')
    s = load_workbook(REPORT, rich_text=True)[RS]
    unmerges = []
    for m in s.merged_cells.ranges:
        assert not (m.min_col <= 11 < m.max_col), str(m)
        if m.min_col > 11:
            unmerges.append(dict(key='criterion', sheet=RS, range=str(m)))
    protected = {str(p):digest(p) for p in DHF.rglob('*')
                 if p.is_file() and p.suffix in ('.xlsx','.docx') and p != TARGET and not p.name.startswith('~$')}
    save('baseline.json', dict(sourceHash=digest(REPORT), targetHash=digest(TARGET),
         protected=protected, rows=s.max_row,
         books=[dict(key='criterion', path=str(TARGET), backup=str(QA / 'report-before.xlsx'),
                     sheet=RS, reviewSheet=CS, views=['J28:K34','G182:H188','K1007:K1009','B2:D24'])]))
    save('plan.json', dict(changes=[dict(key='criterion', sheet=RS, cell='B2',
         value=str(s['B2'].value).replace('유닛 시험 보고서','유닛 시험 기준서'))],
         unmerges=unmerges, clears=[dict(key='criterion',sheet=RS,range=f'L1:M{s.max_row}')],
         heights=[],shades=[]))
    print(json.dumps(dict(sourceHash=digest(REPORT),targetHash=digest(TARGET),
                         rows=s.max_row,removeMerges=len(unmerges)),ensure_ascii=False))


def fixed_font_schemes(root):
    """서로 다른 문서 테마가 Calibri를 다른 글꼴로 바꾸지 않도록 고정한다."""
    for sc in list(root.findall('.//s:scheme', N)):
        parent = sc.getparent()
        name = parent.find('s:name', N)
        if name is None:
            name = parent.find('s:rFont', N)
        assert name is not None and name.get('val') in ('Calibri','Calibri Light'), E.tostring(parent)
        parent.remove(sc)


def build():
    source = native.parts(QA / 'report-before.xlsx')
    target = native.parts(QA / 'criterion-before.xlsx')
    if USE_SOURCE_THEME:
        # 표지의 테마 글꼴 연결은 원래 이름으로 고정한다. 본문 테마 교체가
        # 표지의 맑은 고딕을 Calibri로 바꾸지 않도록 표지에만 적용한다.
        cover=native.parse(target[native.sheet_path(target,'표지')])
        target_styles=native.parse(target['xl/styles.xml'])
        fonts=target_styles.find('s:fonts',N)
        xfs=target_styles.find('s:cellXfs',N)
        shared_target=native.parse(target['xl/sharedStrings.xml'])
        for cell in cover.findall('.//s:c',N):
            font=fonts[int(xfs[int(cell.get('s','0'))].get('fontId','0'))]
            scheme=font.find('s:scheme',N)
            if scheme is not None:
                assert font.find('s:name',N) is not None
                font.remove(scheme)
            if cell.get('t')=='s':
                value=cell.find('s:v',N)
                if value is not None:
                    for scheme in list(shared_target[int(value.text)].findall('.//s:scheme',N)):
                        scheme.getparent().remove(scheme)
        target['xl/styles.xml']=native.serialize(target_styles)
        target['xl/sharedStrings.xml']=native.serialize(shared_target)
    original_styles = native.parse(source['xl/styles.xml'])
    if not USE_SOURCE_THEME:
        fixed_font_schemes(original_styles)
    source['xl/styles.xml'] = native.serialize(original_styles)
    shared = native.parse(source['xl/sharedStrings.xml'])
    if not USE_SOURCE_THEME:
        fixed_font_schemes(shared)
    source['xl/sharedStrings.xml'] = native.serialize(shared)
    report_path = native.sheet_path(source, RS)
    criterion_path = native.sheet_path(target, CS)
    imp = native.Importer(target, source, 'criterion_report_source')
    root = imp.convert_sheet(native.parse(source[report_path]))
    assert not root.findall('.//s:f', N)
    assert not root.findall('.//s:drawing', N)
    assert not root.findall('.//s:conditionalFormatting', N)
    data = root.find('s:sheetData', N)
    for row in data:
        for cell in list(row):
            if native.col(cell.get('r')) > 11:
                row.remove(cell)
        if 'spans' in row.attrib:
            lo, hi = map(int, row.get('spans').split(':'))
            row.set('spans', f'{min(lo,11)}:{min(hi,11)}')
    # 결과 열 삭제 후 Excel 자동 높이 계산으로 줄어드는 행만 원래 화면 높이로 고정한다.
    height_path=QA/'native-heights.json'
    if height_path.exists():
        for item in read('native-heights.json'):
            row=data.find(f's:row[@r="{item["row"]}"]',N)
            assert row is not None,item['row']
            row.set('ht',str(item['height']))
            row.set('customHeight','1')
    merges = root.find('s:mergeCells', N)
    for m in list(merges):
        if range_boundaries(m.get('ref'))[0] > 11:
            merges.remove(m)
    merges.set('count', str(len(merges)))
    cols = root.find('s:cols', N)
    for col in list(cols):
        lo, hi = int(col.get('min')), int(col.get('max'))
        if lo in (12,13) and hi <= 13:
            cols.remove(col)
        elif lo > 13:
            col.set('min', str(lo - 2))
            col.set('max', str(hi - 2 if hi < 16384 else hi))
    rows = read('baseline.json')['rows']
    root.find('s:dimension', N).set('ref', f'A1:K{rows}')
    # 기준서 제목 한 셀 외의 공통 열 문구를 바꾸지 않는다.
    authored = load_workbook(QA / 'criterion-authored.xlsx', rich_text=True)[RS]
    planned_title = read('plan.json')['changes'][0]['value']
    assert str(authored['B2'].value) == planned_title
    title = root.find('.//s:c[@r="B2"]', N)
    for child in list(title):
        title.remove(child)
    title.set('t','inlineStr')
    text = E.SubElement(E.SubElement(title, TAG + 'is'), TAG + 't')
    text.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
    text.text = planned_title
    # 시트 전용 인쇄 설정만 원본 보고서 것으로 교체하고 기준서 표지는 유지한다.
    srcrels = native.parse(source[native.relpath(report_path)])
    dstrels = native.parse(target[native.relpath(criterion_path)])
    assert len(srcrels) == len(dstrels) == 1
    assert srcrels[0].get('Type').endswith('/printerSettings')
    assert dstrels[0].get('Type').endswith('/printerSettings')
    target[native.resolve(criterion_path,dstrels[0].get('Target'))] = source[native.resolve(report_path,srcrels[0].get('Target'))]
    root.find('s:pageSetup',N).set('{'+native.RNS+'}id',dstrels[0].get('Id'))
    # phoneticPr 역시 새 스타일 자원 번호로 연결한다.
    phonetic = root.find('s:phoneticPr', N)
    if phonetic is not None and 'fontId' in phonetic.attrib:
        phonetic.set('fontId', str(imp.maps['fonts'][int(phonetic.get('fontId'))]))
    target[criterion_path] = native.serialize(root)
    imp.flush()
    workbook = native.parse(target['xl/workbook.xml'])
    for name in workbook.findall('s:definedNames/s:definedName', N):
        if name.get('localSheetId') == '1':
            if name.get('name') == '_xlnm.Print_Area':
                name.text = f"'{CS}'!$B$2:$K${rows}"
            elif name.get('name') == '_xlnm.Print_Titles':
                name.text = f"'{CS}'!$28:$29"
    target['xl/workbook.xml'] = native.serialize(workbook)
    if USE_SOURCE_THEME:
        # 보고서의 본문 테마 글꼴까지 같게 맞춘다. 표지의 값/셀 서식은 유지하고
        # 실제 Excel의 표지 글꼴/화면을 별도로 대조한 경우에만 이 옵션을 사용한다.
        target['xl/theme/theme1.xml'] = source['xl/theme/theme1.xml']
    with zipfile.ZipFile(QA / 'criterion-candidate.xlsx','w',zipfile.ZIP_DEFLATED) as z:
        for name, raw in target.items():
            z.writestr(name, raw)
    print('기준서 표지 보존 / 보고서 A~K 이식 / L,M 물리 열 제거')


def font_key(font):
    font = copy.copy(font)
    if font.name in ('Calibri','Calibri Light'):
        font.scheme = None
    return font


def rich_key(value):
    if not isinstance(value, list):
        return str(value) if value is not None else ''
    result=[]
    for segment in value:
        if isinstance(segment,str):
            result.append((segment,None))
        else:
            font=copy.copy(segment.font)
            if font.rFont in ('Calibri','Calibri Light'):
                font.scheme=None
            result.append((segment.text, font))
    return result


def verify():
    old=load_workbook(QA/'criterion-before.xlsx',rich_text=True)
    report=load_workbook(QA/'report-before.xlsx',rich_text=True)
    authored=load_workbook(QA/'criterion-authored.xlsx',rich_text=True)
    new=load_workbook(QA/'criterion-candidate.xlsx',rich_text=True)
    assert old.sheetnames==new.sheetnames
    s,d=report[RS],new[CS]
    source_parts=native.parts(QA/'report-before.xlsx')
    source_root=native.parse(source_parts[native.sheet_path(source_parts,RS)])
    explicit={c.get('r') for c in source_root.findall('.//s:c',N)}
    assert d.max_row==s.max_row and d.max_column==11
    assert set(map(str,d.merged_cells.ranges))=={str(m) for m in s.merged_cells.ranges if m.max_col<=11}
    height_overrides={p['row']:p['height'] for p in read('native-heights.json')} if (QA/'native-heights.json').exists() else {}
    for r,dim in s.row_dimensions.items():
        original=dict(dim)
        if r in height_overrides:
            original['ht']=format(height_overrides[r],'g')
            original['customHeight']='1'
        assert dict(d.row_dimensions[r])==original,(r,'row dimension')
    cells=0
    for row in s.iter_rows(max_col=11):
        for c in row:
            nc=d[c.coordinate]
            text='' if c.value is None else str(c.value)
            expected=text.replace('유닛 시험 보고서','유닛 시험 기준서') if c.coordinate=='B2' else text
            assert ('' if nc.value is None else str(nc.value))==expected,c.coordinate
            if c.coordinate=='B2':
                assert str(authored[RS][c.coordinate].value)==expected,c.coordinate
            if c.coordinate!='B2':
                assert rich_key(c.value)==rich_key(nc.value),(c.coordinate,'rich text')
            # 비어 있고 XML에도 없는 셀/병합의 하위 셀은 문서 기본 글꼴을
            # 상속하지만 표시 문구가 없다. 표시 셀의 실제 글꼴을 대조한다.
            if c.coordinate in explicit and c.__class__.__name__!='MergedCell':
                assert font_key(c.font)==font_key(nc.font),(c.coordinate,'font')
            for attr in ('alignment','fill','border','number_format','protection'):
                assert copy.copy(getattr(c,attr))==copy.copy(getattr(nc,attr)),(c.coordinate,attr)
            cells+=1
    for key,col in s.column_dimensions.items():
        if col.min>11:
            continue
        nc=d.column_dimensions[key]
        assert {k:v for k,v in dict(col).items() if k!='style'}=={k:v for k,v in dict(nc).items() if k!='style'},key
        assert font_key(col.font)==font_key(nc.font),key
        for attr in ('alignment','fill','border','number_format','protection'):
            assert copy.copy(getattr(col,attr))==copy.copy(getattr(nc,attr)),(key,attr)
    for row in old['표지']:
        for c in row:
            nc=new['표지'][c.coordinate]
            if USE_SOURCE_THEME:
                assert str(c.value)==str(nc.value),('표지',c.coordinate)
            else:
                assert c.value==nc.value,('표지',c.coordinate)
            for attr in ('font','alignment','fill','border','number_format','protection'):
                expected,actual=copy.copy(getattr(c,attr)),copy.copy(getattr(nc,attr))
                if USE_SOURCE_THEME and attr=='font':
                    expected.scheme=actual.scheme=None
                assert expected==actual,('표지',c.coordinate,attr)
    target_parts=native.parts(QA/'criterion-before.xlsx')
    saved=native.parts(QA/'criterion-candidate.xlsx')
    assert target_parts['xl/worksheets/sheet1.xml']==saved['xl/worksheets/sheet1.xml']
    expected_theme=source_parts['xl/theme/theme1.xml'] if USE_SOURCE_THEME else target_parts['xl/theme/theme1.xml']
    assert expected_theme==saved['xl/theme/theme1.xml']
    ids=[str(c.value).strip() for row in d.iter_rows(min_col=4,max_col=4) for c in row if 'SW-FW-UnitTest-' in str(c.value)]
    assert len(ids)==len(set(ids))==115
    assert not any(c.value is not None for row in d.iter_rows(min_col=12,max_col=13) for c in row)
    assert not any(c.value is not None for row in authored[RS].iter_rows(min_col=12,max_col=13) for c in row)
    root=native.parse(saved[native.sheet_path(saved,CS)])
    assert not any(native.col(c.get('r'))>11 for c in root.findall('.//s:c',N))
    assert d.print_area==f"'{CS}'!$B$2:$K${s.max_row}"
    for path,h in read('baseline.json')['protected'].items():
        assert digest(path)==h,('보호 파일 변경',path)
    assert digest(TARGET)==read('baseline.json')['targetHash'],'기준서 추가 수정 감지'
    summary=dict(rows=s.max_row,columns=11,testIDs=115,checkedCells=cells,
                 merges=len(d.merged_cells.ranges),sourceValuesPreservedExceptTitle=True,
                 effectiveFontsRichTextBordersMergesWidthsHeightsPreserved=True,
                 criterionCoverPreserved=True,resultColumnsRemoved=True,
                 nativeHeightRowsFixed=len(height_overrides),
                 protectedFiles=len(read('baseline.json')['protected']),
                 hash=digest(QA/'criterion-candidate.xlsx'))
    save('verification.json',summary)
    print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__=='__main__':
    {'prepare':prepare,'build':build,'verify':verify}[sys.argv[1]]()
