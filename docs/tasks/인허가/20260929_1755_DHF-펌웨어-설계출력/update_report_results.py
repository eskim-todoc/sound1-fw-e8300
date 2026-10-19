"""보고서 결과를 채우고 자료 표시/짝수 ID 배경을 갱신한다.

원본을 보존하고 Artifact Tool이 작성한 텍스트와 대조한 뒤, 승인된 셀만
native OOXML에 반영한다. 실제 수집값과 예상값 기준의 시험자 확인 기록을
구분하며 이전 불일치와 실제 계측 부재를 임의의 통과값으로 바꾸지 않는다.
openpyxl은 입력과 검증에만 사용한다.
"""
from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path

from lxml import etree as E
from openpyxl import load_workbook

from preserve_aligned_report_packages import N, NS, parse, parts, serialize, sheet_path

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
DHF = ROOT / 'docs/참고/인허가/사운드1/DHF'
QA = TASK / 'qa/report-result-format'
UPLAN = TASK / 'qa/unit-inline-acceptance/plan.json'
XMLSPACE = '{http://www.w3.org/XML/1998/namespace}space'
LEGEND = (
    '시험자가 문제없다고 확인한 내용은 예상값 기준으로 결과를 기입했다. '
    '이 칸의 RTT 전문과 값은 새로 수집한 로그 원문이 아니다. '
    '범위로 정의된 값은 범위 충족으로 기록하며 임의의 실측값을 만들지 않는다. '
    '이전 불일치와 계측 기록이 없는 항목은 별도로 남긴다.'
)
PHOTO = '(사진필요)'
QCC = '(QCC 검증 결과 필요)'
RED = 'C00000'
BLUE = '0000FF'
BLACK = '000000'
BOOKS = [
    dict(key='unit', file='SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서.xlsx',
         sheet='보고서_유닛 시험', idcol=4, resultcol=12, verdictcol=13,
         common=11, first=30,
         criterion='SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx',
         csheet='기준서_유닛 검증', views=['K161:M168','K266:M268','K125:M131']),
    dict(key='integration', file='SITR-FW-TD24A 소프트웨어 통합시험 보고서.xlsx',
         sheet='보고서_통합시험', idcol=3, resultcol=14, verdictcol=15,
         common=13, first=25,
         criterion='SITC-FW-TD24A 소프트웨어 통합시험 기준서.xlsx',
         csheet='기준서_통합 시험', views=['M25:O26','M27:O27','M35:O35']),
    dict(key='system', file='SSTR-FW-TD24A 소프트웨어 시스템 시험 보고서.xlsx',
         sheet='보고서_시스템 시험', idcol=2, resultcol=11, verdictcol=12,
         common=10, first=34,
         criterion='[작성자 확인] SSTC-FW-TD24A 소프트웨어 시스템 시험 기준서.xlsx',
         csheet='기준서_시스템 시험', views=['J41:L41','J50:L51','J64:L64','J104:L104']),
]


def digest(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def save_json(p, value):
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def text(v):
    return '' if v is None else str(v)


def confirmed(value):
    """정해진 입력/수치를 바꾸지 않고 시험자 확인 기록의 문체로 바꾼다."""
    value = value.replace('·', '/')
    value = value.replace('관찰값 비교 기준:', '관찰 결과:')
    for a, b in [
        ('확인되어야 한다', '확인됨'), ('확인 되어야 한다', '확인됨'),
        ('확인할 수 있어야 한다', '확인됨'), ('보여야 한다', '확인됨'),
        ('들려야 한다', '들림'), ('구분되어야 한다', '구분됨'),
        ('동일해야 한다', '동일함'), ('일치해야 한다', '일치함'),
        ('없어야 한다', '없음'), ('유지되어야 한다', '유지됨'),
        ('나타나야 한다', '나타남'), ('표시되어야 한다', '표시됨'),
        ('전달되어야 한다', '전달됨'), ('발생해야 한다', '발생함'),
        ('처리되어야 한다', '처리됨'), ('충족해야 한다', '충족함'),
        ('않아야 한다', '않음'), ('되어야 한다', '됨'),
        ('이어야 한다', '임'), ('해야 한다', '함'),
        ('예상 결과:', '결과:'), ('예상 RTT 전문:', 'RTT 결과:'),
        ('예상 자극 순서=', '자극 순서='),
    ]:
        value = value.replace(a, b)
    return value.strip()


def attachments(old):
    """기존 자료 종류만 유지한다. 파일명/사진/계측값을 생성하지 않는다."""
    result = []
    for mark in (PHOTO, QCC):
        for line in old.splitlines():
            if mark not in line:
                continue
            chunk = line[line.index(mark):]
            other = QCC if mark == PHOTO else PHOTO
            if other in chunk:
                chunk = chunk[:chunk.index(other)].rstrip(' /;')
            chunk = chunk.replace('보관한다.', '첨부.').replace('구분하여 첨부한다.', '첨부.')
            chunk = chunk.replace('추가한다.', '첨부.')
            if chunk not in result:
                result.append(chunk)
    return '\n'.join(result)


def colored_segments(value):
    """각 자료 표지부터 같은 줄 끝/다음 표지까지 색과 굵기를 부여한다."""
    spans = []
    for m in re.finditer(r'\(사진필요\)|\(QCC 검증 결과 필요\)', value):
        end = value.find('\n', m.start())
        if end < 0:
            end = len(value)
        next_mark = re.search(r'\(사진필요\)|\(QCC 검증 결과 필요\)', value[m.end():end])
        if next_mark:
            end = m.end() + next_mark.start()
        spans.append((m.start(), end, BLUE if m.group() == PHOTO else RED))
    out = []; pos = 0
    for start, end, color in spans:
        if start > pos:
            out.append([value[pos:start], BLACK, False])
        out.append([value[start:end], color, True]); pos = end
    if pos < len(value):
        out.append([value[pos:], BLACK, False])
    return out or [[value, BLACK, False]]


def clean_old_issue(v):
    """대조 요청 대신 이미 확인된 불일치 내용을 남긴다."""
    for a, b in [
        ('결과 대조 필요', '기존 기록 불일치'),
        ('신규 결과 대조 대기', '이전 기록'),
        ('새 결과 대조 대기', '별도 실제 기록 없음'),
        ('신규 실제 결과를 생성하지 않음', '기존 기록 보존'),
    ]:
        v = v.replace(a, b)
    return v


def prepare():
    QA.mkdir(parents=True, exist_ok=True)
    if (QA / 'plan.json').exists():
        raise RuntimeError('이미 준비된 계획을 덮어쓰지 않는다.')
    unit_plan = json.loads(UPLAN.read_text(encoding='utf-8'))
    unit_records = unit_plan['records']
    baselines = []
    for spec in BOOKS:
        b = dict(spec)
        b['path'] = str(DHF / '04_설계 검증/SW/펌웨어' / b['file'])
        b['backup'] = str(QA / (b['key'] + '-before.xlsx'))
        b['sha256'] = digest(b['path'])
        shutil.copyfile(b['path'], b['backup'])
        b['criterionPath'] = str(DHF / '03_설계 출력/SW/펌웨어' / b['criterion'])
        b['criterionHash'] = digest(b['criterionPath'])
        baselines.append(b)

    oldprotected = unit_plan['baseline']['protected']
    protected = {p: digest(p) for p in oldprotected if Path(p).exists()
                 and p not in {b['path'] for b in baselines}}
    for p in DHF.rglob('*'):
        if p.is_file() and (p.name.startswith('[작성자 전용]') or 'EPTC' in p.name):
            protected[str(p)] = digest(p)
    changes = {}; shades = []; groups = {}; counts = {}

    def put(b, sheet, cell, value, reason):
        if len(value) > 32767:
            raise AssertionError((b['key'], cell, 'Excel cell text limit'))
        changes[(b['key'], sheet, cell)] = dict(
            key=b['key'], sheet=sheet, cell=cell, value=value,
            segments=colored_segments(value), reason=reason)

    for b in baselines:
        w = load_workbook(b['backup'], rich_text=True)
        s = w[b['sheet']]
        rc, vc = b['resultcol'], b['verdictcol']
        letter = s.cell(1, rc).column_letter
        put(b, s.title, letter+'5', LEGEND, '시험자 확인/수집 원문/기댓값 기록 구분')
        put(b, s.title, letter+'10',
            '기준서의 시험 행/조건/절차/통과 기준은 유지한다. '
            '결과와 판정만 추가하며 같은 행에서 비교한다.\n'
            '(QCC 검증 결과 필요) 빨간색 굵은 글씨는 QCC 자료 첨부 위치.\n'
            '(사진필요) 파란색 굵은 글씨는 실제 사진/화면 자료 첨부 위치.\n'
            '짝수 시험 ID의 배경은 RGB(245, 245, 245)이다.', '자료 표시 범례')

        if b['key'] == 'unit':
            groups[b['key']] = [dict(id=r['testId'], start=r['start'], end=r['end'])
                                for r in unit_records]
            known = [r for r in unit_records if '결과 대조 필요' in r['verdict']]
            counts['retainedUnitIssues'] = [r['id'] for r in known]
            for rec in unit_records:
                issue = '결과 대조 필요' in rec['verdict']
                all_actual = '\n'.join(text(s.cell(rr['row'], rc).value) for rr in rec['rows'])
                qcc = QCC in all_actual
                for item in rec['rows']:
                    row = item['row']; old = text(s.cell(row, rc).value)
                    expected = text(s.cell(row, b['common']).value)
                    kind = item['kind']
                    if kind == 'vector':
                        value = expected.replace('입력/원소별 기준:', '입력/원소별 결과:')
                        if issue:
                            value = '예상값 기준 기록 / 기존 불일치 판정은 유지\n' + value
                    elif kind == 'observation':
                        useful = [ln for ln in expected.splitlines() if not ln.startswith((
                            '변동 이유:', '하한 인덱스:', '상한 인덱스:'))]
                        value = confirmed('\n'.join(useful))
                        value += '\n범위/관계식 충족 확인(시험자 확인). 단일 실측값을 임의 지정하지 않음.'
                    elif issue:
                        value = clean_old_issue(old)
                    elif kind == 'rtt':
                        # 전문은 기존 시험자 확인 기록과 같은 값이다. 원문 출처는 5행에 표시.
                        value = old.split('\n기록 구분:')[0]
                    elif kind == 'comparison':
                        if any(t in old for t in ('기입 필요', '____', '대조 대기')):
                            value = confirmed(expected)
                        else:
                            value = old
                    elif kind == 'summary':
                        value = confirmed(expected)
                        extra = attachments(old)
                        if extra:
                            value += '\n' + extra
                        if rec['id'] == 'BOOT-015':
                            value += '\n실제 전송 원시 자료 없음. PC 요청 기록과 실제 메모리 전송은 구분.'
                        if rec['id'] == 'BOOT-018':
                            value += '\nPC reset 요청 확인 기록만 존재. 타깃 재시작 원시 자료 없음.'
                    else:
                        value = old
                    if rec['id'] == 'CFX-020' and row == 164:
                        value = '조건 [3]\n정상 OD 경로에 1000 Hz 순음/무음 입력.\n순음 입력: 1000 Hz 순음이 들리고 구분됨.\n무음 입력: 순음이 들리지 않고 무음으로 구분됨.\n청취 결과: 통과(시험자 확인).'
                    put(b, s.title, f'{letter}{row}', value, '조건별 구체 결과/값 기록')

                if issue:
                    verdict = '기존 기록 불일치'
                elif rec['id'] == 'BOOT-015':
                    verdict = 'PC 요청 확인 / 실제 전송 기록 없음'
                elif rec['id'] == 'BOOT-018':
                    verdict = 'PC 요청 확인 / 타깃 재시작 기록 없음'
                elif qcc:
                    verdict = 'E8300 통과(시험자 확인)\n(QCC 검증 결과 필요)'
                else:
                    verdict = '통과(시험자 확인)'
                put(b, s.title, s.cell(rec['start'], vc).coordinate, verdict, '판정 범위 구분')
        else:
            oldqa = TASK / 'qa/validation-reports'
            names = ['integration-results.json'] if b['key'] == 'integration' else [
                'system-results-1.json', 'system-results-2.json']
            entries = {}
            for name in names:
                for entry in json.loads((oldqa / name).read_text(encoding='utf-8')):
                    entries[entry['id']] = entry
            tests = []
            for row in range(b['first'], s.max_row+1):
                ident = text(s.cell(row, b['idcol']).value)
                if not re.fullmatch(r'SW-FW-(?:IntegrationTest|SYS-Test)-ID-\d+', ident):
                    continue
                tests.append(dict(id=ident, start=row, end=row))
                old = text(s.cell(row, rc).value)
                expected = text(s.cell(row, b['common']).value)
                entry = entries[ident]
                number = int(ident.rsplit('-', 1)[1])
                issue = '결과 대조 필요' in text(s.cell(row, vc).value)
                no_timer = b['key'] == 'system' and number in (12, 58)
                qcc = QCC in old

                lines = ['시험자 확인 기록(예상값 기준)\n'+confirmed(expected)]
                for i, comp in enumerate(entry['comparisons'], 1):
                    value = comp['expected'].replace('·', '/')
                    # 현재 기준서가 수정한 OD/PDM 출력 경로를 옛 예시보다 우선한다.
                    if b['key'] == 'system' and number in (8, 16):
                        value = value.replace('스피커 출력', 'OD1 차동 PDM 호환 청취 수신기 출력')
                    if '____' in value:
                        # 미정인 맵의 전극 번호를 임의로 만들어 넣지 않는다.
                        value = value.replace('전극 ____', '선택 맵의 대응 전극')
                    if b['key'] == 'system' and number in (31, 32):
                        # 이전 예시는 1/5/10 입력에 모두 응답 2를 쓴 오류가 있다.
                        # 명령의 전 상태를 특정하지 않은 조건에는 가짜 응답 바이트를 넣지 않는다.
                        label = comp['label']
                        match = re.search(r'(?:오디오|자극) 볼륨\s*(\d+)', label)
                        if match:
                            value = f'설정/표시 볼륨={match.group(1)}. 해당 단계의 자극 출력 변화가 확인됨.'
                    label = comp['label'].split('=____')[0].strip().rstrip(';/ ')
                    if '____' in label:
                        label = label.split('____')[0].strip().rstrip('=;/ ')
                    lines.append(f'[{i}] {label}\n'+confirmed(value))
                extra = attachments(old)
                if extra:
                    lines.append(extra)
                if issue:
                    # 별도 문서의 과거 기록을 수정/삭제하지 않고 현재 행에도 보존한다.
                    lines.append('기존 보조/분석 기록의 불일치(현재 판정에 보존):\n'+clean_old_issue(old))
                    verdict = '기존 기록 불일치'
                elif no_timer:
                    lines[0] = '계측 기준 / 예상값 기준 기록\n'+expected
                    lines.append('현재 타깃의 연속 시작/완료 시각 계측 경로와 원시 시각 배열은 없음. '
                                 '위 수치는 허용 기준이며 새 실측 시각을 생성한 기록이 아님.')
                    verdict = '계측 기록 없음'
                elif qcc:
                    verdict = 'E8300/관찰 통과(시험자 확인)\n(QCC 검증 결과 필요)'
                else:
                    verdict = '통과(시험자 확인)'
                value = '\n\n'.join(lines)
                put(b, s.title, f'{letter}{row}', value, '기준값/조건별 예상 출력/시험자 확인 결과')
                put(b, s.title, s.cell(row, vc).coordinate, verdict, '판정 범위 구분')
            groups[b['key']] = tests

        # 보조/관찰/예상 그림 시트는 원래 내용을 유지하고 자료 안내의 색만 맞춘다.
        for sheet in w:
            if sheet.title == s.title:
                continue
            for cells in sheet:
                for c in cells:
                    value = text(c.value)
                    if PHOTO in value or QCC in value:
                        put(b, sheet.title, c.coordinate, value, '자료 안내 색상만 갱신')
        for group in groups[b['key']]:
            if int(group['id'].rsplit('-',1)[1]) % 2 == 0:
                shades.append(dict(key=b['key'], sheet=s.title,
                                   start=group['start'], end=group['end'],
                                   firstcol=2, lastcol=vc, color='F5F5F5'))
        counts[b['key']] = dict(tests=len(groups[b['key']]),
                               evenTests=sum(int(t['id'].rsplit('-',1)[1]) % 2 == 0 for t in groups[b['key']]))
        w.close()

    save_json(QA/'baseline.json', dict(books=baselines, protected=protected))
    save_json(QA/'plan.json', dict(changes=list(changes.values()), shades=shades,
                                  groups=groups, counts=counts))
    print(json.dumps(dict(changedCells=len(changes), counts=counts), ensure_ascii=False))


def xmltext(package, cell):
    if cell is None:
        return ''
    if cell.get('t') == 's':
        strings = parse(package['xl/sharedStrings.xml']).findall('s:si', N)
        return ''.join(strings[int(cell.find('s:v', N).text)].itertext())
    if cell.get('t') == 'inlineStr':
        return ''.join(cell.find('s:is', N).itertext())
    v = cell.find('s:v', N)
    return '' if v is None else text(v.text)


def compact():
    """행을 추가하거나 글씨를 줄이지 않고 중복 안내를 제거한다."""
    plan = json.loads((QA/'plan.json').read_text(encoding='utf-8'))
    base = json.loads((QA/'baseline.json').read_text(encoding='utf-8'))
    oldqa = TASK/'qa/validation-reports'
    entries = {}
    for name in ('integration-results.json','system-results-1.json','system-results-2.json'):
        for e in json.loads((oldqa/name).read_text(encoding='utf-8')):
            entries[e['id']] = e
    sys_issues = {
        7:'기존 코드: 전하량 제한 매크로 비활성. 300×251=75300 pC에 대한 오류 표시 경로가 활성화되어 있지 않음.',
        18:'관련 기존 기록: 시험 7의 전하량 제한 매크로 비활성. 1~7 전체 통과와 개별 정상 관찰을 구분.',
        24:'기존 경계 기록: CM3-046 E-07의 처리 중 명령은 actual=0 / expected=74(0x4A). 정상 맵 변경과 별도.',
        35:'기존 경계 기록: CM3-049 E-09 최종 TX[0]/TX[3]=0/0, 기대=240/1. 정상 0x67 응답과 별도.',
        39:'기존 경계 기록: CM3-049 E-09 최종 TX[0]/TX[3]=0/0, 기대=240/1. 정상 0x65 종료 응답과 별도.',
        40:'기존 경계 기록: CM3-049 E-09 최종 TX[0]/TX[3]=0/0, 기대=240/1. 정상 0x66 시작/중단과 별도.',
        41:'기존 기록: 연결 완료 표시와 역방향 수신 표시의 LED 요청/우선순위 구분은 미확정.',
        43:'기존 기록: 내부기 연결 완료 표시와 역방향 수신 녹색 점멸 표시의 범위는 미확정.',
        50:'기존 저장 경계 기록: f_sync/f_close 오류를 무시하는 경로가 존재함. 정상 이벤트 기록과 별도.',
        51:'기존 저장 경계 기록: f_sync/f_close 오류를 무시하는 경로가 존재함. 정상 무결성 오류 기록과 별도.',
        54:'기존 BOOT-011 보조 기록: C6/C108/C109 반환0, C110 반환1 / 기대는 모두 -1. 손상/짧은 이미지 오류 처리 불일치.',
        56:'기존 코드: 볼륨1→4의 상한119→150을 한 번에 대입. 시간에 따른 단계적 자극 증가 경로가 없음. 코드값은 uA 실측값이 아님.',
        61:'기존 코드: 미연결 LED는 잔량>=80% 녹색, 10~79% ORANGE 지속, <10% ORANGE 점멸. 전 구간 녹색 조건과 불일치.',
        65:'기존 코드: 모든 오류 플래그가 적색 오류 LED에 연결되는 것은 아님. 오류 종류별 적용 범위 미확정.',
        72:'CRA 인증 실패/페어링 키 실패 대상과 잠금 시간이 현재 자료에서 미정. 수치 잠금 시간을 임의로 지정하지 않음.',
        73:'기존 PC 저장 경계 기록: C24/C25의 f_sync/f_close 실패 반환=0, 기대=-1.',
        78:'코드 식별: AES-128 ECB, 블록16바이트 / CRC16. 적용 사이버보안 규격의 모드/사용 목적 적합성 근거는 없음.',
        79:'기존 코드: ENABLE_UI_CMD의 --dump_log/--write_integrity_err 자체에 인가 확인 없음. RTT 접근 정책과 무선 접근은 별도.',
        80:'기존 코드: 기본값 저장 반환값을 검사하지 않는 경로가 있음. 정상 초기화와 저장 실패 경계를 구분.',
    }
    int_issues = {
        3:'기존 CFX-015 copy_done: 복사 뒤0 / 기대1. 정상 루프 I2S_handle 호출은 주석 처리. RAM 복사/전환 관찰과 별도.',
        6:'기존 CFX-011: FS=31인데 32대역 맵의 재적재 완료=1. 양의 범위 초과 조건도 별도 검사 없음.',
        8:'기존 CFX-041 reset은 preamble/Reset_PCM만0. FIFO 24개=777, mode=2 유지. FIFO/모드 초기화 요구와 불일치.',
        9:'기존 CM3-006: 동일 맵 요청에서도 재계산/완료 플래그 변경. CM3-028/029 전하량 검사 매크로 비활성.',
        10:'기존 처리 중 새 명령 경계: command actual=0 / expected=74(0x4A). 정상 설정/조회와 별도.',
        13:'기존 정상 루프의 I2S_handle 호출 주석 처리. fade 함수 직접 호출 결과를 실제 스트리밍 전환으로 확대하지 않음.',
        14:'일반 연결은 설정 완료 플래그로 통보. 매핑 경로는 별도 상태를 사용하므로 설정 완료 이전 High 통보의 적용 범위는 미확정.',
        16:'기존 BOOT-011 짧은 구획/읽기 오류: 반환0/0/1 / 기대-1. 이미지 오류 후 응용 실행 금지 요구와 불일치.',
        17:'기존 CM3-051 C17/C18: 응답[195,0] / 기대[240,3], [240,4]. 정상 전송과 오류 응답 조건을 분리.',
    }
    review_only_sys = {41,43,72,78,79}
    real_failure_sys = {7,18,54,56,61,65,73}
    books = {b['key']:b for b in base['books']}
    for key,b in books.items():
        w = load_workbook(b['backup'],rich_text=True); s=w[b['sheet']]
        rc,vc=b['resultcol'],b['verdictcol']
        bycell={(p['sheet'],p['cell']):p for p in plan['changes'] if p['key']==key}
        if key=='unit':
            for p in bycell.values():
                if p['sheet']==s.title and p['cell'] in ('L145','L150','L169','L888','L1262'):
                    p['value']=re.sub(r'\(QCC 검증 결과 필요\)[^\n]*',QCC+' QCC 자료 첨부.',p['value'])
                    p['segments']=colored_segments(p['value'])
            w.close(); continue
        for g in plan['groups'][key]:
            row=g['start']; ident=g['id']; number=int(ident.rsplit('-',1)[1])
            expected=text(s.cell(row,b['common']).value)
            old=text(s.cell(row,rc).value)
            e=entries[ident]; out=[]; seen={}
            if key=='system' and number==2:
                out.append('채널1~24의 기댓값: 각 1000 Hz / 허용950~1050 Hz 충족(시험자 확인).')
            elif key=='system' and number in (18,44,45,57,60):
                out.append(confirmed(expected)+' (시험자 확인 기준 기록)')
            elif key=='system' and number in (21,22,62):
                out.append('충전 연결/해제 각각: 잔량80% 녹색 지속, 79%/10% 노란색 지속, 9% 노란색 점멸.\n'
                           '8개 조건의 표시를 확인함(시험자 확인). 코드 색상 enum은 ORANGE, 주석은 노랑. 사진은 별도 첨부.')
            else:
                for i,c in enumerate(e['comparisons'],1):
                    value=c['expected'].replace('·','/')
                    value=value.replace('스피커 출력','OD1 차동 PDM 호환 청취 수신기 출력') if key=='system' and number in (8,16) else value
                    # 시험 절차/주의/자료 수집 지시는 결과 칸에 다시 복제하지 않는다.
                    omit=('원기준','실제 시험 맵','계측창 내','측정창 길이','전극 1 출력 유무를 기록',
                          '전체 양상 펄스','음압/이득의 실제','순서=[명령','실제 시험','별도 ',
                          '시험 맵의','각 모드의 관찰','초기 출력 설정 시점','원절차의')
                    kept=[ln for ln in value.splitlines() if not ln.startswith(omit)]
                    value='\n'.join(kept).replace('전극 ____','대응 전극')
                    value=value.replace('____','시험 입력에 대응한 값')
                    if key=='system' and number in (31,32):
                        m=re.search(r'(?:오디오|자극) 볼륨\s*(\d+)',c['label'])
                        if m: value=f'설정/표시 볼륨={m.group(1)} / 해당 단계의 자극 출력 변화 확인.'
                    if key=='system' and number in (10,11):
                        value=value.replace('맵에 배정된 전극 대응 전극','맵에 배정된 대응 전극')
                    value=confirmed(value)
                    if not value: continue
                    if value in seen:
                        continue
                    seen[value]=i
                    out.append(f'[{i}] '+value)
            # 현재 기준이 새로 정한 동작 경로/계측 범위는 오래된 예시보다 우선한다.
            if key=='system' and number==17:
                out=['[1] FPGA/PCM 통신 정상 / 자극 데이터 전달 확인.', '[2] 정상 내부기 연결: 흰색 LED 지속 점등 확인(시험자 확인).']
            if key=='system' and number in (12,58):
                out=['계측 기준 / 예상값: '+expected,
                     '연속 시작/완료 시각 원시 기록 없음. 허용값이며 실제 시각 배열을 임의 생성하지 않음.']
            note=(sys_issues if key=='system' else int_issues).get(number)
            if note: out.append(note)
            extra=attachments(old)
            for line in extra.splitlines():
                if QCC in line:
                    out.append(QCC+' QCC 자료 첨부.')
                else:
                    short=line.split('첨부.')[0].split('보관한다.')[0]
                    short=short.replace('원시 ','').replace('실제 ','')
                    # 첫 문장 뒤의 수집 메타데이터 설명은 기준서와 보조 기록에 그대로 존재.
                    short=short.split('. ')[0]
                    out.append(short.rstrip('.')+'.')
            value='\n'.join(out)
            assert '____' not in value,(key,ident,'unfilled expected')
            p=bycell[(s.title,s.cell(row,rc).coordinate)]
            p['value']=value; p['segments']=colored_segments(value)
            vp=bycell[(s.title,s.cell(row,vc).coordinate)]
            if key=='system' and number in (12,58):
                verdict='계측 기록 없음'
            elif key=='system' and number in real_failure_sys:
                verdict='기존 기록 불일치'
            elif key=='system' and number in review_only_sys:
                verdict='적용 범위/근거 미확정'
            elif key=='system' and note:
                verdict='정상 조건 확인 / 경계 기록 보존'
            elif key=='integration' and note:
                verdict='기존 기록 불일치'
            else:
                verdict='통과(시험자 확인)'
            if QCC in old:
                verdict+='\n'+QCC
            vp['value']=verdict;vp['segments']=colored_segments(verdict)
        w.close()
    overrides = {
        ('integration','N26'):'[1] 저전력 요청1 → CFX PCM/오디오/I2S/DMIC/FIFO/HEAR 정리 → 요청0 → CM3 절전.\n[2] 대기 전/후 처리 횟수 동일, 증가량0.\n[3] 터치 → 준비/초기화 → 허가 → 정상 입력 처리 재개.',
        ('integration','N29'):'[1] 혼합[0..15]=[0,-32,17,-100,99,0,0,0,0,0,0,0,0,0,0,0], Event0 뒤 maxAudioInput=100.\n'
            '[2] 무음: 혼합[0..15]=0, Event0 뒤 maxAudioInput=0.\n'
            '[3] 무음512표본: CMD1/Event1 뒤 FFT[0..511]=0(512개 모두0).\n'
            '[4] 무음: vMag[0..255]=0(256개 모두0).\n'
            '[5] 1000Hz 입력 한 주기=[0,25080,46341,60547,65536,60547,46341,25080,0,-25080,-46341,-60547,-65536,-60547,-46341,-25080], 32회 반복. 최대 크기 bin=32(1000Hz). 진폭/블록 지수는 기준 고정값 없음.\n'
            '[6] 혼합 → CMD0/Event0 → AGC/FFT 갱신 → CMD1/Event1 → 대역/매핑 → CIS PCM. 완료 이벤트 뒤 결과 사용.\n'
            '[7] 자극 인덱스=[0,1,2,3,4,5,6,7,8,9,10,11], 전극번호=인덱스+1.\n'
            '[8] 자극 인덱스=[12,13,14,15,16,17,18,19,20,21,22,23], 전극번호=인덱스+1.\n'
            '[9] 자극 인덱스=[24,25,26,27,28,29,30,31,0,1,2,3], 전극번호=인덱스+1.\n'
            '[7~9] 각12개 크기=100, 자극 뒤 NOP(0x25555)1개. 12자극×2프레임=24프레임/ms.',
        ('integration','N33'):'[1] 명령 → 맵 요청/CFX 반영 → CM3 계산/공유 → CFX 적용 → 내부기 준비 → Live. 3프레임/자극, 8채널/ms.\n'
            '[2] NofM/14us/32대역, T_uA[0..31]=100, C_uA[0..31]=300。각32개 모두 같은 값. DAC A/Offset DAC A, 역수20871/32768 및14435/32768.\n'
            '[3] CM3 공유와 CFX 수신: T[0..31]=0, C[0..31]=127(각32개). C=floor((300-100)×20871/32768)=127, offset44. 각 인덱스 동일.\n'
            '[4] Specific 명령의24프레임 처리, 종료 후 종료 상태/PCM 모드 전환. 준비 전 Live 재진입 없음.\n'
            '[5] T=100uA/C=2500uA, 폭25us: DAC 범위 실패를 성공으로 게시하지 않음. 실패 후 T/C 정책은 미정.\n'
            '[6] T=100uA/C=300uA, 폭251us: 75300pC>75000pC, 거부 반환 기댓값0.\n'
            '기존 CM3-006: 동일 맵도 재계산/완료 플래그 변경. CM3-028/029 전하량 검사 비활성.\n'+QCC+' QCC 자료 첨부.',
        ('integration','N37'):'[1] DMIC1/I2S[0..15]=320(각16개). 입력>>5=10. 최종 /2 평균 없음.\n'
            '[2] 마이크 게인65536 / I2S 게인131072: 혼합[0..15]=30(16개 모두30), Event0 최대값30.\n'
            '[3] 마이크 게인65536 / I2S unity65536: 혼합[0..15]=20(16개 모두20), 최대값20.\n'
            '[4] 매핑 중 두 게인65536: 혼합[0..15]=20, 최대값20.\n'
            '[5] 매핑 중 저장 A/B 인덱스 유지. 해제 후 크래들 출력16개=30 / 일반 출력16개=20.\n'
            '[6] 전환의 기댓값: 시작은 블록 크기 증가, 중단은 감소 후 마이크 복귀.\n'
            '기존 정상 루프 I2S_handle 호출 주석 처리. fade 직접 호출과 실제 전환을 분리.\n'+QCC+' QCC 자료 첨부.',
        ('integration','N39'):'[1] 기준16개=0. 지연 현재=[0,64,128,192,256,320,384,448,512,576,640,704,768,832,896,960], 직전=[2048,0,0,0,0,0,0,0,0,0,0,0,0,0,0,0]. 왼쪽DMIC2/오른쪽DMIC1 지연.\n'
            '혼합=[1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,32], 최대값32.\n'
            '[2] 입력32개=100: 선택=[0,1,2,3,4,5,6,7,8,9,10,11,12,13,14,15], 크기16개=100.\n'
            '[3] 입력[i]=i(i=0..31): 선택=[16,17,18,19,20,21,22,23,24,25,26,27,28,29,30,31], 각 크기=인덱스 값.\n'
            '[4] 교차 순서=[0,8,1,9,2,10,3,11,4,12,5,13,6,14,7,15], 직전 마지막15에서 유지.\n'
            '[5] 직전 마지막0: 순서=[15,0,8,1,9,2,10,3,11,4,12,5,13,6,14,7].\n'
            '[6] phase0→1: [0,8,1,9,2,10,3,11].\n'
            '[7] phase1→0: [4,12,5,13,6,14,7,15].\n'
            '[8] 회전 뒤 phase0→1: [15,0,8,1,9,2,10,3].\n'
            '[9] 회전 뒤 phase1→0: [11,4,12,5,13,6,14,7].\n'
            '[6~9] 각 크기8개=100 / 전극=인덱스+1 / 각 자극 뒤 NOP(0x25555)2개. 각 구간8자극+16NOP=24프레임.\n'
            '[10] 동률:16..31 미출력 / 증가값:0..15 미출력. 미선택 내부 배열을0으로 만들라는 조건 아님.',
        ('integration','N41'):'[1] START=[요청 명령,0,p[2]], END=[요청 명령,0,p[1]], 길이3. DFU 연결 플래그는 유지.\n'
            '[2] 상태4→DFU1, 상태0→DFU0, 응답[0x35,1](2바이트).\n'
            '[3~4] s=1/2, t=1/2/3/4: /s/의 MANIFEST.TXT/APP000.FEZ/APP001.FEZ/APP002.FEZ. setup ACK=[195,s,t,1,9]. 저장=[18,52], 길이2, 완료 file_type=0. 입력2바이트는 정상 부팅 이미지가 아님.\n'
            '[5] 오류=[240,2], 정상 완료로 처리하지 않음.\n'
            '[6] 슬롯3/type1=[240,3], 슬롯2/type5=[240,4].\n'
            '[7] OTA 지원 슬롯1..2 / 기동 상태 표현1..4를 구분.\n'
            '[8] 처리 횟수3→선택2. [9] 횟수4→선택1 / BOOT_TRY_DONE / FAIL.\n'
            '[10] 성공선택2 / 실패선택1, sub_state=UNKNOWN.\n'
            '[11] 저장ACK와 불완전 이미지의 기동 성공은 별도.\n'
            '[12] 유닛 보조 RTT:\n'
            '[UT-CM3-050] C1 A[0:4]=192,0,7,3,0\n[UT-CM3-050] C1 PASS N=5 F=0\n'
            '[UT-CM3-050] C2 A[0:4]=193,0,9,3,1\n[UT-CM3-050] C2 PASS N=5 F=0\n'
            '[UT-CM3-050] C3 A[0:4]=53,1,2,1,4\n[UT-CM3-050] C3 PASS N=5 F=0\n'
            '[UT-CM3-050] C4 A[0:4]=53,1,2,0,0\n[UT-CM3-050] C4 PASS N=5 F=0\n'
            '기존 C17/C18 응답[195,0] / 기대[240,3],[240,4].\n'+QCC+' QCC 자료 첨부.',
        ('system','K38'):'[1] 전류 기댓값500uA / 허용450~550uA 충족.\n[2] 전류 기댓값1000uA / 허용900~1100uA 충족.\n[3] 전류 기댓값1500uA / 허용1350~1650uA 충족.\nI[uA]=V[V]/R[ohm]×1000000, 오차率=|I-설정값|/설정값×100<=10%.\n'+PHOTO+' 각 설정별 PBS 부하전압 파형과 부하저항/환산전류 자료 첨부.',
        ('system','K43'):'대역1~32: 각 입력 순음에 대해 선택 맵의 해당 전극에서 자극 출력 확인(시험자 확인).\n각 대역과 전극의 배정은 시험에 사용하는 맵의 대응 관계를 적용.\n'+PHOTO+' 대역별 입력 순음/전극 대응과 PBS 자극 출력 파형 첨부.',
        ('system','K63'):'[1] ACK=[0x44,2] / 상태=[0x43,37,2,4,5,1,2,1].\n[2] ACK=[0x44,3] / 상태=[0x43,37,3,4,5,1,2,1].\n[3] ACK=[0x44,4] / 상태=[0x43,37,4,4,5,1,2,1].\n[4] ACK=[0x44,1] / 상태=[0x43,37,1,4,5,1,2,1].\n각 선택 맵의 자극 출력 변화와 앱 맵 번호2/3/4/1 확인.\n'+PHOTO+' 앱 맵 번호 화면과 변경 전/후 PBS 파형 첨부.\n'+QCC+' QCC 자료 첨부.',
        ('system','K66'):'[1] OFF: ACK=[0x49,2], 상태=[0x43,37,1,4,5,2,2,1]. 앱 꺼짐/일반 LED 소등.\n[2] ON: ACK=[0x49,1], 상태=[0x43,37,1,4,5,1,2,1]. 앱 켜짐/현재 동작 LED 표시.\n'+PHOTO+' 앱 LED 설정 화면과 외부기 OFF/ON 사진 첨부.\n'+QCC+' QCC 자료 첨부.',
        ('system','K70'):'[1] Live 시작 응답=[0x66,2], 하위상태9(HoldOn). 입력 음원에 반응한 자극 출력 확인.\n'
            '[2] 전극 배열/CIS 순서=[1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,16,17,18,19,20,21,22,23,24,25,26,27,28,29,30,31,32].\n'
            '[3] T[0..31]=100uA / C[0..31]=300uA(각32개). 음원 ON/크기 변경에 따라 해당 대역 자극 변화 확인.\n'+PHOTO+' 매핑 Live 시작/유지/종료의 PBS 파형 첨부.\n'+QCC+' QCC 자료 첨부.',
    }
    overrides={k:v.replace('。','. ').replace('오차率','오차율') for k,v in overrides.items()}
    for p in plan['changes']:
        value=overrides.get((p['key'],p['cell'])) if p['sheet'].startswith('보고서_') else None
        if value is not None:
            p['value']=value;p['segments']=colored_segments(value)
    # 기존 문서가 명시한 구형/PC 시험 경계를 간결하게 보존한다.
    scopes={
        'L1164':'\n적용 범위: Board_1_2 ISL91128 주소0x1C/레지스터1의 A5/5A 분석은 비활성 구형 경로. 기준12/34의 적용 보드/주소 원시 기록은 없음. 현재 Board_1_5 물리 리셋 성공으로 확대하지 않음.',
        'L1170':'\n구형 Board_1_3 비활성 경로 E06/E08/E10/E12: I2C 실패 반환actual1 / expected0. 현재 Board_1_5 실패를 뜻하지 않음. disable0x05는 물리 출력 차단 증거가 아님.',
        'L1328':'\n40회 ISR/1ms는 주입 조건. PC tick4000/14440은 실제 Timer2 주기나 핀 전압 측정값이 아님.',
    }
    for p in plan['changes']:
        if p['key']=='unit' and p['sheet']=='보고서_유닛 시험' and p['cell'] in scopes:
            p['value']+=scopes[p['cell']];p['segments']=colored_segments(p['value'])
    # 상단 QCC 내부 자료 추가 위치도 빨간 굵은 글씨로 표시한다. 문구/공통 내용은 그대로다.
    for key in ('unit','integration'):
        b=books[key];w=load_workbook(b['backup'],rich_text=True)
        value=text(w[b['sheet']]['D19'].value);w.close()
        existing=next((p for p in plan['changes'] if p['key']==key and p['sheet']==b['sheet'] and p['cell']=='D19'),None)
        item=dict(key=key,sheet=b['sheet'],cell='D19',value=value,
                  segments=[[value,RED,True]],reason='QCC 내부 자료 추가 위치 강조')
        if existing: existing.update(item)
        else: plan['changes'].append(item)
    save_json(QA/'plan.json',plan)
    print('compact plan updated')


def set_rich(cell, segments):
    for child in list(cell):
        cell.remove(child)
    cell.set('t', 'inlineStr')
    ins = E.SubElement(cell, '{'+NS+'}is')
    for value, color, bold in segments:
        if not value:
            continue
        run = E.SubElement(ins, '{'+NS+'}r')
        prop = E.SubElement(run, '{'+NS+'}rPr')
        E.SubElement(prop, '{'+NS+'}color', rgb='FF'+color)
        E.SubElement(prop, '{'+NS+'}b', val='1' if bold else '0')
        t = E.SubElement(run, '{'+NS+'}t')
        t.set(XMLSPACE, 'preserve'); t.text = value


def polish():
    plan=json.loads((QA/'plan.json').read_text(encoding='utf-8'))
    for p in plan['changes']:
        value=p['value']
        if p['sheet'].startswith('보고서_'):
            value=re.sub(r'(을|를)\s*\.',r'\1 첨부.',value)
            if '입력/원소별 결과:' in value:
                value=confirmed(value).replace('남아야 한다','보존됨')
        if value!=p['value']:
            p['value']=value;p['segments']=colored_segments(value)
    save_json(QA/'plan.json',plan)
    print('result wording polished')


def build(expected_hashes=None):
    base = json.loads((QA/'baseline.json').read_text(encoding='utf-8'))
    plan = json.loads((QA/'plan.json').read_text(encoding='utf-8'))
    for b in base['books']:
        expected = (expected_hashes or {}).get(b['path'], b['sha256'])
        assert digest(b['path']) == expected, ('external edit', b['path'])
        old = parts(b['backup']); authored = parts(QA/(b['key']+'-authored.xlsx'))
        roots = {}; ar = {}
        styles = parse(old['xl/styles.xml'])
        fills = styles.find('s:fills', N); xfs = styles.find('s:cellXfs', N)
        fill = E.SubElement(fills, '{'+NS+'}fill')
        pf = E.SubElement(fill, '{'+NS+'}patternFill', patternType='solid')
        E.SubElement(pf, '{'+NS+'}fgColor', rgb='FFF5F5F5')
        E.SubElement(pf, '{'+NS+'}bgColor', indexed='64')
        fill_id = len(fills)-1; fills.set('count', str(len(fills)))
        stylemap = {}
        def root(sn):
            if sn not in roots:
                roots[sn] = parse(old[sheet_path(old, sn)])
                ar[sn] = parse(authored[sheet_path(authored, sn)])
            return roots[sn]
        for p in (p for p in plan['changes'] if p['key'] == b['key']):
            dst = root(p['sheet'])
            c = dst.find(f".//s:c[@r='{p['cell']}']", N)
            ac = ar[p['sheet']].find(f".//s:c[@r='{p['cell']}']", N)
            assert c is not None, (p['cell'], 'missing styled original')
            # Excel/Artifact Tool은 CRLF를 LF로 정규화한다. 줄바꿈 이외의 차이는 허용하지 않는다.
            assert xmltext(authored, ac).replace('\r\n','\n') == p['value'].replace('\r\n','\n'), (p['cell'], 'artifact authored mismatch')
            set_rich(c, p['segments'])
        for m in (p for p in plan.get('merges', []) if p['key'] == b['key']):
            dst = root(m['sheet'])
            authored_merges = {x.get('ref') for x in ar[m['sheet']].findall('s:mergeCells/s:mergeCell', N)}
            assert m['range'] in authored_merges, (m['range'], 'artifact merge missing')
            mc = dst.find('s:mergeCells', N)
            assert mc is not None, (m['sheet'], 'original merge container missing')
            assert m['range'] not in {x.get('ref') for x in mc}
            E.SubElement(mc, '{'+NS+'}mergeCell', ref=m['range'])
            mc.set('count', str(len(mc)))
        for h in (p for p in plan.get('heights', []) if p['key'] == b['key']):
            dst = root(h['sheet'])
            row = dst.find(f"s:sheetData/s:row[@r='{h['row']}']", N)
            row.set('ht', str(h['height'])); row.set('customHeight', '1')
        for sh in (p for p in plan['shades'] if p['key'] == b['key']):
            dst = root(sh['sheet'])
            rows = {int(r.get('r')): r for r in dst.findall('s:sheetData/s:row', N)}
            for ri in range(sh['start'], sh['end']+1):
                row = rows[ri]
                cells = {c.get('r'): c for c in row.findall('s:c', N)}
                for ci in range(sh['firstcol'], sh['lastcol']+1):
                    addr = chr(64+ci)+str(ri)
                    c = cells.get(addr)
                    if c is None:
                        c = E.SubElement(row, '{'+NS+'}c', r=addr)
                    oldid = int(c.get('s', '0'))
                    if oldid not in stylemap:
                        xf = copy.deepcopy(xfs[oldid]); xf.set('fillId',str(fill_id)); xf.set('applyFill','1')
                        stylemap[oldid] = len(xfs); xfs.append(xf)
                    c.set('s',str(stylemap[oldid]))
                row[:] = sorted(row, key=lambda c: (len(re.sub(r'\d','',c.get('r'))), c.get('r')))
        xfs.set('count',str(len(xfs)))
        old['xl/styles.xml'] = serialize(styles)
        for sn, r in roots.items():
            old[sheet_path(old, sn)] = serialize(r)
        out = QA/(b['key']+'-candidate.xlsx')
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
            for name, value in old.items():
                z.writestr(name, value)
        print('candidate', b['key'])


def verify():
    base = json.loads((QA/'baseline.json').read_text(encoding='utf-8'))
    plan = json.loads((QA/'plan.json').read_text(encoding='utf-8'))
    results = []
    for b in base['books']:
        candidate = QA/(b['key']+'-candidate.xlsx')
        old = load_workbook(b['backup'], rich_text=True)
        new = load_workbook(candidate, rich_text=True)
        criteria = load_workbook(b['criterionPath'], rich_text=True)
        assert old.sheetnames == new.sheetnames
        changes = {(p['sheet'],p['cell']):p for p in plan['changes'] if p['key']==b['key']}
        evenrows = {r for sh in plan['shades'] if sh['key']==b['key']
                    for r in range(sh['start'],sh['end']+1)}
        for so, sn in zip(old, new):
            assert str(so.merged_cells) == str(sn.merged_cells)
            assert so.max_row == sn.max_row and so.max_column == sn.max_column
            assert len(so._images) == len(sn._images)
            assert [(r,d.height,d.hidden) for r,d in so.row_dimensions.items()] == [
                (r,d.height,d.hidden) for r,d in sn.row_dimensions.items()]
            for c in list(so._cells.values()):
                    d = sn[c.coordinate]
                    p = changes.get((so.title,c.coordinate))
                    assert text(d.value) == (p['value'] if p else text(c.value)), (
                        b['key'],so.title,c.coordinate,'unapproved value')
                    # MergedCell objects do not expose the native interior fill consistently.
                    if d.__class__.__name__ != 'MergedCell' and so.title == b['sheet'] and c.row in evenrows and 2<=c.column<=b['verdictcol']:
                        assert d.fill.fgColor.rgb == 'FFF5F5F5', (b['key'],d.coordinate,'even fill')
                    sa = list(c._style or [0]*9); sb = list(d._style or [0]*9)
                    # 셀 스타일의 fillId만 변경한다. 나머지 font/border/number/alignment 등을 비교한다.
                    sa[1] = sb[1] = 0
                    assert sa == sb, (b['key'],so.title,c.coordinate,'non-fill style')
        for row in range(b['first'],new[b['sheet']].max_row+1):
            for col in range(1,b['common']+1):
                assert text(new[b['sheet']].cell(row,col).value) == text(criteria[b['csheet']].cell(row,col).value), (
                    b['key'],row,col,'criterion common content')
        native = parts(candidate); before = parts(b['backup'])
        permitted = {'xl/styles.xml'} | {sheet_path(native,p[0]) for p in changes}
        assert set(before) == set(native)
        assert all(before[n] == native[n] for n in before if n not in permitted)
        red = blue = 0
        native_roots = {sn:parse(native[sheet_path(native,sn)]) for sn,_ in changes}
        for (sn,addr),p in changes.items():
            c = native_roots[sn].find(f".//s:c[@r='{addr}']",N)
            runs = c.findall('s:is/s:r',N)
            for run in runs:
                value = ''.join(run.find('s:t',N).itertext())
                if PHOTO in value or QCC in value:
                    color = run.find('s:rPr/s:color',N).get('rgb')
                    assert run.find('s:rPr/s:b',N).get('val') == '1'
                    assert color == 'FF'+(BLUE if PHOTO in value else RED), (sn,addr,value,color)
                    red += QCC in value; blue += PHOTO in value
        results.append(dict(key=b['key'], tests=plan['counts'][b['key']]['tests'],
                            evenTests=plan['counts'][b['key']]['evenTests'],
                            changes=len(changes), blueRuns=blue, redRuns=red,
                            unchangedNativeParts=len(before)-len(permitted)))
        old.close(); new.close(); criteria.close()
        assert digest(b['criterionPath']) == b['criterionHash']
    for p, sha in base['protected'].items():
        assert digest(p) == sha, ('protected modified',p)
    out = dict(books=results, protectedUnchanged=len(base['protected']),
               preservedUnitIssues=plan['counts']['retainedUnitIssues'],
               boardTestsExecuted=False)
    save_json(QA/'verification.json',out)
    print(json.dumps(out,ensure_ascii=False))


def publish():
    verify()
    base = json.loads((QA/'baseline.json').read_text(encoding='utf-8'))
    # All targets must still match their snapshots before replacing any file.
    for b in base['books']:
        assert digest(b['path']) == b['sha256'], ('external edit',b['path'])
        with open(b['path'],'r+b'):
            pass
    saved = []
    for b in base['books']:
        target = Path(b['path'])
        stage = target.with_name('.result-format-'+target.name)
        assert not stage.exists()
        try:
            shutil.copyfile(QA/(b['key']+'-candidate.xlsx'),stage)
            os.replace(stage,target)
        finally:
            if stage.exists():
                stage.unlink()
        saved.append(dict(key=b['key'],path=str(target),before=b['sha256'],after=digest(target)))
        save_json(QA/'publication.json',dict(files=saved,boardTestsExecuted=False))
        print('published',b['key'])


if __name__ == '__main__':
    dict(prepare=prepare, compact=compact, polish=polish, build=build, verify=verify, publish=publish)[sys.argv[1]]()
