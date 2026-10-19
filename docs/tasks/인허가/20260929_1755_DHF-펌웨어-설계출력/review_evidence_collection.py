"""현재 시험 문서 읽기, 편집 입력 준비와 저장본 보존 검사를 수행한다.

openpyxl은 읽기/검증에만 사용한다. 셀 작성은 Artifact Tool에서 수행한다.
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
QA = TASK / "qa/evidence-collection"
DHF = ROOT / "docs/참고/인허가/사운드1/DHF"
OUT = DHF / "03_설계 출력/SW/펌웨어"
VER = DHF / "04_설계 검증/SW/펌웨어"
CONFIG = [
    ("unit-criterion", OUT / "SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx", "기준서_유닛 검증", 4, 11),
    ("unit-report", VER / "SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서.xlsx", "보고서_유닛 시험", 4, 11),
    ("integration-criterion", OUT / "SITC-FW-TD24A 소프트웨어 통합시험 기준서.xlsx", "기준서_통합 시험", 3, 13),
    ("integration-report", VER / "SITR-FW-TD24A 소프트웨어 통합시험 보고서.xlsx", "보고서_통합시험", 3, 13),
    ("system-input", DHF / "02_설계 입력/SW/펌웨어/[작성자 확인] SSTC-FW-TD24A 소프트웨어 시스템 시험 기준서.xlsx", "기준서_시스템 시험", 2, 10),
    ("system-criterion", OUT / "[작성자 확인] SSTC-FW-TD24A 소프트웨어 시스템 시험 기준서.xlsx", "기준서_시스템 시험", 2, 10),
    ("system-report", VER / "SSTR-FW-TD24A 소프트웨어 시스템 시험 보고서.xlsx", "보고서_시스템 시험", 2, 10),
]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(value: object) -> str:
    return "" if value is None else str(value)


def snapshot() -> None:
    QA.mkdir(parents=True, exist_ok=True)
    records = []
    for key, path, sheet_name, id_col, common in CONFIG:
        book = load_workbook(path, rich_text=True)
        assert sheet_name in book.sheetnames, (path, book.sheetnames)
        sheet = book[sheet_name]
        tests = []
        for row in sheet:
            identifier = text(row[id_col - 1].value)
            if re.fullmatch(r"SW-FW-(?:UnitTest-[A-Z0-9]+|IntegrationTest|SYS-Test)-ID-\d+", identifier):
                tests.append({"row": row[0].row, "id": identifier,
                              "cells": {c.coordinate: text(c.value) for c in row}})
        backup = QA / f"{key}-before.xlsx"
        if not backup.exists():
            shutil.copy2(path, backup)
        records.append({"key": key, "path": str(path), "sha256": digest(path),
                        "backup": str(backup), "sheet": sheet_name,
                        "common": common, "idCol": id_col,
                        "sheets": book.sheetnames, "maxRow": sheet.max_row,
                        "maxCol": sheet.max_column, "tests": tests})
        book.close()
    protected = [p for p in DHF.rglob("*") if p.is_file() and not p.name.startswith("~$")
                 and (p.name.startswith("[작성자 전용]") or "필수성능시험 기준서" in p.name)]
    source = [p for p in (ROOT / "src").rglob("*") if p.is_file() and p.suffix in {".c", ".h", ".ld", ".hct"}]
    baseline = {"books": records, "protected": {str(p): digest(p) for p in protected},
                "source": {str(p): digest(p) for p in source}}
    baseline_path = QA / "baseline.json"
    if baseline_path.exists():
        old = json.loads(baseline_path.read_text(encoding="utf-8"))
        assert all(digest(Path(p["path"])) == p["sha256"] for p in old["books"]), "Baseline changed"
    else:
        baseline_path.write_text(json.dumps(baseline, ensure_ascii=False, indent=2), encoding="utf-8")
    for item in records:
        print(item["key"], len(item["tests"]), item["sheets"], item["maxRow"], item["maxCol"])


def show(key: str, pattern: str = ".*") -> None:
    data = json.loads((QA / "baseline.json").read_text(encoding="utf-8"))
    item = next(p for p in data["books"] if p["key"] == key)
    for row in item["tests"]:
        if re.search(pattern, row["id"]):
            print(json.dumps(row, ensure_ascii=False))


if __name__ == "__main__":
    if len(sys.argv) == 1 or sys.argv[1] == "snapshot":
        snapshot()
    else:
        show(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else ".*")
