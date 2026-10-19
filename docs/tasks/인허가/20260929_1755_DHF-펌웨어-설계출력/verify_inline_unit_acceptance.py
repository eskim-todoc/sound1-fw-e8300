"""Verify inline criteria/results against CURRENT preserved source records."""
import json
import re
import sys
import zipfile
from pathlib import Path
import openpyxl
from inline_unit_acceptance import QA, BOOKS, parts, sheet_path, sha, unit_id, save_json, chunks, clean_navigation, text

def compact(s): return re.sub(r'\s+','',text(s))

def verify():
    plan=json.loads((QA/'plan.json').read_text('utf-8'))
    wc=openpyxl.load_workbook(QA/'criterion-candidate.xlsx');wr=openpyxl.load_workbook(QA/'report-candidate.xlsx')
    c=wc[BOOKS['criterion'][1]];r=wr[BOOKS['report'][1]]
    assert c.max_row==r.max_row
    expected_merges={str(x) for x in c.merged_cells.ranges}
    actual_merges={str(x) for x in r.merged_cells.ranges if x.min_col<=11}
    assert expected_merges==actual_merges
    common=0;log_lines=0;vectors=0;observations=0;cells=0
    for rec in plan['records']:
        start,end=rec['start'],rec['end']
        assert c.cell(start,4).value==r.cell(start,4).value==rec['testId']
        assert c.cell(start,7).value==rec['detailId']
        assert r.cell(start,13).value==rec['verdict']
        # Continuation labels are navigation, not payload. A paragraph may span
        # two rows without having that repeated label injected into its value.
        expected='\n'.join(text(c.cell(n,11).value).partition('\n')[2] for n in range(start,end+1))
        actual='\n'.join(text(r.cell(n,12).value).partition('\n')[2] for n in range(start,end+1))
        for n in range(start,end+1):
            for cn in range(1,12):
                assert c.cell(n,cn).value==r.cell(n,cn).value,(n,cn)
                common+=1
            assert c.row_dimensions[n].height==r.row_dimensions[n].height
            for cn in (11,12):
                target=c if cn==11 else r
                assert len(text(target.cell(n,cn).value))<32767
                assert target.cell(n,cn).alignment.wrap_text
                cells+=1
        for _,case,condition,log,result,note in rec['sourceLogs']:
            assert compact(condition) in compact(expected),(rec['id'],case,'input missing')
            for line in log.splitlines():
                assert line in expected,(rec['id'],case,line)
                log_lines+=1
            assert compact(result) in compact(actual),(rec['id'],case,'record missing')
        for heading,value in rec['inputDefinitions']:
            assert compact(value) in compact(expected),(rec['id'],heading,'vector missing')
            vectors+=1
        for _,heading,value in rec['sourceObservations']:
            assert compact(value) in compact(expected),(rec['id'],heading,'observation missing')
            observations+=1
        # Preserve existing warnings, photo/QCC flags and numbered data. Only
        # obsolete navigation to a separate RTT tab may be removed.
        old=clean_navigation(rec['oldResult'])
        for line in old.splitlines():
            line=re.sub(r'^\[\d+\] ','',line)
            assert compact(line) in compact(actual),(rec['id'],'old result line missing',line)
    ids=[unit_id(row[3]) for row in c.values if len(row)>3 and unit_id(row[3])]
    assert len(ids)==len(set(ids))==115
    # Native, unrelated worksheets/objects are byte-for-byte preserved.
    preservation=[]
    for key in ['criterion','report']:
        before=parts(QA/f'{key}-before.xlsx');after=parts(QA/f'{key}-candidate.xlsx')
        changed={sheet_path(before,BOOKS[key][1]),'xl/styles.xml','xl/workbook.xml'}
        if key=='criterion':changed.add(sheet_path(before,'기준서_유닛 검증_변경 비교'))
        for part in before:
            if part not in changed:
                assert after[part]==before[part],(key,'native part changed',part)
        assert set(before)==set(after)
        preservation.append({'key':key,'unchangedParts':len(before)-len(changed)})
    diff=wc['기준서_유닛 검증_변경 비교']
    before_diff=openpyxl.load_workbook(QA/'criterion-before.xlsx')['기준서_유닛 검증_변경 비교']
    assert [row for row in before_diff.values if row[0]=='[-]']==[row for row in diff.values if row[0]=='[-]']
    plus={}
    for row in diff.values:
        if row[0]=='[+]' and unit_id(row[3]):plus[row[3]]=row[10]
    assert len(plus)==115
    for rec in plan['records']:assert plus[rec['testId']]==c.cell(rec['start'],11).value
    # Central QCC anchors are checked against the newly saved row structure.
    qcc=(Path(plan['baseline']['books'][0]['path']).parents[4]/'QCC 자료 추가 위치 관리.md').read_text('utf-8')
    qcc_units=[]
    for rec in plan['records']:
        flag_rows=[x['row'] for x in rec['rows'] if '(QCC검증결과필요)' in compact(x['actual'])]
        if flag_rows:
            line=next(line for line in qcc.splitlines() if line.startswith('| '+rec['id']+' |'))
            assert 'D'+str(rec['start']) in line
            for row in flag_rows:
                assert 'L'+str(row) in line
                assert r.cell(row,12).font.color.rgb=='FFC00000'
            qcc_units.append(rec['id'])
    assert len(qcc_units)==18
    protected=plan['baseline']['protected']
    assert all(sha(Path(p))==value for p,value in protected.items())
    outcome={'units':115,'commonCells':common,'mainRows':c.max_row,
        'exactExpectedRttLines':log_lines,'inlineInputDefinitions':vectors,
        'inlineObservationDefinitions':observations,'conditionRows':plan['counts'],
        'preservedParts':preservation,'protectedFiles':len(protected),
        'verdictsUnchanged':True,'mainCellLimitsPassed':True,
        'historicalNegativeRowsPreserved':True,'qccAnchorsChecked':len(qcc_units)}
    save_json('verification.json',outcome);print(json.dumps(outcome,ensure_ascii=False,indent=2))

if __name__=='__main__':verify()
