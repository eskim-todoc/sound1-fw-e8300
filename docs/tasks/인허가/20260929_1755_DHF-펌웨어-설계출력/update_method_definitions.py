"""기준서의 방법/설명을 준비하고 실제 저장본을 검증한다. Excel 작성은 Artifact Tool을 사용한다."""
from pathlib import Path
from copy import deepcopy, copy
from difflib import SequenceMatcher
import json, re, sys, hashlib, zipfile, shutil
import openpyxl
from lxml import etree as E
from inspect_methods import T, ROOT, D, paths
Q=T/'qa/method-definitions';Q.mkdir(parents=True,exist_ok=True)
NS={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
def tag(s):return '{'+NS['s']+'}'+s
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
LEGEND='\n'.join([
 '[I] Inspection : 제어 또는 시각적 확인',
 '[A] Analysis : 분석적인 증거에 기반한 확인',
 '[D] Demonstration : 정량적인 측정 없이 동작 특성 확인',
 '[T] Test : 정량적인 측정을 통한 정량적인 특성 확인'])
SYSTEM={i:'D' for i in range(1,82)}
for letter,ids in {'T':[1,2,3,4,5,6,11,12,47,55,58,67,68,72],
 'A':[9,18,44,45,57,60,77,78],
 'I':[7,20,21,22,27,29,33,53,61,62,63,64,65,74,75,81]}.items():
 for n in ids:SYSTEM[n]=letter
# Functional flags are D; quantities and exact digital arrays are T. Hardware setup is reviewed as A.
UNIT_A={'CFX':set(range(1,7))|{9,14,21,38,39,41,42},
 'CM3':{1,3,5,8,15,31,34,35,37}, 'BOOT':{1,2,3,5,6,8,11,13,16},'HEAR':set()}
UNIT_T={'CFX':{7,11,12,15,16,17,18,19,20,22,23,24,25,26,27,28,29,30,31,32,33,34,35,36,37,40},
 'CM3':{7,9,16,17,19,25,26,27,28,29,36,38,40,42,45},
 'BOOT':{4,9,14,15,17,20},'HEAR':{1,2}}
INTEGRATION={n:('T' if n in {3,4,5,6,9,11,12,13,15,17} else 'D') for n in range(1,18)}
INTEGRATION[14]='A'
OVERRIDES={
 'CFX-001':{7:'- CFX 실행 허용 플래그(is_enabled_CFX_iteration)를 0(금지)과 1(허용)로 나누어 검토한다.\n- 오디오 블록 수집 완료, HEAR 연산 완료 및 PCM 전송 준비 플래그를 각각 0(미완료)과 1(완료)로 설정한 조건을 독립적으로 검토한다.',
  10:'- 실행 허용 플래그가 0이면 정상 처리를 호출하지 않는다. 1이면 준비된 이벤트만 한 번 처리한다. 처리한 완료 플래그는 0으로 해제하고, 처리하지 않은 이벤트의 플래그는 유지한다.'},
 'CFX-009':{7:'- 오디오와 PCM 출력이 활성인 상태에서 저전력 진입 요청 플래그를 1(요청 있음)로 설정한 분기를 검토한다.'},
 'CFX-015':{7:'- I2S 입력 16표본을 100+i(i=0~15)로 설정한다. FIFO 수신 처리 허용 플래그는 1(처리 가능), 수신 버퍼 위치와 복사 횟수는 0으로 설정한다.\n- 버퍼 복사 완료 플래그는 0(미완료)에서 시작하고 복사, 완료 조회, 완료 해제, 출력 버퍼 초기화를 차례로 실행한다.',
  8:'현재 복사 코드에는 복사 완료 플래그를 1로 설정하는 처리가 없다. 상세 설계의 완료 조건과 비교하여 불일치를 기록한다. 시험에서 완료 플래그를 대신 설정하지 않는다.',
  10:'- 첫 수신 위치의 16표본은 100~115, 복사 후 수신 위치와 복사 횟수는 각각 1이다. 처리한 FIFO 수신 허용 플래그는 0으로 해제된다.\n- 버퍼 복사 완료 조회는 복사 후 1(완료), 완료 해제 후 0(미완료)이어야 한다. 출력 버퍼 초기화 후 16표본은 모두 0이다.'},
 'CFX-019':{7:'- 매 단계에서 16표본이 모두 65536인 새 입력 블록을 준비한다. 완만한 전환의 단계 수를 4로 설정한다. 시작(fade-in)과 중단(fade-out)의 첫 단계, 마지막 단계 및 종료 다음 호출을 각각 실행한다.'},
 'CFX-020':{7:'- AGC 출력 16표본을 i(i=0~15)로 설정한다. 포화 시험에서는 [-131073,-131072,0,131071,131072,나머지 0]을 별도로 입력한다. OD 청취는 1000Hz 순음과 무음을 각각 입력하여 보조 확인한다.',
 10:'- 첫 시험의 FIFO 출력은 [0,64,...,960]이다. 포화 시험의 앞 5개 출력은 [-8388608,-8388608,0,8388544,8388607], 나머지는 0이다. 원래 AGC 출력 배열은 바뀌지 않는다.\n- OD 청취에서는 순음과 무음이 구분되어야 한다. 청취 결과로 전류값이나 정확한 이득을 판정하지 않는다.'},
 'CFX-021':{7:'- I2S 수신 여부, 매핑 연결 여부, 크래들 입력 여부, 활성 마이크 수 및 좌/우 빔 방향을 각각 구분하여 입력 경로 선택 분기를 검토한다.',
 10:'- 매핑 연결 중 I2S 수신은 두 입력에 게인 1.0을 적용한다. 일반 I2S 수신은 I2S에 게인 1.0, 마이크에 Gain A를 적용한다. 크래들 입력은 Gain A/B를 적용한다.\n- I2S 미수신에서는 마이크 한 개이면 단일 입력, 두 개이면 선택 방향의 빔포밍 경로를 사용한다. 해당 조건의 경로만 실행한다.'},
 'CFX-023':{7:'- 오디오 볼륨 1과 10(내부 인덱스 0과 9)을 각각 설정하고 이전 이득은 0으로 초기화한다.\n- 로그 변환 직후 입력 d를 [-7700297,-7700296,-7700295,-6209353,-6209352,-6209351](Q8.16)로 차례로 고정하여 시간 변화 적용 전 목표 이득 g를 기록한다.'},
 'CFX-032':{7:'- 디지털 T 레벨은 50, 변경 후 C 레벨은 150, 오프셋은 0으로 설정하고 자극 알림을 끈다. 상하한 제한 직전 계산값 y를 [40,50,100,160]으로 각각 고정한다.'},
 'CFX-036':{7:'- PCM 자극 헤더의 각 비트 필드를 모두 0으로 초기화한다. 이어서 한 필드씩 1로 설정하여 해당 비트 위치만 변경되는지 비교한다.'},
 'CFX-039':{7:'- 16표본 오디오 블록과 CIS/NofM 처리 자료를 준비한다. HEAR 연산 완료 플래그를 0(미완료)과 1(완료)로 나누어 결과 사용 순서를 코드에서 검토한다.',
 10:'- HEAR 연산 완료 전에는 연산 결과를 사용하지 않는다. 완료 후 AGC, FFT, 대역 대표값, 매핑 및 선택한 자극 전략 순서로 처리한다. 패킷 생성 함수와 전달 인자는 해당 전략의 설계 조건과 같아야 한다.'},
 'CFX-042':{7:'- 내부기 역방향 통신 요청 플래그를 1(요청 있음)로 설정한 조건에서 시작, 채널별 전송, 종료 및 송신 완료 상태의 처리 순서를 코드에서 검토한다.'},
 'CM3-004':{7:'- 공유 메모리 주소와 자료가 준비되지 않은 상태에서는 CFX 실행 허용 플래그(is_enabled_CFX_iteration)를 0(금지)으로 둔다. 준비가 끝난 상태에서는 1(허용)로 설정하고 CFX의 조회 결과를 RTT로 확인한다.',
 10:'- 공유 자료 준비 전에는 실행 허용 플래그가 0이다. 준비 완료 후에는 1이며, 실행 금지 요청 후 다시 0이다. 공유 주소와 자료를 설정하기 전에 실행을 허용하지 않는다.'},
 'CM3-007':{7:'- 정상 사용 상태에서 터치나 저배터리 종료 조건을 제외한다. 내부기 미연결 상태의 누적 카운터를 호출 전 179999와 180000으로 각각 설정한다. 내부기 연결 완료 상태도 별도로 시험한다.',
 10:'- 미연결 상태의 179999는 한 번 처리 후 180000이 되며 자동 종료를 요청하지 않는다. 180000은 180001에서 종료를 요청하고 카운터를 0으로 초기화한다. 내부기 사용 허용 상태(enable_ISD)는 false(사용 중지)로 바뀐다.\n- 내부기 연결 완료 상태에서는 카운터가 0으로 유지된다. 경계값 주입은 종료 분기 시험이며 실제 3분 경과 측정을 대신하지 않는다.'},
 'CM3-009':{7:'- 터치 진입 시간을 0ms로 두고 현재 시간을 2399/2400/2401ms로 각각 설정한다.\n- 부팅 직후 터치 무시 상태에서 연속 미누름 1/2/3회, 센서 읽기 실패 1~20/21회, 지속 누름 30/60/90초 조건을 각각 준비한다. 폴링 주기는 100ms, 읽기 실패 상태 유지 시간은 2000ms, 부팅 후 터치 무시 해제 시간은 300ms다.',
 10:'- 길게 누름 이벤트는 2399ms에 false, 2400ms에 true, 2401ms에 false로 한 번 발생한다. 손을 뗀 후 다시 누르면 재판정한다.\n- 부팅 직후 터치 무시 상태는 연속 미누름 3회에서 해제한다. 읽기 실패 1~20회에는 이전 누름 상태(HOLD)를 유지하고 21회에는 미누름(NOT_TOUCH)으로 바뀐다.\n- 지속 누름 30/60/90초에는 각각 기준값 갱신(RESEED), 센서 재보정(RE_ATI), 센서 리셋(MCLR)을 요청한다. 가상 시간 입력으로 분기를 시험하고 실제 재부팅은 차단한다.'},
 'CM3-012':{7:'- CFX에 전달할 PMIC 제어 요청을 켜짐과 꺼짐으로 각각 설정한다. 저전력 진입 요청 플래그는 0(요청 없음)에서 1(요청 있음)로 변경한다.'},
 'CM3-015':{7:'- CFX 주변장치 정리 완료 상태를 0(미완료)과 1(완료)로 나누어 CM3의 절전 진입 및 복귀 분기를 검토한다.',
 10:'- CFX 정리가 미완료이면 CM3는 저클럭 전환과 최종 절전 진입을 진행하지 않는다. 완료 상태에서 정리와 클럭 변경을 진행하며, 복귀 시 통신 설정과 정상 반복 처리를 복구한다.'},
 'CM3-016':{7:'- 오류 비트 하나를 설정한다. 공유 메모리 주소가 일치하거나 다른 조건과 CFX의 EEPROM 자료 적재 완료 플래그가 0(미적재) 또는 1(적재 완료)인 조건을 각각 준비한다.'},
 'CM3-020':{7:'- CM3 자극 파라미터 계산 완료 플래그를 0(미완료) 또는 1(완료)로 설정한다. 맵 변경 요청 플래그는 0(요청 없음) 또는 1(요청 있음)로 각각 설정한다. 계산 자료의 공유 위치도 조회한다.',
 10:'- 계산 완료, 맵 변경 요청 및 매핑 프로그램 연결 상태는 각각 설정한 상태로 전달된다. 계산 자료 포인터는 해당 공유 자료를 가리킨다. 자료를 쓰기 전에 계산 완료를 통보하지 않는다.'},
 'CM3-026':{10:'- 응답 바이트는 [0x12,0x34] 순서다. 응답 준비 신호가 High이면 버퍼 비어 있음 조회는 false다. 준비 신호 해제 후 Low이며 조회는 true다.\n- 버퍼 초기화 후 모든 데이터는 0이다. 데이터 지우기와 응답 준비 신호 해제는 별도 기능이다.'},
 'CM3-028':{7:'- 내부기 DAC의 자극 전류 상한 A/B/C/D를 400/803/1203/1606uA, 오프셋 상한 A/B를 295/592uA로 고정한다. A 범위 환산계수는 20871/14435(Q1.15), 펄스폭은 25us, 전극 수는 32로 설정한다. 전 채널 T/C 입력 전류는 100/300uA로 설정한다.',
 10:'- 자극 DAC와 오프셋 DAC는 각각 A 범위를 선택한다. 오프셋 전류는 100uA, 오프셋 레지스터 값은 44다. 32채널의 디지털 T/C 값은 각각 0/127이고 계산은 true(성공)를 반환한다.\n- 지원 상한을 넘는 입력은 실패로 반환한다. 이 시험은 변환 계산값을 확인하며 실제 전극 전류를 측정하는 시험은 아니다.'},
 'CM3-029':{7:'- 펄스폭 25us, 전극 수 32, 전 채널 T/C 입력 전류 100/300uA를 설정한다. DAC 자극 상한 A/B/C/D=400/803/1203/1606uA, 오프셋 상한 A/B=295/592uA, A 환산계수=20871/14435(Q1.15)를 적용한다. CIS와 NofM을 각각 선택한다.',
 10:'- NofM 32대역은 자극당 3프레임, 1ms당 8채널이다. CIS는 자극당 2프레임, 1ms당 12채널이다. 공유한 32채널 디지털 T/C는 각각 0/127이다. 계산 실패를 정상 공유 완료로 처리하지 않는다.'},
 'CM3-036':{7:'- FPGA의 시험 레지스터 응답은 0(오류 없음)으로 준비한다. 이후 오류 표시 비트와 FIFO 비어 있음 표시 비트를 한 개씩 1로 설정한 응답을 각각 준비한다. FIFO 읽기 응답은 [0x12,0x34]다.'},
 'CM3-039':{7:'- 자극 전원(VP) 활성 요청, 비활성 요청 및 출력 차단 리셋 요청을 각각 준비한다. 전원 제어 명령의 기록은 시험 응답으로 확인한다.',
 9:'- 각 요청을 실행한 뒤 PMIC에 전달한 활성/차단 명령과 제어 상태를 RTT로 기록한다.\n- 전원부 동작 확인이 필요한 경우 안전한 시험 부하에서 VP 출력이 나타나는지, 차단 요청 후 사라지는지를 추가 확인한다.',
 10:'- VP 활성 요청은 출력 허용 설정을, 비활성 및 차단 리셋은 출력 차단 설정을 전달한다. 각 요청 후 제어 상태가 해당 명령과 같아야 한다. 추가 관찰에서는 활성/차단에 따른 출력 유무를 확인한다. 전압 정확도와 전극 전류는 이 시험으로 판정하지 않는다.'},
 'CM3-046':{7:'- 일반 리모콘 명령 처리 허용 상태를 false(처리 금지)와 true(처리 허용)로 각각 설정한다. 조회와 설정 명령을 입력하고, 연결 확인 명령은 별도로 입력한다.',
 10:'- 일반 명령은 처리 허용 상태에서만 대응 기능을 실행하고 처리 후 명령을 대기 상태로 정리한다. 연결 확인 명령은 별도 경로로 처리한다. 이 허용 상태를 암호학적 패스키 인증 결과로 해석하지 않는다.'},
 'BOOT-013':{7:'- 헤더를 포함한 첫 CRC 계산값을 0(잔여값 정상)과 1(잔여값 오류)로 나누어 코드의 분기를 검토한다. 복호화 후 CRC도 0/1, 복호화된 헤더의 암호화 표시 비트도 0(표시 없음)/1(표시 있음)로 각각 검토한다.\n- 헤더 포인터 없음, AES 초기화/복호화 오류 및 데이터 길이 0/1/3/4/5 조건을 추가 검토한다.',
 10:'- 첫 CRC가 0이면 성공(반환 0)이다. 첫 CRC 오류 후 복호화한 CRC가 0이고 암호화 표시 비트가 1이면 성공이다. 복호화 후 CRC 오류 또는 암호화 표시가 없으면 -1이다.\n- 헤더 포인터가 없으면 -1, AES 오류는 해당 오류값을 반환한다. CRC는 4바이트 워드와 남은 바이트를 처리하며 헤더 포함 잔여값 0으로 판정한다.'},
 'BOOT-017':{7:'- 서비스 RTT 입력 채널에 "dump\\n"과 "unknown\\n"을 각각 보낸다. 시험 파일에는 [0x12,0x34]를 저장한다. 파일 열기/읽기 실패 조건도 독립적으로 준비한다.',
 10:'- dump 명령은 파일 시작/종료 표시 사이에 "12 34"를 출력한다. unknown 명령은 "invalid command."를 출력한다. 파일 접근 실패에는 오류 메시지를 출력하며 정상 자극 처리를 호출하지 않는다. 서비스 명령은 RTT 입력으로 시험한다.'},
 'BOOT-020':{7:'- 기본 기동 슬롯은 1, 대체 슬롯은 2, 기동 시도 상한은 3으로 설정한다. 정상 기동(BOOT), 대체 기동/시도 중(ALT_BOOT/BOOT_TRY), 시도 완료(TRY_DONE) 상태를 각각 준비한다.\n- 시도 중 상태의 호출 전 횟수 2와 3, 시도 완료의 성공/실패 및 지원하지 않는 하위 상태를 독립적으로 적용한다.',
 10:'- 정상 기동은 슬롯 1이다. 호출 전 시도 횟수가 2이면 새 횟수 3, 슬롯 2다. 횟수가 3이면 새 횟수 4, 슬롯 1로 돌아가며 BOOT/TRY_DONE/FAIL 상태를 기록한다.\n- 시도 완료 성공은 슬롯 2, 실패 또는 미지원 상태는 슬롯 1과 UNKNOWN 상태다. 슬롯 번호의 유효 범위 확인은 해당 분기를 별도로 분석한다.'},
}
def natural(s):
 for a,b in [('허가=0','CFX 실행 허용 플래그가 0(금지)'),('허가=1','CFX 실행 허용 플래그가 1(허용)'),('CFX 실행 CFX','CFX'),('정적 분석 시험이다.','코드 분석으로 판정한다.'),('정적 분석 시험','코드 분석 시험'),('독립 초기 상태','서로 독립된 초기 상태'),('계약','설계 조건'),('unity','게인 1.0'),('별도헤더','별도 헤더'),('RTT down채널','RTT 입력 채널'),('부트 상태','기동 상태')]:s=s.replace(a,b)
 # Korean spacing around compact numeric clauses; keep identifiers, formulas and vectors intact.
 s=re.sub(r'(?<=[가-힣])(?=(?:\d|0x))',' ',s)
 s=re.sub(r'(?<=[0-9])(?=[가-힣])',' ',s)
 for a,b in [('는잡음','는 잡음'),('는감쇠','는 감쇠'),('는증폭','는 증폭'),('다음호출','다음 호출'),('후재누름','후 재누름'),('뒤0','뒤 0'),('구획 적재 호출은없다','구획 적재 호출은 없다'),('반환길이','반환 길이'),('원본은불변','원본은 불변'),('으로판정','으로 판정'),('직접비교','직접 비교'),('않는다.','않는다.')]:s=s.replace(a,b)
 return s.replace('·','/').replace('정략적인','정량적인')
def prepare():
 previous=json.loads((Q/'plan.json').read_text(encoding='utf8')) if (Q/'plan.json').exists() else None
 if previous:
  assert all(sha(Path(b['path']))==b['sha'] for b in previous['books']),'Baseline changed'
 protected={str(p):sha(p) for p in D.rglob('*') if p.is_file() and p.name.startswith('[작성자 전용]')}
 sources={str(p):sha(p) for p in (ROOT/'src').rglob('*') if p.is_file() and p.suffix.lower() in {'.c','.h','.hct'}}
 plan={'legend':LEGEND,'protected':protected,'sources':sources,'books':[]}
 for key,path in paths.items():
  assert '[작성자 전용]' not in path.name
  d=Q/key;d.mkdir(exist_ok=True)
  if (d/'before.xlsx').exists():assert sha(d/'before.xlsx')==sha(path)
  else:shutil.copy2(path,d/'before.xlsx')
  w=openpyxl.load_workbook(path,rich_text=True);s=w.worksheets[1]
  cols={'input':10,'system':10,'unit':11,'integration':13}[key]
  idcol={'input':2,'system':2,'unit':4,'integration':3}[key]
  methodcol={'input':4,'system':4,'unit':6,'integration':9}[key]
  rows=[];changes=[]
  for r in range(1,s.max_row+1):
   vals=[str(s.cell(r,c).value) if s.cell(r,c).value is not None else '' for c in range(1,cols+1)]
   if not vals[idcol-1].startswith('SW-FW-'):continue
   before=list(vals);identifier=vals[idcol-1];n=int(identifier.split('-')[-1])
   if key in ('input','system'):
    vals[methodcol-1]=SYSTEM[n]
    if n==57 and key=='system':vals[8]='필수성능시험 보고서의 각 측정 결과와 판정이 해당 요구사항 및 필수성능시험 기준서와 일치하는지 대조한다.'
    if n==58 and key=='system':vals[8]='신호처리 시작/종료 시각을 시험용 타이머로 기록하고 처리 시간을 비교한다. 결과는 처리 후 RTT로 출력한다. 필요한 경우 Calibration 핀의 처리 구간을 오실로스코프로 측정한다.'
    if n in (9,77,78) and not vals[8]:vals[8]='해당 기능의 소스 코드와 설계 조건을 분석하고 통과 기준에 맞는 구현인지 대조한다.'
    if n==45 and key=='input':vals[8]='매핑 명령 송수신 및 제어 시험의 기록을 검토하여 블루투스 연계가 요구사항을 충족하는지 확인한다.'
   elif key=='unit':
    core=vals[1];uid=f'{core}-{n:03}'
    vals[5]='A' if n in UNIT_A[core] else ('T' if n in UNIT_T[core] else 'D')
    if uid=='CM3-014':vals[5]='I'
    for c in (7,8,9,10):vals[c]=natural(vals[c])
    vals[9]=re.sub(r'\n- 대상:.*?(?=\n- 조건별|$)','',vals[9],flags=re.S)
    if vals[5]=='A':
     vals[8]='코드 분석으로 판정한다. 실제 핀 파형, 전원, 처리 시간 및 전극 전류의 실측과 구분한다.'
     vals[9]='- 준비한 각 입력과 초기 상태에 대해 조건 분기, 호출 순서, 인자, 반환값 및 상태 변경을 코드에서 확인한다. 아래 예상 결과와 대조하고 검토한 조건 및 근거를 기록한다.'
    else:
     obs='입력/출력 배열, 계산값과 횟수' if vals[5]=='T' else '상태 변화, 반환 결과와 처리 순서'
     vals[9]=f'- 준비한 조건을 각각 독립 초기화하여 해당 유닛을 실행한다. {obs}를 처리 완료 후 RTT로 기록하여 아래 예상 결과와 비교한다.'
     vals[8]='시험 입력 또는 가상 장치 응답은 유닛 입력 경계에 설정한다. 실행 결과를 미리 예상값으로 덮어쓰지 않는다. 실제 전원이나 RF 연결 확인과 소프트웨어 상태 확인을 구분한다.'
     if core in ('CFX','HEAR'):
      vals[8]+=' CFX/HEAR 결과는 별도 버퍼에 보존한 뒤 CM3 RTT로 읽고 ISR에서 긴 로그를 출력하지 않는다.'
     if core=='BOOT':vals[8]+=' 운영 이미지 대신 격리된 시험 저장소를 사용한다. 초기 RTT 미사용 시 결과를 보존한 뒤 읽는다.'
    for c,v in OVERRIDES.get(uid,{}).items():vals[c]=v
    if uid=='CM3-014':vals[9]='- 각 LED 표시 상태를 요청하여 색상, 점등/점멸 및 소등 여부를 눈으로 확인한다. 선택한 패턴과 종료 상태는 RTT로 함께 확인한다.'
    if uid=='BOOT-019':vals[9]='- 시험 상태 파일을 읽은 뒤 반환 결과와 선택 슬롯을 보존하여 RTT로 확인한다. UART 바이트 파형 측정은 기본 판정에 요구하지 않는다.'
    if vals[5]=='A':vals[7]=vals[7].replace('실행한다.','실행하는 조건을 검토한다.').replace('호출한다.','호출하는 조건을 검토한다.')
   else:
    vals[8]=INTEGRATION[n]
    for c in (9,10,11,12):vals[c]=natural(vals[c])
    if n==14:
     vals[9]='- QCC 동작 모드를 정상/종료로, 내부기 설정 완료 상태를 false(미완료)/true(완료)로 나누어 E8300의 출력 설정 분기를 검토한다.'
     vals[10]='코드와 GPIO 출력 설정값을 대조한다. QCC 내부 동작이나 실제 핀 전압을 판정하는 시험은 아니다.'
     vals[11]='- 초기화와 각 상태 변경 코드에서 QCC_CTRL 및 ISD_CHECK의 출력 레벨과 적용 순서를 확인한다.\n- 내부기 설정 완료 전에 연결 완료를 통보하지 않는지 대조한다.'
     vals[12]='- 초기화에서는 두 출력 설정이 Low다. QCC_CTRL은 정상 모드에서 High, 종료 모드에서 Low다. ISD_CHECK는 내부기 설정 완료에서 High, 미완료에서 Low다.\n- 각 상태와 출력 설정의 대응 및 통보 순서가 일치해야 한다.'
    if n==3:
     vals[10]=vals[10].replace('완료=1 설정문','버퍼 복사 완료 플래그를 1(완료)로 설정하는 처리')
    if n in (1,7):vals[9]=vals[9].replace('공유 실행 CFX 실행 허용','CFX 실행 허용')
    if n==1:
     vals[9]='- 정상 시험 이미지를 준비한다. CFX 실행 허용 플래그(is_enabled_CFX_iteration)는 0(금지)으로 시작하고 기동 준비 상태와 초기화 순서를 기록한다.'
     vals[11]='- CFX 기동과 대기, CM3의 공유 자료 및 주변장치 준비, CFX 실행 허용 플래그를 1(허용)로 설정, CFX 정상 초기화 순서를 RTT로 확인한다.\n- 마이크, HEAR 체인 및 PCM/I2S 설정이 첫 정상 처리보다 먼저 이루어지는지 확인한다.'
     vals[12]='- CFX 실행 허용 플래그가 0이면 정상 오디오/자극 처리를 시작하지 않는다. CM3의 준비 완료 후 1로 바뀌며, 첫 처리 전에 CFX 초기화가 끝난다.\n- HEAR 명령 0/1과 마이크/PCM/I2S 버퍼가 각 처리 경로에 연결된다. 이미지 적재 전에 응용 코어를 실행하지 않는다.'
    if n==3:
     vals[12]='- 두 번째 블록 복사 후 현재 자료는 B, 직전 자료는 A이며 각 16표본이 다른 마이크의 자료와 섞이지 않는다.\n- I2S 수신 중에는 단일 마이크와 I2S 입력 경로를 사용하고, 중단 후 두 마이크 입력 경로로 돌아간다.\n- I2S 버퍼 복사 완료 플래그의 조회값은 복사 후 1(완료), 해제 후 0(미완료)이어야 한다. 실제 복사와 조회 경로로 확인한다.'
    if n==6:
     vals[12]='- 맵 변경 요청 시 재적재 완료 플래그는 0(미완료)이다. CFX 반영 후 완료 플래그는 1(완료), 변경 요청 플래그는 0(요청 없음)으로 바뀐다. 왼쪽 방향은 DMIC2, 오른쪽 방향은 DMIC1이 지연 입력이다.\n- 자극 볼륨 1/4의 C 상한은 각각 119/150이며 T는 50이다. 오디오 볼륨 계수는 선택 값에 맞게 갱신된다.\n- 같은 맵을 다시 선택하면 불필요한 변경 요청을 만들지 않는다.'
    if n==7:
     vals[9]='- CFX 실행 허용 플래그는 0(금지), 공유 오류 상태는 초기값으로 준비한다. CM3 초기화 전후 QCC 제어 출력 설정과 SPI 응답 버퍼 상태를 기록한다.'
     vals[12]='- QCC_CTRL과 ISD_CHECK의 초기 출력 설정은 모두 Low다. 터치 센서 초기화 중에는 정상 터치 입력을 처리하지 않는다.\n- 공유 주소와 자료가 준비된 뒤 CFX 실행 허용 플래그는 0(금지)에서 1(허용)로 바뀌며 CM3와 CFX가 같은 상태를 확인한다. SPI 송신/수신 버퍼를 구분하여 사용한다.'
    if n==11:
     vals[12]='- 미연결 카운터가 호출 전 179999이면 한 번 처리 후 180000이 되며 자동 종료 요청이 없다. 180000이면 180001에서 종료를 요청하고 카운터를 0으로 초기화한다. 내부기 사용 허용 상태(enable_ISD)는 false(사용 중지)로 바뀐다.\n- 내부기 연결 완료 상태에서는 카운터가 0이다. 종료 결정 후 새 실시간 자극 처리를 진행하지 않고 CFX 정리 후 CM3가 대기한다.'
    if n==10:
     vals[9]='- 일반 리모콘 명령 처리 허용 상태를 false(금지)와 true(허용)로 나누어 준비한다. 허용 상태에서는 맵 2, 오디오 볼륨 10, 자극 볼륨 4의 설정/조회 명령을 입력한다.'
     vals[12]=vals[12].replace('허용=true','일반 명령 처리 허용 상태').replace('허용=false','일반 명령 처리 금지 상태')
   rows.append({'row':r,'before':before,'after':vals})
   for c,(a,b) in enumerate(zip(before,vals),1):
    if a!=b:changes.append({'row':r,'col':c,'value':b})
  diff=[]
  for x in rows:
   if x['before']==x['after']:
    row=list(x['after']);row[0]='';diff.append(row)
   else:
    for mark,k in [('[-]','before'),('[+]','after')]:row=list(x[k]);row[0]=mark;diff.append(row)
  start=min(x['row'] for x in rows)
  olddiff=w.worksheets[2];oldlast=max(r for r in range(start,olddiff.max_row+1) if any(str(olddiff.cell(r,c).value or '').startswith('SW-FW-') for c in range(1,cols+1)))
  plan['books'].append({'key':key,'path':str(path),'sha':sha(path),'sheet':s.title,'comparison':olddiff.title,'cols':cols,'start':start,'last':max(x['row'] for x in rows),'old_diff_last':oldlast,'methodcol':methodcol,'idcol':idcol,'rows':rows,'changes':changes,'diff':diff})
 (Q/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8')
 print([(b['key'],len(b['rows']),len(b['changes'])) for b in plan['books']])
def rich():
 p=json.loads((Q/'plan.json').read_text(encoding='utf8'));checks=[]
 for b in p['books']:
  path=Q/b['key']/'candidate.xlsx'
  with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
  original=dict(parts);wb=E.fromstring(parts['xl/workbook.xml']);rels=E.fromstring(parts['xl/_rels/workbook.xml.rels'])
  links={x.get('Id'):x.get('Target') for x in rels};sh=next(x for x in wb.find(tag('sheets')) if x.get('name')==b['comparison']);target=links[sh.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
  name=target.lstrip('/') if target.startswith('/') else 'xl/'+target
  root=E.fromstring(parts[name]);strings=E.fromstring(parts['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in parts else []
  def value(c):
   if c.get('t')=='s':return ''.join(strings[int(c.findtext(tag('v')))].xpath('.//s:t/text()',namespaces=NS))
   return ''.join(c.xpath('.//s:t/text()',namespaces=NS)) if c.get('t')=='inlineStr' else c.findtext(tag('v')) or ''
  cells={c.get('r'):c for c in root.iter(tag('c'))};styles=E.fromstring(parts['xl/styles.xml']);fonts=styles.find(tag('fonts'));xfs=styles.find(tag('cellXfs'));count=0
  def apply(c,mask,mark):
   nonlocal count
   s=value(c)
   if not s:return
   font=fonts[int(xfs[int(c.get('s','0'))].get('fontId','0'))]
   for child in list(c):
    if child.tag in (tag('v'),tag('is')):c.remove(child)
   c.set('t','inlineStr');ins=E.SubElement(c,tag('is'));a=0
   while a<len(s):
    active=mask[a];e=a+1
    while e<len(s) and mask[e]==active:e+=1
    run=E.SubElement(ins,tag('r'));rp=E.SubElement(run,tag('rPr'))
    for prop in font:
     if prop.tag in (tag('b'),tag('color')):continue
     cp=deepcopy(prop)
     if cp.tag==tag('name'):cp.tag=tag('rFont')
     rp.append(cp)
    if active:E.SubElement(rp,tag('b'))
    E.SubElement(rp,tag('color'),rgb=('FF9C0006' if mark=='[-]' else 'FF006100') if active else 'FF333333')
    tt=E.SubElement(run,tag('t'));tt.text=s[a:e];tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');a=e
   assert value(c)==s;count+=1
  for i,row in enumerate(b['diff']):
   for col,intended in enumerate(row,1):
    addr=f'{openpyxl.utils.get_column_letter(col)}{b["start"]+i}'
    if not intended:continue
    c=cells[addr];actual=value(c)
    if actual!=intended:
     assert actual.strip()==intended.strip(),(b['key'],addr,actual,intended)
     for child in list(c):
      if child.tag in (tag('v'),tag('is')):c.remove(child)
     c.set('t','inlineStr');ins=E.SubElement(c,tag('is'));tt=E.SubElement(ins,tag('t'));tt.text=intended;tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
   mark=row[0]
   if mark not in ('[-]','[+]'):continue
   mate=b['diff'][i+1 if mark=='[-]' else i-1]
   for col,s in enumerate(row):
    s=str(s);other=str(mate[col]);mask=[False]*len(s)
    for op,a,e,_,_ in SequenceMatcher(None,s,other,autojunk=False).get_opcodes():
     if op!='equal':mask[a:e]=[True]*(e-a)
    if col==0:mask=[True]*len(s)
    addr=f'{openpyxl.utils.get_column_letter(col+1)}{b["start"]+i}'
    if addr in cells:apply(cells[addr],mask,mark)
  parts[name]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
  with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
   for n,data in parts.items():z.writestr(n,data)
  assert all(parts[n]==v for n,v in original.items() if n!=name)
  checks.append({'key':b['key'],'rich_cells':count,'only_part':name})
 (Q/'rich-checks.json').write_text(json.dumps(checks,indent=2),encoding='utf8');print(checks)
def verify(publish=False):
 p=json.loads((Q/'plan.json').read_text(encoding='utf8'));results=[];maps={}
 for b in p['books']:
  candidate=Q/b['key']/'candidate.xlsx';w=openpyxl.load_workbook(candidate,rich_text=True);s=w[b['sheet']];old=openpyxl.load_workbook(Q/b['key']/'before.xlsx',rich_text=True)
  assert w.sheetnames==old.sheetnames
  def val(v):return '' if v is None else str(v)
  for x in b['rows']:
   actual=[val(s.cell(x['row'],c).value) for c in range(1,b['cols']+1)]
   assert actual==x['after'],(b['key'],x['row'])
  assert val(s['C12'].value)==LEGEND
  assert val(w[b['comparison']]['C12'].value)==LEGEND
  allowed={(x['row'],x['col']) for x in b['changes']}|{(12,c) for c in range(3,b['cols']+1)}
  if b['key']=='unit':allowed|={(23,c) for c in range(4,7)}
  original=old[b['sheet']]
  for row in original:
   for cell in row:
    if (cell.row,cell.column) not in allowed:
     assert val(cell.value)==val(s.cell(cell.row,cell.column).value),(b['key'],cell.coordinate,'Unrequested value change')
  comparison=w[b['comparison']]
  for i,row in enumerate(b['diff']):
   actual=[val(comparison.cell(b['start']+i,c).value) for c in range(1,b['cols']+1)]
   assert actual==row,(b['key'],'comparison',i)
   if row[0] not in ('[-]','[+]'):continue
   mate=b['diff'][i+1 if row[0]=='[-]' else i-1]
   for col,txt in enumerate(row,1):
    if not txt:continue
    mask=[False]*len(txt)
    for op,a,z,_,_ in SequenceMatcher(None,txt,mate[col-1],autojunk=False).get_opcodes():
     if op!='equal':mask[a:z]=[True]*(z-a)
    if col==1:mask=[True]*len(txt)
    cell=comparison.cell(b['start']+i,col);runs=cell.value
    assert isinstance(runs,openpyxl.cell.rich_text.CellRichText),(b['key'],cell.coordinate)
    offset=0
    for run in runs:
     active=mask[offset] if offset<len(mask) else False
     assert all(x==active for x in mask[offset:offset+len(str(run))])
     expected=('FF9C0006' if row[0]=='[-]' else 'FF006100') if active else 'FF333333'
     assert run.font.color.rgb==expected and bool(run.font.b)==active,(b['key'],cell.coordinate,'Diff emphasis')
     offset+=len(str(run))
    assert offset==len(txt)
  for name in w.sheetnames:
   if name in (b['sheet'],b['comparison']):continue
   a=old[name];z=w[name]
   assert [[val(c.value) for c in row] for row in a]==[[val(c.value) for c in row] for row in z],name
   assert set(map(str,a.merged_cells.ranges))==set(map(str,z.merged_cells.ranges)),name
   for row in a:
    for cell in row:
     other=z.cell(cell.row,cell.column)
     assert all(copy(getattr(cell,k))==copy(getattr(other,k)) for k in ('font','fill','alignment','border')),(name,cell.coordinate,'Auxiliary style')
  assert not any(c.data_type=='e' for sheet in w for row in sheet for c in row),'Excel error value'
  maps[b['key']]={x['after'][b['idcol']-1]:x['after'][b['methodcol']-1] for x in b['rows']}
  counts={m:list(maps[b['key']].values()).count(m) for m in 'IADT'}
  results.append({'key':b['key'],'tests':len(b['rows']),'methods':counts,'candidate_sha':sha(candidate)})
 assert maps['input']==maps['system'],'Common test methods mismatch'
 assert set(maps['input'])=={f'SW-FW-SYS-Test-ID-{n}' for n in range(1,82)}
 for group in ('protected','sources'):
  assert all(sha(Path(path))==h for path,h in p[group].items()),group
 (Q/'checks.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf8');print(results)
 if publish:
  reviewed=json.loads((Q/'visual-reviewed.json').read_text(encoding='utf8'))
  for result in results:assert reviewed['hashes'][result['key']]==result['candidate_sha'],'Visual review baseline changed'
  for b in p['books']:
   dest=Path(b['path']);assert sha(dest)==b['sha'],'External edit: '+str(dest)
   with dest.open('r+b'):pass
  for b in p['books']:
   dest=Path(b['path']);shutil.copyfile(Q/b['key']/'candidate.xlsx',dest);assert sha(dest)==sha(Q/b['key']/'candidate.xlsx')
  (Q/'publication.json').write_text(json.dumps(results,indent=2),encoding='utf8')
if __name__=='__main__':
 {'prepare':prepare,'rich':rich,'verify':verify,'publish':lambda:verify(True)}[sys.argv[1]]()
