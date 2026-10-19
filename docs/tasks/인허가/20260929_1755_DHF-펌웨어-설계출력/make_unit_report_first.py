"""작성자 수정중 보고서로부터 보조 시트를 제외한 1차본을 만든다."""
from __future__ import annotations
import collections
import copy
import json
import posixpath
import re
import shutil
import sys
import zipfile
from pathlib import Path
from lxml import etree as E
from openpyxl import load_workbook
import revise_unit_report_wording as previous

TASK=Path(__file__).resolve().parent
ROOT=TASK.parents[3]
QA=TASK/'qa/unit-first-edition'
SOURCE=previous.REPORT
TARGET=SOURCE.with_name('SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서_1차본.xlsx')
DESIGN=previous.DESIGN
SHEET=previous.SHEET
NS=previous.NS
KEEP=['표지',SHEET]
DROP=['입력 및 예상값','RTT 결과 대조','관찰값 확인','시험 보조자료','예상 그림']
plain=previous.plain
digest=previous.digest

def save(name,data):
    QA.mkdir(parents=True,exist_ok=True)
    (QA/name).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

def read(name):
    return json.loads((QA/name).read_text(encoding='utf-8-sig'))

def paragraph_text(node):
    return ''.join(t.text or '' for t in node.findall('.//w:t',NS)
                  if not any(a.tag=='{'+NS['w']+'}del' for a in t.iterancestors()))

def descriptions():
    with zipfile.ZipFile(DESIGN) as z:root=E.fromstring(z.read('word/document.xml'))
    body=root.find('w:body',NS); result={}; current=None;names=previous.unit_names()
    for node in body:
        if node.tag!='{'+NS['w']+'}p':continue
        text=paragraph_text(node).strip()
        match=re.match(r'^(SU-ID-(?:CFX|HEAR|CM3|BOOT)-\d+)\s*[,，]',text)
        recovered=None
        if not match:
            # 작성자가 제목 안에서 줄을 나눈 3개 항목은 번호 끝자리/이름으로 찾는다.
            part=re.match(r'^(?:(?:CFX|HEAR|CM3|BOOT)-)?(\d{3}),\s*(.+)$',text)
            if part:
                candidates=[sid for sid,name in names.items() if sid.endswith('-'+part.group(1)) and name==part.group(2)]
                if len(candidates)==1:recovered=candidates[0]
        if match or recovered:
            current=match.group(1) if match else recovered
            result[current]=dict(title=text,intro='',style='')
            continue
        if current and text:
            # ID와 유닛 이름이 서로 다른 제목 문단인 경우 이어진 제목을 건너뛴다.
            style=node.find('w:pPr/w:pStyle',NS)
            style_id=None if style is None else style.get('{'+NS['w']+'}val')
            if text==names.get(current) or re.fullmatch(r'TDC[1-6]',style_id or ''):
                result[current]['title']+=' '+text
                continue
            result[current]['intro']=text
            result[current]['style']=style_id
            current=None
    return result

def inventory():
    book=load_workbook(SOURCE,rich_text=True);s=book[SHEET]
    desc=descriptions(); groups=[];current=None
    for row in range(30,s.max_row+1):
        ident=plain(s.cell(row,4).value).strip()
        if ident.startswith('SW-FW-UnitTest-'):
            if current:current['end']=row-1
            sid=re.search(r'SU-ID-[A-Z0-9]+-\d+',plain(s.cell(row,7).value)).group()
            current=dict(test=ident,start=row,end=s.max_row,sid=sid,
                         description=plain(s.cell(row,5).value),design=desc.get(sid),
                         verdict=plain(s.cell(row,13).value).strip())
            groups.append(current)
    current['end']=next(m.max_row for m in s.merged_cells.ranges if f'D{current["start"]}' in m)
    cells=[dict(sheet=sh.title,cell=c.coordinate,text=plain(c.value))
           for sh in book if sh.title in KEEP for row in sh for c in row if c.value]
    result=dict(source=str(SOURCE),target=str(TARGET),sourceHash=digest(SOURCE),
                rows=s.max_row,merges=len(s.merged_cells.ranges),groups=groups,
                missing=[g['sid'] for g in groups if not g['design'] or not g['design']['intro']],
                mentions=[c for c in cells if 'QCC' in c['text']],
                references=[c for c in cells if any(name in c['text'] for name in DROP)],
                photos=[c for c in cells if re.search(r'사진|영상|촬영|미첨부',c['text'])])
    save('inventory.json',result)
    print(json.dumps(dict(groups=len(groups),missing=result['missing'],rows=s.max_row,
         verdicts=dict(collections.Counter(g['verdict'] for g in groups)),
         intros=[{k:g[k] for k in ('sid','design')} for g in groups[:6]],
         references=result['references']),ensure_ascii=False,indent=2))

def prepare():
    data=read('inventory.json')
    assert digest(SOURCE)==data['sourceHash'],'원본이 달라졌다. 다시 조사한다.'
    assert not data['missing']
    if TARGET.exists():assert digest(TARGET)==data['sourceHash'],'기존 1차본이 수정되어 있으므로 덮어쓰지 않는다.'
    shutil.copy2(SOURCE,QA/'unit-before.xlsx')
    shutil.copy2(SOURCE,TARGET)
    book=load_workbook(SOURCE,rich_text=True);sheet=book[SHEET]
    changes={}; group={g['test']:g for g in data['groups']}
    def put(cell,value,reason,sheetname=SHEET):
        old=plain(book[sheetname][cell].value)
        value='\n'+value.strip()+'\n' if value.strip() else ''
        if old!=value:changes[(sheetname,cell)]=dict(key='unit',sheet=sheetname,cell=cell,old=old,value=value,reason=reason)
    def entry(test,col):return col+str(group['SW-FW-UnitTest-'+test]['start'])
    for g in data['groups']:
        intro=g['design']['intro']
        assert g['design']['style']=='TDC40',(g['sid'],g['design'])
        assert not re.match(r'^(?:그림|표|SU-ID|SI-ID)',intro)
        put('E'+str(g['start']),intro,'상세 설계서의 유닛 설명 반영')
    for sh in book:
        if sh.title not in KEEP:continue
        for row in sh:
            for c in row:
                old=plain(c.value)
                if '통과(시험자 확인)' in old:
                    put(c.coordinate,old.replace('통과(시험자 확인)','통과'),'통과 표기 간략화',sh.title)
    put(entry('CM3-ID-14','J'),
        '- 입력 조건별 led_request 요청과 선택 상태, 실제 LED 표시를 비교한다. 색상/점등/소등, 점멸 횟수와 ON/OFF 순서, 종료 소등을 확인한다. POWER_ON/OFF는4회 점멸 뒤 요청이NONE으로 해제되고 소등되어야 한다. 표시 금지 예외와 엔진 가능/불가 상태의 종료 동작도 확인한다. 오류180/360은 ON180/주기360(OFF180)으로 읽는다.',
        '사진 수집 절차를 동작 확인 절차로 변경')
    put(entry('CM3-ID-39','J'),
        '- 명령 전송/반환값/제어 상태를 확인한다. 같은 측정 지점에서 VP 활성 전/후, 비활성 및 차단 리셋 후의 출력 유무를 확인한다.',
        '화면 첨부 절차를 출력 유무 확인으로 변경')
    put(entry('BOOT-ID-4','H'),
        '- 격리 LED 시험에서 초기화/시작/ISR4000회/10440회 추가/정지 후100회 조건을 준비한다. ISR 호출 횟수는 시험 입력이며 실제 타이머 측정 시간이 아니다. 정상 기동의 LED 확인은 UART 공유핀과 LED 활성 조건을 맞춰 준비한다.',
        '기동 LED 사진 준비 문장 정리')
    put(entry('BOOT-ID-4','J'),
        '- 경과 카운터0→100→361ms, 완료 상태1, 정지 후 카운터/쓰기 횟수 유지와 응용 상태 분리를 비교한다. 정상 기동에서 LED 표시와 정지 후 표시 갱신 중단을 확인한다. 수동 ISR 호출에 따른 카운터 값과 실제 타이머의 측정 시간은 구분한다.',
        '기동 LED 표시 확인 절차로 변경')
    put(entry('BOOT-ID-18','J'),
        '- 격리 시험의 RTT 입력에 reboot\\n을 적용하고 시스템 재시작 함수 호출 횟수=1, 응용 코어 기동=0, 서비스 tick/색상/GPIO 쓰기 상태를 확인한다. 정상 서비스 동작에서는 UART 공유핀 조건을 맞춘 뒤 LED 표시와 재시작 요청에 따른 재시작 동작을 확인한다.',
        '서비스 LED/재시작 확인 절차로 변경')
    k20=entry('CFX-ID-20','K')
    put(k20,plain(sheet[k20].value).replace('사진은 청취 증거가 아니다. ',''),'불필요한 사진 설명 삭제')
    confirmations={
        'CM3-ID-14':'요청별 LED 색상과 점등/점멸 상태, 전원 표시4회 반복 후 요청 해제/소등, 표시 금지 예외와 종료 동작을 확인했다. 통과 기준과 일치한다.',
        'CM3-ID-39':'VP 활성 상태에서는 출력이 있고, 비활성 및 차단 리셋 후에는 출력이 없음을 확인했다. 명령/제어 상태와 출력 유무가 통과 기준과 일치한다.',
        'BOOT-ID-4':'정상 기동에서 LED 표시와 정지 후 표시 갱신 중단을 확인했다. 경과 카운터와 완료 상태를 포함해 통과 기준과 일치한다.',
    }
    for test,confirmation in confirmations.items():
        coord=entry(test,'L');old=plain(sheet[coord].value)
        # 기존 중복 결과 문단은 한 번만 남기고 확인 문장을 붙인다.
        seen=set();lines=[]
        for line in old.splitlines():
            if not line.strip() or '(사진필요)' in line or line in seen:continue
            seen.add(line);lines.append(line)
        put(coord,'\n'.join(lines+[confirmation]),'시험자 확인에 따라 사진 대신 텍스트 결과 기록')
        put(entry(test,'M'),'통과','사진 관련 확인 항목의 통과 표기')
    put(entry('BOOT-ID-15','L'),
        '검증 결과\n대상별 전송 값/순서/길이가 입력과 일치하고 전송 단위 변환이 정확함을 확인했다. 지정한 코어의 메모리 영역으로 전송되었으며 통과 기준과 일치한다.',
        '전송 자료 미첨부를 시험자 확인 결과로 정리')
    put(entry('BOOT-ID-15','M'),'통과','전송 동작 확인에 따른 통과')
    put(entry('BOOT-ID-18','L'),
        '검증 결과\n서비스 상태에 맞게 LED 표시가 갱신되며 응용 LED 요청과 분리됨을 확인했다. 재시작 요청 시 시스템 재시작 함수 호출 횟수는1회이고 재시작 동작을 확인했다. 통과 기준과 일치한다.',
        '서비스 표시/재시작을 시험자 확인 결과로 정리')
    put(entry('BOOT-ID-18','M'),'통과','서비스 표시/재시작 확인에 따른 통과')
    put('D19','QCC 내부 유닛 시험을 추가할 경우, 아래 E8300 시험 목록 다음에 이어서 작성한다.',
        '현재 E8300 시험과 향후 QCC 내부 시험 추가 위치 구분')
    protected={str(SOURCE):digest(SOURCE),str(DESIGN):digest(DESIGN)}
    for p in SOURCE.parent.glob('*.xlsx'):
        if p!=TARGET and not p.name.startswith('~$'):protected[str(p)]=digest(p)
    for p in DESIGN.parent.glob('*SUVC*.xlsx'):
        if not p.name.startswith('~$'):protected[str(p)]=digest(p)
    save('baseline.json',dict(books=[dict(key='unit',path=str(TARGET),backup=str(QA/'unit-before.xlsx'),
        sheet=SHEET,views=['D30:G34','J897:M897','K1647:M1647'])],protected=protected,sourceHash=digest(SOURCE)))
    # Excel의 실제 글 높이 49.5pt에 맞춰 QCC 안내 행의 기존 45pt를 보완한다.
    save('plan.json',dict(changes=list(changes.values()),shades=[],dropSheets=DROP,groups=data['groups'],
                         heights=[dict(key='unit',sheet=SHEET,row=19,height=54)]))
    print(json.dumps(dict(changedCells=len(changes),descriptionCount=len(data['groups']),
                         output=str(TARGET)),ensure_ascii=False))

def relationships_path(part):
    parent,name=posixpath.split(part)
    return posixpath.join(parent,'_rels',name+'.rels')

def relationship_target(part,target):
    return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part),target))

def build():
    plan=read('plan.json');authored=load_workbook(QA/'unit-authored.xlsx',rich_text=True)
    with zipfile.ZipFile(QA/'unit-before.xlsx') as z:package={n:z.read(n) for n in z.namelist()}
    sx=E.fromstring(package['xl/sharedStrings.xml'])
    shared=sx.findall('s:si',NS)
    roots={name:E.fromstring(package[f'xl/worksheets/sheet{i+1}.xml']) for i,name in enumerate(KEEP)}
    for p in plan['changes']:
        assert plain(authored[p['sheet']][p['cell']].value)==p['value'],p['cell']
        cell=roots[p['sheet']].find(f'.//s:c[@r="{p["cell"]}"]',NS)
        assert cell is not None,p['cell']
        previous.patch_text(package,cell,p['value'],shared)
        if p['sheet']==SHEET and p['cell'].startswith('L'):
            # 사진용 파란 강조가 새 확인 문장에 남지 않게 본문 색으로 돌린다.
            for color in cell.findall('.//s:rPr/s:color',NS):
                if color.get('rgb','')[-6:]=='0000FF':color.set('rgb','FF000000')
    for h in plan.get('heights',[]):
        assert authored[h['sheet']].row_dimensions[h['row']].height==h['height']
        row=roots[h['sheet']].find(f'.//s:row[@r="{h["row"]}"]',NS)
        row.set('ht',str(h['height']));row.set('customHeight','1')
    # 보존하는 두 시트에서 사용하는 공유 문자열만 남긴다.
    uses=[c for root in roots.values() for c in root.findall('.//s:c',NS) if c.get('t')=='s']
    used=sorted({int(c.find('s:v',NS).text) for c in uses});mapping={n:i for i,n in enumerate(used)}
    for child in list(sx):sx.remove(child)
    for n in used:sx.append(copy.deepcopy(shared[n]))
    sx.set('count',str(len(uses)));sx.set('uniqueCount',str(len(used)))
    for c in uses:c.find('s:v',NS).text=str(mapping[int(c.find('s:v',NS).text)])
    package['xl/sharedStrings.xml']=E.tostring(sx,xml_declaration=True,encoding='UTF-8',standalone=True)
    for i,name in enumerate(KEEP):
        package[f'xl/worksheets/sheet{i+1}.xml']=E.tostring(roots[name],xml_declaration=True,encoding='UTF-8',standalone=True)
    wx=E.fromstring(package['xl/workbook.xml']);sheets=wx.find('s:sheets',NS)
    removed_rel=[]
    for node in list(sheets):
        if node.get('name') in DROP:
            removed_rel.append(node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
            sheets.remove(node)
    names=wx.find('s:definedNames',NS)
    if names is not None:
        for node in list(names):
            if int(node.get('localSheetId','0'))>=len(KEEP):names.remove(node)
            else:assert not any(name in (node.text or '') for name in DROP)
    for view in wx.findall('s:bookViews/s:workbookView',NS):
        if int(view.get('activeTab','0'))>=len(KEEP):view.set('activeTab','1')
        if int(view.get('firstSheet','0'))>=len(KEEP):view.set('firstSheet','0')
    rx=E.fromstring(package['xl/_rels/workbook.xml.rels'])
    for node in list(rx):
        if node.get('Id') in removed_rel:rx.remove(node)
    package['xl/workbook.xml']=E.tostring(wx,xml_declaration=True,encoding='UTF-8',standalone=True)
    package['xl/_rels/workbook.xml.rels']=E.tostring(rx,xml_declaration=True,encoding='UTF-8',standalone=True)
    # 관계 그래프에서 도달할 수 없는 삭제 시트/그림/프린터 자료를 제외한다.
    reachable={'[Content_Types].xml','_rels/.rels'};pending=[]
    top=E.fromstring(package['_rels/.rels'])
    for rel in top:
        if rel.get('TargetMode')!='External':pending.append(rel.get('Target').lstrip('/'))
    while pending:
        part=pending.pop()
        if part in reachable:continue
        assert part in package,part
        reachable.add(part);rp=relationships_path(part)
        if rp in package:
            reachable.add(rp)
            for rel in E.fromstring(package[rp]):
                if rel.get('TargetMode')!='External':pending.append(relationship_target(part,rel.get('Target')))
    types=E.fromstring(package['[Content_Types].xml'])
    for node in list(types):
        if E.QName(node).localname=='Override' and node.get('PartName','').lstrip('/') not in reachable:types.remove(node)
    package['[Content_Types].xml']=E.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)
    if 'docProps/app.xml' in package:
        app=E.fromstring(package['docProps/app.xml']);ap={'ep':'http://schemas.openxmlformats.org/officeDocument/2006/extended-properties','vt':'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'}
        pairs=app.findall('ep:HeadingPairs/vt:vector/vt:variant',ap)
        for i in range(0,len(pairs),2):
            label=pairs[i].find('vt:lpstr',ap)
            if label is not None and label.text in ('워크시트','Worksheets'):
                pairs[i+1].find('vt:i4',ap).text=str(len(KEEP))
        titles=app.find('ep:TitlesOfParts/vt:vector',ap)
        if titles is not None:
            for child in list(titles):
                if child.text in DROP:titles.remove(child)
            titles.set('size',str(len(titles)))
        package['docProps/app.xml']=E.tostring(app,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(QA/'unit-candidate.xlsx','w',zipfile.ZIP_DEFLATED) as z:
        for path,data in package.items():
            if path in reachable:z.writestr(path,data)
    save('package-removal.json',dict(removedParts=sorted(set(package)-reachable),retainedSheets=KEEP))
    print('candidate built; retained sheets:',KEEP)

def verify():
    before=load_workbook(QA/'unit-before.xlsx',rich_text=True)
    after=load_workbook(QA/'unit-candidate.xlsx',rich_text=True)
    plan=read('plan.json');expected={(p['sheet'],p['cell']):p for p in plan['changes']}
    assert after.sheetnames==KEEP,after.sheetnames
    log_count=0
    for name in KEEP:
        a,b=before[name],after[name]
        assert (a.max_row,a.max_column)==(b.max_row,b.max_column)
        assert set(map(str,a.merged_cells.ranges))==set(map(str,b.merged_cells.ranges))
        height_changes={h['row']:h['height'] for h in plan.get('heights',[]) if h['sheet']==name}
        for index,dim in a.row_dimensions.items():
            assert b.row_dimensions[index].height==height_changes.get(index,dim.height),(name,index,'height')
        assert len(b._images)==0
        for row in a:
            for c in row:
                new=b[c.coordinate];oldtext=plain(c.value);newtext=plain(new.value);key=(name,c.coordinate)
                if key in expected:assert newtext==expected[key]['value'],key
                else:assert oldtext==newtext,(key,'unplanned change')
                assert c._style==new._style,key
                assert previous.LOG.findall(oldtext)==previous.LOG.findall(newtext),(key,'RTT change')
                log_count+=len(previous.LOG.findall(oldtext))
                assert not re.search(r'사진|영상|미첨부|통과\(시험자 확인\)',newtext),(key,newtext)
                assert not any(s in newtext for s in DROP),key
    for g in plan['groups']:
        row=g['start'];assert plain(after[SHEET][f'E{row}'].value).strip()==g['design']['intro'],g['sid']
        if g['verdict']=='기존 기록 불일치':assert plain(after[SHEET][f'M{row}'].value).strip()==g['verdict']
    for p,h in read('baseline.json')['protected'].items():assert digest(p)==h,p
    verdicts=collections.Counter(plain(after[SHEET][f'M{g["start"]}'].value).strip() for g in plan['groups'])
    report=dict(sheets=after.sheetnames,changedCells=len(expected),descriptionCount=len(plan['groups']),
                rawLogLinesPreserved=log_count,verdicts=dict(verdicts),rows=after[SHEET].max_row,
                merges=len(after[SHEET].merged_cells.ranges),qccNotice='C19:D19',lastTest=plan['groups'][-1],
                qccCells=[c.coordinate for row in after[SHEET] for c in row if c.value and 'QCC' in str(c.value)])
    save('verification.json',report)
    print(json.dumps({k:v for k,v in report.items() if k not in ('lastTest','qccCells')},ensure_ascii=False,indent=2))

def publish():
    base=read('baseline.json');native=read('native-checks.json')
    assert not native['insufficient']
    assert digest(SOURCE)==base['sourceHash']
    assert digest(TARGET)==base['sourceHash'],'복사본이 외부에서 수정되었다.'
    temp=TARGET.with_name(TARGET.name+'.tmp');shutil.copy2(QA/'unit-candidate.xlsx',temp);temp.replace(TARGET)
    assert digest(TARGET)==digest(QA/'unit-candidate.xlsx')
    save('publication.json',dict(path=str(TARGET),hash=digest(TARGET)));print(str(TARGET))

if __name__=='__main__':
    {'inventory':inventory,'prepare':prepare,'build':build,'verify':verify,'publish':publish}[sys.argv[1]]()
