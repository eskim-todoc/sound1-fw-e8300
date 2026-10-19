"""Preserve SRS workbook content while highlighting existing owner prefixes."""
from pathlib import Path
from copy import copy, deepcopy
import hashlib,json,shutil,zipfile
from lxml import etree as E
import openpyxl

ROOT=Path(__file__).resolve().parents[4]
TASK=Path(__file__).resolve().parent
QA=TASK/'qa/srs-role-format-20261012'
TARGET=ROOT/'docs/참고/인허가/사운드1/DHF/02_설계 입력/SW/펌웨어/SRS-FW-TD24A 소프트웨어 요구사항 사양서_rev0(김은수 작성).xlsx'
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
q=lambda n:'{'+S+'}'+n
ns={'s':S}
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def unpack(p):
    with zipfile.ZipFile(p) as z:return {n:z.read(n) for n in z.namelist()}
def prepare():
    QA.mkdir(parents=True,exist_ok=True)
    before=QA/'before.xlsx'
    if before.exists():assert sha(before)==sha(TARGET),'External edit'
    else:shutil.copyfile(TARGET,before)
    wb=openpyxl.load_workbook(before)
    ws=wb['소프트웨어 입력 요구사항 사양서']
    entries=[]
    for row in ws:
        for c in row:
            txt=c.value
            if isinstance(txt,str) and txt.startswith(('[QCC 단독]','[E8300/QCC ')):
                mode='qcc' if txt.startswith('[QCC 단독]') else 'e8300'
                entries.append({'cell':c.coordinate,'role':mode,'value':txt,'size':12 if mode=='qcc' else c.font.sz,'height':ws.row_dimensions[c.row].height,'ID':ws.cell(c.row,7).value})
    assert len(entries)==12 and sum(x['role']=='qcc' for x in entries)==7
    plan={'target':str(TARGET),'sha256':sha(TARGET),'sheet':ws.title,'edits':entries}
    (QA/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(plan,ensure_ascii=False,indent=2))
def preserve():
    plan=json.loads((QA/'plan.json').read_text(encoding='utf8'))
    pkg=unpack(QA/'before.xlsx');art=unpack(QA/'artifact.xlsx')
    styles=E.fromstring(pkg['xl/styles.xml']);as_=E.fromstring(art['xl/styles.xml'])
    root=E.fromstring(pkg['xl/worksheets/sheet2.xml']);ar=E.fromstring(art['xl/worksheets/sheet2.xml'])
    fonts=styles.find(q('fonts'));xfs=styles.find(q('cellXfs'))
    afonts=as_.find(q('fonts'));axfs=as_.find(q('cellXfs'))
    for item in plan['edits']:
        addr=item['cell'];cell=root.find(f'.//s:c[@r="{addr}"]',ns);ac=ar.find(f'.//s:c[@r="{addr}"]',ns)
        authored=afonts[int(axfs[int(ac.get('s','0'))].get('fontId','0'))]
        original_xf=xfs[int(cell.get('s','0'))]
        # Artifact Tool authors supported color/weight/size. Merge just these
        # properties into the original font to preserve every other native part.
        font=deepcopy(fonts[int(original_xf.get('fontId','0'))])
        for tag in ('b','sz','color'):
            for old in font.findall(q(tag)):font.remove(old)
            val=authored.find(q(tag));assert val is not None,(addr,tag)
            font.append(deepcopy(val))
        # Single underline is not exposed by the artifact font API.
        if item['role']=='qcc':
            for old in font.findall(q('u')):font.remove(old)
            E.SubElement(font,q('u'),{'val':'single'})
        font_id=next((i for i,n in enumerate(fonts) if E.tostring(n)==E.tostring(font)),len(fonts))
        if font_id==len(fonts):fonts.append(font)
        xf=deepcopy(original_xf);xf.set('fontId',str(font_id));xf.set('applyFont','1')
        xf_id=next((i for i,n in enumerate(xfs) if E.tostring(n)==E.tostring(xf)),len(xfs))
        if xf_id==len(xfs):xfs.append(xf)
        cell.set('s',str(xf_id))
    fonts.set('count',str(len(fonts)));xfs.set('count',str(len(xfs)))
    # Heights are measured in Excel at the actual width before applying them.
    measure=QA/'native-check.json'
    if measure.exists():
        for item in json.loads(measure.read_text(encoding='utf-8-sig')):
            row=root.find(f'.//s:row[@r="{item["row"]}"]',ns)
            if item['needed']>float(row.get('ht','15'))+1:
                row.set('ht',str(min(409.5,item['needed']+3)));row.set('customHeight','1')
    pkg['xl/styles.xml']=E.tostring(styles,encoding='UTF-8',xml_declaration=True,standalone=True)
    pkg['xl/worksheets/sheet2.xml']=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
    with zipfile.ZipFile(QA/'candidate.xlsx','w',zipfile.ZIP_DEFLATED) as z:
        for n,d in pkg.items():z.writestr(n,d)
    print('Preserved original workbook; changed only 12 cell font references, added fonts/styles and measured row heights.')
def verify(publish=False):
    import pypdfium2 as pdfium
    plan=json.loads((QA/'plan.json').read_text(encoding='utf8'))
    before=openpyxl.load_workbook(QA/'before.xlsx')
    after=openpyxl.load_workbook(QA/'candidate.xlsx')
    assert before.sheetnames==after.sheetnames
    wanted={x['cell']:x for x in plan['edits']}
    changed=[]
    for a,b in zip(before,after):
        assert a.max_row==b.max_row and a.max_column==b.max_column
        assert set(map(str,a.merged_cells.ranges))==set(map(str,b.merged_cells.ranges))
        for row in a:
            for c in row:
                d=b[c.coordinate]
                assert c.value==d.value,(a.title,c.coordinate,'value')
                if c._style!=d._style:
                    assert a.title==plan['sheet'] and c.coordinate in wanted
                    for prop in ('fill','alignment','border','protection','number_format'):
                        assert copy(getattr(c,prop))==copy(getattr(d,prop)),(c.coordinate,prop)
                    changed.append(c.coordinate)
                    x=wanted[c.coordinate]
                    assert d.font.b and d.font.sz==x['size'] and d.font.name==c.font.name
                    assert d.font.color.rgb[-6:]==('C00000' if x['role']=='qcc' else '0070C0')
                    if x['role']=='qcc':assert d.font.u=='single'
    assert set(changed)==set(wanted)
    old,new=unpack(QA/'before.xlsx'),unpack(QA/'candidate.xlsx')
    assert set(old)==set(new)
    changedparts={n for n in old if old[n]!=new[n]}
    assert changedparts=={'xl/styles.xml','xl/worksheets/sheet2.xml'}
    oa,nb=[E.fromstring(p['xl/worksheets/sheet2.xml']) for p in (old,new)]
    for addr in wanted:
        a=oa.find(f'.//s:c[@r="{addr}"]',ns);b=nb.find(f'.//s:c[@r="{addr}"]',ns)
        b.set('s',a.get('s'))
    heightchanges=[]
    for a,b in zip(oa.findall('.//s:row',ns),nb.findall('.//s:row',ns)):
        if a.attrib!=b.attrib:
            assert a.get('r') in ('56','57')
            heightchanges.append([a.get('r'),a.get('ht'),b.get('ht')])
            b.attrib.clear();b.attrib.update(a.attrib)
    assert E.tostring(oa)==E.tostring(nb)
    native=json.loads((QA/'native-check.json').read_text(encoding='utf-8-sig'))
    assert all(x['needed']<=x['current']+1 for x in native)
    result={'cells':changed,'QCC_alone':7,'E8300_related':5,'values_formulas_IDs_unchanged':True,'merges_unchanged':True,'height_changes':heightchanges,'changed_parts':sorted(changedparts),'native_font_fit_checks':len(native)}
    if publish:
        assert (QA/'visual-review.json').exists()
        assert json.loads((QA/'visual-review.json').read_text(encoding='utf8'))['issues']==[]
        assert sha(TARGET)==plan['sha256'],'External edit'
        temp=TARGET.with_name(TARGET.name+'.pending');assert not temp.exists()
        shutil.copyfile(QA/'candidate.xlsx',temp)
        assert sha(TARGET)==plan['sha256']
        temp.replace(TARGET)
        assert sha(TARGET)==sha(QA/'candidate.xlsx')
        result['path']=str(TARGET);result['sha256']=sha(TARGET)
    else:
        for src in QA.glob('native-*.pdf'):
            doc=pdfium.PdfDocument(str(src))
            for i,p in enumerate(doc):p.render(scale=2).to_pil().save(QA/(src.stem+f'-{i+1}.png'))
    (QA/('published-verification.json' if publish else 'verification.json')).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
if __name__=='__main__':
    import sys
    if '--publish' in sys.argv:verify(True)
    elif '--verify' in sys.argv:verify()
    elif '--preserve' in sys.argv:preserve()
    else:prepare()
