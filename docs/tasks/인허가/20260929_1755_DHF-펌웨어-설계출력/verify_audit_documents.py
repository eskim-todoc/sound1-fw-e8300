"""Independent structural checks and page-image comparison for four artifacts."""
from copy import deepcopy
import hashlib
import json
import re
from pathlib import Path
import zipfile
import openpyxl
import pypdfium2 as pdfium
from lxml import etree as E
from inspect_audit_units import QA,FILES,ROOT,TASK
from inspect_sad_unit_trace import N,q,text,blocks
from preserve_audit_workbooks import unpack,S

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def stable_doc(doc):
    root=E.fromstring(unpack(doc)['word/document.xml'])
    return [(E.QName(x).localname,text(x)) for x in root.find(q('body')) if text(x)]

def main():
    import sys
    published='--published' in sys.argv
    base=json.loads((QA/'baseline.json').read_text(encoding='utf8'))
    plan=json.loads((QA/'sheet-edits.json').read_text(encoding='utf8'))
    report={}
    for kind in ('suvc','suvr'):
        before=unpack(QA/(kind+'-before.xlsx'));after=unpack(QA/(kind+'.xlsx'))
        assert set(before)==set(after)
        assert [n for n in before if before[n]!=after[n]]==['xl/worksheets/sheet2.xml']
        wb0=openpyxl.load_workbook(QA/(kind+'-before.xlsx'),rich_text=True)
        wb1=openpyxl.load_workbook(QA/(kind+'.xlsx'),rich_text=True)
        assert wb0.sheetnames==wb1.sheetnames
        w0,w1=wb0.worksheets[1],wb1.worksheets[1]
        assert w0.max_row==w1.max_row==1734
        assert sorted(map(str,w0.merged_cells.ranges))==sorted(map(str,w1.merged_cells.ranges))
        changes={**plan['common'],**(plan['report'] if kind=='suvr' else {})}
        ids0=[str(row[3].value).strip() for row in w0 if str(row[3].value).strip().startswith('SW-FW-UnitTest-')]
        ids1=[str(row[3].value).strip() for row in w1 if str(row[3].value).strip().startswith('SW-FW-UnitTest-')]
        assert ids0==ids1 and len(ids1)==115
        root0=E.fromstring(before['xl/worksheets/sheet2.xml']); root1=E.fromstring(after['xl/worksheets/sheet2.xml'])
        ns={'s':S};cs0={c.get('r'):E.tostring(c) for c in root0.findall('.//s:c',ns)}
        cs1={c.get('r'):E.tostring(c) for c in root1.findall('.//s:c',ns)}
        assert cs0.keys()==cs1.keys()
        changed={k for k in cs0 if cs0[k]!=cs1[k]}
        assert changed==set(changes),(kind,changed^set(changes))
        for addr,value in changes.items():
            assert str(w1[addr].value)==value,(kind,addr)
            assert value.startswith('\n') and value.endswith('\n')
            assert '·' not in value
        for cell in ('C19','D19'):
            r=root1.find(f'.//s:c[@r="{cell}"]/s:is',ns)
            for run in r.findall('s:r',ns):
                pr=run.find('s:rPr',ns)
                assert pr.find('s:color',ns).get('rgb')=='FFC00000'
                assert pr.find('s:sz',ns).get('val')=='12'
                assert pr.find('s:b',ns) is not None and pr.find('s:u',ns).get('val')=='single'
        for r in root1.findall('.//s:r',ns):
            tx=''.join(r.find('s:t',ns).itertext()) if r.find('s:t',ns) is not None else ''
            if '[E8300 ' in tx or ('보완 필요' in tx and kind=='suvr'):
                pr=r.find('s:rPr',ns)
                assert pr.find('s:color',ns).get('rgb')=='FF0070C0'
                assert pr.find('s:b',ns) is not None
        assert all((w1.row_dimensions[int(r)].height or 0)<=409.5 for r in plan['heights'])
        report[kind]={'test_ids':115,'rows':1734,'changed_cells':sorted(changed),'unchanged_package_parts':len(before)-1,'merges_unchanged':True}
    c=openpyxl.load_workbook(QA/'suvc.xlsx',rich_text=True).worksheets[1]
    r=openpyxl.load_workbook(QA/'suvr.xlsx',rich_text=True).worksheets[1]
    mismatches=[f'{openpyxl.utils.get_column_letter(col)}{row}' for row in range(25,1735) for col in range(1,12) if str(c.cell(row,col).value or '')!=str(r.cell(row,col).value or '')]
    assert not mismatches,mismatches[:20]
    assert all(str(r[a].value)!=str(r[a.replace('L','K')].value) for a in ('L792','L811','L1213'))
    for kind in ('sad','sddd'):
        before=unpack(QA/(kind+'-before.docx'));after=unpack(QA/(kind+'.docx'))
        assert [n for n in before if before[n]!=after[n]]==['word/document.xml']
        root=E.fromstring(after['word/document.xml'])
        saved=E.fromstring(unpack(QA/(kind+'-accepted.docx'))['word/document.xml'])
        assert E.fromstring(after['word/settings.xml']).find(q('trackRevisions')) is not None
        qcc=[];blue=[]
        for p in root.findall('.//w:p',N):
            t=text(p)
            if '[QCC 자료 입력]' in t:qcc.append(p)
            if '[E8300 ' in t:blue.append(p)
        assert len(qcc)==(29 if kind=='sad' else 4),(kind,len(qcc))
        assert len(blue)==(1 if kind=='sad' else 3),(kind,len(blue))
        for p in qcc:
            for run in p.xpath('.//w:r[w:t[contains(.,"[QCC 자료 입력]")]]',namespaces=N):
                pr=run.find(q('rPr'))
                assert pr.find(q('color')).get(q('val'))=='C00000'
                assert pr.find(q('sz')).get(q('val'))=='24'
                assert pr.find(q('b')) is not None and pr.find(q('u')) is not None
        if kind=='sddd':
            units=[text(p) for p in saved.findall('.//w:p',N) if re.match(r'^SU-ID-\w+-\d+,',text(p))]
            previous=TASK/'qa/sad-unit-trace-20261012'
            source_hash=json.loads((previous/'inputs.json').read_text(encoding='utf8'))['sources'][str(FILES['sddd'])]
            assert sha(QA/'sddd-before.docx')==source_hash
            before_units=[b['text'] for b in blocks(previous/'sddd-accepted.docx') if b['type']=='p' and re.match(r'^SU-ID-\w+-\d+,',b['text'])]
            assert units==before_units and len(units)==115
        report[kind]={'qcc_markers':len(qcc),'e8300_markers':len(blue),'unchanged_package_parts':len(before)-1}
    assert stable_doc(QA/'sad-rejected.docx')==stable_doc(TASK/'qa/sad-audit-trace-20261012/candidate-rejected.docx')
    assert stable_doc(QA/'sddd-rejected.docx')==stable_doc(QA/'sddd-before-rejected.docx')
    final=json.loads((QA/'published-verification.json').read_text(encoding='utf8')) if published else None
    for key in ('sad','sddd','suvc','suvr'):
        expected=final['outputs'][key]['sha256'] if published else base[key]['sha256']
        assert sha(FILES[key])==expected,'External edit: '+key
        if published:
            candidate=QA/(key+('-publish.docx' if key in ('sad','sddd') else '.xlsx'))
            assert sha(FILES[key])==sha(candidate)
    assert all(sha(Path(p))==h for p,h in base['protected'].items())
    evidence=json.loads((QA/'checks.json').read_text(encoding='utf8'))
    assert evidence['known_failure_records']==2 and evidence['comparison_records']==93
    report['common_body_columns_identical']=True
    report['protected_files_unchanged']=len(base['protected'])
    report['evidence']={'cases':25,'comparisons':93,'normal_matches':91,'known_mismatches':2,'device_execution':False,'extra_readback_and_timer':True}
    (QA/('post-publication-check.json' if published else 'verification.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(report,ensure_ascii=False,indent=2))

def images():
    report={}
    sources={'sad':TASK/'qa/sad-audit-trace-20261012/candidate.pdf','sddd':TASK/'qa/sad-unit-trace-20261012/sddd.pdf'}
    for kind,old in sources.items():
        olddoc=pdfium.PdfDocument(str(old));newdoc=pdfium.PdfDocument(str(QA/(kind+'.pdf')))
        imdir=QA/(kind+'-pages');imdir.mkdir(exist_ok=True)
        oldhash={}
        for i,page in enumerate(olddoc):
            im=page.render(scale=1.5).to_pil().convert('RGB')
            # The body is compared separately from page number/count fields.
            region=im.crop((0,195,im.width,im.height-100))
            oldhash.setdefault(hashlib.sha256(region.tobytes()).hexdigest(),[]).append(i+1)
        changed=[];reused={}
        for i,page in enumerate(newdoc):
            im=page.render(scale=1.5).to_pil().convert('RGB')
            im.save(imdir/f'page-{i+1:03d}.png')
            region=im.crop((0,195,im.width,im.height-100))
            h=hashlib.sha256(region.tobytes()).hexdigest()
            if h in oldhash:reused[str(i+1)]=oldhash[h]
            else:changed.append(i+1)
        report[kind]={'pages':len(newdoc),'prior_pages':len(olddoc),'new_body_pages':changed,'reused_body_pages':reused,'note':'Header/footer and page counters reviewed separately'}
    (QA/'page-comparison.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:{'pages':v['pages'],'new_body_pages':v['new_body_pages'],'unchanged_bodies':len(v['reused_body_pages'])}for k,v in report.items()},ensure_ascii=False))

if __name__=='__main__':
    import sys
    images() if '--images' in sys.argv else main()
