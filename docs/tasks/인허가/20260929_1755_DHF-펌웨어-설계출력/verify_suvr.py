"""Verify report identifiers, source data, logs and workbook layout."""
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import sys
import zipfile
import openpyxl

TASK=Path(__file__).parent
ROOT=TASK.parents[3]
QA=TASK/'qa/suvr'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def main():
    p=read(QA/'plan.json')
    path=Path(sys.argv[1]) if len(sys.argv)>1 else QA/'candidate.xlsx'
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
    w=openpyxl.load_workbook(path,data_only=False)
    s=w['보고서_유닛 시험']
    errors=[]
    for i,t in enumerate(p['tests'],34):
        for c in range(4,12):
            if s.cell(i,c).value!=t['original'][c-1]:errors.append(f'criteria mismatch {i}/{c}')
        if s.cell(i,13).value!=t['result']['judgment']:errors.append(f'judgment mismatch {i}')
        if not s.cell(i,12).value:errors.append(f'missing result {i}')
        if t['result']['exception'] and s.cell(i,13).value=='통과' and t['result'].get('exception_scope')!='추가 조건':errors.append(f'exception passed {i}')
        if t['result']['photo'] or t['result']['qcc']:
            color=s.cell(i,12).font.color
            if not color or color.type!='rgb' or color.rgb[-6:]!='C00000':errors.append(f'placeholder not red {i}')
    ids=[s.cell(i,7).value for i in range(34,149)]
    assert len(set(ids))==115
    for name,key,cols in [('입력 및 예상값','vectors',3),('RTT 결과 대조','logs',6),('관찰값 확인','observations',3)]:
        sh=w[name]
        if sh.max_row!=len(p[key])+4:errors.append(f'row count {name}')
        for i,row in enumerate(p[key],5):
            for c,v in enumerate(row,1):
                if (sh.cell(i,c).value or '')!=(v or ''):errors.append(f'appendix mismatch {name}/{i}/{c}')
    for sh in w:
        for row in sh:
            for c in row:
                if c.data_type=='e':errors.append(f'Excel error {sh.title}/{c.coordinate}')
                if isinstance(c.value,str) and len(c.value)>32767:errors.append(f'cell overflow {sh.title}/{c.coordinate}')
                if isinstance(c.value,str) and re.search('TD20B|OTE_GEN1|임정우|이호승|2022.08.22',c.value):errors.append(f'legacy field {sh.title}/{c.coordinate}')
        for i,d in sh.row_dimensions.items():
            if d.height and d.height>409.5:errors.append(f'height {sh.title}/{i}')
    source_diffs=[path for path,h in read(QA/'source-before.json').items() if hashlib.sha256((ROOT/path).read_bytes()).hexdigest()!=h]
    criteria=ROOT/'docs/참고/인허가/사운드1/DHF/03_설계 출력/SW/펌웨어/SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx'
    criteria_same=hashlib.sha256(criteria.read_bytes()).hexdigest()=='89967fa692c4f51e3be916c73270d75859b4f692c96928c2d73da0cc6c5096fd'
    if not criteria_same:errors.append('criteria hash changed')
    spec=importlib.util.spec_from_file_location('expected_io_preservation',TASK/'스크립트/expected_io_common.py')
    preservation=importlib.util.module_from_spec(spec);spec.loader.exec_module(preservation)
    old=read(TASK/'qa/expected-io/before.json');current=preservation.snapshot()
    changed_old=[k for k,h in old.items() if current.get(k)!=h]
    added=sorted(set(current)-set(old))
    allowed='docs/참고/인허가/사운드1/DHF/04_설계 검증/SW/펌웨어/SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서.xlsx'
    if changed_old:errors.append(f'previous file changed: {changed_old}')
    if any(k!=allowed for k in added):errors.append(f'unexpected file addition: {added}')
    (QA/'preservation.json').write_text(json.dumps(dict(previous_count=len(old),changed_previous=changed_old,added=added,allowed_addition=allowed),ensure_ascii=False,indent=2),encoding='utf-8')
    cover=w['표지'];assert cover['L2'].value=='SUVR-FW-TD24A' and cover['L3'].value=='Rev.0'
    assert cover['H12'].value=='김은수' and all(cover[a].value is None for a in ['M12','D14','H14','J14','M14','D16','H16','J16','M16','D24'])
    counts={j:sum(t['result']['judgment']==j for t in p['tests']) for j in sorted({t['result']['judgment'] for t in p['tests']})}
    result=dict(file=str(path),sha256=hashlib.sha256(path.read_bytes()).hexdigest(),units=115,judgments=counts,
        source_changed=source_diffs,criteria_unchanged=criteria_same,previous_files_unchanged=not changed_old,rtt_lines=p['log_lines'],sheets=w.sheetnames,
        photos=[t['id'] for t in p['tests'] if t['result']['photo']],qcc=[t['id'] for t in p['tests'] if t['result']['qcc']],
        exceptions=[t['id'] for t in p['tests'] if t['result']['exception']],errors=errors)
    (QA/'checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))
    assert not errors and not source_diffs

if __name__=='__main__':main()
