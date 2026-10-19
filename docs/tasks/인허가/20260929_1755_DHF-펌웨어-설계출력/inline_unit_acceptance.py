"""Read current published records; plan inline, condition-aligned unit tables.

openpyxl is used only for inspection. Workbook authoring is performed by the
adjacent Artifact Tool builder; this module also preserves native OOXML features
which that serializer cannot round-trip (rich runs and historical worksheets).
"""
from __future__ import annotations
import copy
import difflib
import hashlib
import json
import math
import re
import shutil
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from lxml import etree as E
import openpyxl

from preserve_aligned_report_packages import NS, N, parse, serialize, parts, sheet_path, col

TASK = Path(__file__).resolve().parent
ROOT = TASK.parents[3]
DHF = ROOT/'docs/참고/인허가/사운드1/DHF'
QA = TASK/'qa/unit-inline-acceptance'
BOOKS = {
    'criterion': (DHF/'03_설계 출력/SW/펌웨어/SUVC-FW-TD24A 소프트웨어 유닛 검증 기준서.xlsx', '기준서_유닛 검증'),
    'report': (DHF/'04_설계 검증/SW/펌웨어/SUVR-FW-TD24A 소프트웨어 유닛 검증 보고서.xlsx', '보고서_유닛 시험'),
}
UID = re.compile(r'^SW-FW-UnitTest-(CFX|HEAR|CM3|BOOT)-ID-(\d+)$')
q = lambda tag: '{'+NS+'}'+tag

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def save_json(name, value):
    (QA/name).write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
def unit_id(value):
    m = UID.fullmatch(str(value))
    return f'{m[1]}-{int(m[2]):03d}' if m else None
def text(value): return '' if value is None else str(value)

def snapshot():
    QA.mkdir(parents=True, exist_ok=True)
    assert not (QA/'baseline.json').exists(), 'Do not replace this phase baseline.'
    baseline = {'books': [], 'protected': {}}
    for key, (path, sheet) in BOOKS.items():
        backup = QA/f'{key}-before.xlsx'
        shutil.copy2(path, backup)
        w = openpyxl.load_workbook(path, data_only=False)
        baseline['books'].append({'key': key, 'path': str(path), 'sheet': sheet,
            'backup': str(backup), 'hash': sha(path), 'sheets': w.sheetnames})
    for folder in [ROOT/'src', DHF]:
        for p in folder.rglob('*'):
            if p.is_file() and ((folder.name == 'src' and p.suffix in {'.c','.h','.hct'})
                or (folder == DHF and (p.name.startswith('[작성자 전용]')
                    or p.name.startswith('EPTC') or p.name.startswith('SIT')
                    or 'SSTC' in p.name or 'SSTR' in p.name))):
                baseline['protected'][str(p)] = sha(p)
    save_json('baseline.json', baseline)

def wrapped_lines(s, width=435):
    # Font 11: Hangul is ~14.7 px; ASCII ~7 px. Use a conservative estimate.
    out = []
    for line in text(s).split('\n'):
        if not line:
            out.append(''); continue
        span = ''; used = 0
        for ch in line:
            pixels = 14.7 if ord(ch) > 255 else 7.2
            if used + pixels > width and span:
                out.append(span); span = ''; used = 0
            span += ch; used += pixels
        out.append(span)
    return out

def chunks(s, max_lines=22):
    """Split only at real newlines where possible; never split an RTT line."""
    groups = []; lines = []; height = 0
    for line in text(s).split('\n'):
        count = len(wrapped_lines(line))
        if height + count > max_lines and lines:
            groups.append('\n'.join(lines)); lines = []; height = 0
        # Long prose can be a single physical paragraph. Use normal character
        # wrap boundaries for this paragraph; output log lines are <=60 bytes.
        if count > max_lines:
            physical = wrapped_lines(line)
            for offset in range(0, len(physical), max_lines):
                if lines: groups.append('\n'.join(lines)); lines = []; height = 0
                groups.append('\n'.join(physical[offset:offset+max_lines]))
        else:
            lines.append(line); height += count
    if lines: groups.append('\n'.join(lines))
    return groups or ['']

def clean_navigation(s):
    # Only obsolete navigation is removed; values, warnings, photos remain.
    s = re.sub(r'\n?RTT 전문: 「RTT 결과 대조」 시트의 [A-Z0-9-]+', '', text(s))
    s = re.sub(r'\n?조건별 원문/RTT는 「시험 보조자료」[^\n]*', '', s)
    return s.strip()

def current_cases(s):
    # Existing numbered result records are the authority, not old JSON actuals.
    matches = list(re.finditer(r'(?m)^\[(\d+)\] ', s))
    if not matches: return {}, s
    lead = s[:matches[0].start()].strip()
    result = {}
    for i, m in enumerate(matches):
        stop = matches[i+1].start() if i+1 < len(matches) else len(s)
        piece = s[m.end():stop].strip()
        if i == len(matches)-1 and '\n' in piece:
            first, tail = piece.split('\n', 1)
            result[int(m[1])] = first
            return result, '\n'.join(p for p in [lead,tail] if p)
        result[int(m[1])] = piece
    return result, lead

def plan():
    baseline = json.loads((QA/'baseline.json').read_text('utf-8'))
    books = {b['key']:openpyxl.load_workbook(b['backup']) for b in baseline['books']}
    c = books['criterion']['기준서_유닛 검증']
    r = books['report']['보고서_유닛 시험']
    old_pairs = {}
    for name in ['pairs-cfx-hear.json','pairs-cm3.json','pairs-boot.json']:
        for p in json.loads((TASK/'qa/suvr'/name).read_text('utf-8')): old_pairs[p['id']] = p
    logs = defaultdict(list); observations = defaultdict(list); vectors = defaultdict(list)
    for row in list(books['report']['RTT 결과 대조'].values)[4:]:
        if re.fullmatch(r'(CFX|HEAR|CM3|BOOT)-\d{3}',text(row[0])):
            logs[row[0]].append([text(v) for v in row[:6]])
    for row in list(books['report']['관찰값 확인'].values)[4:]:
        if re.fullmatch(r'(CFX|HEAR|CM3|BOOT)-\d{3}',text(row[0])):
            observations[row[0]].append([text(v) for v in row[:3]])
    for row in list(books['report']['입력 및 예상값'].values)[4:]:
        uid, heading, content = map(text, row[:3])
        # Move input definitions used by fixed C/R/V cases and by HEAR math.
        # Historical host execution, binary offsets and auditor narratives are
        # not extra acceptance conditions for the unit under test.
        if re.match(r'^[CRV]\d+',heading) or (uid == 'HEAR-002' and
            (heading.startswith(('5개 입력','vector ','512워드','창 계수')) or '창 계수' in heading)):
            if content.startswith(('구현한 전달기로','구현된 전달기로','구현된 RTT','실제 PC','PC 실행','PC 관찰','기존 PC','실제 관찰')):
                continue
            if re.fullmatch(r'\[UT-[^\n]+\] C\d+ (?:PASS|FAIL|OBSERVED).*',content):
                continue  # The complete exact message is already included.
            vectors[uid].append((heading,content))

    records = []; mapping = {}; body = []; cursor = 30
    counts = defaultdict(int)
    for original in range(30, c.max_row+1):
        uid = unit_id(c.cell(original,4).value)
        if not uid: continue
        common = [c.cell(original, i).value for i in range(1,12)]
        old_result = clean_navigation(r.cell(original,12).value)
        verdict = text(r.cell(original,13).value)
        cases, remainder = current_cases(old_result)
        if uid == 'CFX-020' and 3 in cases and remainder.startswith('(시험 결과 기입 필요)'):
            line, _, rest = remainder.partition('\n')
            cases[3] += '\n'+line
            remainder = rest
        blocks = []
        acceptance = clean_navigation(common[10])
        acceptance = acceptance.replace('아래 예상 결과','같은 시험의 조건별 예상 결과')
        blocks.append(('판정 범위', acceptance,
            remainder if cases else old_result, 'summary'))
        for i, pair in enumerate(old_pairs[uid]['comparisons'],1):
            actual = cases.get(i)
            if actual is None:
                # Preserve a record only if this value is present in CURRENT
                # main result; never resurrect an obsolete pass from old pairs.
                if pair['actual'] and pair['actual'] in old_result: actual = pair['actual']
                else: actual = '이 조건의 결과 기입 필요. 기존 종합 기록은 이 시험의 첫 행에 보존했다.'
            # Current input wording incorporates the last collection-scope
            # review. Do not reintroduce obsolete target/reset/FIFO assumptions.
            label = actual.split(' / ',1)[0] if ' / ' in actual else pair['label']
            anticipated = pair['expected']
            if uid == 'CFX-020': anticipated = anticipated.replace('FIFO 출력','RAM sink 출력')
            expected = f'입력/조건 [{i}]: {label}\n예상 결과: {anticipated}'
            blocks.append((f'조건 [{i}]',expected,actual,'comparison'))
        # Group adjacent rows belonging to the same input definition. Per-index
        # values stay explicit; a 32-element table is not replaced by a summary.
        groups = []
        for heading, value in vectors[uid]:
            if groups and groups[-1][0] == heading: groups[-1][1].append(value)
            else: groups.append([heading,[value]])
        for heading, values in groups:
            blocks.append((f'입력/원소별 기준: {heading}',
                '\n'.join(values),
                '입력과 원소별 값은 이 시험의 조건별 결과/RTT 원문과 대조한다. 실제 추가 관찰값은 이 칸에 기입한다.', 'vector'))
        for uid_, case, condition, expected, actual, note in logs[uid]:
            blocks.append((f'RTT {case}',
                f'입력/조건: {condition}\n예상 RTT 전문:\n{expected}',
                f'{actual}\n기록 구분: {note}' if actual else '실제 수집 RTT 원문 기입 필요.', 'rtt'))
        # SDK/HEAR raw observations require comparison ranges/relations rather
        # than a fake immutable numeric transcript. Other snapshot definitions
        # are retained as optional observations, not extra functional pass tests.
        raw_required = uid in {'CFX-023','CFX-031','CFX-032','HEAR-002'}
        has_observation = any('OBSERVED' in a[3] for a in logs[uid])
        if raw_required or has_observation:
            for uid_, heading, expected in observations[uid]:
                mode = '관찰값 비교 기준' if raw_required else '선택 관찰 / 기능 판정과 구분'
                blocks.append((f'{mode}: {heading}',expected,
                    '실제 수집값 기입 필요. 입력/상태 및 이 행의 범위/관계식과 대조한다.' if raw_required
                    else '선택 관찰을 수집한 경우 입력/상태와 수집값을 기록한다. 관찰만으로 기능 통과를 판정하지 않는다.', 'observation'))
        start = cursor
        physical = []
        for label, expected, actual, kind in blocks:
            es = chunks(expected); ac = chunks(actual)
            n = max(len(es),len(ac))
            for k in range(n):
                prefix = label if n==1 else f'{label} ({k+1}/{n})'
                left = prefix+'\n'+(es[k] if k < len(es) else '앞 행의 같은 조건에 대한 결과를 이어 기록한다.')
                right = prefix+'\n'+(ac[k] if k < len(ac) else '')
                height = min(409, math.ceil(max(42, max(len(wrapped_lines(left)),len(wrapped_lines(right))) * 15 + 18)*1.12/3)*3)
                physical.append({'row':cursor,'uid':uid,'kind':kind,'label':label,
                    'expected':left,'actual':right,'height':height})
                body.append([None]*10+[left,right,None]); counts[kind]+=1; cursor+=1
        # Retain criterion metadata without repeating it on every numeric row.
        body[start-30][:10] = common[:10]
        body[start-30][12] = verdict
        mapping[str(original)] = start
        records.append({'id':uid,'testId':common[3],'detailId':common[6],
            'original':original,'start':start,'end':cursor-1,'rows':physical,
            'common':common,'verdict':verdict,'oldResult':text(r.cell(original,12).value),
            'sourceLogs':logs[uid], 'inputDefinitions':vectors[uid],
            'sourceObservations':observations[uid] if raw_required or has_observation else []})
    assert len(records)==115 and len({p['id'] for p in records})==115
    end_old = max(p['original'] for p in records)
    tail_count = c.max_row-end_old
    save_json('plan.json',{'baseline':baseline,'records':records,'body':body,
        'mapping':mapping,'start':30,'endOld':end_old,'endNew':cursor-1,
        'tailCount':tail_count,'counts':dict(counts)})
    print('Units',len(records),'body rows',len(body),'counts',dict(counts))
    print('Expected RTT source chunks',sum(len(v) for v in logs.values()))

def inline_cell(address, value, style):
    c=E.Element(q('c'),r=address,s=str(style),t='inlineStr')
    t=E.SubElement(E.SubElement(c,q('is')),q('t'))
    t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=text(value)
    return c

def rich_cell(address, value, old, marker, style):
    cell=E.Element(q('c'),r=address,s=str(style),t='inlineStr');ins=E.SubElement(cell,q('is'))
    mask=[True]*len(value)
    for m in difflib.SequenceMatcher(None,old,value,autojunk=False).get_matching_blocks():
        mask[m.b:m.b+m.size]=[False]*m.size
    a=0
    while a<len(value):
        b=a+1
        while b<len(value) and mask[b]==mask[a]:b+=1
        run=E.SubElement(ins,q('r'));rp=E.SubElement(run,q('rPr'))
        E.SubElement(rp,q('rFont'),val='맑은 고딕');E.SubElement(rp,q('sz'),val='11')
        E.SubElement(rp,q('color'),rgb=('FF14632A' if marker=='[+]' else 'FF9C2020') if mask[a] else 'FF666666')
        E.SubElement(rp,q('b'),val='1' if mask[a] else '0')
        t=E.SubElement(run,q('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=value[a:b];a=b
    return cell

def plain_value(c, package):
    if c is None:return ''
    if c.get('t')=='s':
        sst=parse(package['xl/sharedStrings.xml']);index=int(c.find('s:v',N).text)
        return ''.join(sst[index].itertext())
    return ''.join(c.find('s:is',N).itertext()) if c.get('t')=='inlineStr' else text(c.findtext('s:v',namespaces=N))

def replace_payload(dst, src):
    # Native style IDs and formula objects remain owned by the target package.
    for child in list(dst):dst.remove(child)
    dst.set('t','inlineStr')
    for child in src:dst.append(copy.deepcopy(child))

def preserve():
    plan=json.loads((QA/'plan.json').read_text('utf-8')); records=plan['records']
    original_c=parts(QA/'criterion-before.xlsx')
    cpath=sheet_path(original_c,BOOKS['criterion'][1]);croot=parse(original_c[cpath])
    crows={int(x.get('r')):x for x in croot.find('s:sheetData',N)}
    cb=next(x for x in plan['baseline']['books'] if x['key']=='criterion')
    for book in plan['baseline']['books']:
        key=book['key'];p=parts(Path(book['backup']));entry=sheet_path(p,book['sheet']);root=parse(p[entry])
        authored=parts(QA/f'{key}-authored.xlsx');ar=parse(authored[sheet_path(authored,book['sheet'])])
        # Artifact export may use its own shared-string index. Resolve payloads
        # before inserting them into native files with unrelated SST indices.
        authored_sst=parse(authored['xl/sharedStrings.xml']) if 'xl/sharedStrings.xml' in authored else None
        for cell in ar.findall('s:sheetData/s:row/s:c',N):
            if cell.get('t')=='s':
                v=cell.find('s:v',N);si=authored_sst[int(v.text)];cell.remove(v)
                ins=E.SubElement(cell,q('is'))
                for child in si:ins.append(copy.deepcopy(child))
                cell.set('t','inlineStr')
            elif cell.get('t')=='str':
                v=cell.find('s:v',N);value=text(v.text) if v is not None else ''
                if v is not None:cell.remove(v)
                t=E.SubElement(E.SubElement(cell,q('is')),q('t'))
                t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=value
                cell.set('t','inlineStr')
        ar_cells={c.get('r'):c for c in ar.findall('s:sheetData/s:row/s:c',N)}
        rows=root.find('s:sheetData',N);oldrows={int(x.get('r')):x for x in rows}
        # Keep all pre-test metadata/title rows, including native rich strings.
        rows[:]=[x for x in rows if int(x.get('r'))<30]
        merges=root.find('s:mergeCells',N)
        merges[:]=[m for m in merges if int(re.search(r'\d+',m.get('ref').split(':')[1]).group())<30]
        styles=parse(p['xl/styles.xml']);xfs=styles.find('s:cellXfs',N)
        top_styles={};red_styles={}
        def top_style(source_id, red=False, black=False):
            cache=red_styles if red else top_styles
            cache_key=(source_id,black)
            if cache_key in cache:return cache[cache_key]
            xf=copy.deepcopy(xfs[source_id]);alignment=xf.find('s:alignment',N)
            if alignment is None:alignment=E.SubElement(xf,q('alignment'))
            alignment.set('vertical','top');alignment.set('wrapText','1')
            if red or black:
                fonts=styles.find('s:fonts',N);font=copy.deepcopy(fonts[int(xf.get('fontId','0'))])
                color=font.find('s:color',N)
                if color is None:color=E.SubElement(font,q('color'))
                color.attrib.clear();color.set('rgb','FFC00000' if red else 'FF000000')
                xf.set('fontId',str(len(fonts)));fonts.append(font);fonts.set('count',str(len(fonts)))
            index=len(xfs);xfs.append(xf);xfs.set('count',str(len(xfs)));cache[cache_key]=index;return index
        for rec in records:
            base=crows[rec['original']];source_cells={col(c.get('r')):c for c in base}
            ownbase=oldrows[rec['original']];own_cells={col(c.get('r')):c for c in ownbase}
            for item in rec['rows']:
                rn=item['row'];row=E.SubElement(rows,q('row'),r=str(rn),ht=str(item['height']),customHeight='1')
                for cn in range(1,14 if key=='report' else 12):
                    letter=openpyxl.utils.get_column_letter(cn);addr=letter+str(rn)
                    template=own_cells.get(cn)
                    style=int(template.get('s','0')) if template is not None else 0
                    tight=re.sub(r'\s+','',item['actual'])
                    red=cn==12 and any(w in tight for w in ['(사진필요)','(QCC검증결과필요)','수집값기입필요','결과기입필요','결과대조필요'])
                    cell=E.SubElement(row,q('c'),r=addr,s=str(top_style(style,red,black=(cn==12 and not red))))
                    if cn<=10:
                        if rn==rec['start'] and template is not None:
                            # Convert source shared strings to inline so two
                            # packages retain their independent SST unchanged.
                            cell.getparent().replace(cell,inline_cell(addr,rec['common'][cn-1],top_style(style)))
                    elif cn in (11,12):
                        replace_payload(cell,ar_cells[addr])
                    elif cn==13 and rn==rec['start']:
                        cell.getparent().replace(cell,inline_cell(addr,rec['verdict'],top_style(style)))
            for cn in range(2,11):
                letter=openpyxl.utils.get_column_letter(cn)
                if rec['end']>rec['start']:E.SubElement(merges,q('mergeCell'),ref=f'{letter}{rec["start"]}:{letter}{rec["end"]}')
            if key=='report' and rec['end']>rec['start']:
                E.SubElement(merges,q('mergeCell'),ref=f'M{rec["start"]}:M{rec["end"]}')
        shift=plan['endNew']-plan['endOld']
        for number in range(plan['endOld']+1,max(oldrows)+1):
            if number not in oldrows:continue
            row=copy.deepcopy(oldrows[number]);row.set('r',str(number+shift))
            for cell in row:cell.set('r',re.sub(r'\d+$',str(number+shift),cell.get('r')))
            rows.append(row)
        for m in parse(p[entry]).findall('s:mergeCells/s:mergeCell',N):
            begin,end=m.get('ref').split(':');bn=int(re.search(r'\d+',begin).group())
            if bn>plan['endOld']:
                shifted=re.sub(r'\d+',lambda match:str(int(match[0])+shift),m.get('ref'))
                E.SubElement(merges,q('mergeCell'),ref=shifted)
        merges.set('count',str(len(merges)))
        if key=='report':
            for dc in root.findall('s:cols/s:col',N):
                if dc.get('min')=='12' and dc.get('max')=='12':
                    dc.set('width','64.625');dc.set('customWidth','1')
        root.find('s:dimension',N).set('ref',f'A1:{"M" if key=="report" else "K"}{plan["endNew"]+plan["tailCount"]}')
        # A large merged metadata block stays readable during scrolling; K/L
        # remain unmerged, so each expected/result pair is selectable/copyable.
        pane=root.find('s:sheetViews/s:sheetView/s:pane',N)
        view=root.find('s:sheetViews/s:sheetView',N)
        if pane is None:
            pane=E.Element(q('pane'));view.insert(0,pane)
        pane.attrib.clear();pane.attrib.update({'xSplit':'4','topLeftCell':'K28','activePane':'topRight','state':'frozen'})
        for selection in view.findall('s:selection',N):
            selection.attrib.clear();selection.attrib.update({'pane':'topRight','activeCell':'K30','sqref':'K30'})
        breaks=root.find('s:rowBreaks',N)
        if breaks is not None:root.remove(breaks)
        # Footer refs/defined print area must follow the shifted table.
        wb=parse(p['xl/workbook.xml']);idx=next(i for i,s in enumerate(wb.find('s:sheets',N)) if s.get('name')==book['sheet'])
        for name in wb.findall('s:definedNames/s:definedName',N):
            if name.get('localSheetId')==str(idx) and name.get('name')=='_xlnm.Print_Area':
                name.text=f"'{book['sheet']}'!$B$2:${'M' if key=='report' else 'K'}${plan['endNew']+plan['tailCount']}"
        p['xl/workbook.xml']=serialize(wb);p['xl/styles.xml']=serialize(styles);p[entry]=serialize(root)
        if key=='criterion':
            diff_entry=sheet_path(p,'기준서_유닛 검증_변경 비교');diff=parse(p[diff_entry]);drows=diff.find('s:sheetData',N)
            old_diff=list(drows);new=[];counter=30;byid={r['testId']:r for r in records}
            for row in old_diff:
                if int(row.get('r'))<30:new.append(row);continue
                vals={col(c.get('r')):plain_value(c,p) for c in row}
                if vals.get(1)=='[+]' and vals.get(4) in byid:
                    rec=byid[vals[4]]
                    for k,item in enumerate(rec['rows']):
                        nr=E.Element(q('row'),r=str(counter),ht=str(item['height']),customHeight='1')
                        for cn in range(1,12):
                            original=next((c for c in row if col(c.get('r'))==cn),None)
                            style=int(original.get('s','0')) if original is not None else 0
                            value='[+]' if cn==1 else (rec['common'][cn-1] if cn<=10 and k==0 else item['expected'] if cn==11 else '')
                            cell=rich_cell(openpyxl.utils.get_column_letter(cn)+str(counter),text(value),vals.get(cn,''),'[+]',top_style(style))
                            nr.append(cell)
                        new.append(nr);counter+=1
                else:
                    cp=copy.deepcopy(row);cp.set('r',str(counter))
                    for cell in cp:cell.set('r',re.sub(r'\d+$',str(counter),cell.get('r')))
                    new.append(cp);counter+=1
            drows[:]=new;diff.find('s:dimension',N).set('ref',f'A1:K{counter-1}')
            # Historical negative rows are left intact. Only current positive
            # records get extra condition rows, with exact character emphasis.
            p[diff_entry]=serialize(diff);p['xl/styles.xml']=serialize(styles)
            didx=next(i for i,s in enumerate(wb.find('s:sheets',N)) if s.get('name')=='기준서_유닛 검증_변경 비교')
            for name in wb.findall('s:definedNames/s:definedName',N):
                if name.get('localSheetId')==str(didx) and name.get('name')=='_xlnm.Print_Area':
                    name.text=f"'기준서_유닛 검증_변경 비교'!$A$2:$K${counter-1}"
            p['xl/workbook.xml']=serialize(wb)
        with zipfile.ZipFile(QA/f'{key}-candidate.xlsx','w',zipfile.ZIP_DEFLATED) as z:
            for name,data in p.items():z.writestr(name,data)
        print('Preserved',key,plan['endNew']+plan['tailCount'],'rows')

if __name__=='__main__':
    {'snapshot':snapshot,'plan':plan,'preserve':preserve}[sys.argv[1]]()
