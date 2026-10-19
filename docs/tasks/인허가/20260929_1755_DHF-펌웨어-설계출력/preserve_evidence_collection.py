"""Artifact Tool이 작성한 셀만 가져오고 원본의 나머지 OOXML은 보존한다.

기존 그림/병합/인쇄 설정을 다시 작성하지 않는다. 글자별 비교 색은 공개 API에서
지원하지 않는 기존 rich text의 native 보존/갱신 단계이며 셀 내용은 authored 파일과 대조한다.
"""
from __future__ import annotations
import copy
import difflib
import hashlib
import json
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from lxml import etree as E
from preserve_aligned_report_packages import Importer, parts, parse, serialize, sheet_path, NS, N, col
from review_evidence_collection import QA, digest

BASE=json.loads((QA/'baseline.json').read_text(encoding='utf-8'))
PLAN=json.loads((QA/'patch-plan.json').read_text(encoding='utf-8'))

def colored(cell, segments):
    """Preserve the authored text while applying per-run evidence/diff emphasis."""
    text=''.join(v for v,_,_ in segments)
    for child in list(cell):cell.remove(child)
    cell.set('t','inlineStr'); ins=E.SubElement(cell,'{'+NS+'}is')
    for value,color,bold in segments:
        if not value:continue
        run=E.SubElement(ins,'{'+NS+'}r');prop=E.SubElement(run,'{'+NS+'}rPr')
        E.SubElement(prop,'{'+NS+'}color',rgb='FF'+color)
        if bold:E.SubElement(prop,'{'+NS+'}b')
        t=E.SubElement(run,'{'+NS+'}t');t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=value
    assert ''.join(cell.itertext())==text

def diff_segments(old,new,side):
    normal,strong=('8C5B5B','9C0006') if side=='old' else ('54835B','006100')
    result=[]
    for op,a,b,c,d in difflib.SequenceMatcher(None,old,new,autojunk=False).get_opcodes():
        value=old[a:b] if side=='old' else new[c:d]
        result.append((value,normal if op=='equal' else strong,op!='equal'))
    return result

def evidence_segments(text):
    result=[]
    for line in text.splitlines(keepends=True):
        marks=['(사진필요)','(QCC 검증 결과 필요)','(시험 결과 기입 필요)','(자료 필요)','[계측 준비 필요]']
        found=[line.find(t) for t in marks if t in line]
        if found:
            start=min(found)
            result.extend([(line[:start],'111111',False),(line[start:],'C00000',False)])
        else:
            red=len(line)<120 and ('필요' in line or '대기' in line)
            result.append((line,'C00000' if red else '111111',False))
    return result

def build():
    for b in BASE['books']:
        changes=[p for p in PLAN if p['key']==b['key']]
        if not changes:continue
        assert digest(Path(b['path']))==b['sha256'],('external edit',b['path'])
        target=parts(b['backup']);authored=parts(QA/f"{b['key']}-authored.xlsx")
        imp=Importer(target,authored,'evidence_'+b['key'])
        roots={};srcroots={}
        for p in changes:
            sn=p['sheet'];path=sheet_path(target,sn)
            if sn not in roots:
                roots[sn]=parse(target[path])
                srcroots[sn]=imp.convert_sheet(parse(authored[sheet_path(authored,sn)]))
            dst=roots[sn];src=srcroots[sn]
            old=dst.find(f".//s:c[@r='{p['cell']}']",N)
            new=src.find(f".//s:c[@r='{p['cell']}']",N)
            assert new is not None,(b['key'],sn,p['cell'])
            item=copy.deepcopy(new)
            assert old is not None,(b['key'],sn,p['cell'],'no original cell/style')
            item.set('s',old.get('s','0'))
            if 'diffOld' in p:
                colored(item,diff_segments(p['diffOld'],str(p['value']),'new'))
                if p.get('diffMinus'):
                    minus=dst.find(f".//s:c[@r='{p['diffMinus']}']",N)
                    assert minus is not None
                    colored(minus,diff_segments(p['diffOld'],str(p['value']),'old'))
            elif p.get('red'):
                colored(item,evidence_segments(str(p['value'])))
            number=int(re.search(r'\d+',p['cell'])[0]);row=dst.find(f".//s:row[@r='{number}']",N)
            if row is None:
                row=E.SubElement(dst.find('s:sheetData',N),'{'+NS+'}row',r=str(number))
            if old is not None:row.replace(old,item)
            else:row.append(item)
            row[:]=sorted(row,key=lambda c:col(c.get('r')))
            if p.get('height'):row.set('ht',str(p['height']));row.set('customHeight','1')
        for sn,root in roots.items():
            rows=root.find('s:sheetData',N);rows[:]=sorted(rows,key=lambda r:int(r.get('r')))
            target[sheet_path(target,sn)]=serialize(root)
        # Only text/inline rich text was imported. Every destination retains its
        # original style, so no global style/content resource needs to be rewritten.
        # Appending unused duplicate number formats would cause parser ID shifts.
        candidate=QA/f"{b['key']}-candidate.xlsx"
        with zipfile.ZipFile(candidate,'w',zipfile.ZIP_DEFLATED) as z:
            for name,data in target.items():z.writestr(name,data)
        print('candidate',b['key'],len(changes))

def verify():
    results=[]
    for b in BASE['books']:
        changes=[p for p in PLAN if p['key']==b['key']]
        if not changes:continue
        candidate=QA/f"{b['key']}-candidate.xlsx"
        old=load_workbook(b['backup'],rich_text=True);new=load_workbook(candidate,rich_text=True)
        assert old.sheetnames==new.sheetnames
        approved={(p['sheet'],p['cell']) for p in changes}
        approved|={(p['sheet'],p['diffMinus']) for p in changes if p.get('diffMinus')}
        for p in changes:
            assert str(new[p['sheet']][p['cell']].value)==str(p['value']),(b['key'],p['sheet'],p['cell'])
        for so,sn in zip(old,new):
            assert str(so.merged_cells)==str(sn.merged_cells),(b['key'],so.title,'merges')
            assert so.max_row==sn.max_row and so.max_column==sn.max_column
            assert len(so._images)==len(sn._images)
            assert [(k,v.width,v.hidden) for k,v in so.column_dimensions.items()]==[(k,v.width,v.hidden) for k,v in sn.column_dimensions.items()]
            for row in so:
                for a in row:
                    v=sn[a.coordinate]
                    assert a._style==v._style,(b['key'],so.title,a.coordinate,'style changed')
                    if (so.title,a.coordinate) not in approved:assert str(a.value)==str(v.value),(b['key'],so.title,a.coordinate,'unexpected value change')
                    if v.data_type=='e':raise AssertionError((b['key'],so.title,a.coordinate,'Excel error'))
        before=parts(b['backup']);after=parts(candidate)
        for name in before:
            if name.startswith(('xl/media/','xl/drawings/','xl/charts/')):assert before[name]==after[name],(b['key'],name,'native object changed')
        results.append({'book':b['key'],'changedCells':len(changes),'tests':len(b['tests']),'sheets':len(new.sheetnames),'pictures':sum(len(s._images) for s in new)})
        old.close();new.close()
    common_count=0
    for kind in ('unit','integration','system'):
        ck=next(b for b in BASE['books'] if b['key']==kind+'-criterion');rk=next(b for b in BASE['books'] if b['key']==kind+'-report')
        a=load_workbook(QA/f"{ck['key']}-candidate.xlsx",rich_text=True);b=load_workbook(QA/f"{rk['key']}-candidate.xlsx",rich_text=True)
        rr={r['id']:r['row'] for r in rk['tests']}
        for row in range(1,ck['maxRow']+1):
            for c in range(1,ck['common']+1):
                if (row,c)==(2,2):continue  # The document title is criterion/report-specific.
                assert str(a[ck['sheet']].cell(row,c).value)==str(b[rk['sheet']].cell(row,c).value),(kind,row,c,'criterion/report mismatch')
                common_count+=1
        a.close();b.close()
    for scope in ('source','protected'):
        for path,sha in BASE[scope].items():assert digest(Path(path))==sha,(scope,path,'modified')
    report={'books':results,'commonCellsEqual':common_count,'sourceFilesUnchanged':len(BASE['source']),'protectedFilesUnchanged':len(BASE['protected'])}
    (QA/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))

def publish():
    verify()
    record={'date':'2026-10-06','files':[],'boardTestsExecuted':False}
    for b in BASE['books']:
        if not any(p['key']==b['key'] for p in PLAN):continue
        path=Path(b['path']);assert digest(path)==b['sha256'],('external edit',path)
        stage=path.with_name('.evidence-'+path.name)
        assert not stage.exists(),stage
        try:
            shutil.copyfile(QA/f"{b['key']}-candidate.xlsx",stage)
            os.replace(stage,path)
        finally:
            if stage.exists():stage.unlink()
        candidate=QA/f"{b['key']}-candidate.xlsx"
        assert digest(path)==digest(candidate),(b['key'],'saved package mismatch')
        record['files'].append({'key':b['key'],'path':str(path),
                                'beforeSha256':b['sha256'],'afterSha256':digest(path)})
        (QA/'publication.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
        print('published',b['key'])
    for scope in ('source','protected'):
        for path,sha in BASE[scope].items():assert digest(Path(path))==sha,(scope,path,'modified')
    record['sourceFilesUnchanged']=len(BASE['source'])
    record['protectedFilesUnchanged']=len(BASE['protected'])
    (QA/'publication.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')

if __name__=='__main__':
    {'build':build,'verify':verify,'publish':publish}[sys.argv[1]]()
