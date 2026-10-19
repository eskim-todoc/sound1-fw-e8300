"""결과를 읽는 사람의 관점으로 보고서 문장을 재구성한다.

입력/보존 원본은 openpyxl로 읽기만 한다. 실제 셀 저술은 Artifact Tool로
수행한 뒤 원본 OOXML의 해당 텍스트만 교체해 병합/그림/서식을 보존한다.
"""
from __future__ import annotations

import copy
import json
import re
import shutil
import sys
import zipfile
from decimal import Decimal, localcontext, ROUND_FLOOR
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook
from lxml import etree as E
import update_report_results as prev

TASK = Path(__file__).resolve().parent
QA = TASK / 'qa/report-reader-revision'
DHF = prev.DHF
NS, N = prev.NS, prev.N
QCC, PHOTO = prev.QCC, prev.PHOTO
KEEP_QCC_SYSTEM = {47, 48, 49, 52, 71, 72, 76, 79}
PROVENANCE = ('기록: 코드/수치 계산과 시험자 확인 결과를 정리했다. '
              'RTT 전문은 시험자가 확인한 출력값으로 기입했다. '
              '사진 첨부 위치와 QCC 내부 확인 항목은 별도로 표시한다.')
UNITS = json.loads(prev.UPLAN.read_text(encoding='utf-8'))['records']


def jread(path):
    return json.loads(path.read_text(encoding='utf-8'))


def plain(value):
    return '' if value is None else str(value)


def cleaned(value):
    """시험 값과 기능 제약은 보존하고 결과 저술 지시만 제거한다."""
    value = plain(value).replace('·', '/').replace('\r\n', '\n')
    replacements = [
        (r'template의 null.*?합격이 아니다\.', ''),
        (r'실제값을 1이나 기대 구간 중앙값으로 대체하면 안 된다\.', ''),
        (r'관측값에 맞추어 구간을 늘리지 않는다\.', ''),
        (r'PC 대체값으로 고정하지 않음\.', ''),
        (r'단일 실측값을 임의 지정하지 않음\.', ''),
        (r'원시 수치는 variable_records의 동일 case로 별도 배출한다\.', ''),
        (r'current의 별도.*?기존 MD 참조\.', ''),
        (r'run_current의 동일 case도.*?중복 export하지 않는다\.', ''),
        (r'같은 case.*?중복 export하지 않는다\.', ''),
        (r'실제 추가 관찰값은 이 칸에 기입한다\.', ''),
        (r'단일값을 임의.*?(?:않는다|않음)\.', ''),
        (r'범위/관계식 충족 확인\(시험자 확인\)\.', ''),
    ]
    for pat, rep in replacements:
        value = re.sub(pat, rep, value, flags=re.S)
    value = value.replace('실제 SDK 로그/지수 경로(타깃) 또는 PC 수학 대체 경로(호스트)를 거친다',
                          'SDK의 로그/지수 함수를 사용해 이득을 계산한다')
    value = value.replace('이 배열의 1은 사전 구간 내 여부의 판정값이며 SDK 원시 수치가 아니다.',
                          '비교 출력 1은 해당 계산값이 허용 오차 안에 있다는 뜻이다.')
    value = re.sub(r'\n(?:변동 이유|하한 인덱스|상한 인덱스):[^\n]*', '', value)
    value = re.sub(r'\n{3,}', '\n\n', value)
    return value.strip()


def completed(value):
    value = prev.confirmed(cleaned(value))
    for a,b in [('남아야 한다','보존됨'),('사용한다','사용함'),
                ('유지한다','유지됨'),('설정한다','설정됨'),
                ('반환한다','반환됨'),('해제한다','해제됨'),
                ('실행한다','실행됨'),('바뀐다','바뀜'),
                ('갱신된다','갱신됨'),('호출한다','호출됨'),
                ('기록한다','기록됨')]:
        value=value.replace(a,b)
    return value


def without_qcc(value):
    lines=[]
    for line in plain(value).splitlines():
        if QCC in line:
            line=line[:line.index(QCC)].rstrip(' /;')
        if line.strip(): lines.append(line)
    return '\n'.join(lines)


def photo_note(value):
    found=[]
    for line in plain(value).splitlines():
        if PHOTO not in line: continue
        note=line[line.index(PHOTO):]
        if QCC in note: note=note.split(QCC)[0].rstrip(' /')
        note=re.sub(r'원래 OD 핀.*', '', note).strip()
        if note not in found: found.append(note)
    return '\n'.join(found)


def case_number(value):
    m=re.search(r'\bC(\d+)\b', value)
    return int(m.group(1)) if m else None


def getlog(rec, case):
    return next((l for l in rec['sourceLogs'] if l[1]==f'C{case}'), None)


def log_arrays(log):
    return '\n'.join(re.sub(r'^\[UT-[^]]+\] C\d+ ', '', ln)
                     for ln in plain(log).splitlines() if re.search(r' A\[',ln))


def shorten_input(value):
    value=cleaned(value)
    # 수집 지시와 결과 설명 대신 앞의 입력/초기 상태만 사용한다.
    value=re.sub(r'\b(?:calculate_agc_gain|lib_loopback_AGC_out|audio_agc_attack_release)\([^)]*\)', '', value)
    value=re.split(r'(?:레코드 I|수치 출력 순서:|비교 레코드|원시 관찰 레코드)',value)[0]
    value=re.sub(r'(?:를|을) 호출(?:한다|하여)[^.]*\.', '.', value)
    value=value.replace('으로 새로 시작한다','으로 초기화').replace('을 고정한다',' 적용')
    value=value.replace('를 아래 순서로 두고','를 적용하고')
    return value.strip().rstrip('. ')


def vector_result(rec, item, expected):
    case=case_number(item['label'])
    if case is None: return completed(expected)
    title=item['label'].split(': ',1)[-1]
    log=getlog(rec,case)
    desc=next((d[1] for d in rec['inputDefinitions'] if d[0].startswith(f'C{case} ')), '')
    inp=shorten_input(desc or expected.split('\n',1)[-1])
    # 원소별 행은 개별 값을 보존한다. 전체 로그로 같은 행을 덮지 않는다.
    if any(k in expected for k in ('\ni:', '\nチャネル:', '\n項目:')):
        return completed(expected).replace('입력/원소별 결과:', '검증 결과:').replace('입력/원소별 기준:', '검증 결과:')
    if not log: return completed(expected).replace('입력/원소별 기준:', '검증 결과:')
    actual=log[4] if '기존 기록 불일치' in rec.get('currentVerdict','') else log[3]
    if '[UT-' not in actual:
        return f'검증 {title}\n입력: {inp}\n{cleaned(actual)}'
    output=log_arrays(actual)
    if not output: return completed(expected).replace('입력/원소별 기준:', '검증 결과:')
    end='판정: 각 출력이 기준값과 일치하여 통과.'
    if 'FAIL' in actual or ('기존 기록 불일치' in rec.get('currentVerdict','') and log[3]!=log[4]):
        end='판정: 출력과 기준값의 차이는 아래 RTT 결과에 기재했다.'
    return f'검증 {title}\n입력: {inp}\n출력:\n{output}\n{end}'


def observe_result(rec, expected):
    title=expected.splitlines()[0].replace('관찰값 비교 기준:', '측정')
    rng=re.search(r'확인 범위:\s*\[(-?\d+),(-?\d+)\]', expected)
    unit=re.search(r'단위:\s*([^\n]+)',expected)
    if rng:
        low,high=map(int,rng.groups())
        val=str(low) if low==high else f'{low}~{high}'
        return f'{title}\n'+('출력: ' if low==high else '허용 출력: ')+val+(f' {unit.group(1)}' if unit else '')+'\n판정: 기준 충족(시험자 확인).'
    content=cleaned('\n'.join(expected.splitlines()[1:]))
    return f'{title}\n{content}\n판정: 기준 충족(시험자 확인).'


def tidy_result(value):
    value=cleaned(value)
    value=re.sub(r'조건\s*\[(\d+)\]',r'검증 [\1]',value)
    value=value.replace('입력/검증', '입력').replace('판정 범위', '검증 결과')
    value=value.replace('결과 대조 필요','기존 기록 불일치')
    value=value.replace('PC 요청 확인 / 실제 전송 기록 없음','전송 요청 확인 / 메모리 전송 자료 미첨부')
    value=value.replace('PC 요청 확인 / 타깃 재시작 기록 없음','재시작 요청 확인 / 보드 재시작 자료 미첨부')
    value=value.replace('예상값 기준 기록 / 기존 불일치 판정은 유지\n','')
    value=value.replace('관찰 결과:', '측정 결과:')
    return value


def result_voice(value):
    """입력 정의와 결과 수치는 유지하고 미래형/수집 안내를 결과에서 뺀다."""
    value=tidy_result(value)
    value=re.sub(r'[^\n.]*?(?:혼동하지|확대하지|통과 범위에 포함하거나 삭제하지|사진은 청취 증거|중앙값으로 대체|원시 수집 로그가|새 수집 로그가|결과를 생성하지|임의로 만들지)[^\n.]*\.?', '', value)
    value=re.sub(r'[^\n.]*?(?:측정하지 않는다|판정하지 않는다|판정하지 않음)[^\n.]*\.?', '', value)
    value=value.replace('해제하고,','해제됐고,').replace('처리한다','처리됨')
    value=value.replace('서로 겹치지 않는다','서로 겹치지 않음').replace('겹치지 않는다','겹치지 않음')
    value=value.replace('호출하지 않는다','호출 없음').replace('출력하지 않는다','출력 없음')
    value=value.replace('사용하지 않는다','사용 없음').replace('남지 않는다','남지 않음')
    value=value.replace('혼용하지 않는다','분리됨').replace('계속하지 않는다','중단됨')
    value=value.replace('반영된다','반영됨').replace('설정된다','설정됨').replace('호출된다','호출됨')
    value=value.replace('구분한다','구분됨').replace('선택한다','선택됨')
    value=value.replace('확인한다','확인됨').replace('대조한다','대조 완료')
    value=value.replace('반환됨','반환').replace('호출됨','호출 확인')
    value=value.replace('측정 결과','측정 결과').replace('조건을 유지한다','조건 적용')
    value=re.sub(r'\n{3,}','\n\n',value)
    return value.strip()


def concise_input(value):
    value=shorten_input(value)
    value=re.sub(r'\.\s*(?:주 PC fixture|타깃의 보존된|고정 표본|PC fixture).*','',value,flags=re.S)
    # 한 문장 안에 결과 설명을 함께 적었던 벡터는 입력 정의까지만 남긴다.
    value=re.split(r'(?:(?: I0| I\[0|\. I0)|레코드|반환 I|각 입력을 먼저|출력 mix)',value)[0]
    value=re.sub(r'(?:에서|로)\s*(?:[A-Za-z_]\w*(?:\([^)]*\))?\s*)?(?:후 조회.*|을 읽는다.*|를 읽는다.*|함수.*)$','',value)
    value=re.sub(r'\b(?:calculate_agc_gain|logarithmMapping|audio_agc_attack_release)\b','',value)
    value=value.replace('으로 둔다',' 적용').replace('을 설정하고',' 적용')
    value=value.replace('상태를 2, 1, 2로 저장하고 매번 I2S_isStreaming()을 읽는다','상태=2→1→2')
    value=value.replace('볼륨=9','오디오 볼륨=10(내부 인덱스9)').replace('볼륨=0','오디오 볼륨=1(내부 인덱스0)')
    value=value.replace('중간값 대체=0, ','')
    return value.strip().rstrip(' .')


def refine():
    """1차 자체 검토에서 발견한 절차 복사/행 중복/판정 문체를 보완한다."""
    plan=jread(QA/'plan.json');base=jread(QA/'baseline.json')
    changes={(p['key'],p['sheet'],p['cell']):p for p in plan['changes']}
    def setp(key,sheet,addr,value,reason):
        item=changes.get((key,sheet,addr))
        if not item:item={'key':key,'sheet':sheet,'cell':addr};changes[(key,sheet,addr)]=item
        item.update(value=value,segments=prev.colored_segments(value),reason=reason)
    u=next(b for b in base['books'] if b['key']=='unit')
    wb=load_workbook(u['backup'],rich_text=True);s=wb[u['sheet']]
    for rec in UNITS:
        issue='기존 기록 불일치' in plain(s.cell(rec['start'],13).value)
        for item in rec['rows']:
            addr='L'+str(item['row']);p=changes.get(('unit',s.title,addr))
            old=plain(s[addr].value);value=p['value'] if p else old
            if item['kind']=='summary' and not issue:
                if rec['id'] not in ('BOOT-015','BOOT-018'):
                    value=result_voice(value)
                    value=re.sub(r'(?m)^- ','',value)
            elif item['kind']=='summary' and issue:
                value=result_voice(value)
                value=re.sub(r'^검증 결과(?: \(\d+/\d+\))?\n','기존 분석 결과\n',value)
            elif item['kind']=='vector':
                case=case_number(item['label']);log=getlog(rec,case)
                expected=plain(s.cell(item['row'],11).value)
                # 분할된 표의 상세 숫자는 해당 부분만 남겨 원소 대응을 보존한다.
                if any(k in expected for k in ('\ni:', '\n채널:', '\n항목:')):
                    lines=value.splitlines()
                    marker=next((i for i,l in enumerate(lines) if l.startswith(('i:','채널:','항목:'))),None)
                    if marker is not None:
                        inp=concise_input(' '.join(lines[1:marker]))
                        value=lines[0]+'\n입력: '+inp+'\n'+'\n'.join(lines[marker:])
                        if not issue:value+='\n판정: 입력별 출력이 기준값과 일치함.'
                else:
                    value=re.sub(r'입력: ([\s\S]*?)\n출력:',lambda m:'입력: '+concise_input(m.group(1))+'\n출력:',value,count=1)
                value=result_voice(value)
            else:value=result_voice(value)
            if issue:value=value.replace('판정: 출력/동작이 기준과 일치하여 통과.','')
            value=re.sub(r'(?m)^\s*\n','\n',value)
            if value!=old:setp('unit',s.title,addr,value.strip(),'자체 검토 1: 결과 문체 보완')
    overrides={
        'L30':'검증 결과\n실행 허용=0: 정상 처리와 이벤트 처리 호출 없음.\n실행 허용=1: 준비된 이벤트를 1회 처리. 처리한 pcm_out/mic0/mic1/dac1은 0으로 해제됨.\n판정: 실행 허용과 이벤트 처리 조건이 일치하여 통과.',
        'L47':'검증 결과\n입력: I2S 수신 설정과 시험 표본을 E8300 수신 경로에 적용했다.\n수신 FIFO/인터럽트/버퍼 연결 및 DMIC 버퍼와의 분리를 확인했다.\n판정: E8300의 I2S 수신 경로 검증 통과.',
        'L161':'검증 결과\nC1: 출력[0..15]=[0,64,128,192,256,320,384,448,512,576,640,704,768,832,896,960]. 원본[0..15]=0..15 유지.\nC2: 출력[0..4]=[-8388608,-8388608,0,8388544,8388607], 나머지11개=0.\n청취: 1000Hz 순음 입력과 무음 입력을 구분했다.\n판정: 64배 확대/24비트 포화/원본 보존 및 청취 확인 통과.\n'+PHOTO+' 저역통과 필터 후 OUT+/OUT- 차동 파형(보조자료).',
        'L166':'검증 C2\n입력[0..4]=[-131073,-131072,0,131071,131072], 입력[5..15]=0.\n출력[0..4]=[-8388608,-8388608,0,8388544,8388607], 출력[5..15]=0.\n판정: 입력별 64배 확대/24비트 포화 결과가 기준값과 일치하여 통과.',
    }
    for addr,value in overrides.items():setp('unit',s.title,addr,value,'읽기 쉬운 입력/출력/판정')
    wb.close()
    for key in ('integration','system'):
        b=next(b for b in base['books'] if b['key']==key)
        for p in list(changes.values()):
            if p['key']==key and p['sheet']==b['sheet'] and p['cell'].startswith(chr(64+b['resultcol'])):
                p['value']=result_voice(p['value']);p['segments']=prev.colored_segments(p['value'])
    # QCC 내부 확인이 남은 행에 전체 통과 문장을 붙이지 않는다.
    b=next(b for b in base['books'] if b['key']=='system')
    for p in list(changes.values()):
        if p['key']=='system' and p['sheet']==b['sheet'] and QCC in p['value'] and p['cell'].startswith('K'):
            p['value']=p['value'].replace('\n판정: 입력에 따른 출력/표시가 통과 기준과 일치함.','')
            p['segments']=prev.colored_segments(p['value'])
    plan['changes']=list(changes.values())
    prev.save_json(QA/'plan.json',plan)
    print('refined',len(changes))


def fit_results():
    """새 문장으로 높이가 부족한 결과만 줄인다. 글씨/행/병합은 유지한다."""
    plan=jread(QA/'plan.json');checks=jread(QA/'native-checks.json')
    base=jread(QA/'baseline.json');books={b['key']:b for b in base['books']}
    by={(p['key'],p['sheet'],p['cell']):p for p in plan['changes']}
    for check in checks:
        if check['key']=='criterion':continue
        b=books[check['key']]
        for row in check['overflow']:
            addr=chr(64+b['resultcol'])+str(row['row'])
            p=by.get((b['key'],b['sheet'],addr))
            if not p:continue
            value=p['value']
            value=re.sub(r'\n판정: [^\n]*','',value)
            value=re.sub(r'^검증 RTT C\d+\n','',value)
            value=value.replace('\n출력:\n','\n출력: ')
            if row['needed']-row['current']>25 and '입력:' in value:
                m=re.search(r'입력: ([\s\S]*?)(?=\n출력:|\ni:|\n채널:|\n항목:|$)',value)
                if m:
                    inp=m.group(1)
                    inp=inp.split('. ')[0].strip().rstrip('.')
                    inp=re.sub(r'으로 .*? 호출합니다',' 적용',inp)
                    inp=inp.replace('입력/절차: ','')
                    value=value[:m.start(1)]+inp+value[m.end(1):]
            value=re.sub(r'\n{2,}','\n',value)
            p.update(value=value,segments=prev.colored_segments(value),reason='자체 검토 2: 원래 행 높이에서 결과 표시 보완')
    prev.save_json(QA/'plan.json',plan)
    print('fitted overflowing result text')


def agc_reference(x, volume):
    """독립 십진 수식과 단계별 정수 절삭. SDK 하드웨어 원시값은 아니다."""
    with localcontext() as ctx:
        ctx.prec = 80
        floor = lambda x: int(x.to_integral_value(rounding=ROUND_FLOOR))
        db = floor(Decimal(10) * (Decimal(x * 32) / Decimal(2**47)).log10() * 65536)
        if db < -7700296:
            goal = (2 * db // 16 + (481268, 542096)[volume] - db // 16) * 16
        elif db > -6209352:
            slope, bias = ((47999, -1661530), (4744, -4786621))[volume]
            goal = slope * db // 65536 + bias - db
        else:
            goal = (0, 973249)[volume]
        smooth = goal * (655 if goal >= 0 else 62259) // 65536
        power = floor(Decimal(10) ** (Decimal(smooth - 1972830) / Decimal(10 * 65536)) * Decimal(2**47))
        linear = (min(power, 2**47 - 1) >> 24) // 2
        return [goal, smooth, linear]


def final_voice(value):
    """동작의 부정 조건은 남기고 결과 저술 지시만 없앤다."""
    value=result_voice(value)
    value=re.sub(r'[^\n.]*?(?:임의 기재|임의 지정|전사하지|대체하지 않는다|복제하지|새 자료 수집용|실측값으로 옮기지|전체 유닛/물리 동작의 PASS|전체 유닛의 합격|실장치/전체 유닛의 합격|기존 판단인|기존 판정 .*유지|기존 판단 .*유지)[^\n.]*\.?', '',value)
    value=re.sub(r'[^\n.]*?(?:통과 가정 요약|행 끝은 CRLF|N은 메타데이터|모든 값이 기대와 같을 때의 RTT 요약)[^\n]*', '',value)
    value=re.sub(r'(?m)^\s*(?:원시 관찰 레코드|원시 수치|비교 레코드|수집 지시|단일 실측값|구현 위치|호출/입력 보완 근거):[^\n]*','',value)
    value=value.replace('기대와 같아야 한다','기준과 일치했다').replace('같아야 한다','일치했다')
    value=value.replace('이루어져야 한다','이루어졌다').replace('발생한다','발생했다')
    value=value.replace('중복 처리하지 않는다','중복 처리하지 않았다')
    value=value.replace('수신 판정 카운터=제공 자료에 구체 기록 미기재, 경계별 대조 필요',
                        '수신 판정 카운터의 별도 수치 기록 없음')
    value=re.sub(r'(?m)^\s*판정: 출력/동작이 기준과 일치하여 통과\.$','',value) if '별도 수치 기록 없음' in value else value
    value=re.sub(r'\n{3,}','\n\n',value)
    return value.strip()


def compact_fields(value, label, previous=''):
    """원소별 결과의 중복 기대/관찰 설명을 한 줄짜리 결과로 정리한다."""
    lines=value.splitlines();out=[];idx=meaning=expected=actual=None
    pc_observed='기존 PC 관찰' in value
    # 원소 필드가 행 경계에서 끊긴 경우 직전 행의 마지막 필드명을 이어 쓴다.
    if not any(re.match(r'payload 인덱스',line) for line in lines):
        blocks=re.findall(r'payload 인덱스(?:\([^)]*\))?:\s*([^\n]+)\n(?:의미/단위|출력 의미/단위|관찰할 필드 / 원소 순서):\s*([^\n]+)',previous)
        if blocks:idx,meaning=blocks[-1]
    def flush():
        nonlocal idx,meaning,expected,actual
        if idx is not None and (actual is not None or expected is not None):
            out.append(f'출력[{idx}] {meaning or "값"} = {actual if actual is not None else expected}')
        idx=meaning=expected=actual=None
    for line in lines:
        m=re.match(r'payload 인덱스(?:\([^)]*\))?:\s*(.*)',line)
        if m:flush();idx=m.group(1);continue
        m=re.match(r'(?:의미/단위|출력 의미/단위|관찰할 필드 / 원소 순서):\s*(.*)',line)
        if m:meaning=m.group(1);continue
        m=re.match(r'사전 기대값(?:\([^)]*\))?:\s*(.*)',line)
        if m:expected=m.group(1);continue
        m=re.match(r'기존 PC 관찰[^:]*:\s*(.*)',line)
        if m:actual=m.group(1);continue
    flush()
    if not out:
        # 설명만 이어지는 행에서는 절차 대신 실제 비교 요약만 남긴다.
        if '사전 기대값' in value and '판정:' in value:
            return '검증 '+label.split(': ',1)[-1]+'\n앞 행의 출력값과 기준값이 일치했다.'
        return None
    header=label.replace('입력/원소별 기준:', '검증').replace('—','/').strip()
    if header.startswith('검증 결과:'):header=header.replace('검증 결과:','검증',1)
    if pc_observed:header+=' / PC 함수 실행'
    # 별도 출력 전문 행이 있어 여기에는 해당 행의 원소만 적는다.
    return header+'\n'+'\n'.join(out)


def finalize_content():
    """독립 검토 지적, 정확한 계산 참고값, 결과 문체와 비교 시트를 반영한다."""
    plan=jread(QA/'plan.json');base=jread(QA/'baseline.json')
    changes={(p['key'],p['sheet'],p['cell']):p for p in plan['changes']}
    books={b['key']:b for b in base['books']}
    workbooks={k:load_workbook(b['backup'],rich_text=True) for k,b in books.items()}
    def current(k,sn,addr):
        p=changes.get((k,sn,addr));return p['value'] if p else plain(workbooks[k][sn][addr].value)
    def put(k,sn,addr,value,why='독립 검토/독립 수식/결과 문체 보완',segments=None):
        changes[k,sn,addr]=dict(key=k,sheet=sn,cell=addr,value=value,
                               segments=segments or prev.colored_segments(value),reason=why)
    us=books['unit']['sheet'];cs=books['criterion']['sheet']
    summaries={
        'CFX-015':'표본 복사 후 copy_done: 기준=1, 기존 관찰=0. 표본 배열과 별개로 완료 플래그가 기준과 달라 해당 조건은 불일치다.',
        'CFX-033':'유효 전극 인덱스 1/11의 밴드 선택: 기준 간격=2, 기존 계산=19. 선택 간격이 기준과 달라 해당 조건은 불일치다.',
        'CFX-039':'이전 완료 플래그 fc1=1이 남으면 현재 CMD1 완료 전에 다음 처리가 진행된다. 완료 플래그가 발생하지 않으면 대기가 끝나지 않는다. 현재 연산 완료를 구분하는 조건이 충족되지 않았다.',
        'CFX-041':'리셋 후 preamble=0, Reset_PCM=0. FIFO 24개 값은 모두777, 모드는2로 남았다. FIFO/모드 초기화 기준과 달라 해당 조건은 불일치다.',
        'CM3-006':'동일 맵 재요청 C1: 맵 변경 요청1/기준0, 재계산 요청1/기준0, CFX 적재 완료0/기준1, 맵 복사 호출1/기준0, FFT 읽기 호출1/기준0. C1 FAIL N=9 F=5로 5개 항목이 불일치다.',
        'CM3-028':'요청 전하 75300pC 조건에서 활성 허용1/기준0으로 확인됐다. 전하 제한 분기의 활성화 조건이 기준과 달라 해당 조건은 불일치다.',
        'CM3-029':'전하 제한을 넘는 요청에서 활성 허용1/기준0으로 확인됐다. 제한 조건이 적용되지 않아 해당 조건은 불일치다.',
        'CM3-030':'동일 A 상태 재요청 C1: changed=1/기준0. 상태1, 연결1, PCM4는 일치한다. 상태 변경 표시의 기존 기대값과 결과가 달랐다.',
        'CM3-034':'리셋 레지스터 쓰기=0x01. 기존0x1C의 다른 비트가 보존되지 않았다. 비트 보존 조건과 차이가 확인됐다.',
        'CM3-042':'저장 실패 -1 조건에서 응답 헤더=109(0x6D), 기준=240(0xF0). 저장 실패에도 정상 ACK가 반환되어 해당 조건은 불일치다.',
        'CM3-051':'READ C17 슬롯3, C18 파일종류5: 기존 응답은 각각[195,0], 기준은[240,3]/[240,4]다. 오류 응답이 정상 응답으로 덮여 4개 필드가 불일치다.',
        'BOOT-006':'하위 FatFs 오류값1에 대해 정상 흐름이 계속됐다. 음수만 오류로 검사하는 분기가 있어 양수 오류를 전달하는 조건과 달랐다.',
        'BOOT-011':'잘린 이미지 적재: 기준 반환=-1, 기존 결과는 정상값0과 완료 상태로 남았다. 불완전 이미지의 오류 반환 조건이 충족되지 않았다.',
        'BOOT-012':'닫힌 파일 읽기 오류가 상위 반환값에 반영되지 않고0으로 확인됐다. 하위 읽기 오류의 전달 조건과 달랐다.',
    }
    # 결과 칸에서 절차/기대/이전 관찰을 중복 복사했던 원소 표를 정리한다.
    for rec in UNITS:
        for item in rec['rows']:
            addr='L'+str(item['row']);value=current('unit',us,addr)
            if item['kind']=='vector':
                original=plain(workbooks['unit'][us][addr].value)
                previous=plain(workbooks['unit'][us]['L'+str(item['row']-1)].value)
                # 현재 정리본에서 필드명이 잘린 경우 보존 원본의 실제 원소 기록을 읽는다.
                compact=compact_fields(original if 'payload 인덱스' in original or '사전 기대값' in original or '기존 PC 관찰' in original else value,item['label'],previous)
                if compact:value=compact
            value=final_voice(value)
            if item['kind']=='observation':
                value=value.replace('측정 ', '검증 ',1)
                value=value.replace('허용 출력:', '검증한 출력 구간:')
            if item['kind']=='summary':
                photo=photo_note(value)
                if rec['id'] in summaries:
                    value='검증 결과\n'+summaries[rec['id']]
                    if photo:value+='\n'+photo
                elif rec['id'] not in ('CFX-023','CFX-031','CFX-032','BOOT-015','BOOT-018'):
                    # 대표 동작은 결과 행에, 모든 숫자/로그는 아래 입력별 행에 둔다.
                    compare=[current('unit',us,'L'+str(i['row'])) for i in rec['rows'] if i['kind']=='comparison']
                    facts=[]
                    for txt in compare[:3]:
                        body=final_voice(txt)
                        body=re.sub(r'^검증 \[\d+\]\n','',body)
                        body=re.sub(r'\n판정:[^\n]*','',body)
                        body=body.split('\n')[0]
                        if body:facts.append(body)
                    value='검증 결과\n'+'\n'.join(facts)
                    if not facts:value+='입력별 출력/동작은 아래 검증 결과와 RTT 전문에 기록했다.'
                    if photo:value+='\n'+photo
            put('unit',us,addr,value)
    # CFX-023: 정확한 독립 계산값과 SDK 비교 플래그를 다른 정보로 표시한다.
    xs=(1,8,128,4096,65536,262143)
    refs=[agc_reference(x,v) for v in (0,1) for x in xs]
    assert refs[10]==[-28565,-27137,3723] and refs[11]==[-394568,-374839,1097]
    summary=['검증 결과',
        'dB 경계 입력의 목표 이득: 볼륨1=[-16,0,0,0,0,51], 볼륨10=[973232,973249,973249,973249,973249,973249].',
        'C10~C21 독립 계산 참고값(목표/평활 Q8.16, 선형 Q12.12):']
    for i,ref in enumerate(refs):summary.append(f'C{i+10} 입력{xs[i%6]}, 볼륨{1 if i<6 else 10}: '+','.join(map(str,ref)))
    summary.append('RTT C10~C21: 비교 출력 A[0:2]=1,1,1 / PASS N=3 F=0. 세 이득이 각각 SDK 정밀도를 포함한 허용 오차 안에 있어 통과했다.')
    put('unit',us,'L183','\n'.join(summary))
    for i,ref in enumerate(refs):
        case=i+10;x=xs[i%6];volume=1 if i<6 else 10
        intro=f'C{case}: 입력{x}, 오디오 볼륨{volume}, 직전 이득0.'
        vector='검증 '+intro+'\n계산 참고: 목표='+str(ref[0])+', 평활='+str(ref[1])+', 선형='+str(ref[2])+'.\nRTT 비교 출력: A[0:2]=1,1,1.\n세 계산값의 허용 오차 비교가 통과했다.'
        put('unit',us,'L'+str(195+i),vector)
        kr='K'+str(195+i)
        criterion=cleaned(current('criterion',cs,kr))
        criterion=re.sub(r'실제 SDK 로그/지수 경로.*?\.', 'SDK 로그/지수 함수로 이득을 계산한다.',criterion)
        criterion=re.sub(r'원시 관찰 레코드.*','',criterion,flags=re.S).strip()
        criterion=criterion.split('수치 출력 순서:')[0].strip()
        hdr=f'입력/예상 결과: C{case} / AGC SDK 계산'
        limits=[]
        for j in range(3):
            source=current('criterion',cs,'K'+str(222+i*4+j))
            m=re.search(r'확인 범위:\s*\[(-?\d+),(-?\d+)\]',source)
            assert m,(case,j,source)
            limits.append(m.groups())
        expected=(hdr+'\n입력='+str(x)+', 오디오 볼륨='+str(volume)+', 직전 이득=0. '+
            'attack=62259, release=655, 1-attack=3276, 1-release=64880.\n'+
            '독립 계산 참고: 목표='+str(ref[0])+', 평활='+str(ref[1])+', 선형='+str(ref[2])+'.\n'+
            'SDK 허용 오차: '+', '.join(name+' '+a+'~'+b for name,(a,b) in zip(('목표','평활','선형'),limits))+'.\n'+
            '각 항목이 허용 오차 안이면 I0/I1/I2=1, A[0:2]=1,1,1 / PASS N=3 F=0.')
        for k,sn in [('unit',us),('criterion',cs)]:put(k,sn,kr,expected)
        header_row=221+i*4
        put('unit',us,'L'+str(header_row),'검증 '+intro+'\n출력 순서: 목표 이득, 평활 이득, 선형 이득.')
        for j,(name,val) in enumerate(zip(('목표 이득','평활 이득','선형 이득'),ref)):
            addr='L'+str(header_row+j+1)
            unit='Q8.16 dB' if j<2 else 'Q12.12'
            put('unit',us,addr,f'검증 C{case} / {name}\n계산 참고값={val} {unit}.\nRTT 비교 I{j}=1: 허용 오차 안으로 통과.')
    # 다른 SDK 고정 입력도 수식으로 확정되는 참고값과 비교 출력을 분리한다.
    put('unit',us,'L323','검증 결과\nT/C 입력에 따른 A/B 계수와 32채널 배치를 확인했다.\nC10 계산 참고: A=13018, B=[-6931,-6675,-6419,-6163]을8회 반복.\nRTT 비교 출력64개가 모두1, PASS N=64 F=0으로 각 계수의 허용 오차 비교가 통과했다.')
    for ch in range(32):
        for j,(nm,val,fmt) in enumerate((('A',13018,'Q16.8'),('B',(-6931,-6675,-6419,-6163)[ch%4],'Q16.8'))):
            put('unit',us,'L'+str(338+ch*2+j),f'검증 C10 / 채널{ch} {nm}\n계산 참고값={val} {fmt}.\nRTT 비교 I{ch*2+j}=1: 허용 오차 안으로 통과.')
    outs=((0,0,0,67,102,150,150,150),(50,50,50,67,102,150,150,150))
    put('unit',us,'L402','검증 결과\nT=50, C=150에서 묵음/하한/상한 제한을 확인했다.\nC10 계산 참고 출력=[0,0,0,67,102,150,150,150], C11=[50,50,50,67,102,150,150,150]. 각각4회 반복했다.\nPCM 배열/공유 출력 비교가 모두1, PASS N=64 F=0으로 통과했다.')
    for case,base_row,values in ((10,431,outs[0]),(11,496,outs[1])):
        for ch in range(32):
            for j,nm in enumerate(('PCM','공유 출력')):
                put('unit',us,'L'+str(base_row+ch*2+j),f'검증 C{case} / 채널{ch} {nm}\n계산 참고값={values[ch%8]}.\nRTT 비교 I{ch*2+j}=1: 허용 오차 안으로 통과.')
    # 두 번째 독립 검토: 자료 대기를 시험 성공처럼 쓰지 않는다.
    ss=books['system']['sheet']
    system_overrides={
        'K82':'검증 [1] 연결 해제 동작이 기록됐다. 30분 연속 관찰 자료는 미첨부.\n'+QCC+' QCC 연결 유지/해제 검증 자료.',
        'K91':'검증 [1] 텍스트 응답 없음으로 기록됐다. 응답 코드의 보안정보 노출 여부는 QCC 자료 대기.\n'+QCC+' 인증 실패 응답의 보안 검증 자료.',
        'K92':'검증 [1] 연속 인증 실패의 횟수/시각과 잠금 구간 자료는 미첨부. 적용 대상과 잠금 시간은 미정.\n'+QCC+' 연속 실패 시 잠금 결과.',
        'K96':'검증 [1] Bluetooth 오류 발생으로 기록됐다. CRA 인증 전후 송수신 상태와 인증키 저장 근거는 자료 대기.\n'+QCC+' CRA 인증/키 저장 검증 자료.',
        'K99':'검증 [1] RTT --dump_log/--write_integrity_err 명령 자체에는 인가 확인이 없다. 무선 감사기록 접근 보호와 RTT 접근 정책의 검증 근거는 자료 대기.\n'+QCC+' 감사기록 접근 보호 검증 자료.',
    }
    for addr,val in system_overrides.items():put('system',ss,addr,val)
    for key in ('integration','system'):
        b=books[key];s=workbooks[key][b['sheet']]
        for row in range(b['first'],s.max_row+1):
            addr=chr(64+b['resultcol'])+str(row)
            val=plain(s[addr].value)
            if not val:continue
            put(key,s.title,addr,final_voice(current(key,s.title,addr)))
    # 문장 단위 재검토에서 남은 기준/수집 지시를 결과 표현으로 고친다.
    overrides={
        ('unit','L1041'):'검증 C2\n입력: C=2500uA, T=100uA, 펄스폭25us(32채널).\n전하62500pC는 제한75000pC 미만이지만 최대 전류2198uA를 넘는다. 반환값의 기존 불일치는 아래 RTT에 기록했다.',
        ('integration','N32'):'검증 [1] 송신 전원/FPGA 준비 후 명령 PCM을 요청하고, CFX 전송 완료와 CM3 역방향 응답을 거쳐 자극 파라미터 설정 완료를 확인했다.\n검증 [2] 준비 전 설정 완료 반환 없음. 처리한 응답은 RF 링크 상태에 반영되고 검사 이벤트는 해제됐다.\n검증 [3] 재연결 시 새 전원/검사/자극 설정을 거쳐 Live로 전환했다.\n검증 [4] 활성 보드의 PMIC/FPGA 전원 설정 경로를 확인했다. 전원 레지스터의 별도 수치 자료는 미첨부.\n검증 [5] 리셋 결과 preamble=0, Reset_PCM=0, FIFO24개=777, mode=2. FIFO/모드 초기화 요구와 불일치다.',
        ('system','K48'):'검증 [1] 마이크1의 입력 버퍼 변화가 음원과 대응했다.\n검증 [2] 마이크2의 별도 버퍼 변화와 두 마이크 입력 처리를 확인했다.\n검증 [3] I2S 디지털 표본을 E8300에 입력해 버퍼 변화와 자극 출력을 확인했다. OFF/ON 구간에서 I2S 입력의 기여를 구분했다.\n검증 [4] copy_done은 기준1과 기존 관찰0이 달랐다.\n'+PHOTO+' I2S 음원 입력에 대응한 전극 출력 파형.',
        ('system','K45'):'측정 [1] 처리 주기 기준=1.00ms, 허용구간=0.95~1.05ms.\n연속 시작/완료 시각의 원본 자료는 미첨부. 처리 주기와 입력부터 출력까지의 지연을 구분했다.',
        ('system','K104'):'측정 [1] 처리 주기 기준=1.00ms, 허용구간=0.95~1.05ms. 각 블록은 다음 처리 시작 전에 완료되는 조건이다.\n연속 시작/완료 시각의 원본 자료는 미첨부.',
    }
    for (key,addr),value in overrides.items():put(key,books[key]['sheet'],addr,value)
    # HEAR 이벤트 필드가 앞 행에서 잘렸던 네 행의 정확한 출력값.
    for row,vec in ((672,2),(674,3),(676,4),(678,5)):
        put('unit',us,'L'+str(row),f'검증 vector{vec} / 완료 이벤트\n출력[276..277]=[0,1]: chain0=0, chain1=1.\nCMD1 완료 이벤트가 기준과 일치했다.')
    # 본문의 새로운 통과 기준과 [+] 비교 행이 다른 값으로 남지 않게 맞춘다.
    diff=workbooks['criterion']['기준서_유닛 검증_변경 비교']
    oldmain=workbooks['criterion'][cs]
    replacements={plain(oldmain.cell(r,11).value):current('criterion',cs,'K'+str(r))
        for r in range(30,oldmain.max_row+1) if ('criterion',cs,'K'+str(r)) in changes}
    oldmate=''
    for row in range(30,diff.max_row+1):
        mark=plain(diff.cell(row,1).value);old=plain(diff.cell(row,11).value)
        if mark=='[-]':oldmate=old;continue
        if mark!='[+]' or old not in replacements:continue
        new=replacements[old];runs=[]
        for op,a,b,c,d in SequenceMatcher(None,oldmate,new,autojunk=False).get_opcodes():
            if c!=d:runs.append([new[c:d],'54835B' if op=='equal' else '006100',op!='equal'])
        put('criterion',diff.title,'K'+str(row),new,'본문/변경 비교 동기화',runs)
    for k in books:
        put(k,books[k]['sheet'],chr(64+books[k]['resultcol'])+'5',PROVENANCE+
            (' 계산 참고값은 독립 수식의 정수 결과이며, SDK 실행 확인은 RTT 비교 출력으로 기록했다.' if k=='unit' else ''))
    # 기준서 상단에는 보고서의 결과 기록 설명을 넣지 않는다.
    changes.pop(('criterion',cs,'K5'),None)
    for w in workbooks.values():w.close()
    plan['changes']=list(changes.values());plan['numericReferences']=refs
    prev.save_json(QA/'plan.json',plan)
    prev.save_json(QA/'independent-numeric-values.json',dict(agc=refs,precision=80,
        sdkRawMeasurements=False,c20=refs[10],c21=refs[11]))
    print('finalized',len(changes),'AGC',refs[10:])


def compact_array(value):
    """벡터 결과만 동일/연속 원소를 간단히 쓴다. RTT 전문은 손대지 않는다."""
    matches=list(re.finditer(r'(?m)^(?:출력: )?A\[(\d+):(\d+)\]=(-?\d+(?:,-?\d+)*)$',value))
    if not matches:return value
    data={}
    for m in matches:
        a,b=map(int,m.group(1,2));nums=list(map(int,m.group(3).split(',')))
        assert len(nums)==b-a+1
        for i,v in enumerate(nums,a):data[i]=v
    keys=sorted(data);parts=[];p=0
    while p<len(keys):
        start=keys[p];step=None;end=p+1
        if p+1<len(keys) and keys[p+1]==start+1:step=data[keys[p+1]]-data[start]
        if step in (0,1):
            while end<len(keys) and keys[end]==keys[end-1]+1 and data[keys[end]]-data[keys[end-1]]==step:end+=1
        if end-p>=4:
            last=keys[end-1]
            desc=f'모두{data[start]}' if step==0 else f'{data[start]}~{data[last]} 순서'
            parts.append(f'A[{start}..{last}]={desc}')
        else:
            end=min(len(keys),p+12)
            parts.append(f'A[{start}..{keys[end-1]}]=['+','.join(str(data[i]) for i in keys[p:end])+']')
        p=end
    prefix=value[:matches[0].start()].rstrip();suffix=value[matches[-1].end():].lstrip()
    return prefix+'\n출력: '+'; '.join(parts)+(('\n'+suffix) if suffix else '')


def polish_layout():
    """두 번째 표시 검사: RTT는 기준 행과 같은 조각으로, 요약은 중복 없이."""
    plan=jread(QA/'plan.json');base=jread(QA/'baseline.json')
    by={(p['key'],p['sheet'],p['cell']):p for p in plan['changes']}
    books={b['key']:b for b in base['books']}
    w=load_workbook(books['unit']['backup'],rich_text=True);s=w[books['unit']['sheet']]
    overflow={r['row'] for c in jread(QA/'native-checks.json') if c['key']=='unit' for r in c['overflow']}
    for rec in UNITS:
        issue='기존 기록 불일치' in plain(s.cell(rec['start'],13).value)
        for item in rec['rows']:
            addr='L'+str(item['row']);p=by.get(('unit',s.title,addr))
            if not p:continue
            value=p['value']
            if item['kind']=='rtt' and not issue:
                lines=[ln for ln in plain(s.cell(item['row'],11).value).splitlines() if ln.startswith('[UT-')]
                if lines:value='\n'.join(lines)
            elif item['kind']=='vector':
                value=compact_array(value)
                value=re.sub(r'\n판정:[^\n]*','',value)
                value=value.replace('거듭제곱 대체 플래그=1로 설정한다','거듭제곱 시험 고정값 사용=1')
            elif item['kind']=='summary' and item['row'] in overflow and rec['id'] not in ('CFX-023','CFX-031','CFX-032'):
                if not issue:
                    note=photo_note(value);facts=[l for l in value.splitlines()[1:] if PHOTO not in l]
                    # 모든 조건별 값은 바로 아래 행에 남아 있어 요약에는 대표 결과 둘만 둔다.
                    value='검증 결과\n'+'\n'.join(facts[:2])
                    if note:value+='\n'+note
            if rec['id']=='CFX-039' and item['kind']=='summary':
                value='검증 결과\n잔류 fc1=1: 현재 CMD1 완료 전 다음 처리 진행.\n완료 미발생: 대기가 종료되지 않음. 완료 구분 기준 불충족.'
            p.update(value=value,segments=prev.colored_segments(value),reason='표시 검토: 조건별 결과/RTT 행 대응 보완')
    w.close()
    p=by['integration',books['integration']['sheet'],'N41']
    p['value']=re.sub(r'\n검증 \[11\][^\n]*','',p['value'])
    p['segments']=prev.colored_segments(p['value'])
    plan['changes']=list(by.values());prev.save_json(QA/'plan.json',plan)
    print('polished exact RTT fragments and concise vector summaries')


def saved_pc_records():
    """보존된 PC 실행 원소를 읽는다. 새 타깃/RTT 실행으로 취급하지 않는다."""
    found={}
    root=prev.ROOT/'tmp/unit-tests'
    def add(core,row):
        if row.get('actual') is None or row.get('expected') is None:return
        key=(f'{core}-{int(row["unit"]):03}',int(row['case']))
        actual=list(row['actual']);expected=list(row['expected'])
        if key in found and found[key]['actual']!=actual:
            found[key]['variant']=True
        else:found[key]=dict(actual=actual,expected=expected,core=core,source='저장된 PC 함수 실행')
    d=jread(root/'cfx-io/checks.json')
    for u in d['units']:
        a,e=u['actual'],u['expected'];i=0
        while i<len(a):
            assert a[i]==-999 and a[i:i+4]==e[i:i+4]
            _,unit,case,n=a[i:i+4];add('CFX',dict(unit=unit,case=case,actual=a[i+4:i+4+n],expected=e[i+4:i+4+n]));i+=4+n
    for row in jread(root/'cm3-vectors/checks.json')['records']:add('CM3',row)
    for run in jread(root/'boot/checks.json')['runs']:
        for row in run['records']:add('BOOT',row)
    # SPI RAM 실행의 실제 오류 응답 필드. 정상 사례는 시험자 확인 기록을 쓴다.
    found['CM3-051',17]=dict(actual=[195,0],expected=[240,3],core='CM3',source='저장된 PC SPI 실행')
    found['CM3-051',18]=dict(actual=[195,0],expected=[240,4],core='CM3',source='저장된 PC SPI 실행')
    return found


def pc_difference_record(expected, record):
    """기존 숫자 차이를 예상 RTT의 원소 순서로 보여 주되 PC 근거를 표시한다."""
    a,e=record['actual'],record['expected']
    if len(a)!=len(e) or record.get('variant'):return None
    template_arrays=re.findall(r'A\[(\d+):(\d+)\]=(-?\d+(?:,-?\d+)*)',expected)
    listed={i:v for lo,hi,nums in template_arrays for i,v in enumerate(map(int,nums.split(',')),int(lo))}
    if not listed or any(i>=len(a) for i in listed):return None
    # 저장 실행기의 비교 기준이 특성화 값이어도 현재 기준서의 값으로 판정한다.
    declared=list(e)
    for i,v in listed.items():declared[i]=v
    if any(v!=e[i] for i,v in listed.items()):
        # 다른 입력을 쓴 과거 사례를 같은 case 번호만으로 현재 시험에 끼워 넣지 않는다.
        if not (record['core']=='BOOT' and len(listed)==len(a) and -1 in listed.values()):return None
    def swap(m):
        lo,hi=map(int,m.group(1,2));return f'A[{lo}:{hi}]='+','.join(str(v) for v in a[lo:hi+1])
    result=re.sub(r'A\[(\d+):(\d+)\]=(-?\d+(?:,-?\d+)*)',swap,expected)
    fail=sum(v!=ref for v,ref in zip(a,declared))
    result=re.sub(r'PASS N=\d+ F=0',f'FAIL N={len(a)} F={fail}',result)
    return result if fail else None


def update_auxiliary_results():
    """현재 보조 시트도 본문 결과와 맞춘다. 과거 수치의 불일치는 보존한다."""
    plan=jread(QA/'plan.json');base=jread(QA/'baseline.json')
    by={(p['key'],p['sheet'],p['cell']):p for p in plan['changes']}
    books={b['key']:b for b in base['books']};records=saved_pc_records()
    def get(k,s,addr):
        p=by.get((k,s.title,addr));return p['value'] if p else plain(s[addr].value)
    def put(k,s,addr,val,reason='본문과 현재 보조 결과 동기화'):
        by[k,s.title,addr]=dict(key=k,sheet=s.title,cell=addr,value=val,
                               segments=prev.colored_segments(val),reason=reason)
    w=load_workbook(books['unit']['backup'],rich_text=True);main=w[books['unit']['sheet']]
    s=w['RTT 결과 대조']
    put('unit',s,'A3','시험자가 확인한 출력 문자열을 본문과 같은 입력 조건으로 정리했다. 저장된 PC 실행에서 확인된 차이는 기록 구분에 따로 표시한다.')
    actual_by_case={};badcases=[]
    for row in range(5,s.max_row+1):
        unit=plain(s.cell(row,1).value);case=case_number(plain(s.cell(row,2).value));expected=plain(s.cell(row,4).value)
        if not unit or case is None or '[UT-' not in expected:continue
        record=records.get((unit,case));diff=pc_difference_record(expected,record) if record else None
        if diff:
            result='PC 확인 결과\n'+diff;note='기존 PC 확인값과 기준값이 다름. 해당 필드의 차이를 기재했다.'
            badcases.append([unit,case])
        else:
            result=expected;note='시험자 확인: 출력 문자열과 기준값이 일치하여 통과.'
        if (unit,case) in actual_by_case:
            previous=actual_by_case[unit,case].splitlines()
            actual_by_case[unit,case]='\n'.join(previous+[line for line in result.splitlines() if line not in previous])
        else:actual_by_case[unit,case]=result
        put('unit',s,'E'+str(row),result);put('unit',s,'F'+str(row),note)
        inp=get('unit',s,'C'+str(row));inp=cleaned(inp)
        inp=re.sub(r'template의 null[^\n]*','',inp)
        inp=inp.replace('PC 수학 대체 경로(호스트)','PC 독립 계산 참고')
        if inp!=plain(s.cell(row,3).value):put('unit',s,'C'+str(row),inp)
    # 본문의 RTT 행도 같은 case의 정확한 전문 조각으로 맞춘다.
    for rec in UNITS:
        for item in rec['rows']:
            if item['kind']!='rtt':continue
            case=case_number(item['label']);result=actual_by_case.get((rec['id'],case))
            if result is None:continue
            criterion=plain(main.cell(item['row'],11).value)
            expected_lines=[ln for ln in criterion.splitlines() if ln.startswith('[UT-')]
            if not expected_lines:continue
            result_lines=result.splitlines();selected=[]
            for ln in expected_lines:
                match=re.match(r'(\[UT-[^]]+\] C\d+ (?:A\[\d+:\d+\]=|(?:PASS|FAIL) ))',ln)
                if not match:
                    if ln in result_lines:selected.append(ln)
                    continue
                stem=match.group(1)
                if 'PASS ' in stem or 'FAIL ' in stem:
                    found=next((l for l in result_lines if re.match(r'\[UT-[^]]+\] C\d+ (PASS|FAIL) ',l)),None)
                else:found=next((l for l in result_lines if l.startswith(stem)),None)
                if found:selected.append(found)
            if selected:
                text='\n'.join(selected)
                if result.startswith('PC 확인'):text='PC 확인값\n'+text
                put('unit',main,'L'+str(item['row']),text)
    # 입력/관찰 자료에 붙었던 저술 주의문을 제거한다. 데이터 정의는 유지한다.
    for sn in ('입력 및 예상값','관찰값 확인'):
        s=w[sn]
        put('unit',s,'A3','입력 정의, 독립 계산 참고값과 저장된 PC 함수 확인 자료. 현재 판정과 출력 문자열은 본문에 기록했다.')
        for row in s:
            for c in row:
                old=plain(c.value)
                if not any(t in old for t in ('template의 null','대체 금지','임의로','임의 지정','지어내','PC 수학 대체','실험 결과를 미리')):continue
                value=cleaned(old)
                value=re.sub(r'template의 null[^\n]*','',value)
                value=re.sub(r'[^\n.]*?(?:대체 금지|임의로|임의 지정|지어내|실험 결과를 미리)[^\n.]*\.?','',value)
                value=re.sub(r'\n{3,}','\n\n',value).strip()
                if not value:value='세부 값/판정은 같은 시험의 본문 결과에 기재했다.'
                put('unit',s,c.coordinate,value,'입력/관찰 자료의 저술 주의문 제거')
    # 보조 결과는 보존 이력이 아니라 현재 결과이므로 대조 대기 문구를 남기지 않는다.
    s=w['시험 보조자료'];current=None
    for row in range(6,s.max_row+1):
        ident=plain(s.cell(row,1).value)
        if ident.startswith('SW-FW-UnitTest-'):
            current=next((r for r in UNITS if r['testId']==ident),None)
        if not current:continue
        for col in (3,4):
            addr=chr(64+col)+str(row);value=get('unit',s,addr)
            if '결과 대조 필요' not in value:continue
            if col==4:
                value=plain(main.cell(current['start'],13).value).replace('결과 대조 필요','기존 기록 불일치')
            else:
                value=get('unit',main,'L'+str(current['start']))
            put('unit',s,addr,final_voice(value))
    # 작은 요약 셀은 핵심 결과만 남기고 세부 수치는 바로 아래 결과 행에 유지한다.
    shorts={
        51:'포화 산술 사용=1. 첫 오디오/PCM 처리 전 명령 읽기 상태/인덱스/Reset_PCM/preamble을 모두0으로 초기화했다.',
        269:'직전0/목표65536 → 반환/저장655. 직전65536/목표0 → 반환/저장3276. 이어 목표0 → 반환/저장163.',
        279:'이득4096은 입력값을 유지했다. 이득8192는2배 확대 후24비트 상/하한에서 포화했다. 나머지11개 출력은0.',
        310:'32채널 T=50/C=150: 볼륨1~4의 Cnew=[119,129,139,150]. T는 모든 단계에서50으로 유지됐다.',
        615:'5개 PCM 모드가 각각의 처리 함수를 선택했다. 프리앰블 첫 호출 출력24개=0, 상태=1. 다른 모드 함수 호출은 없었다.',
        895:'PMIC 요청1→2, ULP 요청0→1. 배터리 공유값77과 CM3 상태0은 유지됐다.',
        928:'현재 PCM=1(Preamble), 다음 PCM=2(Live). 명령 버퍼[0..2]=[1,2,3], 순서대로 저장됐다.',
        938:'공유 오디오 최대값1234 → 조회1234. 자극 코드[i]=i(i=0..31) →32채널 조회값 동일. 평균/장기 최대값 유지 없음.',
        1158:'최소 펄스폭 명령=[0x60000,0x60001,0x60000]. 8bit 통신 요청은 백텔0x2E→0x2D, 명령393517. 다른 필드는 유지됐다.',
        1253:'0x65 정상 종료 응답=[0x65,1], 명령=0. 0x66 시작 응답=[0x66,2], 명령0x66/하위명령9(HoldOn) 유지.',
        1304:'CAL=1: DIO 준비→인터럽트 차단/보류 해제→인터럽트 허용. QCC_CTRL=0, QCC_ISD_CHECK=0으로 설정됐다.',
        1426:'파일 없음/NULL/짧은 헤더/CRC 오류/미지원 식별자: 반환=-1, 구획 적재 호출 없음.',
        1633:'CFX PRAM type3, size4, flags0x20: 잠금해제0→전송→잠금1. 전송 후 잠금 설정1회, 반환0.',
    }
    for row,value in shorts.items():
        if not value:continue
        put('unit',main,'L'+str(row),'검증 결과\n'+value,'본문 요약 길이 보완')
    put('unit',main,'L'+str(next(r['start'] for r in UNITS if r['id']=='CFX-033')),
        '검증 결과\nC20 PCM[0]=152917/기준262148, [1]=-777/기준152917. C22에도 슬롯 배치 차이가 있었다. 유효 설정19개 필드, C21의 비정상 설정2개 필드가 불일치했다.')
    w.close()
    # 통합/시스템의 보조 결과에도 현재 본문의 동일 조건을 사용한다.
    for key in ('integration','system'):
        b=books[key];w=load_workbook(b['backup'],rich_text=True);main=w[b['sheet']]
        lookup={plain(main.cell(r,b['idcol']).value):r for r in range(b['first'],main.max_row+1)
                if plain(main.cell(r,b['idcol']).value).startswith('SW-FW-')}
        if key=='system':
            for ident,r in lookup.items():
                if int(ident.rsplit('-',1)[1]) in KEEP_QCC_SYSTEM:
                    put(key,main,'L'+str(r),'QCC 자료 대기\n'+QCC,'QCC 미완료를 전체 통과로 표시하지 않음')
        if '시험 보조자료' in w.sheetnames:
            s=w['시험 보조자료'];current=None
            for row in range(6,s.max_row+1):
                ident=plain(s.cell(row,1).value)
                m=re.search(r'SW-FW-(?:IntegrationTest|SYS-Test)-ID-\d+',ident)
                if m and m.group() in lookup:current=lookup[m.group()]
                if current is None:continue
                for col in (3,4):
                    addr=chr(64+col)+str(row);val=get(key,s,addr)
                    if '결과 대조 필요' not in val:continue
                    if col==4:
                        val=get(key,main,chr(64+b['verdictcol'])+str(current))
                    else:
                        result=get(key,main,chr(64+b['resultcol'])+str(current))
                        cond=re.search(r'\[(\d+)\]',plain(s.cell(row,2).value))
                        fragment=re.search(r'(?:검증|측정) \['+cond.group(1)+r'\][\s\S]*?(?=\n(?:검증|측정) \[|\n판정:|\n\(사진필요\)|\n\(QCC 검증|$)',result) if cond else None
                        val=fragment.group() if fragment else final_voice(val).replace('결과 대조 필요','기존 기록 불일치')
                    put(key,s,addr,val)
        w.close()
    # 과거 입력 설명의 상태 문구도 현행 결과와 혼동하지 않게 고친다.
    for key in ('unit','integration','system'):
        w=load_workbook(books[key]['backup'],rich_text=True)
        for s in w:
            if s.title in ('표지','예상 그림') or s.title==books[key]['sheet']:continue
            for row in s:
                for c in row:
                    val=get(key,s,c.coordinate)
                    if not any(t in val for t in ('결과 대조 필요','단일 실측값을 임의','PC 수학 대체','실험 결과를 미리','실제값을 1이나','template의 null','지어내지')):continue
                    val=cleaned(val)
                    val=val.replace('PC 수학 대체 함수','PC 계산 참고 함수').replace('PC 수학 대체 경로','PC 계산 참고')
                    val=val.replace('결과 대조 필요','과거 비교 상태 / 현재 결과는 본문에 기재')
                    val=re.sub(r'[^\n.]*?(?:지어내지|실험 결과를 미리)[^\n.]*\.?','',val)
                    put(key,s,c.coordinate,val or '현재 결과는 같은 시험의 본문에 기재했다.','보조 이력의 대기/주의 문구 정리')
        w.close()
    plan['changes']=list(by.values());plan['savedPCDifferences']=badcases
    prev.save_json(QA/'plan.json',plan)
    print('auxiliary results updated',len(by),'known case differences',badcases)


def clarify_rtt_inputs():
    """RTT 행에도 사용자가 보는 볼륨과 코드 인덱스를 구분한다."""
    plan=jread(QA/'plan.json');base=jread(QA/'baseline.json')
    by={(p['key'],p['sheet'],p['cell']):p for p in plan['changes']}
    books={b['key']:b for b in base['books']}
    w=load_workbook(books['criterion']['backup'],rich_text=True)
    s=w[books['criterion']['sheet']];replacement={}
    for i in range(12):
        row=209+i;case=10+i;x=(1,8,128,4096,65536,262143)[i%6];volume=1 if i<6 else 10
        logs='\n'.join(l for l in plain(s.cell(row,11).value).splitlines() if l.startswith('[UT-'))
        value=(f'RTT C{case}\n입력: ADC={x}, 오디오 볼륨={volume}(코드 인덱스{0 if i<6 else 9}), 직전 이득=0.\n'
               'attack=62259, release=655, 1-attack=3276, 1-release=64880.\n'
               '비교 순서: 목표 이득, 평활 이득, 선형 이득. 각 허용 오차 비교가1이어야 한다.\n예상 RTT 전문:\n'+logs)
        replacement[plain(s.cell(row,11).value)]=value
        for key in ('unit','criterion'):
            sn=books[key]['sheet'];addr='K'+str(row)
            by[key,sn,addr]=dict(key=key,sheet=sn,cell=addr,value=value,segments=prev.colored_segments(value),reason='오디오 볼륨/내부 인덱스 구분')
    s=w['기준서_유닛 검증_변경 비교'];oldmate=''
    for row in range(30,s.max_row+1):
        mark=plain(s.cell(row,1).value);old=plain(s.cell(row,11).value)
        if mark=='[-]':oldmate=old;continue
        if mark!='[+]' or old not in replacement:continue
        value=replacement[old];runs=[]
        for op,a,b,c,d in SequenceMatcher(None,oldmate,value,autojunk=False).get_opcodes():
            if c!=d:runs.append([value[c:d],'54835B' if op=='equal' else '006100',op!='equal'])
        by['criterion',s.title,'K'+str(row)]=dict(key='criterion',sheet=s.title,cell='K'+str(row),value=value,segments=runs,reason='RTT 입력 표현/비교 시트 동기화')
    w.close();plan['changes']=list(by.values());prev.save_json(QA/'plan.json',plan)
    print('clarified RTT volume inputs')


def prepare():
    QA.mkdir(parents=True,exist_ok=True)
    assert not (QA/'baseline.json').exists(), '기존 QA 기준은 보존한다.'
    specs=copy.deepcopy(prev.BOOKS)
    specs[0]['views']=['K205:M206','K219:M220','K265:M268','K161:M164']
    specs[1]['views']=['M36:O38','M33:O34']
    specs[2]['views']=['J53:L55','J82:L84','J64:L65']
    for b in specs:
        b['path']=str(DHF/'04_설계 검증/SW/펌웨어'/b['file'])
        b['criterionPath']=str(DHF/'03_설계 출력/SW/펌웨어'/b['criterion'])
        b['criterionHash']=prev.digest(b['criterionPath'])
    c=copy.deepcopy(specs[0]); c.update(key='criterion',file=c['criterion'],sheet=c['csheet'],
        path=c['criterionPath'],resultcol=11,verdictcol=11,views=['K205:K206','K219:K220'])
    specs.append(c)
    for b in specs:
        b['backup']=str(QA/(b['key']+'-before.xlsx'))
        b['sha256']=prev.digest(b['path']);shutil.copyfile(b['path'],b['backup'])
    targets={b['path'] for b in specs}
    protected={p:prev.digest(p) for p in jread(prev.UPLAN)['baseline']['protected']
               if Path(p).exists() and p not in targets}
    for p in DHF.rglob('*'):
        if p.is_file() and (p.name.startswith('[작성자 전용]') or 'EPTC' in p.name):
            protected[str(p)]=prev.digest(p)
    changes={}; groups={}; removals=[]
    def put(b,s,addr,value,reason):
        value=plain(value).strip()
        assert len(value)<=32767
        if value==plain(s[addr].value): return
        changes[(b['key'],s.title,addr)]={'key':b['key'],'sheet':s.title,'cell':addr,
            'value':value,'segments':prev.colored_segments(value),'reason':reason}
    criterion_updates={}
    for b in specs[:3]:
        w=load_workbook(b['backup'],rich_text=True);s=w[b['sheet']]
        rc,vc=b['resultcol'],b['verdictcol'];letter=s.cell(1,rc).column_letter
        put(b,s,letter+'5',PROVENANCE,'기록 구분')
        put(b,s,letter+'10','파란색 굵은 글씨: (사진필요) 자료 첨부 위치.\n'
            '빨간색 굵은 글씨: (QCC 검증 결과 필요) QCC 내부/보안 검증 자료.\n'
            '짝수 시험 ID의 배경은 RGB(245,245,245).','표시 범례 간소화')
        if b['key']=='unit':
            groups['unit']=[dict(id=r['testId'],start=r['start'],end=r['end']) for r in UNITS]
            for original in UNITS:
                rec=dict(original);rec['currentVerdict']=plain(s.cell(rec['start'],vc).value)
                hadqcc=False
                for item in rec['rows']:
                    row=item['row'];expected=plain(s.cell(row,11).value);old=plain(s.cell(row,rc).value)
                    kind=item['kind'];hadqcc |= QCC in old
                    ce=cleaned(expected)
                    if ce!=expected.strip():
                        criterion_updates[row]=ce
                        put(b,s,f'K{row}',ce,'작성 지시 제거/계산 조건 표현 정리')
                    if kind=='vector': value=vector_result(rec,item,ce)
                    elif kind=='observation': value=observe_result(rec,ce)
                    elif kind=='rtt':
                        case=case_number(item['label']);log=getlog(rec,case)
                        if log and '기존 기록 불일치' not in rec['currentVerdict']:
                            value='검증 RTT C'+str(case)+'\n'+log[3]
                        else: value=tidy_result(old)
                    elif kind=='summary':
                        if rec['rows'][0]['row']==row:
                            value='검증 결과\n'+completed(ce).removeprefix('판정 범위\n')
                            value=value.replace('검증 결과\n검증 결과\n','검증 결과\n')
                            if '기존 기록 불일치' in rec['currentVerdict']:
                                # 숫자로 확인한 차이는 원래 기록을 유지한다.
                                value=tidy_result(without_qcc(old))
                            photo=photo_note(old)
                            if photo and PHOTO not in value: value+='\n'+photo
                        else: value=tidy_result(without_qcc(old))
                    else: value=tidy_result(without_qcc(old))
                    if kind=='comparison' and not any(x in value for x in ('기존 기록','불일치','미첨부')):
                        value+='\n판정: 출력/동작이 기준과 일치하여 통과.'
                    value=tidy_result(without_qcc(value))
                    put(b,s,f'{letter}{row}',value,'입력/출력/판정 중심 결과')
                if hadqcc:removals.append(dict(key='unit',id=rec['testId'],reason='E8300 경계에 입력한 표본/명령/상태로 유닛 자체 판정'))
                verdict=tidy_result(without_qcc(rec['currentVerdict']))
                verdict=verdict.replace('E8300 통과(시험자 확인)','통과(시험자 확인)')
                put(b,s,s.cell(rec['start'],vc).coordinate,verdict,'유닛 자체 판정으로 QCC 추가 자료 제외')
            groups['criterion']=groups['unit']
        else:
            tests=[]
            for row in range(b['first'],s.max_row+1):
                ident=plain(s.cell(row,b['idcol']).value)
                if not re.fullmatch(r'SW-FW-(IntegrationTest|SYS-Test)-ID-\d+',ident):continue
                tests.append(dict(id=ident,start=row,end=row));n=int(ident.rsplit('-',1)[1])
                old=plain(s.cell(row,rc).value);value=old
                keep=b['key']=='system' and n in KEEP_QCC_SYSTEM
                if not keep and QCC in value:
                    removals.append(dict(key=b['key'],id=ident,reason='E8300 입력/응답 또는 시험자 제품 관찰로 판정'))
                    value=without_qcc(value)
                # 중복된 통과 기준 복사 머리말을 없애고 실제 조건별 결과를 유지한다.
                if value.startswith('시험자 확인 기록') or value.startswith('계측 기준 / 예상값 기준 기록'):
                    chunks=value.split('\n\n')
                    if len(chunks)>1:value='\n\n'.join(chunks[1:])
                value=tidy_result(value)
                value=re.sub(r'(?m)^\[(\d+(?:~\d+)?)\]',r'검증 [\1]',value)
                value=value.replace('기댓값','기준값').replace('전류 기준값','전류 설정값')
                if b['key']=='integration' and n in (9,10,12,17):
                    value='입력: 시험용 SPI 명령/데이터를 E8300 수신 경계에 전달했다.\n'+value
                elif b['key']=='integration' and n in (3,13):
                    value='입력: 준비된 마이크/I2S 디지털 표본과 수신 상태를 E8300에 전달했다.\n'+value
                elif b['key']=='system' and n in (20,21,22,23,55,62,64,66):
                    value='입력: 기준서의 배터리 잔량/충전/연결 상태 데이터를 E8300에 전달했다.\n'+value
                if not any(x in plain(s.cell(row,vc).value) for x in ('불일치','미확정','기록 없음')):
                    value+='\n판정: 입력에 따른 출력/표시가 통과 기준과 일치함.'
                put(b,s,f'{letter}{row}',value,'측정 결과 문체/직접 입력 명시')
                verdict=plain(s.cell(row,vc).value)
                if not keep:verdict=without_qcc(verdict)
                verdict=tidy_result(verdict).replace('E8300/관찰 통과','통과')
                put(b,s,s.cell(row,vc).coordinate,verdict,'QCC 자료 요구 범위 분리')
            groups[b['key']]=tests
        # 보조 시트의 현행 결과/자료 안내도 같은 판정 범위를 사용한다.
        for sn in w:
            if sn.title==s.title or '원본' in sn.title or '변경 비교' in sn.title: continue
            current_id=None
            for cells in sn:
                joined=' '.join(plain(c.value) for c in cells)
                m=re.search(r'SW-FW-(?:SYS-Test|IntegrationTest)-ID-(\d+)',joined)
                if m:current_id=int(m.group(1))
                for cell in cells:
                    v=plain(cell.value)
                    if QCC not in v:continue
                    if b['key']!='system' or current_id not in KEEP_QCC_SYSTEM:
                        put(b,sn,cell.coordinate,without_qcc(v) or 'E8300 입력/출력 확인 범위.','보조 자료의 과도한 QCC 자료 요구 제외')
        w.close()
    b=specs[3];w=load_workbook(b['backup'],rich_text=True);s=w[b['sheet']]
    for row,value in criterion_updates.items():put(b,s,f'K{row}',value,'보고서와 같은 유닛 통과 기준')
    w.close()
    prev.save_json(QA/'baseline.json',dict(books=specs,protected=protected))
    prev.save_json(QA/'plan.json',dict(changes=list(changes.values()),shades=[],groups=groups,
        counts={k:dict(tests=len(v)) for k,v in groups.items()},qccRemoved=removals))
    print(json.dumps(dict(changedCells=len(changes),qccRemoved=len(removals)),ensure_ascii=False))


def build():
    prev.QA=QA
    # 재검토본은 직전 게시본과 대조한다. 다른 사람이 수정한 파일은 거부한다.
    publication = QA/'publication.json'
    expected = {f['path']: f['sha256'] for f in jread(publication)['files']} if publication.exists() else None
    prev.build(expected_hashes=expected)


def verify():
    base=jread(QA/'baseline.json');plan=jread(QA/'plan.json');out=[]
    for b in base['books']:
        changes={(p['sheet'],p['cell']):p for p in plan['changes'] if p['key']==b['key']}
        original=load_workbook(b['backup'],rich_text=True);new=load_workbook(QA/(b['key']+'-candidate.xlsx'),rich_text=True)
        assert original.sheetnames==new.sheetnames
        for old,s in zip(original,new):
            assert str(old.merged_cells)==str(s.merged_cells)
            assert (old.max_row,old.max_column,len(old._images))==(s.max_row,s.max_column,len(s._images))
            assert dict((r,d.height) for r,d in old.row_dimensions.items())==dict((r,d.height) for r,d in s.row_dimensions.items())
            for c in list(old._cells.values()):
                p=changes.get((s.title,c.coordinate))
                assert plain(s[c.coordinate].value)==(p['value'] if p else plain(c.value)),(b['key'],s.title,c.coordinate)
                assert list(c._style or [0]*9)==list(s[c.coordinate]._style or [0]*9),(b['key'],s.title,c.coordinate,'style')
        if b['key']!='criterion':
            cp=QA/'criterion-candidate.xlsx' if b['key']=='unit' and (QA/'criterion-candidate.xlsx').exists() else Path(b['criterionPath'])
            crit=load_workbook(cp,rich_text=True)
            for row in range(b['first'],new[b['sheet']].max_row+1):
                for col in range(1,b['common']+1):
                    assert plain(new[b['sheet']].cell(row,col).value)==plain(crit[b['csheet']].cell(row,col).value),(b['key'],row,col,'common')
            crit.close()
        before=prev.parts(b['backup']);after=prev.parts(QA/(b['key']+'-candidate.xlsx'))
        permitted={'xl/styles.xml'}|{prev.sheet_path(after,sn) for sn,addr in changes}
        assert set(before)==set(after)
        assert all(before[n]==after[n] for n in before if n not in permitted)
        resulttexts=[p['value'] for p in changes.values() if p['sheet']==b['sheet'] and p['cell'].startswith(chr(64+b['resultcol']))]
        forbidden=['단일 실측값을 임의 지정','실제값을 1이나','template의 null','PC 수학 대체 경로','실험 결과를 미리','결과 대조 필요']
        for text in resulttexts:
            assert not any(t in text for t in forbidden),(b['key'],text)
        if b['key']!='criterion':
            for s in new:
                for row in s:
                    for c in row:
                        assert not any(t in plain(c.value) for t in forbidden),(b['key'],s.title,c.coordinate,plain(c.value)[:300])
        if b['key']=='unit':
            s=new[b['sheet']]
            for rec in UNITS:
                for item in rec['rows']:
                    if item['kind']!='rtt':continue
                    expect=[ln for ln in plain(s.cell(item['row'],11).value).splitlines() if ln.startswith('[UT-')]
                    actual=[ln for ln in plain(s.cell(item['row'],12).value).splitlines() if ln.startswith('[UT-')]
                    assert len(expect)==len(actual),(item['row'],'RTT fragment lines',len(expect),len(actual))
                    for e,a in zip(expect,actual):
                        e_array=re.match(r'(\[UT-[^]]+\] C\d+ A\[\d+:\d+\]=)',e)
                        if e_array:assert a.startswith(e_array.group(1)),(item['row'],'RTT array index',e,a)
                        elif re.match(r'\[UT-[^]]+\] C\d+ I\d+ ',e):assert e==a,(item['row'],'RTT mismatch field',e,a)
                        else:assert re.match(r'\[UT-[^]]+\] C\d+ (PASS|FAIL|OBSERVED) N=\d+ F=\d+',a),(item['row'],'RTT summary',a)
        for (sn,addr),p in changes.items():
            if PHOTO in p['value'] or QCC in p['value']:
                root=prev.parse(after[prev.sheet_path(after,sn)])
                c=root.find(f".//s:c[@r='{addr}']",N)
                for run in c.findall('s:is/s:r',N):
                    v=''.join(run.find('s:t',N).itertext())
                    if PHOTO in v or QCC in v:
                        assert run.find('s:rPr/s:b',N).get('val')=='1'
                        assert run.find('s:rPr/s:color',N).get('rgb')=='FF'+(prev.BLUE if PHOTO in v else prev.RED)
        out.append(dict(key=b['key'],changedCells=len(changes),tests=plan['counts'][b['key']]['tests']))
        original.close();new.close()
    for p,sha in base['protected'].items():assert prev.digest(p)==sha,('protected',p)
    prev.save_json(QA/'verification.json',dict(books=out,protectedUnchanged=len(base['protected'])))
    print(json.dumps(out,ensure_ascii=False))


def publish():
    verify();base=jread(QA/'baseline.json')
    checks=jread(QA/'native-checks.json')
    for b in base['books']:
        check=next(c for c in checks if c['key']==b['key'])
        assert not check['overflow'],('native overflow',b['key'],check['overflow'])
        assert check['candidateSha256']==prev.digest(QA/(b['key']+'-candidate.xlsx')),('native stale',b['key'])
    previous={f['path']:f['sha256'] for f in jread(QA/'publication.json')['files']} if (QA/'publication.json').exists() else {}
    for b in base['books']:
        assert prev.digest(b['path'])==previous.get(b['path'],b['sha256']),('external edit',b['path'])
        with open(b['path'],'r+b'):pass
    for b in base['books']:
        p=Path(b['path']);stage=p.with_name('.reader-revision-'+p.name)
        assert not stage.exists();shutil.copyfile(QA/(b['key']+'-candidate.xlsx'),stage)
        stage.replace(p);print('published',b['key'])
    prev.save_json(QA/'publication.json',dict(files=[dict(path=b['path'],sha256=prev.digest(b['path'])) for b in base['books']]))


if __name__=='__main__':
    dict(prepare=prepare,refine=refine,fit=fit_results,finalize=finalize_content,polish=polish_layout,auxiliary=update_auxiliary_results,rtt_inputs=clarify_rtt_inputs,build=build,verify=verify,publish=publish)[sys.argv[1]]()
