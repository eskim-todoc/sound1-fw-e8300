# -*- coding: utf-8 -*-
"""최종 교차 검증.

만든 표와 문서에 적은 숫자가 서로 맞는지, 그리고 원본 측정값과도 맞는지
독립적으로 다시 확인한다. make_table.py 의 자체 검증과 별개로 돌린다.
"""
import io, os, re, sys, math, json

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = u'E:/workspace/projects/sound1-fw-e8300-backup'
XLSX = ROOT + u'/docs/참고/인허가/사운드1/[참고] Sound1 입력 음압 대비 ADC 출력 테이블.xlsx'
DOC = ROOT + u'/docs/참고/인허가/사운드1/Sound1 입력 음압 대비 ADC 출력 산출 근거.md'
SRC = os.path.join(HERE, u'측정원본.xlsx')

FULL_SCALE = 262143
ATT_PCB = -3.644
ATT_JIG = -8.293


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    from openpyxl import load_workbook

    ok_all = True

    def chk(name, ok, detail=u''):
        nonlocal_ok = ok
        print(u'  %-52s %s   %s' % (name, u'통과' if ok else u'실패', detail))
        return nonlocal_ok

    an = json.load(io.open(os.path.join(HERE, 'analysis.json'), encoding='utf-8'))
    ref = {int(k): v for k, v in an['ref'].items()}

    print(u'=== A. 산출물 존재 ===')
    for p in (XLSX, DOC):
        e = os.path.exists(p)
        ok_all &= chk(os.path.basename(p)[:50], e)
    if not ok_all:
        return

    wb = load_workbook(XLSX, data_only=True)
    print(u'\n=== B. 표 구조 ===')
    ok_all &= chk(u'시트 2개 (표 + 산출 근거)', len(wb.sheetnames) == 2, str(wb.sheetnames))
    ws = wb[wb.sheetnames[0]]

    # 머리 행은 B 열이 «음압» 으로 시작하고 바로 아래 행이 숫자인 자리다
    hdr = None
    for r in range(1, 20):
        v = ws.cell(r, 2).value
        if v and str(v).strip().startswith(u'음압') \
                and isinstance(ws.cell(r + 1, 2).value, (int, float)):
            hdr = r
            break
    ok_all &= chk(u'머리 행을 찾았는가', hdr is not None, u'%s 행' % hdr)

    rows = []
    r = hdr + 1
    while isinstance(ws.cell(r, 2).value, (int, float)):
        rows.append(tuple(ws.cell(r, c).value for c in range(2, 7)))
        r += 1
    ok_all &= chk(u'행 수 92 (29~120 dB SPL)', len(rows) == 92, u'%d 행' % len(rows))
    ok_all &= chk(u'첫 행 29 dB, 끝 행 120 dB',
                  rows[0][0] == 29 and rows[-1][0] == 120,
                  u'%s ~ %s' % (rows[0][0], rows[-1][0]))

    print(u'\n=== C. 1.5세대 두 열이 원본과 같은가 ===')
    d1 = [(a[0], a[1], ref[a[0]][0]) for a in rows if a[0] in ref and a[1] != ref[a[0]][0]]
    d2 = [(a[0], a[2], ref[a[0]][1]) for a in rows if a[0] in ref and a[2] != ref[a[0]][1]]
    ok_all &= chk(u'미적용 열', len(d1) == 0, u'차이 %d건 %s' % (len(d1), d1[:3]))
    ok_all &= chk(u'멤브레인 열', len(d2) == 0, u'차이 %d건 %s' % (len(d2), d2[:3]))

    print(u'\n=== D. Sound1 두 열이 식과 맞는가 ===')
    bad_p = bad_j = []
    bad_p = [(a[0], a[3], int(math.floor(FULL_SCALE * 10 ** ((a[0] - 120 + ATT_PCB) / 20.0))))
             for a in rows
             if a[3] != int(math.floor(FULL_SCALE * 10 ** ((a[0] - 120 + ATT_PCB) / 20.0)))]
    bad_j = [(a[0], a[4], int(math.floor(FULL_SCALE * 10 ** ((a[0] - 120 + ATT_JIG) / 20.0))))
             for a in rows
             if a[4] != int(math.floor(FULL_SCALE * 10 ** ((a[0] - 120 + ATT_JIG) / 20.0)))]
    ok_all &= chk(u'PCB 열 = floor(262143 x 10^((dB-120-3.644)/20))',
                  len(bad_p) == 0, u'차이 %d건 %s' % (len(bad_p), bad_p[:3]))
    ok_all &= chk(u'기구 열 = floor(262143 x 10^((dB-120-8.293)/20))',
                  len(bad_j) == 0, u'차이 %d건 %s' % (len(bad_j), bad_j[:3]))

    print(u'\n=== E. 판정 기준이 목적을 달성하는가 (허용 +-5 dB) ===')
    att = {k: v['flat_mean'] for k, v in an['attenuation'].items()}
    PCB = (u'501C', u'501D', u'501E', u'501F', u'5020')
    JIG = (u'500A', u'500E', u'500F', u'500D', u'500C-2차', u'500C-1차')
    pm = sum(att[n] for n in PCB) / len(PCB)
    ok_all &= chk(u'PCB 열이 실측 5 시료 평균 %.3f 인가' % ATT_PCB,
                  abs(pm - ATT_PCB) < 0.001, u'실제 %.4f dB' % pm)

    # 판정 기준은 «시험 대상(기구 조립) 실측» 의 중앙이다.
    hi = max(att[n] for n in JIG)
    lo = min(att[n] for n in JIG)
    center = (hi + lo) / 2
    ok_all &= chk(u'기구 열이 실측 상·하한의 중앙 %.3f 인가' % center,
                  abs(center - ATT_JIG) < 0.01,
                  u'상한 %.2f(최고) / 하한 %.2f(최저)' % (hi, lo))

    fail = [(n, round(att[n] - ATT_JIG, 2)) for n in JIG if abs(att[n] - ATT_JIG) > 5]
    ok_all &= chk(u'기구 조립 6 시료가 전건 통과하는가', len(fail) == 0,
                  u'탈락 %s' % (fail if fail else u'없음'))

    m_hi = 5 - (hi - ATT_JIG)
    m_lo = 5 + (lo - ATT_JIG)
    ok_all &= chk(u'상한 여유가 남는가 (실측 최고 대비)', m_hi > 0, u'%+.2f dB' % m_hi)
    ok_all &= chk(u'하한 여유가 남는가 (실측 최저 대비)', m_lo > 0, u'%+.2f dB' % m_lo)

    # 조립 감쇠가 0 인 물리적 극단(= PCB 최고)도 담겨야 안전망이 산다
    pcb_hi = max(att[n] for n in PCB)
    m_pcb = 5 - (pcb_hi - ATT_JIG)
    ok_all &= chk(u'조립 감쇠 0 인 극단(PCB 최고)도 담기는가', m_pcb > 0,
                  u'%+.2f dB' % m_pcb)
    ok_all &= chk(u'아래 여유가 위 여유보다 좁지 않은가 (위험은 아래로만 흩어진다)',
                  m_lo >= m_pcb, u'아래 %+.2f / PCB극단 %+.2f' % (m_lo, m_pcb))
    ok_all &= chk(u'상·하한 여유가 대칭인가 (차 0.1 dB 이내)', abs(m_hi - m_lo) < 0.1,
                  u'상 %+.2f / 하 %+.2f' % (m_hi, m_lo))
    ok_all &= chk(u'제외 3 시료도 이 기준에서는 통과하는가',
                  all(abs(att[n] - ATT_JIG) <= 5
                      for n in (u'500D', u'500C-2차', u'500C-1차')),
                  u'%s' % [round(att[n] - ATT_JIG, 2)
                           for n in (u'500D', u'500C-2차', u'500C-1차')])
    print(u'  통과 하한 %.2f dB - 실측 최저보다 %.2f dB 더 나쁜 조립까지 받아준다'
          % (ATT_JIG - 5, lo - (ATT_JIG - 5)))

    print(u'\n=== F. 측정 원본과의 역산 대조 ===')
    # 표의 PCB 값에서 감쇠를 되돌리면 1.5세대 값이 나와야 한다
    back = []
    for a in rows:
        if a[0] < 40 or a[0] not in ref:
            continue          # 작은 값은 내림 오차 비중이 크다
        est = a[3] * 10 ** (-ATT_PCB / 20.0)
        if abs(est - ref[a[0]][0]) > max(1.5, ref[a[0]][0] * 0.002):
            back.append((a[0], round(est, 1), ref[a[0]][0]))
    ok_all &= chk(u'PCB 열을 되돌리면 1.5세대 값인가 (40 dB 이상)',
                  len(back) <= 1, u'어긋남 %d건 %s' % (len(back), back[:3]))

    print(u'\n=== G. 문서에 적은 숫자가 표와 맞는가 ===')
    md = io.open(DOC, encoding='utf-8').read()
    tbl = {a[0]: a for a in rows}

    # 문서의 표본 표를 파싱해서 실제 표와 대조한다. 기댓값을 손으로 적지 않는다.
    mism = []
    n_parsed = 0
    for line in md.split('\n'):
        m = re.match(r'^\|\s*(\d+)\s*\|\s*\*{0,2}([\d,]+)\*{0,2}\s*\|\s*\*{0,2}([\d,]+)\*{0,2}'
                     r'\s*\|\s*\*{0,2}([\d,]+)\*{0,2}\s*\|\s*\*{0,2}([\d,]+)\*{0,2}\s*\|\s*$',
                     line.strip())
        if not m:
            continue
        d = int(m.group(1))
        if d not in tbl:
            continue
        vals = tuple(int(m.group(i).replace(',', '')) for i in (2, 3, 4, 5))
        n_parsed += 1
        if tbl[d][1:] != vals:
            mism.append((d, tbl[d][1:], vals))
    ok_all &= chk(u'문서의 표본 표(%d행)가 실제 표와 같은가' % n_parsed,
                  len(mism) == 0 and n_parsed >= 5, str(mism[:2]))
    for token in (u'-3.644', u'-8.293', u'-11.43', u'-5.15', u'1.86',
                  u'-13.29', u'0.27', u'6564', u'-7.48', u'-6.76'):
        ok_all &= chk(u'문서에 «%s» 가 있는가' % token, token in md)

    print(u'\n=== H. 표 시트 머리 설명문이 실제와 맞는가 ===')
    # 값만 보고 설명문을 읽지 않아 옛 기준의 문구가 남았던 적이 있다.
    note = ws.cell(4, 2).value or u''
    ok_all &= chk(u'설명문이 있는가', len(note) > 50, u'%d 자' % len(note))
    ok_all &= chk(u'내림을 floor 라 적었는가 (round 아님)',
                  u'floor(262143' in note and u'round(262143' not in note)
    ok_all &= chk(u'PCB %.3f 를 «5 시료 평균» 이라 적었는가' % ATT_PCB,
                  (u'%.3f' % ATT_PCB) in note and u'5 시료 평균' in note)
    ok_all &= chk(u'기구 %.3f 를 «6 시료의 상·하한 중앙» 이라 적었는가' % ATT_JIG,
                  (u'%.3f' % ATT_JIG) in note and u'6 시료' in note and u'중앙' in note,
                  u'' if u'6 시료' in note else u'옛 문구(3 시료 평균)가 남았다')
    ok_all &= chk(u'옛 기준값(-7.500 / -5.593)이 남아 있지 않은가',
                  u'-7.500' not in note and u'-5.593' not in note)

    print(u'\n=== I. 원본 88 dB 이상치가 그대로 실렸는가 ===')
    ok_all &= chk(u'표의 88 dB 미적용 값이 원본 6564 인가', tbl[88][1] == 6564,
                  u'실제 %s' % tbl[88][1])

    print(u'\n전체: %s' % (u'통과' if ok_all else u'실패'))


if __name__ == '__main__':
    main()
