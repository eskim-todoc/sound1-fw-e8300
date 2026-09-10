# -*- coding: utf-8 -*-
"""+-5 dB 범위 열·계열과 테이블 시트 설명문을 검증한다.

    python verify_range.py <대상.xlsx> [기준본.xlsx]
"""
import io, sys, re, math, zipfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SHEET = u'Sound1 시료 측정 데이터'
PCB, JIG, FS = -3.644, -8.293, 262143
adc = lambda d, a: int(math.floor(FS * 10 ** ((d - 120 + a) / 20.0)))
DBS = list(range(45, 111, 5))


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    tgt = sys.argv[1]
    base = sys.argv[2] if len(sys.argv) > 2 else None
    ok = [True]

    def chk(name, good, detail=u''):
        ok[0] &= bool(good)
        print(u'  %-48s %s   %s' % (name, u'통과' if good else u'실패', detail))

    wb = load_workbook(tgt, data_only=True)
    ws = wb[SHEET]

    print(u'=== A. 범위 열 머리 ===')
    chk(u'AF3 = «허용 범위 (PCB 상태)»', ws['AF3'].value == u'허용 범위 (PCB 상태)', repr(ws['AF3'].value))
    chk(u'AH3 = «허용 범위 (기구 조립)»', ws['AH3'].value == u'허용 범위 (기구 조립)', repr(ws['AH3'].value))
    mg = [str(m) for m in ws.merged_cells.ranges]
    chk(u'AF3:AG3 · AH3:AI3 병합', 'AF3:AG3' in mg and 'AH3:AI3' in mg,
        u'%s' % [m for m in mg if m.startswith(('AD', 'AF', 'AH'))])
    chk(u'부제가 +5 dB / -5 dB 인가',
        [ws[c].value for c in ('AF4', 'AG4', 'AH4', 'AI4')] == [u'+5 dB', u'-5 dB', u'+5 dB', u'-5 dB'],
        u'%s' % [ws[c].value for c in ('AF4', 'AG4', 'AH4', 'AI4')])

    print(u'\n=== B. 범위 값 ===')
    bad = []
    for r, d in zip(range(5, 19), DBS):
        want = (adc(d, PCB + 5), adc(d, PCB - 5), adc(d, JIG + 5), adc(d, JIG - 5))
        got = tuple(ws.cell(r, c).value for c in (32, 33, 34, 35))
        if got != want:
            bad.append((d, got, want))
    chk(u'AF~AI 14 행이 감쇠 +-5 dB 와 맞는가', len(bad) == 0, u'어긋남 %d %s' % (len(bad), bad[:2]))
    chk(u'상한 > 기준 > 하한 (전 행)',
        all(ws.cell(r, 32).value > ws.cell(r, 30).value > ws.cell(r, 33).value
            and ws.cell(r, 34).value > ws.cell(r, 31).value > ws.cell(r, 35).value
            for r in range(5, 19)))
    # +-5 dB 는 음압 격자가 5 dB 라 정확히 이웃 행이어야 한다
    shift = [(d, ws.cell(r, 34).value, ws.cell(r + 1, 31).value)
             for r, d in zip(range(5, 18), DBS[:-1])
             if ws.cell(r, 34).value != ws.cell(r + 1, 31).value]
    chk(u'상한[d] == 기준[d+5] (한 칸 위)', len(shift) == 0, u'어긋남 %d %s' % (len(shift), shift[:2]))
    chk(u'배율이 10^(+-5/20) 인가 (110 dB)',
        abs(ws['AH18'].value / ws['AE18'].value - 10 ** 0.25) < 0.001
        and abs(ws['AI18'].value / ws['AE18'].value - 10 ** -0.25) < 0.001,
        u'x%.4f / x%.4f' % (ws['AH18'].value / ws['AE18'].value,
                            ws['AI18'].value / ws['AE18'].value))
    chk(u'AJ 열 이후는 비어 있는가',
        all(ws.cell(r, c).value is None for r in range(1, 19) for c in range(36, 40)))

    print(u'\n=== C. 테이블 시트 머리 설명문 ===')
    note = wb[wb.sheetnames[0]]['B4'].value or u''
    chk(u'floor 로 적혀 있는가 (round 아님)',
        u'floor(262143' in note and u'round(262143' not in note)
    chk(u'PCB 를 «5 시료 평균» 이라 적었는가', u'5 시료 평균' in note)
    chk(u'기구 를 «6 시료의 상·하한 중앙» 이라 적었는가', u'6 시료' in note and u'중앙' in note)
    chk(u'옛 기준값(-7.500 / -5.593)이 없는가', u'-7.500' not in note and u'-5.593' not in note)
    chk(u'+-5 dB 안내가 있는가', u'+-5 dB' in note or u'±5 dB' in note)

    print(u'\n=== D. 차트 계열 ===')
    z = zipfile.ZipFile(tgt)
    per = {}
    for c in sorted(n for n in z.namelist() if re.match(r'xl/charts/chart\d+\.xml$', n)):
        x = z.read(c).decode('utf-8')
        names = []
        for m in re.finditer(r'<c:ser>(.*?)</c:ser>', x, re.S):
            b = m.group(1)
            nm = re.search(r'<c:tx>.*?<c:v>(.*?)</c:v>', b, re.S)
            f = re.search(r'<c:tx>.*?<c:f>(.*?)</c:f>', b, re.S)
            yv = re.search(r'<c:yVal>.*?<c:f>(.*?)</c:f>', b, re.S)
            names.append(((nm.group(1) if nm else (f.group(1) if f else u'-')),
                          (yv.group(1) if yv else u'').split('!')[-1]))
        per[c.split('/')[-1]] = names
    for cn, names in sorted(per.items()):
        got = {n for n, _ in names}
        chk(u'%s 에 상·하한 계열이 있는가' % cn,
            u'허용 상한 (+5 dB)' in got and u'허용 하한 (-5 dB)' in got,
            u'%d 계열' % len(names))
        for n, y in names:
            print(u'      %-28s %s' % (n, y))

    if base:
        print(u'\n=== E. 건드리지 않은 곳 ===')
        wb0 = load_workbook(base, data_only=True)
        ws0 = wb0[SHEET]
        diff = [(get_column_letter(c) + str(r), ws0.cell(r, c).value, ws.cell(r, c).value)
                for r in range(1, 19) for c in range(1, 32)
                if ws0.cell(r, c).value != ws.cell(r, c).value]
        chk(u'B~AE 열 (측정 원본 + 기준) 무변경', len(diff) == 0,
            u'차이 %d %s' % (len(diff), diff[:3]))
        s0, s1 = wb0[wb0.sheetnames[0]], wb[wb.sheetnames[0]]
        d2 = [(r, c) for r in range(1, s0.max_row + 1) for c in range(1, s0.max_column + 1)
              if s0.cell(r, c).value != s1.cell(r, c).value]
        chk(u'테이블 시트는 B4 설명문만 바뀌었는가', d2 == [(4, 2)], u'%s' % d2[:4])
        z0 = zipfile.ZipFile(base)
        chk(u'빠진 패키지 파트 없음', len(set(z0.namelist()) - set(z.namelist())) == 0)
        n0 = sum(len(re.findall(r'<c:ser>', z0.read(n).decode('utf-8')))
                 for n in z0.namelist() if re.match(r'xl/charts/chart\d+\.xml$', n))
        n1 = sum(len(v) for v in per.values())
        chk(u'전체 계열 수 %d -> %d (+4)' % (n0, n1), n1 == n0 + 4)

    print(u'\n전체: %s' % (u'통과' if ok[0] else u'실패'))
    return 0 if ok[0] else 1


if __name__ == '__main__':
    sys.exit(main())
