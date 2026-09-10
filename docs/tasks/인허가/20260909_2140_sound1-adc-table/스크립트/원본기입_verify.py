# -*- coding: utf-8 -*-
"""원본 측정 기록에 넣은 Sound1 산출 기준 열·차트 계열을 검증한다.

    python verify_src.py <대상.xlsx> [기준본.xlsx]

기준본은 수정 전 파일이다. 주면 «건드리지 않은 곳이 그대로인가» 까지 본다.
"""
import io, os, sys, re, math, zipfile
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SHEET = u'Sound1 시료 측정 데이터'
PCB, JIG, FS = -3.644, -8.293, 262143
adc = lambda d, a: int(math.floor(FS * 10 ** ((d - 120 + a) / 20.0)))
DBS = list(range(45, 111, 5))                     # 14 점, r5~r18 (헤더보다 한 행 아래)


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    tgt = sys.argv[1]
    base = sys.argv[2] if len(sys.argv) > 2 else None
    ok_all = [True]

    def chk(name, ok, detail=u''):
        ok_all[0] &= bool(ok)
        print(u'  %-46s %s   %s' % (name, u'통과' if ok else u'실패', detail))

    wb = load_workbook(tgt, data_only=True)
    ws = wb[SHEET]

    print(u'=== A. 머리와 형식 ===')
    chk(u'AD3 = «Sound1 산출 기준»', ws['AD3'].value == u'Sound1 산출 기준', repr(ws['AD3'].value))
    chk(u'AD3:AE3 병합 (기존 시료와 같은 형식)',
        'AD3:AE3' in [str(m) for m in ws.merged_cells.ranges])
    chk(u'AD4 = «PCB 상태», AE4 = «기구 조립»',
        ws['AD4'].value == u'PCB 상태' and ws['AE4'].value == u'기구 조립',
        u'%s / %s' % (ws['AD4'].value, ws['AE4'].value))

    print(u'\n=== B. 값이 산출식과 맞는가 ===')
    bad = []
    for r, d in zip(range(5, 19), DBS):
        want = (adc(d, PCB), adc(d, JIG))
        got = (ws.cell(r, 30).value, ws.cell(r, 31).value)
        if got != want:
            bad.append((d, got, want))
    chk(u'AD/AE 14 행 = floor(262143x10^((dB-120+a)/20))', len(bad) == 0,
        u'어긋남 %d %s' % (len(bad), bad[:2]))
    chk(u'AD5 가 45 dBSPL 자리인가 (한 행 아래 정렬)', ws['B4'].value == 45 and ws['AD5'].value == adc(45, PCB),
        u'B4=%s AD5=%s' % (ws['B4'].value, ws['AD5'].value))
    chk(u'AD18 이 110 dBSPL 자리인가', ws['B17'].value == 110 and ws['AD18'].value == adc(110, PCB),
        u'B17=%s AD18=%s' % (ws['B17'].value, ws['AD18'].value))
    chk(u'AD > AE (기구 조립이 더 감쇠)',
        all(ws.cell(r, 30).value > ws.cell(r, 31).value for r in range(5, 19)))
    chk(u'AF 열 이후는 비어 있는가',
        all(ws.cell(r, c).value is None for r in range(1, 19) for c in range(32, 36)))

    print(u'\n=== C. 차트 계열 ===')
    z = zipfile.ZipFile(tgt)
    want_ser = {u'AD': u'Sound1 산출 기준 (PCB 상태)',
                u'AE': u'Sound1 산출 기준 (기구 조립)'}
    found = {}
    for c in sorted(n for n in z.namelist() if re.match(r'xl/charts/chart\d+\.xml$', n)):
        x = z.read(c).decode('utf-8')
        for m in re.finditer(r'<c:ser>(.*?)</c:ser>', x, re.S):
            b = m.group(1)
            nm = re.search(r'<c:tx>.*?<c:v>(.*?)</c:v>', b, re.S)
            yv = re.search(r'<c:yVal>.*?<c:f>(.*?)</c:f>', b, re.S)
            xv = re.search(r'<c:xVal>.*?<c:f>(.*?)</c:f>', b, re.S)
            if not (nm and u'Sound1 산출 기준' in nm.group(1)):
                continue
            col = re.search(r'\$(A[DE])\$5:\$(A[DE])\$18', yv.group(1) if yv else u'')
            found[col.group(1) if col else u'?'] = (
                c.split('/')[-1], nm.group(1),
                xv.group(1) if xv else u'-', yv.group(1) if yv else u'-',
                '<a:prstDash val="dash"' in b)
    for col, title in sorted(want_ser.items()):
        f = found.get(col)
        chk(u'%s 열 계열이 있는가' % col, f is not None and f[1] == title,
            u'%s' % (u'%s / %s' % (f[0], f[1]) if f else u'없음'))
        if f:
            chk(u'  %s x축이 $B$4:$B$17 인가' % col, f[2].endswith(u'$B$4:$B$17'), f[2])
            chk(u'  %s 파선(기준선)으로 그렸는가' % col, f[4])
    chk(u'추가 계열이 정확히 2 개인가', len(found) == 2, u'%d 개' % len(found))

    if base:
        print(u'\n=== D. 건드리지 않은 곳이 그대로인가 ===')
        wb0 = load_workbook(base, data_only=True)
        ws0 = wb0[SHEET]
        diff = [(get_column_letter(c) + str(r), ws0.cell(r, c).value, ws.cell(r, c).value)
                for r in range(1, 19) for c in range(1, 30)
                if ws0.cell(r, c).value != ws.cell(r, c).value]
        chk(u'B~AC 열 (측정 원본) 무변경', len(diff) == 0, u'차이 %d %s' % (len(diff), diff[:3]))
        s0, s1 = wb0[wb0.sheetnames[0]], wb[wb.sheetnames[0]]
        d2 = sum(1 for r in range(1, s0.max_row + 1) for c in range(1, s0.max_column + 1)
                 if s0.cell(r, c).value != s1.cell(r, c).value)
        chk(u'1.5세대 테이블 시트 무변경', d2 == 0, u'차이 %d' % d2)
        z0 = zipfile.ZipFile(base)
        miss = set(z0.namelist()) - set(z.namelist())
        chk(u'빠진 패키지 파트 없음', len(miss) == 0, u'%s' % (miss or u'없음'))
        n0 = sum(len(re.findall(r'<c:ser>', z0.read(n).decode('utf-8')))
                 for n in z0.namelist() if re.match(r'xl/charts/chart\d+\.xml$', n))
        n1 = sum(len(re.findall(r'<c:ser>', z.read(n).decode('utf-8')))
                 for n in z.namelist() if re.match(r'xl/charts/chart\d+\.xml$', n))
        chk(u'전체 계열 수 %d -> %d (+2)' % (n0, n1), n1 == n0 + 2)

    print(u'\n전체: %s' % (u'통과' if ok_all[0] else u'실패'))
    return 0 if ok_all[0] else 1


if __name__ == '__main__':
    sys.exit(main())
