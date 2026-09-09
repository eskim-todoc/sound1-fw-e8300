# -*- coding: utf-8 -*-
"""Sound1 입력 음압 대비 ADC 출력 테이블 - 측정 데이터 분석.

원본: G:/공유 드라이브/20_프로젝트별/01_project sound1/99_공유/인허가 관련/
      Sound1 필수성능시험 준비 시료별 측정 기록.xlsx
      (같은 폴더의 측정원본.xlsx 로 복사해 두었다)

1.5세대와 Sound1 이 같은 스펙의 마이크를 쓰므로, 1.5세대 테이블을 기준으로
Sound1 의 감쇠량만 구하면 테이블을 만들 수 있다.
"""
import io, os, sys, math, json

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, u'측정원본.xlsx')

# ------------------------------------------------------------------ 시료 정의
# (최대값 열, 이름, 그룹)  평균 열은 최대 열 + 1 이다.
SAMPLES = [
    (5,  u'501C', u'PCB'),
    (7,  u'501D', u'PCB'),
    (9,  u'501E', u'PCB'),
    (11, u'501F', u'PCB'),
    (13, u'5020', u'PCB'),
    (16, u'500A', u'기구'),
    (18, u'500C-1차', u'기구'),
    (20, u'500D', u'기구'),
    (22, u'500E', u'기구'),
    (24, u'500F', u'기구'),
    (26, u'500A(중복)', u'중복'),      # 16 열과 값이 완전히 같다
    (28, u'500C-2차', u'기구'),
]

# 은수님 지시로 기준 산출에서 제외하는 시료 (기록은 남긴다)
EXCLUDED = [u'500C-1차', u'500C-2차', u'500D']

# 감쇠가 평탄해지는 구간. 저음압은 잡음 바닥 영향으로 흔들린다.
FLAT_DB = [80, 85, 90, 95, 100, 105, 110]

DB_LIST = [45, 50, 55, 60, 65, 70, 75, 80, 85, 90, 95, 100, 105, 110]
FIRST_DATA_ROW = 5          # 데이터는 헤더보다 한 행 아래에서 시작한다


def load():
    from openpyxl import load_workbook
    wb = load_workbook(SRC, data_only=True)
    ws = wb[u'Sound1 시료 측정 데이터']
    gs = wb[u'1.5세대 입력 음압 대비 ADC 출력 테이블']

    ref = {}
    for r in gs.iter_rows(min_row=2, values_only=True):
        try:
            ref[int(r[1])] = (int(r[2]), int(r[3]))
        except (TypeError, ValueError):
            pass

    data = {}
    for col, name, grp in SAMPLES:
        vals = {}
        for k, d in enumerate(DB_LIST):
            v = ws.cell(FIRST_DATA_ROW + k, col + 1).value      # 평균 열
            mx = ws.cell(FIRST_DATA_ROW + k, col).value          # 최대 열
            if isinstance(v, (int, float)):
                vals[d] = {'avg': v, 'max': mx if isinstance(mx, (int, float)) else None}
        data[name] = {'group': grp, 'vals': vals}
    return ref, data


def check_alignment():
    """데이터가 헤더보다 한 행 아래에서 시작한다는 전제를 구조로 검증한다.

    수치로는 두 가설이 잘 갈리지 않는다(감쇠 변화폭이 비슷하다). 결정적인 것은
    원본 시트의 구조다.
      1. 음압 값은 4~17 행(45~110)에 14 개, 데이터는 5~18 행에 14 개다
      2. 4 행에는 «최대/평균» 부제만 있고 측정값이 없다
      3. 18 행에는 측정값이 있는데 음압 값이 비어 있다
    2 와 3 은 «데이터가 한 행 아래»로만 설명된다. 그대로 읽으면 45 dB 측정이
    통째로 없고 18 행 데이터가 음압 없이 뜬다.
    """
    from openpyxl import load_workbook
    wb = load_workbook(SRC, data_only=True)
    ws = wb[u'Sound1 시료 측정 데이터']

    db_rows = [(r, ws.cell(r, 2).value) for r in range(4, 19)]
    n_db = len([1 for _, v in db_rows if isinstance(v, (int, float))])

    def has_data(r):
        return any(isinstance(ws.cell(r, c).value, (int, float))
                   for c, _, _ in SAMPLES)

    n_data = len([1 for r in range(4, 19) if has_data(r)])
    row4_data = has_data(4)
    row18_data = has_data(18)
    row18_db = isinstance(ws.cell(18, 2).value, (int, float))
    return {
        'n_db': n_db, 'n_data': n_data,
        'row4_has_data': row4_data,
        'row18_has_data': row18_data, 'row18_has_db': row18_db,
        'ok': (n_db == n_data and not row4_data and row18_data and not row18_db),
    }


def attenuation(ref, data):
    """시료마다 1.5세대(멤브레인 미적용) 대비 감쇠를 dB 로 구한다."""
    att = {}
    for name, rec in data.items():
        per_db = {}
        for d, v in rec['vals'].items():
            if d in ref and ref[d][0] > 0:
                per_db[d] = 20 * math.log10(v['avg'] / ref[d][0])
        flat = [per_db[d] for d in FLAT_DB if d in per_db]
        att[name] = {
            'group': rec['group'],
            'per_db': per_db,
            'flat_mean': (sum(flat) / len(flat)) if flat else None,
            'flat_n': len(flat),
            'flat_span': (max(flat) - min(flat)) if flat else None,
        }
    return att


def main():
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
    ref, data = load()

    print(u'=== 1. 행 정렬 검증 ===')
    al = check_alignment()
    print(u'  음압 값 %d개 / 측정값이 있는 행 %d개  -> %s'
          % (al['n_db'], al['n_data'], u'일치' if al['n_db'] == al['n_data'] else u'불일치'))
    print(u'  4 행에 측정값이 있는가  : %s  (없어야 한다. 부제 행이다)'
          % (u'있다' if al['row4_has_data'] else u'없다'))
    print(u'  18 행에 측정값이 있는가 : %s  (있어야 한다)'
          % (u'있다' if al['row18_has_data'] else u'없다'))
    print(u'  18 행에 음압 값이 있는가: %s  (없어야 한다)'
          % (u'있다' if al['row18_has_db'] else u'없다'))
    print(u'  -> 전제 «데이터가 한 행 아래»: %s' % (u'성립' if al['ok'] else u'불성립'))
    if not al['ok']:
        print(u'  [경고] 전제가 깨졌다. 아래 결과를 쓰지 말 것.')

    att = attenuation(ref, data)

    print(u'\n=== 2. 시료별 감쇠 (1.5세대 멤브레인 미적용 대비) ===')
    print(u'  %-12s %-6s %8s %6s %8s' % (u'시료', u'그룹', u'평탄평균', u'표본', u'구간편차'))
    for _, name, _ in SAMPLES:
        a = att[name]
        mark = u'  <- 제외' if name in EXCLUDED else u''
        print(u'  %-12s %-6s %7.2f dB %5d %7.2f dB%s'
              % (name, a['group'], a['flat_mean'], a['flat_n'], a['flat_span'], mark))

    print(u'\n=== 3. 500A 중복 확인 ===')
    same = all(abs(data[u'500A']['vals'][d]['avg'] - data[u'500A(중복)']['vals'][d]['avg']) < 1e-9
               for d in data[u'500A']['vals'])
    print(u'  16 열과 26 열의 값이 완전히 같은가: %s' % (u'예 (중복 기재)' if same else u'아니오'))

    print(u'\n=== 4. 그룹 기준값 ===')
    pcb = [att[n]['flat_mean'] for _, n, g in SAMPLES if g == u'PCB']
    jig = [att[n]['flat_mean'] for _, n, g in SAMPLES if g == u'기구' and n not in EXCLUDED]
    jig_all = [att[n]['flat_mean'] for _, n, g in SAMPLES if g == u'기구']
    pcb_m = sum(pcb) / len(pcb)
    jig_m = sum(jig) / len(jig)
    print(u'  PCB 상태      %d개  평균 %.3f dB  (범위 %.2f ~ %.2f)'
          % (len(pcb), pcb_m, min(pcb), max(pcb)))
    print(u'  기구 조립     %d개  평균 %.3f dB  (범위 %.2f ~ %.2f)  <- 제외 후'
          % (len(jig), jig_m, min(jig), max(jig)))
    print(u'  기구 조립(전체) %d개 평균 %.3f dB  <- 참고'
          % (len(jig_all), sum(jig_all) / len(jig_all)))
    print(u'  기구에 의한 추가 감쇠 = %.3f dB' % (jig_m - pcb_m))

    out = {
        'ref': {str(k): v for k, v in ref.items()},
        'attenuation': {k: {'group': v['group'], 'flat_mean': v['flat_mean'],
                            'flat_n': v['flat_n'], 'flat_span': v['flat_span'],
                            'per_db': {str(d): r for d, r in v['per_db'].items()}}
                        for k, v in att.items()},
        'pcb_mean': pcb_m, 'jig_mean': jig_m,
        'jig_all_mean': sum(jig_all) / len(jig_all),
        'excluded': EXCLUDED, 'flat_db': FLAT_DB, 'db_list': DB_LIST,
    }
    p = os.path.join(HERE, 'analysis.json')
    io.open(p, 'w', encoding='utf-8').write(json.dumps(out, ensure_ascii=False, indent=1))
    print(u'\n저장: %s' % p)


if __name__ == '__main__':
    main()
