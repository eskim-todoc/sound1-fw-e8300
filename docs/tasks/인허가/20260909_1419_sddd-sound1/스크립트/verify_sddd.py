# -*- coding: utf-8 -*-
"""SDDD 자체 검증"""
import io, sys, re, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.oxml.ns import qn

P = u'E:/workspace/projects/sound1-fw-e8300-backup/docs/참고/인허가/사운드1/[02] SDDD-FW-SND1 소프트웨어 상세 설계서.docx'
d = Document(P)
sec = d.sections[0]
BODY = sec.page_width.cm - sec.left_margin.cm - sec.right_margin.cm
print(u'용지 %.1f x %.1f cm / 본문 폭 %.2f cm / 표 %d / 단락 %d'
      % (sec.page_width.cm, sec.page_height.cm, BODY, len(d.tables), len(d.paragraphs)))

paras = [p.text.strip() for p in d.paragraphs]
alltext = '\n'.join(paras)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            alltext += '\n' + c.text

print(u'\n[검증_1] 표 폭 <= 본문 폭')
bad = 0
for i, t in enumerate(d.tables):
    g = t._tbl.find(qn('w:tblGrid'))
    tot = sum(int(gc.get(qn('w:w'))) / 567.0 for gc in g.findall(qn('w:gridCol')))
    if tot > BODY + 0.02:
        bad += 1
        print(u'  X 표%d %.2f cm' % (i, tot))
print(u'  결과: %s' % (u'통과' if bad == 0 else u'실패 %d건' % bad))

print(u'\n[검증_2] 장/절 번호 상호참조 없음 (착탈 안전성)')
pat = re.compile(u'\\d+\\.\\d*\\s*[장절]|[장절]\\s*\\d+\\.\\d')
hits = []
for t in paras:
    if not t or re.match(u'^\\d+(\\.\\d+)*\\.\\s', t):
        continue
    m = pat.search(t)
    if m:
        hits.append((m.group(0), t[:70]))
for h in hits:
    print(u'  X', h)
print(u'  결과: %s (%d건)' % (u'통과' if not hits else u'실패', len(hits)))

print(u'\n[검증_3] SU-ID 연속성 및 4요소 완비')
sus = [t for t in paras if re.match(u'^4\\.\\d+\\.\\d+\\.\\d+\\.\\s+SU-ID-CFX-', t)]
nums = [int(re.search(u'SU-ID-CFX-(\\d+)', t).group(1)) for t in sus]
print(u'  유닛 %d건 / 번호 연속: %s'
      % (len(nums), 'OK' if nums == list(range(1, len(nums) + 1)) else 'X %s' % nums[:6]))
missing = []
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if not re.match(u'^4\\.\\d+\\.\\d+\\.\\d+\\.\\s+SU-ID-CFX-', t):
        continue
    blk = '\n'.join(x.text for x in d.paragraphs[i + 1:i + 7])
    for k in (u'입력 파라미터', u'반환 파라미터', u'수정 파라미터', u'기능 및 동작'):
        if k not in blk:
            missing.append((t[:40], k))
print(u'  4요소 누락: %s' % (u'없음' if not missing else missing))

print(u'\n[검증_4] 파라미터가 실제 서명과 일치')
sys.path.insert(0, '.')
import cfx_units as CU
idx = json.load(io.open('sig_index.json', encoding='utf-8'))


def kcount(s):
    return 0 if s.strip() == u'없음' else len([x for x in s.split(',') if x.strip()])


mis = []
chk = 0
for _, secs in CU.CFX_SECTIONS:
    for b in secs:
        for u in b['units']:
            if not u['fn']:
                continue
            v = idx.get(u['fn'])
            if not v:
                mis.append((u['id'], u['fn'], u'서명 없음'))
                continue
            chk += 1
            if kcount(u['inp']) != v['argc']:
                mis.append((u['id'], u['fn'], u'인자 %d != %d' % (kcount(u['inp']), v['argc'])))
            if (v['ret'].strip() == 'void') != (u['ret'].strip() == u'없음'):
                mis.append((u['id'], u['fn'], u'반환 불일치'))
print(u'  서명 대조 %d건 / 불일치: %s' % (chk, u'없음' if not mis else mis))

print(u'\n[검증_5] 교차참조(U/M 라벨) 반영')
xr = sorted({x for _, ss in CU.CFX_SECTIONS for b in ss for u in b['units']
             if u['xref'] for x in u['xref'].split(',') if x.strip().startswith('U')},
            key=lambda s: int(s.strip()[1:]))
print(u'  문서에 반영된 U 라벨: %s (%d/14)' % (' '.join(x.strip() for x in xr), len(xr)))
lost = [u'U%d' % i for i in range(1, 15) if u'U%d' % i not in [x.strip() for x in xr]]
print(u'  누락: %s' % (u'없음' if not lost else lost))

print(u'\n[검증_6] 폐기 부품 언급')
for b in [u'nRF', u'NRF', u'가속도', u'이어피스', u'USB', u'7150', u'텔레코일']:
    n = [ln.strip()[:70] for ln in alltext.split('\n') if b in ln]
    if n:
        print(u'  ? %s : %s' % (b, n[:2]))
print(u'  (검출 없으면 통과)')

print(u'\n[검증_7] 그림')
imgs = [r for r in d.part.rels.values() if 'image' in r.reltype]
print(u'  이미지 %d개 (블록도 2 + 순서도 6 = 8 기대)' % len(imgs))
