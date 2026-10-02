from pathlib import Path
import json,re,hashlib
import openpyxl,pypdfium2 as pdf
T=Path(__file__).parent;Q=T/'feedback-sstc-20260929';j=json.loads((Q/'plan.json').read_text(encoding='utf8'))
w=openpyxl.load_workbook(Q/'final.xlsx',data_only=True);b=openpyxl.load_workbook(Q/'before.xlsx',data_only=True)
s=w[j['sheet']];d=w[j['diffname']];bs=b[j['sheet']]
assert hashlib.sha256(Path(j['srs']).read_bytes()).hexdigest()==j['srsHash']
for row in b.worksheets[0]:
 for c in row:assert c.value==w.worksheets[0][c.coordinate].value
for r in list(range(1,25))+list(range(105,bs.max_row+1)):
 nr=r if r<25 else r+1
 for c in range(1,11):assert bs.cell(r,c).value==s.cell(nr,c).value,('outside body',r,c)
trace=[];sysids=set();seen=set();testids=set();unlinked=[]
for e in j['body']:
 r=e['row'];n=e['id']
 for c,value in enumerate(e['values'],2):assert s.cell(r,c).value==value,('body',r,c,s.cell(r,c).value,value)
 assert s.cell(r,5).value==f'SW-FW-RID-{n}'
 assert s.cell(r,2).value==f'SW-FW-SYS-Test-ID-{n}'
 expected=j['requirements'][str(n)]['ref'];actual=s.cell(r,6).value
 assert expected==actual,(n,expected,actual)
 refs=re.findall(r'SYS-RID-FW-(\d+)',str(actual));sysids.update(map(int,refs))
 if not refs:unlinked.append(n)
 seen.add(n);testids.add(s.cell(r,2).value)
 trace.append({'SW':n,'TEST':s.cell(r,2).value,'SYS':refs,'row':r})
assert seen==set(range(1,82)) and len(testids)==81
assert sysids==set(range(1,51))
for r,e in enumerate(j['entries'],25):
 assert (d.cell(r,1).value or '')==e['mark']
 for c,value in enumerate(e['values'],2):assert str(d.cell(r,c).value)==str(value),('diff',r,c)
for kind in ['body','diff']:
 doc=pdf.PdfDocument(str(Q/(kind+'.pdf')))
 for i in range(len(doc)):doc[i].render(scale=1.3).to_pil().save(Q/(kind+str(i+1)+'.png'))
summary={'requirements':81,'tests':81,'missing_SW':[],'wrong_SYS_links':[],'SYS_coverage':50,'SRS_local_requirements_without_SYS':unlinked,'SRS_unchanged':True,'cover_and_non_test_sections_preserved':True,'trace':trace}
(Q/'verification.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf8')
print({k:v for k,v in summary.items() if k!='trace'})
