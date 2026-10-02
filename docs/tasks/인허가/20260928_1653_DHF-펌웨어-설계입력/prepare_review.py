"""Read-only workbook comparison, accounting for inserted test rows."""
from pathlib import Path
import json, shutil, hashlib, re
import openpyxl
T=Path(__file__).parent
O=T.parents[2]/'참고/인허가/사운드1/DHF/02_설계 입력/SW/펌웨어'
Q=T/'review-20260929'; Q.mkdir(exist_ok=True)
logs=json.loads((T/'excel-changes.json').read_text(encoding='utf-8'))
jobs=json.loads((T/'workbook-plan.json').read_text(encoding='utf-8'))
result=[]
for job in jobs:
 p=O/job['name']; backup=Q/job['name']; shutil.copy2(p,backup)
 old=openpyxl.load_workbook(job['source']); new=openpyxl.load_workbook(p)
 cache=openpyxl.load_workbook(p,data_only=True)
 reasons={r['location']:r['reason'] for r in logs if r['file']==Path(job['source']).name}
 records=[]; colors=[]
 for si,sh in enumerate(new):
  before=old[sh.title]
  for row in sh:
   for c in row:
    rn=c.row; col=c.column_letter if not isinstance(c,openpyxl.cell.cell.MergedCell) else None
    if col is None: continue
    added=si==1 and ((job['kind']=='SRS' and 80<=rn<=91) or (job['kind']=='SSTC' and 93<=rn<=104))
    source_row=rn-12 if si==1 and job['kind']=='SSTC' and rn>=105 else rn
    ov=None if added else before.cell(source_row,c.column).value
    nv=c.value
    if ov==nv or ov in (None,'') and nv in (None,''):continue
    status='[+] 추가' if ov in (None,'') else '[-] 삭제' if nv in (None,'') else '[~] 수정'
    color='E2F0D9' if status.startswith('[+]') else 'FFF2CC' if status.startswith('[~]') else None
    rid=''
    if si==1:
     rid=cache[sh.title].cell(rn,7 if job['kind']=='SRS' else 2).value or ''
    reason=reasons.get(f'{sh.title}!{c.coordinate}','기존 내용과 현재 Sound1 문서 대조')
    records.append({'sheet':sh.title,'cell':c.coordinate,'id':str(rid),'status':status,'old':ov,'new':nv,'source':f'{sh.title}!{col}{source_row}' if not added else '신규 행','reason':reason})
    if color:colors.append({'sheet':sh.title,'cell':c.coordinate,'color':color})
 # Drawing removal is not captured by cell comparison.
 if sum(len(s._images) for s in old)>sum(len(s._images) for s in new):
  records.append({'sheet':'표지','cell':'J12','id':'서명란','status':'[-] 삭제','old':'이전 제품 작성/검토 서명 그림','new':None,'source':'표지 서명란','reason':'날짜/서명/검토/승인 빈칸 요청 반영'})
 result.append({'kind':job['kind'],'name':job['name'],'source':job['source'],'target':str(p),'backup':str(backup),'hash':hashlib.sha256(p.read_bytes()).hexdigest(),'records':records,'colors':colors})
(Q/'plan.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
for j in result:print(j['kind'],len(j['records']),{s:sum(r['status']==s for r in j['records']) for s in ['[+] 추가','[~] 수정','[-] 삭제']})
