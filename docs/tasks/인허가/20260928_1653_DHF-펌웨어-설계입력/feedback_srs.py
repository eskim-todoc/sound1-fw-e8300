from pathlib import Path
import json,re,shutil,hashlib
import openpyxl
T=Path(__file__).parent;Q=T/'feedback-srs-20260929';Q.mkdir(exist_ok=True)
job=json.loads((T/'workbook-plan.json').read_text(encoding='utf8'))[0]
target=T.parents[2]/'참고/인허가/사운드1/DHF/02_설계 입력/SW/펌웨어'/job['name']
shutil.copy2(target,Q/'before.xlsx')
s=openpyxl.load_workbook(target,data_only=True).worksheets[1]
o=openpyxl.load_workbook(job['source'],data_only=True).worksheets[1]
orig={o.cell(r,7).value:[o.cell(r,c).value for c in range(2,10)] for r in range(6,80) if o.cell(r,7).value}
rowid={s.cell(r,7).value:r for r in range(6,92) if s.cell(r,7).value}
restores=[1,2,6,11,15,16,46]; edits={}
for n in restores:
 ident=f'SW-FW-RID-{n}';edits[ident]=orig[ident][4]
for n in [42,45,53]:edits[f'SW-FW-RID-{n}']=s.cell(rowid[f'SW-FW-RID-{n}'],6).value
edits['SW-FW-RID-42']=edits['SW-FW-RID-42'].replace('IQS323 터치 센서','터치 센서')
edits['SW-FW-RID-45']=edits['SW-FW-RID-45'].replace('Bluetooth LE로 데이터 통신이','블루투스를 사용하여 데이터 통신이')
split=edits['SW-FW-RID-53'].split('\n');assert len(split)==2
edits['SW-FW-RID-53']=split[0];newtext=split[1]
prefixes={**{n:'[E8300/QCC 연계]' for n in [44,45,69,70]},**{n:'[QCC 단독]' for n in [47,48,49,52,71,72,76]},15:'[E8300/QCC 각각 확인, 향후 I2S]'}
for n,p in prefixes.items():
 ident=f'SW-FW-RID-{n}';edits[ident]=p+' '+edits.get(ident,s.cell(rowid[ident],6).value)
body=[]
for r in range(6,92):
 v=[s.cell(r,c).value for c in range(2,10)];ident=v[5]
 if ident in edits:v[4]=edits[ident]
 if ident=='SW-FW-RID-53':v[6]='SYS-RID-FW-30'
 body.append({'row':len(body)+6,'sourceRow':r,'values':v})
 if ident=='SW-FW-RID-53':body.append({'row':len(body)+6,'sourceRow':r,'values':[None,None,None,None,newtext,'SW-FW-RID-81','SYS-RID-FW-31',None]})
entries=[]
norm=lambda v:re.sub(r'SYS-RID-FW:(?=\d)','SYS-RID-FW-',v) if isinstance(v,str) else v
for b in body:
 v=b['values'];ident=v[5];old=orig.get(ident)
 if ident is None:old=[o.cell(b['sourceRow']-12 if b['sourceRow']>=73 else b['sourceRow']-2 if b['sourceRow']>=53 else b['sourceRow'],c).value for c in range(2,10)]
 if old is None:entries.append({'mark':'[+]','values':v,'bodyrow':b['row']})
 elif [norm(x) for x in old]==[norm(x) for x in v]:entries.append({'mark':'','values':v,'bodyrow':b['row']})
 else:
  entries.extend([{'mark':'[-]','values':old,'bodyrow':b['row']},{'mark':'[+]','values':v,'bodyrow':b['row']}])
merges=[]
for m in s.merged_cells:
 if m.min_row<6:continue
 a=m.min_row+(m.min_row>=62);b=m.max_row+(m.max_row>=62)
 merges.append([a,b,m.min_col,m.max_col])
plan={'target':str(target),'hash':hashlib.sha256(target.read_bytes()).hexdigest(),'sheet':s.title,'diffname':s.title+'_변경 비교','edits':edits,'newtext':newtext,'body':body,'entries':entries,'merges':merges}
(Q/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8')
print('Modified IDs:',list(edits));print('New ID:81 -> SYS-RID-FW-31; ID53 -> SYS-RID-FW-30')
