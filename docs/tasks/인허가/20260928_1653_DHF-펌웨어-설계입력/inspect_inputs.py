from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
import openpyxl, json, hashlib

TASK=Path(__file__).parent
ROOT=Path('G:/공유 드라이브/20_프로젝트별/01_project sound1/04_docs/99_참고용DHF/02_설계 입력/SW/펌웨어')
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
records={}
for p in ROOT.rglob('*'):
    if not p.is_file(): continue
    rel=p.relative_to(ROOT).as_posix()
    item={'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'bytes':p.stat().st_size}
    lines=[]
    if p.suffix=='.docx':
        with ZipFile(p) as z:
            item['parts']={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
            for n in z.namelist():
                if n=='word/document.xml' or n.startswith(('word/header','word/footer')) and n.endswith('.xml'):
                    root=E.fromstring(z.read(n))
                    for i,para in enumerate(root.findall('.//w:p',NS)):
                        txt=''.join(para.xpath('.//w:t/text()',namespaces=NS))
                        if txt: lines.append(f'{n} p{i}: {txt}')
    elif p.suffix=='.xlsx':
        wb=openpyxl.load_workbook(p,data_only=False)
        for sh in wb:
            lines.append(f'SHEET {sh.title} {sh.max_row}x{sh.max_column} merges={list(sh.merged_cells.ranges)}')
            for row in sh:
                vals=[f'{c.coordinate}={c.value}' for c in row if c.value is not None]
                if vals: lines.append(' | '.join(vals))
    elif p.suffix=='.pptx':
        with ZipFile(p) as z:
            for n in z.namelist():
                if n.startswith('ppt/slides/slide') and n.endswith('.xml'):
                    root=E.fromstring(z.read(n))
                    lines.append(n+': '+' | '.join(root.xpath('.//a:t/text()',namespaces=NS)))
    item['text']=lines
    records[rel]=item
TASK.joinpath('baseline.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
TASK.joinpath('source-text.txt').write_text('\n\n'.join('FILE '+k+'\n'+'\n'.join(v['text']) for k,v in records.items() if v['text']),encoding='utf-8')
exports=Path('C:/Users/EUNSU/AppData/Local/Temp/browser-use/exports')
matrix=next(exports.glob('*f82bcace-c06f-48c8-aeb1-e66ca1ff8091.xlsx'))
wb=openpyxl.load_workbook(matrix,data_only=True)
sh=wb['설계 입출력 검증 매트릭스']
rows=[]
for i in list(range(1,10))+list(range(100,150)):
    rows.append({'row':i,'cells':{c.column_letter:c.value for c in sh[i] if c.value is not None}})
TASK.joinpath('matrix.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
print('inspected',len(records),'files; matrix rows',len(rows))
