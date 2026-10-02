from pathlib import Path
import json,re
import openpyxl,pypdfium2 as pdf
Q=Path(__file__).parent/'unified-20260929';jobs=json.loads((Q/'plan.json').read_text(encoding='utf8'))
for j in jobs:
 w=openpyxl.load_workbook(Q/(j['kind']+'-final.xlsx'),data_only=True);s=w[j['sheet']];d=w[j['diffname']]
 assert '변경 비교' not in w.sheetnames
 ids=[]
 for r in j['body']:
  for col,value in enumerate(r['values'],2):
   c=s.cell(r['row'],col)
   if isinstance(c,openpyxl.cell.cell.MergedCell):continue
   assert c.value==value,(j['kind'],'body',c.coordinate,c.value,value)
  if r['id']:ids.append(s.cell(r['row'],7 if j['kind']=='SRS' else 2).value)
 assert len(ids)==80 and len(set(ids))==80
 for i,e in enumerate(j['entries'],j['start']):
  assert (d.cell(i,1).value or '')==e['mark'],('marker',i)
  for col,value in enumerate(e['values'],2):
   c=d.cell(i,col)
   if isinstance(c,openpyxl.cell.cell.MergedCell):continue
   assert ('' if c.value is None else str(c.value))==('' if value is None else str(value)),(j['kind'],'diff',c.coordinate,c.value,value)
 assert len(w.sheetnames)==3
 for tag in ['body','diff']:
  doc=pdf.PdfDocument(str(Q/(j['kind']+'-'+tag+'.pdf')))
  doc[0].render(scale=1.5).to_pil().save(Q/(j['kind']+'-'+tag+'.png'))
 print(j['kind'],'80 unique IDs, all body and diff values verified')
