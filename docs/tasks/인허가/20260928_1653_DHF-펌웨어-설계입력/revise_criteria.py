from pathlib import Path
import json, re, shutil, hashlib
import openpyxl
T=Path(__file__).parent
Q=T/'criteria-feedback-20260929'; Q.mkdir(exist_ok=True)
O=T.parents[2]/'참고/인허가/사운드1/DHF/02_설계 입력/SW/펌웨어'
job=json.loads((T/'workbook-plan.json').read_text(encoding='utf8'))[1]
target=O/job['name']; shutil.copy2(target,Q/'before.xlsx')
w=openpyxl.load_workbook(target,data_only=True); s=w.worksheets[1]; d=w.worksheets[2]
old=openpyxl.load_workbook(job['source'],data_only=True).worksheets[1]
original={int(str(old.cell(r,2).value).split('-')[-1]):[old.cell(r,c).value for c in range(2,11)] for r in range(25,93)}
restore=[8,9,10,11,12,13,14,15,16,50,56,57,58,59,61,64,65]
battery='- 배터리 잔량 80% 이상: 녹색 점등\n- 10% 이상 80% 미만: 노란색 점등\n- 10% 미만: 노란색 점멸'
setup='LED 표시를 켜고 내부기, 매핑 프로그램의 연결을 해제한다. 오류 및 업데이트가 없는 일반 동작 상태에서 확인한다.'
updates={
6:{'H':None},
17:{'J':'- FPGA와 PCM 통신이 정상이고, 정상 동작 상태에 해당하는 LED가 표시되어야 한다.'},
21:{'G':setup+' 충전 케이스의 덮개는 열어 둔다.', 'I':'충전 전원을 연결하고 배터리 잔량에 따른 LED를 확인한다.', 'J':battery+'\n- 충전 여부에 관계없이 같은 잔량 표시 기준을 적용한다.'},
22:{'G':setup,'I':'배터리 잔량을 확인하면서 각 잔량 구간의 LED 색상과 점멸 여부를 확인한다.','J':battery},
56:{'J':'- 오디오 볼륨이 한번에 10으로 설정되지 않아야 한다.'},
62:{'G':setup+' 충전 케이스의 덮개는 열어 둔다.','I':'충전 전원을 연결하고 LED를 확인한다.','J':battery+'\n- 충전 여부에 관계없이 같은 잔량 표시 기준을 적용한다.'},
69:{'I':'블루투스로 오디오를 스트리밍하고 자극 출력을 확인한다.','J':'오디오 스트리밍에 따른 자극 출력이 확인되어야 한다.'},
70:{'I':'BLE 통신으로 명령어를 송신하고 응답을 확인한다.','J':'BLE 통신을 통한 명령어 송/수신이 확인되어야 한다.'}}
body=[]; edits=[]
for r in range(25,106):
 v=[s.cell(r,c).value for c in range(2,11)]; before=v[:]; n=int(v[0].split('-')[-1])
 if n in restore:
  for i in [2,7,8]:v[i]=original[n][i]
 for col,val in updates.get(n,{}).items():v[ord(col)-66]=val
 if n in (43,77):
  v=[x.replace('PCMI','PMIC').replace('사요자','사용자') if isinstance(x,str) else x for x in v]
 for i,(a,b) in enumerate(zip(before,v)):
  if a!=b:edits.append({'cell':f'{chr(66+i)}{r}','value':b,'id':n})
 body.append({'row':r,'id':n,'values':v})
entries=[]
for b in body:
 prev=original.get(b['id']); v=b['values']
 if prev==v:entries.append({'mark':'','bodyrow':b['row'],'values':v})
 else:
  if prev is not None:entries.append({'mark':'[-]','bodyrow':b['row'],'values':prev})
  entries.append({'mark':'[+]','bodyrow':b['row'],'values':v})
existing=[r for r in range(25,d.max_row+1) if str(d.cell(r,2).value).startswith('SW-FW-SYS-Test-ID-')]
assert len(existing)==len(entries),(len(existing),len(entries))
plan={'target':str(target.resolve()),'hash':hashlib.sha256(target.read_bytes()).hexdigest(),'sheet':s.title,'diffname':d.title,'body':body,'entries':entries,'edits':edits,'restore':restore,'source':job['source']}
(Q/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf8')
print('Changed cells',len(edits),'comparison rows',len(entries))
