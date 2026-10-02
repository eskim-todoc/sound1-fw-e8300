from pathlib import Path
import json,shutil,hashlib,re
import openpyxl
T=Path(__file__).parent;Q=T/'unified-20260929';Q.mkdir(exist_ok=True)
O=T.parents[2]/'참고/인허가/사운드1/DHF/02_설계 입력/SW/펌웨어'
jobs=json.loads((T/'workbook-plan.json').read_text(encoding='utf8'))
order=list(range(1,46))+[69,70]+list(range(46,55))+list(range(71,81))+list(range(55,69))
def expanded(s,r,c):
 for m in s.merged_cells:
  if m.min_row<=r<=m.max_row and m.min_col<=c<=m.max_col:return s.cell(m.min_row,m.min_col).value
 return s.cell(r,c).value
out=[]
for j in jobs:
 target=O/j['name'];backup=Q/j['name'];shutil.copy2(target,backup)
 wb=openpyxl.load_workbook(backup,data_only=True);s=wb.worksheets[1]
 old=openpyxl.load_workbook(j['source'],data_only=True).worksheets[1]
 if j['kind']=='SRS':
  rows=list(range(6,51))+[80,81]+list(range(51,61))+list(range(82,92))+list(range(61,80));start=6;end=91;cols=9
 else:rows=[n+24 for n in order];start=25;end=104;cols=10
 rowmap={r:start+i for i,r in enumerate(rows)}
 merges=[]
 if j['kind']=='SRS':
  for m in s.merged_cells:
   if m.min_row<6:continue
   if m.min_row>=80:continue
   a=rowmap[m.min_row];b=rowmap[m.max_row]
   if str(m)=='B49:B50':b=rowmap[81]
   if str(m) in ('B52:B60','C52:E60'):b=rowmap[91]
   merges.append(f'{openpyxl.utils.get_column_letter(m.min_col)}{a}:{openpyxl.utils.get_column_letter(m.max_col)}{b}')
  merges.append(f'C{rowmap[80]}:E{rowmap[81]}')
 entries=[];body=[]
 for r in rows:
  ident=s.cell(r,7 if j['kind']=='SRS' else 2).value
  num=int(re.search(r'(\d+)$',str(ident))[1]) if ident else None
  values=[s.cell(r,c).value for c in range(2,cols+1)]
  if j['kind']=='SRS' and r>=80:
   values[0]=None;values[1]='Bluetooth' if r==80 else None
  body.append({'oldrow':r,'row':rowmap[r],'id':num,'values':values})
  new=values[:]
  # Compare displayed values; row-relative ID formulas are not textual changes.
  oldrow=(num+5 if num<=45 else num+6 if num<=54 else num+6 if num<=56 else num+11) if num and num<=68 and j['kind']=='SRS' else num+24 if num and num<=68 else None
  if num is None:oldrow=r
  prev=[old.cell(oldrow,c).value for c in range(2,cols+1)] if oldrow else [None]*(cols-1)
  # Category merge geometry is represented once as current context, not repeated as false deletions.
  changed=[c for c,(a,b) in enumerate(zip(prev,new),2) if (a if a is not None else '')!=(b if b is not None else '')]
  if oldrow is None:
   entries.append({'mark':'[+]','values':new,'changed':list(range(2,cols+1)),'bodyrow':rowmap[r],'id':num,'oldrow':None})
  elif changed:
   entries.append({'mark':'[-]','values':prev,'changed':changed,'bodyrow':rowmap[r],'id':num,'oldrow':oldrow})
   entries.append({'mark':'[+]','values':new,'changed':changed,'bodyrow':rowmap[r],'id':num,'oldrow':oldrow})
  else:entries.append({'mark':'','values':new,'changed':[],'bodyrow':rowmap[r],'id':num,'oldrow':oldrow})
 # Existing sheets are preserved outside body; compare non-test preamble/footer fields as row pairs too.
 extras=[]
 for r in range(1,s.max_row+1):
  if start<=r<=end:continue
  rr=r-12 if j['kind']=='SSTC' and r>=105 else r
  new=[s.cell(r,c).value for c in range(2,cols+1)];prev=[old.cell(rr,c).value for c in range(2,cols+1)]
  changed=[c for c,(a,b) in enumerate(zip(prev,new),2) if (a or '')!=(b or '')]
  if changed:extras.append({'row':r,'oldrow':rr,'old':prev,'new':new,'changed':changed})
 out.append({**{k:j[k] for k in ('name','kind','source')},'target':str(target),'backup':str(backup),'hash':hashlib.sha256(target.read_bytes()).hexdigest(),'sheet':s.title,'diffname':s.title+'_변경 비교','start':start,'end':end,'cols':cols,'body':body,'merges':merges,'entries':entries,'extras':extras})
(Q/'plan.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf8')
for j in out:print(j['kind'],'body',len(j['body']),'diff',len(j['entries']),'extras',len(j['extras']))
