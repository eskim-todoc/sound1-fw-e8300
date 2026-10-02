from pathlib import Path
import json, re, hashlib, shutil
import openpyxl
import pypdfium2 as pdf
Q=Path(__file__).parent/'criteria-feedback-20260929'
j=json.loads((Q/'plan.json').read_text(encoding='utf8'))
before=openpyxl.load_workbook(Q/'before.xlsx',data_only=False)
after=openpyxl.load_workbook(Q/'final.xlsx',data_only=False)
s=after[j['sheet']];d=after[j['diffname']]
allowed={e['cell'] for e in j['edits']}
for name in before.sheetnames[:2]:
 for row in before[name]:
  for c in row:
   if name==j['sheet'] and c.coordinate in allowed:continue
   assert c.value==after[name][c.coordinate].value,(name,c.coordinate)
for b in j['body']:
 assert [s.cell(b['row'],c).value for c in range(2,11)]==b['values'],b['id']
for r,e in enumerate(j['entries'],25):
 assert [d.cell(r,c).value for c in range(2,11)]==e['values'],r
 assert (d.cell(r,1).value or '')==e['mark'],r
 if e['mark']:assert d.cell(r,2).fill.fgColor.rgb in ('00FCE4D6','00E2F0D9','FFFCE4D6','FFE2F0D9')
for row in before[j['diffname']]:
 for c in row:
  if 25<=c.row<=173:continue
  assert c.value==d[c.coordinate].value,c.coordinate
srsfile=Path(j['target']).with_name('SRS-FW-TD24A 소프트웨어 요구사항 사양서.xlsx')
sr=openpyxl.load_workbook(srsfile,data_only=True).worksheets[1]
refs={}
for r in range(6,sr.max_row+1):
 v=sr.cell(r,7).value
 if not str(v).startswith('SW-FW-RID-'):continue
 ref=sr.cell(r,8).value
 for m in sr.merged_cells:
  if m.min_row<=r<=m.max_row and m.min_col<=8<=m.max_col:ref=sr.cell(m.min_row,m.min_col).value
 refs[v]=ref
assert {b['id'] for b in j['body']}==set(range(1,82))
sysids=set()
for b in j['body']:
 v=b['values'];assert v[4]==refs[v[3]],b['id']
 sysids.update(map(int,re.findall(r'SYS-RID-FW-(\d+)',v[4] or '')))
 assert all(x in 'TDIA' for x in str(v[2]).split()),v[2]
assert sysids==set(range(1,51))
for name in ['led','bluetooth','diff']:
 doc=pdf.PdfDocument(str(Q/f'{name}.pdf'))
 for i,page in enumerate(doc):page.render(scale=1.5).to_pil().save(Q/f'{name}{i+1}.png')
report={'changedCells':len(allowed),'tests':81,'softwareRequirements':81,'systemRequirements':50,'untouchedCells':'preserved','comparisonRows':149,'restoredIds':j['restore']}
(Q/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(report)
