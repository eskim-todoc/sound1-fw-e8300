from pathlib import Path
import json,zipfile,hashlib
from lxml import etree as E
import openpyxl
Q=Path(__file__).parent/'short-methods-20260929';j=json.loads((Q/'plan.json').read_text(encoding='utf8'))
with zipfile.ZipFile(Q/'before.xlsx') as z:parts={n:z.read(n) for n in z.namelist()}
ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};S=ns['s'];trees={}
auth=openpyxl.load_workbook(Q/'authored.xlsx')
for edit in j['edits']:
 assert auth[edit['sheet']][edit['addr']].value==edit['value']
 name=f"xl/worksheets/sheet{edit['index']}.xml"
 if name not in trees:trees[name]=E.fromstring(parts[name])
 cell=trees[name].find(f".//s:c[@r='{edit['addr']}']",ns);assert cell is not None
 for child in list(cell):cell.remove(child)
 cell.set('t','inlineStr');E.SubElement(E.SubElement(cell,f'{{{S}}}is'),f'{{{S}}}t').text=edit['value']
for name,tree in trees.items():parts[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)
out=Q/'final.xlsx'
with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in parts.items():z.writestr(n,b)
before=openpyxl.load_workbook(Q/'before.xlsx');after=openpyxl.load_workbook(out)
expected={(e['sheet'],e['addr']):e['value'] for e in j['edits']}
for s in before:
 for row in s:
  for c in row:assert after[s.title][c.coordinate].value==expected.get((s.title,c.coordinate),c.value)
with zipfile.ZipFile(Q/'before.xlsx') as z:
 assert all(z.read(n)==data for n,data in parts.items() if n not in trees)
target=Path(j['target']);assert hashlib.sha256(target.read_bytes()).digest()==hashlib.sha256((Q/'before.xlsx').read_bytes()).digest()
target.write_bytes(out.read_bytes())
print('Verified: only method-cell text changed; all other package parts preserved.')
