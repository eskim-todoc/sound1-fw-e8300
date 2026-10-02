"""Source-preserving Sound1 design-input adaptation. No source writes."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import json,re,shutil,hashlib,sys
from lxml import etree as E
import openpyxl

TASK=Path(__file__).resolve().parent
ROOT=TASK.parents[3]
SRC=Path('G:/공유 드라이브/20_프로젝트별/01_project sound1/04_docs/99_참고용DHF/02_설계 입력/SW/펌웨어')
OUT=ROOT/'docs/참고/인허가/사운드1/DHF/02_설계 입력/SW/펌웨어'
STAGE=TASK/'build'; STAGE.mkdir(exist_ok=True)
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
NS={'w':W,'a':A,'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
base=json.loads((TASK/'baseline.json').read_text(encoding='utf-8'))
matrix={v['row']-99:v['cells'] for v in json.loads((TASK/'matrix.json').read_text(encoding='utf-8')) if v['row']>=100}
changes=[]
def ser(x): return E.tostring(x,encoding='UTF-8',xml_declaration=True,standalone=True)
def readzip(p):
    with ZipFile(p) as z:return {n:z.read(n) for n in z.namelist()}
def writezip(p,parts):
    p.parent.mkdir(parents=True,exist_ok=True)
    with ZipFile(p,'w',ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)
def txt(p):return ''.join(p.xpath('.//w:t/text()',namespaces=NS))
def replace(p,old,new):
    for start in reversed([m.start() for m in re.finditer(re.escape(old),txt(p))]):
        end=start+len(old); pos=0; inserted=False
        for n in p.xpath('.//w:t',namespaces=NS):
            t=n.text or '';stop=pos+len(t)
            if pos<end and stop>start:
                n.text=t[:max(0,start-pos)]+(new if not inserted else '')+(t[max(0,end-pos):] if stop>end else '')
                n.set('{http://www.w3.org/XML/1998/namespace}space','preserve');inserted=True
            pos=stop
def log(file,loc,old,new,why):
    if old!=new:changes.append(dict(file=file,location=loc,before=old,after=new,reason=why))

WORD={
'SGR':{290:'E8300 펌웨어는 별도 운영체제 없이 동작한다.',292:'onsemi E8300 SDK와 제조사 제공 라이브러리를 사용한다.'},
'SR':{
181:'Ezairo 8300은 CFX, ARM Cortex-M3, HEAR, Filter Engine, LPDSP32, NNA를 포함한다. Sound1의 신호처리와 제어는 CFX, HEAR, Cortex-M3를 중심으로 구성한다.',
182:'- CFX : 오디오 입력, 빔포밍 및 자극 데이터 생성',183:'- HEAR : FFT 등 신호처리 연산',184:'- Filter Engine, LPDSP32, NNA : 칩 내 제공 기능으로, 사용 여부는 펌웨어 설계에 따른다.',
185:'- ARM Cortex-M3 : 시스템 제어 및 QCC5181과 통신. QCC5181은 별도의 Bluetooth 칩이다.',
187:'< 그림 3. Ezairo 8300 Chip 블록도 >',
190:'프로그램과 맵 데이터는 외부 비휘발성 메모리에 저장하며, Ezairo 8300 내부 RAM에서 프로그램을 실행한다.',
191:'RAM은 프로그램용, 데이터용 및 코어 간 공유 영역으로 구분한다. 메모리 배치는 각 프로젝트의 링커 설정을 따른다.',
192:'부팅 시 부트로더가 저장된 프로그램과 데이터를 RAM으로 적재한다. 아래 값은 현재 CM3 프로젝트 링커에서 선언한 영역 크기이며, 실제 사용량은 빌드 결과로 확인한다.',
193:'- 비휘발성 메모리',194:'프로그램, 맵 데이터 및 설정 데이터를 저장할 수 있어야 한다.',
195:'- 프로그램 메모리',196:'CFX : 128 kBytes, ARM Cortex-M3 : 384 kBytes',
197:'- 데이터 메모리',198:'ARM Cortex-M3 DRAM : 96 kBytes',199:'공유 메모리 : 신호처리 버퍼와 코어 간 공유 데이터 영역을 중복 없이 배치한다.',
201:'',204:'',207:'',
210:'디지털 마이크 2개로 소리를 입력 받고, IQS323 터치 센서로 사용자 입력을 받는다.',
214:'시동과 종료 중에는 자극 출력을 생성하지 않으며, 맵에서 정한 자극 크기와 펄스폭 범위를 지켜야 한다.',
218:'Bluetooth Classic / LE :',219:'QCC5181을 통해 오디오를 수신하고, 리모콘 앱과 매핑 프로그램의 제어 명령 및 응답 데이터를 송수신한다.',
226:'- CFX : C language, Assembly language / HEAR : Microcode',227:'- ARM Cortex-M3 : C language',
229:'- 각 코어의 실행 코드와 데이터가 링커에서 할당한 메모리 범위를 넘지 않아야 한다.',
241:'검토 결과 :',245:'검토 결과 :',249:'검토 결과 :',253:'검토 결과 :',257:'검토 결과 :',261:'검토 결과 :',
267:'소프트웨어 요구사항 확정 후 위험분석을 재평가하여 새로운 위해요인 및 위험 산정의 변경 여부를 기록한다.'}}

def word_build():
    diagram=ROOT/'docs/tasks/인허가/20260909_1129_sad-sound1/그림/그림1_시스템_블록도.png'
    for p in SRC.glob('*.docx'):
        code=p.name.split('-')[0];parts=readzip(p)
        for name in list(parts):
            if name!='word/document.xml' and not re.match(r'word/(header|footer)\d+\.xml$',name):continue
            x=E.fromstring(parts[name]);paras=x.xpath('//w:p',namespaces=NS)
            for i,para in enumerate(paras):
                old=txt(para);new=old.replace('TD20B','TD24A').replace('Ezairo7150 chip','Ezairo 8300 chip')
                if name=='word/document.xml' and i in WORD[code]:new=WORD[code][i]
                if name=='word/document.xml' and i<135:
                    if old=='임정우':new='김은수'
                    elif old=='사업부장':new='수석연구원'
                    elif old in ('이호승','기술이사') or re.fullmatch(r'20\d\d\.\d\d\.\d\d',old.strip()):new=''
                    # Empty future revision numbers; only revision 0 remains.
                    elif i>66 and old.isdigit() and int(old)>0:new=''
                    if i<40:
                        for n in para.xpath('.//w:drawing|.//w:pict',namespaces=NS):n.getparent().remove(n)
                if 'header' in name and re.fullmatch(r'20\d\d\.\d\d\.\d\d',old.strip()):new=''
                if old!=new:
                    replace(para,old,new);log(p.name,f'{name} p{i}',old,new,'Sound1 식별/현행 설계 반영 또는 과거 승인 기록 제거')
            if code=='SR' and name=='word/document.xml':
                rels=E.fromstring(parts['word/_rels/document.xml.rels'])
                for i in (178,186,200,203,206):
                    para=paras[i]
                    for draw in para.xpath('.//w:drawing|.//w:pict',namespaces=NS):
                        if i in (178,186):
                            ids=draw.xpath('.//@r:embed|.//@r:id',namespaces=NS)
                            for rid in ids:
                                rel=next((r for r in rels if r.get('Id')==rid),None)
                                if rel is not None:
                                    target='word/'+rel.get('Target'); data=diagram.read_bytes() if i==178 else (STAGE/'e8300-chip.png').read_bytes()
                                    # Use a PNG relationship regardless of original image type.
                                    newtarget=f'media/sound1-{i}.png'; rel.set('Target',newtarget);parts['word/'+newtarget]=data
                                    from PIL import Image
                                    import io
                                    im=Image.open(io.BytesIO(data));ratio=im.height/im.width
                                    for ex in draw.xpath('.//*[local-name()="extent" or local-name()="ext"]'):
                                        if ex.get('cx') and ex.get('cy'):ex.set('cy',str(round(int(ex.get('cx'))*ratio)))
                                    log(p.name,f'그림 {2 if i==178 else 3}',target,newtarget,'구형 하드웨어 그림을 현행 구조로 교체')
                        else:
                            draw.getparent().remove(draw);log(p.name,f'p{i} 구형 RAM 그림','Ezairo7150 RAM 그림','','E8300 링커 영역 설명으로 대체')
                parts['word/_rels/document.xml.rels']=ser(rels)
                for i in range(200,209):
                    para=paras[i]
                    if not txt(para).strip():para.getparent().remove(para)
                # Keep a bilingual review question together with its answer field.
                for i in (259,260):
                    pp=paras[i].find(f'{{{W}}}pPr')
                    if pp is None:pp=E.Element(f'{{{W}}}pPr');paras[i].insert(0,pp)
                    if pp.find(f'{{{W}}}keepNext') is None:E.SubElement(pp,f'{{{W}}}keepNext')
            parts[name]=ser(x)
        ct=E.fromstring(parts['[Content_Types].xml'])
        if not any(n.get('Extension')=='png' for n in ct):E.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Default',Extension='png',ContentType='image/png')
        parts['[Content_Types].xml']=ser(ct)
        writezip(STAGE/p.name.replace('TD20B','TD24A'),parts)

# Original row numbers are deliberately retained.
ROW_FW={6:[2],7:[2],8:[1],9:[4],10:[5],11:[3],13:[6],14:[7],15:[8],16:[9],17:[10],18:[11],19:[12],
49:[13],50:[14],52:[17],53:[22],54:[18],55:[26],56:[23],57:[27],58:[19],59:[30,31],60:[32],61:[48],62:[47],
68:[39,40],69:[41],70:[28],71:[42],72:[43],73:[44],75:[45],76:[46],77:[37],78:[49],79:[50]}
EXTRA=[15,16,20,21,24,25,29,33,34,35,36,38]
CUSTOM={6:'전채널 최대 자극률은 24 kHz이어야 한다. 32채널 사용 시 채널당 750 Hz로 자극한다.',7:'24채널 사용 시 채널당 1000 Hz로 자극한다.',
11:'자극 펄스폭은 13~255 us 범위에서 1 us 단위로 조정할 수 있어야 한다.',16:'CIS 및 NofM 자극 기법을 사용할 수 있어야 한다. 선택된 맵의 채널 수와 자극 순서를 따른다.',17:'사운드 처리 주기는 1 ms 이하이어야 한다.',
20:'디지털 마이크 2개와 Bluetooth 오디오를 입력 받을 수 있어야 한다.',21:'16 kHz 표본화로 8 kHz 미만의 오디오 신호를 입력 받을 수 있어야 한다.',
24:'터치 센서의 입력에 따라 전원 모드와 사용자 제어를 처리한다.',47:'I2C 프로토콜을 사용하여 IQS323 터치 센서와 데이터를 송수신할 수 있어야 한다.',61:'배터리 저전압을 주기적으로 반복하여 자극 출력과 LED로 알린다.',
50:'매핑 프로그램이 실행되는 기기와 Bluetooth LE로 데이터 통신이 가능해야 한다.',62:'자극 볼륨은 단계적으로 증가하도록 설계한다.(soft increase)'}

def clean(s):return str(s).strip().replace('·','/').replace('ㆍ','/').replace('\r','')
def matrix_field(ids,key):return '\n'.join(clean(matrix[n][key]) for n in ids if key in matrix[n])
def workbook_plan():
    jobs=[];srs=next(SRC.glob('SRS*.xlsx'));sst=next(SRC.glob('SSTC*.xlsx'))
    srwb=openpyxl.load_workbook(srs);req=srwb.worksheets[1]
    id_by_row={};n=0
    for r in range(6,80):
        if req[f'F{r}'].value:n+=1;id_by_row[r]=f'SW-FW-RID-{n}'
    for r,fw in enumerate(EXTRA,80):n+=1;id_by_row[r]=f'SW-FW-RID-{n}';ROW_FW[r]=[fw]
    for source in (srs,sst):
        wb=openpyxl.load_workbook(source);edits={};kind=source.name.split('-')[0]
        def put(sheet,addr,value,why):
            old=wb[sheet][addr].value if sheet in wb.sheetnames else None
            if old!=value:
                edits.setdefault(sheet,{})[addr]=value;log(source.name,f'{sheet}!{addr}',old,value,why)
        for sh in wb:
            for row in sh:
                for c in row:
                    if isinstance(c.value,str) and 'TD20B' in c.value:put(sh.title,c.coordinate,c.value.replace('TD20B','TD24A'),'Sound1 문서 식별')
        cover=wb.worksheets[0].title
        for cell,value in {'L3':'Rev. 0','D12':'토닥 / 수석연구원','H12':'김은수','M12':None,'D14':None,'H14':None,'J14':None,'M14':None,'D16':None,'H16':None,'J16':None,'M16':None,'J12':None,'D24':None,'B24':0}.items():put(cover,cell,value,'작성자/개정 0, 날짜/서명/검토/승인 빈칸')
        sh=wb.worksheets[1].title
        if kind=='SRS':
            put(sh,'H74',None,'배터리 잔량은 승계 요구사항으로 유지하되 현행 FW-33 보안 항목과 혼동되는 구형 번호 제거')
            for r,fw in ROW_FW.items():
                if r!=7:put(sh,f'H{r}',', '.join(f'SYS-RID-FW-{x}' for x in fw),'현행 매트릭스 식별자 연결')
                if r>=52 and r not in (61,62,68,69,70,71,72,73,75,76,78,79) or r>=80:
                    put(sh,f'F{r}',matrix_field(fw,'F'),'현행 요구사항 반영')
            for r,v in CUSTOM.items():put(sh,f'F{r}',v,'Sound1 하드웨어/필수성능 반영')
            for r in range(6,80+len(EXTRA)):
                put(sh,f'G{r}',f'=IF(COUNTA(F{r}),"SW-FW-RID-"&COUNTA(F$6:F{r}),"")','기존 식별 순서 유지, CONCAT/#REF! 오류 제거')
            for r,fw in enumerate(EXTRA,80):
                put(sh,f'B{r}','c) 시스템 연계' if fw in (15,16) else 'e) 보안 요구사항','누락 요구사항을 기존 필드로 추가')
                put(sh,f'C{r}','Bluetooth' if fw in (15,16) else '사이버 보안','기존 분류 적용')
        else:
            put(sh,'D7','TD24A-A-EXT','대상 제품 식별');put(sh,'D8',None,'시험 버전은 실제 시험 시 기록')
            # The original has 68 tests in rows 25..92. Each corresponds to a nonblank SRS row.
            original_rows=[r for r in range(6,80) if req[f'F{r}'].value]
            assert len(original_rows)==68,len(original_rows)
            tests=[]
            for tr,r in enumerate(original_rows+list(range(80,80+len(EXTRA))),25):
                fw=ROW_FW.get(r,[]);tests.append((tr,r,fw))
                put(sh,f'E{tr}',id_by_row[r],'깨진 외부 파일 참조를 확정 식별자로 연결')
                if fw:
                    if tr>92:put(sh,f'C{tr}',matrix_field(fw,'F'),'누락 요구사항 시험 추가')
                    put(sh,f'D{tr}',matrix_field(fw,'T') or 'Test','매트릭스의 검증 방법')
                    put(sh,f'I{tr}',matrix_field(fw,'U'),'매트릭스의 시험 입력/절차')
                    put(sh,f'J{tr}',matrix_field(fw,'V') or '요구사항에 정한 동작이 확인되어야 한다.','매트릭스의 예상 결과/합격 기준')
                if r in (24,47):
                    put(sh,f'C{tr}','터치 입력 및 제어' if r==24 else 'IQS323 I2C 통신','가속도 센서 대신 터치 센서')
                    put(sh,f'I{tr}','터치 센서를 조작하고 입력 상태 및 제어 동작을 확인한다.','Sound1 사용자 입력')
                    put(sh,f'J{tr}','터치 상태를 정상적으로 읽고, 설정된 터치 동작에 따라 제어되어야 한다.','터치 입력 검증')
                if r in (20,21):
                    put(sh,f'C{tr}','오디오 소스 입력' if r==20 else '오디오 입력 대역','Sound1 입력 구성')
                    put(sh,f'I{tr}','디지털 마이크와 Bluetooth 오디오를 각각 입력한다.' if r==20 else '필수성능시험 기준서의 저주파/중간/고주파 순음을 입력한다.','실제 입력 경로')
                    put(sh,f'J{tr}','선택된 입력 소스에 따라 자극 출력이 발생해야 한다.' if r==20 else '입력 주파수에 해당하는 채널의 자극 출력이 확인되어야 한다.','스피커 출력 대신 내부기 자극 확인')
                if r==16:
                    put(sh,f'I{tr}','CIS 맵과 NofM 맵을 각각 적용한다. NofM은 선택 채널 수보다 많은 대역의 합성음을 입력하고 32채널 테스트 내부기에서 각 전극 출력을 확인한다.','NofM 시험 승계')
                    put(sh,f'J{tr}','CIS는 맵의 자극 순서를 따른다. NofM은 설정된 N개 채널이 선택되고, 선택되지 않은 채널에는 자극이 없어야 하며 맵의 자극 순서를 따라야 한다.','선택 채널과 출력 순서 확인')
                if r==11:
                    put(sh,f'I{tr}','펄스폭을 13, 14, 255 us로 설정하고 출력 펄스폭을 측정한다.','설계 입력값 기준')
                    put(sh,f'J{tr}','각 설정 펄스폭의 ±5% 이내. 매트릭스 합격 기준의 하한 12 us와 설계 입력의 13 us 불일치는 확정 전 검토한다.','원본 매트릭스 불일치 명시')
                if fw==[15]:
                    put(sh,f'D{tr}','Test','기존 시험 방법 적용')
                    put(sh,f'I{tr}','Bluetooth Classic 오디오 송신 기기를 연결하고 1048 Hz 순음을 재생한다. 오디오 입력 상태와 해당 주파수 대역의 자극 출력을 확인한다.','입력/출력 절차 보충')
                    put(sh,f'J{tr}','오디오 스트리밍 연결과 데이터 수신이 유지되고, 입력 주파수에 대응하는 자극 출력이 확인되어야 한다.','오디오 통신 요구사항의 동작 확인')
                if fw==[16]:
                    put(sh,f'D{tr}','Test','기존 시험 방법 적용')
                    put(sh,f'I{tr}','Bluetooth LE 연결과 인증 후 앱에서 맵 번호를 읽고 1~4번으로 변경한다. 응답과 적용된 맵 번호를 확인한다.','입력/출력 절차 보충')
                    put(sh,f'J{tr}','읽기/변경 명령에 정상 응답하며 응답한 맵 번호와 실제 적용된 번호가 일치해야 한다.','데이터 통신 요구사항의 동작 확인')
                if r==59:put(sh,f'C{tr}','소프트웨어 유지보수 및 배포 관리','현행 관리 요구사항')
                if r==60:put(sh,f'C{tr}','부팅 시 펌웨어 무결성 검사','구형 CM3 크기 제한 대체')
                if r==53:put(sh,f'C{tr}','무조작 Bluetooth 세션 종료','현행 요구사항')
                if r==77:put(sh,f'C{tr}','Bluetooth 연결 해제 시 기본 동작','이어피스 항목 대체')
                if fw==[21]:put(sh,f'H{tr}','인증 실패 5회와 페어링 실패 5회가 매트릭스에서 혼용되어 있으므로 인증 단계와 잠금 시간을 확정한 후 시험한다.','시험 조건 불일치 기록')
                if fw==[36]:put(sh,f'I{tr}','권한이 있는 연결과 권한이 없는 연결에서 감사 기록 읽기를 시도한다. 전송 보호 적용 여부와 원격 수정 거부를 확인한다.','열람 권한/기밀성 요구사항에 맞춰 기존 절차 보충')
                if r==61:
                    put(sh,f'I{tr}','자극/LED 알림을 켜고 저전압 상태를 만든 뒤 알림 주기 동안 반복 관찰한다. LED 패턴과 내부기 전극의 알림 자극 출력을 각각 확인한다.','저전압의 두 알림 출력 확인')
                    put(sh,f'J{tr}','저전압 LED와 알림 자극이 설정된 주기에 따라 반복되어야 한다.','FW-48의 LED 및 자극 출력 모두 검증')
                if fw:put(sh,f'F{tr}',', '.join(f'SYS-RID-FW-{x}' for x in fw),'시험별 발생원 추적')
                if tr>92:put(sh,f'B{tr}',f'SW-FW-SYS-Test-ID-{tr-24}','기존 시험 번호 다음 번호')
        jobs.append(dict(source=str(source),name=source.name.replace('TD20B','TD24A').replace('_ver1',''),kind=kind,edits=edits,extra=len(EXTRA)))
    (TASK/'workbook-plan.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2),encoding='utf-8')
    (TASK/'traceability.json').write_text(json.dumps([dict(srs_row=r,id=id_by_row[r],fw=fw) for r,fw in ROW_FW.items()],ensure_ascii=False,indent=2),encoding='utf-8')

def ppt_build():
    p=SRC/'인공달팽이관 시스템 구성도.pptx';parts=readzip(p)
    for name in parts:
        if not re.match(r'ppt/slides/slide\d+\.xml$',name):continue
        x=E.fromstring(parts[name])
        for t in x.xpath('//a:t',namespaces=NS):
            old=t.text or '';new=old
            if name.endswith('slide4.xml'):
                new={'DSP':'E8300 / QCC5181','가속도센서':'터치센서','마이크로폰':'마이크 2개'}.get(old,old)
            if name.endswith('slide2.xml') and old=='외부장치':new='외부장치 Sound1'
            t.text=new;log(p.name,name,old,new,'Sound1 제품/부품 구성 반영')
        parts[name]=ser(x)
    writezip(STAGE/'인공달팽이관 시스템 구성도.pptx',parts)

def copy_reference():
    for rel in base:
        if not rel.startswith('시스템 Data/'):continue
        p=SRC/rel
        if p.suffix.lower() not in ('.docx','.xlsx','.pptx','.pdf','.doc','.xls','.ppt'):continue
        target=OUT/rel;target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,target)
    for name in ('Ezairo 8300 Hardware Reference.pdf','Ezairo 8300 Firmware Reference.pdf'):
        target=OUT/'시스템 Data/Ezairo8300'/name;target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(ROOT/'docs/참고/e8300'/name,target)

if __name__=='__main__':
    assert all(hashlib.sha256((SRC/p).read_bytes()).hexdigest()==v['sha256'] for p,v in base.items())
    word_build();workbook_plan();ppt_build();copy_reference()
    (TASK/'changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Prepared',len(changes),'changes')
