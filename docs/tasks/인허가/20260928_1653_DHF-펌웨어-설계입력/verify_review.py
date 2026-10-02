from pathlib import Path
import json,hashlib,shutil
import openpyxl,pypdfium2 as pdf
from copy import copy
T=Path(__file__).parent;Q=T/'review-20260929'
jobs=json.loads((Q/'plan.json').read_text(encoding='utf-8'))
results=[]
for j in jobs:
 b=openpyxl.load_workbook(j['backup']);p=Q/(j['kind']+'-final.xlsx');a=openpyxl.load_workbook(p)
 assert a.sheetnames==b.sheetnames+['변경 비교']
 for s in b:
  d=a[s.title]
  assert str(s.merged_cells)==str(d.merged_cells)
  assert len(s._images)==len(d._images)
  for row in s:
   for c in row:
    v=d.cell(c.row,c.column)
    assert c.value==v.value,(j['kind'],s.title,c.coordinate,c.value,v.value)
    assert c.number_format==v.number_format,(s.title,c.coordinate,'format')
    assert copy(c.font)==copy(v.font) and copy(c.alignment)==copy(v.alignment) and copy(c.border)==copy(v.border),(s.title,c.coordinate,'style')
  assert s.print_area==d.print_area
 for c in j['colors']:
  assert a[c['sheet']][c['cell']].fill.fgColor.rgb[-6:]==c['color']
 sh=a['변경 비교'];assert sh.auto_filter.ref==f"A7:G{len(j['records'])+7}"
 for i,r in enumerate(j['records'],8):
  assert sh[f'A{i}'].value==r['status']
  assert sh[f'B{i}'].hyperlink.location=="'"+r['sheet'].replace("'","''")+"'!"+r['cell']
  for col,key in [('D','old'),('E','new')]:
   expected='(빈칸)' if r[key] is None else str(r[key])
   assert sh[f'{col}{i}'].value==expected,(i,key,sh[f'{col}{i}'].value,expected)
  assert all(sh.cell(i,k).data_type!='e' for k in range(1,8))
 doc=pdf.PdfDocument(str(Q/(j['kind']+'-preview.pdf')))
 doc[0].render(scale=1.3).to_pil().save(Q/(j['kind']+'-native.png'))
 results.append({'file':j['name'],'body_values_formulas_styles_preserved':True,'colored_cells':len(j['colors']),'comparison_entries':len(j['records']),'hyperlinks_verified':True})
(Q/'verified.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(results,ensure_ascii=False))
