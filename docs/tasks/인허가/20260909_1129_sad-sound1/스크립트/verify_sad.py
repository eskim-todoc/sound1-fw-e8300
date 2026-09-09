# -*- coding: utf-8 -*-
"""생성된 SAD 자체 검증"""
import io, sys, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Cm

P = u'E:/workspace/projects/sound1-fw-e8300-backup/docs/참고/인허가/사운드1/[01] SAD-FW-SND1 소프트웨어 아키텍쳐 기술서.docx'
d = Document(P)
sec = d.sections[0]
BODY = sec.page_width.cm - sec.left_margin.cm - sec.right_margin.cm
print('본문 폭 %.2f cm / 표 %d / 단락 %d' % (BODY, len(d.tables), len(d.paragraphs)))

print('\n[검증_1] 표 폭 <= 본문 폭')
bad = 0
from docx.oxml.ns import qn as _qn
for i, t in enumerate(d.tables):
    grid = t._tbl.find(_qn('w:tblGrid'))
    ws = [int(gc.get(_qn('w:w'))) / 567.0 for gc in grid.findall(_qn('w:gridCol'))]
    tot = sum(ws)
    ok = tot <= BODY + 0.02
    if not ok:
        bad += 1
        print('  X 표%d 폭 %.2f cm (초과 %.2f)' % (i, tot, tot - BODY))
print('  결과: %s (초과 %d건)' % ('통과' if bad == 0 else '실패', bad))

print('\n[검증_2] 본문에 절 번호 상호참조 없음 (착탈 안전성)')
pat = re.compile(u'\\d+\\.\\d+\\s*절|\\d+\\.\\d+\\.\\d+\\s*절|절\\s*\\d+\\.\\d+')
hits = []
for p in d.paragraphs:
    t = p.text.strip()
    if not t or re.match(u'^\\d+(\\.\\d+)*\\.\\s', t):   # 제목 자신은 제외
        continue
    m = pat.search(t)
    if m:
        hits.append((m.group(0), t[:70]))
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            m = pat.search(c.text)
            if m:
                hits.append((m.group(0), c.text[:70]))
for h in hits:
    print('  X', h)
print('  결과: %s (%d건)' % ('통과' if not hits else '실패', len(hits)))

print('\n[검증_3] 폐기 부품 언급 없음')
banned = [u'nRF', u'nrf', u'NRF', u'가속도', u'이어피스', u'USB', u'7150']
found = []
alltext = '\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            alltext += '\n' + c.text
for b in banned:
    for ln in alltext.split('\n'):
        if b in ln:
            found.append((b, ln.strip()[:80]))
for f in found:
    print('  ?', f)
print('  결과: %s (%d건)' % ('통과' if not found else '검토 필요', len(found)))

print('\n[검증_4] 추적성 표 완결성')
tt = None
for t in d.tables:
    if t.rows[0].cells[0].text.strip().startswith('Software requirements'):
        tt = t
if tt is None:
    print('  X 추적성 표를 찾지 못함')
else:
    rows = tt.rows[2:]
    ids = [r.cells[0].text.strip() for r in rows]
    stats = [r.cells[3].text.strip() for r in rows]
    exp = [u'SW-FW-RID-%d' % i for i in range(1, 77)]
    print('  행 수 %d (기대 76)  / ID 연속성 %s' % (len(rows), 'OK' if ids == exp else 'X'))
    from collections import Counter
    cnt = Counter(stats)
    for k in [u'유지', u'변경', u'삭제', u'신규', u'확인 필요']:
        print('    %s : %d' % (k, cnt.get(k, 0)))
    miss = [ids[i] for i, s in enumerate(stats)
            if s in (u'변경', u'삭제', u'확인 필요') and not rows[i].cells[4].text.strip()]
    print('  사유 누락: %s' % ('없음' if not miss else miss))

print('\n[검증_5] 치환 토큰 존재')
for tok in ['SAD-FW-SND1', 'SND1-A-EXT', 'SDP-FW-SND1', 'SGR-FW-SND1', 'SR-FW-SND1', 'YYYY.MM.DD']:
    n = alltext.count(tok)
    hn = 0
    for hp in sec.header.paragraphs:
        hn += hp.text.count(tok)
    for htb in sec.header.tables:
        for row in htb.rows:
            for c in row.cells:
                hn += c.text.count(tok)
    print('  %-14s 본문 %d / 머리글 %d' % (tok, n, hn))

print('\n[검증_6] 그림 삽입')
imgs = [r for r in d.part.rels.values() if 'image' in r.reltype]
print('  이미지 %d개 (기대 2)' % len(imgs))

print('\n[검증_7] 부트로더 절 격리 - 3.4 를 가리키는 문장')
bl_refs = [p.text[:80] for p in d.paragraphs
           if (u'부트로더' in p.text or u'OTA' in p.text) and not p.text.strip().startswith(('3.4', '3.'))]
for b in bl_refs:
    print('  -', b)
