"""Read the current criterion and reconcile all 115 IDs, without editing Excel.

This is a coverage inventory, never an acceptance report. Worker inventories
describe limitations; executable checks under tmp/unit-tests describe actual
host execution. Device execution must be recorded separately.
"""
from __future__ import annotations
from collections import Counter
import hashlib
import json
from pathlib import Path
import re
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
TASK = ROOT / "docs/tasks/인허가/20260929_1755_DHF-펌웨어-설계출력"
CRITERION = ROOT / "docs/tasks/인허가/20260929_1755_DHF-펌웨어-설계출력/qa/evidence-collection/unit-criterion-before.xlsx"
EXPECTED_SHA = "89967fa692c4f51e3be916c73270d75859b4f692c96928c2d73da0cc6c5096fd"
PURE_CFX = {8,13,16,17,18,20,22,23,24,25,28,29,30,31,32,34,35}
ANALYSIS_CFX = {1,2,3,4,5,6,9,21,38,39,41,42}
ANALYSIS_CM3 = {1,3,5,8,15}


def read(relative):
    return json.loads((ROOT / relative).read_text(encoding="utf-8"))


def main():
    sha = hashlib.sha256(CRITERION.read_bytes()).hexdigest()
    assert sha == EXPECTED_SHA, "Criterion changed: review current IDs/inputs before regenerating coverage"
    model = {(v["core"], v["id"]): v for v in read(str((TASK / "qa/unit-reconciliation/model.json").relative_to(ROOT)))}
    io_file = "tests/unit/tdc_ut_cfx_io_coverage.json"
    pure_file = "tests/unit/cm3_pure_control_coverage.json"
    cm3_file = "tests/unit/tdc_ut_io_cm3_coverage.json"
    boot_file = "tests/unit/boot_coverage.json"
    io = {(v["core"],v["unit"]):v for v in read(io_file)["units"]}
    cm3_pure_data = read(pure_file)
    cm3_pure = {v["unit"]:v for v in cm3_pure_data["units"]}
    cm3_extra_data = cm3_pure_data["source_preserving_extension"]["extra_host"]
    cm3_extra = {v["unit"]:v for v in cm3_extra_data["units"]}
    cm3_io_data = read(cm3_file)
    cm3_io = {v["id"]:v for v in cm3_io_data["units"]}
    boot = {v["unit"]:v for v in read(boot_file)["coverage"]}
    analysis_file = "tests/unit/remaining_analysis_coverage.json"
    assert (ROOT / analysis_file).is_file(), "remaining analysis inventory missing"
    analysis={(v["core"],v["unit"]):v for v in read(analysis_file)["units"]}
    workbook = load_workbook(CRITERION, read_only=True, data_only=True)
    sheet = workbook["기준서_유닛 검증"]
    rows = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True),1):
        test_id = row[3] if len(row)>3 else None
        if not isinstance(test_id,str) or not test_id.startswith("SW-FW-UnitTest-"):
            continue
        match = re.fullmatch(r"SW-FW-UnitTest-(CFX|HEAR|CM3|BOOT)-ID-(\d+)",test_id)
        assert match, test_id
        core, number = match[1], int(match[2])
        assert row[1]==core and (core,number) in model
        entry = dict(core=core, unit=number, test_id=test_id, name=row[4], method=row[5],
                     unit_id=row[6], criterion_row=row_number,
                     device_executed=False, whole_unit_pass=False,
                     source_refs=model[(core,number)]["refs"])
        # Inventory ordering has no relationship to verdict or test priority.
        if core=="CFX" and number in PURE_CFX:
            entry.update(kind="fixed_vectors", target="ram_vectors", host="two_pass_vectors",
                         inventory="src/1__cfx/test/tdc_ut_vectors_cfx.c",
                         evidence="tmp/unit-tests/cfx-vectors/checks.json", mismatch_words=0,
                         remaining=("Actual SDK log/power paths and independent predeclared numeric bounds are prepared; host math adapters are not SDK/device execution"
                                    if number in (23,31,32) else "Prepared vector subset only; device execution and physical timing are separate"))
            if number in (23,31,32):
                entry["numeric_evidence"]="tmp/unit-tests/cfx-numeric-reference/checks.json"
        elif core=="CFX" and number in ANALYSIS_CFX:
            entry.update(kind="analysis", target="not_automatically_run", host="source_checks",
                         inventory="tests/unit/check_cfx_analysis.py", evidence="tmp/unit-tests/cfx-analysis/checks.json",
                         remaining="Source/order checks do not certify silicon event timing or every branch")
        elif (core,number) in io:
            data=io[(core,number)]
            entry.update(kind="target_required" if core=="HEAR" else "analysis" if number==14 else "partial_vectors" if data["scope"]=="partial_criterion" else "fixed_vectors",
                         target="hear_start_poll" if core=="HEAR" else "source_checks" if number==14 else "ram_vectors",
                         host=data["status"], mismatch_words=data.get("host_mismatch_count"),
                         inventory=analysis_file if number==14 and core=="CFX" else io_file,
                         evidence=analysis_file if number==14 and core=="CFX" else "tmp/unit-tests/cfx-io/checks.json",
                         remaining=data.get("missing") or data.get("limitation") or "Device execution required")
        elif core=="CM3" and number in cm3_pure:
            data=cm3_pure[number]
            observer=data["scope"]=="safe_target_observer_manual_gateway"
            entry.update(kind="unconfigured_observer" if observer else "partial_vectors" if data["scope"]=="partial" else "fixed_vectors",
                         target="manual_gateway_not_configured" if observer else "ram_vectors",
                         host="no_input_observer_twice" if observer else "two_pass_vectors",
                         mismatch_words=None if observer else data.get("host_mismatches",0), inventory=pure_file,
                         evidence="tmp/unit-tests/cm3-vectors/checks.json",
                         remaining=data.get("missing") or data.get("limitation") or "Device execution required")
        elif core=="CM3" and number in cm3_io:
            data=cm3_io[number]
            target_ram=number in cm3_io_data["target_ram_run_ids"]
            host_run=number in cm3_io_data["host_run_ids"]
            entry.update(kind="partial_vectors" if host_run else "analysis_observer" if entry["method"]=="A" else "observer_only",
                         target="ram_vectors" if target_ram else "manual_observer",
                         host="vectors" if host_run else "source_analysis_and_observer",
                         mismatch_words=data.get("mismatches",data.get("host_mismatches")), inventory=cm3_file,
                         evidence="tmp/unit-tests/cm3-io/checks.json",
                         remaining=data.get("missing") or "Device execution required")
        elif core=="CM3" and number in ANALYSIS_CM3:
            entry.update(kind="analysis", target="not_automatically_run", host="source_checks",
                         inventory=analysis_file, evidence=analysis_file,
                         remaining="Code path checked; no electrical or complete device execution proof")
        elif core=="CM3" and number==14:
            entry.update(kind="visual_required", target="manual_visual", host="not_applicable",
                         inventory="src/2__cm3/source/led/tdc_led_output.c",
                         evidence=None, remaining="Operator must observe the actual LED indication and retain image if required")
        elif core=="BOOT" and number in boot:
            data=boot[number]
            entry.update(kind="analysis" if data["host"]=="UNSUPPORTED" else "host_and_target_subset",
                         target=data["target"],host=data["host"], inventory=boot_file,
                         evidence="tmp/unit-tests/boot/checks.json",
                         mismatch_words=sum(v["unit"]==number for v in read(boot_file)["known_product_failures"]) if data["host"]!="UNSUPPORTED" else None,
                         remaining=data["manual_remaining"])
        else:
            raise AssertionError(f"No implementation/review/observation mapping: {test_id}")
        if core == "CM3" and number in cm3_extra:
            extra = cm3_extra[number]
            if number not in cm3_pure:
                # A PC-only extension must not turn an unisolated device path
                # into an automatically executable target RAM test.
                entry.update(kind="partial_vectors", host="two_process_isolated_passes",
                             inventory=pure_file, evidence=cm3_extra_data["evidence"],
                             mismatch_words=0, remaining=extra["remaining"])
            entry["additional_host_evidence"] = cm3_extra_data["evidence"]
            entry["additional_host_scope"] = extra
        if core == "CFX" and number == 33:
            entry["mismatch_classification"] = {
                "supported_small_band_inputs": io[(core,number)]["supported_input_mismatches"],
                "inconsistent_cm3_combination": io[(core,number)]["inconsistent_cm3_input_mismatches"],
                "criterion_32_band_inputs": io[(core,number)]["criterion_32_band_mismatches"],
            }
        if (core,number) in analysis:
            entry["criterion_issues"]=analysis[(core,number)]["criterion_issues"]
        rows.append(entry)
    workbook.close()
    assert len(rows)==115 and len({(v["core"],v["unit"]) for v in rows})==115
    assert Counter(v["core"] for v in rows)==dict(CFX=42,HEAR=2,CM3=51,BOOT=20)
    assert Counter(v["method"] for v in rows)==dict(I=1,A=31,D=34,T=49)
    for v in rows:
        assert (ROOT/v["inventory"]).is_file(), v
    pc_ids = {
        "CFX": PURE_CFX | {n for c,n in io if c == "CFX" and n != 14},
        "CM3": {n for n,v in cm3_pure.items() if v["scope"] != "safe_target_observer_manual_gateway"}
               | set(cm3_io_data["host_run_ids"]) | set(cm3_extra),
        "BOOT": {n for n,v in boot.items() if v["host"] != "UNSUPPORTED"},
    }
    result=dict(schema_version=1, criterion=str(CRITERION.relative_to(ROOT)),criterion_sha256=sha,
                device_execution=False, acceptance_report=False,
                pc_input_unit_counts={core:len(ids) for core,ids in pc_ids.items()},
                pc_input_units_total=sum(map(len,pc_ids.values())),
                counts=dict(Counter(v["kind"] for v in rows)), units=rows)
    out=ROOT/"tmp/unit-tests/coverage";out.mkdir(parents=True,exist_ok=True)
    (out/"115-units.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print("COVERAGE_MAPPED 115 unique current IDs; not 115 completed/passed tests")
    print(json.dumps(result["counts"],ensure_ascii=False))
    print("PC_INPUT_UNITS", result["pc_input_units_total"], result["pc_input_unit_counts"], "; partial inputs included, not acceptance")


if __name__=="__main__":
    main()
